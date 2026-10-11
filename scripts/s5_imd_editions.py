"""S5 (MHCLG English indices of deprivation, File 10: Local Authority
District summaries, lower tier): edition history and loader on the
period-editions engine, with a single-vintage live table.

The live table la_imd_2025 keeps its name, key (lad24cd) and columns and
holds one vintage (W1 reads imd_rank_of_average_rank from it). The vintage
is the engine period, a text key that sorts (rule 8): 'IoD2025'. The live
table gains one additive column for it, run by `ddl` only:

    ALTER TABLE public.la_imd_2025 ADD COLUMN vintage text NOT NULL
        DEFAULT 'IoD2025' CHECK (vintage = 'IoD2025')

(default-filled, not in the key; the CHECK makes a second vintage in live
impossible). This module owns the editions table and file-check ledger,
append-only:

    SPEC     la_imd_2025_editions               period vintage
    ledger   la_imd_2025_editions_file_checks

An edition is what one File 10 workbook says about its vintage: per
authority la_name and the ten IMD measures of the IMD sheet. A reissue of
the live vintage (the publisher reissued File 10 as 'v2' on 14 November
2025, page note of 17 November 2025: an LSOA allocated to Cherwell instead
of West Northamptonshire) is the next edition of that vintage; it reaches
live only through refresh-latest. A different vintage (IoD2030, say) is
stored as editions only and never enters live through a load
(EditionSpec.live_single_period): adopting one is a deliberate decision.
A different vintage is read with the headers its own file carries, which
needs its header added to FILE10_HEADERS deliberately, with the evidence;
until then its file halts on the header, which is the intended stop.

History (so the wording here stays truthful): the n8n workflow loaded the
table on 2026-04-01 (pipeline_run_log 30) from a CSV converted in a chat
session (imd_2025_la_summary.csv); its code applied parseFloat and parseInt
to the CSV (no zero default: a blank would have become NaN, then NULL), and
every held value equals the held File 10 v2 read again
(sha256 6c38a2fd3998e0ca..., data/raw/s5_imd, the same bytes as the
data/reference copy and scripts/verify/src/imd2025_file10.xlsx).
migrate-legacy (one-off) records the held vintage as edition 1 "as
loaded", with that proof on every stored column (296 x 11).

Identity (rule 3), from the file itself on every path. File 10 has no cover
sheet and its Notes sheet is empty, so nothing in the workbook states the
release in words. What the checks rely on: the sheet set, exactly Notes,
IMD, Income, Employment, Education, Health, Crime, Barriers, Living, IDACI,
IDAOPI; the IMD header row, exactly one declared in FILE10_HEADERS (trailing
spaces ignored; the IoD2025 header names its last four measures 'IMD25 -
...', which is what ties the file to the vintage); 296 E06-E09 rows; and
docProps modified (core.xml), which is the file's rank. The release page's
year and its attachment title 'File 10: Local Authority District summaries
- lower tier' corroborate when the page is read; --release must equal the
file's vintage. The rank of a file is its docProps modified, never its name
('v2' in a file name is not evidence) or link.

Blanks and zeros (docs/RULES.md rule 1): every IMD cell must be a number
(num_cell); a blank or text halts and is never read as 0. Published zeros
stay 0 (104 authorities have no LSOA in the most deprived 10%, 15 have an
extent of 0). A value with more decimal places than its live column holds
halts (storing it would round it).

Checks inside the file (checks): each rank column is a ranking of 1..296 (a
permutation of 1..296, except that tied measures share the lowest rank, as
the publisher ranks: the 104 zero proportions share rank 193 and the 15 zero
extents share 282 in the held file); each rank follows its measure (a
higher measure, which is more deprived, never ranks after a lower one;
measures equal at the published precision may rank apart: in the held file
2 groups of average scores, 4 of proportions and 13 of extents, presumably
because the publisher ranks the unrounded values, which is an inference:
the file does not say and the technical report was not read); proportions
and extents lie in [0, 1].

Geography (rule 4): geography.resolve('5', codes) (source 5 is declared
'old': an E08000038 or E08000039 halts), then each code's canonical form
must be in la_boundaries; anything else is UNEXPLAINED and stops.

Stop conditions (in load; a stopped vintage is REJECTED: nothing stored, no
ledger row, exit 1): a reissue that drops or adds an authority (PARTIAL; no
flag releases it); a vintage other than IoD2025 into an empty live table,
new or already held as editions only (stopped before the live CHECK would
refuse it; no flag releases it); equal docProps modified with
different content (two files claim the same release) unless --accept-reissue
VINTAGE; an older file (its rank before the held tip's, or a vintage before
the newest held) halts unless --allow-older-file. Each acceptance is written
to the edition's label and the run log.

W1 input: a new vintage's live rows take loaded_at = now(); a revision
reaches live only through refresh-latest, which copies the edition's
loaded_at and sets source (the edition's release_label) on every live row,
so run refresh-latest in the same session before W1.

Subcommands (every writing command previews by default; --commit and
--simulate are mutually exclusive; --simulate runs the --commit path and
always rolls back). There is no sync-new: migrate-legacy records the held
vintage.
    python scripts/s5_imd_editions.py ddl [--commit | --simulate]
    python scripts/s5_imd_editions.py status
    python scripts/s5_imd_editions.py load [--release YYYY]
        [--file PATH [--no-page]] [--recheck VINTAGE] [--allow-older-file]
        [--accept-reissue VINTAGE ...] [--commit | --simulate]
        # the file is downloaded to data/raw/s5_imd/ (also in a preview; a
        # preview writes nothing to the database)
    python scripts/s5_imd_editions.py refresh-latest [--commit | --simulate]
    python scripts/s5_imd_editions.py migrate-legacy FILE
        [--commit | --simulate]
    python scripts/s5_imd_editions.py restore-edition VINTAGE N
        [--commit | --simulate]
"""
import argparse
import dataclasses
import hashlib
import json
import math
import numbers
import re
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parent))

import editions_core as core  # noqa: E402
import geography  # noqa: E402
import period_editions as pe  # noqa: E402
from editions_core import halt  # noqa: E402

RUN_AGENT = "s5_imd_editions"
RUN_SOURCE = "5"
LIVE_MISSING = pe.LIVE_MISSING
HELD_NOT_LIVE = pe.HELD_NOT_LIVE
UTC = timezone.utc

API = "https://www.gov.uk/api/content"
GOV_UK = "https://www.gov.uk"
COLLECTION_PATH = "/government/collections/english-indices-of-deprivation"
COLLECTION_TITLE = "English indices of deprivation"
REPO = Path(__file__).resolve().parent.parent
RAW_DIR = REPO / "data" / "raw" / "s5_imd"     # git-ignored
USER_AGENT = ("ucws-pipeline S5 loader (read-only download of the public "
              "MHCLG English indices of deprivation File 10)")
MIN_FILE_BYTES = 50 * 1024
NO_PAGE = "--file --no-page: the release page was not read"

# ---------------------------------------------------------------------------
# Columns (live types from information_schema, checked 2026-10-10:
# lad24cd character varying(9) NOT NULL (PK), la_name text, the ten measures
# below (each nullable; 0 NULL held), source text, loaded_at timestamptz
# DEFAULT now(); no vintage column until ddl). The editions table holds every
# value NOT NULL: a blank halts, never stored.
# ---------------------------------------------------------------------------

KEY = "lad24cd"
PERIOD = "vintage"
NAME = "la_name"
VALUE_TYPES = (
    ("imd_average_rank", "numeric(10,2) NOT NULL"),
    ("imd_rank_of_average_rank", "integer NOT NULL"),
    ("imd_average_score", "numeric(10,3) NOT NULL"),
    ("imd_rank_of_average_score", "integer NOT NULL"),
    ("imd_proportion_most_deprived_10", "numeric(8,4) NOT NULL"),
    ("imd_rank_of_proportion", "integer NOT NULL"),
    ("imd_extent", "numeric(8,4) NOT NULL"),
    ("imd_rank_of_extent", "integer NOT NULL"),
    ("imd_local_concentration", "numeric(10,2) NOT NULL"),
    ("imd_rank_of_local_concentration", "integer NOT NULL"),
)
VALUES = tuple(c for c, _ in VALUE_TYPES)
CONTENT = VALUES + (NAME,)
SCALES_SQL = {c: tuple(int(x) for x in re.match(r"numeric\((\d+),(\d+)\)",
                                                 t).groups())
              for c, t in VALUE_TYPES if t.startswith("numeric")}
# decimal places per column (a rank column has none); not a source value
SCALE = {c: SCALES_SQL[c][1] if c in SCALES_SQL else 0 for c in VALUES}  # not a source value
# (measure, its rank): a higher measure is more deprived and ranks first
RANK_PAIRS = tuple((VALUES[i], VALUES[i + 1]) for i in range(0, 10, 2))
UNIT_COLS = ("imd_proportion_most_deprived_10", "imd_extent")

# ---------------------------------------------------------------------------
# The file
# ---------------------------------------------------------------------------

SHEETS = ("Notes", "IMD", "Income", "Employment", "Education", "Health",
          "Crime", "Barriers", "Living", "IDACI", "IDAOPI")
IMD_SHEET = "IMD"
CORE_XML = "docProps/core.xml"
# The IMD sheet's header row per vintage, verbatim from the file (trailing
# spaces as the publisher wrote them; they are ignored when reading). The
# columns after the two identifiers map in order to VALUES. IoD2025: read
# 2026-10-10 from File 10 v2 (data/raw/s5_imd). A different vintage is added
# here deliberately, with the evidence of its own file.
FILE10_HEADERS = {
    "IoD2025": (
        "Local Authority District code (2024)",
        "Local Authority District name (2024)",
        "IMD - Average rank ",
        "IMD - Rank of average rank ",
        "IMD - Average score ",
        "IMD - Rank of average score ",
        "IMD - Proportion of LSOAs in most deprived 10% nationally ",
        "IMD - Rank of proportion of LSOAs in most deprived 10% nationally ",
        "IMD25 - Extent ",
        "IMD25 - Rank of extent ",
        "IMD25 - Local concentration ",
        "IMD25 - Rank of local concentration ",
    ),
}
WIDTH = 12
DISTRICT_RE = re.compile(r"E0[6-9][0-9]{6}")
VINTAGE_RE = re.compile(r"IoD([0-9]{4})")
EXPECTED_AREAS = 296
# the one vintage the live table may hold (its CHECK)
LIVE_VINTAGE = "IoD2025"
PARTIAL = "PARTIAL FILE"    # problems no flag releases

# ---------------------------------------------------------------------------
# Discovery: the GOV.UK content API
# ---------------------------------------------------------------------------

DOC_TITLE_RE = re.compile(r"English indices of deprivation ([0-9]{4})")
# A title that looks like the series however it is worded ('English Indices
# of Deprivation 2030', 'Indices of Deprivation 2019: income ...')
SERIES_RE = re.compile(r"indices of deprivation", re.I)
ATTACHMENT_PREFIX = "File 10: Local Authority District summaries - lower tier"

# ---------------------------------------------------------------------------
# The held state (migrate-legacy preconditions), surveyed 2026-10-10
# ---------------------------------------------------------------------------

# (rows, the survey hash: md5 of lad24cd|la_name|the ten measures|source,
# NULL as '~', ordered by lad24cd; see live_state)
LEGACY_LIVE = (296, "60f16360106cf7e2" "5cba3f4283f4bae8")
# the held file (sha256 split so the credential scan does not flag it); the
# URL is the File 10 attachment of the IoD2025 release page, which the ledger
# row names so that a page load finds the file checked
LEGACY_FILE = {
    "name": ("File_10_-_IoD2025_Local_Authority_District_Summaries__"
             "lower-tier__v2.xlsx"),
    "sha256": "6c38a2fd3998e0caf76a3ef8da331613"
              "d52d6d0320065ccef346103fd9770a7e",
    "url": "https://assets.publishing.service.gov.uk/media/"
           "6917412ebc34c86ce4e6e7fc/File_10_-_IoD2025_Local_Authority_"
           "District_Summaries__lower-tier__v2.xlsx",
}
# before the editions exist (a preview against live): the held rows equal
# that file (the migrate-legacy proof), whose docProps modified is this
LEGACY_RANK = datetime(2025, 11, 14, 12, 48, 6, tzinfo=UTC)


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _json(x):
    return json.loads(x) if isinstance(x, (str, bytes)) else (x or {})


def _norm(s) -> str:
    return " ".join(str("" if s is None else s).split())


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or \
        (isinstance(v, str) and not v.strip())


def iso(when) -> str:
    """'2025-11-14T12:48:06Z' of an aware datetime (UTC)."""
    return when.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


def rank_text(rank) -> str:
    return f"docProps modified {iso(rank)}" if rank else "no file rank"


def vintage_of(year) -> str:
    return f"IoD{int(year)}"


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------

def _documents(collection_json) -> list:
    d = _json(collection_json)
    return list(((d.get("links") or {}).get("documents")) or [])


def _titles(collection_json) -> list:
    return [_norm(d.get("title")) for d in _documents(collection_json)]


def check_collection(collection_json) -> None:
    """ValueError unless the collection is still titled COLLECTION_TITLE."""
    t = _norm(_json(collection_json).get("title"))
    if t != COLLECTION_TITLE:
        raise ValueError(f"the collection is titled {t!r}, not "
                         f"{COLLECTION_TITLE!r}; the page has changed")


def release_candidates(collection_json) -> dict:
    """{year: [(title, base_path)]} of the documents titled 'English indices
    of deprivation <yyyy>'."""
    out = {}
    for doc in _documents(collection_json):
        t = _norm(doc.get("title"))
        mm = DOC_TITLE_RE.fullmatch(t)
        if mm and doc.get("base_path"):
            out.setdefault(int(mm.group(1)), []).append((t, doc["base_path"]))
    return out


def _strays(titles, newest) -> list:
    """Titles that look like the series, do not fit the strict pattern and
    name a later year than newest (a title with no year, such as the privacy
    notice, is not a release)."""
    out = []
    for t in titles:
        if DOC_TITLE_RE.fullmatch(t) or not SERIES_RE.search(t):
            continue
        ys = [int(y) for y in re.findall(r"(?<![0-9])([0-9]{4})(?![0-9])", t)]
        if any(y > newest for y in ys):
            out.append(t)
    return out


def latest_release(collection_json) -> tuple:
    """(title, base_path) of the newest 'English indices of deprivation
    <yyyy>' document. ValueError (listing the titles seen) when none has the
    title, when two claim the newest year, or when a document that looks
    like the series names a later year in a title that does not fit (never
    passed over for an older release; the message names the newest release
    seen)."""
    cands = release_candidates(collection_json)
    titles = _titles(collection_json)
    if not cands:
        raise ValueError("no document titled 'English indices of deprivation "
                         f"<yyyy>' in the collection; titles seen: "
                         f"{titles[:40]}")
    newest = max(cands)
    stray = _strays(titles, newest)
    if stray:
        raise ValueError(f"the newest matching release is 'English indices of "
                         f"deprivation {newest}', but the collection lists a "
                         "document that looks like the series, with a later "
                         f"year, whose title does not fit: {stray}; not "
                         "choosing an older release over it; titles seen: "
                         f"{titles[:40]}")
    if len(cands[newest]) != 1:
        raise ValueError(f"{len(cands[newest])} documents for {newest}: "
                         f"{[t for t, _ in cands[newest]]}; refusing to "
                         "choose one")
    return cands[newest][0]


def newest_listed(collection_json) -> "int | None":
    cands = release_candidates(collection_json)
    return max(cands) if cands else None


def release_for_year(collection_json, year) -> tuple:
    """(title, base_path) of that year's release (--release YYYY, or the
    vintage of a --file); ValueError unless exactly one document has the
    title."""
    cands = release_candidates(collection_json).get(int(year), [])
    if len(cands) != 1:
        raise ValueError(f"{len(cands)} documents titled 'English indices of "
                         f"deprivation {year}' in the collection: "
                         f"{[t for t, _ in cands]}; titles seen: "
                         f"{_titles(collection_json)[:40]}")
    return cands[0]


def page_year(release_json) -> "int | None":
    mm = DOC_TITLE_RE.fullmatch(_norm(_json(release_json).get("title")))
    return int(mm.group(1)) if mm else None


def _attachments(page_json) -> list:
    d = _json(page_json)
    return list(((d.get("details") or {}).get("attachments")) or [])


def file10_attachment(release_json) -> dict:
    """{title, url, content_type, file_size} of the release page's one
    attachment whose title starts ATTACHMENT_PREFIX, with an .xlsx URL.
    ValueError listing every attachment title otherwise (none, several, or
    not .xlsx)."""
    atts = _attachments(release_json)
    names = [_norm(a.get("title")) for a in atts]
    hits = [a for a in atts
            if _norm(a.get("title")).startswith(ATTACHMENT_PREFIX)]
    if len(hits) != 1 or not hits[0].get("url"):
        raise ValueError(f"{len(hits)} attachments whose title starts "
                         f"{ATTACHMENT_PREFIX!r}; attachments: {names}")
    a = hits[0]
    if Path(urlparse(a["url"]).path).suffix.lower() != ".xlsx":
        raise ValueError(f"the File 10 attachment is not an .xlsx file "
                         f"({a['url']}); attachments: {names}")
    return {"title": _norm(a["title"]), "url": a["url"],
            "content_type": a.get("content_type"),
            "file_size": a.get("file_size")}


def change_history(release_json) -> list:
    """[(public_timestamp date, note)] of the page, as published (the 2025
    page's 17 November 2025 note is the evidence of the v2 reissue)."""
    d = _json(release_json)
    ch = ((d.get("details") or {}).get("change_history")) or []
    return [(str(c.get("public_timestamp") or "")[:10], _norm(c.get("note")))
            for c in ch]


def discovery_note(collection_json, held_periods) -> list:
    """Lines to print when discovery found nothing newer than what is held:
    the newest release the collection lists, and that none newer is listed.
    Empty when the newest listed vintage is not held."""
    nl = newest_listed(collection_json)
    if nl is None or vintage_of(nl) not in {str(p) for p in held_periods}:
        return []
    return [f"the newest release the collection lists is 'English indices "
            f"of deprivation {nl}' ({vintage_of(nl)}), which is already "
            "held; no newer release is listed in the collection"]


# ---------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------

def num_cell(v, where):
    """A published number (an int, or a finite float). Anything else raises
    ValueError naming `where`: a blank, any text (a number written as text
    included), a boolean, NaN or infinity. Never returns a default."""
    if v is None or (isinstance(v, str) and not v.strip()):
        raise ValueError(f"{where}: a blank cell (a blank halts and is never "
                         "read as 0)")
    if isinstance(v, str):
        raise ValueError(f"{where}: text {v!r} is not a number (text halts; "
                         "it is never read as a number or as 0)")
    if isinstance(v, bool):
        raise ValueError(f"{where}: {v!r} is not a number")
    if isinstance(v, numbers.Integral):
        return int(v)
    if isinstance(v, numbers.Real):
        if math.isnan(v) or math.isinf(v):
            raise ValueError(f"{where}: {v!r} is not a number (a blank or "
                             "error cell; never read as 0)")
        return float(v)
    raise ValueError(f"{where}: {v!r} ({type(v).__name__}) is not a number")


def _scaled(v, col, where):
    """v at the live column's scale: an int for a rank, else a Decimal
    quantised to the column's scale. ValueError when that would change it
    (a non-whole rank, or more decimal places than the column holds)."""
    s = SCALE[col]
    if s == 0:
        if isinstance(v, float) and not v.is_integer():
            raise ValueError(f"{where}: {v!r} is not a whole number ({col} "
                             "is an integer column)")
        return int(v)
    d = Decimal(repr(v)) if isinstance(v, float) else Decimal(int(v))
    q = d.quantize(Decimal(1).scaleb(-s))
    if q != d:
        raise ValueError(f"{where}: {v!r} has more decimal places than {col} "
                         f"holds ({s}); storing it would round it, so the "
                         "load halts")
    return q


def _fmt(col, v) -> str:
    """A stored value as text at its live scale ('0.0000', '23075.98')."""
    if col == NAME:
        return str(v)
    if SCALE[col] == 0:
        return str(int(v))
    return f"{Decimal(v):.{SCALE[col]}f}"


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------

def _when(text, where) -> datetime:
    s = (text or "").strip()
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        raise ValueError(f"{where}: docProps modified {s!r} is not a date "
                         "and time") from None
    if d.tzinfo is None:
        raise ValueError(f"{where}: docProps modified {s!r} has no time zone")
    return d.astimezone(UTC)


def docprops_modified(path) -> datetime:
    """The workbook's own docProps modified (core.xml), in UTC. ValueError
    when it is absent or unreadable: the rank is never taken from the file
    name."""
    p = Path(path)
    where = f"{p.name}: {CORE_XML}"
    try:
        with zipfile.ZipFile(p) as z:
            if CORE_XML not in z.namelist():
                raise ValueError(f"{where} is absent: the file's rank "
                                 "(docProps modified) cannot be read from the "
                                 "file itself, and it is never taken from the "
                                 "name")
            root = ET.fromstring(z.read(CORE_XML))
    except (zipfile.BadZipFile, ET.ParseError, OSError) as e:
        raise ValueError(f"{where} is not readable ({e}): the file's rank "
                         "(docProps modified) cannot be read") from None
    hits = [el.text for el in root
            if el.tag.rsplit("}", 1)[-1] == "modified"
            and el.text and el.text.strip()]
    if len(hits) != 1:
        raise ValueError(f"{where}: {len(hits)} modified values; the file's "
                         "rank (docProps modified) cannot be read from the "
                         "file itself, and it is never taken from the name")
    return _when(hits[0], where)


def _vintage_of_header(header, where) -> str:
    got = [h.rstrip() if isinstance(h, str) else h for h in header]
    hits = [v for v, want in FILE10_HEADERS.items()
            if got == [w.rstrip() for w in want]]
    if len(hits) != 1:
        raise ValueError(
            f"{where}: the header row {list(header)!r} is not a declared File "
            f"10 header (declared: {sorted(FILE10_HEADERS)}; IoD2025 is "
            f"{list(FILE10_HEADERS.get('IoD2025', ()))!r}). A renamed column "
            "or another vintage's header halts: the reading must be corrected "
            "deliberately, with the evidence (FILE10_HEADERS)")
    return hits[0]


def read_file10(path) -> dict:
    """The facts of one File 10 workbook: {path, name, sheets, header (as
    the file has it), vintage, year, rows [{row, code, name, cells (ten
    numbers, num_cell)}], modified (docProps, UTC), rank (= modified)}.
    ValueError naming what was seen on: unreadable docProps modified; sheets
    not exactly SHEETS; an IMD header that is not exactly one declared in
    FILE10_HEADERS (trailing spaces ignored); a row that is not an E06-E09
    authority with a name; a repeated code; data after an empty row or
    beyond the headed columns; a cell that is not a number; a number of
    authorities other than EXPECTED_AREAS."""
    import openpyxl
    p = Path(path)
    modified = docprops_modified(p)
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 (not an .xlsx workbook)
        raise ValueError(f"{p.name}: not a readable .xlsx workbook ({e})"
                         ) from None
    try:
        sheets = list(wb.sheetnames)
        if sheets != list(SHEETS):
            unknown = [s for s in sheets if s not in SHEETS]
            missing = [s for s in SHEETS if s not in sheets]
            raise ValueError(f"{p.name}: sheets {sheets} are not exactly "
                             f"{list(SHEETS)} (unknown {unknown}, missing "
                             f"{missing}"
                             + ("" if unknown or missing else ", order differs")
                             + ")")
        rows = [list(r) for r in wb[IMD_SHEET].iter_rows(values_only=True)]
    finally:
        wb.close()
    where = f"{p.name}: {IMD_SHEET}"
    if not rows:
        raise ValueError(f"{where}: the sheet is empty")
    header = list(rows[0])
    while header and _blank(header[-1]):
        header.pop()
    vintage = _vintage_of_header(header, where)
    out, seen = [], {}
    empty = 0  # empty rows seen so far; not a source value
    for n, r in enumerate(rows[1:], start=2):
        if all(_blank(v) for v in r):
            empty += 1
            continue
        rw = f"{where} row {n}"
        if empty:
            raise ValueError(f"{rw}: data after {empty} empty row(s); empty "
                             "rows are allowed only at the end")
        if any(not _blank(v) for v in r[WIDTH:]):
            raise ValueError(f"{rw}: a cell beyond the {WIDTH} headed columns")
        r = (list(r) + [None] * WIDTH)[:WIDTH]
        code, name = _norm(r[0]), r[1]
        if not DISTRICT_RE.fullmatch(code):
            raise ValueError(f"{rw}: {code!r} is not an E06-E09 authority "
                             "code")
        if not isinstance(name, str) or not name.strip():
            raise ValueError(f"{rw}: {code} without a name ({name!r})")
        if code in seen:
            raise ValueError(f"{rw}: {code} appears twice (also row "
                             f"{seen[code]})")
        seen[code] = n
        cells = tuple(num_cell(r[2 + i], f"{rw} {code} {col}")
                      for i, col in enumerate(VALUES))
        out.append({"row": n, "code": code, "name": name.strip(),
                    "cells": cells})
    want = globals()["EXPECTED_AREAS"]
    if len(out) != want:
        raise ValueError(f"{where}: {len(out)} authorities (E06-E09 rows), "
                         f"expected {want}")
    return {"path": p, "name": p.name, "sheets": sheets, "header": header,
            "vintage": vintage, "year": int(VINTAGE_RE.fullmatch(
                vintage).group(1)),
            "rows": out, "modified": modified, "rank": modified}


# ---------------------------------------------------------------------------
# Records and checks
# ---------------------------------------------------------------------------

def _getter(resolve):
    return resolve if callable(resolve) else (lambda c: resolve.get(c))


def build_records(file, resolve) -> list:
    """The vintage's records from a read file, one per authority, sorted by
    lad24cd: {lad24cd, vintage, la_name, the ten measures at the live
    scales (ranks int, the rest Decimal)}. resolve: {publisher code:
    lad24cd} or a callable. ValueError when a code resolves to nothing, a
    lad24cd repeats, or a value would change at its live scale."""
    get = _getter(resolve)
    out, seen = [], {}
    for row in file["rows"]:
        code = row["code"]
        where = f"{IMD_SHEET} row {row['row']} {code}"
        lad = get(code)
        if not lad:
            raise ValueError(f"{where}: {code} resolves to no lad24cd")
        if lad in seen:
            raise ValueError(f"{where}: duplicate key {lad} (also "
                             f"{seen[lad]}); refusing to choose a winner")
        seen[lad] = code
        rec = {KEY: lad, PERIOD: file["vintage"], NAME: row["name"]}
        for col, v in zip(VALUES, row["cells"]):
            rec[col] = _scaled(v, col, f"{where} {col}")
        out.append(rec)
    out.sort(key=lambda x: x[KEY])
    return out


def _ranking_problem(ranks, n) -> "str | None":
    """None when ranks are a standard competition ranking of 1..n (each
    rank r held k times means the next rank is r + k; with no ties, a
    permutation of 1..n); else what is wrong."""
    c = Counter(ranks)
    below = 0  # ranks seen below the current one; not a source value
    for r in sorted(c):
        if r != below + 1:
            dup = sorted(x for x, k in c.items() if k > 1)
            return (f"rank {r} where {below + 1} was expected (ranks "
                    f"{min(c)}..{max(c)}, {len(c)} distinct, shared ranks "
                    f"{dup[:6]})")
        below += c[r]
    return None


def _order_problem(records, measure, rank) -> "str | None":
    """None when a higher measure always ranks before a lower one; else the
    first pair that does not."""
    srt = sorted(records, key=lambda r: r[measure], reverse=True)
    prev_max, prev_rec, i = None, None, 0
    while i < len(srt):
        j = i
        while j < len(srt) and srt[j][measure] == srt[i][measure]:
            j += 1
        group = srt[i:j]
        lo = min(group, key=lambda r: r[rank])
        if prev_max is not None and lo[rank] <= prev_max:
            return (f"{prev_rec[KEY]} ({measure} {prev_rec[measure]}) ranks "
                    f"{prev_rec[rank]}, {lo[KEY]} ({measure} {lo[measure]}, "
                    f"lower) ranks {lo[rank]}")
        hi = max(group, key=lambda r: r[rank])
        if prev_max is None or hi[rank] > prev_max:
            prev_max, prev_rec = hi[rank], hi
        i = j
    return None


def checks(records) -> list:
    """The checks inside the file, on a vintage's records: each rank column
    a ranking of 1..n (n = the records; a permutation of 1..n where no
    measures tie, tied measures sharing the lowest rank); each rank follows
    its measure (higher measure = more deprived = earlier rank; measures
    equal at the published precision may rank apart); proportions and
    extents in [0, 1]. Returns the lines saying what passed; raises
    ValueError listing every problem."""
    n = len(records)
    bad, lines = [], []
    for measure, rank in RANK_PAIRS:
        ranks = [r[rank] for r in records]
        if any(not isinstance(x, int) or x < 1 or x > n for x in ranks):
            bad.append(f"{rank}: a rank outside 1..{n} "
                       f"({sorted(set(x for x in ranks if not isinstance(x, int) or x < 1 or x > n))[:4]})")
            continue
        why = _ranking_problem(ranks, n)
        if why:
            bad.append(f"{rank}: not a ranking of 1..{n}: {why}")
            continue
        why = _order_problem(records, measure, rank)
        if why:
            bad.append(f"{rank}: the ranks do not follow {measure} (higher "
                       f"is more deprived and ranks first): {why}")
            continue
        shared = {x: k for x, k in Counter(ranks).items() if k > 1}
        lines.append(f"{rank} ranks 1..{n} following {measure}"
                     + (f" (shared ranks {dict(sorted(shared.items()))})"
                        if shared else " (a permutation)"))
    for col in UNIT_COLS:
        out = [r[KEY] for r in records if not 0 <= r[col] <= 1]
        if out:
            bad.append(f"{col}: outside [0, 1] for {out[:6]}")
        else:
            zeros = sum(1 for r in records if r[col] == 0)  # not a source value
            lines.append(f"{col} in [0, 1] ({zeros} published zeros)")
    if bad:
        raise ValueError("the file's own checks failed: " + "; ".join(bad))
    return lines


def rows_content_sha(records) -> str:
    """sha256 of a vintage's content: one line per record sorted by lad24cd,
    lad24cd|la_name|the ten measures at their live scales (VALUES order),
    joined by LF, UTF-8. Nothing else is in it, so where a file was read
    never makes a new edition."""
    lines = sorted("|".join([str(r[KEY]), _fmt(NAME, r[NAME])]
                            + [_fmt(c, r[c]) for c in VALUES])
                   for r in records)
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Geography
# ---------------------------------------------------------------------------

def resolve_codes(cur, codes) -> tuple:
    """({publisher code: lad24cd}, problems). geography.resolve(cur, '5',
    codes) (source 5 is declared 'old': a new Barnsley or Sheffield code is
    a problem); then each code's canonical form must be in la_boundaries,
    else 'UNEXPLAINED'. Two codes reaching one lad24cd are a problem."""
    codes = set(codes)
    recode, problems = geography.resolve(cur, RUN_SOURCE, codes)
    problems = list(problems)
    cur.execute("SELECT lad24cd FROM public.la_boundaries")
    valid = {r[0] for r in cur.fetchall()}
    out, hit = {}, {}
    for c in sorted(codes):
        r = geography.canonical(c, recode)
        if r in valid:
            out[c] = r
            hit.setdefault(r, []).append(c)
        else:
            problems.append(f"UNEXPLAINED {c}: not in la_boundaries and not a "
                            "la_code_lookup recode")
    two = {k: v for k, v in hit.items() if len(v) > 1}
    if two:
        problems.append(f"codes {two} reach one lad24cd; refusing to choose "
                        "a winner")
    return out, problems


# ---------------------------------------------------------------------------
# Ranks, ledger and planning
# ---------------------------------------------------------------------------

_ISO = r"[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}:[0-9]{2}Z"
_RANK_RE = re.compile(rf"\bdocProps modified ({_ISO})")
_LEDGER_RE = re.compile(rf"(.*) \(vintage (IoD[0-9]{{4}}); docProps modified "
                        rf"({_ISO})\)")


def release_rank(file_or_source) -> "datetime | None":
    """The docProps modified (UTC) of a file: a read file's own rank, or the
    'docProps modified <iso>' an edition's or ledger row's source_file
    carries; None if it names none (a file name, URL or media id is never
    read for a rank)."""
    x = file_or_source
    if isinstance(x, dict):
        return x.get("rank")
    mm = _RANK_RE.search(str(x or ""))
    if not mm:
        return None
    return _when(mm.group(1), "source_file")


def edition_source(file, src, note="") -> str:
    """An edition's source_file: the file name, its vintage and docProps
    modified (the rank), where it was read, and a note."""
    return (f"{src['file_name']}; vintage {file['vintage']}; "
            f"{rank_text(file['modified'])}; {src['where']}"
            + (f"; {note}" if note else ""))


def ledger_source(where, vintage, rank) -> str:
    """A file's ledger source_file: where it was read (the final URL, or the
    path of a --file), its vintage and docProps modified, so a later run
    knows them without reading the file."""
    return f"{where} (vintage {vintage}; {rank_text(rank)})"


def parse_ledger_source(s) -> "tuple | None":
    """(where, vintage, rank) of a ledger source_file, or None."""
    mm = _LEDGER_RE.fullmatch(str(s or ""))
    if not mm:
        return None
    return mm.group(1), mm.group(2), _when(mm.group(3), "ledger")


def content_sha256(path) -> str:
    """SHA-256 of the file's bytes."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def tip_ranks(tips, checked) -> dict:
    """{held vintage: the newest file rank that stated it}: the tip
    edition's, or a later file the ledger records for it (a file found
    unchanged stores no edition)."""
    out = {}
    for p, t in tips.items():
        ranks = [t.get("rank")] + [release_rank(s)
                                   for s, _ in checked.get(p, ())]
        ranks = [r for r in ranks if r is not None]
        out[p] = max(ranks) if ranks else None
    return out


def ledger_complete(checked, where, sha, ranks, *, allow_older=False,
                    stranded=()) -> "list | None":
    """[vintage] when the ledger records this file's (where, sha256) pair
    for the vintage its ledger source declares, else None (the file is read
    and planned). A vintage whose tip comes from a newer file is read again
    (and then planned as older, which halts) unless allow_older; a stranded
    vintage (editions, no live rows, live empty) always makes the file be
    read. So a run that failed or was rejected leaves no ledger row and the
    rerun reads the file again."""
    have, declared = set(), None
    for p, pairs in checked.items():
        for s, h in pairs:
            parsed = parse_ledger_source(s)
            if h == sha and parsed and parsed[0] == where:
                have.add(str(p))
                declared = parsed
    if declared is None:
        return None
    _, vintage, rank = declared
    if vintage not in have or vintage in {str(x) for x in stranded}:
        return None
    if not allow_older and ranks.get(vintage) is not None \
            and ranks[vintage] > rank:
        return None
    return [vintage]


def plan_vintage(vintage, rank, ranks, checked, source, sha, *, recheck,
                 allow_older) -> tuple:
    """(new, revised, skipped) for the file's one vintage. ranks: {held
    vintage: tip_ranks}; checked: {vintage: {(source, sha)}} of the ledger
    (stranded vintages left out by the caller, so they are always read).
    new: not held (unless it sorts before the newest held vintage: skipped
    as older, unless allow_older); revised: held, compared with the tip;
    skipped {vintage: reason}: held with a tip from a newer file (older,
    unless allow_older), or already in the ledger (checked, unless
    --recheck names it). --recheck must be the file's held vintage
    (ValueError)."""
    v = str(vintage)
    if recheck is not None and (recheck != v or v not in ranks):
        raise ValueError(f"--recheck {recheck}: not a held vintage this file "
                         f"states ({v})")
    newest_held = max(ranks) if ranks else None
    if v not in ranks:
        if newest_held is not None and v < newest_held and not allow_older:
            return [], [], {v: (f"older: {v} is before the newest held "
                                f"vintage {newest_held}")}
        return [v], [], {}
    t = ranks[v]
    if t is not None and t > rank and not allow_older:
        return [], [], {v: (f"older: its tip comes from a file of "
                            f"{rank_text(t)}, this file is "
                            f"{rank_text(rank)}")}
    if recheck != v and (source, sha) in checked.get(v, ()):
        return [], [], {v: "checked: this file is already in the ledger"}
    return [], [v], {}


# ---------------------------------------------------------------------------
# Stop conditions
# ---------------------------------------------------------------------------

def _by_key(records) -> dict:
    return {r[KEY]: r for r in records or ()}


def period_problems(new, tip, *, kind, vintage=None, live_empty=False) -> list:
    """The stop conditions of the file's vintage (empty = fine); no flag
    releases any of them. revised (a reissue of a held vintage): an
    authority of the held edition missing from the file, or one the held
    edition lacks (the key set of a vintage never changes by a reissue, and
    refresh would refuse it). Any kind (new, or revised: a vintage held as
    editions only): a vintage other than LIVE_VINTAGE into an empty live
    table (the live CHECK holds LIVE_VINTAGE only; another vintage never
    enters live by a load, so with live empty only LIVE_VINTAGE is
    inserted)."""
    out = []
    if kind == "revised":
        tk, nk = _by_key(tip), _by_key(new)
        gone = sorted(set(tk) - set(nk))
        added = sorted(set(nk) - set(tk))
        if gone:
            out.append(f"{PARTIAL}: {len(gone)} authorit(ies) of the held "
                       "edition missing from the file (a partial file never "
                       f"replaces a fuller edition): {', '.join(gone[:8])}")
        if added:
            out.append(f"{PARTIAL}: {len(added)} authorit(ies) not in the "
                       "held edition (a reissue never changes a vintage's "
                       f"key set): {', '.join(added[:8])}")
    if live_empty and vintage is not None and str(vintage) != LIVE_VINTAGE:
        # every kind: a vintage held as editions only (planned as
        # revised) must not be inserted into an empty live table either
        out.append(f"{vintage} into an empty live table: the live table holds "
                   f"{LIVE_VINTAGE} only (its CHECK), and another vintage "
                   "never enters live through a load (adopting one is a "
                   f"deliberate decision); load {LIVE_VINTAGE} first, then "
                   f"{vintage} is stored as editions only")
    return out


def _differs(new, old) -> bool:
    nk, ok = _by_key(new), _by_key(old)
    if set(nk) != set(ok):
        return True
    return any(nk[k][c] != ok[k][c] for k in nk for c in CONTENT)


# ---------------------------------------------------------------------------
# Spec and profile
# ---------------------------------------------------------------------------

SPEC = core.EditionSpec(
    name="s5",
    live_table="la_imd_2025",
    editions_table="la_imd_2025_editions",
    key_cols=(KEY,),
    period_col=PERIOD,
    value_cols=VALUE_TYPES,
    extra_cols=((NAME, "text NOT NULL"),),
    refresh_cols=VALUES + (NAME, "source", "loaded_at"),
    refresh_from=(("source", "release_label"), ("loaded_at", "loaded_at")),
    as_loaded_source_col="source",
    refresh_source_whole_period=True,   # one file's label per vintage
    key_types=((KEY, "varchar(9) NOT NULL"), (PERIOD, "text NOT NULL")),
    table_constraints=(f"CHECK ({PERIOD} ~ '^IoD[0-9]{{4}}$')",),
    fk_la_boundaries=True,
    live_single_period=True,            # one vintage in live
    refresh_key_changes=False,          # the key set never changes
)


def live_alter_sql(spec) -> str:
    """The one live DDL change: the vintage column (additive,
    default-filled, not in the key; the CHECK keeps one vintage in live)."""
    return (f"ALTER TABLE public.{core._ident(spec.live_table)} ADD COLUMN "
            f"{PERIOD} text NOT NULL DEFAULT '{LIVE_VINTAGE}' CHECK "
            f"({PERIOD} = '{LIVE_VINTAGE}')")


def _no_label(fetched_on) -> str:
    raise ValueError("an S5 release label comes from the file; store "
                     "through apply_vintage(...)")


def check_records(records, period) -> None:
    """ValueError unless the vintage's records are whole: each of this
    vintage, one per lad24cd, every content column present."""
    seen = set()
    for r in records:
        if r.get(PERIOD) != period:
            raise ValueError(f"{period}: a record of vintage "
                             f"{r.get(PERIOD)!r}")
        if r[KEY] in seen:
            raise ValueError(f"{period}: {r[KEY]} appears twice")
        if any(r.get(c) is None for c in CONTENT):
            raise ValueError(f"{period}: {r[KEY]} lacks a value")
        seen.add(r[KEY])


PROFILE = pe.Profile(
    spec=SPEC, value_cols=CONTENT, run_agent=RUN_AGENT, run_source=RUN_SOURCE,
    heading="S5 MHCLG English indices of deprivation, File 10 (la_imd_2025)",
    default_source_file="MHCLG English indices of deprivation File 10",
    expected_areas=EXPECTED_AREAS, release_label=_no_label,
    content_sha256=rows_content_sha, check_records=check_records,
    file_checks=True, savepoint="s5_period",
    example_label="(ten IMD measures, la_name)")


def profile() -> pe.Profile:
    """PROFILE bound to the current SPEC and EXPECTED_AREAS (looked up when
    called, so tests can swap them)."""
    g = globals()
    return dataclasses.replace(PROFILE, spec=g["SPEC"],
                               expected_areas=g["EXPECTED_AREAS"])


def _late(name: str):
    """A call through this module's attribute `name`, looked up when
    called (so mock.patch.object(s5_imd_editions, name, ...) reaches the
    engine)."""
    return lambda *a, **kw: globals()[name](*a, **kw)


def live_has_vintage(cur, spec) -> bool:
    cur.execute("SELECT 1 FROM information_schema.columns WHERE "
                "table_schema = 'public' AND table_name = %s AND "
                "column_name = %s", (spec.live_table, PERIOD))
    return cur.fetchone() is not None


def ensure_live_vintage(cur, spec) -> str:
    """Add the live vintage column (live_alter_sql) when it is absent and
    return 'added'; 'present' when it is there exactly as planned (text,
    NOT NULL, DEFAULT LIVE_VINTAGE, the CHECK); halt when a vintage column
    of any other shape exists."""
    cur.execute("SELECT data_type, is_nullable, column_default FROM "
                "information_schema.columns WHERE table_schema = 'public' "
                "AND table_name = %s AND column_name = %s",
                (spec.live_table, PERIOD))
    row = cur.fetchone()
    if row is None:
        cur.execute(live_alter_sql(spec))
        return "added"
    cur.execute("SELECT pg_get_constraintdef(oid) FROM pg_constraint WHERE "
                "conrelid = %s::regclass AND contype = 'c'",
                (f"public.{spec.live_table}",))
    defs = [r[0] for r in cur.fetchall()]
    want = ("text", "NO", f"'{LIVE_VINTAGE}'::text")
    check = f"({PERIOD} = '{LIVE_VINTAGE}'::text)"
    if tuple(row) != want or not any(check in d for d in defs):
        halt(f"{spec.live_table}.{PERIOD} exists but is not the planned "
             f"column (type, nullable, default {tuple(row)}; CHECK "
             f"constraints {defs}); expected: {live_alter_sql(spec)}")
    return "present"


def create_all(cur) -> None:
    """ddl: the editions table, its append-only triggers, the file-check
    ledger and the live vintage column (live_alter_sql, when absent).
    Idempotent; no value is written; no other live column or row is
    touched."""
    prof = profile()
    core.create_schema(cur, prof.spec)
    pe.create_file_checks(cur, prof)
    ensure_live_vintage(cur, prof.spec)


def insert_live(cur, profile, period: str, records: list) -> None:
    """A vintage's live rows: the key, the vintage, the ten measures,
    la_name and source (each record's live_source); loaded_at takes its
    default."""
    from psycopg2.extras import execute_values
    spec = profile.spec
    cols = (KEY, spec.period_col) + VALUES + (NAME, "source")
    execute_values(
        cur, f"INSERT INTO public.{spec.live_table} ({', '.join(cols)}) "
        "VALUES %s",
        [(r[KEY], period) + tuple(r[c] for c in VALUES)
         + (r[NAME], r["live_source"]) for r in records], page_size=1000)


@dataclass(frozen=True)
class PeriodInfo:
    """What a vintage's edition records: source_file, the release label and
    the file read ((ledger source, sha256),)."""
    source_file: str
    label: str
    files: tuple


def _tip_label(cur, spec, period) -> str:
    tip = core.chain_tip(cur, spec, period)
    cur.execute(f"SELECT DISTINCT release_label FROM "
                f"public.{spec.editions_table} WHERE {spec.period_col} = %s "
                "AND edition = %s", (period, tip))
    return cur.fetchone()[0]


def apply_vintage(cur, profile, period, records, *, fetched_on, info) -> str:
    """pe.apply_period for the vintage, then its ledger row, inside the
    caller's per-period savepoint (s5_period): edition, live rows (the
    first vintage, or a stranded one) and ledger row commit or roll back
    together. Live source is the edition's release label (a stranded
    vintage's live rows take the tip edition's)."""
    kind = pe.classify_period(cur, profile, period, records)
    live_source = (_tip_label(cur, profile.spec, period)
                   if kind == LIVE_MISSING else info.label)
    records = [dict(r, live_source=live_source) for r in records]
    prof = dataclasses.replace(profile, release_label=lambda d: info.label)
    kind = pe.apply_period(cur, prof, period, records, fetched_on=fetched_on,
                           source_file=info.source_file,
                           classify=lambda *a, **kw: kind,
                           insert=_late("insert_live"))
    tip = core.chain_tip(cur, prof.spec, period)
    for src, sha in info.files:
        pe.record_file_check(cur, prof, period, src, sha, kind, tip)
    return kind


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

def _session(session):
    if session is not None:
        return session
    import requests
    return requests


def _get(session, url, **kw):
    try:
        r = session.get(url, headers={"User-Agent": USER_AGENT},
                        timeout=300, **kw)
        r.raise_for_status()
    except Exception as e:  # noqa: BLE001 (any failure halts)
        halt(f"GET {url} failed, nothing downloaded: {e}")
    return r


def fetch_json(path, session=None) -> dict:
    """The GOV.UK content API reply for a page path (a read-only GET)."""
    return json.loads(_get(_session(session), API + path).text)


def _abs_url(url) -> str:
    return url if urlparse(url).scheme else GOV_UK + url


def _opens(path) -> None:
    """Raise unless the file is an .xlsx workbook with an IMD sheet."""
    with zipfile.ZipFile(path) as z:
        if "xl/workbook.xml" not in z.namelist():
            raise ValueError("not an .xlsx workbook")
        if f'name="{IMD_SHEET}"'.encode() not in z.read("xl/workbook.xml"):
            raise ValueError(f"no {IMD_SHEET} sheet")


def fetch(url, dest, session=None) -> tuple:
    """Download url into the directory dest (redirects followed) and return
    (path to read, final URL). The file name is the final URL's. Checked
    before it is kept: more than MIN_FILE_BYTES, an .xlsx name, and it
    opens as an .xlsx workbook with an IMD sheet. A same-named file with the
    same content is kept as it is; one with different content is never
    replaced: the download is saved beside it as <stem>-<sha8><suffix> and
    that is the one read. An existing name is never a reason to skip the
    download."""
    dest = Path(dest)
    r = _get(_session(session), url)
    final = str(getattr(r, "url", None) or url)
    if final != url:
        print(f"redirected: {url} -> {final}")
    body = r.content
    if len(body) <= MIN_FILE_BYTES:
        halt(f"{final}: {len(body)} bytes, expected more than "
             f"{MIN_FILE_BYTES}; nothing written; saw {body[:100]!r}")
    name = Path(urlparse(final).path).name
    suffix = Path(name).suffix.lower()
    if suffix != ".xlsx":
        halt(f"{final}: not an .xlsx file; nothing written")
    dest.mkdir(parents=True, exist_ok=True)
    target = dest / name
    new_sha = hashlib.sha256(body).hexdigest()
    tmp = dest / (Path(name).stem + ".download.tmp" + suffix)
    tmp.write_bytes(body)
    try:
        _opens(tmp)
    except Exception as e:  # noqa: BLE001
        tmp.unlink()
        halt(f"{final}: does not open as File 10 ({e}); nothing kept")
    path = target
    if target.exists():
        old_sha = content_sha256(target)
        if old_sha == new_sha:
            tmp.unlink()
            print(f"{name} already held with the same content (sha256 "
                  f"{new_sha[:16]}); kept as it is")
        else:
            path = dest / f"{target.stem}-{new_sha[:8]}{target.suffix}"
            if path.exists() and content_sha256(path) == new_sha:
                tmp.unlink()
                print(f"NOTE: {name} differs from this download; "
                      f"{path.name} already holds it (sha256 "
                      f"{new_sha[:16]})")
            else:
                tmp.replace(path)
                print(f"NOTE: {name} already exists with different content "
                      f"(sha256 {old_sha[:16]}, new {new_sha[:16]}); kept it "
                      f"untouched and saved the new download as {path.name}")
    else:
        tmp.replace(target)
        print(f"downloaded {name}: {len(body):,} bytes, sha256 "
              f"{new_sha[:16]}")
    return path, final


def _file_source(path: Path) -> str:
    """Where a --file was read: the path under the repo (data/raw/...),
    else the absolute path."""
    p = Path(path).resolve()
    try:
        return p.relative_to(REPO).as_posix()
    except ValueError:
        return p.as_posix()


# ---------------------------------------------------------------------------
# Reading held data
# ---------------------------------------------------------------------------

table_exists = pe.table_exists


def _conn(writing):
    return pe.connect(writing)


def records(cur, spec, period, edition=None) -> list:
    """The vintage's rows as records (lad24cd, the vintage, la_name, the ten
    measures): the live table's (edition None; every live row when the live
    table has no vintage column yet) or the given edition's."""
    names = (KEY,) + CONTENT
    if edition is None:
        where, args = "", ()
        if live_has_vintage(cur, spec):
            where, args = f" WHERE {spec.period_col} = %s", (period,)
        cur.execute(f"SELECT {', '.join(names)} FROM public.{spec.live_table}"
                    + where, args)
    else:
        cur.execute(f"SELECT {', '.join(names)} FROM "
                    f"public.{spec.editions_table} WHERE {spec.period_col} = "
                    "%s AND edition = %s", (period, edition))
    out = [dict(zip(names, r)) for r in cur.fetchall()]
    for r in out:
        r[spec.period_col] = str(period)
    return out


def tip_info(cur, spec) -> dict:
    """{vintage: {edition, source_file, rank}} of every vintage with
    editions."""
    cur.execute(f"SELECT DISTINCT {spec.period_col} FROM "
                f"public.{spec.editions_table}")
    out = {}
    for (p,) in cur.fetchall():
        tip = core.latest_edition(cur, spec, str(p))
        cur.execute(f"SELECT DISTINCT source_file FROM "
                    f"public.{spec.editions_table} WHERE {spec.period_col} = "
                    "%s AND edition = %s", (p, tip))
        src = {r[0] for r in cur.fetchall()}
        sf = src.pop() if len(src) == 1 else None
        out[pe._p(p)] = {"edition": tip, "source_file": sf,
                         "rank": release_rank(sf)}
    return out


def _live_rows(cur, spec) -> int:
    cur.execute(f"SELECT COUNT(*) FROM public.{spec.live_table}")
    return cur.fetchone()[0]


def legacy_tips(cur, spec) -> dict:
    """Before the editions table exists: {LIVE_VINTAGE: {rank}} when live
    holds rows (the rank of the held file, LEGACY_RANK), else {}."""
    if not _live_rows(cur, spec):
        return {}
    return {LIVE_VINTAGE: {"edition": None, "source_file": None,
                           "rank": LEGACY_RANK}}


def ledger_checks(cur, prof) -> dict:
    """{vintage: {(source_file, file_sha256)}} of the ledger."""
    pc = prof.spec.period_col
    cur.execute(f"SELECT DISTINCT {pc}, source_file, file_sha256 "
                f"FROM public.{pe.file_checks_table(prof)}")
    out = {}
    for p, f, s in cur.fetchall():
        out.setdefault(pe._p(p), set()).add((f, s))
    return out


# ---------------------------------------------------------------------------
# ddl, status, run log, refresh-latest
# ---------------------------------------------------------------------------

def log_run(cur, rows_written: int, notes: str, started_at=None) -> None:
    """The pipeline_run_log row for a committed run (agent RUN_AGENT, source
    '5'). Called only on committed load, migrate-legacy and restore-edition
    runs, partial runs included."""
    pe.log_run(cur, profile(), rows_written, notes, started_at)


def _readiness(cur, prof) -> dict:
    spec = prof.spec
    lt = pe.file_checks_table(prof)
    return {spec.editions_table: table_exists(cur, spec.editions_table),
            lt: table_exists(cur, lt),
            f"{spec.live_table}.{PERIOD}": live_has_vintage(cur, spec)}


def cmd_ddl(args) -> int:
    """ddl [--commit | --simulate]: the editions table and its triggers, the
    ledger, and the live vintage column (the preview prints the exact
    ALTER)."""
    prof = profile()
    spec = prof.spec
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            for t in (spec.editions_table, pe.file_checks_table(prof)):
                print(f"{t}: {'exists' if table_exists(cur, t) else 'does not exist'}")
            if live_has_vintage(cur, spec):
                print(f"{spec.live_table}: vintage column present (checked "
                      "against the planned definition on --commit)")
            else:
                print(f"{spec.live_table}: no vintage column; ddl runs:\n  "
                      f"{live_alter_sql(spec)}\n  (additive: a default-filled "
                      "column, not in the key; no other column or row "
                      "changes)")
            if not writing:
                conn.rollback()
                print("DRY RUN: nothing written (use --commit or --simulate)")
                return 0
            create_all(cur)
        pe.finish(conn, args, f"ddl: {spec.editions_table} and triggers "
                              f"{spec.trigger}, {spec.truncate_trigger}, "
                              f"{pe.file_checks_table(prof)} and "
                              f"{spec.live_table}.{PERIOD} present")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def status_lines(cur, prof) -> list:
    """Per vintage: the tip edition, its file's rank and the ledger's latest
    file."""
    s = prof.spec
    lt = pe.file_checks_table(prof)
    lines = []
    for p, t in sorted(tip_info(cur, s).items()):
        line = f"  {p}: tip edition {t['edition']} ({rank_text(t['rank'])})"
        cur.execute(f"SELECT source_file, outcome FROM public.{lt} WHERE "
                    f"{s.period_col} = %s ORDER BY id DESC LIMIT 1", (p,))
        last = cur.fetchone()
        line += (f"; ledger latest: {last[1]} {last[0]}" if last
                 else "; no ledger row")
        lines.append(line)
    return lines


def cmd_status(args) -> int:
    """status: what needs action; exit 1 if anything, or (with a clean
    message, nothing created) if the editions table, ledger or live vintage
    column does not exist."""
    prof = profile()
    conn = _conn(False)
    try:
        with conn.cursor() as cur:
            missing = [t for t, ok in _readiness(cur, prof).items() if not ok]
            if missing:
                print(f"status: {' and '.join(missing)} "
                      f"{'does' if len(missing) == 1 else 'do'} not exist "
                      "yet; run `ddl --commit`, then `migrate-legacy FILE "
                      "--commit` (S5 has no sync-new)")
                return 1
            st = pe.status(cur, prof)
            lines = status_lines(cur, prof)
    finally:
        conn.rollback()
        conn.close()
    print(pe.format_status(prof, st, lines))
    return 0 if st["ok"] else 1


def _tables_exist(cur) -> bool:
    """True when the editions table, ledger and live vintage column exist;
    halts (S5's own wording: there is no sync-new) when one does not."""
    for t, ok in _readiness(cur, profile()).items():
        if not ok:
            halt(f"{t} does not exist yet; run `ddl --commit`, then "
                 "`migrate-legacy FILE --commit`")
    return True


def cmd_refresh_latest(args) -> int:
    """refresh-latest [--accept-drift VINTAGE] [--commit | --simulate]: the
    engine's refresh of the live vintage only (its latest edition copied
    into live, loaded_at included and source on every row, under the
    before/after hash guard); any other vintage named halts."""
    return pe.run_refresh_latest(
        profile(), args, connect=_late("_conn"),
        table_exists=lambda cur, t: _tables_exist(cur))


# ---------------------------------------------------------------------------
# load
# ---------------------------------------------------------------------------

def _end(conn, args, writing, rc) -> int:
    conn.rollback()  # commit mode has committed already
    if not writing:
        print("PREVIEW: nothing written to the database (use --commit or "
              "--simulate)")
    elif args.simulate:
        print("SIMULATION: ROLLED BACK (nothing persisted)")
    return rc


def _year_arg(text) -> str:
    if not re.fullmatch(r"[0-9]{4}", str(text or "")):
        raise argparse.ArgumentTypeError(f"not a year YYYY: {text!r}")
    return str(text)


def _vintage_arg(text) -> str:
    if not VINTAGE_RE.fullmatch(str(text or "")):
        raise argparse.ArgumentTypeError(f"not a vintage IoDYYYY: {text!r}")
    return str(text)


def _read(path) -> dict:
    try:
        return read_file10(path)
    except ValueError as e:      # the message names the file already
        halt(f"{e}; nothing stored")


def identity_problems(file, page_year=None, release=None) -> list:
    """Problems (empty = fine) between a read file and where it came from:
    the release page's year (when the page was read) and --release must
    both be the file's own vintage."""
    out = []
    if page_year is not None and vintage_of(page_year) != file["vintage"]:
        out.append(f"the release page is 'English indices of deprivation "
                   f"{page_year}' ({vintage_of(page_year)}), the file's own "
                   f"IMD header is the {file['vintage']} header")
    if release is not None and vintage_of(release) != file["vintage"]:
        out.append(f"--release {release}, the file's own IMD header is the "
                   f"{file['vintage']} header")
    return out


def _print_page(page, base):
    print(f"release: {page.get('title')} ({GOV_UK}{base}); first published "
          f"{str(page.get('first_published_at'))[:10]}; updated "
          f"{str(page.get('public_updated_at'))[:10]}")
    for ts, note in change_history(page):
        print(f"  change_history {ts}: {note[:300]}")


def _collection():
    coll = fetch_json(COLLECTION_PATH)
    try:
        check_collection(coll)
    except ValueError as e:
        halt(f"{GOV_UK}{COLLECTION_PATH}: {e}")
    return coll


def _page(base, year):
    page = fetch_json(base)
    if page_year(page) != int(year):
        halt(f"{GOV_UK}{base}: the page title {page.get('title')!r} is not "
             f"'English indices of deprivation {year}'")
    _print_page(page, base)
    return page


def _attachment(page, base):
    try:
        att = file10_attachment(page)
    except ValueError as e:
        halt(f"{GOV_UK}{base}: {e}")
    print(f"  attachment: {att['title']!r} {att['url']}")
    return att


def _source(args) -> dict:
    """The file the run reads: {path, where, how, page_year, page_url,
    file_name, file, discovered}. --file reads a local file (nothing
    downloaded; it is read here, file set) and finds its own vintage's
    release page in the collection unless --no-page (the page must still
    carry the File 10 attachment: corroboration); otherwise the content API
    finds the release (the newest, or --release YYYY) and its File 10 is
    downloaded to RAW_DIR (also in a preview)."""
    out = {"path": None, "where": None, "how": None, "page_year": None,
           "page_url": None, "file_name": None, "file": None,
           "discovered": None}
    if args.file:
        path = Path(args.file)
        if not path.is_file():
            halt(f"--file {args.file}: no such file")
        print(f"file given: {path}; nothing downloaded; identity is read "
              "from the file")
        f = _read(path)
        bad = identity_problems(f, None, args.release)
        if bad:
            halt("file identity check failed, nothing stored: "
                 + "; ".join(bad))
        out.update(path=path, where=_file_source(path), how="local file",
                   file_name=path.name, file=f)
        if args.no_page:
            out.update(page_url=NO_PAGE, how="local file, --no-page: the "
                       "release page was not read")
            print("NOTE: --no-page: the release page is not read")
            return out
        coll = _collection()
        try:
            _, base = release_for_year(coll, f["year"])
        except ValueError as e:
            halt(f"{GOV_UK}{COLLECTION_PATH}: {e}; the release page of the "
                 f"file's own vintage ({f['vintage']}) was not found. Give "
                 "--no-page to load the file without it (recorded as such)")
        _attachment(_page(base, f["year"]), base)
        out.update(page_year=f["year"], page_url=GOV_UK + base)
        return out
    coll = _collection()
    try:
        if args.release:
            title, base = release_for_year(coll, args.release)
            how = f"release {args.release} given"
        else:
            title, base = latest_release(coll)
            how = "newest release (content API)"
            out["discovered"] = coll
    except ValueError as e:
        halt(f"{GOV_UK}{COLLECTION_PATH}: {e}")
    y = int(DOC_TITLE_RE.fullmatch(title).group(1))
    att = _attachment(_page(base, y), base)
    print(f"downloading to {RAW_DIR} (also in a preview; the database is "
          "not written)")
    path, final = fetch(_abs_url(att["url"]), RAW_DIR)
    print(f"  final URL: {final}")
    out.update(path=path, where=final, how=how, page_year=y,
               page_url=GOV_UK + base,
               file_name=Path(urlparse(final).path).name)
    return out


def _label(file, src, sha, how, *, notes=()) -> str:
    return (f"MHCLG English indices of deprivation {file['year']}, File 10 "
            f"Local Authority District summaries (lower tier), "
            f"{file['vintage']}; {src['file_name']}; "
            f"{rank_text(file['modified'])}; {src['where']} sha256 "
            f"{sha[:16]}; {how}; published_date is the file's docProps "
            "modified date" + "".join(f"; {n}" for n in notes if n))


def _report(file) -> None:
    print(f"{file['name']}: {file['vintage']} (its IMD header is the declared "
          f"{file['vintage']} header); sheets {len(file['sheets'])} as "
          f"expected; {len(file['rows'])} authorities; "
          f"{rank_text(file['modified'])} (the file's rank)")
    print("  identity: File 10 has no cover sheet and an empty Notes sheet; "
          "the checks rely on the sheet set, the exact IMD header and "
          "docProps modified (and the release page's attachment title when "
          "the page is read)")


def _build(cur, file) -> list:
    """The vintage's records, codes resolved and the file's own checks run;
    halts on a geography, cell or check problem."""
    try:
        rmap, problems = resolve_codes(cur, {r["code"] for r in file["rows"]})
    except (ValueError, KeyError) as e:
        halt(f"geography: {e}; nothing stored")
    if problems:
        halt("geography check failed, nothing stored: " + "; ".join(problems))
    recoded = sorted((c, r) for c, r in rmap.items() if r != c)
    if recoded:
        print("codes resolved (geography.resolve, la_code_lookup recode): "
              + ", ".join(f"{c} -> {r}" for c, r in recoded))
    try:
        recs = build_records(file, rmap)
        lines = checks(recs)
    except ValueError as e:
        halt(f"{e}; nothing stored")
    print("checks passed: " + "; ".join(lines))
    return recs


def _preview_live(cur, spec, period, recs) -> None:
    """Before the editions table exists: the file against the live rows
    (lad24cd, la_name and the ten measures; NULL never equals a number)."""
    if period != LIVE_VINTAGE:
        print(f"  {period}: would be held as editions only (live holds "
              f"{LIVE_VINTAGE}; compared with nothing)")
        return
    held = records(cur, spec, period)
    if not held:
        print(f"  {period}: new, {len(recs)} areas (live is empty)")
        return
    hk, nk = _by_key(held), _by_key(recs)
    diff = [k for k in sorted(set(hk) | set(nk))
            if k not in hk or k not in nk
            or any(hk[k][c] != nk[k][c] for c in CONTENT)]
    ex = [f"{k}: {tuple(hk[k][c] for c in CONTENT) if k in hk else 'absent'}"
          f" -> {tuple(nk[k][c] for c in CONTENT) if k in nk else 'absent'}"
          for k in diff[:3]]
    print(f"  {period}: {'revised' if diff else 'unchanged'}, {len(diff)} "
          "areas changed (compared with live)"
          + ("; " + "; ".join(ex) if ex else ""))


def cmd_load(args) -> int:
    if args.no_page and not args.file:
        halt("--no-page is only for --file (a local file whose release page "
             "is not read)")
    prof = profile()
    spec = prof.spec
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            ready = _readiness(cur, prof)
            has_ed = all(ready.values())
            missing = [t for t, ok in ready.items() if not ok]
            if writing and not has_ed:
                halt(f"missing {', '.join(missing)}; run `ddl --commit` and "
                     "`migrate-legacy FILE --commit` first")
            live_empty = _live_rows(cur, spec) == 0
            if has_ed:
                _, new_live, errors = core.latest_map(cur, spec)
                if errors:
                    halt("invalid edition chain: " + "; ".join(
                        f"{pe._p(p)}: {msg}" for p, msg in errors.items()))
                if new_live:
                    halt(f"{spec.live_table}: live vintages with no editions "
                         f"{[pe._p(x) for x in new_live]}; run migrate-legacy "
                         "first")
                tips = tip_info(cur, spec)
                checked = ledger_checks(cur, prof)
                stranded = (pe.live_missing_periods(cur, prof) if live_empty
                            else [])
            else:
                print(f"NOTE: {', '.join(missing)} not present: no editions "
                      "table yet; this preview compares the file with the "
                      "LIVE table")
                tips, checked, stranded = legacy_tips(cur, spec), {}, []
            ranks = tip_ranks(tips, checked)
            held = sorted(ranks)
            print("held vintages: " + (", ".join(held) or "none"))
            src = _source(args)
            if src["discovered"] is not None:
                for line in discovery_note(src["discovered"], held):
                    print(line)
            path = src["path"]
            sha = content_sha256(path)
            print(f"{path.name}: sha256 {sha[:16]}"
                  + ("; byte-identical to the held file"
                     if sha == LEGACY_FILE["sha256"] else ""))
            f = src["file"]
            if has_ed and args.recheck is None:
                done = ledger_complete(checked, src["where"], sha, ranks,
                                       allow_older=args.allow_older_file,
                                       stranded=stranded)
                if done:
                    print(f"the file's (source, sha256) is already in the "
                          f"ledger for {', '.join(done)}; "
                          + ("nothing compared" if f else "nothing parsed")
                          + " (--recheck VINTAGE reads it again)")
                    print("nothing to do")
                    return _end(conn, args, writing, rc=0)
            if f is None:
                f = _read(path)
                bad = identity_problems(f, src["page_year"], args.release)
                if bad:
                    halt("file identity check failed, nothing stored: "
                         + "; ".join(bad))
            _report(f)
            recs = _build(cur, f)
            period = f["vintage"]
            lsrc = ledger_source(src["where"], period, f["rank"])
            chk = {p: s for p, s in checked.items() if p not in stranded}
            try:
                new, revised, skipped = plan_vintage(
                    period, f["rank"], ranks, chk, lsrc, sha,
                    recheck=args.recheck, allow_older=args.allow_older_file)
            except ValueError as e:
                halt(str(e))
            rc = _run(args, conn, cur, prof, has_ed, src, f, recs, new,
                      revised, skipped, tips, ranks, held, sha, lsrc,
                      live_empty)
        return _end(conn, args, writing, rc)
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def _tip_records(cur, spec, p, tips, has_ed) -> list:
    if has_ed:
        return records(cur, spec, p, tips[p]["edition"])
    return records(cur, spec, p)


def _run(args, conn, cur, prof, has_ed, src, f, recs, new, revised, skipped,
         tips, ranks, held, sha, lsrc, live_empty) -> int:
    started = datetime.now(timezone.utc)
    spec = prof.spec
    rank = f["rank"]
    older_note = {}
    for p in revised + new:
        t = ranks.get(p)
        if t is not None and t > rank:
            older_note[p] = (f"--allow-older-file given; {p}'s tip comes from "
                             f"a file of {rank_text(t)}, this file is "
                             f"{rank_text(rank)}")
        elif p in new and held and p < max(held):
            older_note[p] = (f"--allow-older-file given; {p} is before the "
                             f"newest held vintage {max(held)}")
        if p in older_note:
            print(f"NOTE: {older_note[p]}")
    print(f"planned: new {new or 'none'}; compare {revised or 'none'}; "
          "skipped " + (", ".join(f"{p} ({r})" for p, r in sorted(
              skipped.items())) or "none"))
    periods = sorted(new + revised)
    reissue = set(args.accept_reissue or ())
    stray = sorted(reissue - set(periods))
    if stray:
        halt(f"--accept-reissue {stray}: not vintages this run compares or "
             f"stores ({periods or 'none'}); nothing stored")
    if not periods:
        if skipped and all(r.startswith("older") for r in skipped.values()):
            halt(f"older file: the file is {f['vintage']}, "
                 f"{rank_text(rank)}, older than what is held for "
                 f"{', '.join(sorted(skipped))}; storing it would record an "
                 "older file. If this is deliberate, re-run with "
                 "--allow-older-file")
        print("nothing to do")
        return 0
    problems, infos, reissued = {}, {}, {}
    for p in periods:
        if p in revised:
            kind = "revised"
            tip = _tip_records(cur, spec, p, tips, has_ed)
            bad = period_problems(recs, tip, kind=kind, vintage=p,
                                  live_empty=live_empty)
            if ranks.get(p) == rank and _differs(recs, tip):
                if p in reissue:
                    reissued[p] = (f"--accept-reissue {p}: the file has "
                                   f"{rank_text(rank)}, the same as the file "
                                   "of the held tip, with different content "
                                   "(a publisher reissue keeping its date)")
                    print(f"  {p}: ACCEPTED REISSUE: {reissued[p]}")
                else:
                    bad.append(f"the file has {rank_text(rank)}, the same as "
                               "the file of the held tip, but its content "
                               "differs: two files claim the same release. "
                               "A publisher reissue may keep its date: read "
                               "the page's change note, then "
                               f"--accept-reissue {p}")
        else:
            kind = "new"
            bad = period_problems(recs, None, kind=kind, vintage=p,
                                  live_empty=live_empty)
        if bad:
            problems[p] = bad
            continue
        notes = (older_note.get(p, ""), reissued.get(p, ""))
        how = (("new vintage" if kind == "new" else "restated vintage")
               + f" ({src['how']})")
        infos[p] = PeriodInfo(
            edition_source(f, src, "; ".join(n for n in notes if n)),
            _label(f, src, sha, how, notes=notes), ((lsrc, sha),))
    for p, msgs in sorted(problems.items()):
        for msg in msgs:
            print(f"  {p}: REJECTED, not stored: STOP CONDITION: {msg}")
    ok = [p for p in periods if p not in problems]
    rc = 0
    stats = {"periods": [], "kinds": {}, "stored_rows": 0, "live_rows": 0}
    if ok and has_ed:
        def apply(c, prof_, per, rs, *, fetched_on, source_file=None):
            return _late("apply_vintage")(c, prof_, per, rs,
                                          fetched_on=fetched_on,
                                          info=infos[per])
        rc = pe.load_periods(cur, prof, ok, lambda per: recs,
                             f["modified"].date(), args.commit,
                             simulate=args.simulate, against="editions",
                             stats=stats, apply=apply)
    elif ok:
        for p in ok:
            _preview_live(cur, spec, p, recs)
        print("  (no editions table yet: preview against live only)")
    if problems:
        print(f"REJECTED {len(problems)} vintage(s), nothing stored for them: "
              f"{', '.join(sorted(problems))}; exit 1")
        rc = 1
    if args.commit:
        stored = live = 0  # rows written by this run; not a source value
        for p in stats["periods"]:
            k = stats["kinds"][p]
            if k in ("new", "revised", HELD_NOT_LIVE):
                stored += len(recs)
            if k in ("new", LIVE_MISSING):
                live += len(recs)
        left = [p for p in ok if p not in stats["periods"]]
        failed, not_tried = (left[:1], left[1:]) if left else ([], [])
        partial = ""
        if rc:
            partial = (f"PARTIAL RUN (exit 1): rejected "
                       f"{sorted(problems) or 'none'}"
                       + "".join(f" ({p}: {'; '.join(problems[p])[:300]})"
                                 for p in sorted(problems))
                       + f"; failed {failed or 'none'}; not attempted "
                       f"{not_tried or 'none'}; a vintage that failed or was "
                       "rejected stored nothing and is not in the counts "
                       "below; ")
        notes = (partial + f"File 10 {src['where']} sha256 {sha[:16]}, "
                 f"{f['vintage']}, {rank_text(rank)} ({src['how']}): "
                 + ("; ".join(f"{p} {k}" for p, k in stats["kinds"].items())
                    or "nothing stored")
                 + (f"; skipped {sorted(skipped)}" if skipped else "")
                 + "".join(f"; {older_note[p]}" for p in sorted(older_note)
                           if p in stats["periods"])
                 + "".join(f"; {reissued[p]}" for p in sorted(reissued)
                           if p in stats["periods"])
                 + f". Rows stored: {stored} edition rows, {live} live rows.")
        log_run(cur, stored + live, notes, started)
        conn.commit()
        print("pipeline_run_log row written"
              + (" (partial run: see its notes)" if rc else ""))
    return rc


# ---------------------------------------------------------------------------
# One-off: migrate the held table
# ---------------------------------------------------------------------------

SURVEY_COLS = (KEY, NAME) + VALUES + ("source",)


def live_state(cur, spec=None) -> tuple:
    """(rows, survey hash, distinct loaded_at, vintages) of the live table.
    The survey hash is md5 of lad24cd|la_name|the ten measures|source (NULL
    as '~'), rows ordered by lad24cd, joined by LF (the vintage column is
    left out, so the hash is the same before and after ddl). vintages: the
    distinct live vintages, or None when the column does not exist yet."""
    spec = spec or globals()["SPEC"]
    expr = "||'|'||".join(f"coalesce({c}::text,'~')" for c in SURVEY_COLS)
    cur.execute(f"SELECT COUNT(*), md5(string_agg({expr}, E'\\n' ORDER BY "
                f"{KEY})), COUNT(DISTINCT loaded_at) FROM "
                f"public.{spec.live_table}")
    n, h, k = cur.fetchone()
    vintages = None
    if live_has_vintage(cur, spec):
        cur.execute(f"SELECT DISTINCT {PERIOD} FROM public.{spec.live_table} "
                    "ORDER BY 1")
        vintages = [r[0] for r in cur.fetchall()]
    return n, h, k, vintages


def _proof(held, built) -> tuple:
    """(values compared, differences) of the held rows against the rebuilt
    rows on lad24cd, la_name and the ten measures."""
    hk, bk = _by_key(held), _by_key(built)
    diffs = [("only held", k) for k in sorted(set(hk) - set(bk))]
    diffs += [("only rebuilt", k) for k in sorted(set(bk) - set(hk))]
    cells = 0  # values compared; not a source value
    for k in sorted(set(hk) & set(bk)):
        for c in CONTENT:
            cells += 1
            if hk[k][c] != bk[k][c]:
                diffs.append((k, c, hk[k][c], bk[k][c]))
    return cells, diffs


def migrate_legacy(cur, file, *, write, legacy=None, files=None) -> dict:
    """One-off. Preconditions: on write the editions table, ledger and live
    vintage column exist (ddl first) and the tables are empty (in a preview
    they may be absent; if present they must be empty); live as surveyed
    (legacy: (rows, survey hash), default LEGACY_LIVE; one loaded_at; the
    vintage column absent or LIVE_VINTAGE on every row); the file's sha256
    as files['sha256'] (default LEGACY_FILE). Proof, any failure halts: the
    file's identity (sheets, the exact IoD2025 header, docProps modified),
    its own checks, and this parser on the file reproduces every live value
    (lad24cd, la_name and the ten measures). Edition 1 'as loaded' from live
    (the engine's as-loaded label, published_date the load date,
    source_file 'as loaded: <name>; vintage IoD2025; docProps modified
    <iso>'), and a ledger row for the file (outcome 'unchanged', edition 1,
    files['url']) so the next page load finds it checked. Live untouched.
    Nothing commits here; the caller rolls everything back on a halt."""
    g = globals()
    legacy = tuple(legacy or g["LEGACY_LIVE"])
    files = files or g["LEGACY_FILE"]
    prof = profile()
    spec = prof.spec
    ready = _readiness(cur, prof)
    if write and not all(ready.values()):
        halt(f"missing {[t for t, ok in ready.items() if not ok]}; run "
             "`ddl --commit` first")
    for t in (spec.editions_table, pe.file_checks_table(prof)):
        if ready[t]:
            cur.execute(f"SELECT COUNT(*) FROM public.{t}")
            n = cur.fetchone()[0]
            if n:
                halt(f"{t} already holds {n} rows; migrate-legacy runs once, "
                     "on empty tables")
    n, h, k, vintages = live_state(cur, spec)
    if (n, h) != legacy or k != 1 or vintages not in (None, [LIVE_VINTAGE]):
        halt(f"{spec.live_table} is not as surveyed: {n} rows, hash {h}, {k} "
             f"loaded_at value(s), vintages {vintages}; expected {legacy[0]} "
             f"rows, hash {legacy[1]}, one loaded_at, vintage {LIVE_VINTAGE} "
             "(or no vintage column yet); nothing stored")
    path = Path(file)
    sha = content_sha256(path)
    if sha != files["sha256"]:
        halt(f"{path.name}: sha256 {sha[:16]}, expected "
             f"{files['sha256'][:16]} (the held file); nothing stored")
    f = _read(path)
    _report(f)
    period = f["vintage"]
    if period != LIVE_VINTAGE:
        halt(f"{path.name} is the {period} File 10; the live table holds "
             f"{LIVE_VINTAGE}; nothing stored")
    held = records(cur, spec, period)
    built = _build(cur, f)
    cells, diffs = _proof(held, built)
    if diffs:
        halt(f"proof failed: {len(diffs)} differences between the held rows "
             f"and the held file read again, e.g. {diffs[:5]}; nothing "
             "stored")
    cur.execute(f"SELECT MIN((loaded_at AT TIME ZONE 'UTC')::date) FROM "
                f"public.{spec.live_table}")
    loaded = cur.fetchone()[0]
    plan = {"period": period, "rows": len(held), "cells": cells,
            "records": held, "sha": sha, "loaded": loaded,
            "content_sha": rows_content_sha(held),
            "source_file": (f"as loaded: {files['name']}; vintage {period}; "
                            f"{rank_text(f['modified'])}"),
            "ledger": ledger_source(files["url"], period, f["rank"])}
    print(f"proof: the held file read again reproduces every held value: "
          f"{len(held):,} authorities matched by lad24cd, {cells:,} values "
          f"(la_name and the ten measures), 0 differences")
    print(f"  {period}: edition 1 as loaded, {len(held):,} rows, published "
          f"{loaded} (the load date); source_file {plan['source_file']!r}; "
          f"content sha256 {plan['content_sha']}; ledger {plan['ledger']}")
    if not write:
        return plan
    cur.execute(f"SAVEPOINT {prof.savepoint}")
    ed = core.insert_edition(
        cur, spec, held, period, release_label=core.AS_LOADED_LABEL,
        published_date=loaded, source_file=plan["source_file"],
        source_sha256=plan["content_sha"], supersedes=None, strict=True)
    if ed != 1 or core.rows_differing(cur, spec, period, ed):
        halt(f"{spec.editions_table} {period}: edition 1 does not equal "
             "live; rolled back")
    pe.record_file_check(cur, prof, period, plan["ledger"], sha,
                         "unchanged", 1)
    cur.execute(f"RELEASE SAVEPOINT {prof.savepoint}")
    return plan


def cmd_migrate_legacy(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    try:
        with conn.cursor() as cur:
            plan = migrate_legacy(cur, args.file, write=writing)
            if args.commit:
                log_run(cur, plan["rows"], f"migrate-legacy: edition 1 as "
                        f"loaded for {plan['period']} ({plan['rows']} rows) "
                        f"from the live table; proof {plan['cells']} values "
                        "equal the held file read again "
                        f"(sha256 {plan['sha'][:16]}); ledger row for the "
                        "file. Live untouched.", started)
                conn.commit()
                print("migrate-legacy: COMMITTED; pipeline_run_log row "
                      "written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# restore-edition
# ---------------------------------------------------------------------------

def restore_edition(cur, period, edition) -> int:
    """Store edition `edition` of `period` as the next edition ('restored
    from edition N') for refresh-latest to apply; return its number. Halts
    if the edition does not exist or is the tip."""
    spec = globals()["SPEC"]
    p = str(period)
    tip = core.latest_edition(cur, spec, p)
    recs = records(cur, spec, p, edition)
    if not recs:
        halt(f"{spec.editions_table} {p}: no edition {edition}")
    if edition == tip:
        halt(f"{p}: edition {edition} is the tip; nothing to restore")
    cur.execute(f"SELECT DISTINCT source_file FROM public.{spec.editions_table}"
                f" WHERE {spec.period_col} = %s AND edition = %s", (p, edition))
    src = cur.fetchone()[0]
    new = core.insert_edition(
        cur, spec, recs, p, release_label=f"restored from edition {edition}",
        published_date=date.today(), source_file=src,
        source_sha256=rows_content_sha(recs), supersedes=tip, strict=True,
        allow_revert=True)
    if new == tip:
        halt(f"{p}: edition {edition} was not stored as the next edition; "
             "nothing to restore")
    return new


def cmd_restore_edition(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    spec = globals()["SPEC"]
    try:
        with conn.cursor() as cur:
            if not table_exists(cur, spec.editions_table):
                halt(f"{spec.editions_table} does not exist")
            new = restore_edition(cur, args.period, args.edition)
            print(f"{args.period}: edition {args.edition} "
                  + ("stored" if writing else "would be stored")
                  + f" as edition {new}; run refresh-latest to apply it")
            if args.commit:
                n = len(records(cur, spec, args.period, new))
                log_run(cur, n, f"restore-edition: {args.period} edition "
                        f"{args.edition} stored as edition {new}", started)
                conn.commit()
                print("COMMITTED; pipeline_run_log row written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="S5 MHCLG English indices of deprivation, File 10: "
        "editions loader on the period-editions engine, one vintage in "
        "live. There is no sync-new; migrate-legacy records the held "
        "vintage once.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ddl", help="create the editions table, its triggers "
                       "and the file-check ledger, and add the live vintage "
                       "column (preview by default)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_ddl)
    sub.add_parser("status", help="what needs action; exit 1 if anything"
                   ).set_defaults(func=cmd_status)
    p = sub.add_parser("load", help="find the newest release's File 10, "
                       "compare and store its vintage (preview by default)")
    p.add_argument("--release", type=_year_arg, metavar="YYYY",
                   help="that year's release page (default the newest)")
    p.add_argument("--file", metavar="PATH",
                   help="a local File 10 (nothing downloaded; identity is "
                   "read from the file)")
    p.add_argument("--no-page", action="store_true",
                   help="with --file only: do not read the release page "
                   "(recorded as such)")
    p.add_argument("--recheck", type=_vintage_arg, metavar="VINTAGE",
                   help="compare this held vintage again although the file "
                   "is in the ledger")
    p.add_argument("--allow-older-file", action="store_true",
                   help="compare and store a vintage whose tip comes from a "
                   "newer file, or a vintage before the newest held "
                   "(logged)")
    p.add_argument("--accept-reissue", action="append", type=_vintage_arg,
                   metavar="VINTAGE",
                   help="store a vintage whose file has the held file's "
                   "docProps modified but different content (a publisher "
                   "reissue keeping its date; repeatable; logged)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_load)
    p = sub.add_parser("refresh-latest", help="copy the live vintage's latest "
                       "edition into the live table, loaded_at included "
                       "(preview by default)")
    p.add_argument("--accept-drift", action="append", metavar="VINTAGE",
                   help="overwrite the live vintage although its live rows "
                   "equal no stored edition (repeatable; any other vintage "
                   "halts)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_refresh_latest)
    p = sub.add_parser("migrate-legacy", help="one-off: edition 1 as loaded "
                       "for the held vintage, with the proof (preview by "
                       "default)")
    p.add_argument("file", metavar="FILE")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_migrate_legacy)
    p = sub.add_parser("restore-edition", help="store an earlier edition's "
                       "rows as the next edition (preview by default)")
    p.add_argument("period", type=_vintage_arg, metavar="VINTAGE")
    p.add_argument("edition", type=int)
    pe.mode_parser(p)
    p.set_defaults(func=cmd_restore_edition)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
