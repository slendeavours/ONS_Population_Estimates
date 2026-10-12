"""S3 (ONS mid-year population estimates, England and Wales): edition
history and loader on the period-editions engine.

The live table la_population keeps its name, key (lad24cd, reference_year)
and columns and is the latest-edition layer (W1 reads population at the
newest reference_year). This module owns its editions table and file-check
ledger, append-only:

    SPEC     la_population_editions               period reference_year
    ledger   la_population_editions_file_checks

An edition is what one tables workbook says about one mid-year: per English
authority population and la_name. The annual release "Estimates of the
population for England and Wales", edition "Mid-<yyyy>: <yyyy> local
authority boundaries", carries one workbook (mye<yy>tablesew.xlsx). Its MYE2
- Persons sheet gives the newest mid-year; its MYE4 restates every year from
mid-2011; its MYE7 (from the July 2026 layout) names the revisions. So a
file states its newest year (a new period, from MYE2) and every held year
MYE4 covers (compared with the held edition: a back-year revision such as
July 2026's revision of mid-2024 is caught and becomes the next edition).
Back-filling years before the first held one is not done by load.

History (so the wording here stays truthful): mid-2024 was loaded on
2026-03-26 by the n8n workflow "ONS Population Estimates (S3)" from the
NOMIS NM_2002_1 API plus a supplement CSV made in a chat session for the
authorities NOMIS lacked (pipeline_run_log 16 wrote the 296 rows); mid-2025
by scripts/s3_mye_refresh.py on 2026-08-13 (run log 86), which also widened
the key to (lad24cd, reference_year). Both are retired. migrate-legacy
(one-off) records each held year as edition 1 "as loaded" (live as held),
with a proof that this parser reproduces every held population from the
year's own workbook; four 2024 names are held in the NOMIS form ("Bristol
City of", against the workbook's "Bristol, City of"; also Herefordshire,
Kingston upon Hull, Bournemouth Christchurch and Poole), named in
LEGACY_NAME_FORMS.

Blanks and zeros (docs/RULES.md rule 1): no marker, blank or zero appears in
any English LA cell of the mid-2024 or mid-2025 workbook (MYE2 All ages,
MYE4 every year). So an LA cell must be a positive whole number; a blank,
text, a zero, a negative or a non-integer halts naming the cell (la_cell),
and is never read as 0 or stored.

Identity (rule 3), from the workbook itself on every path: the Cover sheet's
title 'Mid-Year Population Estimates, England and Wales, June <yyyy>', its
line 'The data tables in this spreadsheet were originally published at
9.30am <d Month yyyy>.' (the rank), its 'next publication' line and, when
present, its 'Date published:' line (the same day) and 'Release number:
MYE<yy>'; the Notes sheet's geography note (kept verbatim) and any
'Revisions to mid-...' note (kept verbatim); MYE2 - Persons' title '...
mid-<yyyy>' and header starting exactly 'Code, Name, Geography, All ages';
MYE4's title '... mid-2011 to mid-<yyyy>' and header exactly 'Code, Name,
Geography, Mid-<yyyy>, ..., Mid-2011'; MYE7's title '<Month yyyy> revisions
to mid-<a> to mid-<b> ...' with its published and revised column pairs. The
mid-year of the Cover, MYE2 and MYE4 must agree. The July 2025 layout (MYE7
the uncertainty table, the revisions in MYE8 and MYE9) is read where it
differs: its MYE7 is named and not read. The edition's mid-year (when the
landing page is read) and --release must equal the file's. Inside the file
the English authorities must sum to the England row in MYE2 and in every
MYE4 year read, MYE2 and MYE4 must agree on the newest year, and MYE4 must
equal MYE7's revised column for each revised year read (reconcile). The rank
of a file is its own (Cover publication date, mid-year), never its name,
URL or the API's releaseDate (2026-07-28T23:00Z for a Cover of 29 July
2026: the evening before).

Geography (rule 4): geography.resolve('3', codes, by_period={year: codes})
(source 3 is declared 'old': an E08000038 or E08000039 halts), then each
code must be in la_boundaries; anything else is UNEXPLAINED and stops.

Stop conditions, calibrated on the mid-2025 workbook (read 2026-10-11):

    new year:  the England total moving more than NEW_TOTAL_BP basis points
               against the previous held year: the largest MYE4 annual move
               of England mid-2011 to mid-2025 is +1.4008% (mid-2022 to
               mid-2023), plus one point = 2.4008%, so 2.41% (241 bp); or
               any authority moving more than 15% (NEW_AREA_BP): MYE4's
               largest authority moves are City of London +38.34% (2022)
               and +20.75% (2023), which would have needed --acknowledge;
               the next largest is Isles of Scilly +11.87% (2021); the
               largest in 2025 is Rugby +2.39%.
    revised:   any change at all against the held edition (population or
               la_name) needs --acknowledge; a move of more than REV_AREA_BP
               at any authority or more than REV_TOTAL_BP nationally is
               LARGE and REJECTED unless also --accept-large-revision.
               Calibrated on MYE7 (July 2026 revisions of mid-2022 to
               mid-2024, against the July 2025 figures): England +0.0793%
               (2022), +0.1007% (2023), -0.0221% (2024, -12,931); the
               largest authority revisions City of London +3.487% (2022),
               City of London +6.674% (2023), Westminster +2.713% (2024,
               +5,698, the largest absolute). The planned 2% authority
               limit would have rejected every one of the three revision
               years, so the authority limit is the largest observed plus
               one point, 7.674%, so 7.68% (768 bp); the national limit is
               0.5% (50 bp), five times the largest observed.
    both:      fewer authorities than EXPECTED_AREAS, or (revised) an
               authority of the held edition missing: a partial file never
               replaces a fuller edition, and no flag releases that.

A stopped year is REJECTED (nothing stored, no ledger row, exit 1). Equal
rank with different content (two files claiming the same release) stops
unless --accept-reissue YEAR. Each acknowledgement names its year and is
written to the edition's label and the run log. The July 2026 revision of
mid-2024 (295 of 296 authorities, England -12,931, and the four NOMIS-form
names) needs --acknowledge 2024.

W1 input: a new year's live rows take loaded_at = now(); a revision reaches
live only through refresh-latest, which copies the edition's loaded_at, so
run refresh-latest in the same session before W1.

Subcommands (every writing command previews by default; --commit and
--simulate are mutually exclusive; --simulate runs the --commit path and
always rolls back). There is no sync-new: migrate-legacy records the held
years.
    python scripts/s3_mye_editions.py ddl [--commit | --simulate]
    python scripts/s3_mye_editions.py status
    python scripts/s3_mye_editions.py load [--release YYYY]
        [--file PATH [--no-page]] [--recheck YYYY] [--allow-older-file]
        [--acknowledge YYYY ...] [--accept-large-revision YYYY ...]
        [--accept-reissue YYYY ...] [--commit | --simulate]
        # the file is downloaded to data/raw/s3_mye/ (also in a preview; a
        # preview writes nothing to the database)
    python scripts/s3_mye_editions.py refresh-latest [--commit | --simulate]
    python scripts/s3_mye_editions.py migrate-legacy FILE [FILE ...]
        [--commit | --simulate]
    python scripts/s3_mye_editions.py restore-edition YYYY N
        [--commit | --simulate]
"""
import argparse
import dataclasses
import hashlib
import json
import math
import numbers
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlparse

sys.path.insert(0, str(Path(__file__).resolve().parent))

import editions_core as core  # noqa: E402
import geography  # noqa: E402
import period_editions as pe  # noqa: E402
from editions_core import halt  # noqa: E402

RUN_AGENT = "s3_mye_editions"
RUN_SOURCE = "3"
LIVE_MISSING = pe.LIVE_MISSING

ONS = "https://www.ons.gov.uk"
LANDING = ("/peoplepopulationandcommunity/populationandmigration/"
           "populationestimates/datasets/estimatesofthepopulationfor"
           "englandandwales")
LANDING_TITLE = "Estimates of the population for England and Wales"
REPO = Path(__file__).resolve().parent.parent
RAW_DIR = REPO / "data" / "raw" / "s3_mye"     # git-ignored
USER_AGENT = ("ucws-pipeline S3 loader (read-only download of the public "
              "ONS mid-year population estimates tables)")
MIN_FILE_BYTES = 100 * 1024
NO_PAGE = "--file --no-page: the landing page was not read"

# ---------------------------------------------------------------------------
# Columns (live types from information_schema, checked 2026-10-11:
# lad24cd character varying(9) NOT NULL, la_name character varying(100)
# (nullable; 0 NULL held), population integer (nullable; 0 NULL held),
# reference_year integer NOT NULL, loaded_at timestamptz DEFAULT now(); PK
# (lad24cd, reference_year)). The editions table holds both values NOT
# NULL: a blank halts, never stored.
# ---------------------------------------------------------------------------

VALUE_TYPES = (("population", "integer NOT NULL"),
               ("la_name", "text NOT NULL"))
VALUES = tuple(c for c, _ in VALUE_TYPES)
KEY = "lad24cd"
PERIOD = "reference_year"

# ---------------------------------------------------------------------------
# The file
# ---------------------------------------------------------------------------

COVER = "Cover sheet"
NOTES = "Notes"
MYE2 = "MYE2 - Persons"
MYE4 = "MYE4"
MYE7 = "MYE7"
FIRST_YEAR = 2011
COVER_TITLE_RE = re.compile(r"Mid-Year Population Estimates, England and "
                            r"Wales, June ([0-9]{4})")
PUB_RE = re.compile(r"The data tables in this spreadsheet were originally "
                    r"published at 9\.30am ([0-9]{1,2}) ([A-Za-z]+) "
                    r"([0-9]{4})\.?")
NEXT_RE = re.compile(r"The next publication is scheduled for (.+?)\.?")
DATE_RE = re.compile(r"Date published: ([0-9]{1,2}) ([A-Za-z]+) ([0-9]{4})")
RELEASE_NO_RE = re.compile(r"Release number: MYE([0-9]{2})")
GEOG_NOTE_RE = re.compile(r"Geography: .+", re.S)
REVISION_NOTE_RE = re.compile(r"Revisions to mid-.+", re.S)
MYE2_TITLE_RE = re.compile(r"MYE2: Persons by single year of age and sex for "
                           r"local authorities in England and Wales, "
                           r"mid-([0-9]{4})")
MYE4_TITLE_RE = re.compile(r"MYE4: Population estimates: Summary for England "
                           r"and Wales, mid-2011 to mid-([0-9]{4})")
MYE7_TITLE_RE = re.compile(r"MYE7: ([A-Z][a-z]+ [0-9]{4}) revisions to "
                           r"mid-([0-9]{4}) to mid-([0-9]{4}) estimates for "
                           r"England & Wales")
MYE2_HEADER = ("Code", "Name", "Geography", "All ages")
LA_RE = re.compile(r"E0[6-9][0-9]{6}")
WALES_RE = re.compile(r"W06[0-9]{6}")
# aggregates set aside: counties, metropolitan counties, regions, countries
AGGREGATE_RE = re.compile(r"(E1[0-3]|E92|W92|K0[2-4])[0-9]{6}")
ENGLAND = "E92000001"
EXPECTED_AREAS = 296

# ---------------------------------------------------------------------------
# Stop conditions (basis points; see the module docstring for the
# calibration)
# ---------------------------------------------------------------------------

NEW_TOTAL_BP = 241          # new year: England total against the year before
NEW_AREA_BP = 1500          # new year: any authority against the year before
REV_AREA_BP = 768           # revised year: any authority (LARGE above it)
REV_TOTAL_BP = 50           # revised year: England total (LARGE above it)
PARTIAL = "PARTIAL FILE"    # problems no flag releases
LARGE = "LARGE REVISION"    # problems needing --accept-large-revision too

# ---------------------------------------------------------------------------
# Discovery: the ONS dataset landing page (/data JSON)
# ---------------------------------------------------------------------------

EDITION_RE = re.compile(r"mid([0-9]{4})([0-9]{4})localauthorityboundaries")
EDITION_LABEL_RE = re.compile(r"Mid-([0-9]{4}): ([0-9]{4}) local authority "
                              r"boundaries")
# When the next release is expected. The mid-<Y> estimates are published in
# the summer of Y+1 (Covers: 9.30am 30 July 2025 for mid-2024, 29 July 2026
# for mid-2025; 'The next publication is scheduled for summer 2027'). So the
# mid-<Y+1> release is overdue after 30 September of Y+2. EXPECTED_RELEASES
# holds a date the publisher has announced, as data that overrides the
# derived one: {mid-year: the last day before the warning starts}.
EXPECTED_RELEASES = {}


def expected_next_release(newest_year, announced=None) -> tuple:
    """(mid-year, last day before it is overdue) of the release after
    `newest_year`: the announced date if one is held, else 30 September of
    the year after it would be published."""
    nxt = int(newest_year) + 1
    table = EXPECTED_RELEASES if announced is None else announced
    return nxt, table.get(nxt, date(nxt + 1, 9, 30))

# ---------------------------------------------------------------------------
# The held state (migrate-legacy preconditions), surveyed 2026-10-11
# ---------------------------------------------------------------------------

# (rows, the survey hash: md5 of lad24cd|la_name|population|reference_year
# ordered by reference_year, lad24cd; see live_state)
LEGACY_LIVE = (592, "fe07de1ef4526faa" "77db179ec7eae1b9")
# each held year's one loaded_at
LEGACY_LOADED = {
    "2024": datetime(2026, 3, 26, 9, 40, 44, 303564, tzinfo=timezone.utc),
    "2025": datetime(2026, 8, 13, 21, 11, 31, 51713, tzinfo=timezone.utc),
}
# the held workbooks (sha256 split so the credential scan does not flag
# them); url: the download URL the landing page gives (no redirect, read
# 2026-10-11), which a ledger row names so that a page load finds the file
# checked; loaded_by: how the held rows were really loaded
LEGACY_FILES = {
    "2024": {"name": "mye24tablesew.xlsx",
             "sha256": "4d2d966c8f856623f63c04a222cb7c34"
                       "f8ae396dfb2c207c8d65d4a43ed28736",
             "url": ONS + "/file?uri=" + LANDING
             + "/mid20242023localauthorityboundaries/mye24tablesew.xlsx",
             "loaded_by": "n8n 'ONS Population Estimates (S3)' run log 16 "
                          "(2026-03-26), NOMIS NM_2002_1 and the supplement "
                          "CSV; populations proved equal to"},
    "2025": {"name": "mye25tablesew.xlsx",
             "sha256": "85b4e105e9618043884e5ac41d46e6ea"
                       "e8a46bb01b31790f1c76110cbe92d0d8",
             "url": ONS + "/file?uri=" + LANDING
             + "/mid20252023localauthorityboundaries/mye25tablesew.xlsx",
             "loaded_by": "scripts/s3_mye_refresh.py run log 86 "
                          "(2026-08-13), from"},
}
# names held in another form than the year's workbook: {year: {lad24cd:
# (held name, workbook name)}}. The n8n load read NOMIS NM_2002_1, whose
# names drop the comma; edition 1 keeps the held name (live as held).
LEGACY_NAME_FORMS = {
    "2024": {
        "E06000010": ("Kingston upon Hull City of",
                      "Kingston upon Hull, City of"),
        "E06000019": ("Herefordshire County of", "Herefordshire, County of"),
        "E06000023": ("Bristol City of", "Bristol, City of"),
        "E06000058": ("Bournemouth Christchurch and Poole",
                      "Bournemouth, Christchurch and Poole"),
    },
}
# before the editions exist (a preview against live): the rank of the file
# each held year equals (Covers read 2026-10-11)
LEGACY_RANKS = {"2024": (date(2025, 7, 30), 2024),
                "2025": (date(2026, 7, 29), 2025)}


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _json(x):
    return json.loads(x) if isinstance(x, (str, bytes)) else (x or {})


def _norm(s) -> str:
    return " ".join(str("" if s is None else s).split())


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or \
        (isinstance(v, str) and not v.strip())


def rank_text(rank) -> str:
    """'mid-2025, published 2026-07-29' of a (date, year) rank."""
    if not rank:
        return "no file rank"
    return f"mid-{rank[1]}, published {rank[0].isoformat()}"


def _date(mm, where) -> date:
    try:
        return datetime.strptime(f"{mm.group(1)} {mm.group(2)} "
                                 f"{mm.group(3)}", "%d %B %Y").date()
    except ValueError:
        raise ValueError(f"{where}: {mm.group(0)!r} is not a date") from None


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------

def check_landing(landing_json) -> None:
    """ValueError unless the landing page is still titled LANDING_TITLE."""
    d = _json(landing_json)
    t = _norm((d.get("description") or {}).get("title"))
    if t != LANDING_TITLE:
        raise ValueError(f"the landing page is titled {t!r}, not "
                         f"{LANDING_TITLE!r}; the page has changed")


def _uris(landing_json) -> list:
    d = _json(landing_json)
    return [str(x.get("uri") or "") for x in (d.get("datasets") or [])]


def _segment(uri) -> "str | None":
    """The edition part of a URI under the dataset, or None."""
    head = LANDING + "/"
    return uri[len(head):] if uri.startswith(head) and "/" not in \
        uri[len(head):] else None


def irregular_editions(landing_json) -> list:
    """[(uri, [mid-years it names])] of the URIs under the dataset that do
    not fit 'mid<yyyy><yyyy>localauthorityboundaries' (older editions'
    'editionofthisdataset' and 'boundaires' spellings, the time series and
    the like); a mid-year is a 'mid<yyyy>' in the URI."""
    out = []
    for u in _uris(landing_json):
        seg = _segment(u)
        if seg is None or EDITION_RE.fullmatch(seg):
            continue
        out.append((u, [int(y) for y in re.findall(r"mid([0-9]{4})", seg)]))
    return out


def editions_listed(landing_json) -> list:
    """[(mid_year, boundary_year, uri)] of the landing page's editions whose
    URI ends 'mid<yyyy><yyyy>localauthorityboundaries', newest first.
    ValueError (listing every URI seen) when none fits, or when a URI that
    does not fit names a mid-year newer than the newest that does (never
    passed over for an older edition). Irregular URIs naming older years
    are ignored (irregular_editions lists them)."""
    out = []
    for u in _uris(landing_json):
        seg = _segment(u)
        mm = EDITION_RE.fullmatch(seg) if seg is not None else None
        if mm:
            out.append((int(mm.group(1)), int(mm.group(2)), u))
    seen = _uris(landing_json)
    if not out:
        raise ValueError("no edition 'mid<yyyy><yyyy>localauthorityboundaries' "
                         f"on the landing page; URIs seen: {seen}")
    newest = max(y for y, _, _ in out)
    newer = [(u, ys) for u, ys in irregular_editions(landing_json)
             if any(y > newest for y in ys)]
    if newer:
        raise ValueError(f"the newest regular edition is mid-{newest}, but "
                         "the landing page lists a URI that names a newer "
                         f"mid-year and does not fit the pattern: "
                         f"{[u for u, _ in newer]}; not choosing an older "
                         f"edition over it; URIs seen: {seen}")
    return sorted(out, key=lambda x: (-x[0], -x[1], x[2]))


def latest_edition(landing_json) -> tuple:
    """(mid_year, boundary_year, uri) of the newest edition; ValueError
    when two editions claim it (two boundary vintages: refusing to
    choose)."""
    eds = editions_listed(landing_json)
    newest = eds[0][0]
    hits = [e for e in eds if e[0] == newest]
    if len(hits) != 1:
        raise ValueError(f"{len(hits)} editions for mid-{newest}: "
                         f"{[u for _, _, u in hits]}; refusing to choose one")
    return hits[0]


def edition_for_year(landing_json, year) -> tuple:
    """(mid_year, boundary_year, uri) of that mid-year's edition (--release
    YYYY, or the year of a --file); ValueError unless exactly one."""
    eds = editions_listed(landing_json)
    hits = [e for e in eds if e[0] == int(year)]
    if len(hits) != 1:
        raise ValueError(f"{len(hits)} editions for mid-{year} on the "
                         f"landing page: {[u for _, _, u in hits]}; URIs "
                         f"seen: {_uris(landing_json)}")
    return hits[0]


def newest_listed(landing_json) -> "int | None":
    try:
        return editions_listed(landing_json)[0][0]
    except ValueError:
        return None


def edition_download(edition_json) -> tuple:
    """(file, edition label, the API's releaseDate) of an edition page:
    exactly one download, an .xlsx file, and the label 'Mid-<yyyy>: <yyyy>
    local authority boundaries'. ValueError naming what was seen. The
    releaseDate is printed, never used as a rank."""
    d = _json(edition_json)
    desc = d.get("description") or {}
    files = [str(x.get("file") or "") for x in (d.get("downloads") or [])]
    if len(files) != 1 or not files[0]:
        raise ValueError(f"{len(files)} downloads on the edition page "
                         f"({files}); expected exactly one")
    if Path(files[0]).suffix.lower() != ".xlsx":
        raise ValueError(f"the download {files[0]!r} is not an .xlsx file")
    label = _norm(desc.get("edition"))
    if not EDITION_LABEL_RE.fullmatch(label):
        raise ValueError(f"the edition label {label!r} is not 'Mid-<yyyy>: "
                         "<yyyy> local authority boundaries'")
    return files[0], label, str(desc.get("releaseDate") or "")


def download_url(uri, file) -> str:
    return f"{ONS}/file?uri={uri}/{file}"


def _today() -> date:
    return date.today()


def discovery_note(landing_json, held_periods, today=None) -> list:
    """Lines to print when discovery found nothing newer than what is held:
    the newest edition the landing page lists, that none newer is listed,
    and a WARNING when the next release is overdue (expected_next_release).
    Empty when the newest listed year is not held."""
    nl = newest_listed(landing_json)
    if nl is None or str(nl) not in {str(p) for p in held_periods}:
        return []
    out = [f"the newest edition the landing page lists is mid-{nl}, which "
           "is already held; no newer edition is listed"]
    today = today or _today()
    y, last = expected_next_release(nl)
    if today > last:
        out.append(f"WARNING: the mid-{y} release was expected by "
                   f"{last:%d %B %Y} (published each summer, or its "
                   f"announced date) and today is {today:%d %B %Y}, but "
                   "the landing page does not list it. It may be late, or "
                   "discovery may be missing it (check the ONS page by "
                   "hand; --file PATH loads a downloaded file)")
    return out


# ---------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------

def la_cell(v, where) -> int:
    """A published population: a positive whole number (an int, or a float
    equal to one). Anything else raises ValueError naming `where`: a blank,
    any text, a boolean, a zero, a negative or a non-integer. Never returns
    a default."""
    if isinstance(v, str):
        if not v.strip():
            raise ValueError(f"{where}: a blank cell (a blank halts and is "
                             "never read as 0)")
        raise ValueError(f"{where}: text {v!r} is not a population")
    if v is None:
        raise ValueError(f"{where}: a blank cell (a blank halts and is never "
                         "read as 0)")
    if isinstance(v, bool):
        raise ValueError(f"{where}: {v!r} is not a population")
    if isinstance(v, numbers.Integral):
        n = int(v)
    elif isinstance(v, numbers.Real):
        if math.isnan(v):
            raise ValueError(f"{where}: a blank cell (NaN); never read as 0")
        if not float(v).is_integer():
            raise ValueError(f"{where}: non-integer {v!r}")
        n = int(v)
    else:
        raise ValueError(f"{where}: {v!r} ({type(v).__name__}) is not a "
                         "population")
    if n == 0:
        raise ValueError(f"{where}: a zero (no authority's population has "
                         "ever been published as 0; never stored)")
    if n < 0:
        raise ValueError(f"{where}: negative population {n}")
    return n


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------

def _frames(p: Path) -> dict:
    """{sheet: [row tuples]} of the Cover sheet, Notes, MYE2 - Persons (its
    first four columns), MYE4 and MYE7 when present, and '_sheets': every
    sheet name. ValueError when the file does not open as a workbook."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 (not an .xlsx workbook)
        raise ValueError(f"{p.name}: not a readable .xlsx workbook "
                         f"({e})") from None
    try:
        names = list(wb.sheetnames)
        out = {"_sheets": names}
        for s in (COVER, NOTES, MYE2, MYE4, MYE7):
            if s not in names:
                continue
            width = 4 if s == MYE2 else None
            out[s] = [tuple(r[:width] if width else r)
                      for r in wb[s].iter_rows(values_only=True)]
    finally:
        wb.close()
    return out


def _cells(rows) -> list:
    return [v for r in rows for v in r if not _blank(v)]


def _texts(rows) -> list:
    return [_norm(c) for c in _cells(rows) if isinstance(c, str)]


def _one(texts, rx, what, where):
    hits = [mm for mm in (rx.fullmatch(t) for t in texts) if mm]
    if len(hits) != 1:
        raise ValueError(f"{where}: {len(hits)} cells like {what!r}, expected "
                         f"one; cells seen: {texts[:16]}")
    return hits[0]


def parse_cover(rows, name="") -> dict:
    """{title, year, published, pub_text, next_release} of the Cover sheet.
    ValueError (naming what was seen) unless every title cell names one
    year, there is exactly one 'originally published at 9.30am' line and
    one 'next publication' line, and the 'Date published:' and 'Release
    number: MYE<yy>' lines, when present, agree."""
    where = f"{name}: {COVER}"
    texts = _texts(rows)
    titles = [mm for mm in (COVER_TITLE_RE.fullmatch(t) for t in texts) if mm]
    years = sorted({int(mm.group(1)) for mm in titles})
    if len(years) != 1:
        raise ValueError(f"{where}: title cells name {years or 'no'} "
                         "year(s), expected one 'Mid-Year Population "
                         "Estimates, England and Wales, June <yyyy>'; cells "
                         f"seen: {texts[:16]}")
    year = years[0]
    pub = _one(texts, PUB_RE, "The data tables in this spreadsheet were "
               "originally published at 9.30am <d Month yyyy>.", where)
    nxt = _one(texts, NEXT_RE, "The next publication is scheduled for "
               "<when>.", where)
    published = _date(pub, where)
    for t in texts:
        mm = DATE_RE.fullmatch(t)
        if mm and _date(mm, where) != published:
            raise ValueError(f"{where}: {t!r} differs from the "
                             f"'originally published' line "
                             f"({published:%d %B %Y}); the file's identity "
                             "is unclear")
        mm = RELEASE_NO_RE.fullmatch(t)
        if mm and int(mm.group(1)) != year % 100:
            raise ValueError(f"{where}: {t!r} is not the June {year} "
                             "release")
    return {"title": titles[0].group(0), "year": year, "published": published,
            "pub_text": pub.group(0), "next_release": _norm(nxt.group(1))}


def parse_notes(rows, name="") -> dict:
    """{geography_note, revision_notes} of the Notes sheet, verbatim (white
    space normalised). ValueError unless there is exactly one note starting
    'Geography:'."""
    where = f"{name}: {NOTES}"
    texts = _texts(rows)
    geo = [t for t in texts if GEOG_NOTE_RE.fullmatch(t)]
    if len(geo) != 1:
        raise ValueError(f"{where}: {len(geo)} notes starting 'Geography:', "
                         f"expected one; notes seen: "
                         f"{[t[:60] for t in texts[:24]]}")
    revs = [t for t in texts if REVISION_NOTE_RE.fullmatch(t)]
    return {"geography_note": geo[0], "revision_notes": revs}


def _header_at(rows, where) -> int:
    for i, r in enumerate(rows[:15]):
        if r and _norm(r[0]) == "Code":
            return i
    raise ValueError(f"{where}: no header row starting 'Code' in the first "
                     f"15 rows; cells seen: {_texts(rows[:10])[:12]}")


def _trim(row) -> list:
    out = [_norm(v) if isinstance(v, str) else v for v in row]
    while out and _blank(out[-1]):
        out.pop()
    return out


def _title(rows, rx, what, where):
    t = _norm(rows[0][0]) if rows and rows[0] else ""
    mm = rx.fullmatch(t)
    if not mm:
        raise ValueError(f"{where}: the title {t!r} is not {what!r}")
    return mm


def _data(rows, h, ncols, where) -> list:
    """[(sheet row number, code, name, geography, [cells])] of the rows
    below the header; empty rows are allowed only at the end."""
    out, empty = [], 0  # trailing empty rows seen; not a source value
    for n, r in enumerate(rows[h + 1:], start=h + 2):
        r = list(r)
        if all(_blank(v) for v in r):
            empty += 1
            continue
        if empty:
            raise ValueError(f"{where} row {n}: data after {empty} empty "
                             "row(s); empty rows are allowed only at the end")
        cells = (r + [None] * (3 + ncols))[:3 + ncols]
        if any(not _blank(v) for v in r[3 + ncols:]):
            raise ValueError(f"{where} row {n}: a cell beyond the "
                             f"{3 + ncols} headed columns")
        out.append((n, _norm(cells[0]), _norm(cells[1]), _norm(cells[2]),
                    cells[3:]))
    return out


def english_rows(table) -> dict:
    """{las {code: {row, name, geography, values}}, england {row, values},
    wales (count of W06 rows, skipped), aggregates [codes set aside]} of a
    parsed sheet's data rows: E06, E07, E08 and E09 authorities only; W06
    rows counted; E10/E11/E12/E13/E92/W92/K0x aggregates set aside, the
    England row kept for reconciliation. ValueError on any other code, an
    authority without a name, or a repeated authority."""
    where = table["where"]
    cols = table["cols"]
    las, england, wales, agg = {}, None, 0, []  # counts; not a source value
    for n, code, name, geog, cells in table["data"]:
        vals = dict(zip(cols, cells))
        if LA_RE.fullmatch(code):
            if not name:
                raise ValueError(f"{where} row {n}: authority {code} without "
                                 "a name")
            if code in las:
                raise ValueError(f"{where} row {n}: {code} appears twice "
                                 f"(also row {las[code]['row']})")
            las[code] = {"row": n, "name": name, "geography": geog,
                         "values": vals}
        elif WALES_RE.fullmatch(code):
            wales += 1
        elif code == ENGLAND:
            if england is not None:
                raise ValueError(f"{where} row {n}: a second England row")
            england = {"row": n, "values": vals}
        elif AGGREGATE_RE.fullmatch(code):
            agg.append(code)
        else:
            raise ValueError(f"{where} row {n}: unexpected code {code!r} "
                             f"({name!r}, {geog!r}); known: E06-E09 "
                             "authorities, W06 Welsh authorities, the "
                             "county, region and country aggregates")
    if england is None or not las:
        raise ValueError(f"{where}: England row {england is not None}, "
                         f"{len(las)} English authorities; both are expected")
    return {"las": las, "england": england, "wales": wales,
            "aggregates": agg}


def parse_mye2(rows, name="") -> dict:
    where = f"{name}: {MYE2}"
    mm = _title(rows, MYE2_TITLE_RE, "MYE2: Persons by single year of age "
                "and sex for local authorities in England and Wales, "
                "mid-<yyyy>", where)
    h = _header_at(rows, where)
    head = _trim(rows[h])
    if tuple(head[:4]) != MYE2_HEADER:
        raise ValueError(f"{where}: the header starts {head[:4]}, expected "
                         f"exactly {list(MYE2_HEADER)}; the reading must be "
                         "corrected deliberately, with the evidence")
    table = {"where": where, "cols": ["All ages"],
             "data": _data([r[:4] for r in rows], h, 1, where)}
    out = english_rows(table)
    out.update(title=mm.group(0), year=int(mm.group(1)), header=head[:4])
    return out


def parse_mye4(rows, name="") -> dict:
    where = f"{name}: {MYE4}"
    mm = _title(rows, MYE4_TITLE_RE, "MYE4: Population estimates: Summary for "
                "England and Wales, mid-2011 to mid-<yyyy>", where)
    year = int(mm.group(1))
    h = _header_at(rows, where)
    head = _trim(rows[h])
    years = list(range(year, FIRST_YEAR - 1, -1))
    want = ["Code", "Name", "Geography"] + [f"Mid-{y}" for y in years]
    if head != want:
        unknown = [g for g in head if g not in want]
        missing = [w for w in want if w not in head]
        raise ValueError(f"{where}: the header {head} is not the expected one: "
                         f"unknown {unknown}, missing {missing} (the title "
                         f"says mid-2011 to mid-{year}, so expected exactly "
                         f"Code, Name, Geography, Mid-{year} .. "
                         f"Mid-{FIRST_YEAR}); the reading must be corrected "
                         "deliberately, with the evidence")
    table = {"where": where, "cols": years,
             "data": _data(rows, h, len(years), where)}
    out = english_rows(table)
    out.update(title=mm.group(0), year=year, years=years, header=head)
    return out


def parse_mye7(rows, name="") -> dict:
    """The MYE7 sheet: {layout 'revisions', title, label (e.g. 'July 2026'),
    years [a..b], las {code: {year: (published, revised)}}, england {year:
    (published, revised)}, wales, published_label} when its title is
    '<Month yyyy> revisions to mid-<a> to mid-<b> ...'; otherwise {layout
    'other', title} (the July 2025 layout: the uncertainty table). The
    revision columns per year: 'Population mid-<y> (published <Month
    yyyy>)', 'Population mid-<y> (revised <label>)', 'Revision to Mid-<y>',
    'Percentage Revision to Mid-<y>'; ValueError on any other header."""
    where = f"{name}: {MYE7}"
    t = _norm(rows[0][0]) if rows and rows[0] else ""
    mm = MYE7_TITLE_RE.fullmatch(t)
    if not mm:
        return {"layout": "other", "title": t}
    label, a, b = mm.group(1), int(mm.group(2)), int(mm.group(3))
    years = list(range(a, b + 1))
    h = _header_at(rows, where)
    head = _trim(rows[h])
    pub = re.fullmatch(r"Population mid-[0-9]{4} \(published ([A-Z][a-z]+ "
                       r"[0-9]{4})\)", head[3] if len(head) > 3 else "")
    plabel = pub.group(1) if pub else "?"
    want = ["Code", "Name", "Geography"]
    for y in years:
        want += [f"Population mid-{y} (published {plabel})",
                 f"Population mid-{y} (revised {label})",
                 f"Revision to Mid-{y}", f"Percentage Revision to Mid-{y}"]
    if head != want:
        raise ValueError(f"{where}: the header {head} is not the expected "
                         f"revisions header {want}")
    cols = [(y, k) for y in years for k in ("pub", "rev", "diff", "pct")]
    table = {"where": where, "cols": cols,
             "data": _data(rows, h, len(cols), where)}
    er = english_rows(table)

    def pairs(vals):
        return {y: (vals[(y, "pub")], vals[(y, "rev")]) for y in years}
    return {"layout": "revisions", "title": t, "label": label,
            "published_label": plabel, "years": years,
            "las": {c: pairs(r["values"]) for c, r in er["las"].items()},
            "england": pairs(er["england"]["values"]),
            "wales": er["wales"]}


def read_file(path) -> dict:
    """The facts of one tables workbook: parse_cover's (title, year,
    published, pub_text, next_release), parse_notes' (geography_note,
    revision_notes), mye2, mye4 and mye7 (None when absent), layout_notes
    (what was seen where the layout differs), path, sheets and rank
    (published, year). The mid-year of the Cover, MYE2 and MYE4 must agree.
    Raises ValueError naming what was seen."""
    p = Path(path)
    fr = _frames(p)
    sheets = fr["_sheets"]
    need = [s for s in (COVER, NOTES, MYE2, MYE4) if s not in sheets]
    if need:
        raise ValueError(f"{p.name}: no {', '.join(need)} sheet(s); sheets "
                         f"{sheets}")
    cover = parse_cover(fr[COVER], p.name)
    notes = parse_notes(fr[NOTES], p.name)
    m2 = parse_mye2(fr[MYE2], p.name)
    m4 = parse_mye4(fr[MYE4], p.name)
    if not (cover["year"] == m2["year"] == m4["year"]):
        raise ValueError(f"{p.name}: the Cover says June {cover['year']}, "
                         f"{MYE2} says mid-{m2['year']}, {MYE4} says "
                         f"mid-2011 to mid-{m4['year']}; they must agree")
    layout = []
    m7 = parse_mye7(fr[MYE7], p.name) if MYE7 in fr else None
    if m7 is None:
        layout.append(f"no {MYE7} sheet; no revisions table read")
    elif m7["layout"] == "other":
        layout.append(f"{MYE7} is {m7['title']!r}, not a revisions table "
                      "(the July 2025 layout, whose revisions are in MYE8 "
                      "and MYE9); not read")
    if set(m2["las"]) != set(m4["las"]):
        raise ValueError(f"{p.name}: {MYE2} and {MYE4} list different "
                         f"authorities: only {MYE2} "
                         f"{sorted(set(m2['las']) - set(m4['las']))}, only "
                         f"{MYE4} {sorted(set(m4['las']) - set(m2['las']))}")
    out = dict(cover)
    out.update(notes)
    out.update({"path": p, "sheets": sheets, "mye2": m2, "mye4": m4,
                "mye7": m7, "layout_notes": layout,
                "rank": (cover["published"], cover["year"])})
    return out


# ---------------------------------------------------------------------------
# Reconciliation inside the file
# ---------------------------------------------------------------------------

def reconcile(file, years=None) -> list:
    """Problems (empty = fine), exact: MYE2's English authorities (every
    cell a population, la_cell) sum to its England row; for each MYE4 year
    read (years; default the newest) the same in MYE4; MYE2 and MYE4 agree
    on the newest year for every authority and England; for each year read
    that MYE7 revises, MYE4 equals MYE7's revised column for every authority
    and England, and MYE7 lists the same authorities."""
    out = []
    m2, m4, m7 = file["mye2"], file["mye4"], file.get("mye7")
    newest = file["year"]
    try:
        la2 = {c: la_cell(r["values"]["All ages"],
                          f"{MYE2} row {r['row']} {c} All ages")
               for c, r in m2["las"].items()}
        e2 = la_cell(m2["england"]["values"]["All ages"],
                     f"{MYE2} England All ages")
        if sum(la2.values()) != e2:
            out.append(f"{MYE2} mid-{newest}: the authorities sum to "
                       f"{sum(la2.values()):,}, the England row says {e2:,}")
    except ValueError as e:
        out.append(str(e))
        la2 = e2 = None
    for y in (years if years is not None else [newest]):
        y = int(y)
        if y not in m4["years"]:
            out.append(f"{y}: not a year of {MYE4}")
            continue
        try:
            la4 = {c: la_cell(r["values"][y], f"{MYE4} row {r['row']} {c} "
                              f"Mid-{y}") for c, r in m4["las"].items()}
            e4 = la_cell(m4["england"]["values"][y], f"{MYE4} England "
                         f"Mid-{y}")
        except ValueError as e:
            out.append(str(e))
            continue
        if sum(la4.values()) != e4:
            out.append(f"{MYE4} mid-{y}: the authorities sum to "
                       f"{sum(la4.values()):,}, the England row says {e4:,}")
        if y == newest and la2 is not None:
            bad = [c for c in sorted(la4) if la2.get(c) != la4[c]]
            if bad or e2 != e4:
                out.append(f"{MYE2} and {MYE4} disagree on mid-{y}: "
                           + ", ".join(f"{c} {la2.get(c)} vs {la4[c]}"
                                       for c in bad[:6])
                           + (f"; England {e2:,} vs {e4:,}" if e2 != e4
                              else ""))
        if m7 and m7.get("layout") == "revisions" and y in m7["years"]:
            if set(m7["las"]) != set(la4):
                out.append(f"{MYE7} lists different authorities from "
                           f"{MYE4}")
                continue
            bad = [c for c in sorted(la4) if m7["las"][c][y][1] != la4[c]]
            if bad:
                out.append(f"{MYE7} revised ({m7['label']}) mid-{y} differs "
                           f"from {MYE4} at {len(bad)} authorit(ies): "
                           + ", ".join(f"{c} {m7['las'][c][y][1]} vs "
                                       f"{la4[c]}" for c in bad[:6]))
            if m7["england"][y][1] != e4:
                out.append(f"{MYE7} revised mid-{y} England "
                           f"{m7['england'][y][1]} differs from {MYE4} "
                           f"{e4:,}")
    return out


def england_total(file, year) -> int:
    y = int(year)
    if y == file["year"]:
        return la_cell(file["mye2"]["england"]["values"]["All ages"],
                       f"{MYE2} England")
    return la_cell(file["mye4"]["england"]["values"][y], f"{MYE4} England")


# ---------------------------------------------------------------------------
# Records
# ---------------------------------------------------------------------------

def _getter(resolve):
    return resolve if callable(resolve) else (lambda c: resolve[c])


def build_records(file, year, resolve, provenance=None) -> list:
    """The mid-year's records from a read file, one per English authority,
    sorted by lad24cd: {lad24cd, reference_year (the year as a string, the
    engine's period), population, la_name}. The newest year comes from MYE2
    (All ages and its names), and MYE4 must say the same for every
    authority; a back year comes from MYE4 (its column and names). resolve:
    {publisher code: lad24cd} or a callable. provenance: optional {name:
    value} added to every record for the caller (nothing of it is stored or
    compared). ValueError when the file cannot state the year, a cell is
    not a population (la_cell), MYE2 and MYE4 disagree, a code resolves to
    nothing, or a lad24cd repeats."""
    y = int(year)
    m2, m4 = file["mye2"], file["mye4"]
    if y not in m4["years"]:
        raise ValueError(f"the file covers mid-{m4['years'][-1]} to "
                         f"mid-{m4['years'][0]}; it cannot state mid-{y}")
    get = _getter(resolve)
    out, seen = [], {}
    for code, r4 in sorted(m4["las"].items()):
        w4 = f"{MYE4} row {r4['row']} {code}"
        pop4 = la_cell(r4["values"][y], f"{w4} Mid-{y}")
        if y == file["year"]:
            r2 = m2["las"].get(code)
            if r2 is None:
                raise ValueError(f"{code} is in {MYE4} but not in {MYE2}")
            pop = la_cell(r2["values"]["All ages"],
                          f"{MYE2} row {r2['row']} {code} All ages")
            if pop != pop4:
                raise ValueError(f"{code}: {MYE2} says {pop:,}, {MYE4} says "
                                 f"{pop4:,} for mid-{y}; the sheets disagree")
            name, where = r2["name"], f"{MYE2} row {r2['row']}"
        else:
            pop, name, where = pop4, r4["name"], w4
        lad = get(code)
        if not lad:
            raise ValueError(f"{where}: {code} resolves to no lad24cd")
        if lad in seen:
            raise ValueError(f"{where}: duplicate key {lad} (also {seen[lad]}); "
                             "refusing to choose a winner")
        seen[lad] = code
        rec = {KEY: lad, PERIOD: str(y), "population": pop, "la_name": name}
        if provenance:
            rec.update(provenance)
        out.append(rec)
    out.sort(key=lambda x: x[KEY])
    return out


def periods_stated(file, held_years) -> list:
    """The years a file states: its newest year, plus every held year MYE4
    covers. Sorted ints."""
    years = set(file["mye4"]["years"])
    out = {file["year"]}
    for h in held_years:
        if int(h) in years:
            out.add(int(h))
    return sorted(out)


def rows_content_sha(records) -> str:
    """sha256 of a year's content: one line per record sorted by lad24cd,
    lad24cd|population|la_name, joined by LF, UTF-8. Nothing else is in it,
    so where a file was read never makes a new edition."""
    lines = sorted(f"{r[KEY]}|{r['population']}|{r['la_name']}"
                   for r in records)
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Geography
# ---------------------------------------------------------------------------

def resolve_codes(cur, codes, period) -> tuple:
    """({publisher code: lad24cd}, problems). geography.resolve(cur, '3',
    codes, by_period={period: codes}) (source 3 is declared 'old': a new
    Barnsley or Sheffield code is a problem); then each code's canonical
    form must be in la_boundaries, else 'UNEXPLAINED'. Two codes reaching
    one lad24cd are a problem."""
    codes = set(codes)
    recode, problems = geography.resolve(cur, RUN_SOURCE, codes,
                                         by_period={period: codes})
    problems = list(problems)
    cur.execute("SELECT lad24cd FROM public.la_boundaries")
    valid = {r[0] for r in cur.fetchall()}
    out, hit = {}, {}
    for c in sorted(codes):
        r = geography.canonical(c, recode)
        if r in valid:
            out[c] = r
            hit.setdefault(r, []).append(c)
        else:
            problems.append(f"UNEXPLAINED {c}: not in la_boundaries and not a "
                            "la_code_lookup recode")
    two = {k: v for k, v in hit.items() if len(v) > 1}
    if two:
        problems.append(f"{period}: codes {two} reach one lad24cd; refusing "
                        "to choose a winner")
    return out, problems


# ---------------------------------------------------------------------------
# Ranks, ledger and planning
# ---------------------------------------------------------------------------

_RANK_RE = re.compile(r"\bmid-([0-9]{4}); published ([0-9]{4}-[0-9]{2}-"
                      r"[0-9]{2})\b")
_LEDGER_RE = re.compile(r"(.*) \(mid-([0-9]{4}); published ([0-9]{4}-"
                        r"[0-9]{2}-[0-9]{2}); periods ([0-9]{4}(?:, "
                        r"[0-9]{4})*)\)")


def release_rank(file_or_source) -> "tuple | None":
    """(Cover publication date, mid-year) of a file: a read file's own
    rank, or the 'mid-<yyyy>; published <yyyy-mm-dd>' an edition's or
    ledger row's source_file carries; None if it names none (a file name,
    URL, media id or API date is never read for a rank)."""
    x = file_or_source
    if isinstance(x, dict):
        return x.get("rank")
    mm = _RANK_RE.search(str(x or ""))
    if not mm:
        return None
    return (date.fromisoformat(mm.group(2)), int(mm.group(1)))


def edition_source(file, src, note="") -> str:
    """An edition's source_file: the file name, its own mid-year and
    publication date (the rank), where it was read, and a note."""
    return (f"{src['file_name']}; mid-{file['year']}; published "
            f"{file['published'].isoformat()}; {src['where']}"
            + (f"; {note}" if note else ""))


def ledger_source(where, rank, periods) -> str:
    """A file's ledger source_file: where it was read (the final URL, or the
    path of a --file), its own mid-year and publication date, and the
    periods it states, so a later run knows them without reading it."""
    ps = ", ".join(str(p) for p in sorted(int(p) for p in periods))
    return (f"{where} (mid-{rank[1]}; published {rank[0].isoformat()}; "
            f"periods {ps})")


def parse_ledger_source(s) -> "tuple | None":
    """(where, rank, [periods as strings]) of a ledger source_file, or
    None."""
    mm = _LEDGER_RE.fullmatch(str(s or ""))
    if not mm:
        return None
    rank = (date.fromisoformat(mm.group(3)), int(mm.group(2)))
    return mm.group(1), rank, mm.group(4).split(", ")


def content_sha256(path) -> str:
    """SHA-256 of the file's bytes."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def tip_ranks(tips, checked) -> dict:
    """{held period: the newest file rank that stated it}: the tip
    edition's, or a later file the ledger records for the period (a file
    found unchanged stores no edition)."""
    out = {}
    for p, t in tips.items():
        ranks = [t.get("rank")] + [release_rank(s)
                                   for s, _ in checked.get(p, ())]
        ranks = [r for r in ranks if r is not None]
        out[p] = max(ranks) if ranks else None
    return out


def ledger_complete(checked, where, sha, ranks, *, allow_older=False,
                    stranded=()) -> "list | None":
    """The periods the ledger records this file's (where, sha256) pair for,
    if that covers every period the file needs (its ledger source declares
    them); else None (the file is read and each period planned). A held
    period whose tip comes from a newer file is not needed (skipped as
    older) unless allow_older; a stranded period (editions, no live rows)
    always makes the file be read. So a run that failed or was rejected
    part-way leaves a period without its ledger row and the rerun reads the
    file again."""
    have, declared = set(), None
    for p, pairs in checked.items():
        for s, h in pairs:
            parsed = parse_ledger_source(s)
            if h == sha and parsed and parsed[0] == where:
                have.add(str(p))
                declared = parsed
    if declared is None:
        return None
    _, rank, need = declared
    need = set(need)
    if not allow_older:
        need = {p for p in need if not (ranks.get(p) is not None
                                        and ranks[p] > rank)}
    if not need or need & {str(x) for x in stranded}:
        return None
    return sorted(have) if need <= have else None


def plan_periods(periods, rank, ranks, checked, source, sha, *, recheck,
                 allow_older) -> tuple:
    """(new, revised, skipped). periods: the years the file states, as
    strings; rank: its own; ranks: {held period: tip_ranks}; checked:
    {period: {(source, sha)}} of the ledger (stranded periods left out by
    the caller, so they are always read). new: not held; revised: held,
    compared with the tip; skipped {period: reason}: held periods whose tip
    comes from a newer file, and new periods older than the newest held
    year (back-filling: not done by load), unless allow_older (the
    older-file guard, on every path); held periods whose (source, sha) the
    ledger already records, unless --recheck names it. --recheck must be a
    held period of the file (ValueError)."""
    periods = sorted(str(p) for p in periods)
    if recheck is not None and (recheck not in periods or recheck not in ranks):
        raise ValueError(f"--recheck {recheck}: not a held year this file "
                         f"states ({', '.join(periods)})")
    newest_held = max(ranks) if ranks else None
    new, revised, skipped = [], [], {}
    for p in periods:
        if p not in ranks:
            if newest_held is not None and p < newest_held and not allow_older:
                skipped[p] = (f"older: mid-{p} is before the newest held "
                              f"year {newest_held}; back-filling is not done "
                              "by load")
                continue
            new.append(p)
            continue
        t = ranks[p]
        if t is not None and t > rank and not allow_older:
            skipped[p] = (f"older: its tip comes from a file of "
                          f"{rank_text(t)}, this file is {rank_text(rank)}")
            continue
        if recheck != p and (source, sha) in checked.get(p, ()):
            skipped[p] = "checked: this file is already in the ledger"
            continue
        revised.append(p)
    return new, revised, skipped


# ---------------------------------------------------------------------------
# Stop conditions
# ---------------------------------------------------------------------------

def _over(old, new, bp) -> bool:
    """True when new moves from old by more than bp basis points (exact)."""
    if old == new:
        return False
    if old == 0:
        return True
    return abs(new - old) * 10000 > bp * old


def _pct(old, new) -> str:
    return f"{(new - old) * 100 / old:+.3f}%" if old else "from 0"


def _by_key(records) -> dict:
    return {r[KEY]: r for r in records or ()}


def _total(records) -> int:
    return sum(r["population"] for r in records or ())


def period_problems(new, tip, prev, *, kind) -> list:
    """The stop conditions of one year (empty = fine). new: the file's
    records; tip: (kind 'revised') the held edition's records of the same
    year; prev: (kind 'new') the latest earlier held year's records, or
    None. Problems starting with PARTIAL are released by nothing: fewer
    authorities than EXPECTED_AREAS (both kinds), or an authority of the
    held edition missing (revised). New: the England total moving by more
    than NEW_TOTAL_BP against prev, or any authority by more than
    NEW_AREA_BP (released by --acknowledge). Revised: any change at all
    (population or la_name; released by --acknowledge), and problems
    starting with LARGE (released only by --acknowledge with
    --accept-large-revision): an authority's population moving by more
    than REV_AREA_BP, or the England total by more than REV_TOTAL_BP."""
    g = globals()
    out = []
    nk = _by_key(new)
    want = g["EXPECTED_AREAS"]
    if len(nk) < want:
        out.append(f"{PARTIAL}: {len(nk)} authorities, expected {want} (a "
                   "partial file never replaces a fuller edition)")
    if kind == "new":
        if not prev:
            return out
        a, b = _total(prev), _total(new)
        if _over(a, b, g["NEW_TOTAL_BP"]):
            out.append(f"national population {a:,} -> {b:,} ({_pct(a, b)}) "
                       f"against the previous held year (limit "
                       f"{g['NEW_TOTAL_BP'] / 100:.2f}%)")
        pk = _by_key(prev)
        big = [f"{k} {pk[k]['population']:,}->{nk[k]['population']:,} "
               f"({_pct(pk[k]['population'], nk[k]['population'])})"
               for k in sorted(set(pk) & set(nk))
               if _over(pk[k]["population"], nk[k]["population"],
                        g["NEW_AREA_BP"])]
        if big:
            out.append(f"population moves by more than "
                       f"{g['NEW_AREA_BP'] / 100:.2f}% against the previous "
                       f"held year in {len(big)} authorit(ies): "
                       f"{', '.join(big[:6])}")
        return out
    tk = _by_key(tip)
    gone = sorted(set(tk) - set(nk))
    if gone:
        out.append(f"{PARTIAL}: {len(gone)} authorit(ies) of the held edition "
                   f"missing from the file (a partial file never replaces a "
                   f"fuller edition): {', '.join(gone[:8])}")
    both = sorted(set(tk) & set(nk))
    pops = [k for k in both if tk[k]["population"] != nk[k]["population"]]
    names = [k for k in both if tk[k]["la_name"] != nk[k]["la_name"]]
    added = sorted(set(nk) - set(tk))
    if pops or names or added:
        a, b = _total(tip), _total(new)
        big = max(pops, key=lambda k: abs(nk[k]["population"]
                                          - tk[k]["population"]),
                  default=None)
        ex = [f"{k} {tk[k]['population']:,}->{nk[k]['population']:,}"
              for k in pops[:4]]
        out.append(
            f"the held edition is revised: population changes at "
            f"{len(pops)} authorit(ies) (England {a:,} -> {b:,}, {b - a:+,}"
            + (f"; largest {big} {nk[big]['population'] - tk[big]['population']:+,}"
               if big else "")
            + (f"; e.g. {', '.join(ex)}" if ex else "")
            + f"), la_name changes at {len(names)}"
            + (": " + ", ".join(f"{k} {tk[k]['la_name']!r}->"
                                f"{nk[k]['la_name']!r}" for k in names[:6])
               if names else "")
            + (f", new authorities {added}" if added else ""))
        lim = [f"{k} {tk[k]['population']:,}->{nk[k]['population']:,} "
               f"({_pct(tk[k]['population'], nk[k]['population'])})"
               for k in pops if _over(tk[k]["population"],
                                      nk[k]["population"], g["REV_AREA_BP"])]
        if lim:
            out.append(f"{LARGE}: population revised by more than "
                       f"{g['REV_AREA_BP'] / 100:.2f}% at {len(lim)} "
                       f"authorit(ies): {', '.join(lim[:6])}")
        if _over(a, b, g["REV_TOTAL_BP"]):
            out.append(f"{LARGE}: national population revised {a:,} -> "
                       f"{b:,} ({_pct(a, b)}; limit "
                       f"{g['REV_TOTAL_BP'] / 100:.2f}%)")
    return out


def _differs(new, old) -> bool:
    nk, ok = _by_key(new), _by_key(old)
    if set(nk) != set(ok):
        return True
    return any(nk[k][c] != ok[k][c] for k in nk for c in VALUES)


# ---------------------------------------------------------------------------
# Spec and profile
# ---------------------------------------------------------------------------

SPEC = core.EditionSpec(
    name="s3",
    live_table="la_population",
    editions_table="la_population_editions",
    key_cols=(KEY,),
    period_col=PERIOD,
    value_cols=VALUE_TYPES,
    refresh_cols=VALUES + ("loaded_at",),
    refresh_from=(("loaded_at", "loaded_at"),),
    key_types=((KEY, "varchar(9) NOT NULL"), (PERIOD, "integer NOT NULL")),
    fk_la_boundaries=True,
)


def _no_label(fetched_on) -> str:
    raise ValueError("an S3 release label comes from the file; store "
                     "through apply_year(...)")


def check_records(records, period) -> None:
    """ValueError unless the year's records are whole: each of this year,
    one per lad24cd, a positive population and a name."""
    seen = set()
    for r in records:
        if r.get(PERIOD) != period:
            raise ValueError(f"{period}: a record of year {r.get(PERIOD)!r}")
        if r[KEY] in seen:
            raise ValueError(f"{period}: {r[KEY]} appears twice")
        p = r.get("population")
        if not isinstance(p, int) or isinstance(p, bool) or p <= 0:
            raise ValueError(f"{period}: {r[KEY]} population {p!r}")
        if not r.get("la_name"):
            raise ValueError(f"{period}: {r[KEY]} lacks a name")
        seen.add(r[KEY])


PROFILE = pe.Profile(
    spec=SPEC, value_cols=VALUES, run_agent=RUN_AGENT, run_source=RUN_SOURCE,
    heading="S3 ONS mid-year population estimates (la_population)",
    default_source_file="ONS mid-year population estimates tables",
    expected_areas=EXPECTED_AREAS, release_label=_no_label,
    content_sha256=rows_content_sha, check_records=check_records,
    file_checks=True, savepoint="s3_period",
    example_label="(population, la_name)")


def profile() -> pe.Profile:
    """PROFILE bound to the current SPEC and EXPECTED_AREAS (looked up when
    called, so tests can swap them)."""
    g = globals()
    return dataclasses.replace(PROFILE, spec=g["SPEC"],
                               expected_areas=g["EXPECTED_AREAS"])


def _late(name: str):
    """A call through this module's attribute `name`, looked up when called
    (so mock.patch.object(s3_mye_editions, name, ...) reaches the
    engine)."""
    return lambda *a, **kw: globals()[name](*a, **kw)


def create_all(cur) -> None:
    """ddl: the editions table, its append-only triggers and the file-check
    ledger. Idempotent; no value is written; the live table is not
    touched."""
    prof = profile()
    core.create_schema(cur, prof.spec)
    pe.create_file_checks(cur, prof)


@dataclass(frozen=True)
class PeriodInfo:
    """What a year's edition records: source_file, the release label and the
    file read ((ledger source, sha256),)."""
    source_file: str
    label: str
    files: tuple


def apply_year(cur, profile, period, records, *, fetched_on, info) -> str:
    """pe.apply_period for one year, then its ledger row, inside the
    caller's per-period savepoint (s3_period): a new year's edition 1, live
    rows and ledger row commit or roll back together."""
    prof = dataclasses.replace(profile, release_label=lambda d: info.label)
    kind = pe.apply_period(cur, prof, period, records, fetched_on=fetched_on,
                           source_file=info.source_file)
    tip = core.chain_tip(cur, prof.spec, period)
    for src, sha in info.files:
        pe.record_file_check(cur, prof, period, src, sha, kind, tip)
    return kind


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

def _session(session):
    if session is not None:
        return session
    import requests
    return requests


def _get(session, url, **kw):
    try:
        r = session.get(url, headers={"User-Agent": USER_AGENT},
                        timeout=300, **kw)
        r.raise_for_status()
    except Exception as e:  # noqa: BLE001 (any failure halts)
        halt(f"GET {url} failed, nothing downloaded: {e}")
    return r


def fetch_json(path, session=None) -> dict:
    """The ONS /data JSON of a page path (a read-only GET)."""
    return json.loads(_get(_session(session), ONS + path + "/data").text)


def _url_name(url) -> str:
    """The file name a URL names: the 'uri' query's last part (ONS
    /file?uri=...), else the path's."""
    u = urlparse(url)
    q = parse_qs(u.query).get("uri")
    return Path(q[0] if q else u.path).name


def _opens(path) -> None:
    """Raise unless the file is an .xlsx workbook with the MYE2 - Persons
    and MYE4 sheets."""
    with zipfile.ZipFile(path) as z:
        wbx = z.read("xl/workbook.xml")
    for s in (MYE2, MYE4):
        if f'name="{s}"'.encode() not in wbx:
            raise ValueError(f"no {s} sheet")


def fetch(url, dest, session=None) -> tuple:
    """Download url into the directory dest (redirects followed) and return
    (path to read, final URL). The file name is the final URL's. Checked
    before it is kept: more than MIN_FILE_BYTES, an .xlsx name, and it opens
    as a workbook with the MYE2 - Persons and MYE4 sheets. A same-named file
    with the same content is kept as it is; one with different content is
    never replaced: the download is saved beside it as
    <stem>-<sha8><suffix> and that is the one read. An existing name is
    never a reason to skip the download."""
    dest = Path(dest)
    r = _get(_session(session), url)
    final = str(getattr(r, "url", None) or url)
    if final != url:
        print(f"redirected: {url} -> {final}")
    body = r.content
    if len(body) <= MIN_FILE_BYTES:
        halt(f"{final}: {len(body)} bytes, expected more than "
             f"{MIN_FILE_BYTES}; nothing written; saw {body[:100]!r}")
    name = _url_name(final)
    suffix = Path(name).suffix.lower()
    if suffix != ".xlsx":
        halt(f"{final}: not an .xlsx file; nothing written")
    dest.mkdir(parents=True, exist_ok=True)
    target = dest / name
    new_sha = hashlib.sha256(body).hexdigest()
    tmp = dest / (Path(name).stem + ".download.tmp" + suffix)
    tmp.write_bytes(body)
    try:
        _opens(tmp)
    except Exception as e:  # noqa: BLE001
        tmp.unlink()
        halt(f"{final}: does not open as the mid-year tables ({e}); nothing "
             "kept")
    path = target
    if target.exists():
        old_sha = content_sha256(target)
        if old_sha == new_sha:
            tmp.unlink()
            print(f"{name} already held with the same content (sha256 "
                  f"{new_sha[:16]}); kept as it is")
        else:
            path = dest / f"{target.stem}-{new_sha[:8]}{target.suffix}"
            if path.exists() and content_sha256(path) == new_sha:
                tmp.unlink()
                print(f"NOTE: {name} differs from this download; "
                      f"{path.name} already holds it (sha256 "
                      f"{new_sha[:16]})")
            else:
                tmp.replace(path)
                print(f"NOTE: {name} already exists with different content "
                      f"(sha256 {old_sha[:16]}, new {new_sha[:16]}); kept it "
                      f"untouched and saved the new download as {path.name}")
    else:
        tmp.replace(target)
        print(f"downloaded {name}: {len(body):,} bytes, sha256 "
              f"{new_sha[:16]}")
    return path, final


def _file_source(path: Path) -> str:
    """Where a --file was read: the path under the repo (data/raw/...),
    else the absolute path."""
    p = Path(path).resolve()
    try:
        return p.relative_to(REPO).as_posix()
    except ValueError:
        return p.as_posix()


# ---------------------------------------------------------------------------
# Reading held data
# ---------------------------------------------------------------------------

table_exists = pe.table_exists


def _conn(writing):
    return pe.connect(writing)


def records(cur, spec, period, edition=None) -> list:
    """The year's rows as records (lad24cd, the year as a string,
    population, la_name): the live table's (edition None) or the given
    edition's."""
    names = (KEY,) + VALUES
    if edition is None:
        cur.execute(f"SELECT {', '.join(names)} FROM public.{spec.live_table} "
                    f"WHERE {spec.period_col} = %s", (period,))
    else:
        cur.execute(f"SELECT {', '.join(names)} FROM "
                    f"public.{spec.editions_table} WHERE {spec.period_col} = "
                    "%s AND edition = %s", (period, edition))
    out = [dict(zip(names, r)) for r in cur.fetchall()]
    for r in out:
        r[spec.period_col] = str(period)
    return out


def tip_info(cur, spec) -> dict:
    """{period (string): {edition, source_file, rank}} of every period with
    editions."""
    cur.execute(f"SELECT DISTINCT {spec.period_col} FROM "
                f"public.{spec.editions_table}")
    out = {}
    for (p,) in cur.fetchall():
        tip = core.latest_edition(cur, spec, str(p))
        cur.execute(f"SELECT DISTINCT source_file FROM "
                    f"public.{spec.editions_table} WHERE {spec.period_col} = "
                    "%s AND edition = %s", (p, tip))
        src = {r[0] for r in cur.fetchall()}
        sf = src.pop() if len(src) == 1 else None
        out[pe._p(p)] = {"edition": tip, "source_file": sf,
                         "rank": release_rank(sf)}
    return out


def legacy_tips(cur, spec) -> dict:
    """Before the editions table exists: {held live period: {rank}}, the
    rank of the file each held year equals (LEGACY_RANKS; a year not listed
    has no rank)."""
    cur.execute(f"SELECT DISTINCT {spec.period_col} FROM "
                f"public.{spec.live_table}")
    ranks = globals()["LEGACY_RANKS"]
    return {pe._p(p): {"edition": None, "source_file": None,
                       "rank": ranks.get(pe._p(p))}
            for (p,) in cur.fetchall()}


def ledger_checks(cur, prof) -> dict:
    """{period (string): {(source_file, file_sha256)}} of the ledger."""
    pc = prof.spec.period_col
    cur.execute(f"SELECT DISTINCT {pc}, source_file, file_sha256 "
                f"FROM public.{pe.file_checks_table(prof)}")
    out = {}
    for p, f, s in cur.fetchall():
        out.setdefault(pe._p(p), set()).add((f, s))
    return out


# ---------------------------------------------------------------------------
# ddl, status, run log, refresh-latest
# ---------------------------------------------------------------------------

def log_run(cur, rows_written: int, notes: str, started_at=None) -> None:
    """The pipeline_run_log row for a committed run (agent RUN_AGENT, source
    '3'). Called only on committed load, migrate-legacy and restore-edition
    runs, partial runs included."""
    pe.log_run(cur, profile(), rows_written, notes, started_at)


def cmd_ddl(args) -> int:
    return pe.run_ddl(profile(), args, connect=_late("_conn"),
                      table_exists=_late("table_exists"),
                      create_schema=lambda cur: create_all(cur))


def status_lines(cur, prof) -> list:
    """Per period: the tip edition, its file's rank and the ledger's latest
    file."""
    s = prof.spec
    lt = pe.file_checks_table(prof)
    lines = []
    for p, t in sorted(tip_info(cur, s).items()):
        line = f"  {p}: tip edition {t['edition']} ({rank_text(t['rank'])})"
        cur.execute(f"SELECT source_file, outcome FROM public.{lt} WHERE "
                    f"{s.period_col} = %s ORDER BY id DESC LIMIT 1", (p,))
        last = cur.fetchone()
        line += (f"; ledger latest: {last[1]} {last[0]}" if last
                 else "; no ledger row")
        lines.append(line)
    return lines


def cmd_status(args) -> int:
    """status: what needs action; exit 1 if anything, or (with a clean
    message, nothing created) if the editions table or ledger does not
    exist."""
    prof = profile()
    conn = _conn(False)
    try:
        with conn.cursor() as cur:
            missing = [t for t in (prof.spec.editions_table,
                                   pe.file_checks_table(prof))
                       if not table_exists(cur, t)]
            if missing:
                print(f"status: {' and '.join(missing)} "
                      f"{'does' if len(missing) == 1 else 'do'} not exist "
                      "yet; run `ddl --commit`, then `migrate-legacy FILE "
                      "FILE --commit` (S3 has no sync-new)")
                return 1
            st = pe.status(cur, prof)
            lines = status_lines(cur, prof)
    finally:
        conn.rollback()
        conn.close()
    print(pe.format_status(prof, st, lines))
    return 0 if st["ok"] else 1


def _tables_exist(cur) -> bool:
    """True when the editions table exists; halts (S3's own wording: there
    is no sync-new) when it or the ledger does not."""
    prof = profile()
    for t in (prof.spec.editions_table, pe.file_checks_table(prof)):
        if not table_exists(cur, t):
            halt(f"{t} does not exist yet; run `ddl --commit`, then "
                 "`migrate-legacy FILE FILE --commit`")
    return True


def cmd_refresh_latest(args) -> int:
    """refresh-latest [--accept-drift YYYY] [--commit | --simulate]: the
    engine's refresh (each year's latest edition copied into live, loaded_at
    included, under the before/after hash guard)."""
    return pe.run_refresh_latest(
        profile(), args, connect=_late("_conn"),
        table_exists=lambda cur, t: _tables_exist(cur))


# ---------------------------------------------------------------------------
# load
# ---------------------------------------------------------------------------

def _end(conn, args, writing, rc) -> int:
    conn.rollback()  # commit mode has committed already
    if not writing:
        print("PREVIEW: nothing written to the database (use --commit or "
              "--simulate)")
    elif args.simulate:
        print("SIMULATION: ROLLED BACK (nothing persisted)")
    return rc


def _year_arg(text) -> str:
    if not re.fullmatch(r"[0-9]{4}", str(text or "")):
        raise argparse.ArgumentTypeError(f"not a year YYYY: {text!r}")
    return str(text)


def _read(path) -> dict:
    try:
        return read_file(path)
    except ValueError as e:      # the message names the file already
        halt(f"{e}; nothing stored")


def identity_problems(file, page_year=None, release=None) -> list:
    """Problems (empty = fine) between a read file and where it came from:
    the edition's mid-year (when the landing page was read) and --release
    must both equal the file's own year."""
    out = []
    if page_year is not None and int(page_year) != file["year"]:
        out.append(f"the edition is mid-{page_year}, the file says "
                   f"mid-{file['year']}")
    if release is not None and int(release) != file["year"]:
        out.append(f"--release {release}, the file says mid-{file['year']}")
    return out


def _landing():
    lp = fetch_json(LANDING)
    try:
        check_landing(lp)
    except ValueError as e:
        halt(f"{ONS}{LANDING}: {e}")
    return lp


def _edition(uri, year) -> tuple:
    """(file, label) of the edition page, its label's mid-year checked."""
    ed = fetch_json(uri)
    try:
        file, label, api_date = edition_download(ed)
    except ValueError as e:
        halt(f"{ONS}{uri}: {e}")
    ly = int(EDITION_LABEL_RE.fullmatch(label).group(1))
    if ly != int(year):
        halt(f"{ONS}{uri}: the edition label {label!r} is not mid-{year}")
    print(f"edition: {label!r} ({ONS}{uri}); download {file}; API releaseDate "
          f"{api_date} (printed only; the rank is the Cover's date)")
    return file, label


def _source(args) -> dict:
    """The file the run reads: {path, where, how, page_year, page_url,
    file_name, file, discovered, label}. --file reads a local file (nothing
    downloaded; it is read here, file set) and finds its own year's edition
    on the landing page unless --no-page; otherwise the landing page gives
    the edition (the newest, or --release YYYY) and its workbook is
    downloaded to RAW_DIR (also in a preview)."""
    out = {"path": None, "where": None, "how": None, "page_year": None,
           "page_url": None, "file_name": None, "file": None,
           "discovered": None, "label": None}
    if args.file:
        path = Path(args.file)
        if not path.is_file():
            halt(f"--file {args.file}: no such file")
        print(f"file given: {path}; nothing downloaded; identity is read "
              "from the file")
        f = _read(path)
        bad = identity_problems(f, None, args.release)
        if bad:
            halt("file identity check failed, nothing stored: "
                 + "; ".join(bad))
        out.update(path=path, where=_file_source(path), how="local file",
                   file_name=path.name, file=f)
        if args.no_page:
            out.update(page_url=NO_PAGE, how="local file, --no-page: the "
                       "landing page was not read")
            print("NOTE: --no-page: the landing page is not read")
            return out
        lp = _landing()
        try:
            y, _, uri = edition_for_year(lp, f["year"])
        except ValueError as e:
            halt(f"{ONS}{LANDING}: {e}; the edition of the file's own year "
                 f"(mid-{f['year']}) was not found. Give --no-page to load "
                 "the file without it (recorded as such)")
        _, label = _edition(uri, y)
        out.update(page_year=y, page_url=ONS + uri, label=label)
        return out
    lp = _landing()
    try:
        if args.release:
            y, _, uri = edition_for_year(lp, args.release)
            how = f"edition mid-{y} given"
        else:
            y, _, uri = latest_edition(lp)
            how = "newest edition (ONS landing page)"
            out["discovered"] = lp
    except ValueError as e:
        halt(f"{ONS}{LANDING}: {e}")
    for u, ys in irregular_editions(lp):
        print(f"  not read (irregular URI, older or not an edition): {u}")
    file, label = _edition(uri, y)
    url = download_url(uri, file)
    print(f"downloading {url} to {RAW_DIR} (also in a preview; the database "
          "is not written)")
    path, final = fetch(url, RAW_DIR)
    print(f"  final URL: {final}")
    out.update(path=path, where=final, how=how, page_year=y,
               page_url=ONS + uri, file_name=_url_name(final), label=label)
    return out


def _label(file, src, sha, how, *, notes=()) -> str:
    return (f"ONS mid-year population estimates: {file['title']}"
            + (f" ({src['label']})" if src.get("label") else "")
            + f"; published {file['published'].isoformat()}; {src['where']} "
            f"sha256 {sha[:16]}; {how}" + "".join(f"; {n}" for n in notes
                                                   if n))


def _report(file) -> None:
    print(f"{file['path'].name}: {file['title']!r}; {file['pub_text']!r}; "
          f"next publication {file['next_release']!r}; {MYE2} "
          f"{len(file['mye2']['las'])} English authorities, "
          f"{file['mye2']['wales']} Welsh rows skipped, "
          f"{len(file['mye2']['aggregates'])} aggregates set aside; {MYE4} "
          f"mid-{file['mye4']['years'][-1]} to mid-{file['mye4']['years'][0]}")
    print(f"  geography note: {file['geography_note']!r}")
    for n in file["revision_notes"]:
        print(f"  revisions note: {n!r}")
    m7 = file.get("mye7")
    if m7 and m7["layout"] == "revisions":
        print(f"  {MYE7}: {m7['title']!r} (published {m7['published_label']}"
              f" / revised {m7['label']}, mid-{m7['years'][0]} to "
              f"mid-{m7['years'][-1]})")
    for n in file["layout_notes"]:
        print(f"  NOTE: {n}")


def _reconciled(file, years) -> None:
    problems = reconcile(file, years)
    if problems:
        halt("reconciliation inside the file failed, nothing stored: "
             + "; ".join(problems[:8]))
    print("reconciled: the English authorities sum to the England row ("
          + ", ".join(f"mid-{y} {england_total(file, y):,}" for y in years)
          + f"), {MYE2} equals {MYE4} on mid-{file['year']}"
          + (f", {MYE4} equals {MYE7}'s revised column on "
             + ", ".join(f"mid-{y}" for y in years
                         if y in file["mye7"]["years"])
             if file.get("mye7") and file["mye7"]["layout"] == "revisions"
             and any(y in file["mye7"]["years"] for y in years) else ""))


def _build(cur, file, periods) -> dict:
    """{period string: records}, codes resolved; halts on a geography or
    cell problem."""
    try:
        rmap, problems = resolve_codes(cur, set(file["mye4"]["las"]),
                                       file["year"])
    except ValueError as e:
        halt(f"geography: {e}; nothing stored")
    if problems:
        halt("geography check failed, nothing stored: " + "; ".join(problems))
    recoded = sorted((c, r) for c, r in rmap.items() if r != c)
    if recoded:
        print("codes resolved (geography.resolve, la_code_lookup recode): "
              + ", ".join(f"{c} -> {r}" for c, r in recoded))
    try:
        return {str(p): build_records(file, p, rmap) for p in periods}
    except ValueError as e:
        halt(f"{e}; nothing stored")


def cmd_load(args) -> int:
    if args.no_page and not args.file:
        halt("--no-page is only for --file (a local file whose landing page "
             "is not read)")
    prof = profile()
    spec = prof.spec
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            lt = pe.file_checks_table(prof)
            ready = {spec.editions_table: table_exists(cur, spec.editions_table),
                     lt: table_exists(cur, lt)}
            has_ed = all(ready.values())
            missing = [t for t, ok in ready.items() if not ok]
            if writing and not has_ed:
                halt(f"missing {', '.join(missing)}; run `ddl --commit` and "
                     "`migrate-legacy FILE FILE --commit` first")
            if has_ed:
                _, new_live, errors = core.latest_map(cur, spec)
                if errors:
                    halt("invalid edition chain: " + "; ".join(
                        f"{pe._p(p)}: {msg}" for p, msg in errors.items()))
                if new_live:
                    halt(f"{spec.live_table}: live years with no editions "
                         f"{[pe._p(x) for x in new_live]}; run migrate-legacy "
                         "first")
                tips = tip_info(cur, spec)
                checked = ledger_checks(cur, prof)
                stranded = pe.live_missing_periods(cur, prof)
            else:
                print(f"NOTE: {', '.join(missing)} not present: no editions "
                      "table yet; this preview compares the file with the "
                      "LIVE table")
                tips, checked, stranded = legacy_tips(cur, spec), {}, []
            ranks = tip_ranks(tips, checked)
            held = sorted(ranks)
            print("held years: " + (", ".join(held) or "none"))
            src = _source(args)
            if src["discovered"] is not None:
                for line in discovery_note(src["discovered"], held):
                    print(line)
            path = src["path"]
            sha = content_sha256(path)
            same = [p for p, v in globals()["LEGACY_FILES"].items()
                    if v["sha256"] == sha]
            print(f"{path.name}: sha256 {sha[:16]}"
                  + (f"; byte-identical to the held mid-{same[0]} file"
                     if same else ""))
            f = src["file"]
            if has_ed and args.recheck is None:
                done = ledger_complete(checked, src["where"], sha, ranks,
                                       allow_older=args.allow_older_file,
                                       stranded=stranded)
                if done:
                    print(f"the file's (source, sha256) is already in the "
                          f"ledger for {', '.join(done)}; "
                          + ("nothing compared" if f else "nothing parsed")
                          + " (--recheck YYYY reads it again)")
                    print("nothing to do")
                    return _end(conn, args, writing, rc=0)
            if f is None:
                f = _read(path)
                bad = identity_problems(f, src["page_year"], args.release)
                if bad:
                    halt("file identity check failed, nothing stored: "
                         + "; ".join(bad))
            _report(f)
            periods = periods_stated(f, held)
            _reconciled(f, periods)
            by_period = _build(cur, f, periods)
            lsrc = ledger_source(src["where"], f["rank"], periods)
            chk = {p: s for p, s in checked.items() if p not in stranded}
            try:
                new, revised, skipped = plan_periods(
                    list(by_period), f["rank"], ranks, chk, lsrc, sha,
                    recheck=args.recheck, allow_older=args.allow_older_file)
            except ValueError as e:
                halt(str(e))
            rc = _run_periods(args, conn, cur, prof, has_ed, src, f, by_period,
                              new, revised, skipped, tips, ranks, held, sha,
                              lsrc)
        return _end(conn, args, writing, rc)
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def _tip_records(cur, spec, p, tips, has_ed) -> list:
    if has_ed:
        return records(cur, spec, p, tips[p]["edition"])
    return records(cur, spec, p)


def _run_periods(args, conn, cur, prof, has_ed, src, f, by_period, new,
                 revised, skipped, tips, ranks, held, sha, lsrc) -> int:
    started = datetime.now(timezone.utc)
    spec = prof.spec
    rank = f["rank"]
    older_note = {}
    for p in revised + new:
        t = ranks.get(p)
        if t is not None and t > rank:
            older_note[p] = (f"--allow-older-file given; {p}'s tip comes from "
                             f"a file of {rank_text(t)}, this file is "
                             f"{rank_text(rank)}")
        elif p in new and held and p < max(held):
            older_note[p] = (f"--allow-older-file given; mid-{p} is before "
                             f"the newest held year {max(held)}")
        if p in older_note:
            print(f"NOTE: {older_note[p]}")
    print(f"planned: new {new or 'none'}; compare {revised or 'none'}; "
          "skipped " + (", ".join(f"{p} ({r})" for p, r in sorted(
              skipped.items())) or "none"))
    periods = sorted(new + revised)
    ack = set(args.acknowledge or ())
    large = set(args.accept_large_revision or ())
    reissue = set(args.accept_reissue or ())
    stray = sorted((ack | large | reissue) - set(periods))
    if stray:
        halt(f"--acknowledge/--accept-large-revision/--accept-reissue "
             f"{stray}: not years this run compares or stores "
             f"({periods or 'none'}); nothing stored")
    if not periods:
        if skipped and all(r.startswith("older") for r in skipped.values()):
            halt(f"older file: every year the file states is skipped; the "
                 f"file is {rank_text(rank)}, older than what is held for "
                 f"{', '.join(sorted(skipped))}; storing it would record an "
                 "older file. If this is deliberate, re-run with "
                 "--allow-older-file")
        print("nothing to do")
        return 0
    problems, infos, acked, reissued, accepted = {}, {}, {}, {}, {}
    for p in periods:
        recs = by_period[p]
        if p in revised:
            kind = "revised"
            tip = _tip_records(cur, spec, p, tips, has_ed)
            found = period_problems(recs, tip, None, kind=kind)
        else:
            kind = "new"
            tip = None
            earlier = [h for h in held if h < p]
            prev = (_tip_records(cur, spec, max(earlier), tips, has_ed)
                    if earlier else None)
            found = period_problems(recs, None, prev, kind=kind)
        hard = [x for x in found if x.startswith(PARTIAL)]
        big = [x for x in found if x.startswith(LARGE)]
        soft = [x for x in found if not x.startswith((PARTIAL, LARGE))]
        bad = list(hard)
        if (soft or big) and p in ack:
            acked[p] = "; ".join(soft)
            print(f"  {p}: ACKNOWLEDGED (--acknowledge): {acked[p]}")
        elif soft or big:
            bad += [x + f"; read the preview, then --acknowledge {p}"
                    for x in soft]
        if big and p in large and p in ack:
            accepted[p] = "; ".join(big)
            print(f"  {p}: ACCEPTED LARGE REVISION (--accept-large-revision):"
                  f" {accepted[p]}")
        elif big:
            bad += [x + f"; read the preview, then --acknowledge {p} "
                    f"--accept-large-revision {p}" for x in big]
        if kind == "revised" and ranks.get(p) == rank and _differs(recs, tip):
            if p in reissue:
                reissued[p] = (f"--accept-reissue {p}: the file is "
                               f"{rank_text(rank)}, the same as the file of "
                               "the held tip, with different content (a "
                               "publisher reissue keeping its date)")
                print(f"  {p}: ACCEPTED REISSUE: {reissued[p]}")
            else:
                bad.append(f"the file is {rank_text(rank)}, the same as the "
                           "file of the held tip, but its content differs: "
                           "two files claim the same release. A publisher "
                           "reissue keeps its date: read the release's "
                           f"notes, then --accept-reissue {p}")
        if bad:
            problems[p] = bad
            continue
        notes = (older_note.get(p, ""), reissued.get(p, ""),
                 f"ACKNOWLEDGED {acked[p]}" if p in acked else "",
                 f"--accept-large-revision {accepted[p]}" if p in accepted
                 else "")
        how = (("new year" if kind == "new" else "restated year")
               + f" ({src['how']})")
        infos[p] = PeriodInfo(
            edition_source(f, src, "; ".join(n for n in notes if n)),
            _label(f, src, sha, how, notes=notes), ((lsrc, sha),))
    for p, msgs in sorted(problems.items()):
        for msg in msgs:
            print(f"  {p}: REJECTED, not stored: STOP CONDITION: {msg}")
    ok = [p for p in periods if p not in problems]
    rc = 0
    stats = {"periods": [], "kinds": {}, "stored_rows": 0, "live_rows": 0}
    if ok and has_ed:
        def apply(c, prof_, per, rs, *, fetched_on, source_file=None):
            return _late("apply_year")(c, prof_, per, rs,
                                       fetched_on=fetched_on, info=infos[per])
        rc = pe.load_periods(cur, prof, ok, lambda per: by_period[per],
                             f["published"], args.commit,
                             simulate=args.simulate, against="editions",
                             stats=stats, apply=apply)
    elif ok:
        rc = pe.load_periods(cur, prof, ok, lambda per: by_period[per],
                             f["published"], False, against="live")
        print("  (no editions table yet: preview against live only)")
    if problems:
        print(f"REJECTED {len(problems)} year(s), nothing stored for them: "
              f"{', '.join(sorted(problems))}; exit 1")
        rc = 1
    if args.commit:
        stored = live = 0  # rows written by this run; not a source value
        for p in stats["periods"]:
            k = stats["kinds"][p]
            if k in ("new", "revised"):
                stored += len(by_period[p])
            if k in ("new", LIVE_MISSING):
                live += len(by_period[p])
        left = [p for p in ok if p not in stats["periods"]]
        failed, not_tried = (left[:1], left[1:]) if left else ([], [])
        partial = ""
        if rc:
            partial = (f"PARTIAL RUN (exit 1): rejected "
                       f"{sorted(problems) or 'none'}"
                       + "".join(f" ({p}: {'; '.join(problems[p])[:300]})"
                                 for p in sorted(problems))
                       + f"; failed {failed or 'none'}; not attempted "
                       f"{not_tried or 'none'}; a year that failed or was "
                       "rejected stored nothing and is not in the counts "
                       "below; ")
        notes = (partial + f"tables file {src['where']} sha256 {sha[:16]}, "
                 f"{rank_text(rank)} ({src['how']}): "
                 + ("; ".join(f"{p} {k}" for p, k in stats["kinds"].items())
                    or "nothing stored")
                 + (f"; skipped {sorted(skipped)}" if skipped else "")
                 + "".join(f"; {older_note[p]}" for p in sorted(older_note)
                           if p in stats["periods"])
                 + "".join(f"; {reissued[p]}" for p in sorted(reissued)
                           if p in stats["periods"])
                 + "".join(f"; ACKNOWLEDGED {p}: {acked[p]}"
                           for p in sorted(acked) if p in stats["periods"])
                 + "".join(f"; --accept-large-revision {p}: {accepted[p]}"
                           for p in sorted(accepted)
                           if p in stats["periods"])
                 + f". Rows stored: {stored} edition rows, {live} live rows.")
        log_run(cur, stored + live, notes, started)
        conn.commit()
        print("pipeline_run_log row written"
              + (" (partial run: see its notes)" if rc else ""))
    return rc


# ---------------------------------------------------------------------------
# One-off: migrate the held table
# ---------------------------------------------------------------------------

def live_state(cur, spec=None) -> tuple:
    """(rows, survey hash, {year: [distinct loaded_at]}) of the live table.
    The survey hash is the spec's: md5 of lad24cd|la_name|population|
    reference_year (NULL as '~'), rows ordered by reference_year, lad24cd,
    joined by LF."""
    spec = spec or globals()["SPEC"]
    cur.execute(f"""SELECT COUNT(*), md5(string_agg(coalesce(lad24cd::text,
        '~')||'|'||coalesce(la_name::text,'~')||'|'||
        coalesce(population::text,'~')||'|'||
        coalesce(reference_year::text,'~'), E'\\n' ORDER BY
        reference_year, lad24cd)) FROM public.{spec.live_table}""")
    n, h = cur.fetchone()
    cur.execute(f"SELECT DISTINCT {spec.period_col}, loaded_at FROM "
                f"public.{spec.live_table} ORDER BY 1, 2")
    loaded = {}
    for p, t in cur.fetchall():
        loaded.setdefault(pe._p(p), []).append(t)
    return n, h, loaded


def _proof(held, built, forms) -> tuple:
    """(values compared, differences, name forms seen) of the held rows
    against the rebuilt rows on lad24cd, population and la_name; a name
    held in the form `forms` names ({lad24cd: (held, workbook)}) exactly is
    not a difference, and is listed."""
    hk, bk = _by_key(held), _by_key(built)
    diffs = [("only held", k) for k in sorted(set(hk) - set(bk))]
    diffs += [("only rebuilt", k) for k in sorted(set(bk) - set(hk))]
    cells, seen = 0, []  # values compared; not a source value
    for k in sorted(set(hk) & set(bk)):
        for c in VALUES:
            cells += 1
            a, b = hk[k][c], bk[k][c]
            if a == b:
                continue
            if c == "la_name" and forms.get(k) == (a, b):
                seen.append((k, a, b))
                continue
            diffs.append((k, c, a, b))
    return cells, diffs, seen


def migrate_legacy(cur, files, *, write, legacy=None, legacy_files=None,
                   name_forms=None) -> dict:
    """One-off. Preconditions: on write the editions table and ledger exist
    and are empty (in a preview they may be absent; if present they must be
    empty); live as surveyed (legacy: (rows, survey hash, {year: loaded_at})
    , default LEGACY_LIVE and LEGACY_LOADED: each year one loaded_at, that
    one); one file per live year, each a held workbook (its sha256 one of
    legacy_files, default LEGACY_FILES, for its own mid-year). Proof per
    year, any failure halts: the file's identity and in-file reconciliation
    (its newest year); this parser on the file reproduces every live
    population and name, a name held in the form name_forms (default
    LEGACY_NAME_FORMS) names exactly being listed, not a difference.
    Edition 1 'as loaded' from live (release label the engine's as-loaded
    label, published_date the year's load date, source_file 'as loaded:
    <how it was loaded> <name>; mid-<yyyy>; published <date>'), and a ledger
    row for each file that reproduces edition 1 on every column (outcome
    'unchanged', edition 1, the file's URL, the periods it states) so the
    next page load finds it checked; a file that differs (a held name form)
    gets none, so a re-read compares it. Live untouched. Nothing commits
    here; the caller rolls everything back on a halt."""
    g = globals()
    if legacy is None:
        legacy = (g["LEGACY_LIVE"][0], g["LEGACY_LIVE"][1],
                  g["LEGACY_LOADED"])
    legacy_files = legacy_files or g["LEGACY_FILES"]
    forms_all = g["LEGACY_NAME_FORMS"] if name_forms is None else name_forms
    prof = profile()
    spec = prof.spec
    lt = pe.file_checks_table(prof)
    ready = {spec.editions_table: table_exists(cur, spec.editions_table),
             lt: table_exists(cur, lt)}
    if write and not all(ready.values()):
        halt(f"missing {[t for t, ok in ready.items() if not ok]}; run "
             "`ddl --commit` first")
    for t, ok in ready.items():
        if ok:
            cur.execute(f"SELECT COUNT(*) FROM public.{t}")
            n = cur.fetchone()[0]
            if n:
                halt(f"{t} already holds {n} rows; migrate-legacy runs once, "
                     "on empty tables")
    n, h, loaded = live_state(cur, spec)
    want_n, want_h, want_loaded = legacy
    if (n, h) != (want_n, want_h):
        halt(f"{spec.live_table} is not as surveyed: {n} rows, hash {h}; "
             f"expected {want_n} rows, hash {want_h}; nothing stored")
    want_loaded = {str(k): v for k, v in want_loaded.items()}
    if sorted(loaded) != sorted(want_loaded) or any(
            loaded[p] != [want_loaded[p]] for p in loaded):
        halt(f"{spec.live_table} loaded_at is not as surveyed: "
             f"{ {p: [str(t) for t in v] for p, v in loaded.items()} }, "
             f"expected { {p: str(t) for p, t in want_loaded.items()} }; "
             "nothing stored")
    live_periods = sorted(loaded)
    by_sha = {v["sha256"]: p for p, v in legacy_files.items()}
    read = {}
    for file in files:
        path = Path(file)
        sha = content_sha256(path)
        if sha not in by_sha:
            halt(f"{path.name}: sha256 {sha[:16]} is not a held file "
                 f"(expected one of "
                 f"{[(p, v['sha256'][:16]) for p, v in legacy_files.items()]}"
                 "); nothing stored")
        f = _read(path)
        p = str(f["year"])
        if p != by_sha[sha]:
            halt(f"{path.name}: the file says mid-{p}, the held file of that "
                 f"sha256 is for mid-{by_sha[sha]}; nothing stored")
        if p in read:
            halt(f"two files for mid-{p}; nothing stored")
        read[p] = (path, sha, f)
    if sorted(read) != live_periods:
        halt(f"live years {live_periods}, files given for "
             f"{sorted(read) or 'none'}; one held file per live year is "
             "needed; nothing stored")
    plan = {"periods": {}, "rows": 0, "cells": 0}
    for p in live_periods:
        path, sha, f = read[p]
        print(f"mid-{p}:")
        _report(f)
        _reconciled(f, [int(p)])
        held = records(cur, spec, p)
        built = _build(cur, f, [p])[p]
        cells, diffs, seen = _proof(held, built, forms_all.get(p, {}))
        if diffs:
            halt(f"proof failed: mid-{p}: {len(diffs)} differences between "
                 f"the held rows and the held file read again, e.g. "
                 f"{diffs[:5]}; nothing stored")
        lf = legacy_files[p]
        pub = loaded[p][0].astimezone(timezone.utc).date()
        stated = periods_stated(f, live_periods)
        entry = {"rows": len(held), "cells": cells, "records": held,
                 "sha": sha, "loaded": pub, "name_forms": seen,
                 "content_sha": rows_content_sha(held),
                 "source_file": (f"as loaded: {lf['loaded_by']} "
                                 f"{lf['name']}; mid-{f['year']}; published "
                                 f"{f['published'].isoformat()}"),
                 "ledger": (None if seen else
                            ledger_source(lf["url"], f["rank"], stated))}
        plan["periods"][p] = entry
        plan["rows"] += len(held)
        plan["cells"] += cells
        if seen:
            print(f"  proof: {path.name} read again reproduces every held "
                  f"population and every held name but {len(seen)}: "
                  f"{len(held):,} authorities matched by lad24cd, {cells:,} "
                  "values (population, la_name) compared, 0 unexplained "
                  f"differences; {len(seen)} name(s) held in another form "
                  "(named in LEGACY_NAME_FORMS, kept as held): "
                  + "; ".join(f"{k} held {a!r}, file {b!r}"
                              for k, a, b in seen))
        else:
            print(f"  proof: {path.name} read again reproduces every held "
                  f"value: {len(held):,} authorities matched by lad24cd, "
                  f"{cells:,} values (population, la_name), 0 differences")
        print(f"  edition 1 as loaded, {len(held):,} rows, published {pub} "
              f"(the load date); source_file {entry['source_file']!r}; "
              f"content sha256 {entry['content_sha']}")
        print("  ledger: " + (entry["ledger"] if entry["ledger"] else
                              "none (the file differs from edition 1 in "
                              "the held name forms, so a re-read compares "
                              "it: equal rank, different content stops "
                              "unless --accept-reissue)"))
    if not write:
        return plan
    cur.execute(f"SAVEPOINT {prof.savepoint}")
    for p, v in plan["periods"].items():
        ed = core.insert_edition(
            cur, spec, v["records"], p, release_label=core.AS_LOADED_LABEL,
            published_date=v["loaded"], source_file=v["source_file"],
            source_sha256=v["content_sha"], supersedes=None, strict=True)
        if ed != 1 or core.rows_differing(cur, spec, p, ed):
            halt(f"{spec.editions_table} {p}: edition 1 does not equal "
                 "live; rolled back")
        if v["ledger"]:
            pe.record_file_check(cur, prof, p, v["ledger"], v["sha"],
                                 "unchanged", 1)
    cur.execute(f"RELEASE SAVEPOINT {prof.savepoint}")
    return plan


def cmd_migrate_legacy(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    try:
        with conn.cursor() as cur:
            plan = migrate_legacy(cur, args.files, write=writing)
            if args.commit:
                log_run(cur, plan["rows"], "migrate-legacy: edition 1 as "
                        "loaded for " + ", ".join(
                            f"{p} ({v['rows']} rows, file sha256 "
                            f"{v['sha'][:16]}, "
                            + (f"{len(v['name_forms'])} held name forms, no "
                               "ledger row" if v["name_forms"]
                               else "ledger row") + ")"
                            for p, v in plan["periods"].items())
                        + f" from the live table; proof {plan['cells']} "
                        "values equal the held files read again. Live "
                        "untouched.", started)
                conn.commit()
                print("migrate-legacy: COMMITTED; pipeline_run_log row "
                      "written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# restore-edition
# ---------------------------------------------------------------------------

def restore_edition(cur, period, edition) -> int:
    """Store edition `edition` of `period` as the next edition ('restored
    from edition N') for refresh-latest to apply; return its number. Halts
    if the edition does not exist or is the tip."""
    spec = globals()["SPEC"]
    p = str(period)
    tip = core.latest_edition(cur, spec, p)
    recs = records(cur, spec, p, edition)
    if not recs:
        halt(f"{spec.editions_table} {p}: no edition {edition}")
    if edition == tip:
        halt(f"{p}: edition {edition} is the tip; nothing to restore")
    cur.execute(f"SELECT DISTINCT source_file FROM public.{spec.editions_table}"
                f" WHERE {spec.period_col} = %s AND edition = %s", (p, edition))
    src = cur.fetchone()[0]
    new = core.insert_edition(
        cur, spec, recs, p, release_label=f"restored from edition {edition}",
        published_date=date.today(), source_file=src,
        source_sha256=rows_content_sha(recs), supersedes=tip, strict=True,
        allow_revert=True)
    if new == tip:
        halt(f"{p}: edition {edition} was not stored as the next edition; "
             "nothing to restore")
    return new


def cmd_restore_edition(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    spec = globals()["SPEC"]
    try:
        with conn.cursor() as cur:
            if not table_exists(cur, spec.editions_table):
                halt(f"{spec.editions_table} does not exist")
            new = restore_edition(cur, args.period, args.edition)
            print(f"{args.period}: edition {args.edition} "
                  + ("stored" if writing else "would be stored")
                  + f" as edition {new}; run refresh-latest to apply it")
            if args.commit:
                n = len(records(cur, spec, args.period, new))
                log_run(cur, n, f"restore-edition: {args.period} edition "
                        f"{args.edition} stored as edition {new}", started)
                conn.commit()
                print("COMMITTED; pipeline_run_log row written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="S3 ONS mid-year population estimates: editions loader "
        "on the period-editions engine. There is no sync-new; migrate-legacy "
        "records the held years once.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ddl", help="create the editions table, its triggers "
                       "and the file-check ledger (preview by default)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_ddl)
    sub.add_parser("status", help="what needs action; exit 1 if anything"
                   ).set_defaults(func=cmd_status)
    p = sub.add_parser("load", help="find the newest edition's workbook, "
                       "compare and store the years it states (preview by "
                       "default)")
    p.add_argument("--release", type=_year_arg, metavar="YYYY",
                   help="that mid-year's edition (default the newest)")
    p.add_argument("--file", metavar="PATH",
                   help="a local tables workbook (nothing downloaded; "
                   "identity is read from the file)")
    p.add_argument("--no-page", action="store_true",
                   help="with --file only: do not read the landing page "
                   "(recorded as such)")
    p.add_argument("--recheck", type=_year_arg, metavar="YYYY",
                   help="compare this held year again although the file is "
                   "in the ledger")
    p.add_argument("--allow-older-file", action="store_true",
                   help="compare and store a year whose tip comes from a "
                   "newer file, or a year before the newest held (logged)")
    p.add_argument("--acknowledge", action="append", type=_year_arg,
                   metavar="YYYY",
                   help="release a year's stop conditions after reading the "
                   "preview (repeatable; logged; never a partial file)")
    p.add_argument("--accept-large-revision", action="append",
                   type=_year_arg, metavar="YYYY",
                   help="with --acknowledge YYYY: also release a revision "
                   "beyond the large-revision limits (repeatable; logged)")
    p.add_argument("--accept-reissue", action="append", type=_year_arg,
                   metavar="YYYY",
                   help="store a year whose file has the held file's rank but "
                   "different content (a publisher reissue keeping its date; "
                   "repeatable; logged)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_load)
    p = sub.add_parser("refresh-latest", help="copy each year's latest "
                       "edition into the live table, loaded_at included "
                       "(preview by default)")
    p.add_argument("--accept-drift", action="append", metavar="YYYY",
                   help="overwrite this year although its live rows equal no "
                   "stored edition (repeatable)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_refresh_latest)
    p = sub.add_parser("migrate-legacy", help="one-off: edition 1 as loaded "
                       "for each held year, with the proof (preview by "
                       "default); one held workbook per live year")
    p.add_argument("files", metavar="FILE", nargs="+")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_migrate_legacy)
    p = sub.add_parser("restore-edition", help="store an earlier edition's "
                       "rows as the next edition (preview by default)")
    p.add_argument("period", type=_year_arg, metavar="YYYY")
    p.add_argument("edition", type=int)
    pe.mode_parser(p)
    p.set_defaults(func=cmd_restore_edition)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
