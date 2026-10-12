"""S21 (ONS clustering similar local authorities and statistical nearest
neighbours in the UK, Table 7a: LTLA global statistical nearest
neighbours): edition history and loader on the period-editions engine, with
a single-vintage live table that is a derived subset of the editions.

The live table la_statistical_neighbours keeps its name, key (lad24cd,
statistical_neighbour_lad24cd) and columns. It holds one release, the
engine period, an ISO date that sorts (rule 8): '2026-03-18'. The live table
gains one additive column for it, run by `ddl` only:

    ALTER TABLE public.la_statistical_neighbours ADD COLUMN release_date
        date NOT NULL DEFAULT '2026-03-18' CHECK (release_date = '2026-03-18')

(default-filled, not in the key; the CHECK makes a second release in live
impossible). This module owns the editions table and file-check ledger,
append-only:

    SPEC     la_statistical_neighbours_editions               period release_date
    ledger   la_statistical_neighbours_editions_file_checks

An edition is the publisher's Table 7a for every English authority: all 20
neighbours, each with the publisher's rank, whatever its nation. Per
(subject, neighbour) row: statistical_neighbour_lad24cd (the resolved code
of an English neighbour, the publisher's code otherwise), and the
editions-only neighbour_code_published, neighbour_name, publisher_rank
(1..20), england_rank (the order among the English neighbours, 1..n; NULL
for a neighbour that is not an English authority, reason in null_reasons),
neighbour_nation, neighbour_rank (the live rank under LIVE_RULE; NULL
outside live, reason in null_reasons) and null_reasons.

The live table is derived from the latest edition by one named rule,
LIVE_RULE (live_rows applies it; nothing else decides what is live): per
subject, the publisher's neighbours whose code is English (E06, E07, E08,
E09; every such code resolves into la_boundaries or the load stops), in
publisher rank order; the first five are live, numbered 1 to 5. This is the
old build's rule (it reproduces all 1,480 held live rows), with
la_boundaries in place of staging_la_signals. Under it, on the 2026 file,
238 of the 1,480 live pairs carry a rank other than the publisher's, for 71
authorities (whose publisher top five include a Welsh, Scottish or Northern
Irish neighbour), and 105 pairs come from publisher ranks 6 to 15 (the
furthest is 15). The alternative (Decisions item 1 of the 2026-10-11 spec:
the publisher's top five as published, English only) is a change of
LIVE_RULE and one `refresh-latest --acknowledge 2026-03-18`, not a reload:
the editions keep publisher_rank and england_rank for every row. Live holds
a subset of the editions rows, so this loader has its own insert_live,
refresh-latest, status and live_equals_tip (live compared with
live_rows(latest edition)); no engine change.

A reissue of the live release (the same Cover date) is the next edition of
it; it reaches live only through refresh-latest, which needs --acknowledge
when it changes a live pair or rank (the comparator sets of the
deliverables that read the table change). A different ONS edition (another
Cover date) is stored as editions only and never enters live through load
or refresh-latest (EditionSpec.live_single_period): adopting it is a
decision.

History (so the wording here stays truthful): the live table was loaded on
2026-07-24 08:56:26 UTC (pipeline_run_log 66) by
scripts/s21_statistical_neighbours_build.py (in the repo from 2026-08-20,
59ae151) from data/raw/ons_statistical_neighbours_2026.xlsx, by TRUNCATE
then insert. That build took its list of authorities from the latest
staging_la_signals run (a W1 output), dropped every non-English neighbour
and every code not in that list, kept the first five left and renumbered
them 1 to 5; it recoded Barnsley and Sheffield with a private dict whose
comment says the file "still carries the pre-recode codes" above a dict
mapping the new codes to the old ones (the file carries the new codes).
migrate-legacy (one-off) records the release as edition 1 "as loaded": the
5,920 rows of the held file (byte-identical to the file the old build read,
data/raw/s21_neighbours), with the proof that LIVE_RULE applied to them
reproduces every live pair and rank.

Identity (rule 3), from the file itself on every path: the Cover sheet's
title 'Clustering similar local authorities and statistical nearest
neighbours in the UK' and its line 'The data tables in this spreadsheet
were originally published at 9.30am <d Month yyyy>.' (the release date,
the period and the rank); the Contents entry of 7.a and the sheet's own
title 'Table 7a: LTLA global statistical nearest neighbours'; its header
exactly 'Local Authority Code, Local Authority Name, Neighbour 1: code,
Neighbour 1: name, ..., Neighbour 20: code, Neighbour 20: name'; Notes 2
(the boundaries, 'as at March 2026') and 12 (the "[x]" marker) captured
verbatim; docProps modified (printed). When the landing page is read the
edition label's year must be the Cover's and its API releaseDate must be
the Cover date (or the evening before, a midnight in British Summer Time).
The rank of a file is its Cover date, never its name, URL or the API date.

Blanks and markers (docs/RULES.md rule 1): every subject row carries 20
code and name pairs; a blank pair, the note 12 marker "[x]", or any other
text in a code cell halts naming the cell (code_cell): none appears in the
2026 file. Nothing is coerced.

Geography (rule 4): geography.resolve('21', codes, by_period={release:
codes}) (source 21 is declared 'new': the file carries E08000038/39, as
subject and as neighbour; an E08000016/19 halts) on every subject code and
every English neighbour code, so E08000038/39 resolve to E08000016/19; then
each canonical code must be in la_boundaries, anything else is UNEXPLAINED
and stops. Welsh, Scottish and Northern Irish neighbour codes are stored
verbatim and never resolved. No private recode dict.

Stop conditions (in load; a stopped release is REJECTED: nothing stored, no
ledger row, exit 1): fewer English subjects than EXPECTED_SUBJECTS, or
(reissue) a subject of the held edition missing (PARTIAL; no flag releases
it); a release other than LIVE_RELEASE into an empty live table (no flag
releases it); a reissue changing any pair, rank, name or nation against the
held edition unless --acknowledge RELEASE; equal Cover date with different
content (two files claiming the same release) unless --accept-reissue
RELEASE; an older file (a release before the newest held) halts unless
--allow-older-file. Each acceptance is written to the edition's label and
the run log.

Not a W1, export or map input. Read by analysis/council_briefing and
analysis/liverpool_paper (comparator sets for one authority).

Subcommands (every writing command previews by default; --commit and
--simulate are mutually exclusive; --simulate runs the --commit path and
always rolls back). There is no sync-new: migrate-legacy records the held
release.
    python scripts/s21_neighbours_editions.py ddl [--commit | --simulate]
    python scripts/s21_neighbours_editions.py status
    python scripts/s21_neighbours_editions.py load [--file PATH [--no-page]]
        [--recheck YYYY-MM-DD] [--allow-older-file]
        [--accept-reissue YYYY-MM-DD ...] [--acknowledge YYYY-MM-DD ...]
        [--commit | --simulate]
        # the file is downloaded to data/raw/s21_neighbours/ (also in a
        # preview; a preview writes nothing to the database)
    python scripts/s21_neighbours_editions.py refresh-latest
        [--acknowledge YYYY-MM-DD] [--commit | --simulate]
    python scripts/s21_neighbours_editions.py migrate-legacy FILE
        [--commit | --simulate]
    python scripts/s21_neighbours_editions.py restore-edition YYYY-MM-DD N
        [--commit | --simulate]
"""
import argparse
import dataclasses
import hashlib
import json
import math
import re
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parent))

import editions_core as core  # noqa: E402
import geography  # noqa: E402
import load_checks  # noqa: E402
import manual_input  # noqa: E402
import period_editions as pe  # noqa: E402
from editions_core import halt  # noqa: E402

RUN_AGENT = "s21_neighbours_editions"
RUN_SOURCE = "21"
LIVE_MISSING = pe.LIVE_MISSING
HELD_NOT_LIVE = pe.HELD_NOT_LIVE
UTC = timezone.utc

ONS = "https://www.ons.gov.uk"
LANDING = ("/peoplepopulationandcommunity/wellbeing/datasets/"
           "clusteringsimilarlocalauthoritiesandstatisticalnearest"
           "neighboursintheuk")
LANDING_TITLE = ("Clustering similar local authorities and statistical "
                 "nearest neighbours in the UK")
REPO = Path(__file__).resolve().parent.parent
RAW_DIR = REPO / "data" / "raw" / "s21_neighbours"     # git-ignored
FILE_ROOTS = manual_input.ROOTS
USER_AGENT = ("ucws-pipeline S21 loader (read-only download of the public "
              "ONS clustering and statistical nearest neighbours tables)")
MIN_FILE_BYTES = 100 * 1024
NO_PAGE = "--file --no-page: the landing page was not read"

# ---------------------------------------------------------------------------
# Columns (live types from information_schema, checked 2026-10-11: lad24cd
# and statistical_neighbour_lad24cd character varying(9) NOT NULL (the
# key), neighbour_rank integer NOT NULL, source text NOT NULL, loaded_at
# timestamptz NOT NULL DEFAULT now(); no release_date column until ddl)
# ---------------------------------------------------------------------------

KEY = ("lad24cd", "statistical_neighbour_lad24cd")
PERIOD = "release_date"
RANK = "neighbour_rank"
NATIONS = {"E": "England", "W": "Wales", "S": "Scotland",
           "N": "Northern Ireland"}
EXTRA_TYPES = (
    ("neighbour_code_published", "varchar(9) NOT NULL CHECK "
     "(neighbour_code_published ~ '^[ENSW][0-9]{8}$')"),
    ("neighbour_name", "text NOT NULL"),
    ("publisher_rank", "integer NOT NULL CHECK (publisher_rank BETWEEN 1 "
     "AND 20)"),
    ("england_rank", "integer CHECK (england_rank BETWEEN 1 AND 20)"),
    ("neighbour_nation", "text NOT NULL CHECK (neighbour_nation IN "
     "('England', 'Wales', 'Scotland', 'Northern Ireland'))"),
    ("null_reasons", "text"),
)
# the editions-only columns (the live table does not have them)
EDITIONS_ONLY = tuple(c for c, _ in EXTRA_TYPES)
# what a record carries and the comparison and content hash use
CONTENT = (RANK,) + EDITIONS_ONLY
NOT_ENGLISH = "not_an_english_authority"
OUTSIDE_LIVE = "outside_live_rule"

# ---------------------------------------------------------------------------
# The file
# ---------------------------------------------------------------------------

COVER = "Cover sheet"
CONTENTS = "Contents"
NOTES = "Notes"
SHEET = "7.a"
COVER_TITLE = LANDING_TITLE
SHEET_TITLE = "Table 7a: LTLA global statistical nearest neighbours"
PUB_RE = re.compile(r"The data tables in this spreadsheet were originally "
                    r"published at 9\.30am ([0-9]{1,2}) ([A-Za-z]+) "
                    r"([0-9]{4})\.")
NOTE2_RE = re.compile(r"Models are based on the latest local authority and "
                      r"combined authority boundaries, as at ([A-Z][a-z]+ "
                      r"[0-9]{4})\..*", re.S)
MARKER = "[x]"
CORE_XML = "docProps/core.xml"
NEIGHBOURS = 20
HEADER = (("Local Authority Code", "Local Authority Name")
          + tuple(x for i in range(1, NEIGHBOURS + 1)
                  for x in (f"Neighbour {i}: code", f"Neighbour {i}: name")))
WIDTH = len(HEADER)
CODE_RE = re.compile(r"[ENSW][0-9]{8}")
EXPECTED_SUBJECTS = 296      # English subjects (LAD May 2024 authorities)
# the one release the live table may hold (its CHECK)
LIVE_RELEASE = "2026-03-18"
PARTIAL = "PARTIAL FILE"     # problems no flag releases
ACK_HINT = "read the preview, then --acknowledge {p}"

# ---------------------------------------------------------------------------
# The live rule (named; live_rows applies it and nothing else decides what
# is live)
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class LiveRule:
    """Which editions rows are live and their neighbour_rank. Per subject,
    the neighbours whose published code starts with one of
    english_prefixes (and, when within_publisher_rank is set, whose
    publisher_rank is at most it), in publisher rank order; the first
    `take` are live, numbered 1.. in that order when renumber, else keeping
    the publisher's rank."""
    name: str
    english_prefixes: tuple
    within_publisher_rank: "int | None"
    take: int
    renumber: bool


LIVE_RULE = LiveRule(
    name="the five nearest English neighbours of the publisher's twenty, "
         "renumbered 1 to 5",
    english_prefixes=("E06", "E07", "E08", "E09"),
    within_publisher_rank=None, take=5, renumber=True)

# ---------------------------------------------------------------------------
# The held state (migrate-legacy preconditions), surveyed 2026-10-11
# ---------------------------------------------------------------------------

# (rows, the survey hash: md5 of lad24cd|statistical_neighbour_lad24cd|
# neighbour_rank|source, NULL as '~', ordered by the key; see live_state)
LEGACY_LIVE = (1480, "23bf09bff006bf75" "6ce1a7eedefe73ef")
LEGACY_LOADED = datetime(2026, 7, 24, 8, 56, 26, 598029, tzinfo=UTC)
LEGACY_SOURCE = ("ONS Clustering similar LAs and statistical nearest "
                 "neighbours, Table 7a (LTLA global), Mar-2026 edition")
# the held file (sha256 split so the credential scan does not flag it); the
# URL is the 2026 edition's download, which the ledger row names so that a
# page load finds the file checked; loaded_by: how the held rows were really
# loaded
LEGACY_FILE = {
    "name": "2026clusteringsimilarlocalauthoritiesandnearestneighbours.xlsx",
    "sha256": "26e9ca037015db12218bb4456344272e"
              "0e4ec373f3e902d69eb89dfb37c972a6",
    "url": (ONS + "/file?uri=" + LANDING + "/2026/2026clusteringsimilar"
            "localauthoritiesandnearestneighbours.xlsx"),
    "loaded_by": ("scripts/s21_statistical_neighbours_build.py run log 66 "
                  "(2026-07-24) from data/raw/ons_statistical_neighbours_"
                  "2026.xlsx, the same bytes as"),
}
# before the editions exist (a preview against live): the held rows come
# from that file, whose Cover says 18 March 2026
LEGACY_RANK = date(2026, 3, 18)


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _json(x):
    return json.loads(x) if isinstance(x, (str, bytes)) else (x or {})


def _norm(s) -> str:
    return " ".join(str("" if s is None else s).split())


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or \
        (isinstance(v, str) and not v.strip())


def rank_text(rank) -> str:
    return f"published {rank.isoformat()}" if rank else "no file rank"


def _date(mm, where) -> date:
    try:
        return datetime.strptime(f"{mm.group(1)} {mm.group(2)} "
                                 f"{mm.group(3)}", "%d %B %Y").date()
    except ValueError:
        raise ValueError(f"{where}: {mm.group(0)!r} is not a date") from None


def _key(r) -> tuple:
    return (r["lad24cd"], r["statistical_neighbour_lad24cd"])


# ---------------------------------------------------------------------------
# Discovery: the ONS dataset landing page (/data JSON)
# ---------------------------------------------------------------------------

EDITION_RE = re.compile(r"[0-9]{4}")


def check_landing(landing_json) -> None:
    """ValueError unless the landing page is still titled LANDING_TITLE."""
    t = _norm((_json(landing_json).get("description") or {}).get("title"))
    if t != LANDING_TITLE:
        raise ValueError(f"the landing page is titled {t!r}, not "
                         f"{LANDING_TITLE!r}; the page has changed")


def _uris(landing_json) -> list:
    return [str(x.get("uri") or "")
            for x in (_json(landing_json).get("datasets") or [])]


def editions_listed(landing_json) -> list:
    """[(edition, uri)] of the landing page's editions, newest first. An
    edition is a URI '<landing>/<yyyy>'. ValueError (listing every URI
    seen) when none fits, or when any URI under the dataset does not fit
    (an unparseable edition is never passed over)."""
    head = LANDING + "/"
    seen = _uris(landing_json)
    out, odd = [], []
    for u in seen:
        seg = u[len(head):] if u.startswith(head) else None
        if seg is not None and EDITION_RE.fullmatch(seg):
            out.append((seg, u))
        else:
            odd.append(u)
    if not out:
        raise ValueError(f"no edition '<landing>/<yyyy>' on the landing page; "
                         f"URIs seen: {seen}")
    out.sort(key=lambda x: x[0], reverse=True)
    if odd:
        raise ValueError(f"the landing page lists a URI that is not an "
                         f"edition '<landing>/<yyyy>': {odd}; the newest "
                         f"edition that fits is {out[0][0]}; not choosing "
                         f"an edition while it is unexplained; URIs seen: "
                         f"{seen}")
    if len({e for e, _ in out}) != len(out):
        raise ValueError(f"an edition is listed twice: {seen}")
    return out


def latest_edition(landing_json) -> tuple:
    """(edition, uri) of the newest edition."""
    return editions_listed(landing_json)[0]


def edition_for_year(landing_json, year) -> tuple:
    """(edition, uri) of the edition labelled `year` (a --file's own Cover
    year); ValueError unless exactly one."""
    hits = [e for e in editions_listed(landing_json) if e[0] == str(year)]
    if len(hits) != 1:
        raise ValueError(f"{len(hits)} editions labelled {year} on the "
                         f"landing page; URIs seen: {_uris(landing_json)}")
    return hits[0]


def edition_download(edition_json) -> tuple:
    """(file, the API's releaseDate) of an edition page: exactly one
    download, an .xlsx file. ValueError naming what was seen. The
    releaseDate corroborates the Cover date; it is never the rank."""
    d = _json(edition_json)
    files = [str(x.get("file") or "") for x in (d.get("downloads") or [])]
    if len(files) != 1 or not files[0]:
        raise ValueError(f"{len(files)} downloads on the edition page "
                         f"({files}); expected exactly one")
    if Path(files[0]).suffix.lower() != ".xlsx":
        raise ValueError(f"the download {files[0]!r} is not an .xlsx file")
    return files[0], str((d.get("description") or {}).get("releaseDate")
                         or "")


def download_url(uri, file) -> str:
    return f"{ONS}/file?uri={uri}/{file}"


def discovery_note(landing_json, held_periods) -> list:
    """Lines naming the newest edition the landing page lists against the
    held releases: held (no newer edition listed), NEWER (read, previewed
    and stored as editions only; never promoted to live), or older than
    what is held. Empty when nothing is held."""
    held = sorted(str(p) for p in held_periods)
    if not held:
        return []
    try:
        eds = editions_listed(landing_json)
    except ValueError:
        return []
    newest = eds[0][0]
    years = {p[:4] for p in held}
    listed = ", ".join(e for e, _ in eds)
    nxt = _norm((_json(landing_json).get("description") or {}).get(
        "nextRelease"))
    tail = (f" (editions listed: {listed}; nextRelease {nxt!r}: the "
            "release is irregular, so no overdue warning is derived)")
    if newest in years:
        return [f"the newest edition the landing page lists is {newest}, "
                f"which is held (release {', '.join(held)}); no newer "
                f"edition is listed" + tail]
    if newest > max(years):
        return [f"NEWER EDITION: the landing page lists edition {newest}, "
                f"newer than every held release ({', '.join(held)}). It is "
                "read and previewed below and load stores it as editions "
                "only (held, not live); it never reaches live through load "
                "or refresh-latest (adopting it is a decision)" + tail]
    return [f"NOTE: the newest edition the landing page lists is {newest}, "
            f"older than the held release(s) {', '.join(held)}; the landing "
            "page may have changed" + tail]


# ---------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------

def code_cell(v, where) -> str:
    """A published ONS area code: nine characters [ENSW] and eight digits,
    exactly as given. Anything else raises ValueError naming `where`: a
    blank, the note 12 marker "[x]" (a missing value), any other text."""
    if _blank(v):
        raise ValueError(f"{where}: a blank cell (every neighbour pair is "
                         "published; a blank halts)")
    if isinstance(v, str) and v.strip() == MARKER:
        raise ValueError(f"{where}: {MARKER} (note 12: a missing value); a "
                         "marker halts and is never stored")
    if not isinstance(v, str) or not CODE_RE.fullmatch(v):
        raise ValueError(f"{where}: {v!r} is not an ONS area code "
                         "(nine characters, [ENSW] and eight digits)")
    return v


def name_cell(v, where) -> str:
    """A published area name: non-blank text other than a marker, as
    given."""
    if _blank(v):
        raise ValueError(f"{where}: a blank cell (every neighbour pair is "
                         "published; a blank halts)")
    if not isinstance(v, str):
        raise ValueError(f"{where}: {v!r} is not a name")
    if v.strip() in (MARKER, "[c]"):
        raise ValueError(f"{where}: {v.strip()} (note 12: a missing or "
                         "suppressed value); a marker halts")
    return v


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------

def docprops_modified(path) -> datetime:
    """The workbook's own docProps modified (core.xml), in UTC. ValueError
    when it is absent or unreadable."""
    p = Path(path)
    where = f"{p.name}: {CORE_XML}"
    try:
        with zipfile.ZipFile(p) as z:
            if CORE_XML not in z.namelist():
                raise ValueError(f"{where} is absent")
            root = ET.fromstring(z.read(CORE_XML))
    except (zipfile.BadZipFile, ET.ParseError, OSError) as e:
        raise ValueError(f"{where} is not readable ({e})") from None
    hits = [el.text for el in root
            if el.tag.rsplit("}", 1)[-1] == "modified"
            and el.text and el.text.strip()]
    if len(hits) != 1:
        raise ValueError(f"{where}: {len(hits)} modified values, expected "
                         "one (docProps modified is part of the file's "
                         "identity)")
    s = hits[0].strip()
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        raise ValueError(f"{where}: docProps modified {s!r} is not a date "
                         "and time") from None
    if d.tzinfo is None:
        raise ValueError(f"{where}: docProps modified {s!r} has no time zone")
    return d.astimezone(UTC)


def _frames(p: Path) -> dict:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 (not an .xlsx workbook)
        raise ValueError(f"{p.name}: not a readable .xlsx workbook "
                         f"({e})") from None
    try:
        names = list(wb.sheetnames)
        need = [s for s in (COVER, CONTENTS, NOTES, SHEET) if s not in names]
        if need:
            raise ValueError(f"{p.name}: no {', '.join(need)} sheet(s); "
                             f"sheets {names}")
        out = {"_sheets": names}
        for s in (COVER, CONTENTS, NOTES, SHEET):
            out[s] = [tuple(r) for r in wb[s].iter_rows(values_only=True)]
    finally:
        wb.close()
    return out


def _texts(rows) -> list:
    return [_norm(v) for r in rows for v in r
            if isinstance(v, str) and v.strip()]


def _cover(rows, name) -> dict:
    where = f"{name}: {COVER}"
    texts = _texts(rows)
    titles = [t for t in texts if t == COVER_TITLE]
    if len(titles) != 1:
        raise ValueError(f"{where}: {len(titles)} cells {COVER_TITLE!r}, "
                         f"expected one; cells seen: {texts[:12]}")
    pubs = [mm for mm in (PUB_RE.fullmatch(t) for t in texts) if mm]
    if len(pubs) != 1:
        raise ValueError(f"{where}: {len(pubs)} cells like 'The data tables "
                         "in this spreadsheet were originally published at "
                         "9.30am <d Month yyyy>.', expected one; cells seen: "
                         f"{texts[:12]}")
    return {"title": titles[0], "pub_text": pubs[0].group(0),
            "published": _date(pubs[0], where)}


def _contents(rows, name) -> str:
    where = f"{name}: {CONTENTS}"
    hits = [r for r in rows if r and _norm(r[0]) == SHEET]
    if len(hits) != 1:
        raise ValueError(f"{where}: {len(hits)} entries for {SHEET!r}, "
                         "expected one")
    t = _norm(hits[0][1] if len(hits[0]) > 1 else None)
    if t != SHEET_TITLE:
        raise ValueError(f"{where}: the entry for {SHEET!r} is {t!r}, not "
                         f"{SHEET_TITLE!r}")
    return t


def _notes(rows, name) -> dict:
    where = f"{name}: {NOTES}"
    found = {}
    for r in rows:
        if r and not _blank(r[0]) and _norm(r[0]) in ("2", "12"):
            n = _norm(r[0])
            if n in found:
                raise ValueError(f"{where}: note {n} appears twice")
            found[n] = r[1] if len(r) > 1 else None
    n2, n12 = found.get("2"), found.get("12")
    if not isinstance(n2, str) or not NOTE2_RE.fullmatch(n2.strip()):
        raise ValueError(f"{where}: note 2 is {n2!r}, not 'Models are based "
                         "on the latest local authority and combined "
                         "authority boundaries, as at <Month yyyy>. ...'")
    if not isinstance(n12, str) or f'"{MARKER}"' not in n12:
        raise ValueError(f"{where}: note 12 is {n12!r}; expected the "
                         f'definition of the "{MARKER}" marker')
    return {"note2": n2, "note12": n12,
            "boundaries": NOTE2_RE.fullmatch(n2.strip()).group(1)}


def _table(rows, name) -> dict:
    where = f"{name}: {SHEET}"
    t = _norm(rows[0][0]) if rows and rows[0] else ""
    if t != SHEET_TITLE:
        raise ValueError(f"{where}: the sheet title is {t!r}, not "
                         f"{SHEET_TITLE!r}")
    heads = [i for i, r in enumerate(rows[:12])
             if r and _norm(r[0]) == HEADER[0]]
    if len(heads) != 1:
        raise ValueError(f"{where}: {len(heads)} header rows starting "
                         f"{HEADER[0]!r} in the first 12 rows")
    h = heads[0]
    header = list(rows[h])
    while header and _blank(header[-1]):
        header.pop()
    if tuple(header) != HEADER:
        unknown = [x for x in header if x not in HEADER]
        missing = [x for x in HEADER if x not in header]
        raise ValueError(f"{where}: the header {header!r} is not exactly "
                         f"{list(HEADER)!r} (unknown {unknown}, missing "
                         f"{missing}); the reading must be corrected "
                         "deliberately, with the evidence")
    subjects, seen = [], {}
    nations = Counter()
    empty = 0  # empty rows seen so far; not a source value
    for n, r in enumerate(rows[h + 1:], start=h + 2):
        if all(_blank(v) for v in r):
            empty += 1
            continue
        rw = f"{where} row {n}"
        if empty:
            raise ValueError(f"{rw}: data after {empty} empty row(s); empty "
                             "rows are allowed only at the end")
        if any(not _blank(v) for v in r[WIDTH:]):
            raise ValueError(f"{rw}: a cell beyond the {WIDTH} headed "
                             "columns")
        r = (list(r) + [None] * WIDTH)[:WIDTH]
        code = code_cell(r[0], f"{rw} {HEADER[0]}")
        subj = name_cell(r[1], f"{rw} {code} {HEADER[1]}")
        if code in seen:
            raise ValueError(f"{rw}: {code} appears twice (also row "
                             f"{seen[code]})")
        seen[code] = n
        nb = []
        for i in range(1, NEIGHBOURS + 1):
            c = code_cell(r[2 * i], f"{rw} {code} {HEADER[2 * i]}")
            nm = name_cell(r[2 * i + 1], f"{rw} {code} {HEADER[2 * i + 1]}")
            nb.append((i, c, nm))
        nations[NATIONS[code[0]]] += 1
        subjects.append({"row": n, "code": code, "name": subj,
                         "neighbours": nb})
    return {"sheet_title": t, "header": header, "subjects": subjects,
            "nations": dict(nations)}


def read_file(path) -> dict:
    """The facts of one clustering workbook: {path, name, sheets, title,
    pub_text, published (the Cover date), release_date (its ISO text, the
    period), rank (= published), year, contents_title, sheet_title, header,
    note2, note12, boundaries, modified (docProps, UTC), subjects [{row,
    code, name, neighbours [(publisher rank, code, name)]}], nations
    {nation: subject rows}}. Raises ValueError naming what was seen: a
    missing sheet; a Cover without exactly one title and one 'originally
    published' line; a Contents entry or sheet title for 7.a other than
    SHEET_TITLE; a header that is not exactly HEADER; notes 2 and 12 not as
    read; docProps modified unreadable; a blank, marker or malformed code
    or name cell; a repeated subject; data after an empty row or beyond the
    headed columns."""
    p = Path(path)
    modified = docprops_modified(p)
    fr = _frames(p)
    cover = _cover(fr[COVER], p.name)
    contents = _contents(fr[CONTENTS], p.name)
    notes = _notes(fr[NOTES], p.name)
    table = _table(fr[SHEET], p.name)
    out = dict(cover)
    out.update(notes)
    out.update(table)
    out.update({"path": p, "name": p.name, "sheets": fr["_sheets"],
                "contents_title": contents, "modified": modified,
                "release_date": cover["published"].isoformat(),
                "rank": cover["published"],
                "year": cover["published"].year})
    return out


# ---------------------------------------------------------------------------
# Records and the live rule
# ---------------------------------------------------------------------------

def _getter(resolve):
    return resolve if callable(resolve) else (lambda c: resolve.get(c))


def _reasons(rec) -> "str | None":
    out = []
    if rec["england_rank"] is None:
        out.append(f"england_rank={NOT_ENGLISH}")
    if rec[RANK] is None:
        out.append(f"{RANK}={OUTSIDE_LIVE}")
    return "; ".join(out) or None


def _is_english(rec, rule=None) -> bool:
    rule = rule or globals()["LIVE_RULE"]
    return str(rec["neighbour_code_published"]).startswith(
        tuple(rule.english_prefixes))


def live_rows(records) -> list:
    """LIVE_RULE applied to a release's editions records: [{lad24cd,
    statistical_neighbour_lad24cd, release_date, neighbour_rank}] sorted by
    subject and rank. Per subject, the neighbours whose published code is
    English (LIVE_RULE.english_prefixes; and within the publisher's top
    LIVE_RULE.within_publisher_rank when set), in publisher rank order; the
    first LIVE_RULE.take, numbered 1.. (renumber) or with the publisher's
    rank. Reads publisher_rank and the codes only (never a stored
    neighbour_rank)."""
    rule = globals()["LIVE_RULE"]
    by = {}
    for r in records:
        by.setdefault(r["lad24cd"], []).append(r)
    out = []
    for subj in sorted(by):
        keep = [r for r in sorted(by[subj], key=lambda x: x["publisher_rank"])
                if _is_english(r, rule)
                and (rule.within_publisher_rank is None
                     or r["publisher_rank"] <= rule.within_publisher_rank)]
        for i, r in enumerate(keep[:rule.take], start=1):
            out.append({"lad24cd": subj,
                        "statistical_neighbour_lad24cd":
                            r["statistical_neighbour_lad24cd"],
                        PERIOD: str(r[PERIOD]),
                        RANK: i if rule.renumber else r["publisher_rank"]})
    return out


def build_records(table, resolve, provenance) -> list:
    """The release's editions records from a read file (read_file), England
    subjects only, 20 per subject, sorted by key: {lad24cd,
    statistical_neighbour_lad24cd (the resolved code of an English
    neighbour, the publisher's code otherwise), release_date,
    neighbour_rank (LIVE_RULE's rank; NULL outside live),
    neighbour_code_published, neighbour_name, publisher_rank, england_rank
    (1..n in publisher order among the English neighbours; NULL for a
    neighbour that is not an English authority), neighbour_nation,
    null_reasons ('col=reason; ...' for each NULL)}. resolve: {publisher
    code: lad24cd} or a callable, asked for English codes only.
    provenance: a mapping (or None) added to every record for the caller;
    it is not stored and not in the content hash. ValueError when an
    English code resolves to nothing, or a subject lists itself or one
    neighbour twice (after resolution)."""
    get = _getter(resolve)
    extra = dict(provenance or {})
    period = table["release_date"]
    out, seen_subj = [], {}
    for s in table["subjects"]:
        if NATIONS[s["code"][0]] != "England":
            continue
        where = f"{SHEET} row {s['row']} {s['code']}"
        lad = get(s["code"])
        if not lad:
            raise ValueError(f"{where}: {s['code']} resolves to no lad24cd")
        if lad in seen_subj:
            raise ValueError(f"{where}: subject {lad} appears twice (also "
                             f"{seen_subj[lad]}); refusing to choose")
        seen_subj[lad] = s["code"]
        recs, seen = [], {}
        eng = 0  # English neighbours counted so far # not a source value
        for rank, code, nm in s["neighbours"]:
            nation = NATIONS[code[0]]
            if nation == "England":
                nb = get(code)
                if not nb:
                    raise ValueError(f"{where} neighbour {rank}: {code} "
                                     "resolves to no lad24cd")
                eng += 1
                er = eng
            else:
                nb, er = code, None
            if nb == lad:
                raise ValueError(f"{where}: the subject lists itself as "
                                 f"neighbour {rank} ({code})")
            if nb in seen:
                raise ValueError(f"{where}: neighbour {nb} listed twice "
                                 f"(ranks {seen[nb]} and {rank})")
            seen[nb] = rank
            rec = {"lad24cd": lad, "statistical_neighbour_lad24cd": nb,
                   PERIOD: period, RANK: None,
                   "neighbour_code_published": code, "neighbour_name": nm,
                   "publisher_rank": rank, "england_rank": er,
                   "neighbour_nation": nation}
            rec.update(extra)
            recs.append(rec)
        ranks = {_key(x): x[RANK] for x in live_rows(recs)}
        for rec in recs:
            rec[RANK] = ranks.get(_key(rec))
            rec["null_reasons"] = _reasons(rec)
        out.extend(recs)
    out.sort(key=_key)
    return out


def rule_counts(records) -> dict:
    """What LIVE_RULE does to a release (printed and checked against the
    spec): live pairs and subjects; pairs whose live rank differs from the
    publisher's and the subjects they belong to; subjects whose publisher
    top five include a neighbour that is not English; live pairs from
    beyond the publisher's top LIVE_RULE.take and the furthest publisher
    rank in live; subjects with fewer live rows than LIVE_RULE.take."""
    rule = globals()["LIVE_RULE"]
    pub = {_key(r): r["publisher_rank"] for r in records}
    live = live_rows(records)
    differ = [x for x in live if x[RANK] != pub[_key(x)]]
    beyond = [x for x in live if pub[_key(x)] > rule.take]
    per = Counter(x["lad24cd"] for x in live)
    subjects = {r["lad24cd"] for r in records}
    top = {r["lad24cd"] for r in records
           if r["publisher_rank"] <= rule.take and not _is_english(r, rule)}
    return {"live_pairs": len(live), "subjects": len(per),
            "rank_differs": len(differ),
            "subjects_affected": len({x["lad24cd"] for x in differ}),
            "subjects_top_non_english": len(top),
            "from_beyond_top": len(beyond),
            "furthest": max((pub[_key(x)] for x in live), default=None),
            "subjects_short": sorted(s for s in subjects
                                     if per.get(s) != rule.take)}


def rule_text(c) -> str:
    return (f"LIVE_RULE ({LIVE_RULE.name}): {c['live_pairs']:,} live pairs "
            f"for {c['subjects']} subjects; {c['rank_differs']} pairs whose "
            f"live rank differs from the publisher's rank, in "
            f"{c['subjects_affected']} subjects ({c['subjects_top_non_english']}"
            f" subjects' publisher top five include a neighbour that is not "
            f"English); {c['from_beyond_top']} pairs from publisher ranks "
            f"beyond {LIVE_RULE.take} (the furthest {c['furthest']}); "
            f"subjects with fewer than {LIVE_RULE.take}: "
            f"{c['subjects_short'] or 'none'}")


def rows_content_sha(records) -> str:
    """sha256 of a release's content: one line per record sorted by key,
    lad24cd|statistical_neighbour_lad24cd|the content columns (CONTENT
    order: neighbour_rank and the editions-only columns; NULL as the empty
    string), joined by LF, UTF-8. Provenance and where the file was read
    are not in it, so they never make a new edition."""
    def cell(v):
        return "" if v is None else str(v)
    lines = sorted("|".join([cell(r[k]) for k in KEY]
                            + [cell(r.get(c)) for c in CONTENT])
                   for r in records)
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Geography
# ---------------------------------------------------------------------------

def english_codes(table) -> set:
    """Every English code the file gives: the English subjects and their
    English neighbours (the codes that are resolved)."""
    out = set()
    for s in table["subjects"]:
        if NATIONS[s["code"][0]] != "England":
            continue
        out.add(s["code"])
        out |= {c for _, c, _ in s["neighbours"] if c[0] == "E"}
    return out


def resolve_codes(cur, codes, period) -> tuple:
    """({publisher code: lad24cd}, problems). geography.resolve(cur, '21',
    codes, by_period={period: codes}) (source 21 is declared 'new': an
    E08000016 or E08000019 is a problem); then each code's canonical form
    must be in la_boundaries, else 'UNEXPLAINED'. Only English codes are
    given; other nations' codes are stored verbatim and never resolved."""
    codes = set(codes)
    recode, problems = geography.resolve(cur, RUN_SOURCE, codes,
                                         by_period={period: codes})
    problems = list(problems)
    cur.execute("SELECT lad24cd FROM public.la_boundaries")
    valid = {r[0] for r in cur.fetchall()}
    out, hit = {}, {}
    for c in sorted(codes):
        r = geography.canonical(c, recode)
        if r in valid:
            out[c] = r
            hit.setdefault(r, []).append(c)
        else:
            problems.append(f"UNEXPLAINED {c}: not in la_boundaries and not a "
                            "la_code_lookup recode")
    two = {k: v for k, v in hit.items() if len(v) > 1}
    if two:
        problems.append(f"{period}: codes {two} reach one lad24cd; refusing "
                        "to choose a winner")
    return out, problems


# ---------------------------------------------------------------------------
# Ranks, ledger and planning
# ---------------------------------------------------------------------------

_ISO = r"[0-9]{4}-[0-9]{2}-[0-9]{2}"
_RANK_RE = re.compile(rf"\bpublished ({_ISO})\b")
_LEDGER_RE = re.compile(rf"(.*) \(release ({_ISO}); published ({_ISO})\)")


def release_rank(file_or_source) -> "date | None":
    """The Cover date of a file: a read file's own rank, or the 'published
    <yyyy-mm-dd>' an edition's or ledger row's source_file carries; None if
    it names none (a file name, URL or the API's releaseDate is never read
    for a rank)."""
    x = file_or_source
    if isinstance(x, dict):
        return x.get("rank")
    mm = _RANK_RE.search(str(x or ""))
    return date.fromisoformat(mm.group(1)) if mm else None


def edition_source(file, src, note="") -> str:
    """An edition's source_file: the file name, its release and Cover date
    (the rank), where it was read, and a note."""
    return (f"{src['file_name']}; release {file['release_date']}; "
            f"{rank_text(file['rank'])}; {src['where']}"
            + (f"; {note}" if note else ""))


def ledger_source(where, rank, period) -> str:
    """A file's ledger source_file: where it was read (the final URL, or the
    path of a --file), the release it states and its Cover date."""
    return f"{where} (release {period}; published {rank.isoformat()})"


def parse_ledger_source(s) -> "tuple | None":
    """(where, rank, release) of a ledger source_file, or None."""
    mm = _LEDGER_RE.fullmatch(str(s or ""))
    if not mm:
        return None
    return mm.group(1), date.fromisoformat(mm.group(3)), mm.group(2)


def content_sha256(path) -> str:
    """SHA-256 of the file's bytes."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def tip_ranks(tips, checked) -> dict:
    """{held release: the newest file rank that stated it}: the tip
    edition's, or a later file the ledger records for it."""
    out = {}
    for p, t in tips.items():
        ranks = [t.get("rank")] + [release_rank(s)
                                   for s, _ in checked.get(p, ())]
        ranks = [r for r in ranks if r is not None]
        out[p] = max(ranks) if ranks else None
    return out


def ledger_complete(checked, where, sha, ranks, *, allow_older=False,
                    stranded=()) -> "list | None":
    """[release] when the ledger records this file's (where, sha256) pair
    for the release its ledger source declares, else None (the file is read
    and planned). A release whose tip comes from a newer file is read again
    (and then planned as older) unless allow_older; a stranded release
    (editions, live empty) always makes the file be read. So a run that
    failed or was rejected leaves no ledger row and the rerun reads the
    file again."""
    have, declared = set(), None
    for p, pairs in checked.items():
        for s, h in pairs:
            parsed = parse_ledger_source(s)
            if h == sha and parsed and parsed[0] == where:
                have.add(str(p))
                declared = parsed
    if declared is None:
        return None
    _, rank, period = declared
    if period not in have or period in {str(x) for x in stranded}:
        return None
    if not allow_older and ranks.get(period) is not None \
            and ranks[period] > rank:
        return None
    return [period]


def plan_periods(periods, rank, ranks, checked, source, sha, *, recheck,
                 allow_older) -> tuple:
    """(new, revised, skipped) for the file's release. ranks: {held
    release: tip_ranks}; checked: {release: {(source, sha)}} of the ledger
    (stranded releases left out by the caller). new: not held (unless it is
    before the newest held release: skipped as older, unless allow_older);
    revised: held, compared with the tip; skipped {release: reason}: held
    with a tip from a newer file (older, unless allow_older), or already in
    the ledger (checked, unless --recheck names it). --recheck must be a
    held release of the file (ValueError)."""
    periods = sorted(str(p) for p in periods)
    if recheck is not None and (recheck not in periods or recheck not in ranks):
        raise ValueError(f"--recheck {recheck}: not a held release this file "
                         f"states ({', '.join(periods)})")
    newest_held = max(ranks) if ranks else None
    new, revised, skipped = [], [], {}
    for p in periods:
        if p not in ranks:
            if newest_held is not None and p < newest_held and not allow_older:
                skipped[p] = (f"older: release {p} is before the newest held "
                              f"release {newest_held}")
                continue
            new.append(p)
            continue
        t = ranks[p]
        if t is not None and t > rank and not allow_older:
            skipped[p] = (f"older: its tip comes from a file {rank_text(t)}, "
                          f"this file is {rank_text(rank)}")
            continue
        if recheck != p and (source, sha) in checked.get(p, ()):
            skipped[p] = "checked: this file is already in the ledger"
            continue
        revised.append(p)
    return new, revised, skipped


# ---------------------------------------------------------------------------
# Stop conditions
# ---------------------------------------------------------------------------

def _by_key(records) -> dict:
    return {_key(r): r for r in records or ()}


def _subjects(records) -> set:
    return {r["lad24cd"] for r in records or ()}


def _changes(new, tip) -> list:
    """Lines naming what differs between two releases' records: pairs added
    and removed, and per pair the content columns that differ."""
    nk, tk = _by_key(new), _by_key(tip)
    added = sorted(set(nk) - set(tk))
    removed = sorted(set(tk) - set(nk))
    changed = [(k, c, tk[k].get(c), nk[k].get(c))
               for k in sorted(set(nk) & set(tk)) for c in CONTENT
               if nk[k].get(c) != tk[k].get(c)]
    out = []
    if added:
        out.append(f"{len(added)} pair(s) added: "
                   + ", ".join("/".join(k) for k in added[:6]))
    if removed:
        out.append(f"{len(removed)} pair(s) removed: "
                   + ", ".join("/".join(k) for k in removed[:6]))
    if changed:
        out.append(f"{len(changed)} cell(s) changed: "
                   + ", ".join(f"{'/'.join(k)} {c} {a!r}->{b!r}"
                               for k, c, a, b in changed[:6]))
    return out


def period_problems(new, tip, *, kind, period=None, live_empty=False) -> list:
    """The stop conditions of the file's release (empty = fine). Problems
    starting with PARTIAL are released by nothing: fewer English subjects
    than EXPECTED_SUBJECTS (both kinds), or (revised) a subject of the held
    edition missing. Revised (a reissue of a held release): any pair added
    or removed or any content cell changed (publisher rank, name, nation,
    the live rank) against the held edition, released by --acknowledge.
    Any kind: a release other than LIVE_RELEASE into an empty live table
    (the live CHECK holds LIVE_RELEASE only; no flag releases it)."""
    g = globals()
    out = []
    want = g["EXPECTED_SUBJECTS"]
    n = len(_subjects(new))
    if n < want:
        out.append(f"{PARTIAL}: {n} English subjects, expected {want} (a "
                   "partial file never replaces a fuller edition)")
    if kind == "revised":
        gone = sorted(_subjects(tip) - _subjects(new))
        if gone:
            out.append(f"{PARTIAL}: {len(gone)} subject(s) of the held "
                       "edition missing from the file: "
                       + ", ".join(gone[:8]))
        ch = _changes(new, tip)
        if ch:
            out.append("the held edition is revised (the comparator sets "
                       "may change): " + "; ".join(ch))
    if live_empty and period is not None and str(period) != LIVE_RELEASE:
        out.append(f"{period} into an empty live table: the live table holds "
                   f"{LIVE_RELEASE} only (its CHECK), and another release "
                   "never enters live through a load (adopting one is a "
                   f"decision); load {LIVE_RELEASE} first, then {period} is "
                   "stored as editions only")
    return out


def _differs(new, old) -> bool:
    return bool(_changes(new, old))


# ---------------------------------------------------------------------------
# Spec and profile
# ---------------------------------------------------------------------------

SPEC = core.EditionSpec(
    name="s21",
    live_table="la_statistical_neighbours",
    editions_table="la_statistical_neighbours_editions",
    key_cols=KEY,
    period_col=PERIOD,
    value_cols=((RANK, "integer CHECK (neighbour_rank BETWEEN 1 AND 20)"),),
    extra_cols=EXTRA_TYPES,
    refresh_cols=(RANK, "source", "loaded_at"),
    refresh_from=(("source", "release_label"), ("loaded_at", "loaded_at")),
    as_loaded_source_col="source",
    refresh_source_whole_period=True,   # one release label per release
    key_types=(("lad24cd", "varchar(9) NOT NULL"),
               ("statistical_neighbour_lad24cd", "varchar(9) NOT NULL"),
               (PERIOD, "date NOT NULL")),
    table_constraints=(
        "CHECK ((england_rank IS NULL) = (neighbour_nation <> 'England'))",
        "CHECK (neighbour_nation = 'England' OR "
        "statistical_neighbour_lad24cd = neighbour_code_published)",
        "CHECK (neighbour_rank IS NULL OR england_rank IS NOT NULL)",
        "CHECK ((england_rank IS NULL OR neighbour_rank IS NULL) = "
        "(null_reasons IS NOT NULL))",
        "CHECK (statistical_neighbour_lad24cd <> lad24cd)"),
    fk_la_boundaries=True,              # on the subject only
    live_single_period=True,            # one release in live
    refresh_key_changes=False,
)


def live_alter_sql(spec) -> str:
    """The one live DDL change: the release_date column (additive,
    default-filled, not in the key; the CHECK keeps one release in live)."""
    return (f"ALTER TABLE public.{core._ident(spec.live_table)} ADD COLUMN "
            f"{PERIOD} date NOT NULL DEFAULT '{LIVE_RELEASE}' CHECK "
            f"({PERIOD} = '{LIVE_RELEASE}')")


def _no_label(fetched_on) -> str:
    raise ValueError("an S21 release label comes from the file; store "
                     "through apply_release(...)")


def check_records(records, period) -> None:
    """ValueError unless the release's records are whole: each of this
    release, one per key, every content column present, every NULL with
    its reason, 20 per subject."""
    seen, per = set(), Counter()
    for r in records:
        if str(r.get(PERIOD)) != period:
            raise ValueError(f"{period}: a record of release "
                             f"{r.get(PERIOD)!r}")
        k = _key(r)
        if k in seen:
            raise ValueError(f"{period}: {k} appears twice")
        seen.add(k)
        missing = [c for c in CONTENT if c not in r]
        if missing:
            raise ValueError(f"{period}: {k} lacks {missing}")
        if _reasons(r) != r["null_reasons"]:
            raise ValueError(f"{period}: {k} null_reasons "
                             f"{r['null_reasons']!r}, expected {_reasons(r)!r}")
        per[r["lad24cd"]] += 1
    bad = {s: n for s, n in per.items() if n != NEIGHBOURS}
    if bad:
        raise ValueError(f"{period}: subjects without {NEIGHBOURS} rows: "
                         f"{dict(list(bad.items())[:6])}")


PROFILE = pe.Profile(
    spec=SPEC, value_cols=CONTENT, run_agent=RUN_AGENT, run_source=RUN_SOURCE,
    heading="S21 ONS statistical nearest neighbours (la_statistical_"
            "neighbours)",
    default_source_file="ONS clustering similar local authorities and "
                        "statistical nearest neighbours",
    expected_areas=None, release_label=_no_label,
    content_sha256=rows_content_sha, check_records=check_records,
    file_checks=True, savepoint="s21_period",
    example_label="(neighbour_rank, code, name, publisher_rank, "
                  "england_rank, nation, null_reasons)")


def profile() -> pe.Profile:
    """PROFILE bound to the current SPEC (looked up when called, so tests
    can swap it)."""
    return dataclasses.replace(PROFILE, spec=globals()["SPEC"])


def _late(name: str):
    """A call through this module's attribute `name`, looked up when
    called (so mock.patch.object(s21_neighbours_editions, name, ...)
    reaches it)."""
    return lambda *a, **kw: globals()[name](*a, **kw)


table_exists = pe.table_exists


def _conn(writing):
    return pe.connect(writing)


def live_has_release(cur, spec) -> bool:
    cur.execute("SELECT 1 FROM information_schema.columns WHERE "
                "table_schema = 'public' AND table_name = %s AND "
                "column_name = %s", (spec.live_table, PERIOD))
    return cur.fetchone() is not None


def ensure_live_release(cur, spec) -> str:
    """Add the live release_date column (live_alter_sql) when it is absent
    and return 'added'; 'present' when it is there exactly as planned
    (date, NOT NULL, DEFAULT LIVE_RELEASE, the CHECK); halt when a
    release_date column of any other shape exists."""
    cur.execute("SELECT data_type, is_nullable, column_default FROM "
                "information_schema.columns WHERE table_schema = 'public' "
                "AND table_name = %s AND column_name = %s",
                (spec.live_table, PERIOD))
    row = cur.fetchone()
    if row is None:
        cur.execute(live_alter_sql(spec))
        return "added"
    cur.execute("SELECT pg_get_constraintdef(oid) FROM pg_constraint WHERE "
                "conrelid = %s::regclass AND contype = 'c'",
                (f"public.{spec.live_table}",))
    defs = [r[0] for r in cur.fetchall()]
    want = ("date", "NO", f"'{LIVE_RELEASE}'::date")
    check = f"({PERIOD} = '{LIVE_RELEASE}'::date)"
    if tuple(row) != want or not any(check in d for d in defs):
        halt(f"{spec.live_table}.{PERIOD} exists but is not the planned "
             f"column (type, nullable, default {tuple(row)}; CHECK "
             f"constraints {defs}); expected: {live_alter_sql(spec)}")
    return "present"


def create_all(cur) -> None:
    """ddl: the editions table, its append-only triggers, the file-check
    ledger and the live release_date column (live_alter_sql, when absent).
    Idempotent; no value is written; no other live column or row is
    touched."""
    prof = profile()
    core.create_schema(cur, prof.spec)
    pe.create_file_checks(cur, prof)
    ensure_live_release(cur, prof.spec)


# ---------------------------------------------------------------------------
# Reading held data
# ---------------------------------------------------------------------------

def records(cur, spec, period, edition) -> list:
    """The edition's rows as records (the key, the release as a string,
    neighbour_rank and the editions-only columns)."""
    names = KEY + CONTENT
    cur.execute(f"SELECT {', '.join(names)} FROM "
                f"public.{spec.editions_table} WHERE {spec.period_col} = %s "
                "AND edition = %s", (period, edition))
    out = [dict(zip(names, r)) for r in cur.fetchall()]
    for r in out:
        r[PERIOD] = str(period)
    return out


def live_pairs(cur, spec, period=None) -> dict:
    """{(subject, neighbour): neighbour_rank} of the live table (of
    `period` when the release_date column exists and a period is given)."""
    where, args = "", None
    if period is not None and live_has_release(cur, spec):
        where, args = f" WHERE {PERIOD} = %s", (period,)
    cur.execute(f"SELECT lad24cd, statistical_neighbour_lad24cd, {RANK} "
                f"FROM public.{spec.live_table}" + where, args)
    return {(a, b): r for a, b, r in cur.fetchall()}


def _want(recs) -> dict:
    return {_key(x): x[RANK] for x in live_rows(recs)}


def _stored(recs) -> dict:
    return {_key(x): x[RANK] for x in recs if x[RANK] is not None}


def tip_info(cur, spec) -> dict:
    """{release: {edition, source_file, rank}} of every release with
    editions."""
    cur.execute(f"SELECT DISTINCT {spec.period_col} FROM "
                f"public.{spec.editions_table}")
    out = {}
    for (p,) in cur.fetchall():
        tip = core.latest_edition(cur, spec, str(p))
        cur.execute(f"SELECT DISTINCT source_file FROM "
                    f"public.{spec.editions_table} WHERE {spec.period_col} = "
                    "%s AND edition = %s", (p, tip))
        src = {r[0] for r in cur.fetchall()}
        sf = src.pop() if len(src) == 1 else None
        out[pe._p(p)] = {"edition": tip, "source_file": sf,
                         "rank": release_rank(sf)}
    return out


def _live_rows_n(cur, spec) -> int:
    cur.execute(f"SELECT COUNT(*) FROM public.{spec.live_table}")
    return cur.fetchone()[0]


def legacy_tips(cur, spec) -> dict:
    """Before the editions table exists: {LIVE_RELEASE: {rank}} when live
    holds rows (the rank of the held file, LEGACY_RANK), else {}."""
    if not _live_rows_n(cur, spec):
        return {}
    return {LIVE_RELEASE: {"edition": None, "source_file": None,
                           "rank": LEGACY_RANK}}


def ledger_checks(cur, prof) -> dict:
    """{release: {(source_file, file_sha256)}} of the ledger."""
    pc = prof.spec.period_col
    cur.execute(f"SELECT DISTINCT {pc}, source_file, file_sha256 "
                f"FROM public.{pe.file_checks_table(prof)}")
    out = {}
    for p, f, s in cur.fetchall():
        out.setdefault(pe._p(p), set()).add((f, s))
    return out


def _readiness(cur, prof) -> dict:
    spec = prof.spec
    lt = pe.file_checks_table(prof)
    return {spec.editions_table: table_exists(cur, spec.editions_table),
            lt: table_exists(cur, lt),
            f"{spec.live_table}.{PERIOD}": live_has_release(cur, spec)}


# ---------------------------------------------------------------------------
# Live: insert, compare, status, refresh (live is a subset of the editions)
# ---------------------------------------------------------------------------

def insert_live(cur, profile, period: str, records: list, *, source,
                loaded_at=None) -> int:
    """LIVE_RULE's rows of the release (live_rows) into live: the key, the
    release, neighbour_rank and source; loaded_at takes its default unless
    given (a refresh copies the edition's). Returns the rows inserted."""
    from psycopg2.extras import execute_values
    spec = profile.spec
    rows = live_rows(records)
    cols = KEY + (PERIOD, RANK, "source") + (("loaded_at",) if loaded_at
                                              else ())
    execute_values(
        cur, f"INSERT INTO public.{spec.live_table} ({', '.join(cols)}) "
        "VALUES %s",
        [(x["lad24cd"], x["statistical_neighbour_lad24cd"], period, x[RANK],
          source) + ((loaded_at,) if loaded_at else ()) for x in rows],
        page_size=1000)
    return len(rows)


def live_equals_tip(cur, spec, periods) -> list:
    """Problems (empty = fine): for each release, the live rows equal
    live_rows(the latest edition) on the key and neighbour_rank, both ways,
    and every live row carries the release."""
    bad = []
    for p in periods:
        p = pe._p(p)
        try:
            ed = core.chain_tip(cur, spec, p)
        except (LookupError, ValueError) as e:
            bad.append(f"{p}: {e}")
            continue
        want = _want(records(cur, spec, p, ed))
        got = live_pairs(cur, spec, p)
        only_ed = sorted(k for k in want if got.get(k, "absent") != want[k])
        only_live = sorted(k for k in got if want.get(k, "absent") != got[k])
        if only_ed or only_live:
            bad.append(f"{p} ed{ed}: {len(only_ed)} LIVE_RULE row(s) not in "
                       f"live, {len(only_live)} live row(s) not LIVE_RULE's "
                       f"(e.g. {(only_ed or only_live)[:3]})")
    return bad


def classify_live(cur, spec, period, tip) -> tuple:
    """(kind, edition) of the live release against its editions: 'current'
    (live equals LIVE_RULE applied to the tip); 'rule-changed' (live equals
    the tip's stored live ranks, built under another LIVE_RULE: a refresh
    applies the current rule); 'pending' (live equals an earlier edition's,
    under the current or its stored rule: a newer edition is not yet live);
    'drift' (live equals none: changed outside the editions machinery)."""
    got = live_pairs(cur, spec, period)
    recs = records(cur, spec, period, tip)
    if got == _want(recs):
        return "current", tip
    if got == _stored(recs):
        return "rule-changed", tip
    cur.execute(f"SELECT DISTINCT edition FROM public.{spec.editions_table} "
                f"WHERE {spec.period_col} = %s AND edition <> %s "
                "ORDER BY 1 DESC", (period, tip))
    for (e,) in cur.fetchall():
        r = records(cur, spec, period, e)
        if got in (_want(r), _stored(r)):
            return "pending", e
    return "drift", None


def status(cur, prof) -> dict:
    """What needs action, in the keys period_editions.format_status reads
    (read-only). For the live release: pending_refresh (a newer edition, or
    a LIVE_RULE change, not yet in live), drift_periods (live equals no
    stored edition under any rule), bad_counts (live rows other than
    LIVE_RULE's count of the tip); new_periods (live with no editions),
    chain_errors; live_missing (editions, live empty); held_not_live
    (another release held as editions only: information, not a fault).
    ok is true only when none needs action."""
    spec = prof.spec
    lm = core.latest_map(cur, spec)
    tips, new, errors = lm
    pending, drift, bad_counts, notes = [], [], {}, {}
    for p, tip in tips.items():
        kind, ed = classify_live(cur, spec, p, tip)
        if kind in ("pending", "rule-changed"):
            pending.append(pe._p(p))
            notes[pe._p(p)] = (kind, ed)
        elif kind == "drift":
            drift.append(pe._p(p))
        want = len(live_rows(records(cur, spec, p, tip)))
        n = len(live_pairs(cur, spec, p))
        if n != want:
            bad_counts[pe._p(p)] = (n, want)
    missing = [] if tips or new else [
        pe._p(x) for x in pe.live_missing_periods(cur, prof)]
    held = [pe._p(x) for x in lm.held_not_live] if (tips or new) else []
    ok = not (new or pending or drift or errors or bad_counts or missing)
    return {"new_periods": [pe._p(x) for x in new],
            "pending_refresh": sorted(pending),
            "drift_periods": sorted(drift),
            "chain_errors": {pe._p(k): v for k, v in errors.items()},
            "bad_counts": bad_counts, "live_missing": missing,
            "held_not_live": held, "periods": len(tips) + len(new),
            "ok": ok, "notes": notes}


def status_lines(cur, prof, st) -> list:
    """Per release: the tip edition, its Cover date and the ledger's latest
    file; the live rule; why a release waits for refresh-latest."""
    s = prof.spec
    lt = pe.file_checks_table(prof)
    lines = [f"  live rule: {LIVE_RULE.name}"]
    for p, t in sorted(tip_info(cur, s).items()):
        line = f"  {p}: tip edition {t['edition']} ({rank_text(t['rank'])})"
        cur.execute(f"SELECT source_file, outcome FROM public.{lt} WHERE "
                    f"{s.period_col} = %s ORDER BY id DESC LIMIT 1", (p,))
        last = cur.fetchone()
        line += (f"; ledger latest: {last[1]} {last[0]}" if last
                 else "; no ledger row")
        lines.append(line)
    for p, (kind, ed) in sorted(st.get("notes", {}).items()):
        lines.append(f"  {p}: " + (
            "live was built under another LIVE_RULE (it equals the tip's "
            "stored live ranks); refresh-latest --acknowledge applies the "
            "current rule" if kind == "rule-changed" else
            f"live equals edition {ed}; refresh-latest --acknowledge applies "
            "the latest"))
    return lines


def _live_diff(got, want) -> dict:
    added = sorted(k for k in want if k not in got)
    removed = sorted(k for k in got if k not in want)
    moved = sorted((k, got[k], want[k]) for k in want
                   if k in got and got[k] != want[k])
    return {"added": added, "removed": removed, "moved": moved}


def _diff_lines(d) -> list:
    out = []
    for what in ("added", "removed"):
        if d[what]:
            out.append(f"{len(d[what])} live pair(s) {what}: "
                       + ", ".join("/".join(k) for k in d[what][:8]))
    if d["moved"]:
        out.append(f"{len(d['moved'])} rank(s) changed: "
                   + ", ".join(f"{'/'.join(k)} {a}->{b}"
                               for k, a, b in d["moved"][:8]))
    return out


def _tip_meta(cur, spec, period, ed) -> tuple:
    cur.execute(f"SELECT DISTINCT release_label, loaded_at FROM "
                f"public.{spec.editions_table} WHERE {spec.period_col} = %s "
                "AND edition = %s", (period, ed))
    rows = cur.fetchall()
    if len(rows) != 1:
        halt(f"{period}: edition {ed} carries {len(rows)} (release_label, "
             "loaded_at) values; expected one")
    return rows[0]


def refresh_plan(cur, prof) -> "dict | None":
    """{period, edition, kind, diff, rows} of the refresh the live release
    needs, or None when live is current (or empty). Halts on drift (no flag
    overrides it: live equals no stored edition under any rule)."""
    spec = prof.spec
    tips, new, errors = core.latest_map(cur, spec)
    if new:
        halt(f"live releases with no editions {[pe._p(x) for x in new]}; run "
             "migrate-legacy first")
    if errors:
        halt("invalid edition chain: " + "; ".join(
            f"{pe._p(p)}: {m}" for p, m in errors.items()))
    for p, tip in tips.items():
        kind, _ = classify_live(cur, spec, p, tip)
        if kind == "current":
            return None
        if kind == "drift":
            halt(f"{pe._p(p)}: live differs from LIVE_RULE applied to the "
                 "latest edition and equals no stored edition (drift: "
                 "changed outside the editions tables); nothing is "
                 "refreshed. Find out how, then restore it deliberately")
        recs = records(cur, spec, p, tip)
        d = _live_diff(live_pairs(cur, spec, p), _want(recs))
        return {"period": pe._p(p), "edition": tip, "kind": kind, "diff": d,
                "records": recs, "rows": len(live_rows(recs))}
    return None


def refresh_latest(cur, prof, *, acknowledge=()) -> dict:
    """Replace the live release's rows with LIVE_RULE applied to its latest
    edition, source the edition's release label and loaded_at the
    edition's, in a savepoint of its own: halts (rolled back) unless the
    release is named in acknowledge when any live pair or rank changes, or
    if afterwards live differs from live_rows(tip), the editions table
    changed, or the row counts are not the plan's. {period, edition,
    deleted, inserted, diff} or {} when nothing is to be done. Another
    release named in acknowledge halts (refuse_non_live)."""
    spec = prof.spec
    core.refuse_non_live(cur, spec, tuple(acknowledge), "--acknowledge")
    plan = refresh_plan(cur, prof)
    if plan is None:
        return {}
    p, ed = plan["period"], plan["edition"]
    lines = _diff_lines(plan["diff"])
    if lines and p not in set(acknowledge):
        halt(f"refresh-latest {p}: the live rows change ("
             + "; ".join(lines) + f"); the comparator sets change. Read the "
             f"preview, then re-run with --acknowledge {p}")
    label, loaded = _tip_meta(cur, spec, p, ed)
    with core._own_savepoint(cur, "s21_refresh"):
        ed_before = core.period_hashes(cur, spec, "editions")
        cur.execute(f"DELETE FROM public.{spec.live_table} WHERE "
                    f"{spec.period_col} = %s", (p,))
        deleted = cur.rowcount
        inserted = insert_live(cur, prof, p, plan["records"], source=label,
                               loaded_at=loaded)
        bad = live_equals_tip(cur, spec, [p])
        if core.period_hashes(cur, spec, "editions") != ed_before:
            bad.append("the editions table changed during the refresh")
        if inserted != plan["rows"]:
            bad.append(f"inserted {inserted} rows, expected {plan['rows']}")
        cur.execute(f"SELECT COUNT(DISTINCT source), COUNT(DISTINCT "
                    f"loaded_at), COUNT(DISTINCT {PERIOD}) FROM "
                    f"public.{spec.live_table}")
        if tuple(cur.fetchone()) != (1, 1, 1):
            bad.append("live source, loaded_at or release_date is not "
                       "uniform after the refresh")
        if bad:
            halt("refresh-latest failed its checks, rolled back: "
                 + "; ".join(bad[:6]))
    return {"period": p, "edition": ed, "deleted": deleted,
            "inserted": inserted, "diff": plan["diff"], "kind": plan["kind"]}


# ---------------------------------------------------------------------------
# Storing a release
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class PeriodInfo:
    """What a release's edition records: source_file, the release label and
    the file read ((ledger source, sha256),)."""
    source_file: str
    label: str
    files: tuple


def apply_release(cur, profile, period, records, *, fetched_on, info) -> str:
    """Store one release inside the caller's savepoint, then its ledger
    row; returns the kind. new: edition 1 and, into an empty live table,
    LIVE_RULE's live rows (live_equals_tip checked); a new release while
    live holds another is HELD_NOT_LIVE (edition 1 only). revised: the next
    edition superseding the tip (live waits for refresh-latest).
    LIVE_MISSING (editions, live empty): LIVE_RULE's rows of the tip with
    its label. unchanged: the ledger row only. Never commits."""
    spec = profile.spec
    kind = pe.classify_period(cur, profile, period, records)
    if kind == LIVE_MISSING:
        if pe.live_holds_other(cur, profile, period):
            halt(f"{period}: live holds a different release; a "
                 "live_single_period table never takes a second one")
        tip = core.chain_tip(cur, spec, period)
        label, _ = _tip_meta(cur, spec, period, tip)
        insert_live(cur, profile, period, _tip_records(cur, spec, period,
                                                        tip), source=label)
    elif kind in ("new", "revised", HELD_NOT_LIVE):
        tip = None if kind != "revised" else core.chain_tip(cur, spec, period)
        ed = core.insert_edition(
            cur, spec, records, period, release_label=info.label,
            published_date=fetched_on, source_file=info.source_file,
            source_sha256=rows_content_sha(records), supersedes=tip,
            strict=True, allow_revert=True)
        bad = load_checks.check_coverage(cur, spec, period, ed, len(records))
        if core.chain_tip(cur, spec, period) != ed:
            bad.append(f"{period}: edition {ed} is not the chain tip")
        if bad:
            halt(f"{period} edition {ed} failed its checks: " + "; ".join(bad))
        if kind == "new" and pe.live_holds_other(cur, profile, period):
            kind = HELD_NOT_LIVE
        if kind == "new":
            insert_live(cur, profile, period, records, source=info.label)
    if kind in ("new", LIVE_MISSING):
        bad = live_equals_tip(cur, spec, [period])
        if bad:
            halt(f"{period}: the live rows failed their checks, rolled back: "
                 + "; ".join(bad))
    tip = core.chain_tip(cur, spec, period)
    for src, sha in info.files:
        pe.record_file_check(cur, profile, period, src, sha, kind, tip)
    return kind


def _tip_records(cur, spec, period, tip) -> list:
    return records(cur, spec, period, tip)


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

def _session(session):
    if session is not None:
        return session
    import requests
    return requests


def _get(session, url, **kw):
    try:
        r = session.get(url, headers={"User-Agent": USER_AGENT},
                        timeout=300, **kw)
        r.raise_for_status()
    except Exception as e:  # noqa: BLE001 (any failure halts)
        halt(f"GET {url} failed, nothing downloaded: {e}")
    return r


def fetch_json(path, session=None) -> dict:
    """The ONS /data JSON of a page path (a read-only GET)."""
    return json.loads(_get(_session(session), ONS + path + "/data").text)


def _url_name(url) -> str:
    """The file name a URL names: the 'uri' query's last part (ONS
    /file?uri=...), else the path's."""
    u = urlparse(url)
    q = parse_qs(u.query).get("uri")
    return Path(q[0] if q else u.path).name


def _opens(path) -> None:
    """Raise unless the file is an .xlsx workbook with the Cover sheet and
    7.a sheets."""
    with zipfile.ZipFile(path) as z:
        wbx = z.read("xl/workbook.xml")
    for s in (COVER, SHEET):
        if f'name="{s}"'.encode() not in wbx:
            raise ValueError(f"no {s} sheet")


def fetch(url, dest, session=None) -> tuple:
    """Download url into the directory dest (redirects followed) and return
    (path to read, final URL). The file name is the final URL's. Checked
    before it is kept: more than MIN_FILE_BYTES, an .xlsx name, and it
    opens as a workbook with the Cover sheet and 7.a sheets. A same-named
    file with the same content is kept as it is; one with different content
    is never replaced: the download is saved beside it as
    <stem>-<sha8><suffix> and that is the one read. An existing name is
    never a reason to skip the download."""
    dest = Path(dest)
    r = _get(_session(session), url)
    final = str(getattr(r, "url", None) or url)
    if final != url:
        print(f"redirected: {url} -> {final}")
    body = r.content
    if len(body) <= MIN_FILE_BYTES:
        halt(f"{final}: {len(body)} bytes, expected more than "
             f"{MIN_FILE_BYTES}; nothing written; saw {body[:100]!r}")
    name = _url_name(final)
    suffix = Path(name).suffix.lower()
    if suffix != ".xlsx":
        halt(f"{final}: not an .xlsx file; nothing written")
    dest.mkdir(parents=True, exist_ok=True)
    target = dest / name
    new_sha = hashlib.sha256(body).hexdigest()
    tmp = dest / (Path(name).stem + ".download.tmp" + suffix)
    tmp.write_bytes(body)
    try:
        _opens(tmp)
    except Exception as e:  # noqa: BLE001
        tmp.unlink()
        halt(f"{final}: does not open as the clustering workbook ({e}); "
             "nothing kept")
    path = target
    if target.exists():
        old_sha = content_sha256(target)
        if old_sha == new_sha:
            tmp.unlink()
            print(f"{name} already held with the same content (sha256 "
                  f"{new_sha[:16]}); kept as it is")
        else:
            path = dest / f"{target.stem}-{new_sha[:8]}{target.suffix}"
            if path.exists() and content_sha256(path) == new_sha:
                tmp.unlink()
                print(f"NOTE: {name} differs from this download; "
                      f"{path.name} already holds it (sha256 "
                      f"{new_sha[:16]})")
            else:
                tmp.replace(path)
                print(f"NOTE: {name} already exists with different content "
                      f"(sha256 {old_sha[:16]}, new {new_sha[:16]}); kept it "
                      f"untouched and saved the new download as {path.name}")
    else:
        tmp.replace(target)
        print(f"downloaded {name}: {len(body):,} bytes, sha256 "
              f"{new_sha[:16]}")
    return path, final


def _file_source(path: Path) -> str:
    """Where a --file was read: the path under the repo (data/raw/...),
    else the absolute path."""
    p = Path(path).resolve()
    try:
        return p.relative_to(REPO).as_posix()
    except ValueError:
        return p.as_posix()


# ---------------------------------------------------------------------------
# ddl, status, run log, refresh-latest
# ---------------------------------------------------------------------------

def log_run(cur, rows_written: int, notes: str, started_at=None) -> None:
    """The pipeline_run_log row for a committed run (agent RUN_AGENT, source
    '21'). Called only on committed load, refresh-latest, migrate-legacy
    and restore-edition runs, partial runs included."""
    pe.log_run(cur, profile(), rows_written, notes, started_at)


def cmd_ddl(args) -> int:
    """ddl [--commit | --simulate]: the editions table and its triggers, the
    ledger, and the live release_date column (the preview prints the exact
    ALTER)."""
    prof = profile()
    spec = prof.spec
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            for t in (spec.editions_table, pe.file_checks_table(prof)):
                print(f"{t}: "
                      f"{'exists' if table_exists(cur, t) else 'does not exist'}")
            if live_has_release(cur, spec):
                print(f"{spec.live_table}: release_date column present "
                      "(checked against the planned definition on --commit)")
            else:
                print(f"{spec.live_table}: no release_date column; ddl runs:"
                      f"\n  {live_alter_sql(spec)}\n  (additive: a "
                      "default-filled column, not in the key; no other "
                      "column or row changes)")
            if not writing:
                conn.rollback()
                print("DRY RUN: nothing written (use --commit or --simulate)")
                return 0
            create_all(cur)
        pe.finish(conn, args, f"ddl: {spec.editions_table} and triggers "
                              f"{spec.trigger}, {spec.truncate_trigger}, "
                              f"{pe.file_checks_table(prof)} and "
                              f"{spec.live_table}.{PERIOD} present")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def cmd_status(args) -> int:
    """status: what needs action; exit 1 if anything, or (with a clean
    message, nothing created) if the editions table, ledger or live
    release_date column does not exist."""
    prof = profile()
    conn = _conn(False)
    try:
        with conn.cursor() as cur:
            missing = [t for t, ok in _readiness(cur, prof).items() if not ok]
            if missing:
                print(f"status: {' and '.join(missing)} "
                      f"{'does' if len(missing) == 1 else 'do'} not exist "
                      "yet; run `ddl --commit`, then `migrate-legacy FILE "
                      "--commit` (S21 has no sync-new)")
                return 1
            st = status(cur, prof)
            lines = status_lines(cur, prof, st)
    finally:
        conn.rollback()
        conn.close()
    print(pe.format_status(prof, st, lines))
    return 0 if st["ok"] else 1


def _tables_exist(cur, prof) -> None:
    for t, ok in _readiness(cur, prof).items():
        if not ok:
            halt(f"{t} does not exist yet; run `ddl --commit`, then "
                 "`migrate-legacy FILE --commit`")


def cmd_refresh_latest(args) -> int:
    """refresh-latest [--acknowledge RELEASE] [--commit | --simulate]: live
    replaced by LIVE_RULE applied to the live release's latest edition
    (source the edition's label, loaded_at the edition's), when it differs;
    a change of any live pair or rank needs --acknowledge RELEASE. Another
    release named halts; drift halts."""
    prof = profile()
    spec = prof.spec
    writing = args.commit or args.simulate
    ack = tuple(args.acknowledge or ())
    started = datetime.now(timezone.utc)
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            _tables_exist(cur, prof)
            core.refuse_non_live(cur, spec, ack, "--acknowledge")
            plan = refresh_plan(cur, prof)
            if plan is None:
                print("rows refresh-latest would write: none (live equals "
                      "LIVE_RULE applied to the latest edition)")
                conn.rollback()
                return 0
            lines = _diff_lines(plan["diff"])
            print(f"rows refresh-latest would write: {plan['period']}="
                  f"{plan['rows']} (live replaced by LIVE_RULE applied to "
                  f"edition {plan['edition']}; {plan['kind']})")
            for line in lines:
                print(f"  {line}")
            if lines and plan["period"] not in ack:
                print(f"  the comparator sets change: --acknowledge "
                      f"{plan['period']} is needed to apply it")
            if not writing:
                conn.rollback()
                print("DRY RUN: nothing written (use --commit or --simulate)")
                return 0
            res = refresh_latest(cur, prof, acknowledge=ack)
            if args.commit:
                log_run(cur, res["inserted"], (
                    f"refresh-latest: {res['period']} live replaced by "
                    f"LIVE_RULE ({LIVE_RULE.name}) applied to edition "
                    f"{res['edition']} ({res['kind']}): {res['deleted']} "
                    f"rows deleted, {res['inserted']} inserted; "
                    + "; ".join(_diff_lines(res["diff"]))
                    + (f"; ACKNOWLEDGED {res['period']} (--acknowledge)"
                       if res["period"] in ack else "")), started)
        pe.finish(conn, args, f"{res['inserted']} live rows of "
                              f"{res['period']} replaced from edition "
                              f"{res['edition']}; checks passed")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# load
# ---------------------------------------------------------------------------

def _end(conn, args, writing, rc) -> int:
    conn.rollback()  # commit mode has committed already
    if not writing:
        print("PREVIEW: nothing written to the database (use --commit or "
              "--simulate)")
    elif args.simulate:
        print("SIMULATION: ROLLED BACK (nothing persisted)")
    return rc


def _period_arg(text) -> str:
    try:
        return date.fromisoformat(str(text)).isoformat()
    except ValueError:
        raise argparse.ArgumentTypeError(f"not a release date YYYY-MM-DD: "
                                         f"{text!r}") from None


def _read(path) -> dict:
    try:
        return read_file(path)
    except ValueError as e:      # the message names the file already
        halt(f"{e}; nothing stored")


def identity_problems(file, edition=None, api_date=None) -> list:
    """Problems (empty = fine) between a read file and the edition page it
    came from: the edition label's year must be the Cover's, and the API's
    releaseDate the Cover date (or, as a UTC time, the evening before: a
    midnight in British Summer Time)."""
    out = []
    if edition is not None and str(edition) != str(file["year"]):
        out.append(f"the edition is {edition}, the file's Cover says "
                   f"{file['published']:%d %B %Y}")
    if api_date:
        try:
            when = datetime.fromisoformat(str(api_date).replace("Z",
                                                                "+00:00"))
        except ValueError:
            out.append(f"the edition page's releaseDate {api_date!r} is not "
                       "a date and time")
        else:
            days = {when.date(), (when + timedelta(hours=1)).date()}
            if file["published"] not in days:
                out.append(f"the edition page's releaseDate {api_date} is "
                           "not the file's Cover date "
                           f"{file['published'].isoformat()}")
    return out


def _landing():
    lp = fetch_json(LANDING)
    try:
        check_landing(lp)
    except ValueError as e:
        halt(f"{ONS}{LANDING}: {e}")
    return lp


def _edition(uri, edition) -> tuple:
    """(file, API releaseDate) of the edition page, its label checked."""
    ed = fetch_json(uri)
    try:
        file, api_date = edition_download(ed)
    except ValueError as e:
        halt(f"{ONS}{uri}: {e}")
    label = _norm((_json(ed).get("description") or {}).get("edition"))
    if label != str(edition):
        halt(f"{ONS}{uri}: the edition label {label!r} is not {edition!r}")
    print(f"edition: {label!r} ({ONS}{uri}); download {file}; API releaseDate "
          f"{api_date} (corroborates the Cover date; never the rank)")
    return file, api_date


def _source(args) -> dict:
    """The file the run reads: {path, where, how, edition, api_date,
    page_url, file_name, file, discovered}. --file reads a local file under
    FILE_ROOTS (manual_input.require_under; nothing downloaded; identity
    from the file) and finds its own year's edition on the landing page
    unless --no-page; otherwise the landing page gives the newest edition
    and its workbook is downloaded to RAW_DIR (also in a preview)."""
    out = {"path": None, "where": None, "how": None, "edition": None,
           "api_date": None, "page_url": None, "file_name": None,
           "file": None, "discovered": None}
    if args.file:
        p = manual_input.require_under(args.file, globals()["FILE_ROOTS"])
        where = _file_source(p)
        manual_input.announce_manual("S21 ONS statistical nearest neighbours "
                                     "(local workbook)", where,
                                     manual_input.file_identity(p))
        f = _read(p)
        out.update(path=p, where=where, how="local file", file_name=p.name,
                   file=f)
        if args.no_page:
            out.update(page_url=NO_PAGE, how="local file, --no-page: the "
                       "landing page was not read")
            print("NOTE: --no-page: the landing page is not read")
            return out
        lp = _landing()
        try:
            edition, uri = edition_for_year(lp, f["year"])
        except ValueError as e:
            halt(f"{ONS}{LANDING}: {e}; the edition of the file's own year "
                 f"({f['year']}) was not found. Give --no-page to load the "
                 "file without it (recorded as such)")
        _, api_date = _edition(uri, edition)
        out.update(edition=edition, api_date=api_date, page_url=ONS + uri,
                   how=f"local file; edition {edition} on the landing page")
        return out
    lp = _landing()
    try:
        eds = editions_listed(lp)
    except ValueError as e:
        halt(f"{ONS}{LANDING}: {e}")
    edition, uri = eds[0]
    print(f"landing page: {ONS}{LANDING}; editions listed: "
          + ", ".join(e for e, _ in eds) + f"; the newest is {edition}")
    file, api_date = _edition(uri, edition)
    url = download_url(uri, file)
    print(f"downloading {url} to {RAW_DIR} (also in a preview; the database "
          "is not written)")
    path, final = fetch(url, RAW_DIR)
    print(f"  final URL: {final}")
    out.update(path=path, where=final, how=f"newest edition {edition} (ONS "
               "landing page)", edition=edition, api_date=api_date,
               page_url=ONS + uri, file_name=_url_name(final),
               discovered=lp)
    return out


def _label(file, src, sha, how, *, notes=()) -> str:
    return (f"ONS {file['title']}, {SHEET_TITLE}"
            + (f", edition {src['edition']}" if src.get("edition") else "")
            + f"; published {file['published'].isoformat()}; "
            f"{src['file_name']}; {src['where']} sha256 {sha[:16]}; {how}; "
            f"live rule: {LIVE_RULE.name}"
            + "".join(f"; {n}" for n in notes if n))


def _report(file) -> None:
    n = file["nations"]
    print(f"{file['name']}: {file['title']!r}; {file['pub_text']!r} (the "
          f"release date and the rank); {SHEET} {file['sheet_title']!r} "
          f"(Contents agrees); header exact ({WIDTH} columns); docProps "
          f"modified {file['modified']:%Y-%m-%dT%H:%M:%SZ}")
    print("  subject rows by nation: "
          + ", ".join(f"{k} {v}" for k, v in sorted(n.items()))
          + f" ({sum(n.values())}); {NEIGHBOURS} neighbour pairs each, no "
          "blank or marker")
    print(f"  note 2 (boundaries {file['boundaries']}): {file['note2']!r}")
    print(f"  note 12: {file['note12']!r}")


def _build(cur, file) -> list:
    """The release's records, codes resolved; halts on a geography or cell
    problem."""
    codes = english_codes(file)
    try:
        rmap, problems = resolve_codes(cur, codes, file["release_date"])
    except (ValueError, KeyError) as e:
        halt(f"geography: {e}; nothing stored")
    if problems:
        halt("geography check failed, nothing stored: " + "; ".join(problems))
    recoded = sorted((c, r) for c, r in rmap.items() if r != c)
    if recoded:
        print("codes resolved (geography.resolve, la_code_lookup recode): "
              + ", ".join(f"{c} -> {r}" for c, r in recoded))
    try:
        recs = build_records(file, rmap, {"file_name": file["name"]})
    except ValueError as e:
        halt(f"{e}; nothing stored")
    print(rule_text(rule_counts(recs)))
    return recs


def _preview_live(cur, spec, period, recs) -> None:
    """Before the editions table exists: LIVE_RULE's rows of the file
    against the live rows (key and neighbour_rank)."""
    if period != LIVE_RELEASE:
        print(f"  {period}: would be held as editions only (live holds "
              f"{LIVE_RELEASE}; compared with nothing)")
        return
    held = live_pairs(cur, spec)
    if not held:
        print(f"  {period}: new, {len(recs):,} editions rows (live is empty)")
        return
    lines = _diff_lines(_live_diff(held, _want(recs)))
    print(f"  {period}: {'revised' if lines else 'unchanged'}, LIVE_RULE's "
          f"{len(live_rows(recs)):,} rows of the file against the "
          f"{len(held):,} live rows (compared with live)"
          + ("; " + "; ".join(lines) if lines else ""))


def _store(cur, prof, period, recs, info, args, fetched_on) -> tuple:
    """Compare (and, on --commit or --simulate, store) the release in its
    own savepoint: committed on --commit, always rolled back otherwise.
    (rc, kind): rc 1 when it failed (nothing stored for it)."""
    sp = prof.savepoint
    write = args.commit or args.simulate
    conn = cur.connection
    in_sp = False
    try:
        check_records(recs, period)
        cur.execute(f"SAVEPOINT {sp}")
        in_sp = True
        cmp = pe.compare_period(cur, prof, period, recs)
        kind = cmp["kind"]
        if write:
            kind = _late("apply_release")(cur, prof, period, recs,
                                          fetched_on=fetched_on, info=info)
        if args.commit:
            cur.execute(f"RELEASE SAVEPOINT {sp}")
            in_sp = False
            conn.commit()
        else:
            cur.execute(f"ROLLBACK TO SAVEPOINT {sp}")
            cur.execute(f"RELEASE SAVEPOINT {sp}")
            in_sp = False
    except (Exception, SystemExit) as e:
        if in_sp:
            cur.execute(f"ROLLBACK TO SAVEPOINT {sp}")
            cur.execute(f"RELEASE SAVEPOINT {sp}")
        if args.commit:
            conn.rollback()
        msg = e.code if isinstance(e, SystemExit) else f"{type(e).__name__}: {e}"
        print(f"  {period}: FAILED, nothing stored for this release ({msg})")
        return 1, None
    n_ed, n_live = f"{len(recs):,}", f"{len(live_rows(recs)):,}"
    done = " COMMITTED" if args.commit else ""
    if kind == "new":
        act = (("stored edition 1 (" if write else "would store edition 1 (")
               + f"{n_ed} rows) and " + ("inserted " if write else
                                         "would insert ")
               + f"{n_live} live rows (LIVE_RULE)")
    elif kind == HELD_NOT_LIVE:
        act = (("stored edition 1 only" if write else
                "would store edition 1 only") + f" ({n_ed} rows); live holds "
               "a different release, not promoted (no live rows)")
    elif kind == "revised":
        act = (("stored" if write else "would store") + " the next edition "
               f"({n_ed} rows); live waits for refresh-latest")
    elif kind == LIVE_MISSING:
        act = ("editions release missing from live: " + ("inserted "
               if write else "would insert ") + f"{n_live} live rows (no new "
               "edition)")
    else:
        act = "nothing to store"
    print(f"  {period}: {kind}, {cmp['changed']} rows changed (compared with "
          f"{cmp['against']})"
          + ("; " + "; ".join(cmp["examples"]) if cmp["examples"] else "")
          + f"; {act}{done}")
    return 0, kind


def cmd_load(args) -> int:
    if args.no_page and not args.file:
        halt("--no-page is only for --file (a local file whose landing page "
             "is not read)")
    prof = profile()
    spec = prof.spec
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            ready = _readiness(cur, prof)
            has_ed = all(ready.values())
            missing = [t for t, ok in ready.items() if not ok]
            if writing and not has_ed:
                halt(f"missing {', '.join(missing)}; run `ddl --commit` and "
                     "`migrate-legacy FILE --commit` first")
            live_empty = _live_rows_n(cur, spec) == 0
            if has_ed:
                _, new_live, errors = core.latest_map(cur, spec)
                if errors:
                    halt("invalid edition chain: " + "; ".join(
                        f"{pe._p(p)}: {msg}" for p, msg in errors.items()))
                if new_live:
                    halt(f"{spec.live_table}: live releases with no editions "
                         f"{[pe._p(x) for x in new_live]}; run migrate-legacy "
                         "first")
                tips = tip_info(cur, spec)
                checked = ledger_checks(cur, prof)
                stranded = (pe.live_missing_periods(cur, prof) if live_empty
                            else [])
            else:
                print(f"NOTE: {', '.join(missing)} not present: no editions "
                      "table yet; this preview compares the file with the "
                      "LIVE table")
                tips, checked, stranded = legacy_tips(cur, spec), {}, []
            ranks = tip_ranks(tips, checked)
            held = sorted(ranks)
            print("held releases: " + (", ".join(held) or "none"))
            src = _source(args)
            if src["discovered"] is not None:
                for line in discovery_note(src["discovered"], held):
                    print(line)
            path = src["path"]
            sha = content_sha256(path)
            print(f"{path.name}: sha256 {sha[:16]}"
                  + ("; byte-identical to the held file"
                     if sha == globals()["LEGACY_FILE"]["sha256"] else ""))
            f = src["file"]
            if has_ed and args.recheck is None:
                done = ledger_complete(checked, src["where"], sha, ranks,
                                       allow_older=args.allow_older_file,
                                       stranded=stranded)
                if done:
                    print(f"the file's (source, sha256) is already in the "
                          f"ledger for {', '.join(done)}; "
                          + ("nothing compared" if f else "nothing parsed")
                          + " (--recheck YYYY-MM-DD reads it again)")
                    print("nothing to do")
                    return _end(conn, args, writing, rc=0)
            if f is None:
                f = _read(path)
            bad = identity_problems(f, src["edition"], src["api_date"])
            if bad:
                halt("file identity check failed, nothing stored: "
                     + "; ".join(bad))
            _report(f)
            recs = _build(cur, f)
            period = f["release_date"]
            if held and period not in held and period > max(held):
                print(f"NOTE: {period} is newer than every held release "
                      f"({', '.join(held)}); it is stored as editions only "
                      "(held, not live) and never reaches live through load "
                      "or refresh-latest")
            lsrc = ledger_source(src["where"], f["rank"], period)
            chk = {p: s for p, s in checked.items() if p not in stranded}
            try:
                new, revised, skipped = plan_periods(
                    [period], f["rank"], ranks, chk, lsrc, sha,
                    recheck=args.recheck, allow_older=args.allow_older_file)
            except ValueError as e:
                halt(str(e))
            rc = _run(args, conn, cur, prof, has_ed, src, f, recs, new,
                      revised, skipped, tips, ranks, held, sha, lsrc,
                      live_empty)
        return _end(conn, args, writing, rc)
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def _run(args, conn, cur, prof, has_ed, src, f, recs, new, revised, skipped,
         tips, ranks, held, sha, lsrc, live_empty) -> int:
    started = datetime.now(timezone.utc)
    spec = prof.spec
    rank = f["rank"]
    older_note = {}
    for p in revised + new:
        t = ranks.get(p)
        if t is not None and t > rank:
            older_note[p] = (f"--allow-older-file given; {p}'s tip comes from "
                             f"a file {rank_text(t)}, this file is "
                             f"{rank_text(rank)}")
        elif p in new and held and p < max(held):
            older_note[p] = (f"--allow-older-file given; release {p} is "
                             f"before the newest held release {max(held)}")
        if p in older_note:
            print(f"NOTE: {older_note[p]}")
    print(f"planned: new {new or 'none'}; compare {revised or 'none'}; "
          "skipped " + (", ".join(f"{p} ({r})" for p, r in sorted(
              skipped.items())) or "none"))
    periods = sorted(new + revised)
    ack = set(args.acknowledge or ())
    reissue = set(args.accept_reissue or ())
    stray = sorted((ack | reissue) - set(periods))
    if stray:
        halt(f"--acknowledge/--accept-reissue {stray}: not releases this run "
             f"compares or stores ({periods or 'none'}); nothing stored")
    if not periods:
        if skipped and all(r.startswith("older") for r in skipped.values()):
            halt(f"older file: the file is release {f['release_date']}, "
                 f"{rank_text(rank)}, older than what is held "
                 f"({', '.join(held)}); storing it would record an older "
                 "file. If this is deliberate, re-run with "
                 "--allow-older-file")
        print("nothing to do")
        return 0
    problems, infos, acked, reissued = {}, {}, {}, {}
    for p in periods:
        kind = "revised" if p in revised else "new"
        tip = (records(cur, spec, p, tips[p]["edition"])
               if kind == "revised" and has_ed else None)
        if kind == "revised" and tip is None:
            # before the editions exist the file is compared with live in
            # the preview (_preview_live); only the file's own checks here
            found = period_problems(recs, None, kind="new", period=p,
                                    live_empty=live_empty)
        else:
            found = period_problems(recs, tip, kind=kind, period=p,
                                    live_empty=live_empty)
        hard = [x for x in found if x.startswith(PARTIAL)
                or "empty live table" in x]
        soft = [x for x in found if x not in hard]
        bad = list(hard)
        if soft and p in ack:
            acked[p] = "; ".join(soft)
            print(f"  {p}: ACKNOWLEDGED (--acknowledge): {acked[p][:400]}")
        elif soft:
            bad += [x + "; " + ACK_HINT.format(p=p) for x in soft]
        if tip is not None and ranks.get(p) == rank and _differs(recs, tip):
            if p in reissue:
                reissued[p] = (f"--accept-reissue {p}: the file is "
                               f"{rank_text(rank)}, the same as the file of "
                               "the held tip, with different content (a "
                               "publisher reissue keeping its date)")
                print(f"  {p}: ACCEPTED REISSUE: {reissued[p]}")
            else:
                bad.append(f"the file is {rank_text(rank)}, the same as the "
                           "file of the held tip, but its content differs: "
                           "two files claim the same release. A publisher "
                           "reissue keeps its date: read the release's notes, "
                           f"then --accept-reissue {p}")
        if bad:
            problems[p] = bad
            continue
        notes = (older_note.get(p, ""), reissued.get(p, ""),
                 f"ACKNOWLEDGED {acked[p]}" if p in acked else "")
        how = (("new release" if kind == "new" else "restated release")
               + f" ({src['how']})")
        infos[p] = PeriodInfo(
            edition_source(f, src, "; ".join(n for n in notes if n)),
            _label(f, src, sha, how, notes=notes), ((lsrc, sha),))
    for p, msgs in sorted(problems.items()):
        for msg in msgs:
            print(f"  {p}: REJECTED, not stored: STOP CONDITION: {msg}")
    ok = [p for p in periods if p not in problems]
    rc = 0
    stats = {}
    if ok and has_ed:
        for p in ok:
            r, kind = _store(cur, prof, p, recs, infos[p], args, rank)
            if r:
                rc = 1
                break
            stats[p] = kind
    elif ok:
        for p in ok:
            _preview_live(cur, spec, p, recs)
        print("  (no editions table yet: preview against live only)")
    if problems:
        print(f"REJECTED {len(problems)} release(s), nothing stored for them: "
              f"{', '.join(sorted(problems))}; exit 1")
        rc = 1
    if args.commit:
        stored = live = 0  # rows written by this run; not a source value
        for p, k in stats.items():
            if k in ("new", "revised", HELD_NOT_LIVE):
                stored += len(recs)
            if k in ("new", LIVE_MISSING):
                live += len(live_rows(recs))
        left = [p for p in ok if p not in stats]
        partial = ""
        if rc:
            partial = (f"PARTIAL RUN (exit 1): rejected "
                       f"{sorted(problems) or 'none'}"
                       + "".join(f" ({p}: {'; '.join(problems[p])[:300]})"
                                 for p in sorted(problems))
                       + f"; failed {left or 'none'}; a release that failed "
                       "or was rejected stored nothing and is not in the "
                       "counts below; ")
        notes = (partial + f"clustering workbook {src['where']} sha256 "
                 f"{sha[:16]}, release {f['release_date']}, {rank_text(rank)} "
                 f"({src['how']}): "
                 + ("; ".join(f"{p} {k}" for p, k in stats.items())
                    or "nothing stored")
                 + (f"; skipped {sorted(skipped)}" if skipped else "")
                 + "".join(f"; {older_note[p]}" for p in sorted(older_note)
                           if p in stats)
                 + "".join(f"; {reissued[p]}" for p in sorted(reissued)
                           if p in stats)
                 + "".join(f"; ACKNOWLEDGED {p}: {acked[p][:300]}"
                           for p in sorted(acked) if p in stats)
                 + f"; live rule: {LIVE_RULE.name}"
                 + f". Rows stored: {stored} edition rows, {live} live rows.")
        log_run(cur, stored + live, notes, started)
        conn.commit()
        print("pipeline_run_log row written"
              + (" (partial run: see its notes)" if rc else ""))
    return rc


# ---------------------------------------------------------------------------
# One-off: migrate the held table
# ---------------------------------------------------------------------------

SURVEY_COLS = KEY + (RANK, "source")


def live_state(cur, spec=None) -> tuple:
    """(rows, survey hash, distinct loaded_at, distinct source, releases) of
    the live table. The survey hash is md5 of lad24cd|statistical_
    neighbour_lad24cd|neighbour_rank|source (NULL as '~'), rows ordered by
    the key, joined by LF (release_date is left out, so the hash is the
    same before and after ddl). releases: the distinct live release dates
    (strings), or None when the column does not exist yet."""
    spec = spec or globals()["SPEC"]
    expr = "||'|'||".join(f"coalesce({c}::text,'~')" for c in SURVEY_COLS)
    cur.execute(f"SELECT COUNT(*), md5(string_agg({expr}, E'\\n' ORDER BY "
                f"{', '.join(KEY)})), COUNT(DISTINCT loaded_at), "
                f"COUNT(DISTINCT source) FROM public.{spec.live_table}")
    n, h, k, s = cur.fetchone()
    rel = None
    if live_has_release(cur, spec):
        cur.execute(f"SELECT DISTINCT {PERIOD}::text FROM "
                    f"public.{spec.live_table} ORDER BY 1")
        rel = [r[0] for r in cur.fetchall()]
    return n, h, k, s, rel


def migrate_legacy(cur, file, *, write, legacy=None, files=None) -> dict:
    """One-off. Preconditions: on write the editions table, ledger and live
    release_date column exist (ddl first) and the tables are empty (in a
    preview they may be absent; if present they must be empty); live as
    surveyed (legacy: (rows, survey hash), default LEGACY_LIVE; one
    loaded_at; one source; release_date absent or LIVE_RELEASE on every
    row); the file's sha256 as files['sha256'] (default LEGACY_FILE). Proof,
    any failure halts: the file's identity (read_file), its geography, and
    LIVE_RULE applied to its records reproduces every live pair and rank
    (both ways). Edition 1 'as loaded' is the file's records (every English
    subject's 20 neighbours; the engine's as-loaded label, published_date
    the load date, source_file 'as loaded: <how it was loaded> <name>;
    release <date>; published <date>'), and a ledger row for the file
    (outcome 'unchanged', edition 1, files['url']) so the next page load
    finds it checked. Live untouched (checked equal to LIVE_RULE applied to
    edition 1 inside the transaction). Nothing commits here; the caller
    rolls everything back on a halt."""
    g = globals()
    legacy = tuple(legacy or g["LEGACY_LIVE"])
    files = files or g["LEGACY_FILE"]
    prof = profile()
    spec = prof.spec
    ready = _readiness(cur, prof)
    if write and not all(ready.values()):
        halt(f"missing {[t for t, ok in ready.items() if not ok]}; run "
             "`ddl --commit` first")
    for t in (spec.editions_table, pe.file_checks_table(prof)):
        if ready[t]:
            cur.execute(f"SELECT COUNT(*) FROM public.{t}")
            n = cur.fetchone()[0]
            if n:
                halt(f"{t} already holds {n} rows; migrate-legacy runs once, "
                     "on empty tables")
    n, h, k, s, rel = live_state(cur, spec)
    if (n, h) != legacy or k != 1 or s != 1 \
            or rel not in (None, [LIVE_RELEASE]):
        halt(f"{spec.live_table} is not as surveyed: {n} rows, hash {h}, {k} "
             f"loaded_at value(s), {s} source value(s), releases {rel}; "
             f"expected {legacy[0]} rows, hash {legacy[1]}, one loaded_at, "
             f"one source, release {LIVE_RELEASE} (or no release_date column "
             "yet); nothing stored")
    path = Path(file)
    sha = content_sha256(path)
    if sha != files["sha256"]:
        halt(f"{path.name}: sha256 {sha[:16]}, expected "
             f"{files['sha256'][:16]} (the held file); nothing stored")
    f = _read(path)
    _report(f)
    period = f["release_date"]
    if period != LIVE_RELEASE:
        halt(f"{path.name} is release {period}; the live table holds "
             f"{LIVE_RELEASE}; nothing stored")
    recs = _build(cur, f)
    bad = period_problems(recs, None, kind="new")
    if bad:
        halt("; ".join(bad) + "; nothing stored")
    held = live_pairs(cur, spec)
    diffs = _diff_lines(_live_diff(held, _want(recs)))
    if diffs:
        halt("proof failed: LIVE_RULE applied to the held file read again "
             "differs from the held live rows: " + "; ".join(diffs)
             + "; nothing stored")
    cur.execute(f"SELECT MIN((loaded_at AT TIME ZONE 'UTC')::date), "
                f"MIN(source) FROM public.{spec.live_table}")
    loaded, source = cur.fetchone()
    c = rule_counts(recs)
    plan = {"period": period, "rows": len(recs), "live": len(held),
            "records": recs, "sha": sha, "loaded": loaded, "counts": c,
            "content_sha": rows_content_sha(recs),
            "source_file": (f"as loaded: {files['loaded_by']} {files['name']}; "
                            f"release {period}; {rank_text(f['rank'])}"),
            "ledger": ledger_source(files["url"], f["rank"], period)}
    print(f"proof: LIVE_RULE ({LIVE_RULE.name}) applied to the held file "
          f"read again reproduces every held live row: {len(held):,} pairs "
          f"and ranks for {len({a for a, _ in held})} subjects, 0 "
          "differences either way")
    print(f"  {c['rank_differs']} pairs whose live rank differs from the "
          f"publisher's rank, in {c['subjects_affected']} subjects; "
          f"{c['from_beyond_top']} pairs from publisher ranks "
          f"{LIVE_RULE.take + 1} and beyond (the furthest {c['furthest']})")
    print(f"  live source (kept as held): {source!r}")
    print(f"  {period}: edition 1 as loaded, {len(recs):,} rows (every "
          f"English subject's {NEIGHBOURS} neighbours), published {loaded} "
          f"(the load date); source_file {plan['source_file']!r}; content "
          f"sha256 {plan['content_sha']}; ledger {plan['ledger']}")
    if not write:
        return plan
    cur.execute(f"SAVEPOINT {prof.savepoint}")
    ed = core.insert_edition(
        cur, spec, recs, period, release_label=core.AS_LOADED_LABEL,
        published_date=loaded, source_file=plan["source_file"],
        source_sha256=plan["content_sha"], supersedes=None, strict=True)
    bad = load_checks.check_coverage(cur, spec, period, ed, len(recs))
    bad += live_equals_tip(cur, spec, [period])
    if ed != 1 or bad:
        halt(f"{spec.editions_table} {period}: edition {ed} failed its "
             "checks, rolled back: " + "; ".join(bad))
    pe.record_file_check(cur, prof, period, plan["ledger"], sha,
                         "unchanged", 1)
    cur.execute(f"RELEASE SAVEPOINT {prof.savepoint}")
    return plan


def cmd_migrate_legacy(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    try:
        with conn.cursor() as cur:
            plan = migrate_legacy(cur, args.file, write=writing)
            if args.commit:
                c = plan["counts"]
                log_run(cur, plan["rows"], (
                    f"migrate-legacy: edition 1 as loaded for "
                    f"{plan['period']} ({plan['rows']} rows, every English "
                    f"subject's {NEIGHBOURS} neighbours, from the held file "
                    f"sha256 {plan['sha'][:16]}); proof: LIVE_RULE "
                    f"({LIVE_RULE.name}) reproduces the {plan['live']} live "
                    f"rows, 0 differences ({c['rank_differs']} pairs whose "
                    f"rank differs from the publisher's, "
                    f"{c['subjects_affected']} subjects, "
                    f"{c['from_beyond_top']} pairs from publisher ranks "
                    f"beyond {LIVE_RULE.take}, furthest {c['furthest']}); "
                    "ledger row for the file. Live untouched."), started)
                conn.commit()
                print("migrate-legacy: COMMITTED; pipeline_run_log row "
                      "written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# restore-edition
# ---------------------------------------------------------------------------

def restore_edition(cur, period, edition) -> int:
    """Store edition `edition` of `period` as the next edition ('restored
    from edition N') for refresh-latest to apply; return its number. Halts
    if the edition does not exist or is the tip."""
    spec = globals()["SPEC"]
    p = str(period)
    tip = core.latest_edition(cur, spec, p)
    recs = records(cur, spec, p, edition)
    if not recs:
        halt(f"{spec.editions_table} {p}: no edition {edition}")
    if edition == tip:
        halt(f"{p}: edition {edition} is the tip; nothing to restore")
    cur.execute(f"SELECT DISTINCT source_file FROM public.{spec.editions_table}"
                f" WHERE {spec.period_col} = %s AND edition = %s", (p, edition))
    src = cur.fetchone()[0]
    new = core.insert_edition(
        cur, spec, recs, p, release_label=f"restored from edition {edition}",
        published_date=date.today(), source_file=src,
        source_sha256=rows_content_sha(recs), supersedes=tip, strict=True,
        allow_revert=True)
    if new == tip:
        halt(f"{p}: edition {edition} was not stored as the next edition; "
             "nothing to restore")
    return new


def cmd_restore_edition(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    spec = globals()["SPEC"]
    try:
        with conn.cursor() as cur:
            if not table_exists(cur, spec.editions_table):
                halt(f"{spec.editions_table} does not exist")
            new = restore_edition(cur, args.period, args.edition)
            print(f"{args.period}: edition {args.edition} "
                  + ("stored" if writing else "would be stored")
                  + f" as edition {new}; run refresh-latest to apply it")
            if args.commit:
                n = len(records(cur, spec, args.period, new))
                log_run(cur, n, f"restore-edition: {args.period} edition "
                        f"{args.edition} stored as edition {new}", started)
                conn.commit()
                print("COMMITTED; pipeline_run_log row written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="S21 ONS statistical nearest neighbours (Table 7a): "
        "editions loader on the period-editions engine, one release in "
        "live, live derived from the latest edition by LIVE_RULE. There is "
        "no sync-new; migrate-legacy records the held release once.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ddl", help="create the editions table, its triggers "
                       "and the file-check ledger, and add the live "
                       "release_date column (preview by default)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_ddl)
    sub.add_parser("status", help="what needs action; exit 1 if anything"
                   ).set_defaults(func=cmd_status)
    p = sub.add_parser("load", help="find the newest edition's workbook, "
                       "compare and store its release (preview by default)")
    p.add_argument("--file", metavar="PATH",
                   help="a local workbook under data/raw or data/reference "
                   "(nothing downloaded; identity is read from the file)")
    p.add_argument("--no-page", action="store_true",
                   help="with --file only: do not read the landing page "
                   "(recorded as such)")
    p.add_argument("--recheck", type=_period_arg, metavar="YYYY-MM-DD",
                   help="compare this held release again although the file "
                   "is in the ledger")
    p.add_argument("--allow-older-file", action="store_true",
                   help="compare and store a release older than the newest "
                   "held (as editions only; logged)")
    p.add_argument("--accept-reissue", action="append", type=_period_arg,
                   metavar="YYYY-MM-DD",
                   help="store a release whose file has the held file's "
                   "Cover date but different content (a publisher reissue "
                   "keeping its date; repeatable; logged)")
    p.add_argument("--acknowledge", action="append", type=_period_arg,
                   metavar="YYYY-MM-DD",
                   help="release a reissue's changed pairs, ranks or names "
                   "after reading the preview (repeatable; logged; never a "
                   "partial file)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_load)
    p = sub.add_parser("refresh-latest", help="replace live by LIVE_RULE "
                       "applied to the live release's latest edition "
                       "(preview by default)")
    p.add_argument("--acknowledge", action="append", type=_period_arg,
                   metavar="YYYY-MM-DD",
                   help="apply a refresh that changes live pairs or ranks "
                   "(the live release only; logged)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_refresh_latest)
    p = sub.add_parser("migrate-legacy", help="one-off: edition 1 as loaded "
                       "for the held release, with the proof (preview by "
                       "default)")
    p.add_argument("file", metavar="FILE")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_migrate_legacy)
    p = sub.add_parser("restore-edition", help="store an earlier edition's "
                       "rows as the next edition (preview by default)")
    p.add_argument("period", type=_period_arg, metavar="YYYY-MM-DD")
    p.add_argument("edition", type=int)
    pe.mode_parser(p)
    p.set_defaults(func=cmd_restore_edition)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
