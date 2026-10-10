"""S6 (Home Office asylum support by local authority: Asy_D11 and Reg_02):
edition history and loader on the period-editions engine.

Two publications, four live tables. They keep their names, keys, columns and
foreign keys and are the latest-edition layers; this module owns one
append-only editions table and one file-check ledger for each:

    SPEC_SUPPORT      la_asylum_support              (England, by lad24cd)
    SPEC_UNALLOCATED  la_asylum_support_unallocated  (no local authority)
    SPEC_NON_ENGLAND  asylum_support_non_england     (Scotland, Wales, NI)
    SPEC_GROUPS       la_immigration_groups          (Reg_02, England)

Asy_D11 ("Asylum seekers in receipt of Home Office support by local
authority detailed datasets", a quarterly time series restating every
quarter since 2014) feeds the first three; they are always applied together,
one period at a time, in one savepoint (s6_d11_period). Reg_02 (one snapshot
per file, "Regional and local authority data on immigration groups") feeds
the fourth (s6_reg02_period). The period is period_ending, the quarter-end
date (the engine's period string is the ISO date); Asy_D11 periods are
loaded from 2018-03-31 (the publisher's notes 14 to 16: no local authority
geography for Section 4 before 2018).

An edition is what one file says about one period of one table. A held
period whose content a file restates unchanged gets a ledger row only (never
a new edition, whatever the path or provenance); a changed one is stored as
the next edition and reaches live only through refresh-latest; a new period
is stored as edition 1 with its live rows in the same transaction.

History (so the wording here stays truthful): the old build
(scripts/s6_asylum_build.py with s6_asylum_verify.py) upserted every period
of the file on every run (INSERT ... ON CONFLICT DO UPDATE, loaded_at =
now()) and deleted and re-inserted asylum_series_breaks; it was run 15 times
(pipeline_run_log 69 to 82 on 25 and 26 July 2026, 98 on 4 September 2026).
The 4 September run rewrote the 2018 to March 2026 rows with values
identical to those it replaced (the March and June 2026 Asy_D11 files agree
cell for cell), so no held value was lost in practice. migrate-legacy
(one-off) records every held period as edition 1 "as loaded", with a proof
that this parser reproduces every held row and every stored column from the
held files.

Live provenance is one label: each live table's source_edition ('year ending
June 2026') is mapped from the editions table's release_label and set on
every row of a refreshed period (whole_period_cols); it is never compared.
loaded_at is copied from the edition on refresh. S6 is not a W1 or map
input; if it is ever wired into W1, run refresh-latest in the same session
before W1 (refresh copies the edition's loaded_at, which refresh_map.py would
otherwise not see as new).

Blanks and zeros (docs/RULES.md rule 1): an Asy_D11 People cell must be a
positive whole number (the publisher documents no marker, and the surveyed
files publish no zero): a blank, a zero, text, a negative or a non-integer
halts naming the sheet, row and column. In Reg_02 the publisher defines one marker, '*'
(fewer than 5, suppressed, note 9): it is stored NULL with suppressed true
and source_marker '*'; any other non-integer cell in a pathway column (a
blank, '-', ':', text) halts. A published 0 stays 0. A Reg_02 cell going
from 0 to '*' or back against the tip is accepted only with --acknowledge
PERIOD.

Identity (rule 3), from the file itself on every path: the cover sheet
(title line, 'year ending <Month yyyy>', publication title, 'Published: <d
Month yyyy>', 'Next update: ...'; a blank leading row allowed); Asy_D11's
Contents 'Period covered' and the data sheet's latest date; Reg_02's title
date; the headers exactly as surveyed. The rank of a file is (year-ending
date, published date) from its cover; for an edition or ledger row, from the
'year ending ...; published ...' text of its source_file. Never the URL, the
media id, the link text or the file name. The old asset URLs redirect to the
newest file, so the final URL is recorded and identity is read from the
file.

Geography (rule 4): geography.resolve(cur, '6', codes, by_period=...) per
table (Barnsley and Sheffield 'mixed': E08000016/19 to 30 September 2025 and
E08000038/39 from 31 December 2025 in Asy_D11; one form per period), then
la_code_lookup 'new_unitary' rows with exactly one target in la_boundaries
(the Cumbria, North Yorkshire and Somerset districts of 2018 to 2023), then
load_checks.check_codes; anything else is UNEXPLAINED and stops.

Subcommands (every writing command previews by default; --commit and
--simulate are mutually exclusive; --simulate runs the --commit path and
always rolls back). There is no sync-new: migrate-legacy records the held
periods.
    python scripts/s6_asylum_editions.py ddl [--commit | --simulate]
    python scripts/s6_asylum_editions.py status
    python scripts/s6_asylum_editions.py load [--only d11|reg02]
        [--release "Month YYYY"] [--file PATH [--no-page]]
        [--d09-file PATH | --no-d09] [--recheck YYYY-MM-DD]
        [--allow-older-file] [--acknowledge YYYY-MM-DD ...]
        [--accept-reissue YYYY-MM-DD] [--commit | --simulate]
        # files are downloaded to data/raw/s6_asylum/ (also in a preview; a
        # preview writes nothing to the database). A period breaking a stop
        # condition is REJECTED (nothing stored for it, exit 1).
    python scripts/s6_asylum_editions.py refresh-latest
        [--accept-key-changes YYYY-MM-DD] [--commit | --simulate]
    python scripts/s6_asylum_editions.py migrate-legacy --d11 FILE
        --reg02 FILE [--reg02 FILE] [--d09 FILE] [--commit | --simulate]
    python scripts/s6_asylum_editions.py restore-edition TABLE YYYY-MM-DD N
        [--commit | --simulate]
"""
import argparse
import calendar
import dataclasses
import hashlib
import json
import math
import numbers
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlparse

sys.path.insert(0, str(Path(__file__).resolve().parent))

import editions_core as core  # noqa: E402
import geography  # noqa: E402
import load_checks  # noqa: E402
import period_editions as pe  # noqa: E402
from editions_core import halt  # noqa: E402

RUN_AGENT = "s6_asylum_editions"
RUN_SOURCE = "6"
LIVE_MISSING = pe.LIVE_MISSING

API = "https://www.gov.uk/api/content"
GOV_UK = "https://www.gov.uk"
TABLES_PATH = ("/government/statistical-data-sets/"
               "immigration-system-statistics-data-tables")
REGIONAL_PATH = ("/government/statistical-data-sets/"
                 "immigration-system-statistics-regional-and-local-authority-"
                 "data")
TABLES_TITLE = "Immigration system statistics data tables"
REGIONAL_TITLE = "Regional and local authority data on immigration groups"
REPO = Path(__file__).resolve().parent.parent
RAW_DIR = REPO / "data" / "raw" / "s6_asylum"            # git-ignored
USER_AGENT = ("ucws-pipeline S6 loader (read-only download of the public Home "
              "Office immigration system statistics tables)")
MIN_FILE_BYTES = 20 * 1024
ODS_MIME = b"application/vnd.oasis.opendocument.spreadsheet"
NO_PAGE = "--file --no-page: the GOV.UK page was not read"

# The next release, as the held files' covers state it ('Next update: 26
# November 2026'). Data, not a claim: a fallback when no ledger row records
# the cover's date. Quarterly, Thursdays 09:30.
EXPECTED_NEXT_UPDATE = date(2026, 11, 26)

FLOOR = date(2018, 1, 1)        # Asy_D11 periods loaded: 2018-03-31 onward

# ---------------------------------------------------------------------------
# Stop conditions (calibrated on 2024-12-31 .. 2026-06-30: national moves
# -2.2% to +5.5% and -9.7%; authority count +26 at the end of the
# subsistence break; largest authority move 619; the June 2024 revision kept
# totals unchanged)
# ---------------------------------------------------------------------------

NEW_TOTAL_PCT = 15          # new Asy_D11 period: England total
NEW_AREAS_DELTA = 30        # new Asy_D11 period: authority count
NEW_AREA_ABS = 1500         # new Asy_D11 period: any authority's total
REV_TOTAL_PCT = 2           # revision: a table's total (England, ...)
REV_MAX_AREAS = 40          # revision: authorities whose totals change
REV_AREA_ABS = 500          # revision: any authority's total
REV_ROWS_PCT = 10           # revision: rows added, or removed
PARTIAL_AREAS = 10          # revision: fewer authorities than the tip
REG02_LAS = 296             # Reg_02: English authorities, exactly
REG02_REV_PCT = 5           # Reg_02 reissue: any pathway's England total
REG02_REV_MAX_AREAS = 30    # Reg_02 reissue: authorities changing
DUPLICATE_HALT = 5          # Asy_D11: same-code duplicate keys in one file
PARTIAL = "PARTIAL FILE"    # problems no --acknowledge releases

# ---------------------------------------------------------------------------
# The held state (migrate-legacy preconditions), surveyed 2026-10-10
# ---------------------------------------------------------------------------

# live table -> (rows, survey hash (live_state), {period: loaded_at}); a
# period not listed carries the '*' entry. Hashes split so the credential
# scan does not flag them.
LEGACY_LIVE = {
    "la_asylum_support": (21953, "5a8a9c162bbbe545" "427c07a754036166",
                          {"*": "2026-09-04 18:20:20.096642+00:00"}),
    "la_asylum_support_unallocated": (
        84, "eb0d207bef3a370a" "7668e04c5718ca82",
        {"*": "2026-09-04 18:20:20.096642+00:00"}),
    "asylum_support_non_england": (
        2553, "1446442a812148e4" "2465e122353da73e",
        {"*": "2026-09-04 18:20:20.096642+00:00"}),
    "la_immigration_groups": (
        7104, "d5b7203bc078c61f" "5bdb84fa95d2bcf9",
        {"2026-03-31": "2026-07-26 19:33:11.029590+00:00",
         "2026-06-30": "2026-09-04 18:20:20.096642+00:00"}),
}
_MEDIA = "https://assets.publishing.service.gov.uk/media/"
# the held files (sha256 split so the credential scan does not flag them),
# where each was published, and its rank from its own cover
LEGACY_FILES = {
    "d11": {"name": "support-local-authority-datasets-jun-2026.xlsx",
            "sha256": "d2207eed53afa18ce2d2afb317b6358d"
                      "9306765735c063e9c3b57b8e249708fa",
            "url": _MEDIA + "6a85c5e654bcee010d514574/"
                   "support-local-authority-datasets-jun-2026.xlsx",
            "rank": (date(2026, 6, 30), date(2026, 8, 27))},
    "d09": {"name": "asylum-seekers-receipt-support-datasets-jun-2026.xlsx",
            "sha256": "72f915e4ae27e78b8d077a3fffc96b45"
                      "a151343b1c8f8ee9d98750b2fd78a8c4",
            "url": _MEDIA + "6a85c2f1c9205b515d421eec/"
                   "asylum-seekers-receipt-support-datasets-jun-2026.xlsx",
            "rank": (date(2026, 6, 30), date(2026, 8, 27))},
    "reg02_mar": {"name": "regional-and-local-authority-dataset-mar-2026.ods",
                  "sha256": "1dd0e75630e1531ade660944ff02b608"
                            "344efa1e00809f0fee92fb727b299bc2",
                  "url": _MEDIA + "6a0f1367f71ef78abbd59dbf/"
                         "regional-and-local-authority-dataset-mar-2026.ods",
                  "rank": (date(2026, 3, 31), date(2026, 5, 21))},
    "reg02_jun": {"name": "regional-and-local-authority-dataset-jun-2026.ods",
                  "sha256": "5d857c6d55bbde81d46e85003f54ba09"
                            "8a39f09715ae1d5ac916d9a4be4dacbb",
                  "url": _MEDIA + "6a8c6c3ca8f84a582b842859/"
                         "regional-and-local-authority-dataset-jun-2026.ods",
                  "rank": (date(2026, 6, 30), date(2026, 8, 27))},
}

# Never written by this loader (the verify script checks them): the old
# build's two constant rows of asylum_series_breaks (break_id ignored).
SERIES_BREAKS = (
    (date(2022, 12, 31), None, "Section 98", "geography",
     "Section 98 gained local authority geography at 2022-12-31. Before that "
     "date all Section 98 people were published as a single national row with "
     "no LA and no accommodation type, and are held in "
     "la_asylum_support_unallocated.",
     "England totals are not comparable across the 2022 Q3 / Q4 boundary. "
     "England rises from 53,749 to 98,375 between 2022-09-30 and 2022-12-31. "
     "That is a reporting change, not 44,000 arrivals."),
    (date(2023, 12, 31), date(2024, 12, 31), "Section 95", "geography",
     "Subsistence Only lost local authority geography for five consecutive "
     "quarters, 2023-12-31 to 2024-12-31 inclusive. All subsistence-only "
     "people in those periods are held in la_asylum_support_unallocated.",
     "Local authority counts and England totals are depressed across those "
     "five periods. 32 English LAs appeared at 2023-09-30 only through "
     "subsistence-only rows; 26 of them are absent at 2023-12-31 and 13 are "
     "absent in all five periods."),
)

# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------


def _json(x):
    return json.loads(x) if isinstance(x, (str, bytes)) else (x or {})


def _norm(s) -> str:
    return " ".join(str("" if s is None else s).split())


def _fold(s) -> str:
    return _norm(s).casefold()


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or \
        (isinstance(v, str) and not v.strip())


def _text(v) -> "str | None":
    if _blank(v):
        return None
    return str(v).strip()


def _month_end(y: int, mo: int) -> date:
    return date(y, mo, calendar.monthrange(y, mo)[1])


def year_ending(text) -> date:
    """The quarter-end date of 'year ending <Month yyyy>' (or '<Month
    yyyy>'); ValueError unless the month ends a calendar quarter."""
    mm = re.fullmatch(r"(?:year ending )?([A-Za-z]+) ([0-9]{4})", _norm(text),
                      re.I)
    if not mm:
        raise ValueError(f"{text!r} is not 'year ending <Month yyyy>'")
    try:
        mo = datetime.strptime(mm.group(1).title(), "%B").month
    except ValueError:
        raise ValueError(f"{text!r}: {mm.group(1)!r} is not a month") \
            from None
    if mo % 3:
        raise ValueError(f"{text!r} does not end a calendar quarter")
    return _month_end(int(mm.group(2)), mo)


def ye_text(d: date) -> str:
    """'year ending June 2026' of a quarter-end date."""
    return f"year ending {d:%B %Y}"


def _long_date(text, where) -> date:
    try:
        return datetime.strptime(_norm(text), "%d %B %Y").date()
    except ValueError:
        raise ValueError(f"{where}: {text!r} is not a date 'd Month yyyy'") \
            from None


def _next_quarter(d: date) -> date:
    y, mo = (d.year + 1, 3) if d.month == 12 else (d.year, d.month + 3)
    return _month_end(y, mo)


def ranges_text(periods) -> str:
    """Quarter-end ISO dates as runs of consecutive quarters
    ('2018-03-31..2022-09-30,2023-12-31'); 'none' when empty."""
    ds = sorted(date.fromisoformat(str(p)) for p in periods)
    if not ds:
        return "none"
    out, start, prev = [], ds[0], ds[0]
    for d in ds[1:]:
        if d == _next_quarter(prev):
            prev = d
            continue
        out.append((start, prev))
        start = prev = d
    out.append((start, prev))
    return ",".join(a.isoformat() if a == b else f"{a}..{b}"
                    for a, b in out)


def parse_ranges(text) -> list:
    """The ISO dates of a ranges_text."""
    text = _norm(text)
    if text == "none":
        return []
    out = []
    for part in text.split(","):
        if ".." in part:
            a, b = (date.fromisoformat(x) for x in part.split(".."))
            d = a
            while d <= b:
                out.append(d.isoformat())
                d = _next_quarter(d)
        else:
            out.append(date.fromisoformat(part).isoformat())
    return out


# ---------------------------------------------------------------------------
# Discovery: the GOV.UK content API
# ---------------------------------------------------------------------------

_YE = r"year ending ([A-Z][a-z]+ [0-9]{4})"
D11_TITLE_RE = re.compile(r"Asylum seekers in receipt of Home Office support "
                          r"by local authority detailed datasets, " + _YE)
D09_TITLE_RE = re.compile(r"Asylum seekers in receipt of Home Office support "
                          r"detailed datasets, " + _YE)
REG02_TITLE_RE = re.compile(r"Regional and local authority data on "
                            r"immigration groups, " + _YE)


def _attachments(page_json) -> list:
    d = _json(page_json)
    return list(((d.get("details") or {}).get("attachments")) or [])


def _titles_like(page_json, words) -> list:
    return [_norm(a.get("title")) for a in _attachments(page_json)
            if any(w in _fold(a.get("title")) for w in words)]


def _page_title(page_json, want, path) -> None:
    got = _norm(_json(page_json).get("title"))
    if got != want:
        raise ValueError(f"{GOV_UK}{path}: the page title is {got!r}, not "
                         f"{want!r}; the page has changed. Attachment titles "
                         "with 'support' or 'local authority': "
                         f"{_titles_like(page_json, ('support', 'local authority'))}")


def _suffix(url) -> str:
    return Path(urlparse(str(url or "")).path).suffix.lower()


def _one_attachment(page_json, rx, loose, what, path) -> dict:
    """The page's one attachment strictly titled rx with an .xlsx URL. A
    title that looks like the series (loose(title)) but does not fit, none,
    several, or a non-.xlsx URL raise ValueError listing every attachment
    title with 'support' or 'local authority'."""
    _page_title(page_json, TABLES_TITLE, path)
    seen = _titles_like(page_json, ("support", "local authority"))
    hits, odd = [], []
    for a in _attachments(page_json):
        t = _norm(a.get("title"))
        mm = rx.fullmatch(t)
        if mm:
            hits.append((a, mm))
        elif loose(_fold(t)):
            odd.append(t)
    if odd:
        raise ValueError(f"{GOV_UK}{path}: attachment title(s) that look like "
                         f"{what} but do not fit the expected wording: {odd}; "
                         f"titles seen: {seen}")
    if len(hits) != 1:
        raise ValueError(f"{GOV_UK}{path}: {len(hits)} attachments titled "
                         f"like {what!r} (expected one); titles seen: {seen}")
    a, mm = hits[0]
    if _suffix(a.get("url")) != ".xlsx":
        raise ValueError(f"{GOV_UK}{path}: the {what} attachment is not an "
                         f".xlsx file ({a.get('url')}); titles seen: {seen}")
    ye = year_ending(mm.group(1))
    return {"title": _norm(a["title"]), "url": a["url"], "year_ending": ye,
            "label": ye_text(ye), "content_type": a.get("content_type")}


def asy_d11_attachment(page_json) -> dict:
    """{title, url, year_ending, label, content_type} of the data tables
    page's one Asy_D11 attachment ('Asylum seekers in receipt of Home Office
    support by local authority detailed datasets, year ending <Month
    yyyy>'). ValueError (listing the titles seen) on a changed page title,
    none, several, a renamed title or a non-.xlsx URL."""
    return _one_attachment(
        page_json, D11_TITLE_RE,
        lambda t: "receipt" in t and "support" in t and "local authority" in t,
        "Asylum seekers in receipt of Home Office support by local authority "
        "detailed datasets", TABLES_PATH)


def asy_d09_attachment(page_json) -> dict:
    """As asy_d11_attachment, for Asy_D09 ('Asylum seekers in receipt of
    Home Office support detailed datasets, year ending <Month yyyy>')."""
    return _one_attachment(
        page_json, D09_TITLE_RE,
        lambda t: ("receipt" in t and "support" in t
                   and "local authority" not in t),
        "Asylum seekers in receipt of Home Office support detailed datasets",
        TABLES_PATH)


def reg02_releases(page_json) -> dict:
    """{year-ending date: {title, url, year_ending, label, suffix}} of the
    regional page's attachments titled 'Regional and local authority data
    on immigration groups, year ending <Month yyyy>' with an .ods or .xlsx
    URL (both formats have been published). ValueError (listing the titles
    seen) on a changed page title, no such attachment, two for one quarter,
    another format, or a title that names immigration groups but does not
    fit."""
    d = _json(page_json)
    got = _norm(d.get("title"))
    seen = _titles_like(page_json, ("immigration groups", "local authority",
                                    "support"))
    if got != REGIONAL_TITLE:
        raise ValueError(f"{GOV_UK}{REGIONAL_PATH}: the page title is {got!r}, "
                         f"not {REGIONAL_TITLE!r}; the page has changed. "
                         f"Titles seen: {seen}")
    out, odd = {}, []
    for a in _attachments(page_json):
        t = _norm(a.get("title"))
        mm = REG02_TITLE_RE.fullmatch(t)
        if not mm:
            if "immigration groups" in _fold(t):
                odd.append(t)
            continue
        sfx = _suffix(a.get("url"))
        if sfx not in (".ods", ".xlsx"):
            raise ValueError(f"{GOV_UK}{REGIONAL_PATH}: {t!r} is not an .ods "
                             f"or .xlsx file ({a.get('url')}); titles seen: "
                             f"{seen}")
        ye = year_ending(mm.group(1))
        if ye in out:
            raise ValueError(f"{GOV_UK}{REGIONAL_PATH}: two attachments for "
                             f"{ye_text(ye)}: {out[ye]['title']!r} and {t!r}; "
                             f"refusing to choose; titles seen: {seen}")
        out[ye] = {"title": t, "url": a["url"], "year_ending": ye,
                   "label": ye_text(ye), "suffix": sfx}
    if odd:
        raise ValueError(f"{GOV_UK}{REGIONAL_PATH}: attachment title(s) naming "
                         f"immigration groups that do not fit the expected "
                         f"wording: {odd}; titles seen: {seen}")
    if not out:
        raise ValueError(f"{GOV_UK}{REGIONAL_PATH}: no attachment titled "
                         "'Regional and local authority data on immigration "
                         f"groups, year ending <Month yyyy>'; titles seen: "
                         f"{seen}")
    return out


def newest_reg02(page_json) -> dict:
    """The newest Reg_02 release by its year-ending date, whatever its
    format (an .xlsx newest is never passed over for an older .ods)."""
    rel = reg02_releases(page_json)
    return rel[max(rel)]


def change_history(page_json) -> list:
    """[(public_timestamp date, note)] of the page, as published."""
    d = _json(page_json)
    ch = ((d.get("details") or {}).get("change_history")) or []
    return [(str(c.get("public_timestamp") or "")[:10], _norm(c.get("note")))
            for c in ch]


def _today() -> date:
    return date.today()


def discovery_note(newest_seen, held_periods, next_update, today=None) -> list:
    """Lines to print after discovery. newest_seen: {publication: quarter-end
    ISO date of the newest release its page lists}; held_periods:
    {publication: periods held}; next_update: the held newest file's 'Next
    update' date (None: EXPECTED_NEXT_UPDATE). One line per publication with
    the newest seen and whether it is held; when every newest release is
    held and today is past next_update, a WARNING naming the date."""
    out, all_held = [], bool(newest_seen)
    for pub, p in sorted(newest_seen.items()):
        held = str(p) in {str(x) for x in held_periods.get(pub, ())}
        all_held = all_held and held
        out.append(f"newest {pub} release the page lists: "
                   f"{ye_text(date.fromisoformat(str(p)))} (period {p})"
                   + ("; already held" if held else "; not held yet"))
    when = next_update or globals()["EXPECTED_NEXT_UPDATE"]
    today = today or _today()
    if all_held and today > when:
        out.append(f"WARNING: the next release was expected on "
                   f"{when:%d %B %Y} (the held file's 'Next update') and today "
                   f"is {today:%d %B %Y}, but the pages list nothing newer than "
                   "what is held. It may be late, or discovery may be missing "
                   "it: check the GOV.UK pages by hand (--file PATH loads a "
                   "downloaded file)")
    return out


# ---------------------------------------------------------------------------
# Reading workbooks (.xlsx with openpyxl, .ods with pandas + odfpy)
# ---------------------------------------------------------------------------


def _pad(rows) -> list:
    width = max((len(r) for r in rows), default=1)
    return [tuple(r) + (None,) * (width - len(r)) for r in rows]


def _read_sheets(path, names, optional=()) -> dict:
    """{sheet: [row tuples]} of the named sheets (blank cells None). Raises
    ValueError naming the sheets seen when one of names is missing."""
    p = Path(path)
    want = tuple(names) + tuple(optional)
    if p.suffix.lower() == ".ods":
        import pandas as pd
        try:
            x = pd.ExcelFile(p, engine="odf")
        except Exception as e:  # noqa: BLE001 (not a spreadsheet)
            raise ValueError(f"{p.name}: not a readable .ods file ({e})") \
                from None
        with x:
            sheets = list(x.sheet_names)
            missing = [s for s in names if s not in sheets]
            if missing:
                raise ValueError(f"{p.name}: sheets {missing} missing; sheets "
                                 f"{sheets}")
            out = {}
            for s in want:
                if s not in sheets:
                    continue
                df = pd.read_excel(x, sheet_name=s, header=None, dtype=object)
                rows = []
                for r in df.itertuples(index=False):
                    rows.append(tuple(None if _blank(v) and
                                      not isinstance(v, str) else v
                                      for v in r))
                out[s] = _pad(rows)
        return out
    import openpyxl
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 (not a workbook)
        raise ValueError(f"{p.name}: not a readable .xlsx workbook ({e})") \
            from None
    try:
        missing = [s for s in names if s not in wb.sheetnames]
        if missing:
            raise ValueError(f"{p.name}: sheets {missing} missing; sheets "
                             f"{wb.sheetnames}")
        return {s: _pad([tuple(r) for r in wb[s].iter_rows(values_only=True)])
                for s in want if s in wb.sheetnames}
    finally:
        wb.close()


def _cells(rows, limit=None) -> list:
    out = [_norm(v) for r in rows for v in r if not _blank(v)]
    return out[:limit] if limit else out


# ---------------------------------------------------------------------------
# The cover sheet
# ---------------------------------------------------------------------------

SERIES_LINE = "Immigration System Statistics"
PUBLICATIONS = {
    "d11": "Asylum - Asylum seekers in receipt of Home Office support by "
           "Local Authority",
    "d09": "Asylum - Asylum seekers in receipt of Home Office support",
    "reg02": "Regional and Local authority data - Immigration groups",
}
COVER = "Cover_sheet"


def parse_cover(rows, name="") -> dict:
    """{series, label ('year ending June 2026'), year_ending (date),
    publication, kind ('d11', 'd09' or 'reg02'), published (date),
    next_update (date or None), next_update_text, home_office (bool)} of a
    cover sheet's rows. The first non-blank cells must be the series line,
    'year ending <Month yyyy>' and the publication title (a blank leading
    row is allowed); exactly one 'Published: <d Month yyyy>' and one 'Next
    update: ...' among the first twelve cells. ValueError naming what was
    seen otherwise."""
    cells = _cells(rows, 12)
    where = f"{name}: {COVER}"
    if len(cells) < 3 or cells[0] != SERIES_LINE:
        raise ValueError(f"{where}: the first line is not {SERIES_LINE!r}; "
                         f"cells seen: {cells}")
    try:
        ye = year_ending(cells[1])
    except ValueError as e:
        raise ValueError(f"{where}: the second line: {e}; cells seen: "
                         f"{cells}") from None
    if not cells[1].lower().startswith("year ending "):
        raise ValueError(f"{where}: {cells[1]!r} is not 'year ending <Month "
                         "yyyy>'")
    pub = cells[2]
    kinds = [k for k, t in PUBLICATIONS.items() if _fold(t) == _fold(pub)]
    if len(kinds) != 1:
        raise ValueError(f"{where}: publication title {pub!r} is not one of "
                         f"{list(PUBLICATIONS.values())}")
    published = [c for c in cells if c.startswith("Published:")]
    nxt = [c for c in cells if c.startswith("Next update:")]
    if len(published) != 1 or len(nxt) != 1:
        raise ValueError(f"{where}: expected one 'Published:' and one 'Next "
                         f"update:' line; cells seen: {cells}")
    pd_ = _long_date(published[0].split(":", 1)[1], where)
    ntext = _norm(nxt[0].split(":", 1)[1])
    try:
        nd = _long_date(ntext, where)
    except ValueError:
        nd = None
    return {"series": cells[0], "label": ye_text(ye), "year_ending": ye,
            "publication": pub, "kind": kinds[0], "published": pd_,
            "next_update": nd, "next_update_text": ntext,
            "home_office": "Home Office" in cells}


def read_cover(path) -> dict:
    """parse_cover of the file's Cover_sheet, plus path and rank
    ((year_ending, published))."""
    rows = _read_sheets(path, (COVER,))[COVER]
    cov = parse_cover(rows, Path(path).name)
    cov["path"] = Path(path)
    cov["rank"] = (cov["year_ending"], cov["published"])
    return cov


# ---------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------


def people_cell(v, where) -> int:
    """An Asy_D11 / Asy_D09 People cell: a positive whole number (an int, or
    a float equal to one). Anything else raises ValueError naming where: a
    blank (the publisher documents no marker; never read as 0), a zero (the
    surveyed files publish none: a 0 would be new and is reviewed, not
    guessed at), text, a boolean, a negative or a non-integer."""
    if v is None or (isinstance(v, str) and not v.strip()):
        raise ValueError(f"{where}: a blank People cell (the publisher "
                         "documents no marker; a blank halts and is never "
                         "read as 0)")
    if isinstance(v, bool):
        raise ValueError(f"{where}: {v!r} is not a count")
    if isinstance(v, numbers.Integral):
        n = int(v)
    elif isinstance(v, numbers.Real):
        if math.isnan(v):
            raise ValueError(f"{where}: a blank People cell (NaN); never read "
                             "as 0")
        if not float(v).is_integer():
            raise ValueError(f"{where}: non-integer People {v!r}")
        n = int(v)
    else:
        raise ValueError(f"{where}: People {v!r} is not a whole number")
    if n < 0:
        raise ValueError(f"{where}: negative People {n}")
    if n == 0:
        raise ValueError(f"{where}: People 0; the surveyed files publish no "
                         "zero in this series (an absent row means not "
                         "published), so a 0 is reviewed with the evidence "
                         "before it is stored")
    return n


def _whole(v, where, what) -> int:
    """A non-negative whole number (int, or a float equal to one)."""
    if isinstance(v, bool) or v is None:
        raise ValueError(f"{where}: {what} {v!r} is not a whole number")
    if isinstance(v, numbers.Integral):
        n = int(v)
    elif isinstance(v, numbers.Real) and not math.isnan(v) \
            and float(v).is_integer():
        n = int(v)
    else:
        raise ValueError(f"{where}: {what} {v!r} is not a whole number")
    if n < 0:
        raise ValueError(f"{where}: negative {what} {n}")
    return n


def reg02_cell(v, where) -> tuple:
    """A Reg_02 pathway cell: (int, None) for a published whole number (a
    published 0 stays 0), (None, '*') for '*' (fewer than 5, suppressed:
    the one marker the publisher defines, note 9). Anything else raises
    ValueError naming where: a blank, '-', ':' or other text (the publisher
    defines none of them in Reg_02), a negative or a non-integer."""
    if isinstance(v, str):
        s = v.strip()
        if s == "*":
            return None, "*"
        raise ValueError(f"{where}: {v!r} is not a count or '*' (Reg_02 "
                         "defines no other marker; a blank or '-' halts)")
    if v is None or (isinstance(v, float) and math.isnan(v)):
        raise ValueError(f"{where}: a blank cell (Reg_02 defines no blank; "
                         "never read as 0)")
    return _whole(v, where, "value"), None


def ratio_cell(v, where) -> Decimal:
    """Reg_02 'Percentage of population (%)': the publisher's ratio (a
    fraction such as 0.004991745, despite the header), rounded to 4
    decimals exactly as the old build stored it (round(float, 4) into
    numeric(8,4)). ValueError on anything but a number."""
    if isinstance(v, bool) or not isinstance(v, numbers.Real) or (
            isinstance(v, float) and math.isnan(v)):
        raise ValueError(f"{where}: percentage {v!r} is not a number")
    return Decimal(repr(round(float(v), 4))).quantize(Decimal("0.0001"))


# ---------------------------------------------------------------------------
# Asy_D11, Asy_D09 and Reg_02
# ---------------------------------------------------------------------------

D11_SHEET = "Data_Asy_D11"
D11_PIVOT = "Asy_D11"
D09_SHEET = "Data_Asy_D09"
REG02_SHEET = "Reg_02"
CONTENTS = "Contents"
D11_HEADERS = ("Date (as at…)", "Support Type", "UK Region / Nation",
               "Local Authority", "LAD Code", "Accommodation Type", "People")
D09_HEADERS = ("Date (as at…)", "Nationality", "Region", "Support Type",
               "Accommodation Type", "UK Region / Nation", "People")
REG02_HEADERS = (
    "Local authority", "Region / Nation", "LTLA (ONS code)",
    "Homes for Ukraine - not including super sponsors (arrivals)",
    "Afghan Resettlement Programme (total) (population)",
    "of which, Afghan Resettlement Programme - transitional (population)",
    "of which, Afghan Resettlement Programme - settled in LA housing "
    "(population)",
    "of which, Afghan Resettlement Programme - settled in PRS housing "
    "(population)",
    "Supported Asylum (total) (population)",
    "of which, Supported Asylum - Initial Accommodation (population)",
    "of which, Supported Asylum - Dispersal Accommodation (population)",
    "of which, Supported Asylum - Contingency Accommodation (population)",
    "of which, Supported Asylum - Other Accommodation (population)",
    "of which, Subsistence only (population)",
    "All 3 pathways (total)", "Population", "Percentage of population (%)")
REG02_TITLE_SHEET_RE = re.compile(r"Reg_02: Immigration groups, by Local "
                                  r"Authority, as at ([0-9]{1,2} [A-Za-z]+ "
                                  r"[0-9]{4})")
PERIOD_COVERED_RE = re.compile(r"([0-9]{4}) to ([0-9]{4})(?: Q([1-4]))?")
SUPPORT_TYPES = ("Section 4", "Section 95", "Section 98")
ACCOMMODATION = ("Dispersal Accommodation", "Initial Accommodation",
                 "Contingency Accommodation - Hotel",
                 "Contingency Accommodation - Other", "Other Accommodation",
                 "Subsistence Only")
NOT_STATED = "not_stated"
ENGLISH_PREFIXES = ("E06", "E07", "E08", "E09")
COUNTRY_BY_PREFIX = {"S12": "Scotland", "W06": "Wales",
                     "N09": "Northern Ireland"}
ENGLISH_REGIONS = ("East Midlands", "East of England", "London", "North East",
                   "North West", "South East", "South West", "West Midlands",
                   "Yorkshire and The Humber")
GRAND_TOTAL = "Grand Total"
# Reg_02 pathway columns (index into the 14 cells after the code) and the
# stored (pathway, sub_pathway), as the old build stored them
REG02_PATHWAYS = (
    (0, "homes_for_ukraine", "total"),
    (1, "afghan_resettlement", "total"),
    (2, "afghan_resettlement", "transitional"),
    (3, "afghan_resettlement", "settled_la_housing"),
    (4, "afghan_resettlement", "settled_prs_housing"),
    (5, "supported_asylum", "total"),
    (6, "supported_asylum", "initial_accommodation"),
    (7, "supported_asylum", "dispersal"),
    (8, "supported_asylum", "contingency"),
    (9, "supported_asylum", "other"),
    (10, "supported_asylum", "subsistence_only"),
    (11, "all_pathways", "total"),
)
REG02_POPULATION = 12
REG02_PERCENT = 13
LOWER_BOUND = ("LOWER BOUND: published total excludes the suppressed {} pathway "
               "(fewer than 5 people, disclosure control). True total is higher "
               "by between 1 and 4 per suppressed pathway.")


def _header(rows, i, want, sheet, name, *, rstrip=False) -> int:
    """Check row i is exactly the headers want (trailing empty columns
    allowed; with rstrip trailing spaces ignored); return its width."""
    if len(rows) <= i:
        raise ValueError(f"{name}: {sheet} has no header row {i + 1}")
    got = [_norm(v) if rstrip else ("" if v is None else str(v))
           for v in rows[i]]
    while got and not got[-1].strip():
        got.pop()
    wanted = [_norm(h) if rstrip else h for h in want]
    if got != wanted:
        unknown = [h for h in got if h not in wanted]
        missing = [h for h in wanted if h not in got]
        raise ValueError(f"{name}: {sheet} header {got} is not the surveyed "
                         f"set: unknown {unknown}, missing {missing} (expected "
                         f"exactly {wanted}, in order); the mapping must be "
                         "corrected deliberately, with the evidence")
    return len(want)


def _body(rows, start, width, sheet, name, *, end_marker=None):
    """(row number, cells) of the body rows from index start: wholly empty
    rows only at the end; no cell beyond width."""
    empty = 0  # trailing empty rows seen; not a source value
    for n, r in enumerate(rows[start:], start=start + 1):
        if all(_blank(v) for v in r):
            empty += 1
            continue
        where = f"{name}: {sheet} row {n}"
        if end_marker and _norm(r[0]) == end_marker and \
                all(_blank(v) for v in r[1:]):
            empty += 1
            continue
        if empty:
            raise ValueError(f"{where}: data after {empty} empty row(s); empty "
                             "rows are allowed only at the end")
        if any(not _blank(v) for v in r[width:]):
            raise ValueError(f"{where}: a cell beyond the {width} headed "
                             "columns")
        yield n, r[:width]


def _d11_date(v, where) -> date:
    if not isinstance(v, str):
        raise ValueError(f"{where}: date {v!r} is not text 'dd Mon yyyy'")
    try:
        return datetime.strptime(v.strip(), "%d %b %Y").date()
    except ValueError:
        raise ValueError(f"{where}: date {v!r} is not 'dd Mon yyyy'") from None


def _period_covered(rows, name) -> dict:
    """{sheet name: quarter-end date} of the Contents table's 'Period
    covered' column ('2014 to 2026 Q2' -> 2026-06-30; a whole year, as the
    December release states it, '2014 to 2025' -> 2025-12-31)."""
    heads = [(i, [_norm(v) for v in r]) for i, r in enumerate(rows)
             if "Period covered" in [_norm(v) for v in r]]
    if len(heads) != 1:
        raise ValueError(f"{name}: {CONTENTS} has {len(heads)} header rows "
                         "with 'Period covered'")
    i, h = heads[0]
    cs, cp = h.index("Sheet"), h.index("Period covered")
    out = {}
    for r in rows[i + 1:]:
        sheet = _norm(r[cs])
        if not sheet:
            continue
        mm = PERIOD_COVERED_RE.fullmatch(_norm(r[cp]))
        if not mm:
            raise ValueError(f"{name}: {CONTENTS}: period covered {r[cp]!r} "
                             f"of {sheet!r} is not '<yyyy> to <yyyy>[ Q<n>]'")
        out[sheet] = _month_end(int(mm.group(2)),
                                3 * int(mm.group(3) or 4))
    return out


def _pivot(rows, name) -> dict:
    """{dates: [quarter-end dates], rows: {label: [int or None]}} of the
    Asy_D11 pivot sheet's cached values (the latest eight quarters), with
    both filters '(All)'."""
    for label in ("Support Type", "Accommodation Type"):
        f = [r for r in rows if _norm(r[0]) == label]
        if len(f) != 1 or _norm(f[0][1]) != "(All)":
            raise ValueError(f"{name}: {D11_PIVOT}: the {label!r} filter is "
                             f"not '(All)' ({[_norm(x[1]) for x in f]}); the "
                             "cached totals would not be the whole table")
    heads = [i for i, r in enumerate(rows) if _norm(r[0]) == "Row Labels"]
    if len(heads) != 1:
        raise ValueError(f"{name}: {D11_PIVOT} has {len(heads)} 'Row Labels' "
                         "rows")
    i = heads[0]
    dates = []
    for v in rows[i][1:]:
        if _blank(v):
            break
        dates.append(_d11_date(v, f"{name}: {D11_PIVOT} row {i + 1}"))
    out = {}
    for n, r in enumerate(rows[i + 1:], start=i + 2):
        label = _norm(r[0])
        if not label:
            raise ValueError(f"{name}: {D11_PIVOT} row {n}: a blank row label "
                             f"before {GRAND_TOTAL!r}")
        vals = []
        for v in r[1:1 + len(dates)]:
            vals.append(None if _blank(v) else
                        _whole(v, f"{name}: {D11_PIVOT} row {n}", "value"))
        out[label] = vals
        if label == GRAND_TOTAL:
            break
    if GRAND_TOTAL not in out:
        raise ValueError(f"{name}: {D11_PIVOT} has no {GRAND_TOTAL!r} row")
    return {"dates": dates, "rows": out}


def read_d11(path) -> dict:
    """The facts of one Asy_D11 workbook: {path, cover, period (the cover's
    quarter end), contents {sheet: date}, header, rows [{row, date, support,
    region, la, code, accommodation, people}], periods (every date stated),
    pivot}. People cells are checked (people_cell) on every row.

    Raises ValueError (naming what was seen) on: a missing sheet; a cover
    that is not Asy_D11's; a Contents 'Period covered' other than the
    cover's quarter; a data header that is not exactly D11_HEADERS; a row
    after an empty row or a cell beyond the header; a date that is not 'dd
    Mon yyyy'; a People cell that is not a positive whole number; a latest
    data date other than the cover's quarter end; a pivot whose filters are
    not '(All)' or whose last quarter is not the cover's."""
    p = Path(path)
    sh = _read_sheets(p, (COVER, CONTENTS, D11_PIVOT, D11_SHEET))
    cov = parse_cover(sh[COVER], p.name)
    if cov["kind"] != "d11":
        raise ValueError(f"{p.name}: the cover says {cov['publication']!r}, "
                         f"not {PUBLICATIONS['d11']!r}")
    period = cov["year_ending"]
    contents = _period_covered(sh[CONTENTS], p.name)
    mine = {s: d for s, d in contents.items() if "asy_d11" in _fold(s).replace(
        "data - ", "data_")}
    if not any(_fold(s).startswith("data") for s in mine):
        raise ValueError(f"{p.name}: {CONTENTS} lists no data sheet for "
                         f"Asy_D11 ({sorted(contents)})")
    off = {s: d.isoformat() for s, d in mine.items() if d != period}
    if off:
        raise ValueError(f"{p.name}: {CONTENTS} period covered {off}, the "
                         f"cover says {cov['label']} ({period})")
    rows = sh[D11_SHEET]
    if not rows or "Asy_D11" not in _norm(rows[0][0]):
        raise ValueError(f"{p.name}: {D11_SHEET} row 1 is not the Asy_D11 "
                         f"title ({_norm(rows[0][0]) if rows else 'empty'})")
    width = _header(rows, 1, D11_HEADERS, D11_SHEET, p.name)
    out = []
    for n, r in _body(rows, 2, width, D11_SHEET, p.name):
        where = f"{p.name}: {D11_SHEET} row {n}"
        out.append({"row": n, "date": _d11_date(r[0], where),
                    "support": _text(r[1]), "region": _text(r[2]),
                    "la": _text(r[3]), "code": _text(r[4]),
                    "accommodation": _text(r[5]),
                    "people": people_cell(r[6], f"{where} People")})
    if not out:
        raise ValueError(f"{p.name}: {D11_SHEET} has no data rows")
    latest = max(r["date"] for r in out)
    if latest != period:
        raise ValueError(f"{p.name}: the latest date in {D11_SHEET} is "
                         f"{latest}, the cover says {cov['label']} ({period})")
    piv = _pivot(sh[D11_PIVOT], p.name)
    if not piv["dates"] or piv["dates"][-1] != period:
        raise ValueError(f"{p.name}: the {D11_PIVOT} pivot's last quarter is "
                         f"{piv['dates'][-1] if piv['dates'] else 'none'}, the "
                         f"cover says {period}")
    cov.update(path=p, rank=(cov["year_ending"], cov["published"]))
    return {"path": p, "cover": cov, "period": period, "contents": contents,
            "header": list(D11_HEADERS), "rows": out,
            "periods": sorted({r["date"] for r in out}), "pivot": piv}


def read_d09(path) -> dict:
    """{path, cover, period, totals {date: {UK region (case-folded): people}}}
    of an Asy_D09 workbook (header exactly D09_HEADERS, People cells
    positive whole numbers, latest date the cover's quarter end)."""
    p = Path(path)
    sh = _read_sheets(p, (COVER, D09_SHEET))
    cov = parse_cover(sh[COVER], p.name)
    if cov["kind"] != "d09":
        raise ValueError(f"{p.name}: the cover says {cov['publication']!r}, "
                         f"not {PUBLICATIONS['d09']!r}")
    rows = sh[D09_SHEET]
    width = _header(rows, 1, D09_HEADERS, D09_SHEET, p.name)
    totals = {}
    for n, r in _body(rows, 2, width, D09_SHEET, p.name):
        where = f"{p.name}: {D09_SHEET} row {n}"
        d = _d11_date(r[0], where)
        reg = _fold(r[5])
        acc = totals.setdefault(d, {})
        acc[reg] = acc.get(reg, 0) + people_cell(r[6], f"{where} People")  # not a source value
    if not totals:
        raise ValueError(f"{p.name}: {D09_SHEET} has no data rows")
    if max(totals) != cov["year_ending"]:
        raise ValueError(f"{p.name}: the latest date in {D09_SHEET} is "
                         f"{max(totals)}, the cover says {cov['label']}")
    cov.update(path=p, rank=(cov["year_ending"], cov["published"]))
    return {"path": p, "cover": cov, "period": cov["year_ending"],
            "totals": totals}


def read_reg02(path) -> dict:
    """The facts of one Reg_02 workbook (.ods or .xlsx): {path, cover,
    period, title_date, header, rows [{row, la, region, code, cells (the 14
    cells after the code)}]}. Raises ValueError on a cover that is not
    Reg_02's, a Reg_02 title whose date is not the cover's quarter end, a
    header that is not the 17 surveyed columns (trailing spaces ignored;
    unknown and missing listed), a row after an empty row, or a row without
    a code."""
    p = Path(path)
    sh = _read_sheets(p, (COVER, REG02_SHEET))
    cov = parse_cover(sh[COVER], p.name)
    if cov["kind"] != "reg02":
        raise ValueError(f"{p.name}: the cover says {cov['publication']!r}, "
                         f"not {PUBLICATIONS['reg02']!r}")
    rows = sh[REG02_SHEET]
    title = _norm(rows[0][0]) if rows else ""
    mm = REG02_TITLE_SHEET_RE.fullmatch(title)
    if not mm:
        raise ValueError(f"{p.name}: {REG02_SHEET} row 1 {title!r} is not "
                         "'Reg_02: Immigration groups, by Local Authority, as "
                         "at <d Month yyyy>'")
    td = _long_date(mm.group(1), f"{p.name}: {REG02_SHEET}")
    if td != cov["year_ending"]:
        raise ValueError(f"{p.name}: the {REG02_SHEET} title says as at "
                         f"{mm.group(1)}, the cover says {cov['label']}")
    width = _header(rows, 1, REG02_HEADERS, REG02_SHEET, p.name, rstrip=True)
    out = []
    for n, r in _body(rows, 2, width, REG02_SHEET, p.name,
                      end_marker="End of table"):
        code = _text(r[2])
        if code is None:
            raise ValueError(f"{p.name}: {REG02_SHEET} row {n}: no LTLA code "
                             f"({[v for v in r if not _blank(v)][:4]})")
        out.append({"row": n, "la": _text(r[0]), "region": _text(r[1]),
                    "code": code, "cells": list(r[3:width])})
    cov.update(path=p, rank=(cov["year_ending"], cov["published"]))
    return {"path": p, "cover": cov, "period": cov["year_ending"],
            "title_date": td, "header": list(REG02_HEADERS), "rows": out}


def read_any(path) -> dict:
    """read_d11, read_d09 or read_reg02, by the file's own cover."""
    kind = read_cover(path)["kind"]
    return {"d11": read_d11, "d09": read_d09, "reg02": read_reg02}[kind](path)


# ---------------------------------------------------------------------------
# Records
# ---------------------------------------------------------------------------

D11_TAGS = ("support", "unallocated", "non_england")
TAGS = D11_TAGS + ("groups",)


def _getter(resolve):
    return resolve if callable(resolve) else (lambda c: resolve[c])


def _is_na(v) -> bool:
    return (v or "").strip().upper().startswith("N/A")


def normalise_accommodation(value, where="") -> tuple:
    """(accommodation type, marker): an 'N/A...' value is NOT_STATED with
    the verbatim text as its marker; any other value is trimmed and each
    word title-cased ('Subsistence only' -> 'Subsistence Only') and must be
    one of ACCOMMODATION (else ValueError naming where)."""
    raw = (value or "").strip()
    if raw.upper().startswith("N/A"):
        return NOT_STATED, raw
    norm = " ".join(w.capitalize() for w in raw.split())
    if norm not in ACCOMMODATION:
        raise ValueError(f"{where}: unknown accommodation type {value!r} "
                         f"(known: {list(ACCOMMODATION)} and 'N/A ...'); the "
                         "mapping must be corrected deliberately, with the "
                         "evidence")
    return norm, None


def d11_codes_by_period(file) -> dict:
    """{period ISO: {English codes}} of the loaded periods (from FLOOR)."""
    out = {}
    for r in file["rows"]:
        if r["date"] < FLOOR:
            continue
        c = r["code"] or ""
        if c[:3] in ENGLISH_PREFIXES and not _is_na(r["la"]) \
                and r["la"] != "Unknown":
            out.setdefault(r["date"].isoformat(), set()).add(c)
    return out


def name_pick(code, name, marker=None) -> tuple:
    """The order of the rows of one key for choosing its published name:
    the highest publisher code wins, then the greatest name, then marker.
    A total order on the row's own values, so the choice is the same
    whatever order the file lists the rows in."""
    return (code or "", name or "", marker or "")


def build_d11_records(file, resolve) -> tuple:
    """(by_period, merges, duplicates) of a read Asy_D11 (read_d11).
    by_period: {period ISO: {'support': [...], 'unallocated': [...],
    'non_england': [...]}} for every period from FLOOR (a list may be
    empty). Rules (as the old build): support type one of SUPPORT_TYPES;
    accommodation normalised (normalise_accommodation); a row whose Local
    Authority or LAD Code starts 'N/A', or whose Local Authority is
    'Unknown', is unallocated with na_reason the verbatim Local Authority
    (or the code when only the code says N/A); S12/W06/N09 codes are
    non-England (country from the prefix); E06-E09 codes are resolved
    (resolve: {code: lad24cd} or a callable); anything else raises. People
    are summed per key; published_la_name (and source_marker, country) are
    those of the key's row with the highest publisher code (ties: the
    greatest name, then marker), so they never depend on row order (a
    reordered release gives the same records; name_pick). This reproduces
    every held name, which the old build took from the last row in file
    order (the held files list a merged key's highest code last; the
    December 2025 file does not). merges: keys reached by distinct
    codes (a reorganisation merge); duplicates: keys reached twice by the
    same code. More than DUPLICATE_HALT duplicates raise ValueError."""
    get = _getter(resolve)
    acc = {t: {} for t in D11_TAGS}
    sources = {}
    periods = sorted({r["date"] for r in file["rows"] if r["date"] >= FLOOR})
    name = Path(file["path"]).name
    for r in file["rows"]:
        if r["date"] < FLOOR:
            continue
        where = f"{name}: {D11_SHEET} row {r['row']}"
        p = r["date"].isoformat()
        support = r["support"]
        if support not in SUPPORT_TYPES:
            raise ValueError(f"{where}: unknown support type {support!r} "
                             f"(known: {list(SUPPORT_TYPES)})")
        acc_type, marker = normalise_accommodation(r["accommodation"], where)
        la, code = r["la"] or "", r["code"] or ""
        if _is_na(la) or _is_na(code) or la == "Unknown":
            reason = la if (_is_na(la) or la == "Unknown") else code
            tag, key = "unallocated", (support, acc_type, reason)
            meta = {}
        elif code[:3] in COUNTRY_BY_PREFIX:
            tag, key = "non_england", (code, support, acc_type)
            meta = {"country": COUNTRY_BY_PREFIX[code[:3]],
                    "published_la_name": la}
        elif code[:3] in ENGLISH_PREFIXES:
            lad = get(code)
            if not lad:
                raise ValueError(f"{where}: {code} resolves to no lad24cd")
            tag, key = "support", (lad, support, acc_type)
            meta = {"published_la_name": la, "source_marker": marker}
        else:
            raise ValueError(f"{where}: LAD Code {code!r} ({la!r}) is not an "
                             "E06-E09, S12, W06 or N09 code and not 'N/A'")
        slot = acc[tag].setdefault((p, key), {"people": 0})  # not a source value
        slot["people"] += r["people"]
        pick = name_pick(code, la, marker)
        if "_pick" not in slot or pick > slot["_pick"]:
            slot.update(meta, _pick=pick)
        sources.setdefault((tag, p, key), []).append(r)
    merges, dups = [], []
    for (tag, p, key), rows in sorted(sources.items(), key=str):
        if len(rows) < 2:
            continue
        codes = sorted({x["code"] for x in rows})
        item = {"table": tag, "period": p, "key": key, "codes": codes,
                "rows": [x["row"] for x in rows],
                "regions": [x["region"] for x in rows],
                "people": [x["people"] for x in rows]}
        (dups if len(codes) == 1 else merges).append(item)
    if len(dups) > globals()["DUPLICATE_HALT"]:
        raise ValueError(f"{name}: {len(dups)} keys published twice under the "
                         f"same code (limit {globals()['DUPLICATE_HALT']}), "
                         f"e.g. {[(d['period'], d['key']) for d in dups[:6]]}")
    by_period = {p.isoformat(): {t: [] for t in D11_TAGS} for p in periods}
    for (p, key), v in sorted(acc["support"].items(), key=str):
        by_period[p]["support"].append({
            "lad24cd": key[0], "support_type": key[1],
            "accommodation_type": key[2], "period_ending": p,
            "people": v["people"], "published_la_name": v["published_la_name"],
            "source_marker": v["source_marker"]})
    for (p, key), v in sorted(acc["unallocated"].items(), key=str):
        by_period[p]["unallocated"].append({
            "support_type": key[0], "accommodation_type": key[1],
            "na_reason": key[2], "period_ending": p, "people": v["people"]})
    for (p, key), v in sorted(acc["non_england"].items(), key=str):
        by_period[p]["non_england"].append({
            "lad_code": key[0], "support_type": key[1],
            "accommodation_type": key[2], "period_ending": p,
            "people": v["people"], "country": v["country"],
            "published_la_name": v["published_la_name"]})
    return by_period, merges, dups


def reg02_counts(file) -> dict:
    """{'england', 'scotland', 'wales', 'northern_ireland', 'unknown'} row
    counts of a read Reg_02 (only England is stored)."""
    out = dict.fromkeys(("england", "scotland", "wales", "northern_ireland",
                         "unknown"), 0)  # not a source value
    names = {"S12": "scotland", "W06": "wales", "N09": "northern_ireland"}
    for r in file["rows"]:
        c = r["code"]
        if c[:3] in ENGLISH_PREFIXES:
            out["england"] += 1
        elif c[:3] in names:
            out[names[c[:3]]] += 1
        elif c == "Unknown":
            out["unknown"] += 1
    return out


def build_reg02_records(file, resolve) -> tuple:
    """(period ISO, records) of a read Reg_02 (read_reg02): twelve records
    per English (E06-E09) row, as the old build stored them: people from
    reg02_cell ('*' -> NULL, suppressed true, source_marker '*'; a published
    0 stays 0), population on every row of the authority, and on
    all_pathways/total the percentage (ratio_cell) and, where a pathway
    total is '*', the LOWER BOUND marker naming the suppressed pathways.
    S12/W06/N09 rows and the 'Unknown' row are read but not stored; any
    other code raises. Two rows reaching one lad24cd raise."""
    get = _getter(resolve)
    name = Path(file["path"]).name
    p = file["period"].isoformat()
    out, seen = [], {}
    for r in file["rows"]:
        code = r["code"]
        if code[:3] in COUNTRY_BY_PREFIX or code == "Unknown":
            continue
        where = f"{name}: {REG02_SHEET} row {r['row']}"
        if code[:3] not in ENGLISH_PREFIXES:
            raise ValueError(f"{where}: LTLA code {code!r} is not E06-E09, "
                             "S12, W06, N09 or 'Unknown'")
        lad = get(code)
        if not lad:
            raise ValueError(f"{where}: {code} resolves to no lad24cd")
        if lad in seen:
            raise ValueError(f"{where}: {code} reaches {lad}, as row "
                             f"{seen[lad]} does; refusing to choose a winner")
        seen[lad] = r["row"]
        cells = r["cells"]
        vals = {}
        for i, pathway, sub in REG02_PATHWAYS:
            vals[(pathway, sub)] = reg02_cell(
                cells[i], f"{where} {REG02_HEADERS[3 + i]}")
        population = _whole(cells[REG02_POPULATION], where, "Population")
        pct = ratio_cell(cells[REG02_PERCENT], where)
        suppressed = [pw for (pw, sub), (v, _) in vals.items()
                      if sub == "total" and pw != "all_pathways" and v is None]
        for _, pathway, sub in REG02_PATHWAYS:
            people, marker = vals[(pathway, sub)]
            top = pathway == "all_pathways" and sub == "total"
            if top and suppressed:
                marker = LOWER_BOUND.format(", ".join(suppressed))
            out.append({"lad24cd": lad, "pathway": pathway,
                        "sub_pathway": sub, "period_ending": p,
                        "people": people, "suppressed": people is None,
                        "source_marker": marker, "population": population,
                        "percentage_of_population": pct if top else None,
                        "published_la_name": r["la"]})
    return p, out


# ---------------------------------------------------------------------------
# Reconciliation, from the files themselves
# ---------------------------------------------------------------------------


def _data_sums(file) -> tuple:
    """({date: {region (case-folded): people}}, {date: people}) of the
    Asy_D11 data sheet."""
    by_reg, tot = {}, {}
    for r in file["rows"]:
        d = r["date"]
        a = by_reg.setdefault(d, {})
        k = _fold(r["region"])
        a[k] = a.get(k, 0) + r["people"]  # not a source value
        tot[d] = tot.get(d, 0) + r["people"]  # not a source value
    return by_reg, tot


def _people(records) -> int:
    return sum(r["people"] for r in records)


def reconcile_d11(file, records) -> list:
    """Problems (empty = fine): for each quarter the pivot caches, each
    region row and the Grand Total equal the data sheet's sums (regions
    compared case-insensitively; a pivot 0 with no data rows agrees); and
    for each period of records (build_d11_records' by_period) England +
    non-England + unallocated equal the data sheet's total for the date."""
    out = []
    by_reg, tot = _data_sums(file)
    piv = file["pivot"]
    for j, d in enumerate(piv["dates"]):
        data = by_reg.get(d, {})
        labels = set()
        for label, vals in piv["rows"].items():
            v = vals[j] if j < len(vals) else None
            if label == GRAND_TOTAL:
                if v != tot.get(d):
                    out.append(f"{d} {GRAND_TOTAL}: the pivot says {v}, the "
                               f"data sheet sums to {tot.get(d)}")
                continue
            k = _fold(label)
            labels.add(k)
            s = data.get(k)
            if s is None:
                if v is not None and v != 0:   # a pivot 0 with no rows agrees
                    out.append(f"{d} {label}: the pivot says {v}, the data "
                               "sheet has no rows")
            elif s != v:
                out.append(f"{d} {label}: the pivot says {v}, the data sheet "
                           f"sums to {s}")
        for k in sorted(set(data) - labels):
            out.append(f"{d} region {k!r}: {data[k]} people in the data sheet, "
                       "no pivot row")
    for p, t in sorted(records.items()):
        d = date.fromisoformat(p)
        s = sum(_people(t[x]) for x in D11_TAGS)
        if s != tot.get(d):
            out.append(f"{p}: England + non-England + unallocated = {s}, the "
                       f"data sheet total is {tot.get(d)}")
    return out


def reconcile_d09(d09, records, d11_cover=None) -> list:
    """Problems (empty = fine): Asy_D09's cover year-ending equals
    Asy_D11's (when d11_cover is given); for each period of records,
    England (la_asylum_support) equals Asy_D09's English regions and
    unallocated equals Asy_D09's 'N/A ...' and 'Unknown' region rows (region
    names compared case-insensitively)."""
    out = []
    if d11_cover is not None and \
            d09["cover"]["year_ending"] != d11_cover["year_ending"]:
        out.append(f"Asy_D09 is {d09['cover']['label']}, Asy_D11 is "
                   f"{d11_cover['label']}")
    eng = {_fold(x) for x in ENGLISH_REGIONS}
    for p, t in sorted(records.items()):
        d = date.fromisoformat(p)
        regs = d09["totals"].get(d)
        if regs is None:
            out.append(f"{p}: not in Asy_D09")
            continue
        e = sum(v for k, v in regs.items() if k in eng)
        u = sum(v for k, v in regs.items()
                if k.startswith("n/a") or k == "unknown")
        if e != _people(t["support"]):
            out.append(f"{p}: England {_people(t['support'])}, Asy_D09 "
                       f"English regions {e}")
        if u != _people(t["unallocated"]):
            out.append(f"{p}: unallocated {_people(t['unallocated'])}, "
                       f"Asy_D09 N/A and Unknown regions {u}")
    return out


def _cell_of(records) -> dict:
    return {(r["lad24cd"], r["pathway"], r["sub_pathway"]): r["people"]
            for r in records}


def d11_la_totals(records) -> dict:
    """{lad24cd: people} of support records."""
    out = {}
    for r in records:
        out[r["lad24cd"]] = out.get(r["lad24cd"], 0) + r["people"]  # not a source value
    return out


def reconcile_reg02(file, records, d11_totals) -> list:
    """Problems (empty = fine) of one Reg_02 period's records: exactly
    REG02_LAS English authorities; each pathway total equals its 'of which'
    columns where all are published; all_pathways equals the published
    pathway totals (a '*' pathway is excluded from it, note 9); and, when
    d11_totals ({lad24cd: Asy_D11 people} of the same period) is given,
    each authority's supported_asylum total equals it (an authority Asy_D11
    does not publish must show 0)."""
    out = []
    cell = _cell_of(records)
    las = sorted({r["lad24cd"] for r in records})
    want = globals()["REG02_LAS"]
    if len(las) != want:
        out.append(f"{len(las)} English authorities, expected {want}")
    subs = {"afghan_resettlement": ("transitional", "settled_la_housing",
                                    "settled_prs_housing"),
            "supported_asylum": ("initial_accommodation", "dispersal",
                                 "contingency", "other", "subsistence_only")}
    for la in las:
        for pw, parts in subs.items():
            tot = cell.get((la, pw, "total"))
            vals = [cell.get((la, pw, s)) for s in parts]
            if tot is not None and None not in vals and tot != sum(vals):
                out.append(f"{la} {pw}: total {tot}, its parts sum to "
                           f"{sum(vals)}")
        allp = cell.get((la, "all_pathways", "total"))
        tops = [cell.get((la, pw, "total")) for pw in
                ("homes_for_ukraine", "afghan_resettlement",
                 "supported_asylum")]
        pub = [v for v in tops if v is not None]
        if allp is not None and allp != sum(pub):
            out.append(f"{la} all_pathways: total {allp}, the published "
                       f"pathway totals sum to {sum(pub)}")
    if d11_totals is not None:
        for la in las:
            reg = cell.get((la, "supported_asylum", "total"))
            d = d11_totals.get(la)
            if reg is None:
                continue
            if d is None:
                if reg != 0:
                    out.append(f"{la}: Reg_02 supported_asylum {reg}, Asy_D11 "
                               "publishes no row")
            elif reg != d:
                out.append(f"{la}: Reg_02 supported_asylum {reg}, Asy_D11 {d}")
        for la in sorted(set(d11_totals) - set(las)):
            out.append(f"{la}: in Asy_D11 ({d11_totals[la]}), not in Reg_02")
    return out


# ---------------------------------------------------------------------------
# Geography
# ---------------------------------------------------------------------------


def resolve_codes(cur, codes, by_period) -> tuple:
    """({publisher code: lad24cd}, problems). geography.resolve(cur, '6',
    codes, by_period=by_period) (Barnsley/Sheffield, declared 'mixed': one
    period carrying both forms of an area is a problem); then la_code_lookup
    'new_unitary' rows with exactly one target (a code with two targets is a
    problem); then load_checks.check_codes; the resolved code must be in
    la_boundaries, else 'UNEXPLAINED'."""
    codes = set(codes)
    for cs in (by_period or {}).values():
        codes |= set(cs)
    recode, problems = geography.resolve(cur, RUN_SOURCE, codes,
                                         by_period=by_period)
    problems = list(problems)
    cur.execute("SELECT old_code, new_code FROM public.la_code_lookup "
                "WHERE change_type = 'new_unitary'")
    succ = {}
    for old, new in cur.fetchall():
        succ.setdefault(old, set()).add(new)
    cur.execute("SELECT lad24cd FROM public.la_boundaries")
    valid = {r[0] for r in cur.fetchall()}
    out = {}
    for c in sorted(codes):
        r = geography.canonical(c, recode)
        if r not in valid and r in succ:
            if len(succ[r]) != 1:
                problems.append(f"{c}: la_code_lookup new_unitary has "
                                f"{len(succ[r])} targets {sorted(succ[r])}; "
                                "refusing to choose one")
                continue
            r = next(iter(succ[r]))
        if r not in valid:
            problems.append(f"UNEXPLAINED {c}: not in la_boundaries and not a "
                            "single-target la_code_lookup new_unitary or recode")
            continue
        out[c] = r
    problems += [x for x in load_checks.check_codes(cur, codes)
                 if not any(x.split()[-1] in y for y in problems)]
    return out, problems


# ---------------------------------------------------------------------------
# Ranks, ledger and planning
# ---------------------------------------------------------------------------

_RANK_RE = re.compile(r"\byear ending ([A-Z][a-z]+ [0-9]{4}); published "
                      r"([0-9]{4}-[0-9]{2}-[0-9]{2})\b")
_LEDGER_RE = re.compile(r"(.*) \(year ending ([A-Z][a-z]+ [0-9]{4}); published "
                        r"([0-9]{4}-[0-9]{2}-[0-9]{2}); next update "
                        r"([0-9]{4}-[0-9]{2}-[0-9]{2}|-); ([^()]*)\)")


def release_rank(source_file_or_cover) -> "tuple | None":
    """(year-ending date, published date) of a file: a read cover's (or a
    read file's) own, or the 'year ending <Month yyyy>; published
    <yyyy-mm-dd>' an edition's or ledger row's source_file carries; None if
    it names none (a name, URL or media id is never read for a rank)."""
    x = source_file_or_cover
    if isinstance(x, dict):
        cov = x.get("cover", x)
        if "year_ending" in cov and "published" in cov:
            return (cov["year_ending"], cov["published"])
        return x.get("rank")
    mm = _RANK_RE.search(str(x or ""))
    if not mm:
        return None
    return (year_ending(mm.group(1)), date.fromisoformat(mm.group(2)))


def rank_text(rank) -> str:
    if not rank:
        return "no file rank"
    return f"{ye_text(rank[0])}, published {rank[1]}"


def edition_source(where, cover, note="") -> str:
    """An edition's source_file: where the file was read, its own
    year-ending and published date (the rank), and a note."""
    return (f"{where}; {cover['label']}; published {cover['published']}"
            + (f"; {note}" if note else ""))


def ledger_source(where, cover, periods_by_tag) -> str:
    """A file's ledger source_file: where it was read (the final URL, the
    path of a --file, or file:<name>), its own year-ending, published date
    and next update, and the periods it states for each table it feeds, so a
    later run knows them without reading the file."""
    nu = cover.get("next_update")
    parts = "; ".join(f"{t} {ranges_text(ps)}"
                      for t, ps in periods_by_tag.items())
    return (f"{where} ({cover['label']}; published {cover['published']}; "
            f"next update {nu.isoformat() if nu else '-'}; {parts})")


def parse_ledger_source(s) -> "tuple | None":
    """(where, rank, next update or None, {tag: [periods]}) of a ledger
    source_file, or None."""
    mm = _LEDGER_RE.fullmatch(str(s or ""))
    if not mm:
        return None
    rank = (year_ending(mm.group(2)), date.fromisoformat(mm.group(3)))
    nu = None if mm.group(4) == "-" else date.fromisoformat(mm.group(4))
    tags = {}
    for part in mm.group(5).split("; "):
        tag, _, rng = part.partition(" ")
        tags[tag] = parse_ranges(rng)
    return mm.group(1), rank, nu, tags


def content_sha256(path) -> str:
    """SHA-256 of the file's bytes."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _cell_text(v) -> str:
    if v is None:
        return "\\N"
    if isinstance(v, bool):
        return "true" if v else "false"
    return str(v)


def rows_content_sha(spec, records) -> str:
    """sha256 of one period's content in a spec: one line per record sorted
    by key, the key columns then the compared columns (spec.compare_cols)
    joined by '|', NULL as \\N, booleans true/false, lines joined by LF,
    UTF-8. Provenance (source_edition, source_file, published_date) is not
    in it, so it never makes a new edition."""
    kc, cc = tuple(spec.key_cols), tuple(spec.compare_cols)
    lines = sorted("|".join(_cell_text(r.get(c)) for c in kc + cc)
                   for r in records)
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


def period_ranks(tips_by_tag, checked_by_tag) -> dict:
    """{held period: the newest file rank that stated it}, over every table:
    the tip editions' ranks and every ledger row's."""
    out = {}
    for tag, tips in tips_by_tag.items():
        for p, t in tips.items():
            r = t.get("rank")
            out.setdefault(p, None)
            if r is not None and (out[p] is None or r > out[p]):
                out[p] = r
    for tag, checked in checked_by_tag.items():
        for p, pairs in checked.items():
            if p not in out:
                continue
            for s, _ in pairs:
                r = release_rank(s)
                if r is not None and (out.get(p) is None or r > out[p]):
                    out[p] = r
    return out


def ledger_complete(checked, where, sha, ranks, *, allow_older=False,
                    stranded=()) -> "list | None":
    """The periods the ledgers record this file's (where, sha256) pair for,
    if that covers every period the file needs in every table it feeds (its
    ledger source declares them per table); else None (the file is read and
    each period planned). checked: {tag: {period: {(source, sha)}}}. A held
    period whose rank (ranks) is newer than the file's is not needed (skipped
    as older) unless allow_older; a stranded period (editions, no live rows)
    always makes the file be read. So a run that failed or was rejected
    part-way leaves a period without its ledger rows and the rerun reads the
    file again."""
    have, declared = {}, None
    for tag, per in checked.items():
        for p, pairs in per.items():
            for s, h in pairs:
                parsed = parse_ledger_source(s)
                if h == sha and parsed and parsed[0] == where:
                    have.setdefault(tag, set()).add(str(p))
                    declared = parsed
    if declared is None:
        return None
    _, rank, _, needs = declared
    total = set()
    for tag, ps in needs.items():
        need = set(ps)
        if not allow_older:
            need = {p for p in need if not (ranks.get(p) is not None
                                            and ranks[p] > rank)}
        if need & set(stranded):
            return None
        if not need <= have.get(tag, set()):
            return None
        total |= need
    if not total:
        return None
    return sorted(total)


def plan_periods(periods, rank, ranks, checked, source, sha, *, recheck,
                 allow_older, needs=None) -> tuple:
    """(new, held, skipped). periods: the file's periods; rank: its own;
    ranks: {held period: period_ranks}; checked: {tag: {period: {(source,
    sha)}}} (stranded periods left out by the caller); needs: {period:
    [tags with records]} (default every tag of checked). new: not held;
    held: compared with the tips; skipped {period: reason}: held periods
    whose rank is newer than the file's (unless allow_older: the older-file
    guard, per period, on every path) or whose (source, sha) every ledger
    the period needs already records (unless recheck is given: then every
    period not older is compared)."""
    periods = sorted(str(p) for p in periods)
    new, held, skipped = [], [], {}
    for p in periods:
        if p not in ranks:
            new.append(p)
            continue
        t = ranks[p]
        if t is not None and t > rank and not allow_older:
            skipped[p] = (f"older: held from a file of {rank_text(t)}, this "
                          f"file is {rank_text(rank)}")
            continue
        tags = (needs or {}).get(p, list(checked))
        if recheck is None and tags and all(
                (source, sha) in checked.get(tag, {}).get(p, ())
                for tag in tags):
            skipped[p] = "checked: this file is already in the ledger"
            continue
        held.append(p)
    return new, held, skipped


# ---------------------------------------------------------------------------
# Stop conditions
# ---------------------------------------------------------------------------

AREA_KEY = {"support": "lad24cd", "non_england": "lad_code",
            "unallocated": "na_reason", "groups": "lad24cd"}
KEY_COLS = {"support": ("lad24cd", "support_type", "accommodation_type"),
            "unallocated": ("support_type", "accommodation_type", "na_reason"),
            "non_england": ("lad_code", "support_type", "accommodation_type"),
            "groups": ("lad24cd", "pathway", "sub_pathway")}


def _area_totals(records, tag) -> dict:
    """{area: people} of an Asy_D11 table's records (people never NULL)."""
    out = {}
    k = AREA_KEY[tag]
    for r in records or ():
        out[r[k]] = out.get(r[k], 0) + r["people"]  # not a source value
    return out


def _over_pct(old, new, pct) -> bool:
    """True when new moves from old by more than pct percent (exact)."""
    if old == new:
        return False
    if not old:
        return True
    return abs(new - old) * 100 > pct * old


def _rkey(r, tag):
    return tuple(r[c] for c in KEY_COLS[tag])


def period_problems(new, tip, prev, *, table, kind,
                    acknowledged=False) -> list:
    """The stop conditions of one period of one table (empty = fine). new:
    the file's records; tip: (kind 'revised') the tip's records of the same
    period; prev: (kind 'new', table 'support') the latest earlier held
    period's records, or None. Problems starting PARTIAL FILE are never
    released; the others are dropped when acknowledged (--acknowledge
    PERIOD).

    support, new: England total moving by more than NEW_TOTAL_PCT, the
    authority count by more than NEW_AREAS_DELTA, or any authority's total
    (on both sides) by more than NEW_AREA_ABS people. support, non_england
    and unallocated, revised: the table's total moving by more than
    REV_TOTAL_PCT (each table within it keeps the UK total within it); more
    than REV_MAX_AREAS authorities' totals changing; any authority's total
    moving by more than REV_AREA_ABS; more than REV_ROWS_PCT of the tip's
    rows added, or removed; PARTIAL FILE: fewer authorities than the tip by
    more than PARTIAL_AREAS, or fewer rows than the tip by more than
    REV_ROWS_PCT. groups: not exactly REG02_LAS authorities (PARTIAL FILE
    when fewer); revised: any pathway's England total (over authorities
    published on both sides) moving by more than REG02_REV_PCT, or more
    than REG02_REV_MAX_AREAS authorities changing."""
    g = globals()
    soft, hard = [], []
    if table == "groups":
        n, want = len({r["lad24cd"] for r in new}), g["REG02_LAS"]
        if n < want:
            hard.append(f"{PARTIAL}: {n} English authorities, expected "
                        f"{want}")
        elif n > want:
            hard.append(f"{n} English authorities, expected {want}")
        if kind == "revised":
            tk = {_rkey(r, table): r for r in tip or ()}
            nk = {_rkey(r, table): r for r in new}
            common = sorted(set(tk) & set(nk))
            for pw in ("homes_for_ukraine", "afghan_resettlement",
                       "supported_asylum", "all_pathways"):
                both = [k for k in common if k[1] == pw and k[2] == "total"
                        and tk[k]["people"] is not None
                        and nk[k]["people"] is not None]
                a = sum(tk[k]["people"] for k in both)
                b = sum(nk[k]["people"] for k in both)
                if both and _over_pct(a, b, g["REG02_REV_PCT"]):
                    soft.append(f"England {pw} total {a:,} -> {b:,} (limit "
                                f"{g['REG02_REV_PCT']}%)")
            cols = ("people", "suppressed", "source_marker", "population",
                    "percentage_of_population", "published_la_name")
            changed = {k[0] for k in set(tk) ^ set(nk)} | {
                k[0] for k in common
                if any(tk[k].get(c) != nk[k].get(c) for c in cols)}
            if len(changed) > g["REG02_REV_MAX_AREAS"]:
                soft.append(f"{len(changed)} authorities change (limit "
                            f"{g['REG02_REV_MAX_AREAS']})")
        return hard + ([] if acknowledged else soft)
    an = _area_totals(new, table)
    if kind == "new":
        if table == "support" and prev:
            ap = _area_totals(prev, table)
            a, b = _people(prev), _people(new)
            if _over_pct(a, b, g["NEW_TOTAL_PCT"]):
                soft.append(f"England total {a:,} -> {b:,} against the "
                            f"previous period (limit {g['NEW_TOTAL_PCT']}%)")
            if abs(len(an) - len(ap)) > g["NEW_AREAS_DELTA"]:
                soft.append(f"authorities {len(ap)} -> {len(an)} against the "
                            f"previous period (limit +/-{g['NEW_AREAS_DELTA']})")
            big = [f"{k} {ap[k]:,}->{an[k]:,}" for k in sorted(set(ap) & set(an))
                   if abs(an[k] - ap[k]) > g["NEW_AREA_ABS"]]
            if big:
                soft.append(f"{len(big)} authorit(ies) move by more than "
                            f"{g['NEW_AREA_ABS']:,} people against the "
                            f"previous period: {', '.join(big[:6])}")
        return [] if acknowledged else soft
    tip = tip or []
    at = _area_totals(tip, table)
    a, b = _people(tip), _people(new)
    if len(at) - len(an) > g["PARTIAL_AREAS"]:
        hard.append(f"{PARTIAL}: {len(an)} authorities, the held edition has "
                    f"{len(at)} (more than {g['PARTIAL_AREAS']} fewer; a "
                    "partial file never replaces a fuller edition)")
    # The row-count partial rule is for the big table only. The unallocated
    # table has 1 to 4 rows a period and non_england about 65: a publisher
    # reassigning an 'Unknown' row to an authority is a key change, released
    # by --acknowledge through the "rows removed" check below, not a partial
    # file. A genuinely partial non_england file still fails PARTIAL_AREAS.
    if table == "support" and len(tip) and \
            (len(tip) - len(new)) * 100 > g["REV_ROWS_PCT"] * len(tip):
        hard.append(f"{PARTIAL}: {len(new)} rows, the held edition has "
                    f"{len(tip)} (more than {g['REV_ROWS_PCT']}% fewer; a "
                    "partial file never replaces a fuller edition)")
    if _over_pct(a, b, g["REV_TOTAL_PCT"]):
        soft.append(f"total {a:,} -> {b:,} (limit {g['REV_TOTAL_PCT']}%)")
    changed = [k for k in sorted(set(at) | set(an), key=str)
               if at.get(k) != an.get(k)]
    if len(changed) > g["REV_MAX_AREAS"]:
        soft.append(f"{len(changed)} authorities' totals change (limit "
                    f"{g['REV_MAX_AREAS']})")
    big = [f"{k} {at[k]:,}->{an[k]:,}" for k in sorted(set(at) & set(an),
                                                       key=str)
           if abs(an[k] - at[k]) > g["REV_AREA_ABS"]]
    if big:
        soft.append(f"{len(big)} authorit(ies) move by more than "
                    f"{g['REV_AREA_ABS']:,} people: {', '.join(big[:6])}")
    tk = {_rkey(r, table) for r in tip}
    nk = {_rkey(r, table) for r in new}
    for what, keys in (("added", nk - tk), ("removed", tk - nk)):
        if tk and len(keys) * 100 > g["REV_ROWS_PCT"] * len(tk):
            soft.append(f"{len(keys)} of {len(tk)} rows {what} (limit "
                        f"{g['REV_ROWS_PCT']}%): "
                        + ", ".join("/".join(map(str, k))
                                    for k in sorted(keys)[:4]))
    return hard + ([] if acknowledged else soft)


def blank_flips(new, tip) -> list:
    """[(key, column, tip value, new value)] of people cells, on keys in
    both, going from 0 to NULL ('*') or from NULL to 0 against the tip (rule
    1.10), sorted."""
    def keyed(rs):
        return {tuple(r[c] for c in KEY_COLS["groups"]): r for r in rs or ()}
    nk, tk = keyed(new), keyed(tip)
    out = []
    for k in sorted(set(nk) & set(tk)):
        a, b = tk[k].get("people"), nk[k].get("people")
        if (a == 0 and b is None) or (a is None and b == 0):
            out.append((k, "people", a, b))
    return out


def blank_report(records, tag) -> dict:
    """{column: {'null', 'star', 'zero'}} of a table's records (rule 1.10)."""
    cols = {"support": ("people", "source_marker"),
            "unallocated": ("people",), "non_england": ("people",),
            "groups": ("people", "population", "percentage_of_population")}[tag]
    out = {}
    for c in cols:
        out[c] = {"null": sum(1 for r in records if r.get(c) is None),
                  "zero": sum(1 for r in records if r.get(c) is not None
                              and not isinstance(r.get(c), str)
                              and r[c] == 0)}
        if c == "people":
            out[c]["star"] = sum(1 for r in records if r.get(c) is None
                                 and r.get("source_marker") == "*")
    return out


# ---------------------------------------------------------------------------
# Specs and profiles
# ---------------------------------------------------------------------------


def tip_row_count(editions_table: str):
    """expected_rows_per_period for status: the tip edition's row count (a
    quarter's authorities are its own)."""
    core._ident(editions_table)

    def count(cur, period):
        cur.execute(f"SELECT DISTINCT edition, supersedes FROM "
                    f"public.{editions_table} WHERE period_ending = %s",
                    (period,))
        try:
            tip = core._chain_tip(cur.fetchall(), str(period))
        except (LookupError, ValueError):
            return None
        cur.execute(f"SELECT COUNT(*) FROM public.{editions_table} WHERE "
                    "period_ending = %s AND edition = %s", (period, tip))
        return cur.fetchone()[0]
    return count


_PROV = (("source_edition", "release_label"), ("loaded_at", "loaded_at"))
_TEXT = "text NOT NULL"
_PERIOD = ("period_ending", "date NOT NULL")
COMPARED = {
    "support": ("people", "published_la_name", "source_marker"),
    "unallocated": ("people",),
    "non_england": ("people", "country", "published_la_name"),
    "groups": ("people", "suppressed", "source_marker", "population",
               "percentage_of_population", "published_la_name"),
}


def _spec(tag, name, live, value_cols, extra_cols, fk):
    keys = KEY_COLS[tag]
    return core.EditionSpec(
        name=name, live_table=live, editions_table=f"{live}_editions",
        key_cols=keys, period_col="period_ending",
        value_cols=value_cols, extra_cols=extra_cols,
        refresh_cols=COMPARED[tag] + ("source_edition", "loaded_at"),
        refresh_from=_PROV, whole_period_cols=("source_edition",),
        key_types=tuple((k, _TEXT) for k in keys) + (_PERIOD,),
        fk_la_boundaries=fk, refresh_key_changes=True,
        expected_rows_per_period=tip_row_count(f"{live}_editions"))


SPEC_SUPPORT = _spec("support", "s6_support", "la_asylum_support",
                     (("people", "integer NOT NULL"),),
                     (("published_la_name", "text NOT NULL"),
                      ("source_marker", "text")), True)
SPEC_UNALLOCATED = _spec("unallocated", "s6_unallocated",
                         "la_asylum_support_unallocated",
                         (("people", "integer NOT NULL"),), (), False)
SPEC_NON_ENGLAND = _spec("non_england", "s6_non_england",
                         "asylum_support_non_england",
                         (("people", "integer NOT NULL"),),
                         (("country", "text NOT NULL"),
                          ("published_la_name", "text NOT NULL")), False)
SPEC_GROUPS = _spec("groups", "s6_groups", "la_immigration_groups",
                    (("people", "integer"), ("suppressed", "boolean NOT NULL"),
                     ("source_marker", "text"), ("population", "integer"),
                     ("percentage_of_population", "numeric(8,4)")),
                    (("published_la_name", "text NOT NULL"),), True)
SPEC_NAMES = {"support": "SPEC_SUPPORT", "unallocated": "SPEC_UNALLOCATED",
              "non_england": "SPEC_NON_ENGLAND", "groups": "SPEC_GROUPS"}


def _no_label(fetched_on) -> str:
    raise ValueError("an S6 release label comes from the file's cover; store "
                     "through apply_d11_period(...) or apply_reg02_period(...)")


def _profile(tag, spec, heading, savepoint):
    return pe.Profile(
        spec=spec, value_cols=COMPARED[tag], run_agent=RUN_AGENT,
        run_source=RUN_SOURCE, heading=heading,
        default_source_file="Home Office immigration system statistics",
        expected_areas=None, release_label=_no_label,
        content_sha256=lambda r, _t=tag: rows_content_sha(
            globals()[SPEC_NAMES[_t]], r),
        file_checks=True, savepoint=savepoint,
        example_label=f"({', '.join(COMPARED[tag])})")


PROFILE_SUPPORT = _profile("support", SPEC_SUPPORT, "S6 Asy_D11 England "
                           "(la_asylum_support)", "s6_d11_period")
PROFILE_UNALLOCATED = _profile("unallocated", SPEC_UNALLOCATED, "S6 Asy_D11 "
                               "unallocated (la_asylum_support_unallocated)",
                               "s6_d11_period")
PROFILE_NON_ENGLAND = _profile("non_england", SPEC_NON_ENGLAND, "S6 Asy_D11 "
                               "Scotland, Wales, Northern Ireland "
                               "(asylum_support_non_england)", "s6_d11_period")
PROFILE_GROUPS = _profile("groups", SPEC_GROUPS, "S6 Reg_02 immigration "
                          "groups (la_immigration_groups)", "s6_reg02_period")
PROFILE_NAMES = {"support": "PROFILE_SUPPORT",
                 "unallocated": "PROFILE_UNALLOCATED",
                 "non_england": "PROFILE_NON_ENGLAND",
                 "groups": "PROFILE_GROUPS"}


def profile(tag) -> pe.Profile:
    """The tag's profile bound to the current spec (looked up when called,
    so tests can swap the specs)."""
    g = globals()
    return dataclasses.replace(g[PROFILE_NAMES[tag]],
                               spec=g[SPEC_NAMES[tag]])


def profiles(tags=TAGS) -> dict:
    return {t: profile(t) for t in tags}


def spec_of(table) -> tuple:
    """(tag, spec) of a live table name or spec name."""
    for t in TAGS:
        s = globals()[SPEC_NAMES[t]]
        if table in (s.live_table, s.name, t):
            return t, s
    raise ValueError(f"{table!r} is not one of "
                     f"{[globals()[SPEC_NAMES[t]].live_table for t in TAGS]}")


def _late(name: str):
    """A call through this module's attribute `name`, looked up when called
    (so mock.patch.object(s6_asylum_editions, name, ...) reaches the
    engine)."""
    return lambda *a, **kw: globals()[name](*a, **kw)


def create_all(cur) -> None:
    """ddl: the four editions tables, their append-only triggers and the
    four file-check ledgers. Idempotent; no live table is touched."""
    for prof in profiles().values():
        core.create_schema(cur, prof.spec)
        pe.create_file_checks(cur, prof)


def insert_live(cur, profile, period: str, records: list) -> None:
    """A new period's live rows: the key, the period, the compared columns
    and source_edition, the same for every row: the edition's release_label
    (profile.release_label, set by the apply functions); loaded_at takes its
    default."""
    from psycopg2.extras import execute_values
    spec = profile.spec
    label = profile.release_label(None)
    cols = (tuple(spec.key_cols) + (spec.period_col,)
            + tuple(profile.value_cols) + ("source_edition",))
    execute_values(
        cur, f"INSERT INTO public.{spec.live_table} ({', '.join(cols)}) "
        "VALUES %s",
        [tuple(r[k] for k in spec.key_cols) + (period,)
         + tuple(r.get(c) for c in profile.value_cols) + (label,)
         for r in records],
        page_size=1000)


@dataclass(frozen=True)
class PeriodInfo:
    """What a period's editions record: source_file, the release label (the
    cover's 'year ending ...'), the published date (the cover's) and the
    file read ((ledger source, sha256),)."""
    source_file: str
    label: str
    published: date
    files: tuple


def _tip_label(cur, spec, period) -> str:
    tip = core.chain_tip(cur, spec, period)
    cur.execute(f"SELECT DISTINCT release_label FROM public.{spec.editions_table}"
                f" WHERE period_ending = %s AND edition = %s", (period, tip))
    rows = cur.fetchall()
    if len(rows) != 1:
        halt(f"{spec.editions_table} {period}: edition {tip} carries "
             f"{len(rows)} release labels")
    return rows[0][0]


def _has_edition(cur, spec, period) -> bool:
    cur.execute(f"SELECT 1 FROM public.{spec.editions_table} WHERE "
                "period_ending = %s LIMIT 1", (period,))
    return cur.fetchone() is not None


def _apply_one(cur, prof, period, records, info) -> str:
    """pe.apply_period for one table's period, then its ledger rows, inside
    the caller's savepoint."""
    kind = pe.classify_period(cur, prof, period, records)
    label = (_tip_label(cur, prof.spec, period) if kind == LIVE_MISSING
             else info.label)
    p2 = dataclasses.replace(prof, release_label=lambda d: label)
    kind = pe.apply_period(cur, p2, period, records, fetched_on=info.published,
                           source_file=info.source_file,
                           classify=lambda *a, **kw: kind,
                           insert=_late("insert_live"))
    tip = core.chain_tip(cur, p2.spec, period)
    for src, sha in info.files:
        pe.record_file_check(cur, p2, period, src, sha, kind, tip)
    return kind


def apply_d11_period(cur, profs, period, recs_by_tag, *, info) -> dict:
    """The three Asy_D11 tables' period in the caller's savepoint
    (s6_d11_period): for each table with records, pe.apply_period (edition 1
    and live rows when new; the next edition when revised; live rows only
    when stranded; nothing when unchanged) and its ledger rows; a table with
    no records stores nothing (it must hold no edition for the period, else
    halt: rows cannot vanish silently). A failure anywhere halts and the
    savepoint rolls back all three. {tag: kind}."""
    kinds = {}
    for tag in D11_TAGS:
        prof = profs[tag]
        recs = recs_by_tag.get(tag) or []
        if not recs:
            if _has_edition(cur, prof.spec, period):
                halt(f"{period}: {prof.spec.live_table} holds an edition but "
                     "the file has no rows for it; rows cannot vanish "
                     "silently; rolled back")
            continue
        kinds[tag] = _apply_one(cur, prof, period, recs, info)
    return kinds


def apply_reg02_period(cur, prof, period, records, *, info) -> str:
    """The Reg_02 period in the caller's savepoint (s6_reg02_period)."""
    return _apply_one(cur, prof, period, records, info)


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------


def _session(session):
    if session is not None:
        return session
    import requests
    return requests


def _get(session, url, **kw):
    try:
        r = session.get(url, headers={"User-Agent": USER_AGENT},
                        timeout=300, **kw)
        r.raise_for_status()
    except Exception as e:  # noqa: BLE001 (any failure halts)
        halt(f"GET {url} failed, nothing downloaded: {e}")
    return r


def fetch_json(path, session=None) -> dict:
    """The GOV.UK content API reply for a page path (a read-only GET)."""
    return json.loads(_get(_session(session), API + path).text)


def _abs_url(url) -> str:
    return url if urlparse(url).scheme else GOV_UK + url


EXPECTED_SHEETS = (D11_SHEET, D09_SHEET, REG02_SHEET)


def _opens(path, suffix, expect) -> None:
    """Raise unless the file opens as the expected workbook: an .xlsx or
    .ods carrying one of the expect sheets."""
    if suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            sheets = list(wb.sheetnames)
        finally:
            wb.close()
    elif suffix == ".ods":
        with zipfile.ZipFile(path) as z:
            if z.read("mimetype").strip() != ODS_MIME:
                raise ValueError("not an OpenDocument spreadsheet")
        import pandas as pd
        with pd.ExcelFile(path, engine="odf") as x:
            sheets = list(x.sheet_names)
    else:
        raise ValueError(f"not an .xlsx or .ods file ({suffix})")
    if not any(s in sheets for s in expect):
        raise ValueError(f"none of the sheets {list(expect)}; sheets {sheets}")


def fetch(url, dest, session=None, expect=EXPECTED_SHEETS) -> tuple:
    """Download url into the directory dest (redirects followed) and return
    (path to read, final URL). The file name is the final URL's. Checked
    before it is kept: more than MIN_FILE_BYTES, an .xlsx or .ods name, and
    it opens as a workbook with one of the expect sheets. A same-named file
    with the same content is kept as it is; one with different content is
    never replaced: the download is saved beside it as <stem>-<sha8><suffix>
    and that is the one read. An existing name is never a reason to skip the
    download."""
    dest = Path(dest)
    r = _get(_session(session), url)
    final = str(getattr(r, "url", None) or url)
    if final != url:
        print(f"redirected: {url} -> {final}")
    body = r.content
    if len(body) <= MIN_FILE_BYTES:
        halt(f"{final}: {len(body)} bytes, expected more than "
             f"{MIN_FILE_BYTES}; nothing written; saw {body[:100]!r}")
    name = Path(urlparse(final).path).name
    suffix = Path(name).suffix.lower()
    if suffix not in (".xlsx", ".ods"):
        halt(f"{final}: not an .xlsx or .ods file; nothing written")
    dest.mkdir(parents=True, exist_ok=True)
    target = dest / name
    new_sha = hashlib.sha256(body).hexdigest()
    tmp = dest / (Path(name).stem + ".download.tmp" + suffix)
    tmp.write_bytes(body)
    try:
        _opens(tmp, suffix, expect)
    except Exception as e:  # noqa: BLE001
        tmp.unlink()
        halt(f"{final}: does not open as the expected workbook ({e}); "
             "nothing kept")
    path = target
    if target.exists():
        old_sha = content_sha256(target)
        if old_sha == new_sha:
            tmp.unlink()
            print(f"{name} already held with the same content (sha256 "
                  f"{new_sha[:16]}); kept as it is")
        else:
            path = dest / f"{target.stem}-{new_sha[:8]}{target.suffix}"
            if path.exists() and content_sha256(path) == new_sha:
                tmp.unlink()
                print(f"NOTE: {name} differs from this download; {path.name} "
                      f"already holds it (sha256 {new_sha[:16]})")
            else:
                tmp.replace(path)
                print(f"NOTE: {name} already exists with different content "
                      f"(sha256 {old_sha[:16]}, new {new_sha[:16]}); kept it "
                      f"untouched and saved the new download as {path.name}")
    else:
        tmp.replace(target)
        print(f"downloaded {name}: {len(body):,} bytes, sha256 "
              f"{new_sha[:16]}")
    return path, final


def _file_source(path: Path) -> str:
    """Where a --file was read: the path under the repo (data/raw/...),
    else the absolute path."""
    p = Path(path).resolve()
    try:
        return p.relative_to(REPO).as_posix()
    except ValueError:
        return p.as_posix()


# ---------------------------------------------------------------------------
# Reading held data
# ---------------------------------------------------------------------------

table_exists = pe.table_exists


def _conn(writing):
    return pe.connect(writing)


def records(cur, spec, period, edition=None) -> list:
    """The period's rows as records (key, period as a string, the compared
    columns): the live table's (edition None) or the given edition's."""
    names = tuple(spec.key_cols) + tuple(spec.compare_cols)
    if edition is None:
        cur.execute(f"SELECT {', '.join(names)} FROM public.{spec.live_table} "
                    "WHERE period_ending = %s", (period,))
    else:
        cur.execute(f"SELECT {', '.join(names)} FROM "
                    f"public.{spec.editions_table} WHERE period_ending = %s "
                    "AND edition = %s", (period, edition))
    out = [dict(zip(names, r)) for r in cur.fetchall()]
    for r in out:
        r["period_ending"] = str(period)
    return out


def tip_info(cur, spec) -> dict:
    """{period (string): {edition, source_file, rank, label}} of every
    period with editions."""
    cur.execute(f"SELECT DISTINCT period_ending FROM "
                f"public.{spec.editions_table}")
    out = {}
    for (p,) in cur.fetchall():
        tip = core.latest_edition(cur, spec, p)
        cur.execute(f"SELECT DISTINCT source_file, release_label, "
                    f"published_date FROM public.{spec.editions_table} WHERE "
                    "period_ending = %s AND edition = %s", (p, tip))
        rows = cur.fetchall()
        sf, label, pub = rows[0] if len(rows) == 1 else (None, None, None)
        out[pe._p(p)] = {"edition": tip, "source_file": sf,
                         "rank": release_rank(sf), "label": label,
                         "published_date": pub}
    return out


def legacy_tips(cur, spec) -> dict:
    """Before the editions tables exist: {held live period: {rank}} from the
    live source_edition and the held files (LEGACY_FILES)."""
    labels = {}
    for f in globals()["LEGACY_FILES"].values():
        labels.setdefault(ye_text(f["rank"][0]), f["rank"])
    cur.execute(f"SELECT period_ending, array_agg(DISTINCT source_edition) "
                f"FROM public.{spec.live_table} GROUP BY 1")
    out = {}
    for p, labs in cur.fetchall():
        ranks = [labels.get(x) for x in labs or ()]
        ranks = [r for r in ranks if r is not None]
        out[pe._p(p)] = {"edition": None, "source_file": None,
                         "rank": max(ranks) if ranks else None,
                         "label": (labs or [None])[0]}
    return out


def ledger_checks(cur, prof) -> dict:
    """{period (string): {(source_file, file_sha256)}} of a ledger."""
    cur.execute(f"SELECT DISTINCT period_ending, source_file, file_sha256 "
                f"FROM public.{pe.file_checks_table(prof)}")
    out = {}
    for p, f, s in cur.fetchall():
        out.setdefault(pe._p(p), set()).add((f, s))
    return out


def _ready(cur, tags) -> dict:
    out = {}
    for t in tags:
        prof = profile(t)
        out[prof.spec.editions_table] = table_exists(cur,
                                                     prof.spec.editions_table)
        lt = pe.file_checks_table(prof)
        out[lt] = table_exists(cur, lt)
    return out


def held_next_update(checked_by_tag) -> "date | None":
    """The 'next update' of the newest file any ledger records."""
    best = None
    for per in checked_by_tag.values():
        for pairs in per.values():
            for s, _ in pairs:
                parsed = parse_ledger_source(s)
                if parsed and parsed[2] is not None and (
                        best is None or parsed[1] > best[0]):
                    best = (parsed[1], parsed[2])
    return best[1] if best else None


# ---------------------------------------------------------------------------
# ddl, status, run log, refresh-latest
# ---------------------------------------------------------------------------


def log_run(cur, rows_written: int, notes: str, started_at=None) -> None:
    """The pipeline_run_log row for a committed run (agent RUN_AGENT,
    source '6'). Called only on committed load, migrate-legacy and
    restore-edition runs, partial runs included."""
    pe.log_run(cur, profile("support"), rows_written, notes, started_at)


def cmd_ddl(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            for t, ok in _ready(cur, TAGS).items():
                print(f"{t}: {'exists' if ok else 'does not exist'}")
            if not writing:
                print("DRY RUN: nothing written (use --commit or --simulate)")
                return 0
            create_all(cur)
        pe.finish(conn, args, "ddl: four editions tables, their triggers and "
                              "four file-check ledgers present")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def status_lines(cur, prof) -> tuple:
    """(lines, ok): per period the tip's edition, release and published
    date, whether live source_edition is uniform and the tip's, and the
    ledger's latest file. ok is false when a period's live source_edition
    is not uniform, or uniform but not the tip's while live equals the
    tip."""
    s = prof.spec
    lines, ok = [], True
    lt = pe.file_checks_table(prof)
    for p, t in sorted(tip_info(cur, s).items()):
        line = (f"  {p}: tip edition {t['edition']} ({t['label']}; "
                f"{rank_text(t['rank'])}; published_date "
                f"{t['published_date']})")
        cur.execute(f"SELECT DISTINCT source_edition FROM public.{s.live_table}"
                    " WHERE period_ending = %s", (p,))
        livep = {r[0] for r in cur.fetchall()}
        current = core.rows_differing(cur, s, p, t["edition"]) == 0
        if livep == {t["label"]}:
            line += "; live source_edition uniform, the tip's"
        elif len(livep) == 1:
            line += "; live source_edition uniform, not the tip's"
            if current:
                ok = False
                line += " (ACTION NEEDED)"
            else:
                line += " (refresh-latest sets it)"
        elif livep:
            ok = False
            line += (f"; live source_edition not uniform ({len(livep)} values):"
                     " ACTION NEEDED")
        cur.execute(f"SELECT source_file, outcome FROM public.{lt} WHERE "
                    "period_ending = %s ORDER BY id DESC LIMIT 1", (p,))
        last = cur.fetchone()
        line += (f"; ledger latest: {last[1]} {last[0]}" if last
                 else "; no ledger row")
        lines.append(line)
    return lines, ok


def cmd_status(args) -> int:
    """status: what needs action; exit 1 if anything, or (with a clean
    message, nothing created) if an editions table or ledger does not
    exist."""
    conn = _conn(False)
    try:
        with conn.cursor() as cur:
            missing = [t for t, ok in _ready(cur, TAGS).items() if not ok]
            if missing:
                print(f"status: {', '.join(missing)} not present yet; run "
                      "`ddl --commit`, then `migrate-legacy --d11 FILE --reg02 "
                      "FILE --reg02 FILE --commit` (S6 has no sync-new)")
                return 1
            ok, texts, checked = True, [], {}
            for t, prof in profiles().items():
                st = pe.status(cur, prof)
                lines, prov_ok = status_lines(cur, prof)
                ok = ok and st["ok"] and prov_ok
                texts.append(pe.format_status(prof, st, lines))
                checked[t] = ledger_checks(cur, prof)
            nu = held_next_update(checked) or EXPECTED_NEXT_UPDATE
            texts.append(f"next expected release: {nu:%d %B %Y} (the held "
                         "file's 'Next update')")
    finally:
        conn.rollback()
        conn.close()
    print("\n".join(texts))
    return 0 if ok else 1


def cmd_refresh_latest(args) -> int:
    """refresh-latest [--accept-key-changes P] [--commit | --simulate]: the
    engine's refresh for each of the four specs in one transaction (preview
    of the rows and the keys added and removed; the refusals;
    editions_core.refresh_latest under its before/after guard;
    source_edition set on every row of each refreshed period and loaded_at
    copied), then load_checks.check_latest_equals_live. A period named in
    --accept-key-changes applies to every table whose plan adds or drops
    keys in it; a name matching none halts."""
    writing = args.commit or args.simulate
    named = tuple(args.accept_key_changes or ())
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            for t, ok in _ready(cur, TAGS).items():
                if not ok:
                    halt(f"{t} does not exist yet; run `ddl --commit`, then "
                         "migrate-legacy")
            plans, keyed_all = {}, set()
            for tag, prof in profiles().items():
                spec = prof.spec
                counts = core.refresh_counts(cur, spec)
                print(f"{spec.live_table}: rows refresh-latest would write: "
                      + (", ".join(f"{pe._p(p)}={n}" for p, n in counts.items())
                         or "none") + f" (total {sum(counts.values())})")
                plan, _ = core._plan(cur, spec)
                keyed = {pe._p(p): p for p in core.key_change_periods(plan)}
                keyed_all |= set(keyed)
                acc = tuple(keyed[n] for n in named if n in keyed)
                lines = core.key_change_lines(plan, acc, name=pe._p)
                print("  keys added and removed:" + ("" if lines else " none"))
                for line in lines:
                    print(f"    {line}")
                plans[tag] = acc
            stray = sorted(set(named) - keyed_all)
            if stray:
                halt(f"--accept-key-changes {stray}: no table adds or drops "
                     "keys in those periods, nothing to accept")
            if not writing:
                print("DRY RUN: nothing written (use --commit or --simulate)")
                return 0
            done = []
            for tag, prof in profiles().items():
                spec = prof.spec
                pe._refresh_guard(cur, spec, (), plans[tag])
                res = core.refresh_latest(cur, spec, (), plans[tag])
                if res["updated"]:
                    bad = load_checks.check_latest_equals_live(
                        cur, spec, sorted(res["updated"]))
                    if bad:
                        halt("refresh-latest: live differs from the latest "
                             "edition after the refresh, rolled back: "
                             + "; ".join(bad[:6]))
                keys = ""
                if res["inserted"] or res["deleted"]:
                    keys = (f" (key changes {sum(res['inserted'].values())} "
                            f"inserted, {sum(res['deleted'].values())} "
                            "deleted)")
                done.append(f"{spec.live_table} {res['rows']} rows in "
                            f"{sorted(pe._p(x) for x in res['updated'])}{keys}")
        pe.finish(conn, args, "refreshed: " + "; ".join(done) + "; "
                  "source_edition set on every row of each refreshed period; "
                  "before/after guard passed")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# load
# ---------------------------------------------------------------------------


def _end(conn, args, writing, rc) -> int:
    conn.rollback()  # commit mode has committed already
    if not writing:
        print("PREVIEW: nothing written to the database (use --commit or "
              "--simulate)")
    elif args.simulate:
        print("SIMULATION: ROLLED BACK (nothing persisted)")
    return rc


def _release_arg(text) -> date:
    try:
        return year_ending(text)
    except ValueError:
        raise argparse.ArgumentTypeError(f"not a release 'Month YYYY' ending "
                                         f"a quarter: {text!r}") from None


def _period_arg(text) -> str:
    try:
        return date.fromisoformat(str(text)).isoformat()
    except ValueError:
        raise argparse.ArgumentTypeError(f"not a period YYYY-MM-DD: "
                                         f"{text!r}") from None


def _read(fn, path):
    try:
        return fn(path)
    except ValueError as e:
        halt(f"{Path(path).name}: {e}; nothing stored")


def _print_page(page, path):
    print(f"page: {page.get('title')} ({GOV_UK}{path}); updated "
          f"{str(page.get('public_updated_at'))[:10]}")
    ch = change_history(page)
    if ch:
        print(f"  latest change note {ch[0][0]}: {ch[0][1][:300]}")


class Ctx:
    """The state one load run shares between its parts."""

    def __init__(self, args, conn, cur, has_ed, writing):
        self.args, self.conn, self.cur = args, conn, cur
        self.has_ed, self.writing = has_ed, writing
        self.d11_records = {}       # {period: support records} of this run
        self.newest_seen = {}
        self.held = {}
        self.checked_all = {}


def _state(ctx, tags) -> tuple:
    """(tips {tag: {period: info}}, checked {tag: {period: pairs}},
    stranded periods) of the tags' tables (live before the editions tables
    exist)."""
    cur = ctx.cur
    tips, checked, stranded = {}, {}, set()
    for t in tags:
        prof = profile(t)
        if ctx.has_ed:
            _, new_live, errors = core.latest_map(cur, prof.spec)
            if errors:
                halt("invalid edition chain: " + "; ".join(
                    f"{pe._p(p)}: {msg}" for p, msg in errors.items()))
            if new_live:
                halt(f"{prof.spec.live_table}: live periods with no editions "
                     f"{[pe._p(x) for x in new_live]}; run migrate-legacy "
                     "first")
            tips[t] = tip_info(cur, prof.spec)
            checked[t] = ledger_checks(cur, prof)
            stranded |= set(pe.live_missing_periods(cur, prof))
        else:
            tips[t] = legacy_tips(cur, prof.spec)
            checked[t] = {}
    return tips, checked, sorted(stranded)


def _tip_records(ctx, tag, period, tips) -> list:
    spec = profile(tag).spec
    if ctx.has_ed:
        return records(ctx.cur, spec, period, tips[tag][period]["edition"])
    return records(ctx.cur, spec, period)


def _resolve(cur, codes_by_period) -> dict:
    try:
        rmap, problems = resolve_codes(cur, set(), codes_by_period)
    except ValueError as e:
        halt(f"geography: {e}; nothing stored")
    if problems:
        halt("geography check failed, nothing stored: " + "; ".join(problems))
    recoded = sorted((c, r) for c, r in rmap.items() if r != c)
    if recoded:
        print("codes resolved (geography.resolve; la_code_lookup recode and "
              "single-target new_unitary): "
              + ", ".join(f"{c} -> {r}" for c, r in recoded))
    return rmap


def _report_blanks(tag, recs) -> None:
    rep = blank_report(recs, tag)
    print(f"  {profile(tag).spec.live_table}: {len(recs):,} rows; "
          + "; ".join(f"{c} NULL {v['null']}"
                      + (f" ('*' {v['star']})" if "star" in v else "")
                      + f", 0 {v['zero']}" for c, v in rep.items()))


def _d11_source(ctx) -> dict:
    """The Asy_D11 (and Asy_D09) files the run reads."""
    args = ctx.args
    out = {"path": None, "where": None, "how": None, "page": None,
           "d09_path": None, "d09_att": None, "att": None}
    if args.file:
        path = Path(args.file)
        out.update(path=path, how="local file")
        if args.no_page:
            out.update(where=f"file:{path.name}",
                       how="local file, --no-page: the page was not read")
            print(f"NOTE: --no-page: the GOV.UK page is not read; source_file "
                  f"is recorded as file:{path.name}")
        else:
            out["where"] = _file_source(path)
            page = fetch_json(TABLES_PATH)
            _print_page(page, TABLES_PATH)
            try:
                att = asy_d11_attachment(page)
            except ValueError as e:
                halt(f"{e}; give --no-page to load the file without reading "
                     "the page (recorded as such)")
            out["page"] = att
            print(f"  newest Asy_D11 the page lists: {att['title']!r}")
        if args.d09_file:
            out["d09_path"] = Path(args.d09_file)
        return out
    page = fetch_json(TABLES_PATH)
    _print_page(page, TABLES_PATH)
    try:
        att = asy_d11_attachment(page)
        d09 = asy_d09_attachment(page) if not args.no_d09 and \
            not args.d09_file else None
    except ValueError as e:
        halt(str(e))
    print(f"  newest Asy_D11 the page lists: {att['title']!r} {att['url']}")
    if d09:
        print(f"  newest Asy_D09 the page lists: {d09['title']!r} {d09['url']}")
    ctx.newest_seen["Asy_D11"] = att["year_ending"].isoformat()
    if args.release is not None and args.release != att["year_ending"]:
        halt(f"--release {ye_text(args.release)}: the data tables page lists "
             f"only the newest Asy_D11 ({att['label']}); for another release "
             "use --file with a copy of that file (identity is read from the "
             "file)")
    print(f"downloading to {RAW_DIR} (also in a preview; the database is not "
          "written)")
    path, final = fetch(_abs_url(att["url"]), RAW_DIR, expect=(D11_SHEET,))
    print(f"  final URL: {final}")
    out.update(path=path, where=final, page=att, att=att, d09_att=d09,
               how="newest release (content API)"
               + ("; --release given" if args.release else ""))
    if args.d09_file:
        out["d09_path"] = Path(args.d09_file)
    return out


def _reg02_source(ctx) -> dict:
    args = ctx.args
    out = {"path": None, "where": None, "how": None, "page": None}
    if args.file:
        path = Path(args.file)
        out.update(path=path, how="local file")
        if args.no_page:
            out.update(where=f"file:{path.name}",
                       how="local file, --no-page: the page was not read")
            print(f"NOTE: --no-page: the GOV.UK page is not read; source_file "
                  f"is recorded as file:{path.name}")
        else:
            out["where"] = _file_source(path)
            page = fetch_json(REGIONAL_PATH)
            _print_page(page, REGIONAL_PATH)
            try:
                rel = reg02_releases(page)
            except ValueError as e:
                halt(f"{e}; give --no-page to load the file without reading "
                     "the page (recorded as such)")
            print(f"  newest Reg_02 the page lists: {rel[max(rel)]['title']!r}")
            out["releases"] = rel
        return out
    page = fetch_json(REGIONAL_PATH)
    _print_page(page, REGIONAL_PATH)
    try:
        rel = reg02_releases(page)
    except ValueError as e:
        halt(str(e))
    newest = rel[max(rel)]
    print(f"  newest Reg_02 the page lists: {newest['title']!r} "
          f"{newest['url']}")
    ctx.newest_seen["Reg_02"] = newest["year_ending"].isoformat()
    if args.release is not None and args.only == "reg02":
        if args.release not in rel:
            halt(f"--release {ye_text(args.release)}: no Reg_02 attachment for "
                 f"it; listed: {[ye_text(d) for d in sorted(rel)]}")
        att, how = rel[args.release], f"release {ye_text(args.release)} given"
    else:
        att, how = newest, "newest release (content API)"
    print(f"downloading to {RAW_DIR} (also in a preview; the database is not "
          "written)")
    path, final = fetch(_abs_url(att["url"]), RAW_DIR, expect=(REG02_SHEET,))
    print(f"  final URL: {final}")
    out.update(path=path, where=final, page=att, how=how)
    return out


def _skip_by_ledger(ctx, src, sha, checked, ranks, stranded) -> bool:
    if not ctx.has_ed or ctx.args.recheck is not None:
        return False
    done = ledger_complete(checked, src["where"], sha, ranks,
                           allow_older=ctx.args.allow_older_file,
                           stranded=stranded)
    if not done:
        return False
    print(f"the file's (source, sha256) is already in the ledgers for every "
          f"period it states ({ranges_text(done)}); nothing parsed "
          "(--recheck YYYY-MM-DD reads it again)")
    return True


def identity_problems(cover, page=None, release=None) -> list:
    """Problems (empty = fine) between a file's cover and where it came
    from: the page attachment's year ending (when the file was found on the
    page) and --release must both equal the cover's."""
    out = []
    if page is not None and page["year_ending"] != cover["year_ending"]:
        out.append(f"the page attachment says {page['label']}, the file says "
                   f"{cover['label']}")
    if release is not None and release != cover["year_ending"]:
        out.append(f"--release {ye_text(release)}, the file says "
                   f"{cover['label']}")
    return out


def _label(cover, older):
    return cover["label"] + (" (older file, stored with --allow-older-file)"
                             if older else "")


def _run_d11(ctx) -> int:
    args, cur = ctx.args, ctx.cur
    started = datetime.now(timezone.utc)
    tips, checked, stranded = _state(ctx, D11_TAGS)
    ranks = period_ranks(tips, checked)
    held = sorted(ranks)
    ctx.held["Asy_D11"] = held
    ctx.checked_all.update(checked)
    print("Asy_D11: held periods: " + (f"{held[0]} .. {held[-1]} ({len(held)})"
                                       if held else "none"))
    src = _d11_source(ctx)
    path = src["path"]
    if not Path(path).is_file():
        halt(f"--file {path}: no such file")
    sha = content_sha256(path)
    print(f"{Path(path).name}: sha256 {sha[:16]}"
          + ("; byte-identical to the held file the old build read"
             if sha == LEGACY_FILES["d11"]["sha256"] else ""))
    if _skip_by_ledger(ctx, src, sha, checked, ranks, stranded):
        print("Asy_D11: nothing to do")
        return 0
    file = _read(read_d11, path)
    cov = file["cover"]
    bad = identity_problems(cov, None if args.file else src["page"],
                            args.release)
    if args.file and src["page"] is not None and \
            src["page"]["year_ending"] != cov["year_ending"]:
        print(f"NOTE: the page lists {src['page']['label']} as the newest "
              f"Asy_D11; the file says {cov['label']} (the page lists only "
              "the newest; identity is read from the file)")
    if bad:
        halt("file identity check failed, nothing stored: " + "; ".join(bad))
    print(f"{Path(path).name}: {cov['publication']!r}; {cov['label']}; "
          f"published {cov['published']}; next update "
          f"{cov['next_update_text']}; Contents period covered "
          f"{sorted({d.isoformat() for d in file['contents'].values()})}; "
          f"{len(file['rows']):,} data rows, {len(file['periods'])} quarters "
          f"{file['periods'][0]} .. {file['periods'][-1]}")
    rmap = _resolve(cur, d11_codes_by_period(file))
    try:
        by_period, merges, dups = build_d11_records(file, rmap)
    except ValueError as e:
        halt(f"{e}; nothing stored")
    print(f"reorganisation merges summed: {len(merges)}; same-code duplicate "
          f"keys summed: {len(dups)} (limit {DUPLICATE_HALT})")
    for d in dups:
        print(f"  duplicate: {d['period']} {d['key']} code {d['codes'][0]} "
              f"rows {d['rows']} regions {d['regions']} people {d['people']} "
              f"-> {sum(d['people'])}")
    allrecs = {t: [r for p in by_period.values() for r in p[t]]
               for t in D11_TAGS}
    for t in D11_TAGS:
        _report_blanks(t, allrecs[t])
    problems = reconcile_d11(file, by_period)
    if problems:
        halt("reconciliation inside the file failed, nothing stored: "
             + "; ".join(problems[:8]))
    piv = file["pivot"]
    print(f"reconciled: the data sheet equals the pivot cache for "
          f"{len(piv['dates'])} quarters ({piv['dates'][0]} .. "
          f"{piv['dates'][-1]}; Grand Total "
          + ", ".join(f"{d} {v:,}" for d, v in zip(
              piv["dates"], piv["rows"][GRAND_TOTAL]) if v is not None)
          + "); England + non-England + unallocated equal the data sheet "
          "total in every period")
    d09_note = "reconciliation 2 (Asy_D09) not done: --no-d09"
    if not args.no_d09:
        d09_path = src.get("d09_path")
        if d09_path is None and src.get("d09_att"):
            d09_path, d09_final = fetch(_abs_url(src["d09_att"]["url"]),
                                        RAW_DIR, expect=(D09_SHEET,))
            print(f"  Asy_D09 final URL: {d09_final}")
        if d09_path is None:
            halt("--file needs --d09-file PATH (the Asy_D09 file of the same "
                 "release) or --no-d09 (reconciliation 2 is then logged as "
                 "not done)")
        d09 = _read(read_d09, d09_path)
        problems = reconcile_d09(d09, by_period, cov)
        if problems:
            halt("reconciliation with Asy_D09 failed, nothing stored: "
                 + "; ".join(problems[:8]))
        d09_note = (f"reconciled with Asy_D09 ({Path(d09_path).name}, sha256 "
                    f"{content_sha256(d09_path)[:16]}, {d09['cover']['label']})"
                    ": England and unallocated equal it in every period")
    print(d09_note)
    for p, t in by_period.items():
        ctx.d11_records[p] = t["support"]
    periods = sorted(by_period)
    gone = sorted(set(held) - set(periods))
    if gone and max(periods) >= max(held):
        halt(f"the file does not state held periods {gone} (a time-series "
             "file restates every quarter); nothing stored")
    needs = {p: [t for t in D11_TAGS if by_period[p][t]] for p in periods}
    lsrc = ledger_source(src["where"], cov, {
        t: [p for p in periods if by_period[p][t]] for t in D11_TAGS})
    chk = {t: {p: s for p, s in c.items() if p not in stranded}
           for t, c in checked.items()}
    new, compared, skipped = plan_periods(
        periods, cov["rank"], ranks, chk, lsrc, sha, recheck=args.recheck,
        allow_older=args.allow_older_file, needs=needs)
    return _run_periods(ctx, "d11", src, file, cov, sha, lsrc, by_period, new,
                        compared, skipped, tips, ranks, held, started,
                        extra_note=d09_note)


def _run_reg02(ctx) -> int:
    args, cur = ctx.args, ctx.cur
    started = datetime.now(timezone.utc)
    tips, checked, stranded = _state(ctx, ("groups",))
    ranks = period_ranks(tips, checked)
    held = sorted(ranks)
    ctx.held["Reg_02"] = held
    ctx.checked_all.update(checked)
    print("Reg_02: held periods: " + (", ".join(held) or "none"))
    src = _reg02_source(ctx)
    path = src["path"]
    if not Path(path).is_file():
        halt(f"--file {path}: no such file")
    sha = content_sha256(path)
    print(f"{Path(path).name}: sha256 {sha[:16]}")
    if _skip_by_ledger(ctx, src, sha, checked, ranks, stranded):
        print("Reg_02: nothing to do")
        return 0
    file = _read(read_reg02, path)
    cov = file["cover"]
    bad = identity_problems(cov, src["page"], args.release)
    if bad:
        halt("file identity check failed, nothing stored: " + "; ".join(bad))
    n = reg02_counts(file)
    print(f"{Path(path).name}: {cov['publication']!r}; {cov['label']}; "
          f"published {cov['published']}; next update "
          f"{cov['next_update_text']}; {REG02_SHEET} as at {file['title_date']}"
          f"; rows: England {n['england']}, Scotland {n['scotland']}, Wales "
          f"{n['wales']}, Northern Ireland {n['northern_ireland']}, Unknown "
          f"{n['unknown']} (only England is stored)")
    p = file["period"].isoformat()
    codes = {r["code"] for r in file["rows"]
             if r["code"][:3] in ENGLISH_PREFIXES}
    rmap = _resolve(cur, {p: codes})
    try:
        p, recs = build_reg02_records(file, rmap)
    except ValueError as e:
        halt(f"{e}; nothing stored")
    _report_blanks("groups", recs)
    d11 = None
    how_d11 = "not checked: no Asy_D11 rows held or read for this period"
    if p in ctx.d11_records:
        d11, how_d11 = d11_la_totals(ctx.d11_records[p]), "this run's Asy_D11"
    else:
        sup = profile("support").spec
        held_sup = (tip_info(cur, sup) if ctx.has_ed else legacy_tips(cur, sup))
        if p in held_sup:
            rs = (records(cur, sup, p, held_sup[p]["edition"]) if ctx.has_ed
                  else records(cur, sup, p))
            d11, how_d11 = d11_la_totals(rs), "the held Asy_D11 rows"
    problems = reconcile_reg02(file, recs, d11)
    if problems:
        halt("Reg_02 reconciliation failed, nothing stored: "
             + "; ".join(problems[:8]))
    print(f"reconciled: {REG02_LAS} English authorities; pathway totals equal "
          "their parts where published; all_pathways equals the published "
          f"pathway totals; supported_asylum against Asy_D11: {how_d11}")
    lsrc = ledger_source(src["where"], cov, {"groups": [p]})
    chk = {t: {q: s for q, s in c.items() if q not in stranded}
           for t, c in checked.items()}
    new, compared, skipped = plan_periods(
        [p], cov["rank"], ranks, chk, lsrc, sha, recheck=args.recheck,
        allow_older=args.allow_older_file, needs={p: ["groups"]})
    if p in new and held and p < max(held) and not args.allow_older_file:
        # a Reg_02 snapshot is one period; an earlier one than anything
        # held is a back-fill, not a new quarter of the series
        halt(f"back-fill: this Reg_02 file is for {p}, earlier than the "
             f"newest held period {max(held)} and not held itself; storing "
             "it would add an earlier snapshot to the editions. If this is "
             "deliberate, re-run with --allow-older-file")
    return _run_periods(ctx, "reg02", src, file, cov, sha, lsrc,
                        {p: {"groups": recs}}, new, compared, skipped, tips,
                        ranks, held, started,
                        extra_note=f"Reg_02 against Asy_D11: {how_d11}")


def _differs(new, old, tag) -> bool:
    cols = COMPARED[tag]
    nk = {_rkey(r, tag): r for r in new}
    ok = {_rkey(r, tag): r for r in old}
    if set(nk) != set(ok):
        return True
    return any(nk[k].get(c) != ok[k].get(c) for k in nk for c in cols)


def _run_periods(ctx, part, src, file, cov, sha, lsrc, by_period, new,
                 compared, skipped, tips, ranks, held, started, *,
                 extra_note="") -> int:
    args, cur = ctx.args, ctx.cur
    tags = D11_TAGS if part == "d11" else ("groups",)
    name = "Asy_D11" if part == "d11" else "Reg_02"
    rank = cov["rank"]
    ack = set(args.acknowledge or ())
    reissue = set(args.accept_reissue or ())
    older_note = {}
    for p in compared:
        t = ranks.get(p)
        if t is not None and t > rank:
            older_note[p] = (f"--allow-older-file given; {p} is held from a "
                             f"file of {rank_text(t)}, this file is "
                             f"{rank_text(rank)}")
            print(f"NOTE: {older_note[p]}")
    if not skipped:
        skip_text = "none"
    elif len(skipped) <= 6:
        skip_text = ", ".join(f"{p} ({r})" for p, r in sorted(skipped.items()))
    else:
        skip_text = (f"{len(skipped)} periods ({ranges_text(skipped)}: "
                     f"{sorted(set(skipped.values()))})")
    print(f"{name}: planned: new {new or 'none'}; compare "
          + (ranges_text(compared) if compared else "none")
          + f"; skipped {skip_text}")
    if args.recheck is not None and args.recheck not in by_period:
        print(f"NOTE: --recheck {args.recheck} is not a period of this file; "
              "the file is read and every period not older is compared")
    periods = sorted(new + compared)
    if not periods:
        if skipped and all(r.startswith("older") for r in skipped.values()):
            halt(f"older file: every period of the {name} file is skipped; the "
                 f"file is {rank_text(rank)}, older than what is held for "
                 f"{ranges_text(skipped)}; storing it would record an older "
                 "file. If this is deliberate, re-run with --allow-older-file")
        print(f"{name}: nothing to do")
        return 0
    problems, infos, acked, reissued = {}, {}, {}, {}
    for p in periods:
        bad, soft_all, flips = [], [], []
        held_p = p in compared
        for tag in tags:
            recs = by_period[p][tag]
            have = p in tips[tag]
            if not recs:
                # an edition with no rows for a table cannot be stored (only
                # tags with records are written), so this stays a stop
                # condition no flag releases
                if have:
                    bad.append(f"{PARTIAL}: {profile(tag).spec.live_table}: the "
                               "file has no rows for this period, the held "
                               "edition has some; rows cannot vanish "
                               "silently")
                continue
            if have:
                tip = _tip_records(ctx, tag, p, tips)
                full = period_problems(recs, tip, None, table=tag,
                                       kind="revised")
                left = period_problems(recs, tip, None, table=tag,
                                       kind="revised",
                                       acknowledged=p in ack)
                if tag == "groups":
                    fl = blank_flips(recs, tip)
                    if fl:
                        msg = (f"{len(fl)} cell(s) go from 0 to '*' or '*' to 0"
                               " against the tip (rule 1.10), e.g. "
                               + ", ".join(f"{'/'.join(k)} {a}->{b}"
                                           for k, _, a, b in fl[:3]))
                        full.append(msg)
                        if p not in ack:
                            left.append(msg + f"; read them, then "
                                        f"--acknowledge {p}")
            else:
                earlier = [h for h in sorted(tips[tag]) if h < p]
                prev = (_tip_records(ctx, tag, earlier[-1], tips)
                        if earlier and tag == "support" else None)
                full = period_problems(recs, None, prev, table=tag, kind="new")
                left = period_problems(recs, None, prev, table=tag, kind="new",
                                       acknowledged=p in ack)
            pref = f"{profile(tag).spec.live_table}: "
            bad += [pref + x for x in left]
            soft_all += [pref + x for x in full if x not in left]
        if held_p and ranks.get(p) == rank:
            diff = [t for t in tags if p in tips[t] and by_period[p][t]
                    and _differs(by_period[p][t],
                                 _tip_records(ctx, t, p, tips), t)]
            if diff:
                if p in reissue:
                    reissued[p] = (f"--accept-reissue {p}: the file is "
                                   f"{rank_text(rank)}, the same as the held "
                                   f"file, with different content in "
                                   f"{diff} (a publisher reissue keeping its "
                                   "cover date)")
                    print(f"  {p}: ACCEPTED REISSUE: {reissued[p]}")
                else:
                    bad.append(f"the file is {rank_text(rank)}, the same as "
                               "the file of the held tip, but its content "
                               f"differs in {diff}: two files claim the same "
                               "release. A publisher reissue keeps its cover "
                               f"date: read the page's change note, then "
                               f"--accept-reissue {p}")
        if soft_all and not bad:
            acked[p] = "; ".join(soft_all)
            print(f"  {p}: ACKNOWLEDGED (--acknowledge): {acked[p]}")
        if bad:
            problems[p] = bad
            continue
        notes = "; ".join(x for x in (older_note.get(p, ""),
                                      reissued.get(p, ""),
                                      f"ACKNOWLEDGED {acked[p]}"
                                      if p in acked else "") if x)
        infos[p] = PeriodInfo(edition_source(src["where"], cov, notes),
                              _label(cov, p in older_note), cov["published"],
                              ((lsrc, sha),))
    stray = sorted((ack | reissue) - set(periods))
    if stray:
        print(f"NOTE: --acknowledge/--accept-reissue {stray}: not compared "
              f"by this {name} run")
    for p, msgs in sorted(problems.items()):
        for msg in msgs:
            print(f"  {p}: REJECTED, not stored: STOP CONDITION: {msg}")
    ok = [p for p in periods if p not in problems]
    rc, outcomes, failed, left = 0, {}, None, []
    if ok:
        rc, outcomes, failed, left = _apply_periods(ctx, part, ok, by_period,
                                                    infos)
    if problems:
        print(f"REJECTED {len(problems)} {name} period(s), nothing stored for "
              f"them: {', '.join(sorted(problems))}; exit 1")
        rc = 1
    if not ctx.has_ed:
        print(f"  ({name}: no editions tables yet: preview against the live "
              "tables only)")
    if args.commit:
        stored = live = 0  # rows written by this run; not a source value
        for p, kinds in outcomes.items():
            for tag, k in kinds.items():
                n = len(by_period[p][tag])
                if k in ("new", "revised"):
                    stored += n
                if k in ("new", LIVE_MISSING):
                    live += n
        partial = ""
        if rc:
            partial = (f"PARTIAL RUN (exit 1): rejected "
                       + (", ".join(f"{p} ({'; '.join(problems[p])[:300]})"
                                    for p in sorted(problems)) or "none")
                       + f"; failed {failed or 'none'}; not attempted "
                       f"{left or 'none'}; a period that failed or was "
                       "rejected stored nothing and is not in the counts "
                       "below; ")
        notes = (partial + f"{name} file {src['where']} sha256 {sha[:16]}, "
                 f"{rank_text(rank)} ({src['how']}): "
                 + ("; ".join(f"{p} " + ", ".join(f"{t} {k}"
                                                  for t, k in kinds.items())
                              for p, kinds in sorted(outcomes.items()))
                    or "nothing stored")
                 + (f"; skipped {ranges_text(skipped)}" if skipped else "")
                 + "".join(f"; {older_note[p]}" for p in sorted(older_note))
                 + "".join(f"; {reissued[p]}" for p in sorted(reissued))
                 + "".join(f"; ACKNOWLEDGED {p}: {acked[p]}"
                           for p in sorted(acked))
                 + (f"; {extra_note}" if extra_note else "")
                 + f". Rows stored: {stored} edition rows, {live} live rows.")
        log_run(cur, stored + live, notes, started)
        ctx.conn.commit()
        print("pipeline_run_log row written"
              + (" (partial run: see its notes)" if rc else ""))
    return rc


def _apply_periods(ctx, part, periods, by_period, infos) -> tuple:
    """Compare and (commit or simulate) store each period in its own
    savepoint, as period_editions.load_periods does: on commit each period
    is committed on its own; on simulate and in preview the savepoint is
    always rolled back. A failure rolls that period back, stops, and is
    reported with the periods not attempted. (rc, {period: {tag: kind}},
    failed period or None, [not attempted])."""
    args, cur = ctx.args, ctx.cur
    commit, simulate = args.commit, args.simulate
    write = commit or simulate
    tags = D11_TAGS if part == "d11" else ("groups",)
    profs = profiles(tags)
    sp = profs[tags[0]].savepoint
    conn = cur.connection
    against = "editions" if ctx.has_ed else "live"
    outcomes = {}
    tally = {}
    for i, p in enumerate(periods):
        in_sp = False
        try:
            cur.execute(f"SAVEPOINT {sp}")
            in_sp = True
            cmps = {}
            for tag in tags:
                recs = by_period[p][tag]
                if recs:
                    cmps[tag] = pe.compare_period(cur, profs[tag], p, recs,
                                                  against)
            kinds = {t: c["kind"] for t, c in cmps.items()}
            if write and ctx.has_ed:
                if part == "d11":
                    kinds = _late("apply_d11_period")(cur, profs, p,
                                                      by_period[p],
                                                      info=infos[p])
                else:
                    kinds = {"groups": _late("apply_reg02_period")(
                        cur, profs["groups"], p, by_period[p]["groups"],
                        info=infos[p])}
            if commit:
                cur.execute(f"RELEASE SAVEPOINT {sp}")
                in_sp = False
                conn.commit()
            else:
                cur.execute(f"ROLLBACK TO SAVEPOINT {sp}")
                cur.execute(f"RELEASE SAVEPOINT {sp}")
                in_sp = False
        except (Exception, SystemExit) as e:
            if in_sp:
                cur.execute(f"ROLLBACK TO SAVEPOINT {sp}")
                cur.execute(f"RELEASE SAVEPOINT {sp}")
            if commit:
                conn.rollback()
            msg = e.code if isinstance(e, SystemExit) else \
                f"{type(e).__name__}: {e}"
            print(f"  {p}: FAILED, nothing stored for this period ({msg})")
            rest = list(periods[i + 1:])
            if rest:
                print(f"  not attempted: {', '.join(rest)}")
            _tally(tally)
            return 1, outcomes, p, rest
        outcomes[p] = kinds
        parts = []
        for tag, c in cmps.items():
            k = kinds.get(tag, c["kind"])
            tally[k] = tally.get(k, 0) + 1  # not a source value
            act = ""
            if k == "new":
                act = (" stored edition 1 and inserted " if write and
                       ctx.has_ed else " would store edition 1 and insert ") \
                    + f"{len(by_period[p][tag]):,} live rows"
            elif k == LIVE_MISSING:
                act = (" inserted " if write and ctx.has_ed else
                       " would insert ") + "the live rows (no new edition)"
            elif k == "revised":
                act = (" stored the next edition" if write and ctx.has_ed
                       else " would store the next edition") + \
                    " (live waits for refresh-latest)"
            parts.append(f"{tag} {k}, {c['changed']} rows "
                         f"{'new' if k == 'new' else 'changed'}{act}"
                         + ("; " + "; ".join(c["examples"][:2])
                            if c["examples"] else ""))
        print(f"  {p}: " + " | ".join(parts) + (" COMMITTED" if commit else ""))
    _tally(tally)
    return 0, outcomes, None, []


def _tally(tally) -> None:
    print("  summary (table-periods): " + (", ".join(
        f"{k} {n}" for k, n in sorted(tally.items())) or "none"))


def cmd_load(args) -> int:
    if args.no_page and not args.file:
        halt("--no-page is only for --file (a local file whose page is not "
             "read)")
    if (args.d09_file or args.no_d09) and args.only == "reg02":
        halt("--d09-file and --no-d09 are for Asy_D11")
    if args.d09_file and args.no_d09:
        halt("--d09-file and --no-d09 exclude each other")
    parts = ["d11", "reg02"] if args.only is None else [args.only]
    if args.file:
        try:
            kind = read_cover(args.file)["kind"]
        except (ValueError, OSError) as e:
            halt(f"--file {args.file}: {e}")
        if kind == "d09":
            halt("--file is an Asy_D09 file; give it as --d09-file with the "
                 "Asy_D11 --file")
        if args.only is not None and args.only != kind:
            halt(f"--only {args.only}, but --file is a {kind} file (from its "
                 "cover)")
        parts = [kind]
    if args.release is not None and args.only is None and not args.file:
        halt("--release needs --only d11 or --only reg02")
    writing = args.commit or args.simulate
    tags = tuple(t for part in parts for t in
                 (D11_TAGS if part == "d11" else ("groups",)))
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            ready = _ready(cur, tags)
            has_ed = all(ready.values())
            missing = [t for t, ok in ready.items() if not ok]
            if writing and not has_ed:
                halt(f"missing {', '.join(missing)}; run `ddl --commit` and "
                     "`migrate-legacy ... --commit` first")
            if not has_ed:
                print(f"NOTE: {', '.join(missing)} not present: no editions "
                      "tables yet; this preview compares each file with the "
                      "LIVE tables")
            ctx = Ctx(args, conn, cur, has_ed, writing)
            rc = 0
            for part in parts:
                rc = max(rc, _run_d11(ctx) if part == "d11" else _run_reg02(ctx))
            if ctx.newest_seen:
                held = {k: v for k, v in ctx.held.items()}
                nu = held_next_update(ctx.checked_all)
                for line in discovery_note(ctx.newest_seen, held, nu):
                    print(line)
        return _end(conn, args, writing, rc)
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# One-off: migrate the held tables
# ---------------------------------------------------------------------------


def live_state(cur, spec) -> tuple:
    """(rows, survey hash, {period: [loaded_at]}, {period: [source_edition]})
    of a live table. The survey hash: md5 of every stored column but
    loaded_at, key columns first then the other columns in table order,
    '|'-joined (NULL as '~'), aggregated by LF ordered by the joined text;
    for la_asylum_support exactly the spec's expression (ordered by key)."""
    cur.execute("SELECT column_name FROM information_schema.columns WHERE "
                "table_schema = 'public' AND table_name = %s ORDER BY "
                "ordinal_position", (spec.live_table,))
    allc = [r[0] for r in cur.fetchall() if r[0] != "loaded_at"]
    keys = ["period_ending"] + list(spec.key_cols)
    cols = keys + [c for c in allc if c not in keys]
    joined = "||'|'||".join(f"coalesce({c}::text,'~')" for c in cols)
    order = "t"
    if tuple(spec.key_cols) == KEY_COLS["support"]:
        order = ", ".join(keys)
        joined = ("period_ending||'|'||lad24cd||'|'||support_type||'|'||"
                  "accommodation_type||'|'||people||'|'||published_la_name||"
                  "'|'||coalesce(source_marker,'~')||'|'||source_edition")
    cur.execute(f"SELECT COUNT(*), md5(string_agg(t, E'\\n' ORDER BY {order}))"
                f" FROM (SELECT {joined} AS t, {', '.join(keys)} FROM "
                f"public.{spec.live_table}) q")
    n, h = cur.fetchone()
    cur.execute(f"SELECT period_ending, array_agg(DISTINCT loaded_at), "
                f"array_agg(DISTINCT source_edition) FROM "
                f"public.{spec.live_table} GROUP BY 1")
    loaded, labels = {}, {}
    for p, la, lb in cur.fetchall():
        loaded[pe._p(p)] = sorted(x.astimezone(timezone.utc) for x in la)
        labels[pe._p(p)] = sorted(lb)
    return n, h, loaded, labels


def _proof(tag, held, built) -> tuple:
    """(cells compared, differences) of held rows against rebuilt rows on
    the key, the period and every compared column."""
    hk = {(r["period_ending"],) + _rkey(r, tag): r for r in held}
    bk = {(r["period_ending"],) + _rkey(r, tag): r for r in built}
    diffs = [("only held", k) for k in sorted(set(hk) - set(bk))]
    diffs += [("only rebuilt", k) for k in sorted(set(bk) - set(hk))]
    cells = 0  # cells compared; not a source value
    for k in sorted(set(hk) & set(bk)):
        for c in COMPARED[tag]:
            cells += 1
            if hk[k].get(c) != bk[k].get(c):
                diffs.append((tag, k, c, hk[k].get(c), bk[k].get(c)))
    return cells, diffs


def migration_plan(cur, d11, reg02_files, *, d09=None, legacy=None,
                   files=None) -> dict:
    """Read-only: the preconditions on live and the files, the proof, and
    the edition 1 and ledger content per table and period. Halts on any
    failure (nothing is written here)."""
    g = globals()
    legacy = legacy or g["LEGACY_LIVE"]
    files = files or g["LEGACY_FILES"]
    profs = profiles()
    for tag, prof in profs.items():
        spec = prof.spec
        n, h, loaded, labels = live_state(cur, spec)
        rows, hsh, want_loaded = legacy[spec.live_table]
        bad = []
        if (n, h) != (rows, hsh):
            bad.append(f"{n} rows, hash {h}; expected {rows} rows, hash {hsh}")
        for p, la in sorted(loaded.items()):
            want = want_loaded.get(p, want_loaded.get("*"))
            if [str(x) for x in la] != [want]:
                bad.append(f"{p}: loaded_at {[str(x) for x in la]}, expected "
                           f"[{want}]")
            if len(labels[p]) != 1:
                bad.append(f"{p}: {len(labels[p])} source_edition values")
        if bad:
            halt(f"{spec.live_table} is not as surveyed: " + "; ".join(bad[:6])
                 + "; nothing stored")
    d11p = Path(d11)
    sha = content_sha256(d11p)
    if sha != files["d11"]["sha256"]:
        halt(f"{d11p.name}: sha256 {sha[:16]}, expected "
             f"{files['d11']['sha256'][:16]} (the held Asy_D11 file); nothing "
             "stored")
    want_reg = {f["sha256"]: k for k, f in files.items()
                if k.startswith("reg02")}
    regs = []
    for rp in reg02_files:
        rs = content_sha256(rp)
        if rs not in want_reg:
            halt(f"{Path(rp).name}: sha256 {rs[:16]}, not one of the held "
                 f"Reg_02 files {[k[:16] for k in want_reg]}; nothing stored")
        regs.append((Path(rp), rs, files[want_reg[rs]]))
    file = _read(read_d11, d11p)
    cov = file["cover"]
    rmap = _resolve(cur, d11_codes_by_period(file))
    try:
        by_period, merges, dups = build_d11_records(file, rmap)
    except ValueError as e:
        halt(f"{e}; nothing stored")
    problems = reconcile_d11(file, by_period)
    d09_note = "not done (no --d09 file)"
    if d09 is not None:
        d9 = _read(read_d09, d09)
        if "d09" in files and content_sha256(d09) != files["d09"]["sha256"]:
            halt(f"{Path(d09).name}: sha256 {content_sha256(d09)[:16]}, "
                 f"expected {files['d09']['sha256'][:16]}; nothing stored")
        problems += reconcile_d09(d9, by_period, cov)
        d09_note = "England and unallocated equal Asy_D09 in every period"
    if problems:
        halt("proof failed: reconciliation: " + "; ".join(problems[:8])
             + "; nothing stored")
    per = {t: {} for t in TAGS}
    proof, diffs = {}, []
    for tag in D11_TAGS:
        spec = profs[tag].spec
        cur.execute(f"SELECT DISTINCT period_ending FROM public.{spec.live_table}")
        live_p = sorted(pe._p(r[0]) for r in cur.fetchall())
        file_p = sorted(p for p, t in by_period.items() if t[tag])
        if live_p != file_p:
            halt(f"{spec.live_table}: live periods {ranges_text(live_p)}, the "
                 f"file states {ranges_text(file_p)}; nothing stored")
        held = [r for p in live_p for r in records(cur, spec, p)]
        built = [r for p in file_p for r in by_period[p][tag]]
        proof[tag] = _proof(tag, held, built)
        diffs += proof[tag][1]
        lsrc = ledger_source(files["d11"]["url"], cov, {
            t: [p for p in sorted(by_period) if by_period[p][t]]
            for t in D11_TAGS})
        for p in live_p:
            per[tag][p] = {"rows": [r for r in held if r["period_ending"] == p],
                           "file": d11p.name, "cover": cov, "ledger": lsrc,
                           "sha": sha}
    d11_tot = {p: d11_la_totals(t["support"]) for p, t in by_period.items()}
    spec = profs["groups"].spec
    cur.execute(f"SELECT DISTINCT period_ending FROM public.{spec.live_table}")
    live_g = sorted(pe._p(r[0]) for r in cur.fetchall())
    seen_g = []
    held_g, built_g = [], []
    for rp, rs, meta in regs:
        f = _read(read_reg02, rp)
        rmap = _resolve(cur, {f["period"].isoformat(): {
            r["code"] for r in f["rows"] if r["code"][:3] in ENGLISH_PREFIXES}})
        try:
            p, recs = build_reg02_records(f, rmap)
        except ValueError as e:
            halt(f"{e}; nothing stored")
        bad = reconcile_reg02(f, recs, d11_tot.get(p))
        if bad:
            halt(f"proof failed: {rp.name}: reconciliation: "
                 + "; ".join(bad[:8]) + "; nothing stored")
        if p not in live_g:
            halt(f"{rp.name} states {p}, not a held Reg_02 period {live_g}; "
                 "nothing stored")
        seen_g.append(p)
        held = records(cur, spec, p)
        held_g += held
        built_g += recs
        per["groups"][p] = {"rows": held, "file": rp.name, "cover": f["cover"],
                            "ledger": ledger_source(meta["url"], f["cover"],
                                                    {"groups": [p]}),
                            "sha": rs}
    if sorted(seen_g) != live_g:
        halt(f"Reg_02 files state {sorted(seen_g)}, live holds {live_g}; "
             "nothing stored")
    proof["groups"] = _proof("groups", held_g, built_g)
    diffs += proof["groups"][1]
    # source_edition: each live period carries its proved file's label
    for tag, ps in per.items():
        _, _, _, labels = live_state(cur, profs[tag].spec)
        for p, info in ps.items():
            if labels[p] != [info["cover"]["label"]]:
                diffs.append((tag, p, "source_edition", labels[p],
                              info["cover"]["label"]))
    if diffs:
        halt(f"proof failed: {len(diffs)} differences between the held rows "
             f"and the held files read again, e.g. {diffs[:5]}; nothing "
             "stored")
    for tag, ps in per.items():
        cur.execute(f"SELECT period_ending, MIN((loaded_at AT TIME ZONE 'UTC')"
                    f"::date), MIN(source_edition) FROM "
                    f"public.{profs[tag].spec.live_table} GROUP BY 1")
        for p, d, lab in cur.fetchall():
            ps[pe._p(p)].update(loaded=d, label=lab)
    return {"periods": per, "proof": {t: v[0] for t, v in proof.items()},
            "merges": len(merges), "duplicates": dups, "d09": d09_note}


def migrate_legacy(cur, d11, reg02_files, *, write, d09=None, legacy=None,
                   files=None) -> dict:
    """One-off. Preconditions: on write the four editions tables and
    ledgers exist and are empty (in a preview they may be absent; present,
    they must be empty); live as surveyed (LEGACY_LIVE: rows, survey hash,
    the one loaded_at per table-period, one source_edition per
    table-period); the files' sha256 as LEGACY_FILES. Proof, any failure
    halts (migration_plan): the new parser on the held files reproduces
    every live row and every stored column of the four tables (the
    source_edition of each period is its proved file's 'year ending ...'),
    and the reconciliations pass (Asy_D09 also, when given). Edition 1 'as
    loaded' from live for every period (release label the live
    source_edition, published_date the load date, source_file 'as loaded:
    <file>; year ending ...; published ...'), and ledger rows for the files
    (outcome 'unchanged', edition 1) so the next load skips them by (URL,
    sha). Live untouched. Nothing commits here; the caller rolls back on a
    halt."""
    profs = profiles()
    ready = _ready(cur, TAGS)
    if write and not all(ready.values()):
        halt(f"missing {[t for t, ok in ready.items() if not ok]}; run "
             "`ddl --commit` first")
    for t, ok in ready.items():
        if ok:
            cur.execute(f"SELECT COUNT(*) FROM public.{t}")
            n = cur.fetchone()[0]
            if n:
                halt(f"{t} already holds {n} rows; migrate-legacy runs once, "
                     "on empty tables")
    plan = migration_plan(cur, d11, reg02_files, d09=d09, legacy=legacy,
                          files=files)
    pr = plan["proof"]
    print("proof: the held files read again reproduce every held row and "
          "every stored column: "
          + ", ".join(f"{profs[t].spec.live_table} {sum(len(v['rows']) for v in plan['periods'][t].values()):,} rows / {pr[t]:,} cells"
                      for t in TAGS)
          + "; 0 differences; source_edition of every period is its file's "
          f"own 'year ending'; reorganisation merges {plan['merges']}, "
          f"same-code duplicates {len(plan['duplicates'])}; Asy_D09: "
          f"{plan['d09']}")
    for t in TAGS:
        ps = plan["periods"][t]
        print(f"  {profs[t].spec.live_table}: edition 1 as loaded for "
              f"{len(ps)} period(s) ({ranges_text(ps)})")
    if not write:
        return plan
    sps = {"support": "s6_d11_period", "groups": "s6_reg02_period"}
    by_p = {}
    for t in TAGS:
        for p, info in plan["periods"][t].items():
            by_p.setdefault(p, []).append((t, info))
    for p in sorted(by_p):
        cur.execute(f"SAVEPOINT {sps['support']}")
        for t, info in by_p[p]:
            spec = profs[t].spec
            cov = info["cover"]
            ed = core.insert_edition(
                cur, spec, info["rows"], p, release_label=info["label"],
                published_date=info["loaded"],
                source_file=(f"as loaded: {info['file']}; {cov['label']}; "
                             f"published {cov['published']}"),
                source_sha256=rows_content_sha(spec, info["rows"]),
                supersedes=None, strict=True)
            if ed != 1 or core.rows_differing(cur, spec, p, ed):
                halt(f"{spec.editions_table} {p}: edition 1 does not equal "
                     "live; rolled back")
            pe.record_file_check(cur, profs[t], p, info["ledger"],
                                 info["sha"], "unchanged", 1)
        cur.execute(f"RELEASE SAVEPOINT {sps['support']}")
    bad = []
    for t in TAGS:
        bad += load_checks.check_latest_equals_live(cur, profs[t].spec)
    if bad:
        halt("migrate-legacy: an edition 1 differs from live, rolled back: "
             + "; ".join(bad[:6]))
    return plan


def cmd_migrate_legacy(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    try:
        with conn.cursor() as cur:
            plan = migrate_legacy(cur, args.d11, args.reg02, write=writing,
                                  d09=args.d09)
            if args.commit:
                per = plan["periods"]
                n = sum(len(v["rows"]) for ps in per.values()
                        for v in ps.values())
                log_run(cur, n, "migrate-legacy: edition 1 as loaded for "
                        + "; ".join(f"{t} {ranges_text(per[t])}" for t in TAGS)
                        + f" ({n} rows) from the live tables; proof "
                        + ", ".join(f"{t} {c} cells" for t, c in
                                    plan["proof"].items())
                        + " equal the held files read again; Asy_D09 "
                        f"{plan['d09']}; ledger rows for the files. Live "
                        "untouched.", started)
                conn.commit()
                print("migrate-legacy: COMMITTED; pipeline_run_log row "
                      "written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# restore-edition
# ---------------------------------------------------------------------------


def restore_edition(cur, table, period, edition) -> int:
    """Store edition `edition` of one table's period as the next edition
    for refresh-latest to apply (its release label and source_file kept,
    with 'restored from edition N' added to the source_file); return its
    number. Halts if the edition does not exist or is the tip."""
    tag, spec = spec_of(table)
    p = str(period)
    tip = core.latest_edition(cur, spec, p)
    recs = records(cur, spec, p, edition)
    if not recs:
        halt(f"{spec.editions_table} {p}: no edition {edition}")
    if edition == tip:
        halt(f"{p}: edition {edition} is the tip; nothing to restore")
    cur.execute(f"SELECT DISTINCT source_file, release_label FROM "
                f"public.{spec.editions_table} WHERE period_ending = %s AND "
                "edition = %s", (p, edition))
    src, label = cur.fetchone()
    new = core.insert_edition(
        cur, spec, recs, p, release_label=label, published_date=date.today(),
        source_file=f"{src}; restored from edition {edition}",
        source_sha256=rows_content_sha(spec, recs), supersedes=tip,
        strict=True, allow_revert=True)
    if new == tip:
        halt(f"{p}: edition {edition} was not stored as the next edition; "
             "nothing to restore")
    return new


def cmd_restore_edition(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    try:
        tag, spec = spec_of(args.table)
    except ValueError as e:
        conn.close()
        halt(str(e))
    try:
        with conn.cursor() as cur:
            if not table_exists(cur, spec.editions_table):
                halt(f"{spec.editions_table} does not exist")
            new = restore_edition(cur, args.table, args.period, args.edition)
            print(f"{spec.live_table} {args.period}: edition {args.edition} "
                  + ("stored" if writing else "would be stored")
                  + f" as edition {new}; run refresh-latest to apply it")
            if args.commit:
                n = len(records(cur, spec, args.period, new))
                log_run(cur, n, f"restore-edition: {spec.live_table} "
                        f"{args.period} edition {args.edition} stored as "
                        f"edition {new}", started)
                conn.commit()
                print("COMMITTED; pipeline_run_log row written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="S6 Home Office asylum support by local authority "
        "(Asy_D11) and immigration groups (Reg_02): editions loader on the "
        "period-editions engine. There is no sync-new; migrate-legacy "
        "records the held periods once.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ddl", help="create the four editions tables, their "
                       "triggers and the four ledgers (preview by default)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_ddl)
    sub.add_parser("status", help="what needs action; exit 1 if anything"
                   ).set_defaults(func=cmd_status)
    p = sub.add_parser("load", help="find the newest Asy_D11 (with Asy_D09) "
                       "and Reg_02 releases, compare and store their periods "
                       "(preview by default)")
    p.add_argument("--only", choices=("d11", "reg02"),
                   help="one publication only (default both)")
    p.add_argument("--release", type=_release_arg, metavar='"Month YYYY"',
                   help="with --only reg02: that release (default the "
                   "newest); with --only d11 it must name the newest listed")
    p.add_argument("--file", metavar="PATH",
                   help="a local Asy_D11 or Reg_02 file (nothing downloaded; "
                   "the kind and identity are read from its cover)")
    p.add_argument("--d09-file", metavar="PATH",
                   help="the Asy_D09 file of the same release (reconciliation "
                   "2)")
    p.add_argument("--no-d09", action="store_true",
                   help="skip the Asy_D09 reconciliation (logged as not done)")
    p.add_argument("--no-page", action="store_true",
                   help="with --file only: do not read the GOV.UK page "
                   "(source_file recorded as file:<name>)")
    p.add_argument("--recheck", type=_period_arg, metavar="YYYY-MM-DD",
                   help="read the file again although the ledgers record it, "
                   "and compare every period not older")
    p.add_argument("--allow-older-file", action="store_true",
                   help="compare and store periods held from a newer file "
                   "(logged, and said in the release label)")
    p.add_argument("--acknowledge", action="extend", nargs="+",
                   type=_period_arg, metavar="YYYY-MM-DD",
                   help="release a period's movement thresholds and 0 <-> '*' "
                   "changes after reading the preview (never a partial file)")
    p.add_argument("--accept-reissue", action="append", type=_period_arg,
                   metavar="YYYY-MM-DD",
                   help="store a period whose file has the held file's rank "
                   "but different content (a publisher reissue keeping its "
                   "cover date; repeatable; logged)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_load)
    p = sub.add_parser("refresh-latest", help="copy each period's latest "
                       "edition into the live tables, source_edition and "
                       "loaded_at included (preview by default)")
    p.add_argument("--accept-key-changes", action="append",
                   type=_period_arg, metavar="YYYY-MM-DD",
                   help="apply this period although rows are added or removed "
                   "(repeatable; each must be named)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_refresh_latest)
    p = sub.add_parser("migrate-legacy", help="one-off: edition 1 as loaded "
                       "for every held period, with the proof (preview by "
                       "default)")
    p.add_argument("--d11", required=True, metavar="FILE")
    p.add_argument("--reg02", required=True, action="append", metavar="FILE")
    p.add_argument("--d09", metavar="FILE")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_migrate_legacy)
    p = sub.add_parser("restore-edition", help="store an earlier edition's "
                       "rows as the next edition (preview by default)")
    p.add_argument("table", metavar="TABLE")
    p.add_argument("period", type=_period_arg, metavar="YYYY-MM-DD")
    p.add_argument("edition", type=int)
    pe.mode_parser(p)
    p.set_defaults(func=cmd_restore_edition)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
