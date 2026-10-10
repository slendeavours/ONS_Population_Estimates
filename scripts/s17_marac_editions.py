"""S17 (SafeLives Marac data, England and Wales, by police force area):
edition history and loader on the period-editions engine.

The live table marac_cases keeps its name, key (pfa_name_safelives,
financial_year) and columns and is the latest-edition layer. This module owns
its editions table and file-check ledger, append-only:

    SPEC     marac_cases_editions               period financial_year
    ledger   marac_cases_editions_file_checks

An edition is what one SafeLives workbook says about one financial year.
SafeLives publishes one workbook per year on its Marac data page
(DATA_PAGE; the registry URL REGISTRY_URL redirects there). Five years are
.xls (2018-19 to 2022-23, read with xlrd 2.0.1), three .xlsx. Each holds a
Notes sheet, a Cases sheet (one row per police force area, the region rows
and the England and Wales and England totals) and a Referral routes sheet whose 'Housing' 'Number' column is housing_referrals.

History (so the wording here stays truthful): n8n run-log rows 22 to 24
(2026-03-30) loaded 2018-19 to 2024-25 from a CSV converted from the
SafeLives files in a chat session that no longer exists; its code read every
value with parseFloat(x) || null, which turned published zeros into NULL.
2025-26 was loaded directly on 2026-08-20 (commit 33c5ec4 records it; no
run-log row, no committed code) and kept its published zeros. The held table
has no row for the West Midlands police force area in any year: in the files
the force shares its name with the West Midlands region. migrate-legacy
(one-off) records each held year as edition 1 "as loaded", with a proof that
this parser reproduces every held value from the held file, or lists it.

Blanks and zeros (docs/RULES.md rule 1): a value cell is a number, or one of
the markers the files use: 'No Data' (NULL, value_flag not_submitted; the
2022-23 file writes 'No data', a named alias) or 'No population data' (NULL,
value_flag no_population). Anything else, a blank included, halts naming the
sheet, row and column. A published 0 stays 0, except under the named rules
in ZERO_RULES, each with its evidence and listed in every preview:
NORFOLK_2023_24_NOT_SUBMITTED and WEST_MIDLANDS_2023_24_NOT_SUBMITTED (a
row published as 0 MARACs and 0 cases with 'No population data' and
'#DIV/0!': a non-submission, every measure NULL not_submitted). There is no
other zero rule. Lancashire's housing referrals, published as 0 in 2023-24,
2024-25 and 2025-26, stay published zeros (Scott, 2026-10-10): none of the
three files carries a note about Lancashire (every cell of every sheet read
on 2026-10-10); the only Lancashire note in any held file is 2022-23's (one
Marac did not submit, July 2022 to March 2023), a year whose housing figure
is 13. The n8n-era note that the 2023-24 and 2024-25 figures were
footnoted as incomplete is not borne out by the files; do not reintroduce a
not-counted rule on that claim. A rule whose published form no longer matches the file halts. value_flag is
an editions-only column: the reason for the NULLs of a row (one reason per
row; two in one row halt). Values are rounded half up to the live column's
scale before they are compared or stored, exactly as the live columns hold
them.

Identity (rule 3), from the file itself on every path: the Notes title
'SafeLives Marac data England and Wales April <y1> - March <y2>' and the
Cases title '... year ending March <y2>' must agree with each other and with
the year asked for (or linked); the Cases header exactly the eight columns
(a differing header only through a named alias in HEADER_ALIASES); the
Housing column found by its header text. The 2022-23 file's Notes title says
April 2021 - March 2022 while every sheet title in it says year ending March
2023: a named erratum for those bytes only (NOTES_TITLE_ERRATA). Inside the
file the English police force areas must sum to the England row on five
measures wherever every part is published (reconcile). The rank of a file is
its own (title year, document properties modified), never its name or link;
the .xls files carry no document properties that xlrd reads, so their rank
is (title year, unknown) and their identity is the titles and the sha256.

Discovery: the data page with a plain User-Agent naming the pipeline (no
browser-like header). Links ending .xlsx or .xls whose names carry two years
'YYYY-YYYY' or 'YYYY-to-YYYY'; two links for one year, none for a held year,
or a workbook link whose years do not fit, halt listing every link. A 403, a
Cloudflare challenge page or a body that is not a workbook halts with the
hand-download message (REFUSED); then download the file by hand into
data/raw/s17_marac/ and load it with --file (manual_input.require_under:
under data/raw or data/reference, a regular file, no symlink).

Geography (rule 4): S17 is keyed by police force area (geography.py
declares '17' 'none'); lad24cd comes through la_pfa_mapping, which is read,
never written. The English police force areas are exactly
la_pfa_mapping.pfa_name_safelives, 39 of them (the West Midlands force
included; the held table carried 38 until the first load, Scott 2026-10-10):
a missing one, an unknown one or a Welsh force taken as English halts.

marac_cases is a W1 input (05_la_signals.sql reads cases_discussed and
cases_per_10k_adult_females at MAX(financial_year) via la_pfa_mapping): run
refresh-latest in the same session before W1. load prints a MAP line when a
stored year is the newest held year.

Subcommands (every writing command previews by default; --commit and
--simulate are mutually exclusive; --simulate runs the --commit path and
always rolls back). There is no sync-new: migrate-legacy records the held
years.
    python scripts/s17_marac_editions.py ddl [--commit | --simulate]
    python scripts/s17_marac_editions.py status
    python scripts/s17_marac_editions.py load [--year YYYY-YY]
        [--file PATH ... [--no-page]] [--recheck YYYY-YY]
        [--allow-older-file] [--acknowledge YYYY-YY] [--accept-reissue
        YYYY-YY] [--acknowledge-flips NAME] [--commit | --simulate]
        # the workbooks are downloaded to data/raw/s17_marac/ (also in a
        # preview; a preview writes nothing to the database). A year
        # breaking a stop condition is REJECTED (nothing stored, exit 1).
    python scripts/s17_marac_editions.py refresh-latest
        [--accept-key-changes YYYY-YY] [--commit | --simulate]
    python scripts/s17_marac_editions.py migrate-legacy [--commit | --simulate]
    python scripts/s17_marac_editions.py restore-edition YYYY-YY N
        [--commit | --simulate]
"""
import argparse
import dataclasses
import hashlib
import html as html_lib
import math
import numbers
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

sys.path.insert(0, str(Path(__file__).resolve().parent))

import editions_core as core  # noqa: E402
import manual_input  # noqa: E402
import period_editions as pe  # noqa: E402
from editions_core import halt  # noqa: E402

RUN_AGENT = "s17_marac_editions"
RUN_SOURCE = "17"
LIVE_MISSING = pe.LIVE_MISSING

REGISTRY_URL = ("https://safelives.org.uk/practice-support/resources-marac-"
                "meetings/latest-marac-data/")
DATA_PAGE = ("https://safelives.org.uk/research-policy/practitioner-datasets/"
             "marac-data/")
REPO = Path(__file__).resolve().parent.parent
RAW_DIR = REPO / "data" / "raw" / "s17_marac"           # git-ignored
FILE_ROOTS = manual_input.ROOTS
USER_AGENT = ("ucws-pipeline S17 loader (read-only download of the public "
              "SafeLives Marac data)")
NO_PAGE = "--no-page: the data page was not read"
REFUSED = ("SafeLives refused the request: download {url} by hand into "
           "data/raw/s17_marac/ and rerun with --file")
CHALLENGE_MARKS = (b"Just a moment", b"cf-chl", b"challenge-platform",
                   b"Attention Required! | Cloudflare", b"cf_chl_opt")
XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# ---------------------------------------------------------------------------
# Columns (live types from information_schema, checked 2026-10-10)
# ---------------------------------------------------------------------------

KEY = "pfa_name_safelives"
PERIOD = "financial_year"
VALUE_TYPES = (
    ("marac_count", "integer"),
    ("cases_discussed", "numeric(10,2)"),
    ("recommended_cases", "numeric(10,2)"),
    ("cases_per_10k_adult_females", "numeric(10,6)"),
    ("repeat_cases", "numeric(10,2)"),
    ("repeat_cases_pct", "numeric(8,4)"),
    ("children_in_household", "numeric(10,2)"),
    ("housing_referrals", "numeric(10,2)"),
)
VALUES = tuple(c for c, _ in VALUE_TYPES)
SCALES = {"marac_count": 0, "cases_discussed": 2, "recommended_cases": 2,
          "cases_per_10k_adult_females": 6, "repeat_cases": 2,
          "repeat_cases_pct": 4, "children_in_household": 2,
          "housing_referrals": 2}
FLAG_TYPES = (("value_flag", "text"),)          # editions-only
COMPARED = VALUES + ("value_flag",)
FLAGS = ("not_submitted", "no_population")
CASES_COLUMNS = VALUES[:7]
# the five measures the English police force areas must sum to England on
RECONCILED = ("marac_count", "cases_discussed", "recommended_cases",
              "repeat_cases", "children_in_household")
MAP_COLUMNS = ("cases_discussed", "cases_per_10k_adult_females")

# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------

SHEETS = ("Notes", "Cases", "Referral routes")
NOTES_TITLE_RE = re.compile(r"SafeLives Marac data England and Wales April "
                            r"([0-9]{4}) - March ([0-9]{4})")
CASES_TITLE_RE = re.compile(r"Cases discussed at multi-agency risk assessment "
                            r"conferences \((?:MARACs|Maracs)\), by police "
                            r"force area and region, year ending March "
                            r"([0-9]{4})")
# the .xls years append footnote 1 to the year ('March 20191')
REFERRAL_TITLE_RE = re.compile(r"Sources of referrals to multi-agency risk "
                               r"assessment conferences \((?:MARACs|Maracs)\), "
                               r"by police force area and region, year ending "
                               r"March ([0-9]{4})[0-9]?")
AREA = "Area Name"
# the Cases header (the .xls years' layout, read first) -> column
CASES_HEADER = (
    ("Area Name", None),
    ("Number of MARACs", "marac_count"),
    ("Number of cases discussed", "cases_discussed"),
    ("Recommended number of cases", "recommended_cases"),
    ("Number of cases per 10,000 adult females", "cases_per_10k_adult_females"),
    ("Number of repeat cases", "repeat_cases"),
    ("Percentage of repeat cases", "repeat_cases_pct"),
    ("Number of children in household", "children_in_household"),
)
CASES_HEADERS = tuple(h for h, _ in CASES_HEADER)
HEADER_OF = {c: h for h, c in CASES_HEADER if c}
HEADER_OF["housing_referrals"] = "Referral routes Housing Number"
# A header written differently, accepted only by name, with the file it was
# seen in (anything else halts).
HEADER_ALIASES = {
    "Number of Maracs": {
        "means": "Number of MARACs",
        "seen": "Marac-data-2019-2020-England-and-Wales.xls (2019-20), Cases "
                "sheet row 4"},
}
HOUSING = "Housing"
NUMBER = "Number"
FOOTNOTE_RE = re.compile(r"([0-9]+)\.\s*(.*)", re.S)
TOTALS = ("England and Wales", "England")
WALES = "Wales"
WELSH_FORCES = ("Dyfed Powys", "Gwent", "North Wales", "South Wales")
KNOWN_REGIONS = ("North East", "North West", "Yorkshire and The Humber",
                 "East Midlands", "West Midlands", "East", "London",
                 "South East", "South West", WALES)

# The markers the files use for a value that is not published.
MARKERS = {"No Data": "not_submitted", "No population data": "no_population"}
MARKER_ALIASES = {
    "No data": {
        "means": "No Data",
        "seen": "Marac-data-2022-2023.xls (2022-23), Cases row 46 and "
                "Referral routes row 47, Norfolk"},
}

# A file whose Notes title disagrees with its own sheet titles, accepted for
# those bytes only (sha256, split so the credential scan does not flag it).
NOTES_TITLE_ERRATA = {
    "b78d639b77366c5a404abecce46e276d" "b604f22c0c19ad872abef8b9f6a8971f": {
        "file": "Marac-data-2022-2023.xls",
        "year": "2022-23",
        "notes_says": (2021, 2022),
        "evidence": (
            "The Notes title reads 'April 2021 - March 2022', but the Notes "
            "sheet's own list of tables, and the titles of the Cases, "
            "Referral routes, Gender and Demographics sheets, all say 'year "
            "ending March 2023'; the data page links the file as 2022-23 "
            "(Marac-data-2022-2023.xls); its figures are not the 2021-22 "
            "file's (England cases discussed 102,011 against 104,743); and "
            "the held 2022-23 rows reproduce from it. The Notes title was "
            "not updated from the year before."),
    },
}

# ---------------------------------------------------------------------------
# The named zero rules (rule 1: not submitted or not counted is NULL, never
# 0). Each turns exactly the published cells it names into NULL with its
# flag, only where the file still publishes them in the form it names;
# otherwise the load halts (the rule must be reviewed with the evidence).
# ---------------------------------------------------------------------------

NOT_SUBMITTED_FORM = {"marac_count": 0, "cases_discussed": 0,
                      "recommended_cases": 0,
                      "cases_per_10k_adult_females": "No population data",
                      "repeat_cases": 0, "repeat_cases_pct": "#DIV/0!",
                      "children_in_household": 0, "housing_referrals": 0}

NORFOLK_2023_24_NOT_SUBMITTED = {
    "name": "NORFOLK_2023_24_NOT_SUBMITTED",
    "pfa": "Norfolk", "years": ("2023-24",), "columns": VALUES,
    "flag": "not_submitted", "published": NOT_SUBMITTED_FORM,
    "decided": "the 2026-10-10 design (spec finding 10); listed for Scott",
    "evidence": (
        "Marac-data-2023-2024.xlsx publishes Norfolk as 0 MARACs, 0 cases "
        "discussed, 0 recommended cases, 'No population data' for the rate, "
        "0 repeat cases, '#DIV/0!' for the repeat percentage, 0 children and "
        "0 in every Referral routes count. SafeLives published 'No Data' for "
        "Norfolk in 2022-23 ('No data'), 2024-25 and 2025-26, and Norfolk "
        "had 3 MARACs and 1,004 to 1,791 cases a year from 2018-19 to "
        "2021-22: the zeros are a non-submission, not a count."),
}

WEST_MIDLANDS_2023_24_NOT_SUBMITTED = {
    "name": "WEST_MIDLANDS_2023_24_NOT_SUBMITTED",
    "pfa": "West Midlands", "years": ("2023-24",), "columns": VALUES,
    "flag": "not_submitted", "published": NOT_SUBMITTED_FORM,
    "decided": "Scott, 2026-10-10 (the West Midlands force is stored; its "
               "2023-24 row is a non-submission like Norfolk's)",
    "evidence": (
        "Marac-data-2023-2024.xlsx publishes the West Midlands police force "
        "area in exactly the form of Norfolk 2023-24 (0 MARACs, 0 cases, "
        "'No population data', '#DIV/0!', 0 in every Referral routes count). "
        "The West Midlands region row (17 MARACs, 4,191 cases) equals "
        "Staffordshire, Warwickshire and West Mercia alone (10 + 3 + 4 "
        "MARACs; 1,942 + 986 + 1,263 cases), so none of the force's MARACs "
        "is counted anywhere in the year. The force had 7 MARACs and 7,810 "
        "cases in 2025-26."),
}

ZERO_RULES = (NORFOLK_2023_24_NOT_SUBMITTED,
              WEST_MIDLANDS_2023_24_NOT_SUBMITTED)

# The NULL/value changes against the tip (rule 1.10) a load may store, each
# named and decided: --acknowledge-flips NAME releases exactly the cells it
# lists per year ('<force>/<column>': (tip value, new value), at the live
# scale, None for NULL) and nothing else.
ACKNOWLEDGED_FLIPS = {
    "s17-restored-zeros-2026-10": {
        "decided": "Scott, 2026-10-10 (a published 0 stays 0; only the "
                   "Norfolk and West Midlands 2023-24 rules make NULL)",
        "why": ("edition 1 (as loaded) holds as NULL the published zeros the "
                "n8n code read as NULL (parseFloat(x) || null); the files "
                "publish them as 0 and no rule names them, so they are "
                "restored to 0"),
        "periods": {
            "2018-19": {"City of London/repeat_cases": (None, "0.00"),
                        "City of London/repeat_cases_pct": (None, "0.0000"),
                        "City of London/housing_referrals": (None, "0.00"),
                        "Gloucestershire/housing_referrals": (None, "0.00"),
                        "Leicestershire/housing_referrals": (None, "0.00")},
            "2019-20": {"City of London/housing_referrals": (None, "0.00")},
            "2020-21": {"City of London/housing_referrals": (None, "0.00"),
                        "Gloucestershire/housing_referrals": (None, "0.00")},
            "2021-22": {"Leicestershire/housing_referrals": (None, "0.00")},
            "2022-23": {"City of London/children_in_household":
                        (None, "0.00")},
            "2023-24": {"Lancashire/housing_referrals": (None, "0.00")},
            "2024-25": {"City of London/housing_referrals": (None, "0.00"),
                        "Lancashire/housing_referrals": (None, "0.00")},
        },
    },
}

# ---------------------------------------------------------------------------
# Stop conditions (percent; calibrated on the eight held files, 2018-19 to
# 2025-26: England cases discussed, as the sum of the forces, moved by at
# most 12.3% a year (2023-24 to 2024-25, the West Midlands force not
# submitted in 2023-24); at most 4 forces moved by more than 40% in one year
# (2022-23 to 2023-24))
# ---------------------------------------------------------------------------

NEW_ENGLAND_PCT = 25        # new year: England cases discussed
NEW_PFA_PCT = 40            # new year: a force's cases discussed ...
NEW_MAX_PFAS = 6            # ... moving by more than that, in more forces

# ---------------------------------------------------------------------------
# The held state (migrate-legacy preconditions), surveyed 2026-10-10
# ---------------------------------------------------------------------------

# (rows, survey hash, distinct loaded_at, rows per year); the survey hash is
# md5 of the columns but loaded_at, each coalesce(col::text,'~'), '|'-joined,
# rows ordered by financial_year, pfa_name_safelives, joined by LF.
LEGACY_LIVE = (304, "d885fc22b822d6bb" "a698c239498f8e2d", 2,
               tuple((y, 38) for y in ("2018-19", "2019-20", "2020-21",
                                       "2021-22", "2022-23", "2023-24",
                                       "2024-25", "2025-26")))
# la_pfa_mapping (read, never written): (rows, md5 of lad24cd, la_name,
# pfa_name, pfa_name_safelives, population_weight, source ordered by lad24cd)
LEGACY_PFA_MAPPING = (296, "590fc32499fef391" "089927c8c52e427e")
# its 39 English police force area names (pfa_name_safelives)
PFA_NAMES_2026_10 = (
    "Avon and Somerset", "Bedfordshire", "Cambridgeshire", "Cheshire",
    "City of London", "Cleveland", "Cumbria", "Derbyshire", "Devon & Cornwall",
    "Dorset", "Durham", "Essex", "Gloucestershire", "Greater Manchester",
    "Hampshire", "Hertfordshire", "Humberside", "Kent", "Lancashire",
    "Leicestershire", "Lincolnshire", "Merseyside", "Metropolitan Police",
    "Norfolk", "North Yorkshire", "Northamptonshire", "Northumbria",
    "Nottinghamshire", "South Yorkshire", "Staffordshire", "Suffolk", "Surrey",
    "Sussex", "Thames Valley", "Warwickshire", "West Mercia", "West Midlands",
    "West Yorkshire", "Wiltshire")
_UPLOADS = "https://safelives.org.uk/wp-content/uploads/"
# the eight held files (sha256 of the bytes, split for the credential scan)
# and the data page link each was downloaded from on 2026-10-10
LEGACY_FILES = {
    y: {"path": RAW_DIR / name, "sha256": a + b, "url": _UPLOADS + link}
    for y, name, a, b, link in (
        ("2018-19", "Marac-data-2018-2019-England-and-Wales.xls",
         "0f92d67dfe9c8efc80671f50da36b003", "d2036a8b8191e919ca79ba3168360f39",
         "2023/12/Marac-data-2018-2019-England-and-Wales.xls"),
        ("2019-20", "Marac-data-2019-2020-England-and-Wales.xls",
         "8d031fea4f37c3dc6f55649ff55fe1a8", "1d91e6b2635b28df674b18697f2cacca",
         "2023/12/Marac-data-2019-2020-England-and-Wales.xls"),
        ("2020-21", "Marac-data-2020-to-2021.xls",
         "474aa4dff5fbd2d0c907c9be635590a2", "3e73f072195eb249c892a188d7694c3e",
         "2024/01/Marac-data-2020-to-2021.xls"),
        ("2021-22", "Marac-data-2021-2022-for-publication.xls",
         "d2b4e0f7bf2f5acd9b9ae58d7e43958c", "38fccb293ef90ead55e02ff9c90b9258",
         "2023/12/Marac-data-2021-2022-for-publication.xls"),
        ("2022-23", "Marac-data-2022-2023.xls",
         "b78d639b77366c5a404abecce46e276d", "b604f22c0c19ad872abef8b9f6a8971f",
         "2023/12/Marac-data-2022-2023.xls"),
        ("2023-24", "Marac-data-2023-2024.xlsx",
         "60a18547cfff9b7139b7adc228e61b80", "555a118cf8f7dd9a1c16e84e8145111c",
         "Marac-data-2023-2024.xlsx"),
        ("2024-25", "Marac-data-2024-2025.xlsx",
         "59bd0e870a0ed0516aa397878ba86ad4", "a822847b1bfcb8d1f90365004bf6b794",
         "Marac-data-2024-2025.xlsx"),
        ("2025-26", "MARAC-DATA-2025-2026.xlsx",
         "3c1f06380fce49f530df8f20ab5eb26d", "0348d7902694b68916abaf8254bfa758",
         "MARAC-DATA-2025-2026.xlsx"))}
# keys every held file has and the held table lacks, each explained: not
# in edition 1 (as loaded = live as held); the first load adds them to the
# next edition and refresh-latest inserts them (--accept-key-changes)
LEGACY_EXPLAINED_KEYS = {
    "West Midlands": (
        "the held table has no row for the West Midlands police force area "
        "in any year, while every file publishes it and la_pfa_mapping maps "
        "seven authorities (Birmingham, Coventry, Dudley, Sandwell, "
        "Solihull, Walsall, Wolverhampton) to it. Neither the n8n load "
        "(2018-19 to 2024-25) nor the direct load of 2026-08-20 (2025-26) "
        "stored it; in the files the force shares its name with the West "
        "Midlands region row"),
}
# edition 1 'as loaded' source_file, by the held rows' load date
AS_LOADED_SOURCES = {
    date(2026, 3, 30): (
        "as loaded: n8n run-log 22 to 24 (2026-03-30) from a CSV converted "
        "from the SafeLives files in a chat session (parseFloat(x) || null: "
        "published zeros held as NULL); value_flag read from the held file"),
    date(2026, 8, 20): (
        "as loaded: the direct load of 2026-08-20 (commit 33c5ec4; no "
        "run-log row, no committed code); value_flag read from the held "
        "file"),
}


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _norm(s) -> str:
    return " ".join(str("" if s is None else s).split())


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or \
        (isinstance(v, str) and not v.strip())


def _is_number(v) -> bool:
    return isinstance(v, numbers.Real) and not isinstance(v, bool) and \
        not (isinstance(v, float) and math.isnan(v))


def _label(y2: int) -> str:
    """'2025-26' of the year ending March 2026."""
    return f"{y2 - 1}-{y2 % 100:02d}"


def label_ok(text) -> str:
    mm = re.fullmatch(r"([0-9]{4})-([0-9]{2})", str(text or ""))
    if not mm or (int(mm.group(1)) + 1) % 100 != int(mm.group(2)):
        raise ValueError(f"not a financial year YYYY-YY: {text!r}")
    return mm.group(0)


def _text(v) -> "str | None":
    """A value at its scale, as text (None for NULL)."""
    return None if v is None else str(v)


def _show(v) -> str:
    return "NULL" if v is None else str(v)


def _same(a, b) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return a == b


def content_sha256(path) -> str:
    """SHA-256 of the file's bytes."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _iso(dt) -> str:
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _when(text) -> "datetime | None":
    if not text:
        return None
    try:
        dt = datetime.fromisoformat(str(text).replace("Z", "+00:00"))
    except ValueError:
        return None
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def _file_source(path) -> str:
    """Where a file was read: the path under the repo (data/...), else the
    absolute path."""
    p = Path(path).resolve()
    try:
        return p.relative_to(REPO).as_posix()
    except ValueError:
        return p.as_posix()


# ---------------------------------------------------------------------------
# Discovery: the data page
# ---------------------------------------------------------------------------

_HREF_RE = re.compile(r"""href\s*=\s*["']([^"']+)["']""", re.I)
_TWO_YEARS_RE = re.compile(r"(?<![0-9])([0-9]{4})-(?:to-)?([0-9]{4})(?![0-9])",
                           re.I)
_ANY_YEAR_RE = re.compile(r"(?<![0-9])(?:19|20)[0-9]{2}(?![0-9])")


class Links(dict):
    """{yyyy-yy: url}, with .all (every workbook link, in page order) and
    .ignored (workbook links that name no year)."""
    all: list
    ignored: list


def data_page_links(html, held=(), base=DATA_PAGE) -> Links:
    """{yyyy-yy: url} of the page's workbook links (.xlsx or .xls) whose file
    names carry two years 'YYYY-YYYY' or 'YYYY-to-YYYY' (relative links
    resolved against base; a link repeated on the page counts once).
    ValueError listing every workbook link when: there is none; two links
    claim one year; a link names a year but not one financial year in that
    form; or a held year has no link."""
    urls = []
    for h in _HREF_RE.findall(html or ""):
        u = urljoin(base, html_lib.unescape(h).strip())
        if urlparse(u).path.lower().endswith((".xlsx", ".xls")) and \
                u not in urls:
            urls.append(u)
    listing = "; workbook links on the page: " + (", ".join(urls) or "none")
    if not urls:
        raise ValueError("the data page lists no .xlsx or .xls link" + listing)
    by_year, odd, ignored = {}, [], []
    for u in urls:
        name = Path(unquote(urlparse(u).path)).name
        hits = _TWO_YEARS_RE.findall(name)
        if len(hits) == 1 and int(hits[0][1]) == int(hits[0][0]) + 1:
            by_year.setdefault(_label(int(hits[0][1])), []).append(u)
        elif hits or _ANY_YEAR_RE.search(name):
            odd.append(u)
        else:
            ignored.append(u)
    if odd:
        raise ValueError(f"workbook link(s) naming a year but not one "
                         f"financial year as YYYY-YYYY or YYYY-to-YYYY: "
                         f"{odd}; not choosing around them" + listing)
    two = {y: us for y, us in by_year.items() if len(us) > 1}
    if two:
        raise ValueError(f"two or more links for one year: {two}; refusing "
                         "to choose one" + listing)
    gone = sorted(set(held) - set(by_year))
    if gone:
        raise ValueError(f"no link for the held year(s) {gone}" + listing)
    out = Links({y: us[0] for y, us in by_year.items()})
    out.all, out.ignored = urls, ignored
    return out


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

def _session(session):
    if session is not None:
        return session
    import requests
    return requests


def _get(session, url):
    try:
        return session.get(url, headers={"User-Agent": USER_AGENT},
                           timeout=120, allow_redirects=True)
    except Exception as e:  # noqa: BLE001 (any failure halts)
        halt(f"GET {url} failed, nothing downloaded: {e}")


def refusal(resp) -> "str | None":
    """Why a reply is a refusal (a 403, 429 or 503, or a Cloudflare challenge
    page), or None."""
    status = getattr(resp, "status_code", 200)
    if status in (403, 429, 503):
        return f"HTTP {status}"
    body = resp.content or b""
    if isinstance(body, str):
        body = body.encode("utf-8", "replace")
    head = body[:20000]
    if any(mark in head for mark in CHALLENGE_MARKS):
        return "a Cloudflare challenge page"
    return None


def page_refused_message(url, reason="refused") -> str:
    return (f"SafeLives refused the request for the data page {url} "
            f"({reason}): download the year's workbook by hand from "
            f"{DATA_PAGE} into data/raw/s17_marac/ and rerun with --file PATH "
            "--no-page")


def _keep(dest, name, body, what) -> Path:
    """The same-name rule: save body as dest/name; a same-named file with
    the same content is kept as it is; one with different content is never
    replaced (body is saved beside it as <stem>-<sha8><suffix>). Returns the
    path holding body."""
    dest.mkdir(parents=True, exist_ok=True)
    target = dest / name
    new_sha = hashlib.sha256(body).hexdigest()
    if not target.exists():
        target.write_bytes(body)
        print(f"downloaded {what} {name}: {len(body):,} bytes, sha256 "
              f"{new_sha[:16]}")
        return target
    old_sha = content_sha256(target)
    if old_sha == new_sha:
        print(f"{name} already held with the same content (sha256 "
              f"{new_sha[:16]}); kept as it is")
        return target
    alt = dest / f"{target.stem}-{new_sha[:8]}{target.suffix}"
    if not (alt.exists() and content_sha256(alt) == new_sha):
        alt.write_bytes(body)
    print(f"NOTE: {name} already exists with different content (sha256 "
          f"{old_sha[:16]}, new {new_sha[:16]}); kept it untouched and saved "
          f"this download as {alt.name}")
    return alt


def fetch_page(url=REGISTRY_URL, session=None, dest=None) -> tuple:
    """(html, final URL) of the data page (redirects followed and printed),
    with the plain User-Agent. Saved under the same-name rule as
    marac-data-page_<today>.html in dest (default RAW_DIR). A refusal halts
    with the hand-download message."""
    r = _get(_session(session), url)
    final = str(getattr(r, "url", None) or url)
    for h in getattr(r, "history", None) or []:
        print(f"redirect: {getattr(h, 'status_code', '?')} "
              f"{getattr(h, 'url', '?')}")
    if final != url:
        print(f"redirected: {url} -> {final}")
    why = refusal(r)
    if why:
        halt(page_refused_message(url, why))
    if getattr(r, "status_code", 200) >= 400:
        halt(f"GET {url} failed: HTTP {r.status_code}; nothing read")
    body = r.content if isinstance(r.content, bytes) else \
        str(r.content).encode("utf-8")
    _keep(Path(dest or globals()["RAW_DIR"]),
          f"marac-data-page_{date.today():%Y-%m-%d}.html", body, "data page")
    return body.decode("utf-8", "replace"), final


def fetch(url, dest, session=None) -> tuple:
    """Download a workbook link into the directory dest (redirects followed)
    and return (path to read, final URL). The file name is the final URL's.
    A refusal (403, 429, 503, a Cloudflare challenge page), a name that is
    not .xlsx or .xls, or a body that is not a readable workbook of that
    kind halts with REFUSED, nothing kept. The same-name rule: a same-named
    file with the same content is kept; one with different content is never
    replaced (the download is saved beside it as <stem>-<sha8><suffix> and
    that is the one read)."""
    dest = Path(dest)
    r = _get(_session(session), url)
    final = str(getattr(r, "url", None) or url)
    if final != url:
        print(f"redirected: {url} -> {final}")
    msg = REFUSED.format(url=url)
    why = refusal(r)
    if why:
        halt(f"{msg} ({why}; nothing kept)")
    if getattr(r, "status_code", 200) >= 400:
        halt(f"GET {url} failed: HTTP {r.status_code}; nothing downloaded")
    body = r.content or b""
    name = Path(unquote(urlparse(final).path)).name
    suffix = Path(name).suffix.lower()
    magic = {".xlsx": XLSX_MAGIC, ".xls": XLS_MAGIC}.get(suffix)
    if magic is None or not body.startswith(magic):
        halt(f"{msg} (the reply to {final} is not an .xlsx or .xls "
             "workbook body; nothing kept)")
    dest.mkdir(parents=True, exist_ok=True)
    tmp = dest / (Path(name).stem + ".download.tmp" + suffix)
    tmp.write_bytes(body)
    try:
        _sheets(tmp)
    except ValueError as e:
        tmp.unlink()
        halt(f"{msg} (the reply does not open as a workbook: {e}; nothing "
             "kept)")
    tmp.unlink()
    return _keep(dest, name, body, "workbook"), final


# ---------------------------------------------------------------------------
# Reading a workbook
# ---------------------------------------------------------------------------

def _clean(v):
    if isinstance(v, str) and not v.strip():
        return None
    return v


def _sheets(path) -> dict:
    """{sheet name: [row tuples]} of an .xlsx (openpyxl, values) or .xls
    (xlrd 2.0.1) workbook: empty cells None, text as published, numbers as
    numbers, error cells as their text ('#DIV/0!'), rows padded to one
    width. ValueError if it is neither or does not open."""
    p = Path(path)
    ext = p.suffix.lower()
    out = {}
    if ext == ".xlsx":
        import openpyxl
        try:
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"{p.name}: not a readable .xlsx workbook "
                             f"({e})") from None
        try:
            for name in wb.sheetnames:
                out[name] = [tuple(_clean(v) for v in r)
                             for r in wb[name].iter_rows(values_only=True)]
        finally:
            wb.close()
    elif ext == ".xls":
        try:
            import xlrd
        except ImportError:
            raise ValueError("xlrd is not installed (pip install "
                             "xlrd==2.0.1); the .xls years cannot be read "
                             "without it") from None
        try:
            book = xlrd.open_workbook(str(p))
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"{p.name}: not a readable .xls workbook "
                             f"({e})") from None
        try:
            for sh in book.sheets():
                rows = []
                for i in range(sh.nrows):
                    row = []
                    for j in range(sh.ncols):
                        t, v = sh.cell_type(i, j), sh.cell_value(i, j)
                        if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            v = None
                        elif t == xlrd.XL_CELL_ERROR:
                            v = xlrd.error_text_from_code.get(v, f"#ERR{v}")
                        elif t == xlrd.XL_CELL_BOOLEAN:
                            v = bool(v)
                        elif t == xlrd.XL_CELL_DATE:
                            v = f"(a date cell {v!r})"
                        row.append(_clean(v))
                    rows.append(tuple(row))
                out[sh.name] = rows
        finally:
            book.release_resources()
    else:
        raise ValueError(f"{p.name}: not an .xlsx or .xls file")
    for name, rows in out.items():
        width = max((len(r) for r in rows), default=0)
        out[name] = [r + (None,) * (width - len(r)) for r in rows]
    return out


def _cells(rows) -> list:
    return [_norm(v) for r in rows for v in r if isinstance(v, str)
            and not _blank(v)]


def _one(cells, rx, what, where):
    hits = [mm for mm in (rx.fullmatch(c) for c in cells) if mm]
    if len(hits) != 1:
        raise ValueError(f"{where}: {len(hits)} cells like {what!r}, expected "
                         f"one; cells seen: {cells[:12]}")
    return hits[0]


def _header_row(rows, where) -> tuple:
    hits = [(i, j) for i, r in enumerate(rows) for j, v in enumerate(r)
            if _norm(v) == AREA]
    if len(hits) != 1:
        raise ValueError(f"{where}: {len(hits)} '{AREA}' header cells, "
                         "expected one")
    return hits[0]


def _force_name(written, footnotes) -> tuple:
    """(name, [footnote text]) of a written area name: a trailing number is
    a footnote marker only when the sheet defines that footnote; otherwise
    the name is kept as written (and is then an unknown name)."""
    w = _norm(written)
    mm = re.fullmatch(r"(.*\D)([0-9]+)", w)
    if mm and int(mm.group(2)) in footnotes:
        return mm.group(1).strip(), [footnotes[int(mm.group(2))]]
    return w, []


def _blocks(rows, start, area_j, last_j, where) -> tuple:
    """(blocks, footnotes) of the rows after a header: blocks of
    consecutive rows with an area name, split by wholly empty rows, each a
    list of (sheet row number, written name, row); footnotes {n: text} from
    the rows that start with '<n>.' (after them only footnotes may follow).
    ValueError on a row with values but no name, a cell beyond the headed
    columns in a data row, or data after the footnotes."""
    blocks, cur, foot = [], [], {}
    for i in range(start, len(rows)):
        r = rows[i]
        name = r[area_j] if area_j < len(r) else None
        if all(_blank(v) for v in r):
            if cur:
                blocks.append(cur)
                cur = []
            continue
        mm = FOOTNOTE_RE.fullmatch(_norm(name)) if isinstance(name, str) \
            else None
        if mm:
            foot[int(mm.group(1))] = _norm(name)
            continue
        if foot:
            raise ValueError(f"{where} row {i + 1}: data after the footnotes "
                             f"({[v for v in r if not _blank(v)][:4]})")
        if _blank(name):
            raise ValueError(f"{where} row {i + 1}: values but no area name "
                             f"({[v for v in r if not _blank(v)][:4]})")
        stray = [j for j, v in enumerate(r)
                 if not _blank(v) and (j < area_j or j > last_j)]
        if stray:
            raise ValueError(f"{where} row {i + 1}: a cell beyond the headed "
                             f"columns (column {stray[0] + 1})")
        cur.append((i + 1, name, r))
    if cur:
        blocks.append(cur)
    return blocks, foot


def _regions(blocks, foot, where, *, check_regions) -> tuple:
    """(totals {name: row}, regions [{name, row, forces: [{name, written,
    row, cells, footnotes}]}]) of a table's blocks: the England and Wales
    block, the England block, then one block per region (its first row the
    region, the rest its forces)."""
    names = [[_norm(n) for _, n, _ in b] for b in blocks]
    if len(blocks) < 3 or names[0] != [TOTALS[0]] or names[1] != [TOTALS[1]]:
        raise ValueError(f"{where}: the table does not start with the "
                         f"'{TOTALS[0]}' and '{TOTALS[1]}' rows, each on its "
                         f"own (blocks seen: {[b[:2] for b in names[:4]]})")
    totals = {TOTALS[0]: blocks[0][0][2], TOTALS[1]: blocks[1][0][2]}
    regions = []
    for b in blocks[2:]:
        rnum, rwritten, rrow = b[0]
        rname, _ = _force_name(rwritten, foot)
        if check_regions and rname not in KNOWN_REGIONS:
            raise ValueError(f"{where} row {rnum}: {rname!r} heads a block "
                             f"but is not a known region ({KNOWN_REGIONS})")
        forces = []
        for num, written, row in b[1:]:
            name, notes = _force_name(written, foot)
            forces.append({"name": name, "written": _norm(written),
                           "row": num, "cells": row, "footnotes": notes})
        if not forces:
            raise ValueError(f"{where} row {rnum}: region {rname!r} has no "
                             "police force area rows")
        regions.append({"name": rname, "row": rnum, "cells": rrow,
                        "forces": forces})
    return totals, regions


def _read_cases(rows, fname) -> dict:
    where = f"{fname}: Cases"
    hi, aj = _header_row(rows, where)
    raw = [_norm(v) for v in rows[hi][aj:]]
    while raw and not raw[-1]:
        raw.pop()
    aliases, header = [], []
    for h in raw:
        if h in HEADER_ALIASES:
            aliases.append(h)
            h = HEADER_ALIASES[h]["means"]
        header.append(h)
    if tuple(header) != CASES_HEADERS:
        unknown = [h for h in header if h not in CASES_HEADERS]
        missing = [h for h in CASES_HEADERS if h not in header]
        raise ValueError(f"{where} header {raw} is not the expected set: "
                         f"unknown {unknown}, missing {missing} (expected "
                         f"exactly {list(CASES_HEADERS)}, in order); record "
                         "a differing header as a named alias in "
                         "HEADER_ALIASES with its file as evidence, or "
                         "correct the mapping deliberately")
    cols = {c: aj + k for k, (_, c) in enumerate(CASES_HEADER) if c}
    last = aj + len(CASES_HEADER) - 1
    blocks, foot = _blocks(rows, hi + 1, aj, last, where)
    totals, regions = _regions(blocks, foot, where, check_regions=True)

    def cells_of(row):
        return {c: row[j] for c, j in cols.items()}
    for reg in regions:
        reg["values"] = cells_of(reg["cells"])
        for f in reg["forces"]:
            f["values"] = cells_of(f["cells"])
    return {"header": raw, "header_aliases": aliases, "cols": cols,
            "totals": {k: cells_of(v) for k, v in totals.items()},
            "regions": regions, "footnotes": foot}


def _read_referrals(rows, fname) -> dict:
    where = f"{fname}: Referral routes"
    hi, aj = _header_row(rows, where)
    if hi + 1 >= len(rows):
        raise ValueError(f"{where}: no sub-header row under the header")
    head, sub = rows[hi], rows[hi + 1]
    hits = [j for j in range(len(head)) if _norm(head[j]) == HOUSING
            and _norm(sub[j]) == NUMBER]
    if len(hits) != 1:
        raise ValueError(f"{where}: {len(hits)} '{HOUSING}' columns with a "
                         f"'{NUMBER}' sub-header, expected one (the column is "
                         "located by its header text, never by position)")
    hj = hits[0]
    blocks, foot = _blocks(rows, hi + 2, aj, len(head) - 1, where)
    _, regions = _regions(blocks, foot, where, check_regions=False)
    forces = {}
    for reg in regions:
        for f in reg["forces"]:
            if f["name"] in forces:
                raise ValueError(f"{where} row {f['row']}: {f['name']} "
                                 f"twice (also row {forces[f['name']]['row']})")
            forces[f["name"]] = {"row": f["row"], "region": reg["name"],
                                 "housing": f["cells"][hj],
                                 "footnotes": f["footnotes"]}
    return {"housing_col": hj, "forces": forces, "footnotes": foot}


def read_workbook(path, year=None) -> dict:
    """The facts of one SafeLives workbook: {path, name, identity
    (manual_input.file_identity), sha256, year ('yyyy-yy'), notes_years
    (y1, y2), cases_year, notes_title, cases_title, erratum (the
    NOTES_TITLE_ERRATA entry used, or None), modified, rank ((year,
    modified)), header, header_aliases, housing_col, cases, referrals}.

    ValueError (naming what was seen) on: a missing Notes, Cases or Referral
    routes sheet; not exactly one Notes title, Cases title or Referral
    routes title; a Notes title not spanning one year; Notes and Cases years
    that disagree (unless a named erratum covers these bytes); a Referral
    routes year other than the Cases year; `year` given and not the file's;
    a Cases header that is not exactly the eight columns (aliases by name
    only); not exactly one Housing 'Number' column; a broken table layout."""
    p = Path(path)
    ident = manual_input.file_identity(p)
    sheets = _sheets(p)
    missing = [s for s in SHEETS if s not in sheets]
    if missing:
        raise ValueError(f"{p.name}: sheet(s) {missing} missing; sheets "
                         f"{list(sheets)}")
    nt = _one(_cells(sheets["Notes"]), NOTES_TITLE_RE,
              "SafeLives Marac data England and Wales April <yyyy> - March "
              "<yyyy>", f"{p.name}: Notes")
    ct = _one(_cells(sheets["Cases"][:4]), CASES_TITLE_RE,
              "Cases discussed at ... year ending March <yyyy>",
              f"{p.name}: Cases")
    rt = _one(_cells(sheets["Referral routes"][:4]), REFERRAL_TITLE_RE,
              "Sources of referrals to ... year ending March <yyyy>",
              f"{p.name}: Referral routes")
    ny = (int(nt.group(1)), int(nt.group(2)))
    cy = int(ct.group(1))
    if ny[1] != ny[0] + 1:
        raise ValueError(f"{p.name}: the Notes title spans April {ny[0]} to "
                         f"March {ny[1]}, not one year")
    if int(rt.group(1)) != cy:
        raise ValueError(f"{p.name}: the Referral routes title says year "
                         f"ending March {rt.group(1)}, the Cases title March "
                         f"{cy}")
    label = _label(cy)
    erratum = None
    if ny != (cy - 1, cy):
        e = globals()["NOTES_TITLE_ERRATA"].get(ident["sha256"])
        if e and tuple(e["notes_says"]) == ny and e["year"] == label:
            erratum = e
        else:
            raise ValueError(f"{p.name}: the Notes title says April {ny[0]} - "
                             f"March {ny[1]}, the Cases title says year ending "
                             f"March {cy}: the file's own titles disagree")
    if year is not None and year != label:
        raise ValueError(f"{p.name}: the file is {label} (its Notes and Cases "
                         f"titles), not {year}")
    cases = _read_cases(sheets["Cases"], p.name)
    refs = _read_referrals(sheets["Referral routes"], p.name)
    modified = _when(ident.get("modified"))
    return {"path": p, "name": p.name, "identity": ident,
            "sha256": ident["sha256"], "year": label, "notes_years": ny,
            "cases_year": cy, "notes_title": nt.group(0),
            "cases_title": ct.group(0), "erratum": erratum,
            "modified": modified, "rank": (label, modified),
            "header": cases["header"],
            "header_aliases": cases["header_aliases"],
            "housing_col": refs["housing_col"], "cases": cases,
            "referrals": refs}


# ---------------------------------------------------------------------------
# Cells and records
# ---------------------------------------------------------------------------

def cell(v, where) -> tuple:
    """(value, flag) of a published value cell: a number -> (number, None);
    'No Data' -> (None, 'not_submitted'); 'No population data' -> (None,
    'no_population'); a named alias of a marker -> the marker's. Anything
    else (a blank, other text, an error value such as '#DIV/0!', a boolean)
    raises ValueError naming `where`: it is never read as 0 or as NULL."""
    if _is_number(v):
        return v, None
    if isinstance(v, str):
        s = _norm(v)
        if s in MARKERS:
            return None, MARKERS[s]
        if s in MARKER_ALIASES:
            return None, MARKERS[MARKER_ALIASES[s]["means"]]
    raise ValueError(f"{where}: {v!r} is neither a number nor a SafeLives "
                     f"marker ({', '.join(repr(k) for k in MARKERS)}); it is "
                     "never read as 0 or as NULL")


def to_scale(v, col, where):
    """A published number at the live column's scale (half up): marac_count
    a whole number (a fraction halts), the others Decimal at SCALES[col]. A
    negative value halts."""
    d = Decimal(repr(float(v)))
    if d < 0:
        raise ValueError(f"{where}: negative value {v!r}")
    if col == "marac_count":
        if d != d.to_integral_value():
            raise ValueError(f"{where}: {v!r} MARACs is not a whole number")
        return int(d)
    return d.quantize(Decimal(1).scaleb(-SCALES[col]), rounding=ROUND_HALF_UP)


class Records(list):
    """A year's records, with .applied (the named rules' cells: {rule, pfa,
    year, column, published, flag}), .markers ({pfa: flag} of the rows a
    marker made NULL) and .footnotes ({pfa: [footnote text]})."""
    applied: list
    markers: dict
    footnotes: dict


def _rules_for(name, year) -> list:
    return [r for r in globals()["ZERO_RULES"]
            if r["pfa"] == name and year in r["years"]]


def _published_as(value, form) -> bool:
    if _is_number(form):
        return _is_number(value) and float(value) == float(form)
    return isinstance(value, str) and _norm(value) == _norm(form)


def zero_rules(records, year) -> list:
    """Apply the named rules (ZERO_RULES) to a year's records in place: for
    each rule naming the year, every cell it names must be published exactly
    in the rule's form (record['_raw']); it becomes NULL and the rule's flag
    is added to the row's reasons. ValueError when a rule's force is not in
    the records or a cell is published otherwise (the rule no longer matches
    the file). Returns the cells applied ({rule, pfa, year, column,
    published, flag})."""
    by = {r[KEY]: r for r in records}
    applied = []
    for rule in globals()["ZERO_RULES"]:
        if year not in rule["years"]:
            continue
        r = by.get(rule["pfa"])
        if r is None:
            raise ValueError(f"{rule['name']}: {rule['pfa']} is not in the "
                             f"{year} records")
        for col in rule["columns"]:
            got, form = r["_raw"][col], rule["published"][col]
            if not _published_as(got, form):
                raise ValueError(
                    f"{rule['name']}: {rule['pfa']} {year} {col} is published "
                    f"as {got!r}, the rule names {form!r}; the rule no longer "
                    "matches the file and must be reviewed with the evidence")
            r[col] = None
            r["_flags"].add(rule["flag"])
            applied.append({"rule": rule["name"], "pfa": rule["pfa"],
                            "year": year, "column": col, "published": got,
                            "flag": rule["flag"]})
    return applied


def _english(wb, pfa_names) -> list:
    """The English force rows of the Cases sheet, checked against pfa_names
    (la_pfa_mapping.pfa_name_safelives) and the Welsh block."""
    fname = wb["name"]
    welsh_named = sorted(set(pfa_names) & set(WELSH_FORCES))
    if welsh_named:
        raise ValueError(f"the English police force area names include Welsh "
                         f"forces {welsh_named}: Welsh rows would be taken "
                         "as English")
    english, seen = [], {}
    for reg in wb["cases"]["regions"]:
        for f in reg["forces"]:
            where = f"{fname}: Cases row {f['row']}"
            if reg["name"] == WALES:
                if f["name"] not in WELSH_FORCES:
                    raise ValueError(f"{where}: {f['name']!r} in the Wales "
                                     f"block is not a Welsh force "
                                     f"({WELSH_FORCES})")
                continue
            if f["name"] in WELSH_FORCES:
                raise ValueError(f"{where}: the Welsh force {f['name']!r} in "
                                 f"the English region {reg['name']!r}; Welsh "
                                 "rows are never taken as English")
            if f["name"] not in pfa_names:
                raise ValueError(f"{where}: {f['written']!r} is not an English "
                                 "police force area of la_pfa_mapping "
                                 "(pfa_name_safelives); names are matched "
                                 "exactly")
            if f["name"] in seen:
                raise ValueError(f"{where}: {f['name']} twice (also row "
                                 f"{seen[f['name']]})")
            seen[f["name"]] = f["row"]
            english.append(f)
    gone = sorted(set(pfa_names) - set(seen))
    if gone:
        raise ValueError(f"{fname}: English police force area(s) {gone} of "
                         "la_pfa_mapping are not in the Cases sheet")
    return english


def records(wb, year, pfa_names) -> Records:
    """The year's records from a read workbook: one per English police force
    area (exactly pfa_names), {pfa_name_safelives, financial_year, the eight
    values at the live scales, value_flag}, sorted by name. The named rules
    (zero_rules) are applied; .applied lists their cells. ValueError on a
    year other than the file's, a force missing, unknown, repeated or Welsh,
    a force missing from Referral routes, a cell that is neither a number
    nor a marker (cell), a fractional MARAC count, or two different reasons
    for the NULLs of one row."""
    if year != wb["year"]:
        raise ValueError(f"{wb['name']} is {wb['year']}, not {year}")
    english = _english(wb, pfa_names)
    refs = wb["referrals"]["forces"]
    ref_english = {n for n, f in refs.items() if f["region"] != WALES}
    names = {f["name"] for f in english}
    if ref_english != names:
        raise ValueError(f"{wb['name']}: Referral routes English forces "
                         f"differ from the Cases sheet's: only in Referral "
                         f"routes {sorted(ref_english - names)}, only in "
                         f"Cases {sorted(names - ref_english)}")
    out = Records()
    out.footnotes, out.markers = {}, {}
    for f in english:
        name = f["name"]
        raw = dict(f["values"])
        raw["housing_referrals"] = refs[name]["housing"]
        covered = {c for rule in _rules_for(name, year)
                   for c in rule["columns"]}
        rec = {KEY: name, PERIOD: year, "_raw": raw, "_flags": set()}
        for c in VALUES:
            sheet = ("Referral routes row " + str(refs[name]["row"])
                     if c == "housing_referrals" else f"Cases row {f['row']}")
            where = f"{wb['name']}: {sheet} {name} {HEADER_OF[c]}"
            if c in covered:
                rec[c] = None
                continue
            v, flag = cell(raw[c], where)
            rec[c] = None if v is None else to_scale(v, c, where)
            if flag:
                rec["_flags"].add(flag)
        if rec["_flags"]:
            out.markers[name] = sorted(rec["_flags"])
        notes = list(dict.fromkeys(f["footnotes"] + refs[name]["footnotes"]))
        if notes:
            out.footnotes[name] = notes
        out.append(rec)
    out.applied = zero_rules(out, year)
    for r in out:
        if len(r["_flags"]) > 1:
            raise ValueError(f"{wb['name']}: {r[KEY]} {year} has two reasons "
                             f"for its NULLs {sorted(r['_flags'])}; one "
                             "value_flag per row: review the row")
        r["value_flag"] = next(iter(r["_flags"]), None)
    out.sort(key=lambda r: r[KEY])
    return out


def reconcile(wb) -> tuple:
    """(problems, not_checked): the English police force areas of the Cases
    sheet must sum to the England row on RECONCILED (exact at 2 dp),
    wherever every part (and the England cell) is a published number; a
    measure with an unpublished part is listed in not_checked."""
    england = wb["cases"]["totals"][TOTALS[1]]
    forces = [f for reg in wb["cases"]["regions"] if reg["name"] != WALES
              for f in reg["forces"]]
    problems, not_checked = [], []
    q = Decimal("0.01")
    for c in RECONCILED:
        parts = [f["values"][c] for f in forces]
        if not all(_is_number(v) for v in parts) or \
                not _is_number(england[c]):
            not_checked.append(c)
            continue
        s = sum(Decimal(repr(float(v))) for v in parts)
        e = Decimal(repr(float(england[c])))
        if s.quantize(q) != e.quantize(q):
            problems.append(f"{c}: the {len(forces)} English police force "
                            f"areas sum to {s.quantize(q)}, the England row "
                            f"says {e.quantize(q)}")
    return problems, not_checked


def published_sums(wb, cols=RECONCILED) -> dict:
    """{column: (sum of the published English parts, England row)} (for the
    preview, where a measure is not checked)."""
    england = wb["cases"]["totals"][TOTALS[1]]
    forces = [f for reg in wb["cases"]["regions"] if reg["name"] != WALES
              for f in reg["forces"]]
    return {c: (sum(Decimal(repr(float(f["values"][c]))) for f in forces
                    if _is_number(f["values"][c])),
                england[c]) for c in cols}


# ---------------------------------------------------------------------------
# Ranks, hashes, ledger and planning
# ---------------------------------------------------------------------------

_RANK_RE = re.compile(r"\bSafeLives ([0-9]{4}-[0-9]{2}); modified "
                      r"([0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9:]{8}Z|unknown)")
_LEDGER_RE = re.compile(r"(.*) \(SafeLives ([0-9]{4}-[0-9]{2}); modified "
                        r"([0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9:]{8}Z|unknown)\)")


def rank_token(rank) -> str:
    """'SafeLives 2025-26; modified 2026-07-02T14:06:06Z' (or 'modified
    unknown': an .xls file has no document properties xlrd reads)."""
    y, mod = rank
    return f"SafeLives {y}; modified {_iso(mod) if mod else 'unknown'}"


def rank_text(rank) -> str:
    if rank is None:
        return "no file rank (as loaded)"
    return rank_token(rank)


def release_rank(x) -> "tuple | None":
    """(year, modified datetime or None) of a read workbook (its own rank),
    or the rank an edition's or ledger row's source_file names; None when it
    names none (a bare file name is never read for a rank)."""
    if isinstance(x, dict):
        return x.get("rank")
    mm = _RANK_RE.search(str(x or ""))
    if not mm:
        return None
    return mm.group(1), (None if mm.group(2) == "unknown"
                         else _when(mm.group(2)))


def edition_source(wb, src) -> str:
    """An edition's source_file: the file name, its rank, sha256 and where
    it was read."""
    return (f"{src['file_name']}; {rank_token(wb['rank'])}; sha256 "
            f"{wb['sha256'][:16]}; {src['where']}")


def ledger_source(where, rank) -> str:
    return f"{where} ({rank_token(rank)})"


def parse_ledger_source(s) -> "tuple | None":
    mm = _LEDGER_RE.fullmatch(str(s or ""))
    if not mm:
        return None
    return mm.group(1), (mm.group(2), None if mm.group(3) == "unknown"
                         else _when(mm.group(3)))


def rows_content_sha(records) -> str:
    """sha256 of a year's content: per record sorted by name,
    name|the eight values at the live scales|value_flag (NULL as ''),
    joined by LF, UTF-8."""
    lines = []
    for r in sorted(records, key=lambda x: x[KEY]):
        lines.append("|".join([r[KEY]] + ["" if r.get(c) is None
                                          else str(r[c]) for c in VALUES]
                              + [r.get("value_flag") or ""]))
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


def _older(a, b) -> bool:
    """True when rank a is known to be older than rank b (both modified
    dates known)."""
    return bool(a and b and a[1] and b[1] and a[1] < b[1])


def tip_ranks(tips, checked) -> dict:
    """{held year: the newest rank that stated it}: the tip edition's or a
    file the ledger records for the year; (year, None) when only unknown
    dates; None when nothing has a rank (as loaded)."""
    out = {}
    for p, t in tips.items():
        ranks = [t.get("rank")] + [release_rank(s) for s, _ in
                                   checked.get(p, ())]
        ranks = [r for r in ranks if r is not None]
        known = [r for r in ranks if r[1] is not None]
        out[p] = (max(known, key=lambda r: r[1]) if known
                  else (ranks[0] if ranks else None))
    return out


def ledger_complete(checked, where, sha, period) -> bool:
    for s, h in checked.get(period, ()):
        parsed = parse_ledger_source(s)
        if h == sha and parsed and parsed[0] == where:
            return True
    return False


def plan_periods(files, ranks, checked, *, recheck=None, allow_older=False,
                 only=None, stranded=()) -> tuple:
    """(new, revised, skipped {year: reason}) of the files read ({year:
    {where, sha, rank}}): a year not held is new; a held year whose tip
    comes from a newer file is skipped as older unless allow_older (the
    older-file guard, on every path); a file the ledger records is skipped
    as unchanged unless --recheck names the year or it is stranded; only
    (--year) limits the run. ValueError when --recheck names a year not read
    or not held, or --year a year not read."""
    if recheck is not None and (recheck not in files or recheck not in ranks):
        raise ValueError(f"--recheck {recheck}: not a held year read in this "
                         f"run ({', '.join(sorted(files)) or 'none'})")
    if only is not None and only not in files:
        raise ValueError(f"--year {only}: no file read for it "
                         f"({', '.join(sorted(files)) or 'none'})")
    new, revised, skipped = [], [], {}
    for y in sorted(files):
        f = files[y]
        if only is not None and y != only:
            skipped[y] = f"not requested (--year {only})"
            continue
        if y not in ranks:
            new.append(y)
            continue
        t = ranks[y]
        if _older(f["rank"], t) and not allow_older:
            skipped[y] = (f"older: its tip comes from a file of "
                          f"{rank_text(t)}, this file is "
                          f"{rank_text(f['rank'])}")
            continue
        if y != recheck and y not in stranded and \
                ledger_complete(checked, f["where"], f["sha"], y):
            skipped[y] = ("unchanged: this file is already in the ledger "
                          "(--recheck YEAR reads it again)")
            continue
        revised.append(y)
    return new, revised, skipped


# ---------------------------------------------------------------------------
# Stop conditions
# ---------------------------------------------------------------------------

def _by_key(records) -> dict:
    return {r[KEY]: r for r in records or ()}


def _cases_total(records):
    return sum((r["cases_discussed"] for r in records
                if r.get("cases_discussed") is not None), Decimal(0))


def _over(old, new, pct) -> bool:
    if old == new:
        return False
    if not old:
        return True
    return abs(new - old) * 100 > pct * old


def period_problems(new, tip, prev, *, kind, cols=COMPARED) -> tuple:
    """(hard, soft) of one year. New (prev: the latest earlier held year's
    records, or None): England cases discussed (the sum of the forces)
    moving by more than NEW_ENGLAND_PCT, or more than NEW_MAX_PFAS forces'
    cases discussed moving by more than NEW_PFA_PCT (soft). Revised (tip:
    the held edition's records): a force of the tip missing (hard: a partial
    file never replaces a fuller edition; nothing releases it); any change
    on cols, or a force new to the year (soft: no revision has been
    observed, so every change is read and acknowledged)."""
    g = globals()
    hard, soft = [], []
    if kind == "new":
        if not prev:
            return hard, soft
        a, b = _cases_total(prev), _cases_total(new)
        if _over(a, b, g["NEW_ENGLAND_PCT"]):
            soft.append(f"England cases discussed (the sum of the forces) "
                        f"{a} -> {b} against the previous held year (limit "
                        f"{g['NEW_ENGLAND_PCT']}%)")
        pk, nk = _by_key(prev), _by_key(new)
        moved = [f"{k} {pk[k]['cases_discussed']}->{nk[k]['cases_discussed']}"
                 for k in sorted(set(pk) & set(nk))
                 if pk[k]["cases_discussed"] is not None
                 and nk[k]["cases_discussed"] is not None
                 and _over(pk[k]["cases_discussed"],
                           nk[k]["cases_discussed"], g["NEW_PFA_PCT"])]
        if len(moved) > g["NEW_MAX_PFAS"]:
            soft.append(f"{len(moved)} police force areas move by more than "
                        f"{g['NEW_PFA_PCT']}% in cases discussed against the "
                        f"previous held year (limit {g['NEW_MAX_PFAS']}): "
                        + ", ".join(moved))
        return hard, soft
    nk, tk = _by_key(new), _by_key(tip)
    gone = sorted(set(tk) - set(nk))
    if gone:
        hard.append(f"PARTIAL FILE: {len(gone)} police force area(s) of the "
                    f"held edition missing from the file: {', '.join(gone)} "
                    "(a partial file never replaces a fuller edition)")
    changed = []
    for k in sorted(set(nk) & set(tk)):
        d = [f"{c} {_show(tk[k].get(c))}->{_show(nk[k].get(c))}"
             for c in cols if not _same(tk[k].get(c), nk[k].get(c))]
        if d:
            changed.append(f"{k} " + ", ".join(d))
    if changed:
        soft.append(f"{len(changed)} police force area(s) change against the "
                    "held edition: " + "; ".join(changed))
    added = sorted(set(nk) - set(tk))
    if added:
        soft.append(f"police force area(s) new to the year: "
                    f"{', '.join(added)} (refresh-latest then needs "
                    "--accept-key-changes for the year)")
    return hard, soft


def flips(new, tip) -> dict:
    """{'<force>/<column>': (tip value text or None, new value text or
    None)} of the values going to or from NULL against the tip (rule
    1.10), on forces in both."""
    nk, tk = _by_key(new), _by_key(tip)
    out = {}
    for k in sorted(set(nk) & set(tk)):
        for c in VALUES:
            a, b = tk[k].get(c), nk[k].get(c)
            if (a is None) != (b is None):
                out[f"{k}/{c}"] = (_text(a), _text(b))
    return out


def flip_problems(name, period, fl) -> list:
    """Problems (empty = released) of a year's flips under
    ACKNOWLEDGED_FLIPS[name]: exactly the listed cells, nothing else."""
    if not fl:
        return []
    msg = (f"{len(fl)} value(s) go to or from NULL against the held edition "
           "(rule 1.10): " + ", ".join(f"{k} {a or 'NULL'}->{b or 'NULL'}"
                                       for k, (a, b) in sorted(fl.items())))
    if not name:
        return [msg + "; released only by a named entry: --acknowledge-flips "
                "NAME (ACKNOWLEDGED_FLIPS in the loader)"]
    want = globals()["ACKNOWLEDGED_FLIPS"][name]["periods"].get(period, {})
    if dict(want) != fl:
        return [msg + f"; --acknowledge-flips {name} lists "
                f"{dict(want) or 'nothing'} for {period}, not exactly these"]
    return []


def restored(new, tip) -> list:
    """['<force> <column>'] of the cells the tip holds NULL and the file
    publishes as 0 (no rule): the restored zeros."""
    nk, tk = _by_key(new), _by_key(tip)
    return [f"{k} {c}" for k in sorted(set(nk) & set(tk)) for c in VALUES
            if tk[k].get(c) is None and nk[k].get(c) is not None
            and nk[k][c] == 0]


def _differs(new, old, cols) -> bool:
    nk, ok = _by_key(new), _by_key(old)
    if set(nk) != set(ok):
        return True
    return any(not _same(nk[k].get(c), ok[k].get(c)) for k in nk for c in cols)


# ---------------------------------------------------------------------------
# Spec and profile
# ---------------------------------------------------------------------------

def tip_row_count(editions_table: str):
    """expected_rows_per_period for status: the tip edition's row count."""
    core._ident(editions_table)

    def count(cur, period):
        cur.execute(f"SELECT DISTINCT edition, supersedes FROM "
                    f"public.{editions_table} WHERE financial_year = %s",
                    (period,))
        try:
            tip = core._chain_tip(cur.fetchall(), str(period))
        except (LookupError, ValueError):
            return None
        cur.execute(f"SELECT COUNT(*) FROM public.{editions_table} WHERE "
                    "financial_year = %s AND edition = %s", (period, tip))
        return cur.fetchone()[0]
    return count


_FLAG_LIST = ", ".join(f"'{f}'" for f in FLAGS)
_ANY_NULL = " OR ".join(f"{c} IS NULL" for c in VALUES)

SPEC = core.EditionSpec(
    name="s17",
    live_table="marac_cases",
    editions_table="marac_cases_editions",
    key_cols=(KEY,),
    period_col=PERIOD,
    value_cols=VALUE_TYPES,
    extra_cols=FLAG_TYPES,
    refresh_cols=VALUES + ("source", "loaded_at"),
    refresh_from=(("source", "release_label"), ("loaded_at", "loaded_at")),
    key_types=((KEY, "varchar(100) NOT NULL"),
               (PERIOD, "varchar(7) NOT NULL")),
    fk_la_boundaries=False,
    table_constraints=(
        f"CONSTRAINT marac_cases_editions_value_flag_chk CHECK (value_flag "
        f"IS NULL OR (value_flag IN ({_FLAG_LIST}) AND ({_ANY_NULL})))",),
    refresh_key_changes=True,
    expected_rows_per_period=tip_row_count("marac_cases_editions"),
)


def _no_label(fetched_on) -> str:
    raise ValueError("an S17 release label comes from the file; store "
                     "through apply_year(...)")


def check_records(records, period) -> None:
    """ValueError unless the year's records are whole: each of this year,
    one per force, value_flag a known flag or NULL, a flag only on a row
    with a NULL value."""
    seen = set()
    for r in records:
        if r.get(PERIOD) != period:
            raise ValueError(f"{period}: a record of {r.get(PERIOD)!r}")
        if r[KEY] in seen:
            raise ValueError(f"{period}: {r[KEY]} appears twice")
        seen.add(r[KEY])
        fl = r.get("value_flag")
        if fl is not None and (fl not in FLAGS or
                               all(r.get(c) is not None for c in VALUES)):
            raise ValueError(f"{period}: {r[KEY]} value_flag {fl!r} without "
                             "a NULL value, or unknown")


PROFILE = pe.Profile(
    spec=SPEC, value_cols=COMPARED, run_agent=RUN_AGENT,
    run_source=RUN_SOURCE, heading="S17 SafeLives Marac data (marac_cases)",
    default_source_file="SafeLives Marac data workbook",
    expected_areas=None, release_label=_no_label,
    content_sha256=rows_content_sha, check_records=check_records,
    file_checks=True, savepoint="s17_period",
    example_label="(eight values, value_flag)")


def profile() -> pe.Profile:
    """PROFILE bound to the current SPEC (looked up when called, so tests
    can swap the spec)."""
    return PROFILE.with_spec(globals()["SPEC"])


def refresh_spec(spec):
    """The spec refresh-latest uses: value_flag is not a live column, so the
    key-change insert must not name it (nothing compared changes)."""
    return dataclasses.replace(spec, extra_cols=())


def _late(name: str):
    return lambda *a, **kw: globals()[name](*a, **kw)


def create_all(cur) -> None:
    """ddl: the editions table, its append-only triggers and the file-check
    ledger. Idempotent; no value is written; the live table is not
    touched."""
    prof = profile()
    core.create_schema(cur, prof.spec)
    pe.create_file_checks(cur, prof)


def insert_live(cur, profile, period: str, records: list) -> None:
    """A new year's live rows: key, period, the eight values and source
    (each record's live_source); value_flag is not a live column; loaded_at
    takes its default."""
    from psycopg2.extras import execute_values
    spec = profile.spec
    cols = (KEY, PERIOD) + VALUES + ("source",)
    execute_values(
        cur, f"INSERT INTO public.{spec.live_table} ({', '.join(cols)}) "
        "VALUES %s",
        [(r[KEY], period) + tuple(r[c] for c in VALUES) + (r["live_source"],)
         for r in records], page_size=1000)


@dataclass(frozen=True)
class PeriodInfo:
    """What a year's edition records: source_file, the release label, the
    published date and the files read ((ledger source, sha256),)."""
    source_file: str
    label: str
    published: date
    files: tuple


def _tip_label(cur, spec, period) -> "str | None":
    try:
        tip = core.chain_tip(cur, spec, period)
    except (LookupError, ValueError):
        return None
    cur.execute(f"SELECT DISTINCT release_label FROM public.{spec.editions_table}"
                f" WHERE {spec.period_col} = %s AND edition = %s", (period, tip))
    r = cur.fetchone()
    return r[0] if r else None


def apply_year(cur, profile, period, records, *, fetched_on, info) -> str:
    """pe.apply_period for one year, then its ledger row, inside the
    caller's per-year savepoint (s17_period): a new year's edition 1, live
    rows and ledger row commit or roll back together. Each record's
    live_source is the edition's label; a stranded year gets the tip
    edition's."""
    kind = pe.classify_period(cur, profile, period, records)
    label = info.label
    if kind == LIVE_MISSING:
        label = _tip_label(cur, profile.spec, period) or info.label
    records = [dict(r, live_source=label) for r in records]
    prof = dataclasses.replace(profile, release_label=lambda d: info.label)
    kind = pe.apply_period(cur, prof, period, records,
                           fetched_on=info.published,
                           source_file=info.source_file,
                           classify=lambda *a, **kw: kind,
                           insert=_late("insert_live"))
    tip = core.chain_tip(cur, prof.spec, period)
    for src, sha in info.files:
        pe.record_file_check(cur, prof, period, src, sha, kind, tip)
    return kind


# ---------------------------------------------------------------------------
# Reading held data
# ---------------------------------------------------------------------------

table_exists = pe.table_exists


def _conn(writing):
    return pe.connect(writing)


def pfa_names(cur) -> frozenset:
    """The English police force area names: la_pfa_mapping's distinct
    pfa_name_safelives (read, never written)."""
    cur.execute("SELECT DISTINCT pfa_name_safelives FROM public.la_pfa_mapping")
    return frozenset(r[0] for r in cur.fetchall())


def mapping_state(cur) -> tuple:
    """(rows, survey hash) of la_pfa_mapping (LEGACY_PFA_MAPPING's form)."""
    cols = ("lad24cd", "la_name", "pfa_name", "pfa_name_safelives",
            "population_weight", "source")
    expr = "||'|'||".join(f"coalesce({c}::text,'~')" for c in cols)
    cur.execute(f"SELECT COUNT(*), md5(string_agg({expr}, E'\\n' ORDER BY "
                "lad24cd)) FROM public.la_pfa_mapping")
    return tuple(cur.fetchone())


def live_state(cur, spec=None) -> tuple:
    """(rows, survey hash, distinct loaded_at, ((year, rows), ...)) of the
    live table (LEGACY_LIVE's form)."""
    spec = spec or globals()["SPEC"]
    cols = (KEY, PERIOD) + VALUES + ("source",)
    expr = "||'|'||".join(f"coalesce({c}::text,'~')" for c in cols)
    t = spec.live_table
    cur.execute(f"SELECT COUNT(*), md5(string_agg({expr}, E'\\n' ORDER BY "
                f"{PERIOD}, {KEY})), COUNT(DISTINCT loaded_at) FROM "
                f"public.{t}")
    n, h, nl = cur.fetchone()
    cur.execute(f"SELECT {PERIOD}, COUNT(*) FROM public.{t} GROUP BY 1 "
                "ORDER BY 1")
    return (n, h, nl, tuple((pe._p(p), c) for p, c in cur.fetchall()))


def live_type_problems(cur, spec=None) -> list:
    """Problems (empty = fine): the live value columns' types from
    information_schema against VALUE_TYPES."""
    spec = spec or globals()["SPEC"]
    cur.execute("SELECT column_name, data_type, numeric_precision, "
                "numeric_scale FROM information_schema.columns WHERE "
                "table_schema = 'public' AND table_name = %s",
                (spec.live_table,))
    got = {r[0]: r[1:] for r in cur.fetchall()}
    bad = []
    for c, t in VALUE_TYPES:
        g = got.get(c)
        if g is None:
            bad.append(f"{spec.live_table}.{c} missing")
            continue
        if t == "integer":
            ok = g[0] == "integer"
        else:
            p, s = map(int, re.fullmatch(r"numeric\((\d+),(\d+)\)", t).groups())
            ok = g[0] == "numeric" and (g[1], g[2]) == (p, s)
        if not ok:
            bad.append(f"{spec.live_table}.{c} is {g}, expected {t}")
    return bad


def held_records(cur, spec, period, edition=None) -> list:
    """The year's rows as records: the given edition's (the eight values,
    value_flag, release_label as live_source), or the live table's (edition
    None: source as live_source, value_flag None)."""
    if edition is None:
        names = (KEY,) + VALUES + ("source",)
        cur.execute(f"SELECT {', '.join(names)} FROM public.{spec.live_table} "
                    f"WHERE {PERIOD} = %s", (period,))
        out = []
        for row in cur.fetchall():
            r = dict(zip(names, row))
            r["live_source"] = r.pop("source")
            r["value_flag"] = None
            out.append(r)
    else:
        names = (KEY,) + VALUES + ("value_flag", "release_label")
        cur.execute(f"SELECT {', '.join(names)} FROM "
                    f"public.{spec.editions_table} WHERE {PERIOD} = %s AND "
                    "edition = %s", (period, edition))
        out = []
        for row in cur.fetchall():
            r = dict(zip(names, row))
            r["live_source"] = r.pop("release_label")
            out.append(r)
    for r in out:
        r[PERIOD] = str(period)
    return sorted(out, key=lambda r: r[KEY])


def tip_info(cur, spec) -> dict:
    """{year: {edition, source_file, label, rank}} of every year with
    editions (rank from the tip's source_file; None for 'as loaded')."""
    cur.execute(f"SELECT DISTINCT {PERIOD} FROM public.{spec.editions_table}")
    out = {}
    for (p,) in cur.fetchall():
        tip = core.latest_edition(cur, spec, str(p))
        cur.execute(f"SELECT DISTINCT source_file, release_label FROM "
                    f"public.{spec.editions_table} WHERE {PERIOD} = %s AND "
                    "edition = %s", (p, tip))
        rows = cur.fetchall()
        sf, lb = rows[0] if len(rows) == 1 else (None, None)
        out[pe._p(p)] = {"edition": tip, "source_file": sf, "label": lb,
                         "rank": release_rank(sf)}
    return out


def legacy_tips(cur, spec) -> dict:
    """Before the editions table exists: {held live year: {rank None}}."""
    cur.execute(f"SELECT DISTINCT {PERIOD} FROM public.{spec.live_table}")
    return {pe._p(p): {"edition": None, "source_file": None, "label": None,
                       "rank": None} for (p,) in cur.fetchall()}


def ledger_checks(cur, prof) -> dict:
    """{year: {(source_file, file_sha256)}} of the ledger."""
    cur.execute(f"SELECT DISTINCT {PERIOD}, source_file, file_sha256 FROM "
                f"public.{pe.file_checks_table(prof)}")
    out = {}
    for p, f, s in cur.fetchall():
        out.setdefault(pe._p(p), set()).add((f, s))
    return out


# ---------------------------------------------------------------------------
# ddl, status, run log, refresh-latest
# ---------------------------------------------------------------------------

def log_run(cur, rows_written: int, notes: str, started_at=None) -> None:
    """The pipeline_run_log row for a committed run (agent RUN_AGENT, source
    '17'). Called only on committed load, migrate-legacy and restore-edition
    runs, partial runs included."""
    pe.log_run(cur, profile(), rows_written, notes, started_at)


def _missing_tables(cur) -> list:
    prof = profile()
    return [t for t in (prof.spec.editions_table, pe.file_checks_table(prof))
            if not table_exists(cur, t)]


def cmd_ddl(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    prof = profile()
    try:
        with conn.cursor() as cur:
            bad = live_type_problems(cur)
            print("live types (information_schema): "
                  + ("as expected" if not bad else "; ".join(bad)))
            if bad:
                halt("the live columns are not as surveyed; nothing created")
            for t in (prof.spec.editions_table, pe.file_checks_table(prof)):
                print(f"{t}: {'exists' if table_exists(cur, t) else 'does not exist'}")
            if not writing:
                print("DRY RUN: nothing written (use --commit or --simulate)")
                return 0
            create_all(cur)
        pe.finish(conn, args, f"ddl: {prof.spec.editions_table}, its triggers "
                              "and ledger present")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def status_lines(cur) -> list:
    prof = profile()
    s = prof.spec
    lt = pe.file_checks_table(prof)
    lines = []
    for p, t in sorted(tip_info(cur, s).items()):
        line = f"  {p}: tip edition {t['edition']} ({rank_text(t['rank'])})"
        cur.execute(f"SELECT source_file, outcome FROM public.{lt} WHERE "
                    f"{PERIOD} = %s ORDER BY id DESC LIMIT 1", (p,))
        last = cur.fetchone()
        line += (f"; ledger latest: {last[1]} {last[0]}" if last
                 else "; no ledger row")
        lines.append(line)
    return lines


def cmd_status(args) -> int:
    """status: what needs action; exit 1 if anything, or (with a clean
    message, nothing created) if the editions table or ledger does not
    exist."""
    prof = profile()
    conn = _conn(False)
    try:
        with conn.cursor() as cur:
            missing = _missing_tables(cur)
            if missing:
                print(f"{prof.heading}: {' and '.join(missing)} not present "
                      "yet; run `ddl --commit`, then `migrate-legacy "
                      "--commit` (S17 has no sync-new)")
                return 1
            st = pe.status(cur, prof)
            lines = status_lines(cur)
    finally:
        conn.rollback()
        conn.close()
    print(pe.format_status(prof, st, lines))
    return 0 if st["ok"] else 1


def live_equals_tip(cur, spec, periods) -> list:
    """Problems (empty = fine): for each year, the live rows equal the tip
    edition on the key, the year and the eight values, NULL-safe, both
    ways."""
    cols = ", ".join((KEY, PERIOD) + VALUES)
    bad = []
    for p in periods:
        try:
            ed = core.chain_tip(cur, spec, p)
        except (LookupError, ValueError) as e:
            bad.append(f"{pe._p(p)}: {e}")
            continue
        a = (f"SELECT {cols} FROM public.{spec.editions_table} "
             f"WHERE {PERIOD} = %s AND edition = %s")
        b = f"SELECT {cols} FROM public.{spec.live_table} WHERE {PERIOD} = %s"
        counts = []
        for x, xa, y, ya in ((a, (p, ed), b, (p,)), (b, (p,), a, (p, ed))):
            cur.execute(f"SELECT COUNT(*) FROM (({x}) EXCEPT ALL ({y})) q",
                        xa + ya)
            counts.append(cur.fetchone()[0])
        if any(counts):
            bad.append(f"{pe._p(p)} ed{ed}: {counts[0]} edition-only, "
                       f"{counts[1]} live-only rows")
    return bad


def cmd_refresh_latest(args) -> int:
    """refresh-latest [--accept-drift P] [--accept-key-changes P] [--commit
    | --simulate]: the engine's refresh (each year's latest edition copied
    into live: the eight values, source from the edition's release label and
    loaded_at, under the before/after hash guard); forces added or removed
    only for a year named in --accept-key-changes; then live_equals_tip."""
    writing = args.commit or args.simulate
    spec = profile().spec
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            missing = _missing_tables(cur)
            if missing:
                halt(f"{', '.join(missing)} not present yet; run `ddl "
                     "--commit`, then `migrate-legacy --commit`")
            accept = pe._accept_periods(cur, spec,
                                        tuple(args.accept_drift or ()))
            accept_kc = pe._accept_key_change_periods(
                cur, spec, pe._accept_key_change_args(args))
            counts = core.refresh_counts(cur, spec, accept)
            print("rows refresh-latest would write: "
                  + (", ".join(f"{pe._p(p)}={n}" for p, n in counts.items())
                     or "none") + f" (total {sum(counts.values())})")
            plan, _ = core._plan(cur, spec, accept)
            lines = core.key_change_lines(plan, accept_kc, name=pe._p)
            print("forces added and removed:" + ("" if lines else " none"))
            for line in lines:
                print(f"  {line}")
            if not writing:
                print("DRY RUN: nothing written (use --commit or --simulate)")
                return 0
            pe._refresh_guard(cur, spec, accept, accept_kc)
            res = core.refresh_latest(cur, refresh_spec(spec), accept,
                                      accept_kc)
            if res["updated"]:
                bad = live_equals_tip(cur, spec, sorted(res["updated"]))
                if bad:
                    halt("refresh-latest: live differs from the latest edition "
                         "after the refresh, rolled back: " + "; ".join(bad[:6]))
        keys = ""
        if res["inserted"] or res["deleted"]:
            keys = (f"; forces {sum(res['inserted'].values())} inserted, "
                    f"{sum(res['deleted'].values())} deleted in "
                    f"{sorted(pe._p(x) for x in res['inserted'])}")
        pe.finish(conn, args, f"{res['rows']} live rows refreshed in "
                              f"{sorted(pe._p(x) for x in res['updated'])}"
                              f"{keys}; before/after guard passed")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# load
# ---------------------------------------------------------------------------

@dataclass
class Plan:
    """A prepared run, before the acknowledgements are applied."""
    prof: object
    has_ed: bool
    files: dict = dataclasses.field(default_factory=dict)
    by_period: dict = dataclasses.field(default_factory=dict)
    held_recs: dict = dataclasses.field(default_factory=dict)
    new: list = dataclasses.field(default_factory=list)
    revised: list = dataclasses.field(default_factory=list)
    skipped: dict = dataclasses.field(default_factory=dict)
    hard: dict = dataclasses.field(default_factory=dict)
    soft: dict = dataclasses.field(default_factory=dict)
    flips: dict = dataclasses.field(default_factory=dict)
    reissue: dict = dataclasses.field(default_factory=dict)
    older: dict = dataclasses.field(default_factory=dict)
    run_notes: list = dataclasses.field(default_factory=list)
    notes: dict = dataclasses.field(default_factory=dict)
    rank_halt: "str | None" = None

    @property
    def periods(self) -> list:
        return sorted(self.new + self.revised)


def _end(conn, args, writing, rc) -> int:
    conn.rollback()  # commit mode has committed already
    if not writing:
        print("PREVIEW: nothing written to the database (use --commit or "
              "--simulate)")
    elif args.simulate:
        print("SIMULATION: ROLLED BACK (nothing persisted)")
    return rc


def _year_arg(text) -> str:
    try:
        return label_ok(text)
    except ValueError as e:
        raise argparse.ArgumentTypeError(str(e)) from None


def _read(path, year=None, why="") -> dict:
    try:
        return read_workbook(path, year)
    except ValueError as e:
        halt(f"file identity check failed{why}: {e}; nothing stored")


def _links(held) -> tuple:
    html, final = fetch_page(REGISTRY_URL)
    try:
        links = data_page_links(html, held)
    except ValueError as e:
        halt(f"the data page {final}: {e}. If the page has changed, download "
             "the year's workbook by hand into data/raw/s17_marac/ and rerun "
             "with --file PATH --no-page")
    print(f"data page: {final} (reached from {REGISTRY_URL}); workbook links "
          f"by year: " + ", ".join(f"{y} {u}" for y, u in sorted(links.items())))
    for u in links.ignored:
        print(f"  workbook link naming no year (not used): {u}")
    return links, final


def _sources(args, held) -> dict:
    """{year: {wb, where, how, file_name, sha, rank, url}} of the files this
    run reads: --file files (under FILE_ROOTS, announced as manual input;
    unless --no-page the data page is read and must link each file's year),
    or every year the data page links (or --year only), downloaded to
    RAW_DIR with the plain User-Agent."""
    out = {}
    if args.file:
        links = None
        if not args.no_page:
            links, final = _links(held)
        for f in args.file:
            p = manual_input.require_under(f, globals()["FILE_ROOTS"])
            ident = manual_input.file_identity(p)
            where = _file_source(p)
            manual_input.announce_manual("S17 SafeLives Marac data "
                                         "(hand-placed workbook)", where,
                                         ident)
            wb = _read(p, args.year)
            y = wb["year"]
            if y in out:
                halt(f"--file: two files for {y}")
            if links is not None and y not in links:
                halt(f"--file {p.name}: the data page {final} lists no link "
                     f"for {y} (the file's own titles); give --no-page to "
                     "load the file without it (recorded as such)")
            how = "hand-placed file" + (f"; {NO_PAGE}" if args.no_page else
                                        f"; its year linked on {final}")
            out[y] = {"wb": wb, "where": where, "how": how,
                      "file_name": p.name, "sha": wb["sha256"],
                      "rank": wb["rank"],
                      "url": links.get(y) if links else None}
        return out
    links, final = _links(held)
    years = sorted(links)
    if args.year is not None:
        if args.year not in links:
            halt(f"--year {args.year}: the data page lists no link for it")
        years = [args.year]
    print(f"downloading to {globals()['RAW_DIR']} (also in a preview; the "
          "database is not written)")
    for y in years:
        path, furl = fetch(links[y], globals()["RAW_DIR"])
        wb = _read(path, y, f" (the data page links {links[y]} as {y})")
        print(f"  {y}: {links[y]}" + (f" -> {furl}" if furl != links[y]
                                      else "") + f"; sha256 "
              f"{wb['sha256'][:16]}")
        out[y] = {"wb": wb, "where": furl, "how": "data page link (plain "
                  "request)", "file_name": Path(path).name,
                  "sha": wb["sha256"], "rank": wb["rank"], "url": links[y]}
    return out


def _report(f) -> None:
    wb = f["wb"]
    print(f"{f['file_name']}: {wb['notes_title']!r}; {wb['cases_title']!r}; "
          f"{rank_text(wb['rank'])}; sha256 {wb['sha256'][:16]}")
    if wb["erratum"]:
        print(f"  NOTE (named erratum NOTES_TITLE_ERRATA): {wb['erratum']['evidence']}")
    for h in wb["header_aliases"]:
        a = HEADER_ALIASES[h]
        print(f"  header alias: {h!r} read as {a['means']!r} (seen: {a['seen']})")


def _held_records(cur, spec, tips, has_ed) -> dict:
    return {p: held_records(cur, spec, p, tips[p]["edition"] if has_ed
                            else None) for p in tips}


def _prepare(args, cur, writing) -> Plan:
    prof = profile()
    spec = prof.spec
    missing = _missing_tables(cur)
    has_ed = not missing
    if writing and not has_ed:
        halt(f"missing {', '.join(missing)}; run `ddl --commit` and "
             "`migrate-legacy --commit` first")
    if has_ed:
        _, new_live, errors = core.latest_map(cur, spec)
        if errors:
            halt("invalid edition chain: " + "; ".join(
                f"{pe._p(p)}: {msg}" for p, msg in errors.items()))
        if new_live:
            halt(f"{spec.live_table}: live years with no editions "
                 f"{[pe._p(x) for x in new_live]}; run migrate-legacy first")
        tips = tip_info(cur, spec)
        checked = ledger_checks(cur, prof)
        stranded = pe.live_missing_periods(cur, prof)
    else:
        print(f"NOTE: {', '.join(missing)} not present: this preview compares "
              f"with the LIVE table {spec.live_table} (the eight values)")
        tips, checked, stranded = legacy_tips(cur, spec), {}, []
    plan = Plan(prof, has_ed)
    ranks = tip_ranks(tips, checked)
    held = sorted(ranks)
    print(f"held years: {', '.join(held) or 'none'}")
    names = pfa_names(cur)
    files = _sources(args, held)
    plan.files = files
    cols = COMPARED if has_ed else VALUES
    for y in sorted(files):
        f = files[y]
        _report(f)
        problems, not_checked = reconcile(f["wb"])
        if problems:
            halt(f"{y}: reconciliation inside the file failed, nothing "
                 "stored: " + "; ".join(problems))
        sums = published_sums(f["wb"], not_checked)
        done = [c for c in RECONCILED if c not in not_checked]
        print(f"  {y}: "
              + (f"reconciled: the English police force areas sum to the "
                 f"England row on {', '.join(done)}" if done else
                 "reconciliation")
              + (("; " if done else " ") + "not checked (a part is not "
                 "published) for " + ", ".join(
                     f"{c} (the published parts sum to {sums[c][0]}, England "
                     f"row {sums[c][1]})" for c in not_checked)
                 if not_checked else ""))
        try:
            recs = records(f["wb"], y, names)
        except ValueError as e:
            halt(f"{e}; nothing stored")
        plan.by_period[y] = recs
        for who, fl in sorted(recs.markers.items()):
            print(f"  {y}: {who}: marker(s) -> NULL {', '.join(fl)}")
        for a in recs.applied:
            print(f"  {y}: RULE {a['rule']}: {a['pfa']} {a['column']} "
                  f"published {a['published']!r} -> NULL {a['flag']}")
        for who, notes in sorted(recs.footnotes.items()):
            for n in notes:
                print(f"  {y}: footnote on {who}: {n}")
    for rule in ZERO_RULES:
        if any(a["rule"] == rule["name"] for r in plan.by_period.values()
               for a in r.applied):
            print(f"  rule {rule['name']} ({rule['decided']}): "
                  f"{rule['evidence']}")
    info = {y: {"where": f["where"], "sha": f["sha"], "rank": f["rank"]}
            for y, f in files.items()}
    chk = {p: s for p, s in checked.items() if p not in stranded}
    try:
        new, revised, skipped = plan_periods(
            info, ranks, chk, recheck=args.recheck,
            allow_older=args.allow_older_file, only=args.year,
            stranded=stranded)
    except ValueError as e:
        halt(str(e))
    plan.new, plan.revised, plan.skipped = new, revised, skipped
    for y, reason in sorted(skipped.items()):
        print(f"  {y}: {reason}")
    if not plan.periods and skipped and all(
            r.startswith("older") for r in skipped.values()):
        plan.rank_halt = (
            "older file: every year read is skipped; the files are older "
            "than what is held for " + ", ".join(sorted(skipped))
            + "; storing them would record an older file. If this is "
            "deliberate, re-run with --allow-older-file")
    held_recs = _held_records(cur, spec, tips, has_ed)
    plan.held_recs = held_recs
    newest = max(held + plan.periods, default=None)
    for p in plan.periods:
        recs = plan.by_period[p]
        kind = "new" if p in new else "revised"
        tip = held_recs.get(p) if kind == "revised" else None
        earlier = [h for h in held if h < p]
        prev = held_recs.get(max(earlier)) if (kind == "new" and earlier) \
            else None
        hard, soft = period_problems(recs, tip, prev, kind=kind, cols=cols)
        plan.hard[p], plan.soft[p] = hard, soft
        if kind == "revised":
            plan.flips[p] = flips(recs, tip)
            back = restored(recs, tip)
            if back:
                print(f"  {p}: restored (a published 0 held as NULL, no rule "
                      f"names it): {', '.join(back)}")
            t = ranks.get(p)
            if t is not None and _differs(recs, tip, cols):
                plan.reissue[p] = (f"the tip comes from a file of "
                                   f"{rank_text(t)}; this file "
                                   f"({rank_text(files[p]['rank'])}) differs "
                                   "in content: a reissue for a held year")
            if _older(files[p]["rank"], t):
                plan.older[p] = (f"--allow-older-file given; {p}'s tip comes "
                                 f"from a file of {rank_text(t)}, this file "
                                 f"is {rank_text(files[p]['rank'])}")
                print(f"NOTE: {plan.older[p]}")
        if p == newest:
            _map_line(p, recs, tip)
    return plan


def _map_line(p, recs, tip) -> None:
    """MAP: what W1 reads (cases_discussed and cases_per_10k_adult_females
    at the newest year) changes if this year is stored and refreshed."""
    if tip is None:
        print(f"  MAP: {p} would be the newest year (W1 reads its "
              f"{', '.join(MAP_COLUMNS)} via la_pfa_mapping)")
        return
    nk, tk = _by_key(recs), _by_key(tip)
    changed = [k for k in sorted(set(nk) & set(tk))
               if any(not _same(nk[k][c], tk[k][c]) for c in MAP_COLUMNS)]
    added = sorted(set(nk) - set(tk))
    print(f"  MAP: {p} is the newest year: {', '.join(MAP_COLUMNS)} change "
          f"for {', '.join(changed) or 'no force'}; forces gaining a row: "
          f"{', '.join(added) or 'none'} (after refresh-latest)")


def _info(p, f, notes) -> PeriodInfo:
    wb = f["wb"]
    extra = "".join(f"; {n}" for n in notes if n)
    label = (f"SafeLives Marac data {p}: {f['file_name']}; "
             f"{rank_token(wb['rank'])}; {f['where']} sha256 "
             f"{wb['sha256'][:16]}; {f['how']}" + extra)
    published = wb["modified"].date() if wb["modified"] else date.today()
    return PeriodInfo(edition_source(wb, f) + extra, label, published,
                      ((ledger_source(f["where"], wb["rank"]), wb["sha256"]),))


def _finalize(plan, args, ack, reissue) -> tuple:
    """(problems {year: [msg]}, infos {year: PeriodInfo}, notes {year: [..]})
    with the acknowledgements applied."""
    problems, infos, notes = {}, {}, {}
    flips_name = getattr(args, "acknowledge_flips", None)
    for p in plan.periods:
        bad = list(plan.hard.get(p, []))
        soft = plan.soft.get(p, [])
        n = [plan.older.get(p, "")]
        if soft and p in ack:
            n.append("ACKNOWLEDGED " + "; ".join(soft))
            print(f"  {p}: ACKNOWLEDGED (--acknowledge): {'; '.join(soft)}")
        elif soft:
            bad += [x + f"; read the preview, then --acknowledge {p}"
                    for x in soft]
        fl = flip_problems(flips_name, p, plan.flips.get(p, {}))
        bad += fl
        if plan.flips.get(p) and not fl:
            a = ACKNOWLEDGED_FLIPS[flips_name]
            n.append(f"ACKNOWLEDGED FLIPS {flips_name} ({a['decided']})")
            print(f"  {p}: ACKNOWLEDGED (--acknowledge-flips {flips_name}): "
                  f"{a['why']}")
        if p in plan.reissue:
            if p in reissue:
                msg = f"--accept-reissue {p}: {plan.reissue[p]}"
                n.append(msg)
                print(f"  {p}: ACCEPTED REISSUE: {msg}")
            else:
                bad.append(plan.reissue[p] + "; read what changed, then "
                           f"--accept-reissue {p}")
        rules = sorted({a["rule"] for a in plan.by_period[p].applied})
        if rules:
            n.append("rules " + ", ".join(rules))
        if bad:
            problems[p] = bad
            continue
        notes[p] = [x for x in n if x]
        infos[p] = _info(p, plan.files[p], notes[p])
    return problems, infos, notes


def _execute(plan, args, conn, cur, problems, infos, writing) -> int:
    started = datetime.now(timezone.utc)
    prof = plan.prof
    for p, msgs in sorted(problems.items()):
        for msg in msgs:
            print(f"  {p}: REJECTED, not stored: STOP CONDITION: {msg}")
    ok = [p for p in plan.periods if p not in problems]
    rc = 0
    stats = {"periods": [], "kinds": {}, "stored_rows": 0, "live_rows": 0}
    if ok and plan.has_ed:
        def apply(c, prof_, per, rs, *, fetched_on, source_file=None):
            return _late("apply_year")(c, prof_, per, rs,
                                       fetched_on=fetched_on, info=infos[per])
        rc = pe.load_periods(cur, prof, ok, lambda per: plan.by_period[per],
                             date.today(), args.commit,
                             simulate=args.simulate, against="editions",
                             stats=stats, apply=apply)
    elif ok:
        live_prof = dataclasses.replace(prof, value_cols=VALUES,
                                        example_label="(eight values)")
        rc = pe.load_periods(cur, live_prof, ok,
                             lambda per: plan.by_period[per], date.today(),
                             False, against="live")
        print("  (no editions table yet: preview against live only, on the "
              "eight values)")
    if problems:
        print(f"REJECTED {len(problems)} year(s), nothing stored for them: "
              f"{', '.join(sorted(problems))}; exit 1")
        rc = 1
    if args.commit and (stats["periods"] or problems or rc):
        stored = live = 0  # rows written by this run; not a source value
        for p in stats["periods"]:
            k = stats["kinds"][p]
            if k in ("new", "revised"):
                stored += len(plan.by_period[p])
            if k in ("new", LIVE_MISSING):
                live += len(plan.by_period[p])
        left = [p for p in ok if p not in stats["periods"]]
        failed, not_tried = (left[:1], left[1:]) if left else ([], [])
        partial = ""
        if rc:
            partial = (f"PARTIAL RUN (exit 1): rejected "
                       f"{sorted(problems) or 'none'}"
                       + "".join(f" ({p}: {'; '.join(problems[p])[:300]})"
                                 for p in sorted(problems))
                       + f"; failed {failed or 'none'}; not attempted "
                       f"{not_tried or 'none'}; a year that failed or was "
                       "rejected stored nothing and is not in the counts "
                       "below; ")
        how = sorted({plan.files[p]["how"] for p in plan.files})
        notes = (partial + "SafeLives Marac workbooks ("
                 + "; ".join(how) + "): "
                 + ("; ".join(f"{p} {k}" for p, k in stats["kinds"].items())
                    or "nothing stored")
                 + (f"; skipped {sorted(plan.skipped)}" if plan.skipped else "")
                 + "".join(f"; {r}" for r in plan.run_notes)
                 + "".join(f"; {p}: {'; '.join(plan.notes.get(p, []))}"
                           for p in stats["periods"] if plan.notes.get(p))
                 + f". Rows stored: {stored} edition rows, {live} live rows.")
        log_run(cur, stored + live, notes, started)
        conn.commit()
        print("pipeline_run_log row written"
              + (" (partial run: see its notes)" if rc else ""))
    return rc


def cmd_load(args) -> int:
    if args.no_page and not args.file:
        halt("--no-page is only for --file (a hand-placed workbook read "
             "without the data page)")
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            plan = _prepare(args, cur, writing)
            if plan.rank_halt:
                halt(plan.rank_halt)
            ack = set(args.acknowledge or ())
            reissue = set(args.accept_reissue or ())
            stray = sorted((ack | reissue) - set(plan.periods))
            if stray:
                halt(f"--acknowledge/--accept-reissue {', '.join(stray)}: not "
                     f"years this run compares or stores "
                     f"({', '.join(plan.periods) or 'none'}); nothing stored")
            if args.acknowledge_flips and not any(plan.flips.values()):
                halt(f"--acknowledge-flips {args.acknowledge_flips}: no value "
                     "goes to or from NULL in this run; nothing stored")
            print(f"planned: new {plan.new or 'none'}; compare "
                  f"{plan.revised or 'none'}; skipped "
                  f"{sorted(plan.skipped) or 'none'}")
            if not plan.periods:
                print("nothing to do")
                return _end(conn, args, writing, rc=0)
            problems, infos, notes = _finalize(plan, args, ack, reissue)
            plan.notes = notes
            rc = _execute(plan, args, conn, cur, problems, infos, writing)
        return _end(conn, args, writing, rc)
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# One-off: migrate the held table
# ---------------------------------------------------------------------------

def _derived_flag(held, rec) -> "str | None":
    """Edition 1's value_flag for a held row: the file's reason when every
    NULL the row holds is a NULL of the file's record (a marker or a named
    rule); None otherwise (a NULL the file publishes as a number)."""
    nulls = [c for c in VALUES if held[c] is None]
    if rec is None or not nulls:
        return None
    if all(rec[c] is None for c in nulls):
        return rec["value_flag"]
    return None


def migrate_legacy(cur, *, write, legacy=None, files=None, mapping=None,
                   explained=None) -> dict:
    """One-off. Preconditions: on write the editions table and ledger exist
    and are empty (in a preview they may be absent; present, they must be
    empty); the live types as VALUE_TYPES; live as surveyed (legacy,
    default LEGACY_LIVE); la_pfa_mapping as surveyed (mapping, default
    LEGACY_PFA_MAPPING); each held file's sha256 as listed (files, default
    LEGACY_FILES).

    Proof per year (any failure halts), the held file read by this loader
    (identity, reconciliation, records with the named rules) against live:
    every held value is reproduced at the column scale, or listed in one of
    these groups: (a) restored: held NULL, published 0, no rule names it
    (the n8n parseFloat(x) || null); (b) a named rule's cell (held value or
    NULL, rule NULL); keys only in the file that `explained` (default
    LEGACY_EXPLAINED_KEYS) explains. Anything else (c) halts.

    Edition 1 'as loaded' of every year is live as held, value_flag the
    file's reason where every NULL of the row is a NULL of the file
    (_derived_flag). A year with nothing in any group gets a ledger row for
    its file (outcome 'unchanged', edition 1). Live untouched. Nothing
    commits here."""
    g = globals()
    legacy = tuple(legacy or g["LEGACY_LIVE"])
    files = files or g["LEGACY_FILES"]
    mapping = tuple(mapping or g["LEGACY_PFA_MAPPING"])
    explained = g["LEGACY_EXPLAINED_KEYS"] if explained is None else explained
    prof = profile()
    spec = prof.spec
    lt = pe.file_checks_table(prof)
    ready = {spec.editions_table: table_exists(cur, spec.editions_table),
             lt: table_exists(cur, lt)}
    if write and not all(ready.values()):
        halt(f"missing {[t for t, ok in ready.items() if not ok]}; run "
             "`ddl --commit` first")
    for t, ok in ready.items():
        if ok:
            cur.execute(f"SELECT COUNT(*) FROM public.{t}")
            n = cur.fetchone()[0]
            if n:
                halt(f"{t} already holds {n} rows; migrate-legacy runs once, "
                     "on empty tables")
    bad = live_type_problems(cur)
    if bad:
        halt("the live columns are not as surveyed: " + "; ".join(bad))
    state = tuple(live_state(cur, spec))
    if state != legacy:
        halt(f"{spec.live_table} is not as surveyed: (rows, hash, loaded_at "
             f"values, years) {state}, expected {legacy}; nothing stored")
    ms = mapping_state(cur)
    if ms != mapping:
        halt(f"la_pfa_mapping is not as surveyed: {ms}, expected {mapping}; "
             "nothing stored")
    names = pfa_names(cur)
    years = [y for y, _ in legacy[3]]
    if sorted(files) != years:
        halt(f"held files for {sorted(files)}, live years {years}; nothing "
             "stored")
    plan = {"years": {}, "rows": 0}
    told = set()
    print("proof: each held file read by this loader against live; groups: "
          "(a) restored (held NULL, published 0), (b) a named rule's cell, "
          "keys only in the file (explained); anything else halts")
    for y in years:
        f = files[y]
        path = Path(f["path"])
        sha = content_sha256(path)
        if sha != f["sha256"]:
            halt(f"{path.name}: sha256 {sha[:16]}, expected "
                 f"{f['sha256'][:16]} (the held file); nothing stored")
        try:
            wb = read_workbook(path, y)
            problems, not_checked = reconcile(wb)
            if problems:
                raise ValueError("reconciliation: " + "; ".join(problems))
            recs = records(wb, y, names)
        except ValueError as e:
            halt(f"proof failed: {path.name}: {e}; nothing stored")
        held = {r[KEY]: r for r in held_records(cur, spec, y)}
        got = {r[KEY]: r for r in recs}
        rule_cells = {(a["pfa"], a["column"]): a for a in recs.applied}
        group_a, group_b, keep_null, other, extra = [], [], [], [], []
        cells = 0  # cells compared; not a source value
        for k in sorted(set(held) - set(got)):
            other.append(f"{k}: held, not in the file")
        for k in sorted(set(got) - set(held)):
            if k in explained:
                extra.append(k)
            else:
                other.append(f"{k}: in the file, not held (unexplained)")
        for k in sorted(set(held) & set(got)):
            for c in VALUES:
                cells += 1
                h, r = held[k][c], got[k][c]
                if _same(h, r):
                    if h is None and (k, c) in rule_cells:
                        keep_null.append(f"{k} {c} (published "
                                         f"{rule_cells[(k, c)]['published']!r};"
                                         f" {rule_cells[(k, c)]['rule']})")
                    continue
                if h is None and r is not None and r == 0:
                    group_a.append(f"{k} {c}")
                elif r is None and (k, c) in rule_cells:
                    group_b.append(f"{k} {c} held {_show(h)} -> NULL "
                                   f"({rule_cells[(k, c)]['rule']})")
                else:
                    other.append(f"{k} {c}: held {_show(h)}, file {_show(r)}")
        if other:
            halt(f"proof failed: {y}: {len(other)} difference(s) outside the "
                 f"groups: " + "; ".join(other[:8]) + "; nothing stored")
        n_diff = len(group_a) + len(group_b) + len(extra)
        print(f"  {y}: {path.name} ({rank_text(wb['rank'])}): {len(held)} "
              f"rows matched by name, {cells} cells, {n_diff} differences"
              + (f"; reconciliation not checked for {', '.join(not_checked)}"
                 if not_checked else ""))
        if group_a:
            print(f"    (a) restored (held NULL, published 0): "
                  f"{', '.join(group_a)}")
        if group_b:
            print(f"    (b) named rule: {', '.join(group_b)}")
        if keep_null:
            print(f"    held NULL that a named rule keeps NULL: "
                  f"{', '.join(keep_null)}")
        for k in extra:
            print(f"    in the file, not held: {k}"
                  + ("" if k in told else f": {explained[k]}"))
            told.add(k)
        for who, fl in sorted(recs.markers.items()):
            print(f"    marker(s) -> NULL {', '.join(fl)}: {who}")
        ed1 = []
        for k in sorted(held):
            r = dict(held[k])
            r["value_flag"] = _derived_flag(held[k], got.get(k))
            r[PERIOD] = y
            ed1.append(r)
        cur.execute(f"SELECT DISTINCT (loaded_at AT TIME ZONE 'UTC')::date "
                    f"FROM public.{spec.live_table} WHERE {PERIOD} = %s", (y,))
        loaded = sorted(r[0] for r in cur.fetchall())
        if len(loaded) != 1:
            halt(f"{y}: live rows carry {len(loaded)} load dates; nothing "
                 "stored")
        src = AS_LOADED_SOURCES.get(
            loaded[0], f"as loaded on {loaded[0]}; value_flag read from the "
                       "held file")
        where = f.get("url") or _file_source(path)
        plan["years"][y] = {
            "records": ed1, "published": loaded[0], "source_file": src,
            "content_sha": rows_content_sha(ed1),
            "ledger": None if n_diff else (ledger_source(where, wb["rank"]),
                                           sha)}
        plan["rows"] += len(ed1)
        print(f"    edition 1 as loaded: {len(ed1)} rows, published "
              f"{loaded[0]}; content sha256 "
              f"{plan['years'][y]['content_sha'][:16]}"
              + ("; ledger row for the file (it reproduces the year)"
                 if not n_diff else "; load stores the file as edition 2"))
    if not write:
        return plan
    cur.execute("SAVEPOINT s17_migrate")
    for y, v in plan["years"].items():
        ed = core.insert_edition(
            cur, spec, v["records"], y, release_label=core.AS_LOADED_LABEL,
            published_date=v["published"], source_file=v["source_file"],
            source_sha256=v["content_sha"], supersedes=None, strict=True)
        if ed != 1 or core.rows_differing(cur, spec, y, ed):
            halt(f"{spec.editions_table} {y}: edition 1 does not equal live; "
                 "rolled back")
        if v["ledger"]:
            pe.record_file_check(cur, prof, y, v["ledger"][0], v["ledger"][1],
                                 "unchanged", 1)
    cur.execute("RELEASE SAVEPOINT s17_migrate")
    return plan


def cmd_migrate_legacy(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    try:
        with conn.cursor() as cur:
            plan = migrate_legacy(cur, write=writing)
            if args.commit:
                led = [y for y, v in plan["years"].items() if v["ledger"]]
                log_run(cur, plan["rows"], f"migrate-legacy: edition 1 as "
                        f"loaded for marac_cases {', '.join(plan['years'])} "
                        f"({plan['rows']} rows) from the live table; proof "
                        "against the eight held SafeLives files; ledger rows "
                        f"for {', '.join(led) or 'no year'}. Live untouched.",
                        started)
                conn.commit()
                print("migrate-legacy: COMMITTED; pipeline_run_log row "
                      "written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# restore-edition
# ---------------------------------------------------------------------------

def restore_edition(cur, period, edition) -> int:
    """Store edition `edition` of `period` as the next edition ('restored
    from edition N') for refresh-latest to apply; return its number. Halts
    if the edition does not exist or is the tip."""
    spec = globals()["SPEC"]
    p = str(period)
    tip = core.latest_edition(cur, spec, p)
    recs = held_records(cur, spec, p, edition)
    if not recs:
        halt(f"{spec.editions_table} {p}: no edition {edition}")
    if edition == tip:
        halt(f"{p}: edition {edition} is the tip; nothing to restore")
    cur.execute(f"SELECT DISTINCT source_file FROM public.{spec.editions_table}"
                f" WHERE {PERIOD} = %s AND edition = %s", (p, edition))
    src = cur.fetchone()[0]
    new = core.insert_edition(
        cur, spec, recs, p, release_label=f"restored from edition {edition}",
        published_date=date.today(), source_file=src,
        source_sha256=rows_content_sha(recs), supersedes=tip, strict=True,
        allow_revert=True)
    if new == tip:
        halt(f"{p}: edition {edition} was not stored as the next edition; "
             "nothing to restore")
    return new


def cmd_restore_edition(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    spec = globals()["SPEC"]
    try:
        with conn.cursor() as cur:
            if not table_exists(cur, spec.editions_table):
                halt(f"{spec.editions_table} does not exist")
            new = restore_edition(cur, args.period, args.edition)
            print(f"{args.period}: edition {args.edition} "
                  + ("stored" if writing else "would be stored")
                  + f" as edition {new}; run refresh-latest to apply it")
            if args.commit:
                n = len(held_records(cur, spec, args.period, new))
                log_run(cur, n, f"restore-edition: {args.period} edition "
                        f"{args.edition} stored as edition {new}", started)
                conn.commit()
                print("COMMITTED; pipeline_run_log row written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="S17 SafeLives Marac data: editions loader on the "
        "period-editions engine. There is no sync-new; migrate-legacy "
        "records the held years once.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ddl", help="create the editions table, its triggers "
                       "and the file-check ledger (preview by default)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_ddl)
    sub.add_parser("status", help="what needs action; exit 1 if anything"
                   ).set_defaults(func=cmd_status)
    p = sub.add_parser("load", help="read the data page's workbooks (or "
                       "--file), compare and store each year (preview by "
                       "default)")
    p.add_argument("--year", type=_year_arg, metavar="YYYY-YY",
                   help="read, compare and store only this year")
    p.add_argument("--file", action="append", metavar="PATH",
                   help="a hand-placed workbook under data/raw or "
                   "data/reference (repeatable; nothing downloaded; identity "
                   "is read from the file)")
    p.add_argument("--no-page", action="store_true",
                   help="with --file only: do not read the data page "
                   "(recorded as such)")
    p.add_argument("--recheck", type=_year_arg, metavar="YYYY-YY",
                   help="compare this held year again although its file is "
                   "in the ledger")
    p.add_argument("--allow-older-file", action="store_true",
                   help="compare and store a year whose tip comes from a "
                   "newer file (logged)")
    p.add_argument("--acknowledge", action="append", type=_year_arg,
                   metavar="YYYY-YY",
                   help="release a year's soft stop conditions after reading "
                   "the preview (repeatable; logged; never a partial file)")
    p.add_argument("--accept-reissue", action="append", type=_year_arg,
                   metavar="YYYY-YY",
                   help="store a held year from a file that differs from the "
                   "one its tip came from (repeatable; logged)")
    p.add_argument("--acknowledge-flips", metavar="NAME",
                   choices=sorted(ACKNOWLEDGED_FLIPS),
                   help="release exactly the values going to or from NULL "
                   "that a named entry lists (ACKNOWLEDGED_FLIPS; logged)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_load)
    p = sub.add_parser("refresh-latest", help="copy each year's latest "
                       "edition into the live table, loaded_at included "
                       "(preview by default)")
    p.add_argument("--accept-drift", action="append", type=_year_arg,
                   metavar="YYYY-YY",
                   help="overwrite this year although its live rows equal no "
                   "stored edition (repeatable)")
    p.add_argument("--accept-key-changes", action="append", type=_year_arg,
                   metavar="YYYY-YY",
                   help="apply this year's forces added and removed "
                   "(repeatable)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_refresh_latest)
    p = sub.add_parser("migrate-legacy", help="one-off: edition 1 as loaded "
                       "for every held year, with the proof (preview by "
                       "default)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_migrate_legacy)
    p = sub.add_parser("restore-edition", help="store an earlier edition's "
                       "rows as the next edition (preview by default)")
    p.add_argument("period", type=_year_arg, metavar="YYYY-YY")
    p.add_argument("edition", type=int)
    pe.mode_parser(p)
    p.set_defaults(func=cmd_restore_edition)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
