"""Pure tests for s17_marac_editions: discovery, download refusals, workbook
identity, cells, records, the named zero rules, reconciliation, ranks,
hashes and stop conditions. No database, no network.

Fixture workbooks are written by the tests (write_marac) with the layout of
the held SafeLives files: a Notes sheet with the title row, a Cases sheet
(title, header row, England and Wales, England, then one block per region:
the region row and its police force area rows, blank rows between) and a
Referral routes sheet (header row with a 'Number'/'%' sub-header row). The
five held .xls years cannot be written here (no xlwt); the real held files
under data/raw/s17_marac are read where present (read-only).
"""
import io
import re
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import s17_marac_editions as m  # noqa: E402

HEADER = ["Area Name", "Number of MARACs", "Number of cases discussed",
          "Recommended number of cases",
          "Number of cases per 10,000 adult females",
          "Number of repeat cases", "Percentage of repeat cases",
          "Number of children in household"]
CASES_COLS = ("marac_count", "cases_discussed", "recommended_cases",
              "cases_per_10k_adult_females", "repeat_cases",
              "repeat_cases_pct", "children_in_household")
REGIONS = (("North West", ("Cumbria", "Lancashire")),
           ("West Midlands", ("Staffordshire", "West Midlands")),
           ("East", ("Norfolk", "Suffolk")),
           ("London", ("Metropolitan Police", "City of London")))
WALES = ("Wales", ("Gwent", "North Wales"))
PFAS = frozenset(n for _, ns in REGIONS for n in ns)
HELD = Path(__file__).resolve().parents[2] / "data" / "raw" / "s17_marac"
HELD_FILES = {
    "2018-19": "Marac-data-2018-2019-England-and-Wales.xls",
    "2019-20": "Marac-data-2019-2020-England-and-Wales.xls",
    "2020-21": "Marac-data-2020-to-2021.xls",
    "2021-22": "Marac-data-2021-2022-for-publication.xls",
    "2022-23": "Marac-data-2022-2023.xls",
    "2023-24": "Marac-data-2023-2024.xlsx",
    "2024-25": "Marac-data-2024-2025.xlsx",
    "2025-26": "MARAC-DATA-2025-2026.xlsx",
}


def base_values(i):
    """Published numbers of the i-th police force area (0-based)."""
    cases = 100 * (i + 1) + 0.5
    repeat = 30 * (i + 1)
    return {"marac_count": i + 1, "cases_discussed": cases,
            "recommended_cases": 90 * (i + 1),
            "cases_per_10k_adult_females": 40.1234567 + i,
            "repeat_cases": repeat, "repeat_cases_pct": repeat / cases * 100,
            "children_in_household": 120 * (i + 1),
            "housing_referrals": 5 * (i + 1) + 0.25}


def fixture_values(values=None):
    """{name: {column: published}} of every force (English and Welsh), the
    City of London's and Lancashire's housing referrals a published 0 (as
    in the held 2024-25 and 2025-26 files; both stay 0), then `values`
    overrides."""
    out = {}
    names = [n for _, ns in REGIONS for n in ns] + list(WALES[1])
    for i, n in enumerate(names):
        out[n] = base_values(i)
    out["City of London"]["housing_referrals"] = 0
    out["Lancashire"]["housing_referrals"] = 0
    for n, cols in (values or {}).items():
        out.setdefault(n, base_values(len(out)))
        out[n].update(cols)
    return out


def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def set_modified(path, iso):
    """Rewrite the workbook's docProps/core.xml dcterms:modified (openpyxl
    stamps the save time)."""
    p = Path(path)
    with zipfile.ZipFile(p) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(<)",
                              rb"\g<1>" + iso.encode() + rb"\g<2>", data)
            z.writestr(info, data)
    p.write_bytes(buf.getvalue())


def write_marac(path, year=2026, *, notes_year=None, cases_year=None,
                values=None, header=None, housing_col=26, regions=None,
                wales=WALES, england_delta=0, names=None, footnotes=(),
                modified="2026-07-02T14:06:06Z", referral_names=None,
                sheets=("Notes", "Cases", "Referral routes")):
    """A SafeLives-style workbook for the year ending March `year`.
    values: {name: {column: published}} overrides (a column of CASES_COLS
    or housing_referrals; a string is written as published). names: {force:
    name as written in Cases} (e.g. a footnote marker). referral_names: the
    same for Referral routes. footnotes: lines written under both tables.
    england_delta: added to the England row's cases discussed."""
    import openpyxl
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    regions = REGIONS if regions is None else regions
    vals = fixture_values(values)
    names = names or {}
    referral_names = referral_names or names
    notes_year = notes_year or year
    cases_year = cases_year or year
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    blocks = list(regions) + ([wales] if wales else [])
    for _, forces in blocks:
        for n in forces:
            vals.setdefault(n, base_values(len(vals)))
    if "Notes" in sheets:
        ws = wb.create_sheet("Notes")
        ws.cell(2, 2, f"SafeLives Marac data England and Wales April "
                      f"{notes_year - 1} - March {notes_year}")
        ws.cell(5, 2, "Cases")
        ws.cell(5, 3, "Cases discussed at multi-agency risk assessment "
                      "conferences (MARACs), by police force area and "
                      f"region, year ending March {cases_year}")
        ws.cell(6, 2, "Referral routes")
        ws.cell(6, 3, "Sources of referrals to multi-agency risk assessment "
                      "conferences (MARACs), by police force area and "
                      f"region, year ending March {cases_year}")

    def total(force_names, col):
        s = 0  # a fixture sum; not a source value
        for n in force_names:
            v = vals[n][col]
            if _num(v):
                s += v
        return s

    english = [n for _, ns in regions for n in ns]
    welsh = list(wales[1]) if wales else []
    if "Cases" in sheets:
        ws = wb.create_sheet("Cases")
        ws.cell(1, 2, "Marac cases")
        ws.cell(2, 2, "Cases discussed at multi-agency risk assessment "
                      "conferences (MARACs), by police force area and region, "
                      f"year ending March {cases_year}")
        for j, h in enumerate(header or HEADER):
            ws.cell(4, 2 + j, h)
        r = 6
        for label, forces in (("England and Wales", english + welsh),
                              ("England", english)):
            ws.cell(r, 2, label)
            for j, c in enumerate(CASES_COLS):
                v = total(forces, c)
                if label == "England" and c == "cases_discussed":
                    v += england_delta
                ws.cell(r, 3 + j, v)
            r += 2
        for region, forces in blocks:
            ws.cell(r, 2, region)
            for j, c in enumerate(CASES_COLS):
                ws.cell(r, 3 + j, total(forces, c))
            r += 1
            for n in forces:
                ws.cell(r, 2, names.get(n, n))
                for j, c in enumerate(CASES_COLS):
                    ws.cell(r, 3 + j, vals[n][c])
                r += 1
            r += 1
        for line in footnotes:
            ws.cell(r, 2, line)
            r += 1
    if "Referral routes" in sheets:
        ws = wb.create_sheet("Referral routes")
        ws.cell(1, 2, "Marac referral routes")
        ws.cell(2, 2, "Sources of referrals to multi-agency risk assessment "
                      "conferences (MARACs), by police force area and region, "
                      f"year ending March {cases_year}1")
        ws.cell(4, 2, "Area Name")
        cols = list(range(3, housing_col + 5, 2))
        hcol = housing_col + 1          # 0-based index -> 1-based column
        for k, c in enumerate(cols):
            ws.cell(4, c, "Housing" if c == hcol else f"Source {k}")
            ws.cell(5, c, "Number ")
            ws.cell(5, c + 1, "%")
        r = 7
        for label, forces in (("England and Wales", english + welsh),
                              ("England", english)):
            ws.cell(r, 2, label)
            ws.cell(r, hcol, total(forces, "housing_referrals"))
            r += 2
        for region, forces in blocks:
            ws.cell(r, 2, region)
            ws.cell(r, hcol, total(forces, "housing_referrals"))
            r += 1
            for n in forces:
                ws.cell(r, 2, referral_names.get(n, n))
                for c in cols:
                    ws.cell(r, c, vals[n]["housing_referrals"] if c == hcol
                            else 1.5)
                    ws.cell(r, c + 1, 2.5)
                r += 1
            r += 1
        for line in footnotes:
            ws.cell(r, 2, line)
            r += 1
    wb.save(path)
    if modified:
        set_modified(path, modified)
    return path


PAGE = """<html><head><title>Our quarterly Marac data - SafeLives</title>
<link rel="canonical" href="{base}" /></head><body>
{links}
<a href="https://safelives.org.uk/guidelines.pdf">Guidelines</a>
</body></html>"""


def page_html(links, base=None):
    return PAGE.format(base=base or m.DATA_PAGE, links="\n".join(
        f'<a href="{u}">Download the data (excel file)</a>' for u in links))


HELD_LINKS = [
    "https://safelives.org.uk/wp-content/uploads/MARAC-DATA-2025-2026.xlsx",
    "https://safelives.org.uk/wp-content/uploads/Marac-data-2024-2025.xlsx",
    "/wp-content/uploads/2024/01/Marac-data-2020-to-2021.xls",
]


class Tmp(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def wb(self, name="f.xlsx", **kw):
        return m.read_workbook(write_marac(self.root / name, **kw))


class DataPageLinks(unittest.TestCase):

    def test_links_by_year_in_both_name_forms(self):
        html = page_html(HELD_LINKS + [HELD_LINKS[0]])   # a repeated link
        got = m.data_page_links(html, held=("2020-21",))
        self.assertEqual(got, {
            "2025-26": HELD_LINKS[0], "2024-25": HELD_LINKS[1],
            "2020-21": "https://safelives.org.uk/wp-content/uploads/2024/01/"
                       "Marac-data-2020-to-2021.xls"})

    def test_none_for_a_held_year_halts_listing_every_link(self):
        with self.assertRaises(ValueError) as e:
            m.data_page_links(page_html(HELD_LINKS), held=("2019-20",))
        msg = str(e.exception)
        self.assertIn("2019-20", msg)
        for u in HELD_LINKS[:2]:
            self.assertIn(u, msg)

    def test_no_workbook_link_at_all_halts(self):
        with self.assertRaises(ValueError):
            m.data_page_links(page_html([]))

    def test_two_links_for_one_year_halt(self):
        other = "https://safelives.org.uk/wp-content/uploads/Marac-data-" \
                "2024-2025-v2.xlsx"
        with self.assertRaises(ValueError) as e:
            m.data_page_links(page_html(HELD_LINKS + [other]))
        self.assertIn(other, str(e.exception))
        self.assertIn(HELD_LINKS[1], str(e.exception))

    def test_a_workbook_link_that_does_not_fit_halts(self):
        for odd in ("https://safelives.org.uk/x/Marac-data-2026.xlsx",
                    "https://safelives.org.uk/x/Marac-data-2024-2026.xlsx"):
            with self.subTest(odd=odd), self.assertRaises(ValueError) as e:
                m.data_page_links(page_html(HELD_LINKS + [odd]))
            self.assertIn(odd, str(e.exception))


class Resp:
    def __init__(self, body, status=200, url=None, history=()):
        self.content = body
        self.text = body.decode("utf-8", "replace") if isinstance(
            body, bytes) else body
        self.status_code = status
        self.url = url
        self.history = list(history)
        self.headers = {}

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"{self.status_code} error")


class Session:
    def __init__(self, resp):
        self.resp = resp
        self.calls = []

    def get(self, url, **kw):
        self.calls.append((url, kw))
        if self.resp.url is None:
            self.resp.url = url
        return self.resp


URL = "https://safelives.org.uk/wp-content/uploads/MARAC-DATA-2025-2026.xlsx"


class Fetch(Tmp):

    def refused(self, resp):
        with self.assertRaises(SystemExit) as e, \
                mock.patch("sys.stdout", io.StringIO()):
            m.fetch(URL, self.root / "dl", session=Session(resp))
        msg = str(e.exception.code)
        self.assertIn(f"SafeLives refused the request: download {URL} by "
                      "hand into data/raw/s17_marac/ and rerun with --file",
                      msg)
        self.assertEqual(list((self.root / "dl").glob("*.xlsx")), [])

    def test_a_403_halts_with_the_hand_download_message(self):
        self.refused(Resp(b"<html>Forbidden</html>", status=403))

    def test_a_cloudflare_challenge_page_halts(self):
        self.refused(Resp(b"<!DOCTYPE html><html><head><title>Just a "
                          b"moment...</title></head><body>challenge-platform"
                          b"</body></html>"))

    def test_a_non_workbook_body_halts(self):
        self.refused(Resp(b"<html><body>Our quarterly Marac data</body></html>"
                          + b" " * 60000))

    def test_the_plain_user_agent_and_the_final_url(self):
        body = write_marac(self.root / "src.xlsx").read_bytes()
        final = URL.replace("MARAC-DATA", "Marac-Data")
        s = Session(Resp(body, url=final))
        with mock.patch("sys.stdout", io.StringIO()):
            path, got = m.fetch(URL, self.root / "dl", session=s)
        self.assertEqual(got, final)
        self.assertEqual(path.name, "Marac-Data-2025-2026.xlsx")
        self.assertEqual(path.read_bytes(), body)
        ua = s.calls[0][1]["headers"]["User-Agent"]
        self.assertEqual(ua, m.USER_AGENT)
        self.assertIn("ucws-pipeline", ua)
        self.assertNotIn("Mozilla", ua)

    def test_the_same_name_rule(self):
        dl = self.root / "dl"
        a = write_marac(self.root / "a.xlsx").read_bytes()
        b = write_marac(self.root / "b.xlsx", values={
            "Cumbria": {"cases_discussed": 7}}).read_bytes()
        with mock.patch("sys.stdout", io.StringIO()):
            p1, _ = m.fetch(URL, dl, session=Session(Resp(a)))
            p2, _ = m.fetch(URL, dl, session=Session(Resp(a)))
            p3, _ = m.fetch(URL, dl, session=Session(Resp(b)))
        self.assertEqual(p1, p2)
        self.assertEqual(p1.read_bytes(), a)            # never replaced
        self.assertNotEqual(p3, p1)
        self.assertTrue(p3.name.startswith("MARAC-DATA-2025-2026-"))
        self.assertEqual(p3.read_bytes(), b)

    def test_the_data_page_refused_halts_with_the_hand_download_message(self):
        with self.assertRaises(SystemExit) as e, \
                mock.patch("sys.stdout", io.StringIO()):
            m.fetch_page(m.REGISTRY_URL, session=Session(
                Resp(b"denied", status=403)), dest=self.root)
        msg = str(e.exception.code)
        self.assertIn("SafeLives refused the request", msg)
        self.assertIn("--file", msg)


class Identity(Tmp):

    def test_identity_from_the_file(self):
        wb = self.wb()
        self.assertEqual(wb["year"], "2025-26")
        self.assertEqual(wb["notes_years"], (2025, 2026))
        self.assertEqual(wb["cases_year"], 2026)
        self.assertEqual(wb["rank"], ("2025-26", datetime(
            2026, 7, 2, 14, 6, 6, tzinfo=timezone.utc)))
        self.assertEqual(len(wb["identity"]["sha256"]), 64)

    def test_notes_and_cases_title_years_disagreeing_raise(self):
        with self.assertRaises(ValueError) as e:
            self.wb(notes_year=2025)
        self.assertIn("Notes", str(e.exception))
        with self.assertRaises(ValueError):
            self.wb("g.xlsx", cases_year=2025)

    def test_the_requested_year_must_be_the_files(self):
        path = write_marac(self.root / "f.xlsx")
        with self.assertRaises(ValueError) as e:
            m.read_workbook(path, year="2024-25")
        self.assertIn("2024-25", str(e.exception))
        self.assertEqual(m.read_workbook(path, year="2025-26")["year"],
                         "2025-26")

    def test_a_named_notes_erratum_passes_only_for_its_own_bytes(self):
        path = write_marac(self.root / "f.xlsx", notes_year=2025)
        sha = m.manual_input.file_identity(path)["sha256"]
        with mock.patch.dict(m.NOTES_TITLE_ERRATA, {sha: {
                "file": path.name, "year": "2025-26",
                "notes_says": (2024, 2025), "evidence": "test"}}):
            wb = m.read_workbook(path)
            self.assertEqual(wb["year"], "2025-26")
            self.assertTrue(wb["erratum"])
            other = write_marac(self.root / "g.xlsx", notes_year=2025,
                                values={"Cumbria": {"marac_count": 9}})
            with self.assertRaises(ValueError):
                m.read_workbook(other)

    def test_a_changed_header_raises(self):
        h = list(HEADER)
        h[2] = "Number of cases heard"
        with self.assertRaises(ValueError) as e:
            self.wb(header=h)
        self.assertIn("Number of cases heard", str(e.exception))

    def test_a_named_header_alias_is_recorded(self):
        h = list(HEADER)
        h[1] = "Number of Maracs"
        wb = self.wb(header=h)
        self.assertEqual(wb["header_aliases"], ["Number of Maracs"])
        self.assertIn("Number of Maracs", m.HEADER_ALIASES)

    def test_a_missing_sheet_raises(self):
        with self.assertRaises(ValueError) as e:
            self.wb(sheets=("Notes", "Cases"))
        self.assertIn("Referral routes", str(e.exception))

    def test_housing_found_by_header_text_when_columns_move(self):
        a = m.records(self.wb("a.xlsx", housing_col=26), "2025-26", PFAS)
        b = m.records(self.wb("b.xlsx", housing_col=18), "2025-26", PFAS)
        self.assertEqual([(r["pfa_name_safelives"], r["housing_referrals"])
                          for r in a],
                         [(r["pfa_name_safelives"], r["housing_referrals"])
                          for r in b])
        self.assertEqual(self.wb("c.xlsx", housing_col=18)["housing_col"], 18)


class Cells(unittest.TestCase):

    def test_numbers_and_markers(self):
        self.assertEqual(m.cell(12, "x"), (12, None))
        self.assertEqual(m.cell(0, "x"), (0, None))
        self.assertEqual(m.cell(3.5, "x"), (3.5, None))
        self.assertEqual(m.cell("No Data", "x"), (None, "not_submitted"))
        self.assertEqual(m.cell("No population data", "x"),
                         (None, "no_population"))
        self.assertEqual(m.cell("No data", "x"), (None, "not_submitted"))
        self.assertIn("No data", m.MARKER_ALIASES)

    def test_anything_else_halts_naming_where(self):
        for v in ("#DIV/0!", "", None, "n/a", "-", True, "12",
                  float("nan")):
            with self.subTest(v=v), self.assertRaises(ValueError) as e:
                m.cell(v, "Cases row 9 Norfolk")
            self.assertIn("Cases row 9 Norfolk", str(e.exception))


def by_name(recs):
    return {r["pfa_name_safelives"]: r for r in recs}


class Records(Tmp):

    def test_the_english_forces_rounded_to_the_column_scales(self):
        recs = m.records(self.wb(), "2025-26", PFAS)
        self.assertEqual(sorted(by_name(recs)), sorted(PFAS))
        r = by_name(recs)["Cumbria"]
        self.assertEqual(r["financial_year"], "2025-26")
        self.assertEqual(r["marac_count"], 1)
        self.assertEqual(r["cases_discussed"], Decimal("100.50"))
        self.assertEqual(r["cases_per_10k_adult_females"],
                         Decimal("40.123457"))
        self.assertEqual(r["repeat_cases_pct"], Decimal("29.8507"))
        self.assertEqual(r["housing_referrals"], Decimal("5.25"))
        self.assertIsNone(r["value_flag"])

    def test_a_published_zero_stays_zero(self):
        r = by_name(m.records(self.wb(), "2025-26", PFAS))["City of London"]
        self.assertEqual(r["housing_referrals"], Decimal("0.00"))
        self.assertIsNotNone(r["housing_referrals"])
        self.assertIsNone(r["value_flag"])

    def test_no_data_and_no_population_data(self):
        nd = {c: "No Data" for c in CASES_COLS + ("housing_referrals",)}
        recs = by_name(m.records(self.wb(values={
            "Norfolk": nd, "Suffolk": {
                "cases_per_10k_adult_females": "No population data"}}),
            "2025-26", PFAS))
        self.assertTrue(all(recs["Norfolk"][c] is None for c in m.VALUES))
        self.assertEqual(recs["Norfolk"]["value_flag"], "not_submitted")
        s = recs["Suffolk"]
        self.assertIsNone(s["cases_per_10k_adult_females"])
        self.assertEqual(s["value_flag"], "no_population")
        self.assertIsNotNone(s["cases_discussed"])

    def test_other_text_halts(self):
        with self.assertRaises(ValueError) as e:
            m.records(self.wb(values={"Suffolk": {"repeat_cases": "n/a"}}),
                      "2025-26", PFAS)
        self.assertIn("Suffolk", str(e.exception))

    def test_two_reasons_in_one_row_halt(self):
        with self.assertRaises(ValueError):
            m.records(self.wb(values={"Suffolk": {
                "cases_per_10k_adult_females": "No population data",
                "children_in_household": "No Data"}}), "2025-26", PFAS)

    def test_a_fractional_marac_count_halts(self):
        with self.assertRaises(ValueError):
            m.records(self.wb(values={"Suffolk": {"marac_count": 2.5}}),
                      "2025-26", PFAS)

    NORFOLK_2324 = {"marac_count": 0, "cases_discussed": 0,
                    "recommended_cases": 0,
                    "cases_per_10k_adult_females": "No population data",
                    "repeat_cases": 0, "repeat_cases_pct": "#DIV/0!",
                    "children_in_household": 0, "housing_referrals": 0}

    def y2324(self, **extra):
        """2023-24 overrides: Norfolk and the West Midlands force in the
        published form their rules name."""
        return {"Norfolk": self.NORFOLK_2324,
                "West Midlands": self.NORFOLK_2324, **extra}

    def test_the_norfolk_rule_fires_only_on_its_key_and_is_listed(self):
        recs = m.records(self.wb(year=2024, values=self.y2324()), "2023-24",
                         PFAS)
        n = by_name(recs)["Norfolk"]
        self.assertTrue(all(n[c] is None for c in m.VALUES))
        self.assertEqual(n["value_flag"], "not_submitted")
        names = {a["rule"] for a in recs.applied}
        self.assertEqual(names, {"NORFOLK_2023_24_NOT_SUBMITTED",
                                 "WEST_MIDLANDS_2023_24_NOT_SUBMITTED"})
        self.assertEqual(len([a for a in recs.applied
                              if a["pfa"] == "Norfolk"]), len(m.VALUES))
        self.assertEqual(by_name(recs)["Suffolk"]["value_flag"], None)
        # the same published form in another year, or another force: halts
        for year, label, vals in (
                (2025, "2024-25", {"Norfolk": self.NORFOLK_2324}),
                (2024, "2023-24", self.y2324(Suffolk=self.NORFOLK_2324))):
            with self.subTest(year=year), self.assertRaises(ValueError) as e:
                m.records(self.wb(f"x{year}.xlsx", year=year, values=vals),
                          label, PFAS)
            self.assertIn("#DIV/0!", str(e.exception))

    def test_a_rule_whose_published_form_changed_halts(self):
        form = dict(self.NORFOLK_2324, cases_discussed=12)
        with self.assertRaises(ValueError) as e:
            m.records(self.wb(year=2024, values=self.y2324(Norfolk=form)),
                      "2023-24", PFAS)
        self.assertIn("NORFOLK_2023_24_NOT_SUBMITTED", str(e.exception))
        wm = dict(self.NORFOLK_2324, marac_count=7)
        with self.assertRaises(ValueError) as e:
            m.records(self.wb("w.xlsx", year=2024, values=self.y2324(**{
                "West Midlands": wm})), "2023-24", PFAS)
        self.assertIn("WEST_MIDLANDS_2023_24_NOT_SUBMITTED", str(e.exception))

    def test_lancashire_housing_zeros_stay_published_zeros(self):
        """Scott, 2026-10-10: no not-counted rule for Lancashire; the files
        carry no Lancashire note in 2023-24 to 2025-26."""
        z = {"housing_referrals": 0}
        for year, label, extra in ((2024, "2023-24", self.y2324()),
                                   (2025, "2024-25", {}),
                                   (2026, "2025-26", {})):
            with self.subTest(year=label):
                recs = m.records(self.wb(f"z{year}.xlsx", year=year, values={
                    "Lancashire": z, **extra}), label, PFAS)
                r = by_name(recs)["Lancashire"]
                self.assertEqual(r["housing_referrals"], Decimal("0.00"))
                self.assertIsNone(r["value_flag"])
                self.assertFalse([a for a in recs.applied
                                  if a["pfa"] == "Lancashire"])
        self.assertFalse([r for r in m.ZERO_RULES
                          if r["pfa"] == "Lancashire"])
        self.assertNotIn("not_counted", m.FLAGS)

    def test_exactly_the_named_rules(self):
        self.assertEqual([r["name"] for r in m.ZERO_RULES],
                         ["NORFOLK_2023_24_NOT_SUBMITTED",
                          "WEST_MIDLANDS_2023_24_NOT_SUBMITTED"])
        for r in m.ZERO_RULES:
            self.assertTrue(r["evidence"])
            self.assertTrue(r["decided"])
            self.assertIn(r["flag"], m.FLAGS)

    def test_a_missing_or_unknown_english_force_halts(self):
        regions = (("North West", ("Cumbria",)),) + REGIONS[1:]
        with self.assertRaises(ValueError) as e:
            m.records(self.wb(regions=regions), "2025-26", PFAS)
        self.assertIn("Lancashire", str(e.exception))
        regions = (("North West", ("Cumbria", "Lancashire", "Blackpool")),) \
            + REGIONS[1:]
        with self.assertRaises(ValueError) as e:
            m.records(self.wb("u.xlsx", regions=regions), "2025-26", PFAS)
        self.assertIn("Blackpool", str(e.exception))

    def test_welsh_forces_are_never_stored(self):
        recs = m.records(self.wb(), "2025-26", PFAS)
        self.assertFalse({"Gwent", "North Wales"} & set(by_name(recs)))
        # a Welsh force in an English region halts
        regions = (("North West", ("Cumbria", "Lancashire", "Gwent")),) \
            + REGIONS[1:]
        with self.assertRaises(ValueError) as e:
            m.records(self.wb("w.xlsx", regions=regions,
                              wales=("Wales", ("North Wales",))),
                      "2025-26", PFAS | {"Gwent"})
        self.assertIn("Gwent", str(e.exception))

    def test_a_region_and_a_force_of_the_same_name(self):
        recs = by_name(m.records(self.wb(), "2025-26", PFAS))
        # the West Midlands force row, never the region row
        self.assertEqual(recs["West Midlands"]["marac_count"], 4)

    def test_a_footnote_marker_on_a_name(self):
        fn = ("1. Data from one MARAC within the Lancashire police force is "
              "not included",)
        recs = m.records(self.wb(names={"Lancashire": "Lancashire1"},
                                 footnotes=fn), "2025-26", PFAS)
        self.assertIn("Lancashire", by_name(recs))
        self.assertEqual(recs.footnotes["Lancashire"], [fn[0]])
        with self.assertRaises(ValueError):
            m.records(self.wb("n.xlsx", names={"Lancashire": "Lancashire7"},
                              footnotes=fn), "2025-26", PFAS)


class Reconcile(Tmp):

    def test_the_english_forces_sum_to_the_england_row(self):
        problems, not_checked = m.reconcile(self.wb())
        self.assertEqual(problems, [])
        self.assertEqual(not_checked, [])

    def test_off_by_one_is_a_problem(self):
        problems, _ = m.reconcile(self.wb(england_delta=1))
        self.assertEqual(len(problems), 1)
        self.assertIn("cases_discussed", problems[0])

    def test_an_unpublished_part_is_not_checked(self):
        nd = {c: "No Data" for c in CASES_COLS + ("housing_referrals",)}
        problems, not_checked = m.reconcile(self.wb(values={"Norfolk": nd}))
        self.assertEqual(problems, [])
        self.assertEqual(sorted(not_checked),
                         sorted(["marac_count", "cases_discussed",
                                 "recommended_cases", "repeat_cases",
                                 "children_in_household"]))


class Rank(Tmp):

    def test_rank_from_the_workbook_never_the_name(self):
        old = write_marac(self.root / "MARAC-DATA-2025-2026-v9.xlsx",
                          modified="2026-07-01T09:00:00Z")
        new = write_marac(self.root / "old-name.xlsx",
                          modified="2026-07-03T09:00:00Z")
        ro = m.release_rank(m.read_workbook(old))
        rn = m.release_rank(m.read_workbook(new))
        self.assertLess(ro[1], rn[1])
        self.assertEqual(ro[0], "2025-26")

    def test_rank_from_a_source_file_string(self):
        wb = self.wb()
        s = m.edition_source(wb, {"where": "x", "file_name": "f.xlsx"})
        self.assertEqual(m.release_rank(s), wb["rank"])
        self.assertIsNone(m.release_rank("as loaded: n8n"))
        self.assertIsNone(m.release_rank("MARAC-DATA-2025-2026.xlsx"))
        s = m.ledger_source("https://x/y.xls", ("2019-20", None))
        self.assertEqual(m.release_rank(s), ("2019-20", None))
        self.assertEqual(m.parse_ledger_source(s),
                         ("https://x/y.xls", ("2019-20", None)))


def rec(name, year="2025-26", flag=None, **vals):
    r = {"pfa_name_safelives": name, "financial_year": year,
         "value_flag": flag}
    for c in m.VALUES:
        r[c] = vals.get(c, Decimal("1.00") if c != "marac_count" else 1)
    return r


class Hashes(unittest.TestCase):

    def test_null_and_zero_differ_and_order_does_not_matter(self):
        a = [rec("A", housing_referrals=Decimal("0.00")), rec("B")]
        b = [rec("A", housing_referrals=None, flag="not_counted"), rec("B")]
        c = [rec("A", housing_referrals=None), rec("B")]
        self.assertNotEqual(m.rows_content_sha(a), m.rows_content_sha(b))
        self.assertNotEqual(m.rows_content_sha(b), m.rows_content_sha(c))
        self.assertEqual(m.rows_content_sha(a),
                         m.rows_content_sha(list(reversed(a))))


class StopConditions(unittest.TestCase):

    def recs(self, n=10, cases=Decimal("100.00"), **over):
        out = [rec(f"P{i}", cases_discussed=cases) for i in range(n)]
        for r in out:
            r.update(over.get(r["pfa_name_safelives"], {}))
        return out

    def test_a_new_year_moving_england_cases_by_more_than_25pc(self):
        prev = self.recs()
        hard, soft = m.period_problems(self.recs(cases=Decimal("126.00")),
                                       None, prev, kind="new")
        self.assertEqual(hard, [])
        self.assertTrue(any("England" in s for s in soft))
        hard, soft = m.period_problems(self.recs(cases=Decimal("124.00")),
                                       None, prev, kind="new")
        self.assertFalse(any("England" in s for s in soft))

    def test_a_new_year_with_more_than_6_forces_moving_40pc(self):
        prev = self.recs()
        big = {f"P{i}": {"cases_discussed": Decimal("141.00")}
               for i in range(7)}
        small = {f"P{i}": {"cases_discussed": Decimal("59.00")}
                 for i in range(7)}
        over = {**{k: v for k, v in big.items() if k < "P4"},
                **{k: v for k, v in small.items() if k >= "P4"}}
        _, soft = m.period_problems(self.recs(**over), None, prev,
                                    kind="new")
        self.assertTrue(any("7 police force" in s for s in soft))
        six = {k: v for k, v in big.items() if k != "P6"}
        _, soft = m.period_problems(self.recs(**six), None, prev, kind="new")
        self.assertFalse(any("police force" in s for s in soft))

    def test_a_revised_year(self):
        tip = self.recs()
        hard, soft = m.period_problems(self.recs(), tip, None, kind="revised")
        self.assertEqual((hard, soft), ([], []))
        hard, _ = m.period_problems(self.recs(n=9), tip, None, kind="revised")
        self.assertTrue(any("P9" in h for h in hard))
        hard, soft = m.period_problems(self.recs(
            P3={"repeat_cases": Decimal("2.00")}), tip, None, kind="revised")
        self.assertEqual(hard, [])
        self.assertTrue(any("P3" in s for s in soft))
        hard, soft = m.period_problems(self.recs(n=11), tip, None,
                                       kind="revised")
        self.assertEqual(hard, [])
        self.assertTrue(any("P10" in s for s in soft))

    def test_flips_need_exactly_the_named_entry(self):
        tip = [rec("A", housing_referrals=None), rec("B")]
        new = [rec("A", housing_referrals=Decimal("0.00")), rec("B")]
        fl = m.flips(new, tip)
        self.assertEqual(fl, {"A/housing_referrals": (None, "0.00")})
        self.assertTrue(m.flip_problems(None, "2025-26", fl))
        entry = {"decided": "t", "why": "t",
                 "periods": {"2025-26": {"A/housing_referrals":
                                         (None, "0.00")}}}
        with mock.patch.dict(m.ACKNOWLEDGED_FLIPS, {"t": entry}):
            self.assertEqual(m.flip_problems("t", "2025-26", fl), [])
            self.assertTrue(m.flip_problems("t", "2024-25", fl))
            more = dict(fl, **{"B/repeat_cases": ("1.00", None)})
            self.assertTrue(m.flip_problems("t", "2025-26", more))
        self.assertEqual(m.restored(new, tip), ["A housing_referrals"])

    def test_the_first_load_entry_lists_the_known_cells(self):
        self.assertEqual(sorted(m.ACKNOWLEDGED_FLIPS),
                         ["s17-restored-zeros-2026-10"])
        e = m.ACKNOWLEDGED_FLIPS["s17-restored-zeros-2026-10"]
        self.assertNotIn("2025-26", e["periods"])
        self.assertEqual(e["periods"]["2023-24"],
                         {"Lancashire/housing_referrals": (None, "0.00")})
        self.assertEqual(e["periods"]["2024-25"],
                         {"City of London/housing_referrals": (None, "0.00"),
                          "Lancashire/housing_referrals": (None, "0.00")})
        self.assertEqual(len(e["periods"]["2018-19"]), 5)
        # every entry restores a NULL to 0: nothing goes to NULL
        self.assertTrue(all(old is None and new is not None
                            for cells in e["periods"].values()
                            for old, new in cells.values()))


@unittest.skipUnless(all((HELD / f).exists() for f in HELD_FILES.values()),
                     "held SafeLives files not present")
class HeldFiles(unittest.TestCase):
    """The eight held files, read-only."""

    def test_each_file_reads_as_its_own_year(self):
        for year, f in HELD_FILES.items():
            with self.subTest(year=year):
                wb = m.read_workbook(HELD / f, year=year)
                self.assertEqual(wb["year"], year)
                self.assertEqual(bool(wb["erratum"]), year == "2022-23")
                self.assertEqual(wb["header_aliases"],
                                 ["Number of Maracs"] if year == "2019-20"
                                 else [])
                self.assertEqual(wb["rank"][1] is None, f.endswith(".xls"))
                problems, _ = m.reconcile(wb)
                self.assertEqual(problems, [])

    def test_2023_24_rules_and_markers(self):
        wb = m.read_workbook(HELD / HELD_FILES["2023-24"])
        recs = by_name(m.records(wb, "2023-24", self.pfas()))
        for who in ("Norfolk", "West Midlands"):
            self.assertTrue(all(recs[who][c] is None for c in m.VALUES))
            self.assertEqual(recs[who]["value_flag"], "not_submitted")
        self.assertEqual(recs["Lancashire"]["housing_referrals"],
                         Decimal("0.00"))
        self.assertIsNone(recs["Lancashire"]["value_flag"])
        self.assertEqual(len(recs), 39)
        self.assertEqual(recs["City of London"]["housing_referrals"],
                         Decimal("2.00"))

    def test_2022_23_no_data_alias(self):
        wb = m.read_workbook(HELD / HELD_FILES["2022-23"])
        recs = by_name(m.records(wb, "2022-23", self.pfas()))
        self.assertEqual(recs["Norfolk"]["value_flag"], "not_submitted")
        self.assertEqual(recs["Lancashire"]["housing_referrals"],
                         Decimal("13.00"))
        self.assertEqual(recs["City of London"]["children_in_household"],
                         Decimal("0.00"))

    @staticmethod
    def pfas():
        return frozenset(m.PFA_NAMES_2026_10)


if __name__ == "__main__":
    unittest.main()
