"""Pure tests for s3_mye_editions: discovery, the workbook reader, LA
cells, reconciliation, records, ranks, ledger and planning, the stop
conditions and the download rules.

No database and no network: every .xlsx file is written by the tests
(openpyxl, the July 2026 layout and the July 2025 layout), and the ONS
replies and downloads are stubbed. The fixture writers are reused by
test_s3_mye_loader. Nothing here imports or runs a retired script.
"""
import contextlib
import hashlib
import io
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import s3_mye_editions as m  # noqa: E402

# the fixture's English authorities, published on the earlier
# Barnsley/Sheffield codes (the form the mid-2024 and mid-2025 files use)
CODES = ["E06000001", "E06000002", "E08000016", "E08000019", "E09000033"]
NAMES = {"E06000001": "Hartlepool", "E06000002": "Middlesbrough",
         "E08000016": "Barnsley", "E08000019": "Sheffield",
         "E08000038": "Barnsley", "E08000039": "Sheffield",
         "E09000033": "Westminster", "E07000999": "Nowhere"}
GEOG = {"E06": "Unitary Authority", "E08": "Metropolitan District",
        "E09": "London Borough", "E07": "Non-metropolitan District"}
WALES = [("W06000001", "Isle of Anglesey"), ("W06000002", "Gwynedd")]
LANDING = m.LANDING
ED_URI = LANDING + "/mid{}2023localauthorityboundaries"


def quiet():
    return contextlib.redirect_stdout(io.StringIO())


def value(code, year):
    """A plausible published population for a fixture authority."""
    i = CODES.index(code) if code in CODES else 7
    return 90000 + 25000 * i + 400 * (year - 2011)


def welsh(code, year):
    return 60000 + 300 * (year - 2011) + (7 if code.endswith("2") else 0)


def _month(d) -> str:
    return f"{d.day} {d:%B %Y}"


def write_file(path, year=2025, *, published=None, codes=None, values=None,
               mye2=None, england=None, cover=None, notes=None,
               mye2_title_year=None, mye4_title_year=None, mye2_header=None,
               mye4_header=None, mye7="revisions", revised=None,
               extra_rows=(), first=2011):
    """A mid-year tables workbook in the July 2026 layout (mye7
    'revisions'), or the July 2025 layout (mye7 'other': MYE7 is the
    uncertainty table). values: {(code, year): cell} overrides of MYE4 (and
    of MYE2 for the newest year unless mye2 overrides it); mye2: {code:
    cell} overrides of MYE2 All ages; england: {('mye2'|'mye4', year):
    cell} overrides of the England row; cover: {title_year, pub_line,
    next_line, date_line} overrides; revised: {(code, year): cell} of MYE7's
    revised column (default MYE4's value)."""
    import openpyxl
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    published = published or date(year + 1, 7, 29)
    codes = list(codes or CODES)
    values = values or {}
    mye2 = mye2 or {}
    england = england or {}
    cover = cover or {}
    years = list(range(year, first - 1, -1))

    def v4(c, y):
        return values[(c, y)] if (c, y) in values else value(c, y)

    def v2(c):
        return mye2[c] if c in mye2 else v4(c, year)

    def total(cells):
        return sum(x for x in cells if isinstance(x, int)
                   and not isinstance(x, bool))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cover sheet"
    ty = cover.get("title_year", year)
    lines = [f"Mid-Year Population Estimates, England and Wales, June {ty}",
             "This spreadsheet contains a selection of the data tables.",
             f"Mid-Year Population Estimates, England and Wales, June {ty}",
             "\nPublication dates"]
    if cover.get("pub_line", True):
        lines.append(cover.get("pub_text") or (
            "The data tables in this spreadsheet were originally published "
            f"at 9.30am {_month(published)}."))
    if cover.get("next_line", True):
        lines.append(f"The next publication is scheduled for summer "
                     f"{year + 2}.")
    lines += ["\nData supplier:", f"Release number: MYE{ty % 100:02d}"]
    if cover.get("date_line", True):
        lines.append(cover.get("date_text")
                     or f"Date published: {_month(published)}")
    for i, t in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=t)
    ws = wb.create_sheet("Contents")
    ws["A1"] = "Table of contents"
    ws = wb.create_sheet("Notes")
    nl = notes if notes is not None else [
        (4, "Geography: where subnational level data are presented, the "
            "population estimates reflect boundaries in place as of April "
            "2024."),
        (6, f"Revisions to mid-{year - 3} to mid-{year - 1} population "
            f"estimates: Population estimates for mid-{year - 3}, "
            f"mid-{year - 2} and mid-{year - 1} have been revised.")]
    ws.append(["Notes related to the data in this spreadsheet"])
    ws.append(["Issues affecting all tables", "Note text ", "Related Links"])
    for n, t in nl:
        ws.append([n, t, "No Link"])

    def lead(ws, title):
        ws.append([title])
        for t in ("This worksheet contains one table.", "Freeze panes.",
                  "Please choose from the links presented below:",
                  "This met my needs", "I need something slightly different",
                  "This is not what I need at all"):
            ws.append([t])

    def body(ws, cell_of, ncols):
        """The country, England, region, county, authority and Welsh
        rows; cell_of(code_or_None, col) -> the cell (None = aggregate)."""
        las = [(c, NAMES[c], GEOG[c[:3]]) for c in codes]
        ws.append(["K04000001", "ENGLAND AND WALES", "Country"]
                  + [cell_of("K", j) for j in range(ncols)])
        ws.append(["E92000001", "ENGLAND", "Country"]
                  + [cell_of("E", j) for j in range(ncols)])
        ws.append(["E12000001", "NORTH EAST", "Region"]
                  + [cell_of("R", j) for j in range(ncols)])
        ws.append(["E11000001", "Tyne and Wear", "Metropolitan County"]
                  + [cell_of("R", j) for j in range(ncols)])
        for c, n, g in las:
            ws.append([c, n, g] + [cell_of(c, j) for j in range(ncols)])
        for r in extra_rows:
            ws.append(list(r))
        ws.append(["W92000004", "WALES", "Country"]
                  + [cell_of("W", j) for j in range(ncols)])
        for c, n in WALES:
            ws.append([c, n, "Unitary Authority"]
                      + [cell_of(c, j) for j in range(ncols)])

    # MYE2 - Persons: All ages and two single years of age
    ws = wb.create_sheet("MYE2 - Persons")
    lead(ws, "MYE2: Persons by single year of age and sex for local "
             "authorities in England and Wales, "
             f"mid-{mye2_title_year or year}")
    ws.append(mye2_header or ["Code", "Name", "Geography", "All ages", "0",
                              "1"])

    def c2(code, j):
        if j:
            return 1000 + j
        if code == "E":
            return england.get(("mye2", year),
                               total(v2(c) for c in codes))
        if code in ("K", "W", "R"):
            return 12345
        if code.startswith("W"):
            return welsh(code, year)
        return v2(code)
    body(ws, c2, 3)
    for name in ("MYE2 - Females", "MYE2 - Males", "MYE3"):
        wb.create_sheet(name)["A1"] = name
    # MYE4: mid-<year> down to mid-2011
    ws = wb.create_sheet("MYE4")
    lead(ws, "MYE4: Population estimates: Summary for England and Wales, "
             f"mid-2011 to mid-{mye4_title_year or year}")
    ws.append(mye4_header or (["Code", "Name", "Geography"]
                              + [f"Mid-{y}" for y in years]))

    def c4(code, j):
        y = years[j]
        if code == "E":
            return england.get(("mye4", y), total(v4(c, y) for c in codes))
        if code in ("K", "W", "R"):
            return 54321
        if code.startswith("W"):
            return welsh(code, y)
        return v4(code, y)
    body(ws, c4, len(years))
    for name in ("MYE5", "MYE6"):
        wb.create_sheet(name)["A1"] = name
    ws = wb.create_sheet("MYE7")
    if mye7 == "revisions":
        ry = [year - 3, year - 2, year - 1]
        ws.append([f"MYE7: July {year + 1} revisions to mid-{ry[0]} to "
                   f"mid-{ry[-1]} estimates for England & Wales "])
        for t in ("This worksheet contains one table.", "Freeze panes.",
                  "This met my needs", "I need something slightly different",
                  "This is not what I need at all"):
            ws.append([t])
        head = ["Code", "Name", "Geography"]
        for y in ry:
            head += [f"Population mid-{y} (published July {year})",
                     f"Population mid-{y} (revised July {year + 1})",
                     f"Revision to Mid-{y}",
                     f"Percentage Revision to Mid-{y} "]
        ws.append(head)
        rev = revised or {}

        def c7(code, j):
            y, k = ry[j // 4], j % 4
            if code in ("K", "W", "R"):
                return 77
            if code == "E":
                new = total(rev.get((c, y), v4(c, y)) for c in codes)
                old = new - 50
            elif code.startswith("W"):
                new = old = welsh(code, y)
            else:
                new = rev.get((code, y), v4(code, y))
                old = new - 10
            return [old, new, new - old, "0.1%"][k]
        body(ws, c7, 12)
    elif mye7 == "other":
        ws.append(["MYE7: Research-based statistical measures of uncertainty "
                   "for local authority mid-year population estimates from "
                   f"2021 to {year} for England and Wales"])
        ws.append(["Code", "Name", f"{year} Mid-year estimate"])
        for t in ("MYE8", "MYE9"):
            wb.create_sheet(t)["A1"] = f"{t}: revisions"
    wb.save(path)
    wb.close()
    return path


def landing(*years, extra=(), title="Estimates of the population for "
            "England and Wales"):
    ds = [{"uri": ED_URI.format(y)} for y in years]
    ds += [{"uri": LANDING + "/" + seg} for seg in extra]
    return {"type": "dataset_landing_page", "uri": LANDING,
            "description": {"title": title,
                            "releaseDate": "2026-07-28T23:00:00.000Z"},
            "datasets": ds}


def edition(year, files=None, label=None, release=None):
    files = [f"mye{year % 100:02d}tablesew.xlsx"] if files is None else files
    return {"type": "dataset", "uri": ED_URI.format(year),
            "description": {
                "edition": label or (f"Mid-{year}: 2023 local authority "
                                     "boundaries"),
                "releaseDate": release or f"{year + 1}-07-28T23:00:00.000Z"},
            "downloads": [{"file": f} for f in files],
            "supplementaryFiles": [], "versions": []}


class _Resp:
    def __init__(self, body, url):
        self.content = body
        self.url = url
        self.text = body.decode("utf-8", "replace") if isinstance(
            body, bytes) else body

    def raise_for_status(self):
        pass


class _Session:
    def __init__(self, body, final):
        self.body, self.final = body, final

    def get(self, url, **kw):
        return _Resp(self.body, self.final)


class Tmp(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.n = 0

    def write(self, year=2025, **kw):
        self.n += 1
        return write_file(self.root / f"f{self.n}.xlsx", year, **kw)

    def read(self, year=2025, **kw):
        return m.read_file(self.write(year, **kw))

    def raises(self, *needles, **kw):
        path = self.write(**kw)
        with self.assertRaises(ValueError) as e:
            m.read_file(path)
        for n in needles:
            self.assertIn(n, str(e.exception))
        return str(e.exception)


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------

class Discovery(unittest.TestCase):

    REAL_EXTRA = ("mid2011tomid2025detailedtimeseries",
                  "englandandwalespopulationestimates1838to2025",
                  "mid20232023localauthorityboundarieseditionofthisdataset",
                  "mid20222023localauthorityboundaires",
                  "mid20222021localauthorityboundaires",
                  "regionalpopulationestimatesforenglandandwales1971to2022")

    def test_the_newest_edition_is_chosen(self):
        lp = landing(2024, 2025, extra=self.REAL_EXTRA)
        self.assertEqual(m.editions_listed(lp),
                         [(2025, 2023, ED_URI.format(2025)),
                          (2024, 2023, ED_URI.format(2024))])
        self.assertEqual(m.latest_edition(lp), (2025, 2023,
                                                ED_URI.format(2025)))
        self.assertEqual(m.edition_for_year(lp, 2024)[2],
                         ED_URI.format(2024))
        self.assertEqual(m.newest_listed(lp), 2025)

    def test_older_irregular_uris_are_listed_and_ignored(self):
        lp = landing(2024, 2025, extra=self.REAL_EXTRA)
        seen = m.irregular_editions(lp)
        self.assertIn((LANDING + "/mid20222023localauthorityboundaires",
                       [2022]), seen)
        self.assertEqual(m.latest_edition(lp)[0], 2025)

    def test_an_irregular_uri_naming_a_newer_mid_year_halts(self):
        for seg in ("mid20262023localauthorityboundaires",
                    "mid20262023localauthorityboundarieseditionofthisdataset",
                    "mid2011tomid2026detailedtimeseries"):
            with self.subTest(seg=seg):
                lp = landing(2024, 2025, extra=(seg,))
                with self.assertRaises(ValueError) as e:
                    m.latest_edition(lp)
                self.assertIn(seg, str(e.exception))
                self.assertIn("mid-2025", str(e.exception))

    def test_none_halts_listing_what_was_seen(self):
        lp = landing(extra=("ukpopulationestimates1838to2022",))
        with self.assertRaises(ValueError) as e:
            m.latest_edition(lp)
        self.assertIn("ukpopulationestimates1838to2022", str(e.exception))

    def test_two_editions_for_one_year_halt(self):
        lp = landing(2025)
        lp["datasets"].append(
            {"uri": LANDING + "/mid20252024localauthorityboundaries"})
        with self.assertRaises(ValueError) as e:
            m.latest_edition(lp)
        self.assertIn("refusing", str(e.exception))
        with self.assertRaises(ValueError):
            m.edition_for_year(lp, 2025)
        with self.assertRaises(ValueError) as e:
            m.edition_for_year(landing(2025), 2023)
        self.assertIn("0 editions", str(e.exception))

    def test_landing_title_checked(self):
        m.check_landing(landing(2025))
        with self.assertRaises(ValueError) as e:
            m.check_landing(landing(2025, title="Population estimates"))
        self.assertIn("Population estimates", str(e.exception))

    def test_edition_download_exactly_one(self):
        self.assertEqual(m.edition_download(edition(2025)),
                         ("mye25tablesew.xlsx",
                          "Mid-2025: 2023 local authority boundaries",
                          "2026-07-28T23:00:00.000Z"))
        for files in ([], ["a.xlsx", "b.xlsx"]):
            with self.subTest(files=files), self.assertRaises(ValueError) as e:
                m.edition_download(edition(2025, files=files))
            self.assertIn("download", str(e.exception))
        with self.assertRaises(ValueError) as e:
            m.edition_download(edition(2025, files=["mye25tablesew.csv"]))
        self.assertIn(".xlsx", str(e.exception))
        with self.assertRaises(ValueError) as e:
            m.edition_download(edition(2025, label="Mid-2025 tables"))
        self.assertIn("Mid-2025 tables", str(e.exception))

    def test_download_url(self):
        self.assertEqual(
            m.download_url(ED_URI.format(2025), "mye25tablesew.xlsx"),
            "https://www.ons.gov.uk/file?uri=" + ED_URI.format(2025)
            + "/mye25tablesew.xlsx")

    def test_discovery_note_names_the_newest_and_warns_when_overdue(self):
        lp = landing(2024, 2025)
        lines = m.discovery_note(lp, ["2024", "2025"], today=date(2027, 9, 30))
        self.assertIn("mid-2025", lines[0])
        self.assertEqual(len(lines), 1)
        lines = m.discovery_note(lp, ["2024", "2025"], today=date(2027, 10, 1))
        self.assertIn("WARNING", lines[-1])
        self.assertIn("mid-2026", lines[-1])
        self.assertEqual(m.discovery_note(lp, ["2024"]), [])
        self.assertEqual(m.expected_next_release(2025),
                         (2026, date(2027, 9, 30)))


# ---------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------

class Cells(unittest.TestCase):

    def test_a_positive_integer_passes(self):
        self.assertEqual(m.la_cell(539146, "x"), 539146)
        self.assertEqual(m.la_cell(539146.0, "x"), 539146)

    def test_blank_text_zero_negative_and_others_halt_naming_the_cell(self):
        for v, needle in ((None, "blank"), ("", "blank"), ("  ", "blank"),
                          ("x", "text"), ("[x]", "text"), ("1,000", "text"),
                          (0, "zero"), (-5, "negative"), (1.5, "non-integer"),
                          (float("nan"), "blank"), (True, "not a")):
            with self.subTest(v=v), self.assertRaises(ValueError) as e:
                m.la_cell(v, "MYE2 row 12 E06000047")
            self.assertIn("MYE2 row 12 E06000047", str(e.exception))
            self.assertIn(needle, str(e.exception))


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------

class Read(Tmp):

    def test_the_july_2026_layout(self):
        f = self.read()
        self.assertEqual(f["year"], 2025)
        self.assertEqual(f["published"], date(2026, 7, 29))
        self.assertEqual(f["rank"], (date(2026, 7, 29), 2025))
        self.assertEqual(f["next_release"], "summer 2027")
        self.assertIn("boundaries in place as of April 2024",
                      f["geography_note"])
        self.assertEqual(len(f["revision_notes"]), 1)
        self.assertTrue(f["revision_notes"][0].startswith(
            "Revisions to mid-2022 to mid-2024"))
        self.assertEqual(sorted(f["mye2"]["las"]), sorted(CODES))
        self.assertEqual(f["mye2"]["wales"], 2)
        self.assertIn("E12000001", f["mye2"]["aggregates"])
        self.assertIn("E11000001", f["mye2"]["aggregates"])
        self.assertEqual(f["mye4"]["years"], list(range(2025, 2010, -1)))
        self.assertEqual(f["mye7"]["layout"], "revisions")
        self.assertEqual(f["mye7"]["label"], "July 2026")
        self.assertEqual(f["mye7"]["years"], [2022, 2023, 2024])
        self.assertEqual(f["mye7"]["las"]["E06000002"][2024],
                         (value("E06000002", 2024) - 10,
                          value("E06000002", 2024)))

    def test_the_july_2025_layout_is_read_naming_what_it_saw(self):
        f = self.read(2024, mye7="other")
        self.assertEqual(f["year"], 2024)
        self.assertEqual(f["mye7"]["layout"], "other")
        self.assertIn("Research-based statistical measures of uncertainty",
                      f["mye7"]["title"])
        self.assertTrue(any("not a revisions table" in n
                            for n in f["layout_notes"]))
        self.assertEqual(m.reconcile(f), [])

    def test_cover_mye2_and_mye4_disagreeing_on_the_year_halt(self):
        self.raises("June 2024", "mid-2025",
                    cover={"title_year": 2024})
        self.raises("MYE2", "mid-2024", mye2_title_year=2024)
        self.raises("MYE4", "mid-2024", mye4_title_year=2024)

    def test_the_cover_lines_are_required_once(self):
        self.raises("originally published", cover={"pub_line": False})
        self.raises("next publication", cover={"next_line": False})
        self.raises("29 July 2026", "1 August 2026",
                    cover={"date_text": "Date published: 1 August 2026"})

    def test_a_changed_header_halts(self):
        self.raises("MYE2", "All persons",
                    mye2_header=["Code", "Name", "Geography", "All persons"])
        hdr = ["Code", "Name", "Geography"] + [
            f"Mid-{y}" for y in range(2025, 2011, -1)]   # mid-2011 missing
        self.raises("MYE4", "Mid-2011", mye4_header=hdr)
        hdr = ["Code", "Name", "Area"] + [
            f"Mid-{y}" for y in range(2025, 2010, -1)]
        self.raises("MYE4", "Area", mye4_header=hdr)

    def test_an_unexpected_row_halts(self):
        self.raises("S12000033", extra_rows=[("S12000033", "Aberdeen City",
                                              "Council Area", 1, 2, 3)])

    def test_a_repeated_authority_halts(self):
        self.raises("E06000002", "twice",
                    extra_rows=[("E06000002", "Middlesbrough",
                                 "Unitary Authority", 5, 6, 7)])

    def test_a_missing_geography_note_halts(self):
        self.raises("Geography:", notes=[(1, "Rounding: estimates are "
                                             "presented unrounded.")])

    def test_not_a_workbook_halts(self):
        p = self.root / "x.xlsx"
        p.write_bytes(b"not a zip")
        with self.assertRaises(ValueError) as e:
            m.read_file(p)
        self.assertIn("x.xlsx", str(e.exception))

    def test_the_rank_comes_from_the_cover_never_the_name_or_api_date(self):
        path = self.write(2025, published=date(2026, 7, 29))
        named = self.root / "mye26tablesew_published_2027-01-01.xlsx"
        named.write_bytes(path.read_bytes())
        f = m.read_file(named)
        self.assertEqual(m.release_rank(f), (date(2026, 7, 29), 2025))
        # the API's releaseDate (the evening before) is never a rank
        self.assertIsNone(m.release_rank("2026-07-28T23:00:00.000Z"))
        self.assertIsNone(m.release_rank(
            "https://www.ons.gov.uk/file?uri=/x/mid20252023local"
            "authorityboundaries/mye25tablesew.xlsx"))
        src = m.edition_source(f, {"file_name": named.name,
                                   "where": "data/raw/x.xlsx"})
        self.assertEqual(m.release_rank(src), (date(2026, 7, 29), 2025))


# ---------------------------------------------------------------------------
# Reconciliation and records
# ---------------------------------------------------------------------------

class Records(Tmp):

    def test_reconciliation_exact_and_off_by_one(self):
        f = self.read()
        self.assertEqual(m.reconcile(f, [2025, 2024]), [])
        total = sum(value(c, 2025) for c in CODES)
        f = self.read(england={("mye2", 2025): total + 1})
        probs = m.reconcile(f)
        self.assertTrue(any("MYE2" in p and f"{total:,}" in p
                            for p in probs), probs)
        f = self.read(england={("mye4", 2024): 1})
        self.assertEqual(m.reconcile(f, [2025]), [])
        self.assertTrue(m.reconcile(f, [2025, 2024]))

    def test_mye7_revised_must_equal_mye4(self):
        f = self.read(revised={("E06000002", 2024): 1})
        probs = m.reconcile(f, [2025, 2024])
        self.assertTrue(any("MYE7" in p and "E06000002" in p
                            for p in probs), probs)
        self.assertEqual(m.reconcile(f, [2025]), [])

    def test_a_blank_text_zero_or_negative_la_cell_halts(self):
        for cell, needle in ((None, "blank"), ("[x]", "text"), (0, "zero"),
                             (-3, "negative")):
            with self.subTest(cell=cell):
                f = self.read(mye2={"E06000002": cell})
                self.assertTrue(any(needle in p for p in m.reconcile(f)))
                with self.assertRaises(ValueError) as e:
                    m.build_records(f, 2025, lambda c: c)
                self.assertIn("E06000002", str(e.exception))
                self.assertIn(needle, str(e.exception))

    def test_a_welsh_row_is_counted_and_skipped(self):
        f = self.read()
        recs = m.build_records(f, 2025, lambda c: c)
        self.assertEqual([r["lad24cd"] for r in recs], sorted(CODES))
        self.assertEqual(f["mye4"]["wales"], 2)

    def test_newest_year_from_mye2_back_year_from_mye4(self):
        f = self.read()
        new = m.build_records(f, 2025, lambda c: c)
        back = m.build_records(f, 2024, lambda c: c)
        by = {r["lad24cd"]: r for r in new}
        self.assertEqual(by["E06000002"], {
            "lad24cd": "E06000002", "reference_year": "2025",
            "population": value("E06000002", 2025),
            "la_name": "Middlesbrough"})
        self.assertEqual({r["lad24cd"]: r["population"] for r in back},
                         {c: value(c, 2024) for c in CODES})
        self.assertTrue(all(r["reference_year"] == "2024" for r in back))

    def test_mye2_and_mye4_disagreeing_on_the_newest_year_halts(self):
        f = self.read(mye2={"E06000002": value("E06000002", 2025) + 1})
        with self.assertRaises(ValueError) as e:
            m.build_records(f, 2025, lambda c: c)
        self.assertIn("MYE2", str(e.exception))
        self.assertIn("MYE4", str(e.exception))
        self.assertTrue(any("MYE2" in p and "MYE4" in p
                            for p in m.reconcile(f)))

    def test_a_year_the_file_does_not_state_raises(self):
        f = self.read()
        with self.assertRaises(ValueError):
            m.build_records(f, 2026, lambda c: c)
        with self.assertRaises(ValueError):
            m.build_records(f, 2010, lambda c: c)

    def test_duplicate_key_after_resolution_raises(self):
        f = self.read()
        with self.assertRaises(ValueError) as e:
            m.build_records(f, 2025, lambda c: "E06000001")
        self.assertIn("duplicate", str(e.exception))

    def test_periods_stated_include_held_back_years(self):
        f = self.read()
        self.assertEqual(m.periods_stated(f, []), [2025])
        self.assertEqual(m.periods_stated(f, ["2024", "2025"]), [2024, 2025])
        self.assertEqual(m.periods_stated(f, ["2010", "2026"]), [2025])
        self.assertEqual(m.periods_stated(f, ["2019"]), [2019, 2025])

    def test_rows_content_sha_covers_key_population_and_name_only(self):
        a = [{"lad24cd": "E06000001", "reference_year": "2025",
              "population": 5, "la_name": "Hartlepool", "where": "x"}]
        b = [dict(a[0], where="y", reference_year="2024")]
        self.assertEqual(m.rows_content_sha(a), m.rows_content_sha(b))
        self.assertNotEqual(m.rows_content_sha(a), m.rows_content_sha(
            [dict(a[0], la_name="Hartlepool UA")]))
        self.assertNotEqual(m.rows_content_sha(a), m.rows_content_sha(
            [dict(a[0], population=6)]))
        self.assertEqual(m.rows_content_sha(a), hashlib.sha256(
            b"E06000001|5|Hartlepool").hexdigest())


class Geography(unittest.TestCase):

    def test_old_codes_pass_and_the_new_codes_halt(self):
        cur = mock.MagicMock()
        cur.fetchall.side_effect = [
            [("E08000038", "E08000016"), ("E08000039", "E08000019")],
            [(c,) for c in CODES], [(c,) for c in CODES]]
        rmap, probs = m.resolve_codes(cur, set(CODES), 2025)
        self.assertEqual(probs, [])
        self.assertEqual(rmap, {c: c for c in CODES})
        cur.fetchall.side_effect = [
            [("E08000038", "E08000016"), ("E08000039", "E08000019")],
            [(c,) for c in CODES], [(c,) for c in CODES]]
        codes = set(CODES) - {"E08000016"} | {"E08000038"}
        rmap, probs = m.resolve_codes(cur, codes, 2025)
        self.assertTrue(any("E08000038" in p and "'old'" in p
                            for p in probs), probs)

    def test_an_unknown_code_is_unexplained(self):
        cur = mock.MagicMock()
        cur.fetchall.side_effect = [
            [("E08000038", "E08000016"), ("E08000039", "E08000019")],
            [(c,) for c in CODES], [(c,) for c in CODES]]
        _, probs = m.resolve_codes(cur, set(CODES) | {"E07000999"}, 2025)
        self.assertTrue(any(p.startswith("UNEXPLAINED E07000999")
                            for p in probs), probs)


# ---------------------------------------------------------------------------
# Ranks, ledger and planning
# ---------------------------------------------------------------------------

R24 = (date(2025, 7, 30), 2024)
R25 = (date(2026, 7, 29), 2025)


class Ledger(unittest.TestCase):

    def test_ledger_source_round_trip(self):
        s = m.ledger_source("https://x/a.xlsx", R25, [2025, 2024])
        self.assertEqual(s, "https://x/a.xlsx (mid-2025; published "
                            "2026-07-29; periods 2024, 2025)")
        self.assertEqual(m.parse_ledger_source(s),
                         ("https://x/a.xlsx", R25, ["2024", "2025"]))
        self.assertEqual(m.release_rank(s), R25)

    def test_ledger_complete_needs_the_pair_for_every_period(self):
        s = m.ledger_source("u", R25, [2024, 2025])
        only25 = {"2025": {(s, "sha")}}
        self.assertIsNone(m.ledger_complete(only25, "u", "sha",
                                            {"2024": R24, "2025": R25}))
        both = {"2025": {(s, "sha")}, "2024": {(s, "sha")}}
        self.assertEqual(m.ledger_complete(both, "u", "sha",
                                           {"2024": R25, "2025": R25}),
                         ["2024", "2025"])
        self.assertIsNone(m.ledger_complete(both, "u", "other", {}))
        self.assertIsNone(m.ledger_complete(both, "v", "sha", {}))
        self.assertIsNone(m.ledger_complete(both, "u", "sha", {},
                                            stranded=["2024"]))

    def test_plan_new_revised_and_skipped(self):
        new, rev, skip = m.plan_periods(["2024", "2025"], R25, {"2024": R24},
                                        {}, "src", "sha", recheck=None,
                                        allow_older=False)
        self.assertEqual((new, rev, skip), (["2025"], ["2024"], {}))
        new, rev, skip = m.plan_periods(["2024"], R24, {"2024": R25},
                                        {}, "src", "sha", recheck=None,
                                        allow_older=False)
        self.assertEqual((new, rev), ([], []))
        self.assertTrue(skip["2024"].startswith("older"))
        new, rev, skip = m.plan_periods(["2024"], R24, {"2024": R25},
                                        {}, "src", "sha", recheck=None,
                                        allow_older=True)
        self.assertEqual(rev, ["2024"])
        new, rev, skip = m.plan_periods(["2023"], R24, {"2024": R24},
                                        {}, "src", "sha", recheck=None,
                                        allow_older=False)
        self.assertTrue(skip["2023"].startswith("older"))
        with self.assertRaises(ValueError):
            m.plan_periods(["2025"], R25, {}, {}, "s", "h", recheck="2025",
                           allow_older=False)

    def test_tip_ranks_take_the_newest_file_seen(self):
        s = m.ledger_source("u", R25, [2024, 2025])
        self.assertEqual(m.tip_ranks({"2024": {"rank": R24}},
                                     {"2024": {(s, "x")}}), {"2024": R25})


# ---------------------------------------------------------------------------
# Stop conditions
# ---------------------------------------------------------------------------

def recs(pops, names=None, year="2025"):
    return [{"lad24cd": k, "reference_year": year, "population": v,
             "la_name": (names or {}).get(k, NAMES.get(k, k))}
            for k, v in pops.items()]


class Stops(unittest.TestCase):

    def setUp(self):
        p = mock.patch.object(m, "EXPECTED_AREAS", 3)
        p.start()
        self.addCleanup(p.stop)
        self.base = {"E06000001": 100000, "E06000002": 200000,
                     "E09000033": 300000}

    def kinds(self, probs):
        return [p.split(":")[0] if p.startswith((m.PARTIAL, m.LARGE))
                else "soft" for p in probs]

    def test_new_year_national_limit_both_sides(self):
        prev = recs(self.base, year="2024")
        total = sum(self.base.values())
        lim = m.NEW_TOTAL_BP
        # move the total to exactly the limit, then one person over
        at = total * lim // 10000
        inside = dict(self.base, E09000033=300000 + at)
        self.assertEqual(m.period_problems(recs(inside), None, prev,
                                           kind="new"), [])
        over = dict(self.base, E09000033=300000 + at + 1)
        probs = m.period_problems(recs(over), None, prev, kind="new")
        self.assertTrue(any("national" in p for p in probs), probs)
        self.assertEqual(self.kinds(probs), ["soft"] * len(probs))

    def test_new_year_authority_limit_both_sides(self):
        prev = recs(dict(self.base, E09000033=10**8), year="2024")
        at = 100000 * m.NEW_AREA_BP // 10000
        ok = dict(self.base, E06000001=100000 + at, E09000033=10**8)
        self.assertEqual(m.period_problems(recs(ok), None, prev,
                                           kind="new"), [])
        bad = dict(ok, E06000001=100000 + at + 1)
        probs = m.period_problems(recs(bad), None, prev, kind="new")
        self.assertTrue(any("E06000001" in p for p in probs), probs)
        self.assertEqual(self.kinds(probs), ["soft"])

    def test_revised_year_any_change_needs_acknowledge(self):
        tip = recs(self.base)
        self.assertEqual(m.period_problems(recs(self.base), tip, None,
                                           kind="revised"), [])
        probs = m.period_problems(recs(dict(self.base, E06000002=200001)),
                                  tip, None, kind="revised")
        self.assertEqual(self.kinds(probs), ["soft"])
        self.assertIn("E06000002", probs[0])
        probs = m.period_problems(recs(self.base,
                                       names={"E06000002": "Mbro"}),
                                  tip, None, kind="revised")
        self.assertEqual(self.kinds(probs), ["soft"])
        self.assertIn("la_name", probs[0])

    def test_revised_year_authority_limit_both_sides(self):
        big = dict(self.base, E09000033=10**9)
        tip = recs(big)
        at = 200000 * m.REV_AREA_BP // 10000
        probs = m.period_problems(recs(dict(big, E06000002=200000 + at)),
                                  tip, None, kind="revised")
        self.assertEqual(self.kinds(probs), ["soft"])
        probs = m.period_problems(recs(dict(big, E06000002=200000 + at + 1)),
                                  tip, None, kind="revised")
        self.assertIn(m.LARGE, self.kinds(probs))
        self.assertTrue(any(p.startswith(m.LARGE) and "E06000002" in p
                            for p in probs))

    def test_revised_year_national_limit_both_sides(self):
        many = {f"E07{i:06d}": 100000 for i in range(1000)}
        with mock.patch.object(m, "EXPECTED_AREAS", 1000):
            tip = recs(many)
            total = 100000 * 1000
            at = total * m.REV_TOTAL_BP // 10000
            # spread the move over the authorities, each well under the
            # authority limit
            each = at // 1000
            inside = {k: v + each for k, v in many.items()}
            probs = m.period_problems(recs(inside), tip, None,
                                      kind="revised")
            self.assertNotIn(m.LARGE, self.kinds(probs))
            over = dict(inside)
            k0 = sorted(over)[0]
            over[k0] += at - each * 1000 + 1
            probs = m.period_problems(recs(over), tip, None, kind="revised")
            self.assertTrue(any(p.startswith(m.LARGE) and "national" in p
                                for p in probs), probs)

    def test_fewer_authorities_than_the_tip_is_rejected_always(self):
        tip = recs(self.base)
        fewer = dict(self.base)
        fewer.pop("E06000001")
        probs = m.period_problems(recs(fewer), tip, None, kind="revised")
        self.assertIn(m.PARTIAL, self.kinds(probs))
        probs = m.period_problems(recs(fewer), None, None, kind="new")
        self.assertIn(m.PARTIAL, self.kinds(probs))

    def test_the_calibration_constants_are_documented(self):
        doc = m.__doc__
        for needle in ("1.4008%", "6.674%", "City of London", "0.1007%",
                       "38.34%", "2.713%"):
            self.assertIn(needle, doc)


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

class Download(Tmp):

    def test_same_name_different_content_is_kept_beside(self):
        held = self.write(2025)
        new = self.write(2025, published=date(2026, 8, 5))
        raw = self.root / "raw"
        raw.mkdir()
        (raw / "mye25tablesew.xlsx").write_bytes(held.read_bytes())
        url = m.download_url(ED_URI.format(2025), "mye25tablesew.xlsx")
        with quiet() as out, mock.patch.object(m, "MIN_FILE_BYTES", 10):
            path, final = m.fetch(url, raw, session=_Session(
                new.read_bytes(), url))
        self.assertEqual(final, url)
        self.assertEqual((raw / "mye25tablesew.xlsx").read_bytes(),
                         held.read_bytes())
        self.assertNotEqual(path.name, "mye25tablesew.xlsx")
        self.assertTrue(path.name.startswith("mye25tablesew-"))
        self.assertEqual(path.read_bytes(), new.read_bytes())
        self.assertIn("kept it untouched", out.getvalue())
        # the same bytes again: kept as they are, still downloaded
        with quiet() as out, mock.patch.object(m, "MIN_FILE_BYTES", 10):
            path2, _ = m.fetch(url, raw, session=_Session(
                held.read_bytes(), url))
        self.assertEqual(path2.name, "mye25tablesew.xlsx")
        self.assertIn("same content", out.getvalue())

    def test_a_redirect_records_the_final_url_and_its_name(self):
        f = self.write(2025)
        raw = self.root / "raw2"
        final = "https://download.example/x/mye25tablesew.xlsx"
        with quiet(), mock.patch.object(m, "MIN_FILE_BYTES", 10):
            path, got = m.fetch("https://www.ons.gov.uk/file?uri=/a/b/"
                                "mye25tablesew.xlsx", raw,
                                session=_Session(f.read_bytes(), final))
        self.assertEqual(got, final)
        self.assertEqual(path.name, "mye25tablesew.xlsx")

    def test_small_or_not_a_workbook_halts(self):
        raw = self.root / "raw3"
        url = "https://www.ons.gov.uk/file?uri=/a/mye25tablesew.xlsx"
        with quiet(), self.assertRaises(SystemExit):
            m.fetch(url, raw, session=_Session(b"tiny", url))
        with quiet(), mock.patch.object(m, "MIN_FILE_BYTES", 2), \
                self.assertRaises(SystemExit):
            m.fetch(url, raw, session=_Session(b"not a workbook", url))
        self.assertEqual([p for p in raw.glob("*.xlsx")], [])


if __name__ == "__main__":
    unittest.main()
