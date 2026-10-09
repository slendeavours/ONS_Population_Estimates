"""Pure tests for s22_ctb_editions: discovery, cells, the CTB and Table 615
readers, reconciliation, records, ranks, ledger and planning, stop
conditions, the 615-to-CTB cross check and the download rules.

No database and no network: every workbook is written by the tests (openpyxl
for the CTB layout, pandas + odfpy for Table 615) and the content API replies
and downloads are stubbed. The fixture writers are reused by
test_s22_ctb_loader.
"""
import contextlib
import hashlib
import io
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import s22_ctb_editions as m  # noqa: E402

CODES = ["E06000001", "E06000002", "E06000003", "E07000096", "E08000038"]
NAMES = {"E06000001": "Hartlepool", "E06000002": "Middlesbrough",
         "E06000003": "Redcar and Cleveland", "E07000096": "Dacorum",
         "E08000038": "Barnsley", "E08000016": "Barnsley",
         "E07000004": "Aylesbury Vale"}
PUB_FIRST = ("Published by the Ministry of Housing, Communities and Local "
             "Government (MHCLG) on 6 November 2025")
PUB_REVISED = PUB_FIRST + " (originally) and revised on 21 January 2026"
BLOCK_TITLES = {
    "1.01": "Table 1.01 Total Number of Dwellings on Valuation List (Line 1)",
    "1.02": "Table 1.02. Number of dwellings on valuation list exempt on 6 "
            "October 2025 (Line 2)",
    "1.11": "Table 1.11. Number of dwellings in line 7 classed as second "
            "homes on 6 October 2025 (Line 11) ",
    "1.17": "Table 1.17. Number of dwellings in line 7 classed as empty and "
            "being charged the Empty Homes Premium on 6 October 2025 (Line "
            "14)",
    "1.18": "Table 1.18. Total number of dwellings in line 7 classed as empty "
            "on 6 October 2025 (Line 15)",
    "1.19": "Table 1.19. Number of dwellings that are classed as empty on 6 "
            "October 2025 and have been for more than 6 months (Line 16)",
    "1.22": "Table 1.22. Number of dwellings that are classed as empty and "
            "have been empty for more than 6 months excluding those (Line 18)",
}
FIELD_OF = {"1.01": "total_dwellings", "1.11": "second_homes",
            "1.17": "empty_homes_premium_count", "1.18": "empty_total",
            "1.19": "empty_6_months_plus"}
ALL_CLASSES = [chr(c) for c in range(ord("A"), ord("W") + 1)]


def quiet():
    return contextlib.redirect_stdout(io.StringIO())


def ctb_values(codes, bump=0):
    """{code: {field: v, 'class X': v}} of plausible published values."""
    out = {}
    for i, c in enumerate(codes):
        v = {"total_dwellings": 40000 + 1000 * i + bump,
             "second_homes": 100 + i, "empty_homes_premium_count": 50 + i,
             "empty_total": 1200 + 10 * i, "empty_6_months_plus": 500 + i}
        for j, k in enumerate(ALL_CLASSES):
            v[f"class {k}"] = (None if k in ("A", "C") else 10 * i + j)
        out[c] = v
    return out


def write_ctb(path, *, year=2025, codes=None, values=None, title=None,
              pub=PUB_REVISED, titles=None, drop_class=None, england=None,
              extra_rows=(), names=None):
    """A CTB local authority level workbook in the published layout:
    Cover (title, publication line), Council Taxbase Data and Supplementary
    Data (labels row 5, headers row 6, England row 7, authorities from row
    8). values: ctb_values overrides {code: {key: v}}; a value may be '[x]'
    and so on. England is the sum of the published authority values unless
    england {key: v} overrides."""
    import openpyxl
    codes = list(codes or CODES)
    vals = ctb_values(codes)
    for c, over in (values or {}).items():
        vals.setdefault(c, {}).update(over)
    titles = dict(BLOCK_TITLES, **(titles or {}))
    names = names or NAMES
    wb = openpyxl.Workbook()
    cov = wb.active
    cov.title = "Cover"
    cov.append([title or f"Council Taxbase: Local Authority Level Data for "
                f"{year}"])
    cov.append([pub])
    cov.append(["Purpose"])

    def eng(key):
        if england and key in england:
            return england[key]
        nums = [vals[c][key] for c in codes
                if isinstance(vals[c].get(key), int)]
        return sum(nums)

    ws = wb.create_sheet("Council Taxbase Data")
    label = [None] * 5
    header = ["E-code", "ONS Code", "Region ", "Local Authority", "Notes"]
    cols = {}
    for number in ("1.01", "1.02", "1.11", "1.17", "1.18", "1.19", "1.22"):
        label += [titles[number], None, None]
        header += ["Band A", "Total", None]
        cols[number] = len(header) - 2
    rows = [["Council Taxbase Data"], ["text"], ["text"], ["text"], label,
            header]

    def main_row(ecode, code, name, get):
        r = [ecode, code, "R", name, None] + [None] * (len(header) - 5)
        for number, col in cols.items():
            f = FIELD_OF.get(number)
            v = get(f) if f else 7
            r[col - 1] = v
            r[col] = v
        return r
    rows.append(main_row("ENG", "E92000001", "England", eng))
    for c in codes:
        rows.append(main_row("E1", c, names.get(c, c),
                             lambda f, c=c: vals[c][f]))
    for r in extra_rows:
        rows.append(r)
    for r in rows:
        ws.append(r)
    ss = wb.create_sheet("Supplementary Data")
    classes = [k for k in ALL_CLASSES if k != drop_class]
    slabel = [None] * 5 + [
        "Table 2.01. Number of dwellings on the Valuation List on 10 "
        "September 2025 that were in exempt classes B, D to W"] + \
        [None] * len(classes) + [None, "Table 2.02. Reduced Council Tax "
                                       "discount granted", None]
    sheader = ["E-code", "ONS Code", "Region ", "Local Authority", "Notes"]
    for k in classes:
        sheader.append(f"Class {k}" + (" - category not in use"
                                       if k in ("A", "C") else
                                       (" " if k == "G" else "")))
    sheader += ["Total exemptions", None, "Has local authority used power",
                "Plan"]
    srows = [["Supplementary Data"], ["t"], ["t"], ["t"], slabel, sheader]

    def supp_row(ecode, code, name, get):
        r = [ecode, code, "R", name, None]
        for k in classes:
            r.append(get(f"class {k}"))
        r += [1, None, "Yes", "No"]
        return r
    srows.append(supp_row("ENG", "E92000001", "England", eng))
    for c in codes:
        srows.append(supp_row("E1", c, names.get(c, c),
                              lambda k, c=c: vals[c][k]))
    for r in srows:
        ss.append(r)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


T615_TITLE = ("Live Table 615: Vacant Dwellings by Local Authority District, "
              "England, from 2004 to {end}")
DATES = {2004: "01/11/2004", 2019: "07/10/2019", 2020: "05/10/2020",
         2021: "04/10/2021", 2022: "03/10/2022", 2023: "02/10/2023",
         2024: "07/10/2024", 2025: "06/10/2025", 2026: "05/10/2026"}
T615_CODES = ["E06000001", "E06000002", "E07000096", "E08000016",
              "E08000038"]


def t615_value(sheet, code, year):
    """The default cell: Barnsley on E08000016 to 2024 and E08000038 from
    2025, as published."""
    if code == "E08000016" and year >= 2025:
        return "[x]"
    if code == "E08000038" and year < 2025:
        return "[x]"
    i = T615_CODES.index(code) if code in T615_CODES else 9
    base = 300 + 20 * i + (year - 2000)
    return base if sheet == "All_vacants" else base // 3


def t615_rows(years=(2023, 2024, 2025), *, codes=None, values=None,
              dates=None, title=None, latest="25 June 2026", england=None,
              end=None, names=None):
    """(cover rows, {sheet: rows}) in the published layout: header row 3
    'ONS code', 'Area', one snapshot date per year; England, one region,
    then the districts. values {(sheet, code, year): v} override the
    defaults; England is the sum of the published district values unless
    england {(sheet, year): v} overrides."""
    codes = list(codes or T615_CODES)
    names = names or NAMES
    end = end or years[-1]
    hdr = dates or [DATES.get(y, f"01/10/{y}") for y in years]
    cover = [[(title or T615_TITLE).format(end=end)], ["Background"],
             ["Symbols Used"], ["[x] = not available, [z] = not applicable, "
                                "[i] = imputed"], ["Latest Update"], [latest],
             ["Next Update"], ["November 2026 to February 2027"]]
    sheets = {}
    for sheet, _ in m.T615_SHEETS:
        def v(code, y, sheet=sheet):
            key = (sheet, code, y)
            if values and key in values:
                return values[key]
            return t615_value(sheet, code, y)
        rows = [[f"{sheet}: table"] + [None] * (len(years) + 1),
                ["This worksheet contains one table."]
                + [None] * (len(years) + 1),
                ["ONS code", "Area"] + list(hdr)]
        eng = []
        for y in years:
            if england and (sheet, y) in england:
                eng.append(england[(sheet, y)])
            else:
                eng.append(sum(x for x in (v(c, y) for c in codes)
                               if isinstance(x, (int, float))))
        rows.append(["E92000001", "England"] + eng)
        rows.append(["E12000001", "North East"] + [1] * len(years))
        for c in codes:
            rows.append([c, names.get(c, c)] + [v(c, y) for y in years])
        sheets[sheet] = rows
    return cover, sheets


def write_615(path, years=(2023, 2024, 2025), **kw):
    """t615_rows written as an .ods file."""
    import pandas as pd
    cover, sheets = t615_rows(years, **kw)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="odf") as w:
        pd.DataFrame(cover).to_excel(w, sheet_name="Cover", header=False,
                                     index=False)
        pd.DataFrame([["Contents"]]).to_excel(w, sheet_name="Contents",
                                              header=False, index=False)
        for sheet, rows in sheets.items():
            pd.DataFrame(rows).to_excel(w, sheet_name=sheet, header=False,
                                        index=False)
    return path


def collection(*titles):
    return {"links": {"documents": [
        {"title": t, "base_path": "/government/statistics/" +
         t.lower().replace(" ", "-").replace("(", "").replace(")", "")}
        for t in titles]}}


def release_page(year, attachments, *, title=None, history=()):
    return {"title": title or f"Council Taxbase {year} in England",
            "first_published_at": f"{year}-11-06T00:00:00+00:00",
            "public_updated_at": f"{year + 1}-01-21T09:30:00+00:00",
            "details": {"attachments": [
                {"title": t, "url": u} for t, u in attachments],
                "change_history": [{"public_timestamp": ts, "note": n}
                                   for ts, n in history]}}


def tables_page(attachments):
    return {"title": "Live tables on dwelling stock (including vacants)",
            "public_updated_at": "2026-06-25T09:30:03+01:00",
            "details": {"attachments": [{"title": t, "url": u}
                                        for t, u in attachments]}}


class Tmp(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)


class Discovery(unittest.TestCase):

    def test_newest_release_chosen_and_an_older_revised_duplicate_ignored(self):
        c = collection("Council Taxbase 2024 in England",
                       "Council Taxbase 2025 in England",
                       "Council Taxbase 2010 England (revised)",
                       "Council Taxbase 2010 England", "Something else")
        y, base = m.latest_release(c)
        self.assertEqual(y, 2025)
        self.assertIn("2025", base)
        self.assertEqual(m.release_for_year(c, 2024)[0], 2024)
        with self.assertRaises(ValueError):
            m.release_for_year(c, 2010)          # two documents for 2010

    def test_a_revised_duplicate_of_the_newest_year_raises(self):
        c = collection("Council Taxbase 2025 in England",
                       "Council Taxbase 2025 England (revised)")
        with self.assertRaisesRegex(ValueError, "newest year 2025"):
            m.latest_release(c)
        with self.assertRaises(ValueError):
            m.latest_release(collection("Unrelated"))

    def test_la_attachment_exactly_one_with_or_without_revised(self):
        u = "https://assets.example/media/1/x.xlsx"
        a = m.la_attachment(release_page(2025, [
            ("Tables 1 to 5: Council Taxbase in England 2025 (revised)", "t"),
            ("Council Taxbase: Local authority level data for 2025 (revised)",
             u)]))
        self.assertEqual((a["year"], a["revised"], a["url"]), (2025, True, u))
        self.assertFalse(m.la_attachment(release_page(2025, [
            ("Council Taxbase: Local authority level data for 2025", u)]))
            ["revised"])

    def test_two_la_level_attachments_raise_listing_titles(self):
        page = release_page(2025, [
            ("Council Taxbase: Local authority level data for 2025", "a"),
            ("Council Taxbase: Local authority level data for 2025 (revised)",
             "b")])
        with self.assertRaisesRegex(ValueError, "2 attachments.*revised"):
            m.la_attachment(page)
        with self.assertRaisesRegex(ValueError, "0 attachments"):
            m.la_attachment(release_page(2025, [("Technical notes", "x")]))

    def test_table_615_attachment(self):
        page = tables_page([("Table 100: dwellings", "a"),
                            ("Table 615: vacant dwellings by local authority "
                             "district: England, from 2004", "b")])
        self.assertEqual(m.table_615_attachment(page)["url"], "b")
        with self.assertRaises(ValueError):
            m.table_615_attachment(tables_page([("Table 100", "a")]))
        with self.assertRaises(ValueError):
            m.table_615_attachment(tables_page([("Table 615: a", "a"),
                                                ("Table 615: b", "b")]))

    def test_cover_year_differs_from_the_page_year(self):
        f = {"year": 2025}
        self.assertEqual(m.ctb_identity_problems(f, 2025, 2025), [])
        self.assertTrue(m.ctb_identity_problems(f, 2024))
        self.assertTrue(m.ctb_identity_problems(f, None, 2024))
        self.assertEqual(m.page_year(release_page(2025, [])), 2025)


class Cells(unittest.TestCase):

    def test_markers_give_null_with_reasons(self):
        self.assertEqual(m.cell("[x]"), (None, "suppressed_or_not_available"))
        self.assertEqual(m.cell(" [z] "), (None, "not_applicable"))
        self.assertEqual(m.cell(0), (0, None))
        self.assertEqual(m.cell(12.0), (12, None))

    def test_blank_imputed_unknown_negative_and_non_integer_halt(self):
        for v in (None, "", float("nan"), "[i]", "[c]", "n/a", -1, 3.5,
                  True):
            with self.subTest(v=v), self.assertRaises(ValueError):
                m.cell(v)

    def test_half_up(self):
        self.assertEqual(m.half_up(1414.5), 1415)       # banker's gives 1414
        self.assertEqual(m.half_up(1415.578), 1416)
        self.assertEqual(m.half_up(494.578), 495)


class ReadCTB(Tmp):

    def test_reads_cover_blocks_and_rows(self):
        f = m.read_ctb(write_ctb(self.root / "a.xlsx"))
        self.assertEqual((f["year"], f["first_published"], f["revised"],
                          f["rank"]),
                         (2025, date(2025, 11, 6), date(2026, 1, 21),
                          date(2026, 1, 21)))
        self.assertEqual(sorted(f["rows"]), sorted(CODES))
        self.assertEqual(f["rows"]["E06000001"]["values"]["total_dwellings"],
                         (40000, None))
        self.assertEqual(f["rows"]["E06000002"]["classes"]["Q"], (26, None))
        self.assertEqual(set(f["blocks"]), set(FIELD_OF) | {"2.01"})
        self.assertEqual(m.reconcile(f), [])

    def test_first_published_only_ranks_by_it(self):
        f = m.read_ctb(write_ctb(self.root / "a.xlsx", pub=PUB_FIRST))
        self.assertEqual((f["revised"], f["rank"]), (None, date(2025, 11, 6)))

    def test_cover_title_must_match(self):
        p = write_ctb(self.root / "a.xlsx", title="Council Taxbase tables")
        with self.assertRaisesRegex(ValueError, "cover title"):
            m.read_ctb(p)
        p = write_ctb(self.root / "b.xlsx", pub="Contact us")
        with self.assertRaisesRegex(ValueError, "publication line"):
            m.read_ctb(p)

    def test_a_block_found_by_number_with_the_wrong_title_raises(self):
        p = write_ctb(self.root / "a.xlsx", titles={
            "1.19": "Table 1.19. Number of dwellings in line 7 classed as "
                    "second homes (Line 16)"})
        with self.assertRaisesRegex(ValueError, "Table 1.19 is titled"):
            m.read_ctb(p)
        p = write_ctb(self.root / "b.xlsx", titles={
            "1.18": "Table 1.81. Total number of dwellings classed as empty"})
        with self.assertRaisesRegex(ValueError, "0 blocks numbered Table 1.18"):
            m.read_ctb(p)

    def test_a_missing_class_raises(self):
        p = write_ctb(self.root / "a.xlsx", drop_class="K")
        with self.assertRaisesRegex(ValueError, r"lacks exemption classes \['K'\]"):
            m.read_ctb(p)

    def test_blank_unknown_negative_and_non_integer_cells_halt(self):
        for v in (None, "[c]", -3, 2.5, "[i]"):
            with self.subTest(v=v):
                p = write_ctb(self.root / f"a{abs(hash(str(v)))}.xlsx",
                              values={"E06000003": {"second_homes": v}})
                with self.assertRaisesRegex(ValueError, "E06000003 second_homes"):
                    m.read_ctb(p)

    def test_duplicate_and_unknown_codes_raise(self):
        p = write_ctb(self.root / "a.xlsx", codes=CODES + ["E06000001"])
        with self.assertRaisesRegex(ValueError, "duplicate code E06000001"):
            m.read_ctb(p)
        p = write_ctb(self.root / "b.xlsx",
                      extra_rows=[["X", "W06000001", "R", "Wales"]])
        with self.assertRaisesRegex(ValueError, "W06000001"):
            m.read_ctb(p)

    def test_england_not_equal_to_the_sum_halts(self):
        p = write_ctb(self.root / "a.xlsx",
                      england={"empty_total": 1, "class D": 2})
        bad = m.reconcile(m.read_ctb(p))
        self.assertEqual(len(bad), 2, bad)
        self.assertIn("empty_total", bad[0])
        self.assertIn("class D", bad[1])


def _resolve(c):
    return {"E08000038": "E08000016"}.get(c, c)


class BuildCTB(Tmp):

    def build(self, **kw):
        f = m.read_ctb(write_ctb(self.root / "a.xlsx", **kw))
        f["source"] = "https://assets.example/media/1/2025.xlsx"
        return m.build_ctb(f, _resolve)

    def test_rows_derived_columns_and_source(self):
        main, cls = self.build()
        self.assertEqual(len(main), 5)
        self.assertEqual(len(cls), 5 * 11)
        r = {x["lad24cd"]: x for x in main}["E06000002"]
        self.assertEqual(r["empty_under_6_months"], 1210 - 501)
        self.assertEqual(r["unoccupied_exemptions_total"],
                         sum(10 + j for j, k in enumerate(ALL_CLASSES)
                             if k in m.UNOCCUPIED_CLASSES))
        self.assertIsNone(r["null_reasons"])
        self.assertIn("E08000016", {x["lad24cd"] for x in main})
        self.assertEqual(r["source_publication"],
                         "Council Taxbase: Local Authority Level Data for "
                         "2025; first published 2025-11-06; revised "
                         "2026-01-21; https://assets.example/media/1/2025.xlsx")
        self.assertEqual(r["taxbase_year"], "2025")

    def test_x_and_z_give_null_with_their_reasons(self):
        main, cls = self.build(values={
            "E06000001": {"second_homes": "[x]"},
            "E06000002": {"empty_homes_premium_count": "[z]"}})
        r = {x["lad24cd"]: x for x in main}
        self.assertIsNone(r["E06000001"]["second_homes"])
        self.assertEqual(r["E06000001"]["null_reasons"],
                         "second_homes=suppressed_or_not_available")
        self.assertEqual(r["E06000002"]["null_reasons"],
                         "empty_homes_premium_count=not_applicable")

    def test_one_suppressed_class_makes_the_total_null_and_a_zero_stays(self):
        main, cls = self.build(values={"E06000003": {"class K": "[x]"},
                                       "E07000096": {"class Q": 0}})
        r = {x["lad24cd"]: x for x in main}
        self.assertIsNone(r["E06000003"]["unoccupied_exemptions_total"])
        self.assertEqual(r["E06000003"]["null_reasons"],
                         "unoccupied_exemptions_total=built_from_null")
        k = [c for c in cls if c["lad24cd"] == "E06000003"
             and c["exemption_class"] == "K"][0]
        self.assertEqual((k["dwellings"], k["null_reasons"]),
                         (None, "dwellings=suppressed_or_not_available"))
        q = [c for c in cls if c["lad24cd"] == "E07000096"
             and c["exemption_class"] == "Q"][0]
        self.assertEqual((q["dwellings"], q["null_reasons"]), (0, None))
        self.assertIsNotNone(r["E07000096"]["unoccupied_exemptions_total"])

    def test_empty_under_6_months_null_when_a_part_is_null(self):
        main, _ = self.build(values={
            "E06000001": {"empty_6_months_plus": "[x]"}})
        r = {x["lad24cd"]: x for x in main}["E06000001"]
        self.assertIsNone(r["empty_under_6_months"])
        self.assertEqual(r["null_reasons"],
                         "empty_6_months_plus=suppressed_or_not_available;"
                         "empty_under_6_months=built_from_null")

    def test_two_codes_on_one_key_raise(self):
        f = m.read_ctb(write_ctb(self.root / "a.xlsx",
                                 codes=CODES + ["E08000016"]))
        f["source"] = "u"
        with self.assertRaisesRegex(ValueError, "both resolve to E08000016"):
            m.build_ctb(f, _resolve)


class Read615(Tmp):

    def test_reads_rows_dates_and_rank(self):
        f = m.read_615(write_615(self.root / "t.ods"))
        self.assertEqual((f["end_year"], f["rank"], f["years"]),
                         (2025, date(2026, 6, 25), [2023, 2024, 2025]))
        self.assertEqual(f["dates"][2025], date(2025, 10, 6))
        self.assertEqual(f["rows"]["E06000001"]["vacant_dwellings"][2024],
                         (324, None))
        self.assertEqual(m.reconcile(f), [])

    def test_rank_from_either_date_style(self):
        cover, sheets = t615_rows(latest="January 2026")
        self.assertEqual(m.parse_615(cover, sheets)["rank"], date(2026, 1, 1))
        cover, sheets = t615_rows(latest="21 January 2026")
        self.assertEqual(m.parse_615(cover, sheets)["rank"],
                         date(2026, 1, 21))
        cover, sheets = t615_rows(latest="soon")
        with self.assertRaisesRegex(ValueError, "Latest Update"):
            m.parse_615(cover, sheets)

    def test_cover_title_and_its_end_year(self):
        cover, sheets = t615_rows(title="Table 615 something {end}")
        with self.assertRaisesRegex(ValueError, "cover title"):
            m.parse_615(cover, sheets)
        cover, sheets = t615_rows(end=2026)
        with self.assertRaisesRegex(ValueError, "last snapshot year"):
            m.parse_615(cover, sheets)

    def test_date_headers_out_of_order_or_two_in_one_year_halt(self):
        cover, sheets = t615_rows(dates=["07/10/2024", "02/10/2023",
                                         "06/10/2025"])
        with self.assertRaisesRegex(ValueError, "out of order"):
            m.parse_615(cover, sheets)
        cover, sheets = t615_rows(dates=["02/10/2023", "01/03/2023",
                                         "06/10/2025"])
        with self.assertRaisesRegex(ValueError, "two snapshot dates in one "
                                                "year"):
            m.parse_615(cover, sheets)
        cover, sheets = t615_rows(dates=["02/10/2023", "2024", "06/10/2025"])
        with self.assertRaisesRegex(ValueError, "dd/mm/yyyy"):
            m.parse_615(cover, sheets)

    def test_dacorum_2012_rounds_half_up_as_held_and_other_fractions_halt(self):
        years = (2011, 2012, 2013)
        vals = {("All_vacants", "E07000096", 2012): 1414.5,
                ("All_long_term_vacants", "E07000096", 2012): 494.578}
        cover, sheets = t615_rows(years, values=vals)
        f = m.parse_615(cover, sheets)
        self.assertEqual(f["rows"]["E07000096"]["vacant_dwellings"][2012],
                         (1415, None))
        self.assertEqual(
            f["rows"]["E07000096"]["long_term_vacant_dwellings"][2012],
            (495, None))
        self.assertEqual(len(f["non_integer"]), 2)
        self.assertEqual(m.reconcile(f), [])     # unrounded against England
        cover, sheets = t615_rows(years, values={
            ("All_vacants", "E06000001", 2012): 10.5})
        with self.assertRaisesRegex(ValueError, "non-integer"):
            m.parse_615(cover, sheets)
        cover, sheets = t615_rows(years, values={
            ("All_vacants", "E07000096", 2013): 10.5})
        with self.assertRaisesRegex(ValueError, "non-integer"):
            m.parse_615(cover, sheets)

    def test_markers_blanks_and_imputed(self):
        for v in (None, "[i]", "[w]", -2):
            with self.subTest(v=v):
                cover, sheets = t615_rows(values={
                    ("All_vacants", "E06000002", 2024): v})
                with self.assertRaises(ValueError):
                    m.parse_615(cover, sheets)

    def test_rows_stored_only_where_a_sheet_has_a_number(self):
        cover, sheets = t615_rows(values={
            ("All_vacants", "E06000002", 2023): "[x]",
            ("All_long_term_vacants", "E06000002", 2023): "[x]",
            ("All_long_term_vacants", "E06000001", 2023): "[x]"})
        f = m.parse_615(cover, sheets)
        b = m.build_615(f, lambda c: (_resolve(c), "direct"))
        k23 = {r["published_la_code"]: r for r in b[2023]}
        self.assertNotIn("E06000002", k23)
        self.assertIsNone(k23["E06000001"]["long_term_vacant_dwellings"])
        self.assertEqual(k23["E06000001"]["null_reasons"],
                         "long_term_vacant_dwellings="
                         "suppressed_or_not_available")
        # Barnsley: E08000016 to 2024, E08000038 in 2025 (rows of [x] are
        # not stored and the code is not a use in that year)
        self.assertNotIn("E08000038", k23)
        self.assertIn("E08000038", {r["published_la_code"] for r in b[2025]})
        uses = m.codes_with_numbers(f)
        self.assertNotIn("E08000038", uses[2024])
        self.assertNotIn("E08000016", uses[2025])
        self.assertNotIn("E06000002", uses[2023])

    def test_england_not_equal_to_the_sum_halts(self):
        cover, sheets = t615_rows(england={("All_long_term_vacants", 2024): 5})
        bad = m.reconcile(m.parse_615(cover, sheets))
        self.assertEqual(len(bad), 1)
        self.assertIn("All_long_term_vacants 2024", bad[0])

    def test_sheets_must_agree(self):
        cover, sheets = t615_rows()
        sheets["All_long_term_vacants"][2][4] = "05/10/2025"
        with self.assertRaisesRegex(ValueError, "different snapshot dates"):
            m.parse_615(cover, sheets)
        cover, sheets = t615_rows()
        sheets["All_long_term_vacants"].pop()
        with self.assertRaisesRegex(ValueError, "missing from sheet"):
            m.parse_615(cover, sheets)
        cover, sheets = t615_rows()
        sheets["All_vacants"].append(["E06000001", "Hartlepool", 1, 1, 1])
        with self.assertRaisesRegex(ValueError, "duplicate code"):
            m.parse_615(cover, sheets)


class RanksLedgerPlan(unittest.TestCase):

    def test_release_rank(self):
        self.assertEqual(m.release_rank(
            "Council Taxbase: ...; first published 2025-11-06; revised "
            "2026-01-21; https://x/y.xlsx"), date(2026, 1, 21))
        self.assertEqual(m.release_rank(
            "as loaded: rows loaded 2026-08-13; file dated 2026-06-25"),
            date(2026, 6, 25))
        self.assertIsNone(m.release_rank(None))

    def test_ledger_source_round_trip(self):
        s = m.ledger_source("https://x/t.ods", date(2026, 6, 25),
                            [2004, 2005, 2006])
        self.assertEqual(s, "https://x/t.ods (file dated 2026-06-25; years "
                            "2004-2006)")
        self.assertEqual(m.parse_ledger_source(s),
                         ("https://x/t.ods", date(2026, 6, 25),
                          [2004, 2005, 2006]))
        self.assertEqual(m.parse_ledger_source(
            m.ledger_source("u", date(2026, 1, 21), [2025]))[2], [2025])

    def test_ledger_complete_needs_the_pair_for_every_period(self):
        r = date(2026, 6, 25)
        src = m.ledger_source("u", r, [2023, 2024, 2025])
        full = {y: {(src, "s")} for y in (2023, 2024, 2025)}
        ranks = {y: r for y in (2023, 2024, 2025)}
        self.assertEqual(m.ledger_complete(full, "u", "s", ranks),
                         [2023, 2024, 2025])
        part = {y: {(src, "s")} for y in (2023, 2024)}
        self.assertIsNone(m.ledger_complete(part, "u", "s", ranks))
        self.assertIsNone(m.ledger_complete(full, "u", "other", ranks))
        self.assertIsNone(m.ledger_complete(full, "v", "s", ranks))
        self.assertIsNone(m.ledger_complete(full, "u", "s", ranks,
                                            stranded=[2024]))
        # a period whose tip is from a newer file is not needed ...
        newer = {**ranks, 2025: date(2026, 11, 1)}
        self.assertEqual(m.ledger_complete(part, "u", "s", newer),
                         [2023, 2024])
        # ... unless --allow-older-file
        self.assertIsNone(m.ledger_complete(part, "u", "s", newer,
                                            allow_older=True))

    def test_plan_periods(self):
        r = date(2026, 1, 21)
        src = m.ledger_source("u", r, [2025])
        new, rev, skip = m.plan_periods([2025], r, {}, {}, src, "s",
                                        recheck=None, allow_older=False)
        self.assertEqual((new, rev, skip), ([2025], [], {}))
        ranks = {2025: r}
        new, rev, skip = m.plan_periods([2025], r, ranks,
                                        {2025: {(src, "s")}}, src, "s",
                                        recheck=None, allow_older=False)
        self.assertEqual(rev, [])
        self.assertTrue(skip[2025].startswith("checked"))
        new, rev, skip = m.plan_periods([2025], r, ranks,
                                        {2025: {(src, "s")}}, src, "s",
                                        recheck=2025, allow_older=False)
        self.assertEqual(rev, [2025])
        older = date(2025, 11, 6)
        new, rev, skip = m.plan_periods([2025], older, ranks, {}, src, "x",
                                        recheck=None, allow_older=False)
        self.assertTrue(skip[2025].startswith("older"))
        new, rev, skip = m.plan_periods([2025], older, ranks, {}, src, "x",
                                        recheck=None, allow_older=True)
        self.assertEqual(rev, [2025])
        with self.assertRaises(ValueError):
            m.plan_periods([2025], r, {}, {}, src, "s", recheck=2025,
                           allow_older=False)

    def test_tip_ranks_take_a_later_ledger_file(self):
        later = m.ledger_source("u", date(2026, 3, 1), [2025])
        self.assertEqual(m.tip_ranks({2025: {"rank": date(2026, 1, 21)}},
                                     {2025: {(later, "s")}}),
                         {2025: date(2026, 3, 1)})


def ctb_rec(code, **kw):
    r = {"lad24cd": code, "taxbase_year": "2025",
         "total_dwellings": 10000, "empty_under_6_months": 600,
         "empty_6_months_plus": 400, "empty_total": 1000,
         "empty_homes_premium_count": 100, "second_homes": 50,
         "unoccupied_exemptions_total": 200, "la_name": code,
         "null_reasons": None}
    r.update(kw)
    return r


def t_rec(code, v=1000, lt=400, **kw):
    r = {"published_la_code": code, "year": "2025", "vacant_dwellings": v,
         "long_term_vacant_dwellings": lt, "published_la_name": code,
         "lad24cd": code, "mapping_status": "direct", "null_reasons": None}
    r.update(kw)
    return r


AUTH = [f"E0600{i:04d}" for i in range(1, 101)]


class StopConditions(unittest.TestCase):

    def setUp(self):
        p = mock.patch.object(m, "CTB_AUTHORITIES", len(AUTH))
        p.start()
        self.addCleanup(p.stop)

    def tip(self):
        return [ctb_rec(c) for c in AUTH]

    def test_ctb_count_and_missing_authorities(self):
        self.assertEqual(m.period_problems(self.tip(), [], kind="new",
                                           part="ctb"), [])
        bad = m.period_problems(self.tip()[:-1], self.tip(), kind="new",
                                part="ctb")
        self.assertEqual(len(bad), 2, bad)
        self.assertIn("99 authorities, expected 100", bad[0])
        self.assertIn("missing from the file", bad[1])

    def test_ctb_national_change_both_sides(self):
        # one authority's empty_total moves: 100 x 1000 nationally
        ok = self.tip()
        ok[0]["empty_total"] = 1000 + 1999           # +1.999%
        self.assertEqual(m.period_problems(ok, self.tip(), kind="revised",
                                           part="ctb"), [])
        bad = self.tip()
        bad[0]["empty_total"] = 1000 + 2001          # +2.001%
        out = m.period_problems(bad, self.tip(), kind="revised", part="ctb")
        self.assertEqual(len(out), 1, out)
        self.assertIn("national empty_total", out[0])

    def test_ctb_changed_authorities_both_sides(self):
        thirty = self.tip()
        for r in thirty[:30]:
            r["second_homes"] = 51
        self.assertEqual(m.period_problems(thirty, self.tip(), kind="revised",
                                           part="ctb"), [])
        many = self.tip()
        for r in many[:31]:
            r["second_homes"] = 51
        out = m.period_problems(many, self.tip(), kind="revised", part="ctb")
        self.assertEqual(len(out), 1, out)
        self.assertIn("31 authorities changed", out[0])

    def test_ctb_total_dwellings_both_sides(self):
        ok = self.tip()
        ok[0]["total_dwellings"] = 11000              # +10%
        self.assertEqual(m.period_problems(ok, self.tip(), kind="revised",
                                           part="ctb"), [])
        bad = self.tip()
        bad[0]["total_dwellings"] = 11001
        out = m.period_problems(bad, self.tip(), kind="revised", part="ctb")
        self.assertEqual(len(out), 1, out)
        self.assertIn("total_dwellings moves", out[0])

    def test_615_new_year_needs_every_boundary_authority_on_both_sheets(self):
        rows = [t_rec(c) for c in AUTH[:3]]
        self.assertEqual(m.period_problems(rows, [], kind="new", part="615",
                                           boundaries=set(AUTH[:3])), [])
        rows[1]["long_term_vacant_dwellings"] = None
        out = m.period_problems(rows, [], kind="new", part="615",
                                boundaries=set(AUTH[:3]))
        self.assertIn(AUTH[1], out[0])

    def test_615_codes_appearing_or_disappearing(self):
        tip = [t_rec(c) for c in AUTH]
        self.assertEqual(m.period_problems(tip, tip, kind="revised",
                                           part="615"), [])
        out = m.period_problems(tip[1:] + [t_rec("E07000001")], tip,
                                kind="revised", part="615")
        self.assertEqual(len(out), 2, out)

    def test_615_england_change_both_sides(self):
        tip = [t_rec(c) for c in AUTH]                # 100,000 nationally
        ok = [t_rec(c) for c in AUTH]
        ok[0]["vacant_dwellings"] = 1000 + 999        # +0.999%
        self.assertEqual(m.period_problems(ok, tip, kind="revised",
                                           part="615"), [])
        bad = [t_rec(c) for c in AUTH]
        bad[0]["vacant_dwellings"] = 1000 + 1001
        out = m.period_problems(bad, tip, kind="revised", part="615")
        self.assertEqual(len(out), 1, out)
        self.assertIn("England vacant_dwellings", out[0])

    def test_615_big_authority_changes_both_sides(self):
        tip = [t_rec(c) for c in AUTH]
        ten = [t_rec(c) for c in AUTH]
        for i in range(10):
            ten[i]["long_term_vacant_dwellings"] = 501       # +25.25%
            ten[i + 50]["long_term_vacant_dwellings"] = 299  # keeps England
        out = m.period_problems(ten, tip, kind="revised", part="615")
        self.assertEqual(len(out), 1, out)               # 20 authorities
        nine = [t_rec(c) for c in AUTH]
        for i in range(5):
            nine[i]["long_term_vacant_dwellings"] = 501
            nine[i + 50]["long_term_vacant_dwellings"] = 299
        self.assertEqual(m.period_problems(nine, tip, kind="revised",
                                           part="615"), [])
        exact = [t_rec(c) for c in AUTH]
        for i in range(11):
            exact[i]["long_term_vacant_dwellings"] = 500     # exactly +25%
            exact[i + 50]["long_term_vacant_dwellings"] = 300
        self.assertEqual(m.period_problems(exact, tip, kind="revised",
                                           part="615"), [])

    def test_zero_null_flips(self):
        tip = [ctb_rec("E06000001", second_homes=0),
               ctb_rec("E06000002", second_homes=None)]
        new = [ctb_rec("E06000001", second_homes=None),
               ctb_rec("E06000002", second_homes=0)]
        self.assertEqual(m.zero_null_flips(new, tip), 2)
        self.assertEqual(m.zero_null_flips(tip, tip), 0)
        cls_t = [{"lad24cd": "E06000001", "exemption_class": "K",
                  "dwellings": 0}]
        cls_n = [{"lad24cd": "E06000001", "exemption_class": "K",
                  "dwellings": None}]
        self.assertEqual(m.zero_null_flips(cls_n, cls_t), 1)


class CrossCheck(unittest.TestCase):

    def test_615_equals_ctb_empty_plus_unoccupied_and_catches_a_difference(self):
        ctb = [ctb_rec(c) for c in AUTH[:5]]
        t615 = [t_rec(c, v=1200) for c in AUTH[:5]]
        self.assertEqual(m.cross_check_615_ctb(t615, ctb, 2025), [])
        t615[2]["vacant_dwellings"] = 1201
        out = m.cross_check_615_ctb(t615, ctb, 2025)
        self.assertEqual(len(out), 1)
        self.assertIn(AUTH[2], out[0])
        out = m.cross_check_615_ctb(t615[:4], ctb, 2025)
        self.assertTrue(any("only in CTB" in x for x in out))


class _Resp:
    def __init__(self, body, url):
        self.content = body
        self.url = url
        self.text = body.decode("latin-1")

    def raise_for_status(self):
        pass


class _Session:
    def __init__(self, body, final):
        self.body, self.final, self.calls = body, final, []

    def get(self, url, **kw):
        self.calls.append(url)
        return _Resp(self.body, self.final)


class Download(Tmp):

    def body(self, **kw):
        return write_ctb(self.root / "src" / f"b{len(kw)}.xlsx",
                         **kw).read_bytes()

    def test_redirect_records_the_final_url(self):
        body = self.body()
        s = _Session(body, "https://assets.example/media/2/final.xlsx")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), quiet() as out:
            path, final = m.fetch("https://assets.example/media/1/old.xlsx",
                                  self.root / "raw", session=s)
        self.assertEqual(final, "https://assets.example/media/2/final.xlsx")
        self.assertEqual(path.name, "final.xlsx")
        self.assertEqual(path.read_bytes(), body)
        self.assertIn("redirected", out.getvalue())

    def test_same_name_different_content_saved_beside_never_overwritten(self):
        raw = self.root / "raw"
        raw.mkdir()
        old = raw / "f.xlsx"
        old.write_bytes(self.body())
        before = old.read_bytes()
        new_body = self.body(year=2026)
        s = _Session(new_body, "https://assets.example/media/3/f.xlsx")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), quiet():
            path, _ = m.fetch("https://assets.example/media/3/f.xlsx", raw,
                              session=s)
        sha8 = hashlib.sha256(new_body).hexdigest()[:8]
        self.assertEqual(path.name, f"f-{sha8}.xlsx")
        self.assertEqual(path.read_bytes(), new_body)
        self.assertEqual(old.read_bytes(), before)

    def test_an_existing_same_named_file_does_not_skip_the_download(self):
        raw = self.root / "raw"
        raw.mkdir()
        body = self.body()
        (raw / "f.xlsx").write_bytes(body)
        s = _Session(body, "https://assets.example/media/3/f.xlsx")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), quiet() as out:
            path, _ = m.fetch("https://assets.example/media/3/f.xlsx", raw,
                              session=s)
        self.assertEqual(s.calls, ["https://assets.example/media/3/f.xlsx"])
        self.assertEqual(path, raw / "f.xlsx")
        self.assertIn("same content", out.getvalue())

    def test_size_floor_and_format_check_keep_nothing(self):
        raw = self.root / "raw"
        s = _Session(b"x" * 50, "https://assets.example/media/3/f.xlsx")
        with self.assertRaises(SystemExit), quiet():
            m.fetch("https://assets.example/media/3/f.xlsx", raw, session=s)
        s = _Session(b"x" * 5000, "https://assets.example/media/3/f.ods")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), \
                self.assertRaises(SystemExit), quiet():
            m.fetch("https://assets.example/media/3/f.ods", raw, session=s)
        self.assertEqual(sorted(p.name for p in raw.iterdir()), [])


if __name__ == "__main__":
    unittest.main()
