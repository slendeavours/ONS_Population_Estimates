"""Pure tests for s24_rsh_register_editions: discovery, identity, cells,
records, ranks, the ledger rule and the stop conditions.

No database and no network: every workbook is written by the test (each
register workbook with a veryHidden sentinel sheet, as the publisher's files
carry one) and the content API replies are dicts built here.
"""
import contextlib
import io
import re
import sys
import tempfile
import unittest
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import s24_rsh_register_editions as m  # noqa: E402

# A string shaped like the publisher-tooling token; it sits only in the
# veryHidden sheet (cell and sheet title) and must reach no output.
SENTINEL = "tokZZ" + "f00dfeed" * 4
HIDDEN_TITLE = "hid" + "f00dfeed" * 3
REGISTER_SHEET = "Registered Providers - list..."
REGISTER_HEADER = ["Organisation name", "Registration number",
                   "Registration date", "Designation", "Corporate form",
                   "Notes "]
REG_ROWS = [
    ("Alpha Homes", "4001", datetime(2001, 1, 2), "Non-profit", "Company ",
     None),
    ("Barnsley Council", "LA01", datetime(1998, 4, 1), "Local authority",
     "Local authority", None),
    ("Pivotal Housing Association", "4747", datetime(2010, 5, 6),
     "Non-profit", "Society", "Subject to an enforcement notice"),
    ("Zen Rented Limited", "5235", datetime(2026, 3, 5), "Profit",
     "Company", None),
]
MEDIA = "https://assets.example/media/{}/{}"
REGISTER_PAGE_TITLE = "List of registered providers – {}"
JUDGEMENTS_TITLE = "Regulatory Judgements and Enforcement Notices Table"


def set_modified(path, when: datetime):
    """Rewrite docProps/core.xml's modified date (openpyxl stamps now())."""
    path = Path(path)
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    stamp = when.strftime("%Y-%m-%dT%H:%M:%SZ")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(<)",
                              rb"\g<1>" + stamp.encode() + rb"\g<2>", data)
            z.writestr(info, data)
    return path


def write_book(path, sheets, modified):
    """sheets: [(title, rows, state)]; returns path."""
    import openpyxl
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows, state in sheets:
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(list(r))
        ws.sheet_state = state
    wb.save(path)
    return set_modified(path, modified)


def write_register(path, rows=None, *, modified=datetime(2026, 9, 17, 14, 47),
                   header=None, title=REGISTER_SHEET, hidden=True):
    sheets = [(title, [header or REGISTER_HEADER] + list(
        REG_ROWS if rows is None else rows), "visible")]
    if hidden:
        sheets.append((HIDDEN_TITLE, [[SENTINEL], ["token", SENTINEL]],
                       "veryHidden"))
    return write_book(path, sheets, modified)


def jrow(**kw):
    """A judgements row (21 cells) with defaults, overridden by name."""
    d = dict(reg="L0001", landlord="Alpha Homes ",
             type="Private Registered Provider", name_change="None",
             other="None", status="Current", c="C1", cchg="First grading",
             cdate=datetime(2026, 2, 25), g="G1",
             gchg="Assessed and unchanged", gdate=datetime(2026, 2, 25),
             v="V2", vchg="Assessed and unchanged",
             vdate=datetime(2026, 2, 25), rent="-", rdate="-", rchg="-",
             pubtype="Narrative RJ", pubdate=datetime(2026, 2, 25),
             eng="Inspection")
    d.update(kw)
    return (d["reg"], d["landlord"], d["type"], d["name_change"], d["other"],
            d["status"], d["c"], d["cchg"], d["cdate"], d["g"], d["gchg"],
            d["gdate"], d["v"], d["vchg"], d["vdate"], d["rent"], d["rdate"],
            d["rchg"], d["pubtype"], d["pubdate"], d["eng"])


LA = dict(type="Local Authority", g="-", gchg="-", gdate="-", v="-",
          vchg="-", vdate="-")
JROWS = [
    jrow(),
    jrow(reg="LA01", landlord="Barnsley Council", c="C2", **LA),
    jrow(reg="L0002", landlord="Merged Homes", c="-",
         cchg="Not assessed yet (as the merged entity)", cdate="-"),
    jrow(reg="L0003", landlord="Quiet Homes", c="-", cchg="-", cdate="-",
         g="-", gchg="-", gdate="-", v="-", vchg="-", vdate="-"),
    jrow(reg="L0004", landlord="Rent Homes", rent="Y",
         rdate=datetime(2026, 1, 9), rchg="First grading", vdate=46000,
         other="None ", name_change="Was L0099"),
    jrow(reg="L0005", landlord="Gap Homes", gdate=None, other=None),
]
NOTICE_HEADER = list(m.NOTICE_HEADERS)
NROWS = [(4747, "Pivotal Housing Association", "None", "None", "Current",
          "Enforcement Notice", datetime(2025, 4, 16), "Reactive Engagement",
          "Economic Standards", datetime(2025, 4, 16))]


def write_judgements(path, jrows=None, nrows=None, *,
                     modified=datetime(2026, 9, 30, 11, 21), jheader=None,
                     nheader=None):
    return write_book(path, [
        ("Regulatory Judgements", [jheader or list(m.JUDGEMENT_HEADERS)]
         + list(JROWS if jrows is None else jrows), "visible"),
        ("Enforcement Notices", [nheader or NOTICE_HEADER]
         + list(NROWS if nrows is None else nrows), "visible")], modified)


def register_page(day="17 September 2026", url=None, *, extra=()):
    url = url or MEDIA.format("r1", "List_of_registered_providers.xlsx")
    atts = [
        {"title": f"List of registered providers – {day} (accessible "
                  "version)", "url": "/government/publications/x/accessible",
         "attachment_type": "html"},
        {"title": REGISTER_PAGE_TITLE.format(day), "url": url,
         "attachment_type": "file"},
        {"title": " Registrations and de-registrations 2026",
         "url": "/government/publications/x/registrations-2026",
         "attachment_type": "html"},
    ] + list(extra)
    return {"title": "Registered providers of social housing – updated "
                     "monthly",
            "public_updated_at": "2026-09-21T09:30:03+01:00",
            "details": {"attachments": atts, "change_history": [
                {"public_timestamp": "2026-09-21T08:30:03Z",
                 "note": f"List updated as of {day}"}]}}


def judgements_page(url, changes=("2026-09-30", "2026-09-09", "2026-08-12"),
                    updated="2026-09-30T09:30:04+01:00", title=None):
    return {"title": "Regulatory judgements and enforcement notices",
            "public_updated_at": updated,
            "details": {"attachments": [{"title": title or JUDGEMENTS_TITLE,
                                         "url": url,
                                         "attachment_type": "file"}],
                        "change_history": [
                            {"public_timestamp": f"{d}T08:30:04Z",
                             "note": "New regulatory judgements added"}
                            for d in changes]}}


class Tmp(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)


class Discovery(unittest.TestCase):

    def test_register_attachment_picks_the_one_dated_xlsx(self):
        a = m.register_attachment(register_page())
        self.assertEqual(a["title_date"], date(2026, 9, 17))
        self.assertTrue(a["url"].endswith(".xlsx"))
        self.assertEqual(len(a["ignored"]), 2)
        self.assertTrue(any("accessible" in t for t in a["ignored"]))
        self.assertTrue(any("Registrations" in t for t in a["ignored"]))

    def test_register_attachment_none_or_several_halts_listing_titles(self):
        page = register_page()
        page["details"]["attachments"] = [
            a for a in page["details"]["attachments"]
            if not a["url"].endswith(".xlsx")]
        with self.assertRaises(ValueError) as e:
            m.register_attachment(page)
        self.assertIn("accessible version", str(e.exception))
        two = register_page(extra=[{
            "title": REGISTER_PAGE_TITLE.format("17 August 2026"),
            "url": MEDIA.format("r0", "old.xlsx")}])
        with self.assertRaises(ValueError) as e:
            m.register_attachment(two)
        self.assertIn("17 August 2026", str(e.exception))
        # an .xlsx with another title is never passed over silently
        odd = register_page(extra=[{"title": "Registered providers list",
                                    "url": MEDIA.format("r9", "x.xlsx")}])
        with self.assertRaises(ValueError):
            m.register_attachment(odd)

    def test_judgements_attachment_exact_title(self):
        a = m.judgements_attachment(judgements_page(
            MEDIA.format("j1", "20260930_RegulatoryJudgementsNotices.xlsx")))
        self.assertEqual(a["file_name"],
                         "20260930_RegulatoryJudgementsNotices.xlsx")
        with self.assertRaises(ValueError) as e:
            m.judgements_attachment(judgements_page(
                MEDIA.format("j1", "x.xlsx"), title="Judgements table"))
        self.assertIn("Judgements table", str(e.exception))
        page = judgements_page(MEDIA.format("j1", "a.xlsx"))
        page["details"]["attachments"].append(
            dict(page["details"]["attachments"][0],
                 url=MEDIA.format("j2", "b.xlsx")))
        with self.assertRaises(ValueError):
            m.judgements_attachment(page)

    def test_overdue_warning(self):
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            m._overdue("register", date(2026, 9, 17), date(2026, 10, 27), 40)
        self.assertEqual(out.getvalue(), "")
        with contextlib.redirect_stdout(out):
            m._overdue("register", date(2026, 9, 17), date(2026, 10, 28), 40)
        self.assertIn("WARNING", out.getvalue())
        self.assertIn("2026-09-17", out.getvalue())

    def test_change_history(self):
        ch = m.change_history(judgements_page("u"))
        self.assertEqual(ch[0][0], "2026-09-30")
        self.assertEqual(len(ch), 3)


class Identity(Tmp):

    def test_register_identity_reads_the_visible_sheet_and_docprops(self):
        p = write_register(self.root / "r.xlsx")
        ident = m.register_identity(p, date(2026, 9, 17))
        self.assertEqual(ident["period"], "2026-09-17")
        self.assertEqual(ident["sheet"], REGISTER_SHEET)
        self.assertEqual(len(ident["rows"]), 4)
        self.assertEqual(ident["rank"][0], date(2026, 9, 17))
        self.assertEqual(m.release_rank(ident), ident["rank"])

    def test_register_title_date_and_docprops_disagreeing_halts(self):
        for mod in (datetime(2026, 9, 16, 23), datetime(2026, 10, 2, 1)):
            p = write_register(self.root / f"r{mod:%d}.xlsx", modified=mod)
            with self.assertRaises(ValueError) as e:
                m.register_identity(p, date(2026, 9, 17))
            self.assertIn("2026-09-17", str(e.exception))
            self.assertIn(f"{mod:%Y-%m-%d}", str(e.exception))
        ok = write_register(self.root / "ok.xlsx",
                            modified=datetime(2026, 10, 1, 12))
        self.assertEqual(m.register_identity(ok, "2026-09-17")["period"],
                         "2026-09-17")

    def test_register_header_changed_or_no_sheet_halts(self):
        hdr = list(REGISTER_HEADER)
        hdr[3] = "Designation type"
        p = write_register(self.root / "h.xlsx", header=hdr)
        with self.assertRaises(ValueError) as e:
            m.register_identity(p, date(2026, 9, 17))
        self.assertIn("Designation type", str(e.exception))
        p = write_register(self.root / "s.xlsx", title="Providers")
        with self.assertRaises(ValueError) as e:
            m.register_identity(p, date(2026, 9, 17))
        self.assertIn("Providers", str(e.exception))
        self.assertNotIn(HIDDEN_TITLE, str(e.exception))

    def test_judgements_identity(self):
        p = write_judgements(self.root / "20260930_RJ.xlsx")
        page = judgements_page("u")
        ident = m.judgements_identity(p, p.name, page)
        self.assertEqual(ident["period"], "2026-09-30")
        self.assertTrue(ident["page_checked"])
        self.assertEqual(len(ident["judgement_rows"]), len(JROWS))
        self.assertEqual(len(ident["notice_rows"]), 1)
        nopage = m.judgements_identity(p, p.name, None)
        self.assertFalse(nopage["page_checked"])
        # public_updated_at alone is enough
        only = judgements_page("u", changes=("2026-09-09",))
        self.assertEqual(m.judgements_identity(p, p.name, only)["period"],
                         "2026-09-30")

    def test_judgements_file_date_matching_no_page_date_halts(self):
        p = write_judgements(self.root / "20260930_RJ.xlsx")
        page = judgements_page("u", changes=("2026-09-09",),
                               updated="2026-09-10T09:30:00+01:00")
        with self.assertRaises(ValueError) as e:
            m.judgements_identity(p, p.name, page)
        self.assertIn("2026-09-30", str(e.exception))
        self.assertIn("2026-09-09", str(e.exception))
        with self.assertRaises(ValueError):
            m.judgements_identity(p, "RJ_published.xlsx", None)

    def test_judgements_docprops_and_header_and_sheets(self):
        p = write_judgements(self.root / "20260930_a.xlsx",
                             modified=datetime(2026, 10, 8, 9))
        with self.assertRaises(ValueError) as e:
            m.judgements_identity(p, p.name, None)
        self.assertIn("2026-10-08", str(e.exception))
        hdr = list(m.JUDGEMENT_HEADERS)
        hdr[15] = "Rent grade"
        p = write_judgements(self.root / "20260930_b.xlsx", jheader=hdr)
        with self.assertRaises(ValueError) as e:
            m.judgements_identity(p, p.name, None)
        self.assertIn("Rent grade", str(e.exception))
        nh = list(m.NOTICE_HEADERS)
        nh[7] = "Path"
        p = write_judgements(self.root / "20260930_c.xlsx", nheader=nh)
        with self.assertRaises(ValueError):
            m.judgements_identity(p, p.name, None)
        p = write_book(self.root / "20260930_d.xlsx", [
            ("Regulatory Judgements", [list(m.JUDGEMENT_HEADERS)], "visible")],
            datetime(2026, 9, 30))
        with self.assertRaises(ValueError) as e:
            m.judgements_identity(p, p.name, None)
        self.assertIn("Enforcement Notices", str(e.exception))


class Cells(unittest.TestCase):

    def test_each_reason_on_its_exact_case(self):
        g = m.grade_cell
        self.assertEqual(g("-", "governance_grade", "Local Authority", "-"),
                         (None, "not_applicable_local_authority"))
        self.assertEqual(g("-", "viability_grade", "Local Authority", "-"),
                         (None, "not_applicable_local_authority"))
        self.assertEqual(
            g("-", "consumer_grade", "Private Registered Provider",
              "Not assessed yet (as the merged entity)"),
            (None, "not_assessed_yet"))
        self.assertEqual(g("-", "consumer_grade", "Private Registered "
                           "Provider", "Not assessed yet"),
                         (None, "not_assessed_yet"))
        self.assertEqual(g("-", "governance_grade",
                           "Private Registered Provider", "-"),
                         (None, "marked_dash_no_reason"))
        self.assertEqual(g("-", "consumer_grade", "Local Authority", "-"),
                         (None, "marked_dash_no_reason"))
        self.assertEqual(g("-", "rent_grade", "Local Authority", "-"),
                         (None, "marked_dash_no_reason"))
        self.assertEqual(g(None, "rent_grade", "x", None), (None, "blank"))
        self.assertEqual(g("Y", "rent_grade", "x", "First grading"),
                         ("Y", None))
        for v, col in (("C1", "consumer_grade"), ("C2*", "consumer_grade"),
                       ("G4", "governance_grade"), ("V1*", "viability_grade"),
                       (" G2 ", "governance_grade")):
            self.assertEqual(g(v, col, "x", "y"), (v.strip(), None))

    def test_an_unseen_grade_halts_naming_the_cell(self):
        for v, col in (("G5", "governance_grade"), ("C2+", "consumer_grade"),
                       ("V1", "governance_grade"), ("N", "rent_grade"),
                       ("C0", "consumer_grade")):
            with self.assertRaises(ValueError) as e:
                m.grade_cell(v, col, "x", "y", where="row 7")
            self.assertIn(v, str(e.exception))
            self.assertIn("row 7", str(e.exception))

    def test_text_cell(self):
        self.assertEqual(m.text_cell("-"), (None, "marked_dash"))
        self.assertEqual(m.text_cell(None), (None, "blank"))
        self.assertEqual(m.text_cell("   "), (None, "blank"))
        self.assertEqual(m.text_cell("None"), (None, "publisher_text_none"))
        self.assertEqual(m.text_cell("None "), (None, "publisher_text_none"))
        self.assertEqual(m.text_cell(" Was L0099 "), ("Was L0099", None))

    def test_date_cell(self):
        w = "row 3"
        self.assertEqual(m.date_cell(datetime(2026, 2, 25, 0, 0), w),
                         (date(2026, 2, 25), None))
        self.assertEqual(m.date_cell(date(2026, 2, 25), w),
                         (date(2026, 2, 25), None))
        self.assertEqual(m.date_cell(46000, w),
                         (date(1899, 12, 30) + timedelta(days=46000), None))
        self.assertEqual(m.date_cell(46000.0, w)[0],
                         date(1899, 12, 30) + timedelta(days=46000))
        self.assertEqual(m.date_cell(20000, w)[1], None)
        self.assertEqual(m.date_cell(60000, w)[1], None)
        self.assertEqual(m.date_cell("-", w), (None, "marked_dash"))
        self.assertEqual(m.date_cell(None, w), (None, "blank"))
        self.assertEqual(m.date_cell("", w), (None, "blank"))
        for bad in (19999, 60001, 46000.5, "2026-02-25", "soon", True):
            with self.assertRaises(ValueError) as e:
                m.date_cell(bad, w)
            self.assertIn(w, str(e.exception))


class Records(Tmp):

    def judgements(self):
        p = write_judgements(self.root / "20260930_RJ.xlsx")
        ident = m.judgements_identity(p, p.name, None)
        return ident, m.build_judgements(ident["judgement_rows"], ident)

    def test_build_judgements_reasons(self):
        ident, recs = self.judgements()
        by = {r["registration_number"]: r for r in recs}
        self.assertEqual(by["LA01"]["governance_grade"], None)
        self.assertEqual(by["LA01"]["null_reasons"]["governance_grade"],
                         "not_applicable_local_authority")
        self.assertEqual(by["LA01"]["null_reasons"]["governance_grade_date"],
                         "marked_dash")
        self.assertEqual(by["L0002"]["null_reasons"]["consumer_grade"],
                         "not_assessed_yet")
        self.assertEqual(by["L0003"]["null_reasons"]["governance_grade"],
                         "marked_dash_no_reason")
        self.assertEqual(by["L0001"]["null_reasons"]["rent_grade"],
                         "marked_dash_no_reason")
        self.assertEqual(by["L0005"]["null_reasons"]["governance_grade_date"],
                         "blank")
        self.assertEqual(by["L0005"]["null_reasons"]["other_landlords"],
                         "blank")
        self.assertEqual(by["L0004"]["null_reasons"]["other_landlords"],
                         "publisher_text_none")
        self.assertEqual(by["L0004"]["name_or_code_change"], "Was L0099")
        self.assertEqual(by["L0004"]["rent_grade"], "Y")
        self.assertEqual(by["L0004"]["viability_grade_date"],
                         date(1899, 12, 30) + timedelta(days=46000))
        self.assertEqual(by["L0001"]["landlord_name"], "Alpha Homes")
        self.assertEqual(by["L0001"]["edition_date"], "2026-09-30")
        self.assertEqual(by["L0001"]["publication_date"], date(2026, 2, 25))
        # every NULL value carries a reason, and only a NULL does
        for r in recs:
            nulls = {c for c in m.JUDGEMENT_VALUES if r[c] is None}
            self.assertEqual(nulls, set(r["null_reasons"]), r)

    def test_duplicate_key_halts(self):
        p = write_judgements(self.root / "20260930_d.xlsx",
                             jrows=[jrow(), jrow(landlord="Again")])
        ident = m.judgements_identity(p, p.name, None)
        with self.assertRaises(ValueError) as e:
            m.build_judgements(ident["judgement_rows"], ident)
        self.assertIn("L0001", str(e.exception))
        p = write_register(self.root / "d.xlsx",
                           rows=[REG_ROWS[0], REG_ROWS[0]])
        ident = m.register_identity(p, date(2026, 9, 17))
        with self.assertRaises(ValueError):
            m.build_register(ident["rows"], ident)

    def test_unseen_grade_in_a_file_halts_naming_the_row(self):
        p = write_judgements(self.root / "20260930_g.xlsx",
                             jrows=[jrow(g="G5")])
        ident = m.judgements_identity(p, p.name, None)
        with self.assertRaises(ValueError) as e:
            m.build_judgements(ident["judgement_rows"], ident)
        self.assertIn("G5", str(e.exception))
        self.assertIn("row 2", str(e.exception))

    def test_serial_out_of_range_in_a_file_halts(self):
        p = write_judgements(self.root / "20260930_s.xlsx",
                             jrows=[jrow(vdate=70000)])
        ident = m.judgements_identity(p, p.name, None)
        with self.assertRaises(ValueError) as e:
            m.build_judgements(ident["judgement_rows"], ident)
        self.assertIn("70000", str(e.exception))

    def test_build_register_and_old_rules_agree_on_values(self):
        p = write_register(self.root / "r.xlsx")
        ident = m.register_identity(p, date(2026, 9, 17))
        new = m.build_register(ident["rows"], ident)
        old = m.build_register(ident["rows"], ident, rules="old")
        by = {r["registration_number"]: r for r in new}
        self.assertEqual(by["4001"]["notes"], None)
        self.assertEqual(by["4001"]["null_reasons"], {"notes": "blank"})
        self.assertEqual(by["4001"]["corporate_form"], "Company")
        self.assertEqual(by["4747"]["null_reasons"], {})
        self.assertEqual(by["4747"]["registration_date"], date(2010, 5, 6))
        for a, b in zip(new, old):
            for c in m.REGISTER_VALUES:
                self.assertEqual(a[c], b[c])

    def test_build_notices(self):
        ident, _ = self.judgements()
        recs = m.build_notices(ident["notice_rows"], ident)
        (r,) = recs
        self.assertEqual(r["registration_number"], "4747")
        self.assertEqual(r["publication_date"], date(2025, 4, 16))
        self.assertEqual(r["null_reasons"],
                         {"name_or_code_change": "publisher_text_none",
                          "other_providers": "publisher_text_none"})


class HiddenSheet(Tmp):

    def test_the_sentinel_reaches_no_output_record_or_exception(self):
        out, err = io.StringIO(), io.StringIO()
        texts = []
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            p = write_register(self.root / "r.xlsx")
            self.assertEqual(m.visible_sheets(p), [REGISTER_SHEET])
            ident = m.register_identity(p, date(2026, 9, 17))
            recs = m.build_register(ident["rows"], ident)
            texts += [repr(ident), repr(recs), m.rows_content_sha(recs)]
            for kw in ({"header": REGISTER_HEADER[:5] + ["Remarks"]},
                       {"title": "Other"},
                       {"modified": datetime(2026, 1, 1)}):
                bad = write_register(self.root / f"b{len(texts)}.xlsx", **kw)
                try:
                    i = m.register_identity(bad, date(2026, 9, 17))
                    m.build_register(i["rows"], i)
                except ValueError as e:
                    texts.append(str(e))
                else:
                    self.fail(f"{kw} did not halt")
        texts += [out.getvalue(), err.getvalue()]
        blob = "\n".join(texts)
        self.assertNotIn(SENTINEL, blob)
        self.assertNotIn(HIDDEN_TITLE, blob)
        self.assertNotIn("f00dfeed", blob)

    def test_the_module_never_iterates_sheets(self):
        src = Path(m.__file__).read_text(encoding="utf-8")
        self.assertNotIn(".worksheets", src)
        self.assertNotIn(".sheetnames", src)


class Ranks(Tmp):

    def test_rank_from_the_source_never_the_url(self):
        url = MEDIA.format("same", "List_of_registered_providers.xlsx")
        a = write_register(self.root / "a" / "List_of_registered_providers"
                           ".xlsx", modified=datetime(2026, 8, 17, 10))
        b = write_register(self.root / "b" / "List_of_registered_providers"
                           ".xlsx", rows=REG_ROWS[:3],
                           modified=datetime(2026, 9, 17, 10))
        self.assertNotEqual(m.content_sha256(a), m.content_sha256(b))
        ia = m.register_identity(a, date(2026, 8, 17))
        ib = m.register_identity(b, date(2026, 9, 17))
        self.assertNotEqual(m.release_rank(ia), m.release_rank(ib))
        self.assertIsNone(m.release_rank(url))
        # the September bytes at the August URL: the file says September
        with self.assertRaises(ValueError):
            m.register_identity(b, date(2026, 8, 17))
        src = m.ledger_source(url, "register", ia["period"], ia["modified"])
        self.assertEqual(m.parse_ledger_source(src)[1], ia["rank"])
        self.assertNotEqual(
            src, m.ledger_source(url, "register", ib["period"],
                                 ib["modified"]))

    def test_release_rank_from_a_source_file(self):
        s = ("List.xlsx; register 2026-09-17; modified 2026-09-17T14:47:33Z; "
             "https://x/y.xlsx")
        self.assertEqual(m.release_rank(s),
                         (date(2026, 9, 17),
                          datetime(2026, 9, 17, 14, 47, 33,
                                   tzinfo=m.timezone.utc)))
        self.assertIsNone(m.release_rank("as loaded: x.xlsx; no file held"))
        self.assertIsNone(m.release_rank(None))

    def test_ledger_source_round_trip_and_as_loaded(self):
        mod = datetime(2026, 9, 30, 11, 21, 9, tzinfo=m.timezone.utc)
        s = m.ledger_source("https://x/20260930.xlsx", "judgements",
                            "2026-09-30", mod, as_loaded=True)
        self.assertIn("as loaded", s)
        where, rank, periods = m.parse_ledger_source(s)
        self.assertEqual(where, "https://x/20260930.xlsx")
        self.assertEqual(rank, (date(2026, 9, 30), mod))
        self.assertEqual(periods, ["2026-09-30"])

    def test_ledger_complete(self):
        mod = datetime(2026, 9, 17, 10, tzinfo=m.timezone.utc)
        src = m.ledger_source("U", "register", "2026-09-17", mod)
        checked = {"2026-09-17": {(src, "S1")}}
        self.assertEqual(m.ledger_complete(checked, "U", "S1", {}),
                         ["2026-09-17"])
        self.assertIsNone(m.ledger_complete(checked, "U", "S2", {}))
        self.assertIsNone(m.ledger_complete(checked, "V", "S1", {}))
        self.assertIsNone(m.ledger_complete(checked, "U", "S1", {},
                                            stranded=["2026-09-17"]))

    def test_plan_periods(self):
        r = (date(2026, 9, 30), datetime(2026, 9, 30, 11,
                                         tzinfo=m.timezone.utc))
        older = (date(2026, 9, 30), datetime(2026, 9, 29,
                                             tzinfo=m.timezone.utc))
        newer = (date(2026, 9, 30), datetime(2026, 10, 1,
                                             tzinfo=m.timezone.utc))
        P = "2026-09-30"
        self.assertEqual(m.plan_periods([P], r, {}, {}, "s", "h",
                                        recheck=None, allow_older=False),
                         ([P], [], {}))
        self.assertEqual(m.plan_periods([P], r, {P: older}, {}, "s", "h",
                                        recheck=None, allow_older=False),
                         ([], [P], {}))
        new, rev, skip = m.plan_periods([P], r, {P: newer}, {}, "s", "h",
                                        recheck=None, allow_older=False)
        self.assertIn("older", skip[P])
        self.assertEqual(m.plan_periods([P], r, {P: newer}, {}, "s", "h",
                                        recheck=None, allow_older=True)[1],
                         [P])
        # a new judgements period older than the newest held one
        new, rev, skip = m.plan_periods(
            ["2026-08-12"], r, {P: r}, {}, "s", "h", recheck=None,
            allow_older=False, newest_only=True)
        self.assertIn("older", skip["2026-08-12"])
        self.assertEqual(m.plan_periods(
            ["2026-08-12"], r, {P: r}, {}, "s", "h", recheck=None,
            allow_older=True, newest_only=True)[0], ["2026-08-12"])
        # the ledger already holds the pair
        new, rev, skip = m.plan_periods([P], r, {P: r}, {P: {("s", "h")}},
                                        "s", "h", recheck=None,
                                        allow_older=False)
        self.assertIn("ledger", skip[P])
        self.assertEqual(m.plan_periods([P], r, {P: r}, {P: {("s", "h")}},
                                        "s", "h", recheck=P,
                                        allow_older=False)[1], [P])
        with self.assertRaises(ValueError):
            m.plan_periods([P], r, {}, {}, "s", "h", recheck="2026-01-01",
                           allow_older=False)


def reg(n, *, start=0):
    return [{"registration_number": f"R{i:04d}", "organisation_name": f"P{i}",
             "registration_date": date(2001, 1, 1), "designation": "Profit",
             "corporate_form": "Company", "notes": None,
             "null_reasons": m.Reasons(notes="blank"),
             "snapshot_date": "2026-09-17"}
            for i in range(start, start + n)]


def jud(n):
    return [{"registration_number": f"L{i:04d}",
             "publication_date": date(2026, 1, 1), "landlord_name": "x",
             "edition_date": "2026-09-30", "null_reasons": m.Reasons(),
             **{c: None for c in m.JUDGEMENT_VALUES if c != "landlord_name"}}
            for i in range(n)]


class StopConditions(unittest.TestCase):

    def test_register_count_threshold_both_sides(self):
        prev = reg(100)
        self.assertEqual(m.period_problems(reg(95), None, prev,
                                           kind="register"), [])
        bad = m.period_problems(reg(94), None, prev, kind="register")
        self.assertTrue(any("95%" in b for b in bad), bad)

    def test_register_added_and_removed_threshold_both_sides(self):
        prev = reg(100)
        self.assertEqual(m.period_problems(reg(130), None, prev,
                                           kind="register"), [])
        self.assertTrue(m.period_problems(reg(131), None, prev,
                                          kind="register"))
        # 30 removed and 30 added (same count) passes; 31 removed does not
        self.assertEqual(m.period_problems(reg(100, start=30), None, prev,
                                           kind="register"), [])
        bad = m.period_problems(reg(100, start=31), None, prev,
                                kind="register")
        self.assertTrue(any("removed" in b for b in bad), bad)
        self.assertTrue(any("added" in b for b in bad), bad)

    def test_judgements_count_threshold_both_sides(self):
        prev = jud(100)
        self.assertEqual(m.period_problems(jud(90), None, prev,
                                           kind="judgements"), [])
        bad = m.period_problems(jud(89), None, prev, kind="judgements")
        self.assertTrue(any("90%" in b for b in bad), bad)
        self.assertEqual(m.period_problems(jud(1), None, None,
                                           kind="judgements"), [])

    def test_a_reissue_that_changes_any_row_needs_acknowledge(self):
        tip = reg(10)
        self.assertEqual(m.period_problems(reg(10), tip, None,
                                           kind="register"), [])
        one = reg(10)
        one[3]["designation"] = "Non-profit"
        bad = m.period_problems(one, tip, None, kind="register")
        self.assertTrue(any("--acknowledge" in b for b in bad), bad)
        why = reg(10)
        why[0]["null_reasons"] = m.Reasons()
        self.assertTrue(m.period_problems(why, tip, None, kind="register"))
        self.assertTrue(m.period_problems(jud(2), jud(3), None,
                                          kind="notices"))


class ContentSha(unittest.TestCase):

    def test_provenance_excluded_reasons_included_order_free(self):
        a = reg(3)
        b = [dict(r, file_name="other.xlsx", file_source_url="u")
             for r in reversed(reg(3))]
        self.assertEqual(m.rows_content_sha(a), m.rows_content_sha(b))
        c = reg(3)
        c[1]["null_reasons"] = m.Reasons()
        self.assertNotEqual(m.rows_content_sha(a), m.rows_content_sha(c))
        self.assertEqual(m.rows_content_sha(a, reasons=False),
                         m.rows_content_sha(c, reasons=False))
        self.assertEqual(len(m.rows_content_sha(a)), 64)


if __name__ == "__main__":
    unittest.main()
