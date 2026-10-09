"""Pure tests for s9a_drd_editions: links, identity, parsing, records,
precision-only. No database, no network; workbooks are built in the test."""
import sys
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import s9a_drd_editions as m  # noqa: E402

BASE = "https://www.england.nhs.uk/statistics/wp-content/uploads/sites/2/"
FILE = "Discharge-Ready-Date-monthly-data-webfile-"


def url(y, mo, name):
    return f"{BASE}{y}/{mo:02d}/{FILE}{name}.xlsx"


def html_of(*hrefs):
    return "<html>" + "".join(f'<a href="{h}">x</a>' for h in hrefs) + "</html>"


# ---------------------------------------------------------------------------
# A workbook with the real layout: 'Cover Sheet', 'UTLA Acceptable' (header
# row 15, 49 columns), 'UTLA Unacceptable'
# ---------------------------------------------------------------------------

TITLE = "Timeliness of Acute Hospital Discharges completed within the month"
WIDTH = 49


def header_row():
    h = [None] * WIDTH
    h[0], h[1], h[2] = "Summary Type", "UTLA Code", "UTLA"
    for _, pos, _, text in m.COLUMNS:
        h[pos] = text
    return h


def data_row(code, name, v=None, kind="UTLA Aggregate"):
    """A data row; v {field: value} overrides the defaults."""
    vals = {"pct_acceptable_trust_coverage": 0.5, "total_discharges": "-",
            "total_discharges_acceptable_trusts": 100,
            "total_bed_days_lost": 300, "pct_same_day_discharge": 0.9,
            "pct_delayed_1plus_days": 0.1, "discharged_no_delay": 90,
            "discharged_1_day": 5, "discharged_2_3_days": 3,
            "discharged_4_6_days": 1, "discharged_7_13_days": 1,
            "discharged_14_20_days": 0, "discharged_21_plus_days": 0,
            "avg_days_drd_to_discharge_inc_zero": 1.25,
            "avg_days_drd_to_discharge_exc_zero": 2.5}
    vals.update(v or {})
    row = [None] * WIDTH
    row[0], row[1], row[2] = kind, code, name
    for field, pos, kind_, _ in m.COLUMNS:
        if kind_ in ("count", "ratio"):
            row[pos] = vals[field]
    return row


def write_workbook(path, period, rows, *, revised=None, published="8th  October 2026",
                   title=TITLE, header=None, utla_period=None,
                   utla_revised=...):
    """rows: list of data_row(...) lists. period: a date."""
    import openpyxl
    wb = openpyxl.Workbook()
    notes = wb.active
    notes.title = "Notes"
    notes.append(["Title:", title])
    cover = wb.create_sheet("Cover Sheet")
    cover.append([])
    for label, value in (("Title:", title), ("Period:", datetime.combine(
            period, datetime.min.time())), ("Published:", published),
            ("Revised:", revised if revised is not None else "-")):
        cover.append([None, label, value])
    sheet = wb.create_sheet("UTLA Acceptable")
    sheet.append([])
    sheet.append(["Title:", title])
    sheet.append(["Summary:", "x"])
    up = utla_period or period
    sheet.append(["Period:", datetime.combine(up, datetime.min.time())])
    sheet.append(["Source:", "x"])
    sheet.append(["Basis:", "Provider"])
    sheet.append(["Published:", published])
    ur = revised if utla_revised is ... else utla_revised
    sheet.append(["Revised:", ur if ur is not None else "-"])
    while sheet.max_row < 13:
        sheet.append([" "])
    sheet.append([" "] * WIDTH)                        # row 14 (block labels)
    sheet.append(header or header_row())               # row 15 (index 14)
    sheet.append(data_row(None, None, kind="National"))
    for r in rows:
        sheet.append(r)
    sheet.append(data_row("E09000002", "x", kind="Provider activity mapped "
                          "to UTLA"))
    wb.create_sheet("UTLA Unacceptable")
    wb.save(path)
    return path


class Links(unittest.TestCase):
    def test_period_from_link(self):
        cases = {
            "...-webfile-April-2024-Revised.xlsx": "2024-04-01",
            "Discharge-Ready-Date-monthly-data-webfile-August-2026.xlsx":
                "2026-08-01",
            "x-webfile-March-2024-revised.xlsx": "2024-03-01",
            "x-webfile-November2023.xlsx": "2023-11-01",
            "x-webfile-Sept-2024.xlsx": "2024-09-01",
            "x-webfile-Sep-2024.xlsx": "2024-09-01",
            "x-webfile-Jan-2025-Revised.xlsx": "2025-01-01",
            "x-webfile-DECEMBER_2025.xlsx": "2025-12-01",
            "x-webfile-May2026.xlsx": "2026-05-01",
        }
        for link, want in cases.items():
            self.assertEqual(m.period_from_link(link), want, link)
        self.assertIsNone(m.period_from_link("x-webfile-timeseries.xlsx"))
        # the upload path never names the month
        self.assertEqual(m.period_from_link(url(2025, 7, "June-2025-Revised")),
                         "2025-06-01")

    def test_find_month_links_webfiles_only(self):
        a = url(2026, 10, "August-2026")
        csv = f"{BASE}2026/10/Discharge-Ready-Date-monthly-data-csv-August-2026.csv"
        ts = f"{BASE}2026/10/Discharge-Ready-Date-timeseries-webfile-Aug-2026.xlsx"
        other = f"{BASE}2026/10/something-else-April-2026.xlsx"
        pre = url(2024, 5, "March-2024-revised")
        got = m.find_month_links(html_of(a, csv, ts, other, pre, a))
        self.assertEqual(got, {"2026-08-01": [a], "2024-03-01": [pre]})

    def test_rank_and_choice(self):
        orig = url(2025, 6, "April-2025")
        rev = url(2025, 7, "April-2025-Revised")
        later = url(2026, 7, "April-2025-Revised")
        self.assertGreater(m.link_rank(rev), m.link_rank(orig))
        self.assertEqual(m.choose_link([orig, rev]), rev)        # -Revised
        self.assertEqual(m.choose_link([rev, later]), later)     # later upload
        self.assertEqual(m.choose_link([later, orig, rev]), later)
        # same upload month: -Revised beats the original
        self.assertEqual(m.choose_link([url(2025, 7, "May-2025"),
                                        url(2025, 7, "May-2025-Revised")]),
                         url(2025, 7, "May-2025-Revised"))
        self.assertEqual(m.choose_link([orig]), orig)

    def test_tie_raises_naming_both(self):
        a = url(2025, 7, "May-2025-Revised")
        b = url(2025, 7, "May-2025-revised")
        with self.assertRaises(ValueError) as cm:
            m.choose_link([a, b])
        self.assertIn(a, str(cm.exception))
        self.assertIn(b, str(cm.exception))

    def test_link_older(self):
        orig = url(2025, 6, "April-2025")
        rev = url(2025, 7, "April-2025-Revised")
        self.assertTrue(m.link_older(orig, rev))
        self.assertFalse(m.link_older(rev, orig))
        self.assertFalse(m.link_older(rev, rev))
        self.assertFalse(m.link_older(rev, None))
        # a hand-built live source (no upload path) never makes a link older
        self.assertFalse(m.link_older(orig, "https://x/page file.csv"))
        # no upload path: only -Revised is compared
        self.assertTrue(m.link_older("April-2025.xlsx", rev))

    def test_csv_reissue_notes(self):
        web = url(2026, 6, "April-2026")
        csvv2 = (f"{BASE}2026/06/Discharge-Ready-Date-monthly-data-csv-"
                 "April-2026v2.csv")
        csv = (f"{BASE}2026/06/Discharge-Ready-Date-monthly-data-csv-"
               "May-2026.csv")
        notes = m.csv_reissue_notes(html_of(web, csvv2, csv),
                                    {"2026-04-01": web})
        self.assertEqual(len(notes), 1)
        self.assertIn("April-2026v2.csv", notes[0])
        self.assertIn("2026-04-01", notes[0])
        self.assertEqual(m.csv_reissue_notes(html_of(web), {}), [])


class Identity(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def wb(self, name="a.xlsx", period=date(2026, 8, 1), **kw):
        return write_workbook(Path(self.tmp.name) / name, period,
                              [data_row("E09000002", "Barking")], **kw)

    def test_read_cover(self):
        c = m.read_cover(self.wb(published="13th  August 2026"))
        self.assertEqual(c["period"], date(2026, 8, 1))
        self.assertEqual(c["published"], date(2026, 8, 13))
        self.assertIsNone(c["revised"])
        self.assertEqual(c["title"], m.DRD_TITLE + " completed within the month")
        c = m.read_cover(self.wb("b.xlsx", revised="9th July 2026"))
        self.assertEqual(c["revised"], date(2026, 7, 9))
        with self.assertRaises(ValueError):
            m.read_cover(self.wb("c.xlsx", published="whenever"))

    def test_identity_ok_and_mismatches(self):
        c = m.read_cover(self.wb())
        self.assertEqual(m.check_identity("2026-08-01", c, "x-August-2026.xlsx"), [])
        # cover period differs from the link's month
        probs = m.check_identity("2026-07-01", c, "x-July-2026.xlsx")
        self.assertTrue(any("Period" in p for p in probs), probs)
        # UTLA sheet period differs
        c2 = m.read_cover(self.wb("d.xlsx", utla_period=date(2026, 7, 1)))
        self.assertTrue(m.check_identity("2026-08-01", c2, "x-August-2026.xlsx"))
        # -Revised in the name but no Revised: date, and the reverse
        probs = m.check_identity("2026-08-01", c, "x-August-2026-Revised.xlsx")
        self.assertTrue(any("Revised" in p for p in probs), probs)
        cr = m.read_cover(self.wb("e.xlsx", revised="9th July 2026"))
        probs = m.check_identity("2026-08-01", cr, "x-August-2026.xlsx")
        self.assertTrue(any("not named -Revised" in p for p in probs), probs)
        self.assertEqual(m.check_identity("2026-08-01", cr,
                                          "x-August-2026-Revised.xlsx"), [])
        # the title
        ct = m.read_cover(self.wb("f.xlsx", title="Something else"))
        self.assertTrue(any("title" in p for p in
                            m.check_identity("2026-08-01", ct, "x-August-2026.xlsx")))


class Parse(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.n = 0

    def wb(self, rows, **kw):
        self.n += 1
        return write_workbook(Path(self.tmp.name) / f"w{self.n}.xlsx",
                              date(2026, 8, 1), rows, **kw)

    def test_parse_rows_types_and_drops(self):
        rows, dropped = m.parse_workbook(self.wb([
            data_row("E09000002", "Barking and Dagenham"),
            data_row("E10000003", "Cambridgeshire"),
            data_row("NULL", "NULL"), data_row("#N/A", "#N/A"),
            data_row("#N/A", "#N/A")]))
        self.assertEqual([r["utla_code"] for r in rows],
                         ["E09000002", "E10000003"])
        self.assertEqual(dropped, {"NULL": 1, "#N/A": 2})
        r = rows[0]
        self.assertEqual(r["utla_name"], "Barking and Dagenham")
        self.assertIsNone(r["total_discharges"])         # '-' is NULL
        self.assertEqual(r["total_bed_days_lost"], 300)
        self.assertIsInstance(r["total_bed_days_lost"], int)
        self.assertEqual(r["avg_days_drd_to_discharge_inc_zero"],
                         Decimal("1.25"))
        self.assertIsInstance(r["pct_same_day_discharge"], Decimal)
        self.assertEqual(set(r), {"utla_code", "utla_name"} | set(m.VALUE_COLUMNS))

    def test_float_is_decimal_of_repr(self):
        v = 0.9287441449555547
        rows, _ = m.parse_workbook(self.wb([data_row(
            "E09000002", "x", {"pct_same_day_discharge": v})]))
        self.assertEqual(rows[0]["pct_same_day_discharge"], Decimal(repr(v)))

    def test_dash_is_null_not_zero_and_zero_stays_zero(self):
        rows, _ = m.parse_workbook(self.wb([data_row(
            "E09000002", "x", {"discharged_14_20_days": "-",
                               "avg_days_drd_to_discharge_exc_zero": "-",
                               "discharged_21_plus_days": 0,
                               "pct_delayed_1plus_days": 0})]))
        r = rows[0]
        self.assertIsNone(r["discharged_14_20_days"])
        self.assertIsNone(r["avg_days_drd_to_discharge_exc_zero"])
        self.assertEqual(r["discharged_21_plus_days"], 0)
        self.assertIsNotNone(r["discharged_21_plus_days"])
        self.assertEqual(r["pct_delayed_1plus_days"], 0)
        self.assertIsNotNone(r["pct_delayed_1plus_days"])

    def test_unknown_marker_halts(self):
        for bad in ("*", "[c]", "n/a", "x", ".."):
            with self.assertRaises(ValueError, msg=bad):
                m.parse_workbook(self.wb([data_row(
                    "E09000002", "x", {"discharged_1_day": bad})]))

    def test_non_whole_count_halts(self):
        with self.assertRaises(ValueError):
            m.parse_workbook(self.wb([data_row(
                "E09000002", "x", {"total_bed_days_lost": 300.5})]))
        # a whole float is a whole number
        rows, _ = m.parse_workbook(self.wb([data_row(
            "E09000002", "x", {"total_bed_days_lost": 300.0})]))
        self.assertEqual(rows[0]["total_bed_days_lost"], 300)

    def test_other_codes_halt(self):
        for code in ("W06000001", "E07000008", "S12000005", "", None):
            with self.assertRaises(ValueError, msg=str(code)):
                m.parse_workbook(self.wb([data_row(code, "x")]))

    def test_header_shift_halts(self):
        h = header_row()
        h[11], h[12] = None, h[11]            # bed days column moved one right
        with self.assertRaises(ValueError) as cm:
            m.parse_workbook(self.wb([data_row("E09000002", "x")], header=h))
        self.assertIn("column 11", str(cm.exception))
        h = header_row()
        h[46] = "Average days from Discharge Ready Date to date of discharge"
        with self.assertRaises(ValueError):
            m.parse_workbook(self.wb([data_row("E09000002", "x")], header=h))
        # whitespace differences alone are not a shift
        h = header_row()
        h[8] = h[8] + " "
        rows, _ = m.parse_workbook(self.wb([data_row("E09000002", "x")],
                                           header=h))
        self.assertEqual(len(rows), 1)

    def test_unacceptable_sheet_is_not_read(self):
        # the sheet match is exact: 'UTLA Unacceptable' alone does not match
        import openpyxl
        p = self.wb([data_row("E09000002", "x")])
        wb = openpyxl.load_workbook(p)
        del wb["UTLA Acceptable"]
        wb.save(p)
        with self.assertRaises(ValueError):
            m.parse_workbook(p)


class Records(unittest.TestCase):
    def rows(self, *codes):
        return [{"utla_code": c, "utla_name": c,
                 **{k: None for k in m.VALUE_COLUMNS}} for c in codes]

    def test_build_records_recodes_sorts_and_stamps_period(self):
        recs = m.build_records(self.rows("E08000019", "E06000001"),
                               {"E08000019": "E08000099"}, "2026-08-01")
        self.assertEqual([r["utla_code"] for r in recs],
                         ["E06000001", "E08000099"])
        self.assertEqual({r["reporting_period"] for r in recs}, {"2026-08-01"})

    def test_duplicate_key_raises(self):
        with self.assertRaises(ValueError):
            m.build_records(self.rows("E06000001", "E06000001"), {},
                            "2026-08-01")
        with self.assertRaises(ValueError):
            m.build_records(self.rows("E08000016", "E08000038"),
                            {"E08000038": "E08000016"}, "2026-08-01")


class Misc(unittest.TestCase):
    def test_precision_only(self):
        self.assertTrue(m.precision_only(Decimal("0.123456789012"),
                                         Decimal("0.123456789")))
        self.assertTrue(m.precision_only(Decimal("1"), Decimal("1")))
        self.assertFalse(m.precision_only(Decimal("0.5"), Decimal("0.50000001")))
        self.assertFalse(m.precision_only(Decimal("0.5"), Decimal("0.51")))
        self.assertFalse(m.precision_only(None, Decimal("0")))
        self.assertFalse(m.precision_only(Decimal("0"), None))
        self.assertFalse(m.precision_only(None, None))

    def test_content_sha256(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "a.bin"
            p.write_bytes(b"abc")
            self.assertEqual(
                m.content_sha256(p),
                "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad")

    def test_plan_months(self):
        held = ["2026-05-01", "2026-06-01", "2026-07-01"]
        r = "https://x/2026/08/f-Revised.xlsx"
        page = {"2026-03-01": "l3", "2026-05-01": "l5", "2026-06-01": r,
                "2026-07-01": "l7", "2026-08-01": "l8"}
        tips = {"2026-05-01": "l5", "2026-06-01": "csv-string",
                "2026-07-01": "l7"}
        new, redo, earlier = m.plan_months(held, page, tips, {},
                                           recheck_all=False, recheck=(),
                                           min_period=None)
        self.assertEqual((new, redo, earlier),
                         (["2026-08-01"], ["2026-06-01"], ["2026-03-01"]))
        # a link already in the ledger is not rechecked
        _, redo, _ = m.plan_months(held, page, tips, {"2026-06-01": {r}},
                                   recheck_all=False, recheck=(),
                                   min_period=None)
        self.assertEqual(redo, [])
        _, redo, _ = m.plan_months(held, page, tips, {"2026-06-01": {r}},
                                   recheck_all=True, recheck=(),
                                   min_period=None)
        self.assertEqual(redo, held)
        _, redo, _ = m.plan_months(held, page, tips, {"2026-06-01": {r}},
                                   recheck_all=False,
                                   recheck=("2026-05-01",), min_period=None)
        self.assertEqual(redo, ["2026-05-01"])
        # --min-period widens: the early month becomes new
        new, _, earlier = m.plan_months(held, page, tips, {},
                                        recheck_all=False, recheck=(),
                                        min_period="2026-03-01")
        self.assertEqual((new, earlier), (["2026-03-01", "2026-08-01"], []))

    def test_stop_problems(self):
        s = {"per_period": {
            "a": {"null_areas": 6, "big_areas": 0, "zero_null": 0},
            "b": {"null_areas": 5, "big_areas": 10, "zero_null": 0},
            "c": {"null_areas": 0, "big_areas": 11, "zero_null": 0},
            "d": {"null_areas": 0, "big_areas": 0, "zero_null": 1}}}
        self.assertEqual(sorted(m.stop_problems(s)), ["a", "c", "d"])
        self.assertEqual(sorted(m.stop_problems(s, ("d", "a"))), ["c"])


if __name__ == "__main__":
    unittest.main()
