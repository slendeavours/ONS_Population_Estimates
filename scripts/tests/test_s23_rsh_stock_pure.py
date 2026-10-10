"""Pure tests for s23_rsh_stock_editions: discovery, the look-up tool reader,
stock cells, reconciliation, records, ranks, ledger and planning, the stop
conditions and the download rules.

No database and no network: every workbook is written by the tests
(openpyxl, the 2025 STOCK_BY_LA layout and a 2024-style header), and the
content API replies and downloads are stubbed. The fixture writers are
reused by test_s23_rsh_stock_loader.
"""
import contextlib
import hashlib
import io
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import s23_rsh_stock_editions as m  # noqa: E402

HEADERS = ["RP_Name", "RP_Code", "RP_Type", "SDR_Size", "Survey_Status",
           "LA_Nm", "LA_Code", "Concat", "Total Social Stock", "LA_GN_SC_Own",
           "LA_GN_BSp_Own", "LA_SHHOP", "LA_LCHO_Less_100_Eqty_Own"]
HEADERS_2024 = HEADERS[:9] + ["CLCRR025_LA_GN_SC_Own",
                              "CLCRR026_LA_GN_BSp_Own",
                              "CLCRR027_LA_Sup/Hop_Own",
                              "CLCHO012_LA_LCHO_Less_100_Eqty_Own"]
PARTS = HEADERS[9:]
TOTAL = HEADERS[8]
CODES = ["E06000001", "E06000002", "E08000016", "E08000019"]
LA_NAMES = {"E06000001": "Hartlepool", "E06000002": "Middlesbrough",
            "E08000016": "Barnsley", "E08000019": "Sheffield",
            "E08000038": "Barnsley", "E07000223": "Adur"}
REGION_OF = {"E06000001": "North East", "E06000002": "North East",
             "E08000016": "Yorkshire and The Humber",
             "E08000019": "Yorkshire and The Humber",
             "E08000038": "Yorkshire and The Humber",
             "E07000223": "South East"}
# (name, code, RP_Type, SDR_Size, Survey_Status): six providers in every
# authority; the Small one holds a few units only
PROVIDERS = [
    ("Clarion Housing Association Limited", "4865", "Large", "Long Form",
     "Signed_Off"),
    ("Home Group Limited", "L3076", "Large", "Long Form", "Signed_Off"),
    ("Housing 21", "L0055", "Large", "Long Form", "Signed_Off"),
    ("Hyde Housing Association Limited", "LH0032", "Large", "Long Form",
     "Signed_Off"),
    ("Habinteg Housing Association Limited", "LH0459", "Large", "Long Form",
     "Signed_Off"),
    ("Abbeyfield South Downs Limited", "H0375", "Small", "Short Form",
     "Signed_Off"),
]
ROWS_PER_AREA = len(PROVIDERS) + 1          # and the authority's own LARP


def quiet():
    return contextlib.redirect_stdout(io.StringIO())


def larp_code(code):
    return f"00C{CODES.index(code) if code in CODES else 9}"


def provider_rows(codes=None):
    """The default provider rows of a fixture tool: [name, rp_code, type,
    size, status, la_code, (gn_sc, gn_bsp, shhop, lcho)]. The first
    authority's LARP owns no stock (published zeros)."""
    out = []
    for i, code in enumerate(codes or CODES):
        for j, (name, rp, typ, size, status) in enumerate(PROVIDERS):
            vals = ((2, 0, 1, 0) if typ == "Small"
                    else (100 + 10 * i + j, j % 2, 20 + i, 5 + j))
            out.append([name, rp, typ, size, status, code, vals])
        vals = (0, 0, 0, 0) if i == 0 else (1000 + i, 0, 50, 3)
        out.append([f"{LA_NAMES.get(code, code)} Council", larp_code(code),
                    "LARP", "LARP", "Signed-Off", code, vals])
    return out


def history_for(version, month, year):
    if version in ("1", "1.0"):
        return [(1, month, "Original release.")]
    return [(1, f"October {year}", "Original release."),
            (float(version), month, "Corrected an issue.")]


def write_tool(path, *, year=2025, version="1.1", month="November 2025",
               title=None, source=None, history=None, history_tail=None,
               header=None, codes=None, rows=None, values=None, drop=(),
               add=(), subtotal=None, region=None, extra_rows=(),
               trailing=3, after_trailing=(), sheets=None, lcho_note=True):
    """A look-up tool workbook in the 2025 layout: Introduction and Contents
    (title, source line, publication month, version), Version History, Area
    Summary (with the 2025 note "Unit counts for LCHO are for LARPs and
    Large PRPs only." unless lcho_note is false) and STOCK_BY_LA (provider
    rows, one LA subtotal row per authority, one
    Region row per region, then wholly empty rows). values {(rp_code,
    la_code): {header: v}} overrides cells (the total is recomputed from the
    parts unless given); drop: (rp_code, la_code) rows left out; add: more
    provider rows (provider_rows form). The subtotals are the sums of the
    provider rows, the regions the sums of the subtotals, unless subtotal
    {la_code: {header: v}} or region {name: {header: v}} override."""
    import openpyxl
    codes = list(codes or CODES)
    base = [list(r) for r in (rows if rows is not None
                              else provider_rows(codes))]
    base = [r for r in base if (r[1], r[5]) not in set(drop)] + \
        [list(r) for r in add]
    header = list(header or HEADERS)
    wb = openpyxl.Workbook()
    intro = wb.active
    intro.title = "Introduction and Contents"
    for r in (["J", "K", "L"], ["All needs met   ", "Some needs met"],
              [title or f"RP social housing by local authority area (SDR "
               f"and LADR data) {year}"],
              ["This tool collates regulatory data."], ["Contents"],
              ["Source: ", source or ("Statistical Data Return (SDR)/Local "
                                      "Authority Data Return (LADR) 1 April "
                                      f"{year - 1} to 31 March {year}")],
              [f"Publication date: {month}"], [f"Version: {version}"]):
        intro.append(r)
    wb.create_sheet("Glossary").append(["Term", "Meaning"])
    vh = wb.create_sheet("Version History")
    vh.append(["RP social housing by local authority area"])
    vh.append(["Version", "Publication Date", "Changes "])
    for r in (history if history is not None
              else history_for(version, month, year)):
        vh.append(list(r))
    vh.append(["Contact"])
    for r in (history_tail if history_tail is not None
              else [[f"Publication date: {month}"], [f"Version: {version}"]]):
        vh.append(r)
    area = wb.create_sheet("Area Summary")
    area.append([None, f"RP social housing by local authority area {year}"])
    area.append([None, "Tables 1 & 2 comprise all LARPs & PRPs - unweighted. "
                 + ("Unit counts for LCHO are for LARPs and Large PRPs only. "
                    if lcho_note else "")
                 + "Figures for GN and SH/HOP include intermediate and "
                 "Affordable Rent units."])
    area.append([None, "Owned stock.  LARPs and Large PRPs only - "
                 "unweighted."])
    ws = wb.create_sheet("STOCK_BY_LA")
    ws.append(header)
    def num(v):
        return isinstance(v, int) and not isinstance(v, bool)

    sums = {}
    for name, rp, typ, size, status, code, vals in base:
        la = LA_NAMES.get(code, code)
        row = dict(zip(HEADERS, [name, rp, typ, size, status, la, code,
                                 la + name, None] + list(vals)))
        over = dict((values or {}).get((rp, code), {}))
        row.update(over)
        if TOTAL not in over:
            row[TOTAL] = sum(row[k] for k in PARTS if num(row[k]))
        ws.append([row[h] for h in HEADERS])
        s = sums.setdefault(code, dict.fromkeys(HEADERS[8:], 0))
        for k in HEADERS[8:]:
            if num(row[k]):
                s[k] += row[k]
    regions = {}
    for code in codes:
        true = sums.get(code, dict.fromkeys(HEADERS[8:], 0))
        s = dict(true, **(subtotal or {}).get(code, {}))
        reg = REGION_OF.get(code, "South East")
        la = LA_NAMES.get(code, code)
        ws.append([la, code, "LA", "LA", "NA", reg, code, reg + la]
                  + [s[k] for k in HEADERS[8:]])
        rs = regions.setdefault(reg, dict.fromkeys(HEADERS[8:], 0))
        for k in HEADERS[8:]:
            rs[k] += true[k]          # the region rows stay the true sums
    for reg, s in regions.items():
        s = dict(s, **(region or {}).get(reg, {}))
        ws.append([reg, "England", "Region", "Region", "Region", "England",
                   "England", "England" + reg] + [s[k] for k in HEADERS[8:]])
    for r in extra_rows:
        ws.append(list(r))
    for _ in range(trailing):
        ws.append([None] * len(header))
    for r in after_trailing:
        ws.append(list(r))
    for s in (sheets or {}).get("remove", ()):
        del wb[s]
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


TITLE = "Registered provider social housing stock and rents in England {} to {}"
BASE = ("/government/statistics/registered-provider-social-housing-stock-"
        "and-rents-in-england-{}-to-{}")


def collection(*titles):
    docs = []
    for t in titles:
        slug = t.lower().replace(" ", "-").replace("(", "").replace(")", "")
        docs.append({"title": t, "base_path": "/government/statistics/" + slug})
    return {"links": {"documents": docs}}


def release_page(y1, y2, attachments, *, title=None, history=(
        ("2025-10-28T00:00:00Z", "First published."),)):
    return {"title": title or TITLE.format(y1, y2),
            "first_published_at": f"{y2}-10-28T00:00:00+00:00",
            "public_updated_at": f"{y2}-10-28T00:00:00+00:00",
            "details": {"attachments": [{"title": t, "url": u}
                                        for t, u in attachments],
                        "change_history": [{"public_timestamp": ts,
                                            "note": n} for ts, n in history]}}


TOOL_TITLE = "Registered providers look-up tool"


class Tmp(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def tool(self, name="t.xlsx", **kw):
        return write_tool(self.root / name, **kw)

    def read(self, **kw):
        return m.read_tool(self.tool(**kw))


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------

class Discovery(unittest.TestCase):

    def test_newest_release_is_chosen(self):
        c = collection(TITLE.format(2023, 2024), TITLE.format(2024, 2025),
                       TITLE.format(2019, 2020), "Something else")
        y1, y2, base = m.latest_release(c)
        self.assertEqual((y1, y2), (2024, 2025))
        self.assertTrue(base.endswith("2024-to-2025"))
        self.assertEqual(m.release_for_years(c, 2023, 2024)[:2], (2023, 2024))

    def test_two_pages_for_one_year_raise_listing_titles(self):
        c = collection(TITLE.format(2024, 2025), TITLE.format(2024, 2025))
        with self.assertRaises(ValueError) as e:
            m.latest_release(c)
        self.assertIn(TITLE.format(2024, 2025), str(e.exception))
        with self.assertRaises(ValueError):
            m.release_for_years(c, 2024, 2025)

    def test_a_renamed_series_raises_listing_titles(self):
        c = collection("Registered provider social housing statistics 2025")
        with self.assertRaises(ValueError) as e:
            m.latest_release(c)
        self.assertIn("social housing statistics 2025", str(e.exception))
        # a newer release whose title no longer fits is not passed over
        c = collection(TITLE.format(2024, 2025),
                       "Registered provider social housing stock and rents "
                       "in England 2025/26")
        with self.assertRaises(ValueError) as e:
            m.latest_release(c)
        self.assertIn("2025/26", str(e.exception))
        with self.assertRaises(ValueError):
            m.release_for_years(collection(TITLE.format(2023, 2024)), 2024,
                                2025)

    def test_every_newer_title_variant_halts_naming_it(self):
        held = TITLE.format(2024, 2025)
        variants = [
            "Registered providers social housing stock and rents in England "
            "2025 to 2026",
            "Registered provider social housing stock and rents in England "
            "2025-26",
            "Registered provider social housing stock and rents in England "
            "2025/26",
            "Registered provider social housing stock and rents in England "
            "2025–26",
            "Registered provider social housing stock and rents in England "
            "2025-2026",
            "Social housing stock and rents in England 2025 to 2026",
            "Registered providers: social housing stock 2026",
            "Registered provider social housing stock and rents in England",
        ]
        for v in variants:
            with self.subTest(title=v):
                with self.assertRaises(ValueError) as e:
                    m.latest_release(collection(held, TITLE.format(2023, 2024),
                                                v))
                self.assertIn(v, str(e.exception))
                self.assertIn("looks like the series", str(e.exception))

    def test_unrelated_and_older_titles_do_not_halt(self):
        c = collection(TITLE.format(2024, 2025), "Something else",
                       "Registered providers social housing 2019/20 archive",
                       "Guidance 2022")
        self.assertEqual(m.latest_release(c)[:2], (2024, 2025))

    def test_release_for_years_does_not_fall_back_to_a_variant(self):
        v = ("Registered providers social housing stock and rents in England "
             "2025 to 2026")
        c = collection(TITLE.format(2024, 2025), v)
        with self.assertRaises(ValueError) as e:
            m.release_for_years(c, 2025, 2026)
        self.assertIn(v, str(e.exception))
        self.assertIn("not used", str(e.exception))
        self.assertEqual(m.release_for_years(c, 2024, 2025)[:2], (2024, 2025))

    def test_nothing_newer_says_which_release_it_saw(self):
        c = collection(TITLE.format(2023, 2024), TITLE.format(2024, 2025))
        out = m.discovery_note(c, ["2025-03-31"], today=date(2026, 10, 9))
        self.assertEqual(len(out), 1, out)
        self.assertIn("2024 to 2025", out[0])
        self.assertIn("2025-03-31", out[0])
        self.assertIn("no newer release is listed", out[0])
        self.assertEqual(m.discovery_note(c, ["2024-03-31"]), [])
        self.assertEqual(m.discovery_note(collection("x"), ["2025-03-31"]), [])

    def test_past_the_announced_date_with_nothing_newer_warns(self):
        c = collection(TITLE.format(2024, 2025))
        for today in (date(2026, 10, 27), date(2026, 11, 3)):
            out = m.discovery_note(c, ["2025-03-31"], today=today)
            self.assertEqual(len(out), 2, out)
            self.assertTrue(out[1].startswith("WARNING"), out)
            self.assertIn("27 October 2026", out[1])
            self.assertIn("2025 to 2026", out[1])
        # the day before: the note, no warning
        out = m.discovery_note(c, ["2025-03-31"], today=date(2026, 10, 26))
        self.assertEqual(len(out), 1)
        # the new release is listed: nothing to warn about
        c2 = collection(TITLE.format(2024, 2025), TITLE.format(2025, 2026))
        self.assertEqual(m.discovery_note(c2, ["2025-03-31"],
                                          today=date(2026, 11, 3)), [])

    def test_the_look_up_tool_attachment(self):
        page = release_page(2024, 2025, [
            ("Registered provider social housing in England (PDF)", "a.pdf"),
            (TOOL_TITLE, "https://assets.example/media/1/RP_TOOL.xlsx"),
            ("Registered providers additional tables", "b.xlsx")])
        att = m.lookup_tool_attachment(page)
        self.assertEqual(att["url"], "https://assets.example/media/1/RP_TOOL.xlsx")
        self.assertEqual(m.page_years(page), (2024, 2025))
        self.assertEqual(m.change_history(page),
                         [("2025-10-28", "First published.")])

    def test_no_or_two_look_up_tools_raise_listing_titles(self):
        for atts in ([("Registered providers additional tables", "b.xlsx")],
                     [(TOOL_TITLE, "x/a.xlsx"), (TOOL_TITLE, "x/b.xlsx")],
                     [(TOOL_TITLE, "x/a.ods")]):
            with self.subTest(atts=atts), self.assertRaises(ValueError) as e:
                m.lookup_tool_attachment(release_page(2024, 2025, atts))
            self.assertIn(atts[0][0], str(e.exception))


# ---------------------------------------------------------------------------
# The look-up tool
# ---------------------------------------------------------------------------

class ReadTool(Tmp):

    def test_identity_rows_and_grains(self):
        t = self.read(trailing=5)
        self.assertEqual((t["y1"], t["y2"]), (2024, 2025))
        self.assertEqual(t["stock_date"], "2025-03-31")
        self.assertEqual(t["edition"], "2024 to 2025")
        self.assertEqual(t["published"], date(2025, 11, 1))
        self.assertEqual(t["version"], "1.1")
        self.assertEqual(t["rank"], (date(2025, 11, 1), (1, 1)))
        self.assertEqual(len(t["providers"]), len(CODES) * ROWS_PER_AREA)
        self.assertEqual(sorted(t["subtotals"]), sorted(CODES))
        self.assertEqual(sorted(t["regions"]),
                         ["North East", "Yorkshire and The Humber"])
        p = t["providers"][0]
        self.assertEqual((p["rp_code"], p["la_code"], p["rp_type"]),
                         ("4865", "E06000001", "Large"))
        self.assertEqual(t["empty_rows"], 5)

    def test_introduction_year_differs_from_the_source_line(self):
        with self.assertRaises(ValueError) as e:
            self.read(title="RP social housing by local authority area (SDR "
                      "and LADR data) 2024")
        self.assertIn("2024", str(e.exception))
        with self.assertRaises(ValueError):
            self.read(source="Statistical Data Return (SDR)/Local Authority "
                      "Data Return (LADR) 1 April 2023 to 31 March 2025")

    def test_version_history_last_row_differs_from_the_introduction(self):
        with self.assertRaises(ValueError) as e:
            self.read(history=[(1, "October 2025", "Original release.")])
        self.assertIn("Version History", str(e.exception))
        with self.assertRaises(ValueError):
            self.read(history=[(1, "October 2025", "x"),
                               (1.1, "December 2025", "y")])
        with self.assertRaises(ValueError):
            self.read(history_tail=[["Publication date: November 2025"],
                                    ["Version: 1.2"]])

    def test_unknown_and_missing_headers_raise(self):
        with self.assertRaises(ValueError) as e:
            self.read(header=HEADERS_2024)
        self.assertIn("CLCRR025_LA_GN_SC_Own", str(e.exception))
        self.assertIn("LA_GN_SC_Own", str(e.exception))
        with self.assertRaises(ValueError) as e:
            self.read(header=HEADERS[:-1] + ["Extra"])
        self.assertIn("Extra", str(e.exception))
        self.assertIn("LA_LCHO_Less_100_Eqty_Own", str(e.exception))

    def test_a_provider_row_with_la_code_na_raises(self):
        add = [["Clarion Housing Association Limited", "4865", "Large",
                "Long Form", "Signed_Off", "N/A", (1, 0, 0, 0)]]
        with self.assertRaises(ValueError) as e:
            self.read(add=add)
        self.assertIn("N/A", str(e.exception))

    def test_an_unknown_rp_type_raises(self):
        add = [["X Limited", "X1", "Medium", "Long Form", "Signed_Off",
                "E06000001", (1, 0, 0, 0)]]
        with self.assertRaises(ValueError) as e:
            self.read(add=add)
        self.assertIn("Medium", str(e.exception))

    def test_a_non_empty_row_with_no_rp_code_raises(self):
        with self.assertRaises(ValueError) as e:
            self.read(extra_rows=[["Stray note"] + [None] * 12])
        self.assertIn("RP_Code", str(e.exception))
        # an empty row followed by data is not trailing
        with self.assertRaises(ValueError):
            self.read(after_trailing=[["Stray"] + [None] * 12])

    def test_bad_stock_cells_halt_and_a_published_zero_stays_zero(self):
        for v in (None, "x", "", -1, 1.5, True):
            with self.subTest(v=v), self.assertRaises(ValueError):
                self.read(values={("L3076", "E06000002"): {"LA_SHHOP": v,
                                                           TOTAL: 5}})
        t = self.read()
        zero = [p for p in t["providers"] if p["rp_code"] == larp_code(CODES[0])]
        self.assertEqual(zero[0]["values"]["total_social_stock"], 0)
        self.assertEqual(m.stock_cell(0, "w"), 0)
        self.assertEqual(m.stock_cell(7.0, "w"), 7)

    def test_stock_cell_has_no_none_path(self):
        for v in (None, float("nan"), " ", "[x]"):
            with self.subTest(v=v), self.assertRaises(ValueError) as e:
                m.stock_cell(v, "STOCK_BY_LA row 9 LA_SHHOP")
            self.assertIn("row 9", str(e.exception))

    def test_missing_sheets_raise(self):
        for s in ("Version History", "STOCK_BY_LA", "Introduction and Contents"):
            with self.subTest(s=s), self.assertRaises(ValueError) as e:
                self.read(sheets={"remove": [s]})
            self.assertIn(s, str(e.exception))

    def test_the_2024_file_shape_halts_on_its_header(self):
        """The real 2024 tool (held, not loaded) if present: halts at the
        header check, naming the prefixed headers."""
        real = m.RAW_DIR / "RP_COMBINED_TOOL_2024_FINAL_V1_Locked.xlsx"
        if not real.exists():
            self.skipTest("2024 tool not on disk")
        with self.assertRaises(ValueError) as e:
            m.read_tool(real)
        self.assertIn("CLCRR025_LA_GN_SC_Own", str(e.exception))


class Reconcile(Tmp):

    def test_the_fixture_reconciles(self):
        with mock.patch.object(m, "EXPECTED_AREAS", len(CODES)), \
                mock.patch.object(m, "EXPECTED_REGIONS", 2):
            t = self.read()
            self.assertEqual(m.reconcile(t), [])
            eng = m.england_totals(t)
            self.assertEqual(eng["total_social_stock"],
                             sum(p["values"]["total_social_stock"]
                                 for p in t["providers"]))

    def test_a_planted_provider_cell_off_by_one_halts(self):
        t = self.read()
        p = [x for x in t["providers"] if x["rp_code"] == "L0055"
             and x["la_code"] == "E06000002"][0]
        p["values"]["general_needs_self_contained"] += 1
        with mock.patch.object(m, "EXPECTED_AREAS", len(CODES)), \
                mock.patch.object(m, "EXPECTED_REGIONS", 2):
            out = m.reconcile(t)
        self.assertEqual(len(out), 1, out)
        self.assertIn("E06000002", out[0])

    def test_one_la_subtotal_off_halts(self):
        with mock.patch.object(m, "EXPECTED_AREAS", len(CODES)), \
                mock.patch.object(m, "EXPECTED_REGIONS", 2):
            t = self.read(subtotal={"E08000019": {"LA_SHHOP": 1}})
            out = m.reconcile(t)
        self.assertTrue(any("E08000019" in x for x in out), out)
        self.assertTrue(any("Yorkshire and The Humber" in x for x in out), out)

    def test_wrong_number_of_authorities_or_regions(self):
        t = self.read()
        self.assertTrue(any("296" in x for x in m.reconcile(t)))


# ---------------------------------------------------------------------------
# Records and geography
# ---------------------------------------------------------------------------

PROV = {"file_edition": "2024 to 2025",
        "file_publication_date": date(2025, 11, 1),
        "file_source_url": "https://assets.example/media/1/t.xlsx",
        "file_name": "t.xlsx",
        "file_release_page_url": "https://www.gov.uk" + BASE.format(2024, 2025)}


class Records(Tmp):

    def test_build_rows(self):
        t = self.read(values={("H0375", "E06000001"): {"SDR_Size": None,
                                                       "Survey_Status": " "}})
        recs = m.build_rows(t, lambda c: c, PROV)
        self.assertEqual(len(recs), len(CODES) * ROWS_PER_AREA)
        r = [x for x in recs if x["rp_code"] == "4865"
             and x["lad24cd"] == "E06000001"][0]
        self.assertEqual(r["provider_type"], "PRP")
        self.assertEqual(r["rp_size_band"], "Long Form")
        self.assertEqual(r["stock_date"], "2025-03-31")
        self.assertEqual(r["publisher_la_code"], "E06000001")
        self.assertEqual(r["la_name"], "Hartlepool")
        self.assertEqual(r["file_name"], "t.xlsx")
        larp = [x for x in recs if x["rp_code"] == larp_code(CODES[1])][0]
        self.assertEqual(larp["provider_type"], "LARP")
        small = [x for x in recs if x["rp_code"] == "H0375"
                 and x["lad24cd"] == "E06000001"][0]
        self.assertIsNone(small["rp_size_band"])
        self.assertIsNone(small["survey_status"])

    def test_total_not_equal_to_the_parts_halts(self):
        t = self.read(values={("4865", "E06000002"): {TOTAL: 3037,
                                                      "LA_GN_SC_Own": 3000,
                                                      "LA_GN_BSp_Own": 0,
                                                      "LA_SHHOP": 36,
                                                      "LA_LCHO_Less_100_Eqty_Own": 0}})
        with self.assertRaises(ValueError) as e:
            m.build_rows(t, lambda c: c, PROV)
        self.assertIn("3,037", str(e.exception))

    def test_duplicate_key_halts(self):
        t = self.read(add=[["Clarion Housing Association Limited", "4865",
                            "Large", "Long Form", "Signed_Off", "E06000001",
                            (1, 0, 0, 0)]])
        with self.assertRaises(ValueError) as e:
            m.build_rows(t, lambda c: c, PROV)
        self.assertIn("4865", str(e.exception))
        # two publisher codes reaching one lad24cd: a duplicate too
        t = self.read(codes=["E06000001", "E08000016"],
                      add=[["Clarion Housing Association Limited", "4865",
                            "Large", "Long Form", "Signed_Off", "E08000038",
                            (1, 0, 0, 0)]])
        with self.assertRaises(ValueError):
            m.build_rows(t, lambda c: {"E08000038": "E08000016"}.get(c, c),
                         PROV)

    def test_rows_content_sha_ignores_provenance(self):
        t = self.read()
        a = m.build_rows(t, lambda c: c, PROV)
        b = m.build_rows(t, lambda c: c, dict(PROV, file_name="other.xlsx",
                                              file_source_url="file:x"))
        self.assertEqual(m.rows_content_sha(a), m.rows_content_sha(b))
        b[0]["general_needs_bedspaces"] += 1
        self.assertNotEqual(m.rows_content_sha(a), m.rows_content_sha(b))
        self.assertEqual(m.rows_content_sha(list(reversed(a))),
                         m.rows_content_sha(a))


# ---------------------------------------------------------------------------
# Not counted (rule 1): a Small PRP's LCHO cell (decided 2026-10-10)
# ---------------------------------------------------------------------------

LCHO = "low_cost_home_ownership"
REASON = "low_cost_home_ownership=not_counted_for_this_provider_type"


def _vals(lcho=0, gn=5):
    return {"total_social_stock": gn + 1 + lcho,
            "general_needs_self_contained": gn,
            "general_needs_bedspaces": 0,
            "supported_housing_and_older_people": 1, LCHO: lcho}


class NotCounted(Tmp):

    def test_constants_name_the_publisher_note_and_the_reason(self):
        self.assertEqual(m.LCHO_NOTE, "Unit counts for LCHO are for LARPs "
                                      "and Large PRPs only.")
        self.assertEqual(m.NOT_COUNTED_TYPES, ("Small",))
        self.assertEqual(m.NOT_COUNTED_COLUMN, LCHO)
        self.assertEqual(m.NOT_COUNTED_REASON, REASON)
        self.assertEqual(m.ROW_COLS, m.COMPARED + ("null_reasons",))

    def test_a_small_prp_zero_becomes_null_with_the_reason(self):
        v, why = m.not_counted("Small", _vals(0), "row 9", True)
        self.assertIsNone(v[LCHO])
        self.assertEqual(why, REASON)
        # the other columns, the total included, are as published
        for c in m.STOCK_COLUMNS[:4]:
            self.assertEqual(v[c], _vals(0)[c], c)

    def test_covered_types_keep_their_published_zero(self):
        for typ in ("Large", "LARP"):
            v, why = m.not_counted(typ, _vals(0), "row 9", True)
            self.assertEqual(v[LCHO], 0, typ)
            self.assertIsNone(why)
            v, why = m.not_counted(typ, _vals(7), "row 9", True)
            self.assertEqual(v[LCHO], 7)
            self.assertIsNone(why)

    def test_a_genuine_null_stays_null(self):
        # covered: NULL stays NULL, no reason invented; not counted: NULL
        # stays NULL with the reason
        v, why = m.not_counted("Large", dict(_vals(), **{LCHO: None}), "r", True)
        self.assertIsNone(v[LCHO])
        self.assertIsNone(why)
        v, why = m.not_counted("Small", dict(_vals(), **{LCHO: None}), "r", True)
        self.assertIsNone(v[LCHO])
        self.assertEqual(why, REASON)

    def test_only_the_lcho_column_is_touched(self):
        zeros = {c: 0 for c in m.STOCK_COLUMNS}
        v, _ = m.not_counted("Small", zeros, "r", True)
        self.assertEqual([c for c in m.STOCK_COLUMNS if v[c] is None], [LCHO])

    def test_a_small_prp_with_a_published_lcho_number_halts(self):
        with self.assertRaises(ValueError) as e:
            m.not_counted("Small", _vals(3), "STOCK_BY_LA row 9", True)
        self.assertIn("row 9", str(e.exception))
        self.assertIn("LARPs and Large PRPs only", str(e.exception))

    def test_the_rule_needs_the_note_in_the_file(self):
        with self.assertRaises(ValueError) as e:
            m.not_counted("Small", _vals(0), "row 9", False)
        self.assertIn("Area Summary", str(e.exception))
        # a covered row does not need it
        self.assertEqual(m.not_counted("Large", _vals(0), "r", False)[0][LCHO],
                         0)

    def test_read_tool_finds_the_note(self):
        self.assertTrue(self.read()["lcho_note"])
        self.assertFalse(self.read(name="n.xlsx", lcho_note=False)["lcho_note"])
        self.assertFalse(self.read(name="s.xlsx",
                                   sheets={"remove": ["Area Summary"]})
                         ["lcho_note"])

    def test_build_rows_applies_the_rule_to_small_rows_only(self):
        recs_ = m.build_rows(self.read(), lambda c: c, PROV)
        small = [r for r in recs_ if r["rp_code"] == "H0375"]
        self.assertEqual(len(small), len(CODES))
        for r in small:
            self.assertIsNone(r[LCHO])
            self.assertEqual(r["null_reasons"], REASON)
            # the published total, which leaves out the provider's LCHO
            self.assertEqual(r["total_social_stock"], 3)
        others = [r for r in recs_ if r["rp_code"] != "H0375"]
        self.assertTrue(all(r[LCHO] is not None and r["null_reasons"] is None
                            for r in others))
        # a covered published zero stays 0 (the first LARP owns nothing)
        larp0 = [r for r in others if r["rp_code"] == larp_code(CODES[0])][0]
        self.assertEqual((larp0[LCHO], larp0["total_social_stock"]), (0, 0))

    def test_build_rows_without_the_rule_is_as_published(self):
        recs_ = m.build_rows(self.read(), lambda c: c, PROV, rule=False)
        self.assertTrue(all(r[LCHO] is not None and r["null_reasons"] is None
                            for r in recs_))
        self.assertEqual({r[LCHO] for r in recs_ if r["rp_code"] == "H0375"},
                         {0})

    def test_build_rows_halts_on_a_small_lcho_number_or_a_missing_note(self):
        t = self.read(values={("H0375", "E06000002"):
                              {"LA_LCHO_Less_100_Eqty_Own": 4}})
        with self.assertRaises(ValueError) as e:
            m.build_rows(t, lambda c: c, PROV)
        self.assertIn("H0375", str(e.exception))
        with self.assertRaises(ValueError):
            m.build_rows(self.read(name="n.xlsx", lcho_note=False),
                         lambda c: c, PROV)
        # without the rule, both read as published
        self.assertTrue(m.build_rows(t, lambda c: c, PROV, rule=False))

    def test_the_total_check_reads_the_published_cells(self):
        # the total is checked before the rule: a Small row whose total is
        # not its parts halts as before
        t = self.read(values={("H0375", "E06000002"): {TOTAL: 9}})
        with self.assertRaises(ValueError) as e:
            m.build_rows(t, lambda c: c, PROV)
        self.assertIn("not the sum", str(e.exception))

    def test_content_sha_without_reasons_is_unchanged(self):
        # a record set with no null_reasons hashes exactly as before the
        # rule (edition 1's stored sha256 stays reproducible)
        recs_ = m.build_rows(self.read(), lambda c: c, PROV, rule=False)
        lines = ["|".join([str(r["rp_code"]), str(r["lad24cd"])]
                          + ["" if r[c] is None else str(r[c])
                             for c in m.COMPARED])
                 for r in sorted(recs_, key=lambda r: (r["rp_code"],
                                                       r["lad24cd"]))]
        want = hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()
        self.assertEqual(m.rows_content_sha(recs_), want)
        ruled = m.build_rows(self.read(), lambda c: c, PROV)
        self.assertNotEqual(m.rows_content_sha(ruled), want)
        # the reason is part of the content
        other = [dict(r) for r in ruled]
        for r in other:
            if r["null_reasons"]:
                r["null_reasons"] = "low_cost_home_ownership=x"
        self.assertNotEqual(m.rows_content_sha(other),
                            m.rows_content_sha(ruled))

    def test_sums_leave_a_null_out(self):
        recs_ = m.build_rows(self.read(), lambda c: c, PROV)
        plain = m.build_rows(self.read(), lambda c: c, PROV, rule=False)
        self.assertEqual(m._national(recs_), m._national(plain))
        self.assertEqual(m._area_sums(recs_), m._area_sums(plain))


class Flips(Tmp):
    """0/NULL changes against the tip, and the named acknowledgement."""

    def setUp(self):
        super().setUp()
        self.plain = m.build_rows(self.read(), lambda c: c, PROV, rule=False)
        self.ruled = m.build_rows(self.read(), lambda c: c, PROV)
        p = mock.patch.dict(m.ACKNOWLEDGED_FLIPS, {"t": dict(
            m.ACKNOWLEDGED_FLIPS["not-counted-lcho-2025"],
            cells=len(CODES))})
        p.start()
        self.addCleanup(p.stop)

    def test_the_real_acknowledgement(self):
        a = m.ACKNOWLEDGED_FLIPS["not-counted-lcho-2025"]
        self.assertEqual((a["stock_date"], a["column"], a["cells"],
                          a["reason"]), ("2025-03-31", LCHO, 2738, REASON))

    def test_flips_are_the_small_lcho_cells(self):
        fl = m._flips(self.ruled, self.plain)
        self.assertEqual(len(fl), len(CODES))
        self.assertTrue(all(k[0] == "H0375" and c == LCHO and a == 0
                            and b is None for k, c, a, b in fl))
        self.assertEqual(m._flips(self.plain, self.ruled)[0][2:], (None, 0))
        self.assertEqual(m._flips(self.ruled, self.ruled), [])

    def test_a_value_change_is_not_a_flip(self):
        other = [dict(r) for r in self.plain]
        other[0]["general_needs_self_contained"] += 1
        self.assertEqual(m._flips(other, self.plain), [])

    def test_the_acknowledgement_covers_exactly_its_cells(self):
        fl = m._flips(self.ruled, self.plain)
        self.assertEqual(m.ack_problems("t", "2025-03-31", fl, self.ruled),
                         [])
        # the count must match exactly
        self.assertTrue(m.ack_problems("t", "2025-03-31", fl[:-1],
                                       self.ruled))
        with mock.patch.dict(m.ACKNOWLEDGED_FLIPS["t"],
                             {"cells": len(CODES) + 1}):
            out = m.ack_problems("t", "2025-03-31", fl, self.ruled)
        self.assertTrue(any("expects" in x for x in out), out)
        # another stock date
        self.assertTrue(m.ack_problems("t", "2024-03-31", fl, self.ruled))
        # no flips at all
        self.assertTrue(m.ack_problems("t", "2025-03-31", [], self.ruled))

    def test_the_acknowledgement_refuses_other_flips(self):
        # NULL -> 0 (the other direction)
        back = m._flips(self.plain, self.ruled)
        self.assertTrue(m.ack_problems("t", "2025-03-31", back, self.plain))
        # a 0 -> NULL in the right column without the rule's reason
        bad = [dict(r) for r in self.ruled]
        for r in bad:
            r["null_reasons"] = None
        fl = m._flips(bad, self.plain)
        out = m.ack_problems("t", "2025-03-31", fl, bad)
        self.assertTrue(any("reason" in x for x in out), out)
        # a flip in another column
        fl2 = [(fl[0][0], "general_needs_bedspaces", 0, None)] + fl[1:]
        self.assertTrue(m.ack_problems("t", "2025-03-31", fl2, self.ruled))

    def test_revert_puts_the_acknowledged_cells_back(self):
        fl = m._flips(self.ruled, self.plain)
        back = m._revert(self.ruled, fl, self.plain)
        self.assertFalse(m._differs(back, self.plain))
        self.assertTrue(m._differs(self.ruled, self.plain))
        # the inputs are not changed
        self.assertTrue(any(r[LCHO] is None for r in self.ruled))


class _Cur:
    """A cursor answering la_boundaries and la_code_lookup reads."""

    def __init__(self, bounds):
        self.bounds = bounds
        self.result = []

    def execute(self, sql, args=None):
        s = " ".join(sql.split())
        if "change_type = 'recode'" in s:
            self.result = [("E08000038", "E08000016"),
                           ("E08000039", "E08000019")]
        elif "FROM public.la_boundaries" in s or "FROM la_boundaries" in s:
            self.result = [(c,) for c in self.bounds]
        else:
            raise AssertionError(f"unexpected SQL {s}")

    def fetchall(self):
        return list(self.result)


class Geography(unittest.TestCase):

    def test_old_codes_resolve_and_new_codes_halt(self):
        cur = _Cur(CODES + ["E07000223"])
        rmap, problems = m.resolve_codes(cur, set(CODES), "2025-03-31")
        self.assertEqual(problems, [])
        self.assertEqual(rmap["E08000016"], "E08000016")
        _, problems = m.resolve_codes(cur, {"E06000001", "E08000038"},
                                      "2025-03-31")
        self.assertTrue(any("E08000038" in p and "'old'" in p
                            for p in problems), problems)

    def test_an_unknown_code_is_unexplained(self):
        cur = _Cur(CODES)
        _, problems = m.resolve_codes(cur, {"E06000001", "E07000999"},
                                      "2025-03-31")
        self.assertTrue(any("UNEXPLAINED E07000999" in p for p in problems),
                        problems)


# ---------------------------------------------------------------------------
# Ranks, ledger and planning
# ---------------------------------------------------------------------------

class Ranks(Tmp):

    def test_rank_from_the_file_never_the_name(self):
        path = write_tool(self.root / "RP_COMBINED_TOOL_2025_FINAL_V9.9.xlsx",
                          version="1.0", month="October 2025")
        t = m.read_tool(path)
        self.assertEqual(m.release_rank(t), (date(2025, 10, 1), (1,)))
        self.assertEqual(m.release_rank(
            "as loaded: RP_COMBINED_TOOL_2025_FINAL_V1.1.xlsx; version 1.1; "
            "dated 2025-11"), (date(2025, 11, 1), (1, 1)))
        self.assertIsNone(m.release_rank("RP_COMBINED_TOOL_2025_FINAL_V1.1.xlsx"))
        self.assertIsNone(m.release_rank(None))
        self.assertLess(m.release_rank("version 1.0; dated 2025-10"),
                        m.release_rank("version 1.1; dated 2025-11"))
        self.assertLess(m.release_rank("version 1.9; dated 2025-11"),
                        m.release_rank("version 1.10; dated 2025-11"))
        self.assertEqual(m.release_rank("version 1; dated 2025-10"),
                         m.release_rank("version 1.0; dated 2025-10"))

    def test_ledger_source_round_trip(self):
        r = (date(2025, 11, 1), (1, 1))
        s = m.ledger_source("https://x/y.xlsx", r, "1.1", "2025-03-31")
        self.assertEqual(m.parse_ledger_source(s),
                         ("https://x/y.xlsx", r, ["2025-03-31"]))
        self.assertEqual(m.release_rank(s), r)

    def test_content_sha256(self):
        p = self.root / "f.bin"
        p.write_bytes(b"abc")
        self.assertEqual(m.content_sha256(p), hashlib.sha256(b"abc").hexdigest())


class Ledger(unittest.TestCase):
    R = (date(2025, 11, 1), (1, 1))
    OLD = (date(2025, 10, 1), (1,))
    P = "2025-03-31"

    def src(self, where="u", rank=None, period=None, version="1.1"):
        return m.ledger_source(where, rank or self.R, version,
                               period or self.P)

    def test_ledger_skip_needs_the_pair_for_the_period(self):
        s = self.src()
        self.assertEqual(m.ledger_complete({self.P: {(s, "h")}}, "u", "h",
                                           {self.P: self.R}), [self.P])
        # another sha, another place, or the pair recorded under another
        # period: read the file
        self.assertIsNone(m.ledger_complete({self.P: {(s, "h2")}}, "u", "h",
                                            {self.P: self.R}))
        self.assertIsNone(m.ledger_complete({self.P: {(s, "h")}}, "v", "h",
                                            {self.P: self.R}))
        self.assertIsNone(m.ledger_complete({"2024-03-31": {(s, "h")}}, "u",
                                            "h", {self.P: self.R}))
        # a stranded period always makes the file be read
        self.assertIsNone(m.ledger_complete({self.P: {(s, "h")}}, "u", "h",
                                            {self.P: self.R},
                                            stranded=[self.P]))

    def test_plan_periods(self):
        s = self.src()
        new, rev, skip = m.plan_periods([self.P], self.R, {}, {}, s, "h",
                                        recheck=None, allow_older=False)
        self.assertEqual((new, rev, skip), ([self.P], [], {}))
        new, rev, skip = m.plan_periods([self.P], self.R, {self.P: self.R},
                                        {self.P: {(s, "h")}}, s, "h",
                                        recheck=None, allow_older=False)
        self.assertEqual(rev, [])
        self.assertIn("checked", skip[self.P])
        new, rev, skip = m.plan_periods([self.P], self.R, {self.P: self.R},
                                        {self.P: {(s, "h")}}, s, "h",
                                        recheck=self.P, allow_older=False)
        self.assertEqual(rev, [self.P])
        new, rev, skip = m.plan_periods([self.P], self.OLD, {self.P: self.R},
                                        {}, s, "h", recheck=None,
                                        allow_older=False)
        self.assertTrue(skip[self.P].startswith("older"))
        new, rev, skip = m.plan_periods([self.P], self.OLD, {self.P: self.R},
                                        {}, s, "h", recheck=None,
                                        allow_older=True)
        self.assertEqual(rev, [self.P])
        with self.assertRaises(ValueError):
            m.plan_periods([self.P], self.R, {}, {}, s, "h",
                           recheck="2024-03-31", allow_older=False)

    def test_tip_ranks_take_a_newer_ledger_file(self):
        newer = (date(2025, 12, 1), (1, 2))
        s = self.src(rank=newer, version="1.2")
        out = m.tip_ranks({self.P: {"rank": self.R}}, {self.P: {(s, "h")}})
        self.assertEqual(out[self.P], newer)


# ---------------------------------------------------------------------------
# Stop conditions
# ---------------------------------------------------------------------------

AREAS = [f"E06{i:06d}" for i in range(1, 51)]


def recs(base=(150, 10, 30, 10), *, small=False, change=None, drop=(),
         add=()):
    """One provider P1 per authority (base parts), plus a provider P2 of two
    units when small; change {(rp, area): (parts)}; drop keys; add keys
    (one unit)."""
    out = []
    for a in AREAS:
        rows = [("P1", base)] + ([("P2", (2, 0, 0, 0))] if small else [])
        for rp, parts in rows:
            if (rp, a) in drop:
                continue
            parts = (change or {}).get((rp, a), parts)
            out.append(_rec(rp, a, parts))
    for rp, a in add:
        out.append(_rec(rp, a, (1, 0, 0, 0)))
    return out


def _rec(rp, a, parts):
    g, b, s, lc = parts
    return {"rp_code": rp, "lad24cd": a, "total_social_stock": g + b + s + lc,
            "general_needs_self_contained": g, "general_needs_bedspaces": b,
            "supported_housing_and_older_people": s,
            "low_cost_home_ownership": lc}


class Stops(unittest.TestCase):

    def setUp(self):
        p = mock.patch.object(m, "EXPECTED_AREAS", len(AREAS))
        p.start()
        self.addCleanup(p.stop)

    def new(self, n, prev=None):
        return m.period_problems(n, None, prev, kind="new")

    def rev(self, n, tip=None):
        return m.period_problems(n, tip or recs(), None, kind="revised")

    def test_constants(self):
        self.assertEqual((m.NEW_TOTAL_PCT, m.NEW_SUPPORTED_PCT,
                          m.NEW_AREA_PCT, m.REV_NATIONAL_PCT, m.REV_MAX_AREAS,
                          m.REV_AREA_PCT, m.REV_KEYS_PCT),
                         (5, 10, 25, 1, 30, 10, 5))

    def test_new_fewer_authorities(self):
        self.assertEqual(self.new(recs()), [])
        out = self.new([r for r in recs() if r["lad24cd"] != AREAS[0]])
        self.assertTrue(any("authorities" in x for x in out), out)

    def test_new_national_total(self):
        prev = recs()                                   # total 10,000
        ok = recs(change={("P1", a): (160, 10, 30, 10) for a in AREAS})
        self.assertEqual(self.new(ok, prev), [])        # +5.0%
        over = recs(change={**{("P1", a): (160, 10, 30, 10) for a in AREAS},
                            ("P1", AREAS[0]): (161, 10, 30, 10)})
        out = self.new(over, prev)
        self.assertTrue(any("national total_social_stock" in x for x in out),
                        out)

    def test_new_national_supported(self):
        prev = recs()                                   # supported 1,500
        ok = recs(change={("P1", a): (150, 10, 33, 10) for a in AREAS})
        self.assertEqual(self.new(ok, prev), [])        # +10.0%
        over = recs(change={**{("P1", a): (150, 10, 33, 10) for a in AREAS},
                            ("P1", AREAS[0]): (150, 10, 34, 10)})
        out = self.new(over, prev)
        self.assertTrue(any("supported_housing_and_older_people" in x
                            for x in out), out)

    def test_new_authority_move(self):
        prev = recs()                                   # 200 each
        ok = recs(change={("P1", AREAS[3]): (200, 10, 30, 10)})
        self.assertEqual(self.new(ok, prev), [])        # 250: +25%
        over = recs(change={("P1", AREAS[3]): (201, 10, 30, 10)})
        out = self.new(over, prev)
        self.assertTrue(any(AREAS[3] in x for x in out), out)

    def test_revised_national(self):
        ok = recs(change={("P1", a): (153, 10, 30, 10) for a in AREAS[:25]})
        self.assertEqual(self.rev(ok), [])              # gn_sc +75: 1.0%
        over = recs(change={**{("P1", a): (153, 10, 30, 10)
                               for a in AREAS[:25]},
                            ("P1", AREAS[0]): (154, 10, 30, 10)})
        out = self.rev(over)
        self.assertTrue(any("national general_needs_self_contained" in x
                            for x in out), out)

    def test_revised_authorities_changed(self):
        ok = recs(change={("P1", a): (151, 10, 30, 10) for a in AREAS[:30]})
        self.assertEqual(self.rev(ok), [])
        over = recs(change={("P1", a): (151, 10, 30, 10) for a in AREAS[:31]})
        out = self.rev(over)
        self.assertTrue(any("31 authorities" in x for x in out), out)

    def test_revised_authority_move(self):
        ok = recs(change={("P1", AREAS[5]): (170, 10, 30, 10)})
        self.assertEqual(self.rev(ok), [])              # 220: +10%
        over = recs(change={("P1", AREAS[5]): (171, 10, 30, 10)})
        out = self.rev(over)
        self.assertTrue(any(AREAS[5] in x for x in out), out)

    def test_revised_provider_rows_added_or_removed(self):
        tip = recs(small=True)                          # 100 rows
        ok = recs(small=True, drop={("P2", a) for a in AREAS[:5]})
        self.assertEqual(self.rev(ok, tip), [])         # 5% removed
        over = recs(small=True, drop={("P2", a) for a in AREAS[:6]})
        out = self.rev(over, tip)                       # 6% fewer: partial
        self.assertTrue(any("removed" in x for x in out), out)
        ok = recs(small=True, add=[("P3", a) for a in AREAS[:5]])
        self.assertEqual(self.rev(ok, tip), [])
        over = recs(small=True, add=[("P3", a) for a in AREAS[:6]])
        self.assertTrue(any("added" in x for x in self.rev(over, tip)))

    def test_revised_fewer_authorities_is_a_partial_file(self):
        out = self.rev([r for r in recs() if r["lad24cd"] != AREAS[0]])
        self.assertTrue(any("authorities" in x for x in out), out)


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

class _Resp:
    def __init__(self, body, url):
        self.content = body
        self.url = url
        self.text = body.decode("latin-1")

    def raise_for_status(self):
        pass


class _Session:
    def __init__(self, body, final):
        self.body, self.final, self.calls = body, final, []

    def get(self, url, **kw):
        self.calls.append(url)
        return _Resp(self.body, self.final)


class Download(Tmp):

    def body(self, **kw):
        return write_tool(self.root / "src" / f"b{len(kw)}.xlsx",
                          **kw).read_bytes()

    def test_redirect_records_the_final_url(self):
        body = self.body()
        s = _Session(body, "https://assets.example/media/2/final.xlsx")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), quiet() as out:
            path, final = m.fetch("https://assets.example/media/1/old.xlsx",
                                  self.root / "raw", session=s)
        self.assertEqual(final, "https://assets.example/media/2/final.xlsx")
        self.assertEqual(path.name, "final.xlsx")
        self.assertEqual(path.read_bytes(), body)
        self.assertIn("redirected", out.getvalue())

    def test_same_name_different_content_saved_beside_never_overwritten(self):
        raw = self.root / "raw"
        raw.mkdir()
        old = raw / "f.xlsx"
        old.write_bytes(self.body())
        before = old.read_bytes()
        new_body = self.body(version="1.2", month="December 2025")
        s = _Session(new_body, "https://assets.example/media/3/f.xlsx")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), quiet():
            path, _ = m.fetch("https://assets.example/media/3/f.xlsx", raw,
                              session=s)
        sha8 = hashlib.sha256(new_body).hexdigest()[:8]
        self.assertEqual(path.name, f"f-{sha8}.xlsx")
        self.assertEqual(path.read_bytes(), new_body)
        self.assertEqual(old.read_bytes(), before)

    def test_an_existing_same_named_file_does_not_skip_the_download(self):
        raw = self.root / "raw"
        raw.mkdir()
        body = self.body()
        (raw / "f.xlsx").write_bytes(body)
        s = _Session(body, "https://assets.example/media/3/f.xlsx")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), quiet() as out:
            path, _ = m.fetch("https://assets.example/media/3/f.xlsx", raw,
                              session=s)
        self.assertEqual(s.calls, ["https://assets.example/media/3/f.xlsx"])
        self.assertEqual(path, raw / "f.xlsx")
        self.assertIn("same content", out.getvalue())

    def test_size_floor_format_and_sheet_checks_keep_nothing(self):
        raw = self.root / "raw"
        s = _Session(b"x" * 50, "https://assets.example/media/3/f.xlsx")
        with self.assertRaises(SystemExit), quiet():
            m.fetch("https://assets.example/media/3/f.xlsx", raw, session=s)
        s = _Session(b"x" * 5000, "https://assets.example/media/3/f.xlsx")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), \
                self.assertRaises(SystemExit), quiet():
            m.fetch("https://assets.example/media/3/f.xlsx", raw, session=s)
        no_sheet = self.body(sheets={"remove": ["STOCK_BY_LA"]})
        s = _Session(no_sheet, "https://assets.example/media/3/f.xlsx")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), \
                self.assertRaises(SystemExit) as e, quiet():
            m.fetch("https://assets.example/media/3/f.xlsx", raw, session=s)
        self.assertIn("STOCK_BY_LA", str(e.exception.code))
        s = _Session(self.body(), "https://assets.example/media/3/f.ods")
        with mock.patch.object(m, "MIN_FILE_BYTES", 10), \
                self.assertRaises(SystemExit), quiet():
            m.fetch("https://assets.example/media/3/f.ods", raw, session=s)
        self.assertEqual(sorted(p.name for p in raw.iterdir())
                         if raw.exists() else [], [])


class NoZero(unittest.TestCase):

    def test_no_none_to_zero_in_the_module(self):
        import check_loaders
        reasons = check_loaders.check_loader(Path(m.__file__),
                                             {"revises_back_series": True})
        self.assertFalse([r for r in reasons if "zero coercion" in r
                          or "commit" in r or "recode dict" in r], reasons)


if __name__ == "__main__":
    unittest.main()
