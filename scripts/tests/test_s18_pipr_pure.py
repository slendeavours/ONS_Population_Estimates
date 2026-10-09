"""Pure tests for s18_pipr_editions: edition names, the landing-page link,
Table 1 parsing on a small workbook built here, records, identity and hashes.
No database, no network.
"""
import sys
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402

import s18_pipr_editions as m  # noqa: E402

D = Decimal
SUFFIXES = [s for s, _, _ in m.BLOCKS]


def header(drop=()):
    cols = ["Time period", "Area code", "Area name", "Region or country name"]
    for s in SUFFIXES:
        for base in ("Index", "Monthly change", "Annual change",
                     "Rental price"):
            name = f"{base}{s}"
            if name not in drop:
                cols.append(name)
    return cols


def make_row(cols, period, code, *, rent=1000, index=100.0, change=1.0,
             overrides=None):
    """One Table 1 row: block b (0..8) gets rent+b, index+b, change+b."""
    out = {"Time period": period, "Area code": code, "Area name": "n",
           "Region or country name": "r"}
    for b, s in enumerate(SUFFIXES):
        out[f"Index{s}"] = index + b
        out[f"Monthly change{s}"] = 0.1
        out[f"Annual change{s}"] = change + b
        out[f"Rental price{s}"] = rent + b
    out.update(overrides or {})
    return [out.get(c) for c in cols]


def write_workbook(path, rows, *, drop=(), cover=True,
                   cover_text="The data tables in this spreadsheet were "
                   "originally published at 9:30am on 16 September 2026. "
                   "Crown copyright."):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cover sheet"
    if cover:
        ws["A1"] = "Price Index of Private Rents"
        ws["A5"] = cover_text
    wb.create_sheet("Contents")
    t = wb.create_sheet("Table 1")
    t.append(["Price Index of Private Rents, UK"])
    t.append(["This worksheet contains one table."])
    cols = header(drop)
    t.append(cols)
    for r in rows:
        t.append(make_row(cols, *r[:2], **(r[2] if len(r) > 2 else {})))
    wb.save(path)
    return path


P = datetime
ROWS = [
    (P(2024, 2, 1), "E06000001"),                 # before the window
    (P(2024, 3, 1), "E06000001"),
    (P(2024, 3, 1), "E08000038"),                 # new Barnsley code
    (P(2024, 3, 1), "W06000001"),                 # Wales: not read
    (P(2024, 3, 1), "K02000001"),                 # UK: not read
    (P(2024, 3, 1), "[z]"),                       # NI BRMA: not read
    (P(2024, 4, 1), "E06000001", {"rent": 1100}),
    (P(2024, 4, 1), "E08000038"),
]


class Names(unittest.TestCase):
    def test_edition_from_link_and_names(self):
        url = ("https://www.ons.gov.uk/file?uri=/economy/inflationandprice"
               "indices/datasets/priceindexofprivaterentsukmonthlypricestat"
               "istics/16september2026/priceindexofprivaterentsukmonthlyprice"
               "statistics13.xlsx")
        self.assertEqual(m.edition_from_link(url), "16september2026")
        self.assertEqual(m.edition_from_link("pipr_17june2026.xlsx"),
                         "17june2026")
        self.assertEqual(m.edition_from_link("pipr_01july2026.xlsx"),
                         "1july2026")
        for bad in ("pipr.xlsx", "", None, "x_32foo2026.xlsx"):
            with self.assertRaises(ValueError):
                m.edition_from_link(bad)

    def test_release_label_and_date(self):
        self.assertEqual(m.release_label("22july2026"),
                         "ONS PIPR 22july2026 edition")
        self.assertEqual(str(m.edition_date("16september2026")), "2026-09-16")
        with self.assertRaises(ValueError):
            m.release_label("2026-07")

    def test_find_workbook_link_first_and_never_hardcoded(self):
        html = ('<a href="/other">x</a>'
                '<a href="/file?uri=/economy/inflationandpriceindices/'
                'datasets/priceindexofprivaterentsukmonthlypricestatistics/'
                '16september2026/priceindexofprivaterentsukmonthlyprice'
                'statistics14.xlsx">new</a>'
                '<a href="/file?uri=/economy/inflationandpriceindices/'
                'datasets/priceindexofprivaterentsukmonthlypricestatistics/'
                '19august2026/priceindexofprivaterentsukmonthlyprice'
                'statistics13.xlsx">old</a>')
        link = m.find_workbook_link(html)
        self.assertTrue(link.startswith("https://www.ons.gov.uk/file?uri="))
        self.assertIn("16september2026", link)
        self.assertTrue(link.endswith("statistics14.xlsx"))
        with self.assertRaises(ValueError) as cm:
            m.find_workbook_link("<html>nothing here</html>")
        self.assertIn("nothing here", str(cm.exception))


class Parse(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def path(self, name="pipr_16september2026.xlsx"):
        return Path(self.tmp.name) / name

    def test_blocks_offsets_and_england_only(self):
        rows = m.parse_workbook(write_workbook(self.path(), ROWS))
        # 5 English rows x 9 blocks; Wales, UK and the NI marker are skipped
        self.assertEqual(len(rows), 5 * 9)
        self.assertEqual({r["area_code"] for r in rows},
                         {"E06000001", "E08000038"})
        by = {(r["area_code"], r["period"], r["breakdown_type"],
               r["category"]): r for r in rows}
        # block b carries rent+b, index+b, change+b: a shifted offset shows
        for b, (_, bt, cat) in enumerate(m.BLOCKS):
            r = by[("E06000001", "2024-03-01", bt, cat)]
            self.assertEqual((r["mean_rent"], r["rent_index"],
                              r["annual_pct_change"]),
                             (D(1000 + b), D(100 + b), D(1 + b)))
        self.assertEqual(by[("E06000001", "2024-04-01", "all", "all")]
                         ["mean_rent"], D(1100))

    def test_provisional_is_the_files_latest_month_only(self):
        rows = m.parse_workbook(write_workbook(self.path(), ROWS))
        flags = {(r["period"]): r["provisional"] for r in rows}
        self.assertEqual(flags, {"2024-02-01": False, "2024-03-01": False,
                                 "2024-04-01": True})

    def test_blank_and_markers_are_null_not_zero(self):
        rows = [(P(2024, 3, 1), "E06000001",
                 {"overrides": {"Rental price one bed": None,
                                "Index one bed": "[x]",
                                "Annual change one bed": "",
                                "Rental price two bed": 0,
                                "Annual change two bed": "[z]"}})]
        out = m.parse_workbook(write_workbook(self.path(), rows))
        one = next(r for r in out if r["category"] == "1_bed")
        self.assertEqual((one["mean_rent"], one["rent_index"],
                          one["annual_pct_change"]), (None, None, None))
        two = next(r for r in out if r["category"] == "2_bed")
        self.assertEqual(two["mean_rent"], D(0))      # a published 0 stays 0
        self.assertIsNone(two["annual_pct_change"])

    def test_rounds_half_up_on_the_published_decimal(self):
        rows = [(P(2024, 3, 1), "E06000001",
                 {"overrides": {"Annual change": 5.845,
                                "Index": 117.905,
                                "Annual change one bed": -1.505,
                                "Rental price": 1286}})]
        out = m.parse_workbook(write_workbook(self.path(), rows))
        r = next(r for r in out if r["category"] == "all")
        self.assertEqual((r["annual_pct_change"], r["rent_index"]),
                         (D("5.85"), D("117.91")))
        one = next(r for r in out if r["category"] == "1_bed")
        self.assertEqual(one["annual_pct_change"], D("-1.51"))

    def test_missing_block_column_raises_naming_it(self):
        p = write_workbook(self.path(), ROWS,
                           drop=("Rental price flat maisonette",))
        with self.assertRaises(ValueError) as cm:
            m.parse_workbook(p)
        self.assertIn("Rental price flat maisonette", str(cm.exception))
        self.assertIn("block", str(cm.exception))

    def test_no_table_1_raises(self):
        wb = openpyxl.Workbook()
        wb.active.title = "Cover sheet"
        p = self.path("nosheet.xlsx")
        wb.save(p)
        with self.assertRaises(ValueError):
            m.parse_workbook(p)

    def test_cover_edition_and_content_hash(self):
        p = write_workbook(self.path(), ROWS)
        self.assertEqual(m.cover_edition(p), "16september2026")
        q = write_workbook(self.path("b.xlsx"), ROWS, cover=False)
        self.assertIsNone(m.cover_edition(q))
        self.assertEqual(len(m.content_sha256(p)), 64)
        self.assertEqual(m.content_sha256(p), m.content_sha256(p))
        self.assertNotEqual(m.content_sha256(p), m.content_sha256(q))


class Records(unittest.TestCase):
    def rows(self):
        with tempfile.TemporaryDirectory() as d:
            return m.parse_workbook(write_workbook(Path(d) / "a.xlsx", ROWS))

    def test_resolution_both_forms_window_and_sort(self):
        rows = self.rows()
        # the new codes resolve through the map; old codes pass unchanged
        recs = m.build_records(rows, {"E08000038": "E08000016"})
        self.assertEqual({r["lad24cd"] for r in recs},
                         {"E06000001", "E08000016"})
        self.assertNotIn("2024-02-01", {r["period"] for r in recs})
        self.assertEqual(len(recs), 4 * 9)
        self.assertEqual(recs, sorted(recs, key=lambda x: (
            x["period"], x["lad24cd"], x["breakdown_type"], x["category"])))
        # the old form of the same area resolves to itself
        old = [dict(r, area_code="E08000016") if r["area_code"] ==
               "E08000038" else r for r in rows]
        recs2 = m.build_records(old, {"E08000038": "E08000016"})
        self.assertEqual({r["lad24cd"] for r in recs2},
                         {"E06000001", "E08000016"})
        # no map: the publisher code is kept
        self.assertIn("E08000038", {r["lad24cd"]
                                    for r in m.build_records(rows, {})})

    def test_min_period_widens_the_window(self):
        recs = m.build_records(self.rows(), {}, min_period="2024-02-01")
        self.assertIn("2024-02-01", {r["period"] for r in recs})

    def test_duplicate_key_raises(self):
        rows = self.rows()
        with self.assertRaises(ValueError) as cm:
            m.build_records(rows + [rows[9]], {})
        self.assertIn("duplicate", str(cm.exception))
        # two source codes meeting on one canonical code
        both = rows + [dict(r, area_code="E08000016") for r in rows
                       if r["area_code"] == "E08000038"]
        with self.assertRaises(ValueError):
            m.build_records(both, {"E08000038": "E08000016"})

    def test_blank_stays_none_through_build(self):
        rows = self.rows()
        rows[0] = dict(rows[0], mean_rent=None)
        recs = m.build_records(rows, {}, min_period="2024-02-01")
        self.assertIn(None, [r["mean_rent"] for r in recs])
        self.assertNotIn(D(0), [r["mean_rent"] for r in recs])


class Identity(unittest.TestCase):
    def test_good(self):
        self.assertEqual(m.check_identity(
            "16september2026", "2026-08-01", "pipr_16september2026.xlsx",
            "16september2026"), [])
        self.assertEqual(m.check_identity(
            "16january2027", "2026-12-01", "pipr_16january2027.xlsx"), [])

    def test_mismatches(self):
        bad = m.check_identity("16september2026", "2026-08-01",
                               "pipr_22july2026.xlsx")
        self.assertTrue(any("22july2026" in p for p in bad))
        bad = m.check_identity("16september2026", "2026-07-01",
                               "pipr_16september2026.xlsx")
        self.assertTrue(any("should end in 2026-08" in p for p in bad))
        bad = m.check_identity("16september2026", "2026-08-01",
                               "pipr_16september2026.xlsx", "22july2026")
        self.assertTrue(any("Cover sheet" in p for p in bad))
        bad = m.check_identity("16september2026", "2026-08-01",
                               "pipr_16september2026.xlsx", None)
        self.assertTrue(any("Cover sheet" in p for p in bad))
        self.assertTrue(m.check_identity("nonsense", "2026-08-01", "x"))


class Hash(unittest.TestCase):
    def rec(self, **kw):
        base = {"lad24cd": "E06000001", "period": "2026-07-01",
                "breakdown_type": "all", "category": "all",
                "mean_rent": D("1000.00"), "rent_index": D("100.00"),
                "annual_pct_change": D("1.00"), "provisional": False}
        base.update(kw)
        return base

    def test_null_differs_from_zero_and_provisional_counts(self):
        h = m.records_sha256([self.rec()])
        self.assertEqual(h, m.records_sha256([self.rec()]))
        self.assertNotEqual(h, m.records_sha256([self.rec(
            annual_pct_change=None)]))
        self.assertNotEqual(
            m.records_sha256([self.rec(annual_pct_change=None)]),
            m.records_sha256([self.rec(annual_pct_change=D("0.00"))]))
        self.assertNotEqual(h, m.records_sha256([self.rec(provisional=True)]))

    def test_period_checks(self):
        def full(areas, **kw):
            return [self.rec(lad24cd=a, breakdown_type=bt, category=c, **kw)
                    for a in areas for _, bt, c in m.BLOCKS]
        m.check_period(full(["A1", "A2"]), "2026-07-01", 2, {"A1"})
        with self.assertRaises(ValueError):
            m.check_period(full(["A1"]), "2026-07-01", 2, set())
        with self.assertRaises(ValueError) as cm:
            m.check_period(full(["A1"])[:-1], "2026-07-01", 1, set())
        self.assertIn("block", str(cm.exception))
        with self.assertRaises(ValueError) as cm:
            m.check_period(full(["A1"], mean_rent=D(0)), "2026-07-01", 1, set())
        self.assertIn("zero or negative", str(cm.exception))
        with self.assertRaises(ValueError):
            m.check_period(full(["A1"]) + full(["A1"])[:1], "2026-07-01",
                           1, set())
        # NULL rent is allowed (suppressed), a missing required area is not
        m.check_period(full(["A1"], mean_rent=None), "2026-07-01", 1, set())
        with self.assertRaises(ValueError):
            m.check_period(full(["A1"]), "2026-07-01", None, {"A2"})


if __name__ == "__main__":
    unittest.main()
