"""S23 (Regulator of Social Housing, registered provider social housing stock
by local authority): edition history and loader on the period-editions
engine.

The live table rsh_rp_stock_by_la keeps its name, key (stock_date, rp_code,
lad24cd), columns and CHECKs and is the latest-edition layer. This module
owns its editions table and file-check ledger, append-only:

    SPEC     rsh_rp_stock_by_la_editions             period stock_date
    ledger   rsh_rp_stock_by_la_editions_file_checks

An edition is what one look-up tool file says about one stock date (31
March). The annual release "Registered provider social housing stock and
rents in England <yyyy> to <yyyy>" carries one look-up tool workbook; its
STOCK_BY_LA sheet holds one row per provider and authority (the key inside a
period is (rp_code, lad24cd), so a revision can add or drop providers), the
publisher's own LA subtotal rows and region rows.

History (so the wording here stays truthful): the old build
(scripts/historical/s23_rsh_stock_build.py, committed in ad8e349) upserted by design
(INSERT ... ON CONFLICT DO UPDATE ... loaded_at = now()). The held data comes
from its single run of 2026-08-14 21:24 UTC (pipeline_run_log id 95,
10,171 rows) and has not been rewritten since; there was no earlier S23
load, so nothing was overwritten in practice. Its num() read a blank stock
cell as 0 (rule 1); that never fired on the held data: every stock cell of
the held 2025 file is a published integer. migrate-legacy (one-off) records
the held stock date as edition 1 "as loaded", with a proof that this parser
reproduces every held row from the held file.

Blanks and zeros (docs/RULES.md rule 1): the publisher documents no marker
for the stock columns, so a stock cell must be a non-negative whole number;
a blank, any text, a negative or a non-integer halts naming the sheet, row
and column (stock_cell). A published 0 stays 0 (12 LARP rows own no stock).
A blank SDR_Size or Survey_Status is stored NULL, never ''.

Identity (rule 3), from the file itself on every path: the Introduction's
title year, its source line "1 April <y1> to 31 March <y2>", its publication
month and version, which must equal the last row of the Version History
sheet; the release page's years when the page is read; the STOCK_BY_LA
header exactly the 2025 set. Inside the file the provider rows must sum to
the LA subtotal rows on all five measures for every authority, the
subtotals to the region rows, with 296 authorities (reconcile). The rank of
a file is its own (publication month, version), never its name or link.

Provenance: the five live columns edition, publication_date, source_url,
source_file and release_page_url are kept per edition as file_edition,
file_publication_date, file_source_url, file_name and file_release_page_url
(not compared, not in the content hash: they never decide whether a file is
new) and are set on every live row of a period (uniform) when it is
inserted or refreshed. A load writes the file's own facts: edition
'<y1> to <y2>', publication_date the first day of the file's publication
month (the file states a month only), source_url the final download URL
(file:<name> for --file), source_file the file name, release_page_url the
release page read (with --file --no-page: 'not read (--file --no-page)').

Geography (rule 4): geography.resolve('23', codes, by_period={stock_date:
codes}) with every LA code of the provider and subtotal rows; the canonical
code must then be in la_boundaries, anything else is UNEXPLAINED and stops.

S23 is not a W1 or map input today. If it is ever wired into W1, run
refresh-latest in the same session before W1 (refresh copies the edition's
loaded_at, which refresh_map.py would otherwise not see as new).

Subcommands (every writing command previews by default; --commit and
--simulate are mutually exclusive; --simulate runs the --commit path and
always rolls back). There is no sync-new: migrate-legacy records the held
stock date.
    python scripts/s23_rsh_stock_editions.py ddl [--commit | --simulate]
    python scripts/s23_rsh_stock_editions.py status
    python scripts/s23_rsh_stock_editions.py load [--release Y1-Y2]
        [--file PATH [--no-page]] [--recheck YYYY-MM-DD] [--allow-older-file]
        [--commit | --simulate]
        # the file is downloaded to data/raw/s23_rsh/ (also in a preview; a
        # preview writes nothing to the database). A stock date breaking a
        # stop condition is REJECTED (nothing stored, exit 1).
    python scripts/s23_rsh_stock_editions.py refresh-latest
        [--accept-key-changes YYYY-MM-DD] [--commit | --simulate]
    python scripts/s23_rsh_stock_editions.py migrate-legacy FILE
        [--commit | --simulate]
    python scripts/s23_rsh_stock_editions.py restore-edition YYYY-MM-DD N
        [--commit | --simulate]
"""
import argparse
import dataclasses
import hashlib
import json
import math
import numbers
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

sys.path.insert(0, str(Path(__file__).resolve().parent))

import editions_core as core  # noqa: E402
import geography  # noqa: E402
import period_editions as pe  # noqa: E402
from editions_core import halt  # noqa: E402

RUN_AGENT = "Source 23 - RSH RP stock by local authority"   # run log id 95
RUN_SOURCE = "23"
LIVE_MISSING = pe.LIVE_MISSING

API = "https://www.gov.uk/api/content"
GOV_UK = "https://www.gov.uk"
COLLECTION_PATH = ("/government/collections/registered-provider-social-"
                   "housing-stock-and-rents-in-england")
REPO = Path(__file__).resolve().parent.parent
RAW_DIR = REPO / "data" / "raw" / "s23_rsh"            # git-ignored
USER_AGENT = ("ucws-pipeline S23 loader (read-only download of the public "
              "RSH registered providers look-up tool)")
MIN_FILE_BYTES = 500 * 1024
NO_PAGE = "not read (--file --no-page)"

# ---------------------------------------------------------------------------
# Columns (live types from information_schema, checked 2026-10-09: the five
# stock columns integer NOT NULL; rp_size_band and survey_status the only
# nullable columns)
# ---------------------------------------------------------------------------

VALUE_TYPES = (
    ("total_social_stock", "integer NOT NULL"),
    ("general_needs_self_contained", "integer NOT NULL"),
    ("general_needs_bedspaces", "integer NOT NULL"),
    ("supported_housing_and_older_people", "integer NOT NULL"),
    ("low_cost_home_ownership", "integer NOT NULL"),
)
EXTRA_TYPES = (
    ("rp_name", "text NOT NULL"),
    ("provider_type", "text NOT NULL"),
    ("rp_size_band", "text"),
    ("survey_status", "text"),
    ("publisher_la_code", "varchar(9) NOT NULL"),
    ("la_name", "text NOT NULL"),
)
# editions-only provenance (live column <- editions column)
PROVENANCE_TYPES = (
    ("file_edition", "text NOT NULL"),
    ("file_publication_date", "date NOT NULL"),
    ("file_source_url", "text NOT NULL"),
    ("file_name", "text NOT NULL"),
    ("file_release_page_url", "text NOT NULL"),
)
PROVENANCE_PAIRS = (("edition", "file_edition"),
                    ("publication_date", "file_publication_date"),
                    ("source_url", "file_source_url"),
                    ("source_file", "file_name"),
                    ("release_page_url", "file_release_page_url"))


def _names(types) -> tuple:
    return tuple(c for c, _ in types)


STOCK_COLUMNS = _names(VALUE_TYPES)
PARTS = STOCK_COLUMNS[1:]
COMPARED = STOCK_COLUMNS + _names(EXTRA_TYPES)       # the 11 compared columns
PROVENANCE = _names(PROVENANCE_TYPES)
LIVE_PROVENANCE = tuple(lc for lc, _ in PROVENANCE_PAIRS)
KEY = ("rp_code", "lad24cd")

# ---------------------------------------------------------------------------
# The look-up tool
# ---------------------------------------------------------------------------

INTRO_SHEET = "Introduction and Contents"
HISTORY_SHEET = "Version History"
SHEET = "STOCK_BY_LA"
TITLE_RE = re.compile(r"RP social housing by local authority area \(SDR and "
                      r"LADR data\) ([0-9]{4})")
SOURCE_RE = re.compile(r"Statistical Data Return \(SDR\)/Local Authority Data "
                       r"Return \(LADR\) 1 April ([0-9]{4}) to 31 March "
                       r"([0-9]{4})")
PUB_RE = re.compile(r"Publication date: ([A-Za-z]+ [0-9]{4})")
VERSION_RE = re.compile(r"Version: ([0-9]+(?:\.[0-9]+)*)")
HISTORY_HEADER = ["Version", "Publication Date", "Changes"]
# STOCK_BY_LA header (exactly the 2025 set, in order) -> field
HEADER_MAP = (
    ("RP_Name", "rp_name"),
    ("RP_Code", "rp_code"),
    ("RP_Type", "rp_type"),
    ("SDR_Size", "sdr_size"),
    ("Survey_Status", "survey_status"),
    ("LA_Nm", "la_nm"),
    ("LA_Code", "la_code"),
    ("Concat", None),                     # the workbook's search-box key
    ("Total Social Stock", "total_social_stock"),
    ("LA_GN_SC_Own", "general_needs_self_contained"),
    ("LA_GN_BSp_Own", "general_needs_bedspaces"),
    ("LA_SHHOP", "supported_housing_and_older_people"),
    ("LA_LCHO_Less_100_Eqty_Own", "low_cost_home_ownership"),
)
HEADERS = tuple(h for h, _ in HEADER_MAP)
HEADER_OF = {f: h for h, f in HEADER_MAP if f}
PROVIDER_TYPES = {"Large": "PRP", "Small": "PRP", "LARP": "LARP"}
SUBTOTAL_TYPE = "LA"
REGION_TYPE = "Region"
DISTRICT_RE = re.compile(r"E0[6-9][0-9]{6}")
EXPECTED_AREAS = 296
EXPECTED_REGIONS = 9

# ---------------------------------------------------------------------------
# Stop conditions (percent; calibrated on 2024 to 2025: national +0.96%,
# supported -1.0%, largest authority move +10.0%; the publisher's LARP
# revisions to earlier years were under 0.4% of total stock)
# ---------------------------------------------------------------------------

NEW_TOTAL_PCT = 5           # new stock date: national total_social_stock
NEW_SUPPORTED_PCT = 10      # new stock date: national supported housing
NEW_AREA_PCT = 25           # new stock date: any authority's total
REV_NATIONAL_PCT = 1        # revision: national total of any stock column
REV_MAX_AREAS = 30          # revision: authorities whose totals change
REV_AREA_PCT = 10           # revision: any authority's total
REV_KEYS_PCT = 5            # revision: provider rows added, or removed

# ---------------------------------------------------------------------------
# The held state (migrate-legacy preconditions), surveyed 2026-10-09
# ---------------------------------------------------------------------------

# (rows, the survey hash: md5 of every stored column, see live_state)
LEGACY_LIVE = (10171, "4967e5b59580d98f" "0edc3c6f8d908ac9")
# the file the old build read (sha256 split so the credential scan does not
# flag it) and the URL it was read from
LEGACY_FILE = {
    "name": "RP_COMBINED_TOOL_2025_FINAL_V1.1.xlsx",
    "sha256": "6c237c79861c5ca82fea48b67039556e"
              "f0585c37149c3a951146188faa40a705",
    "url": "https://assets.publishing.service.gov.uk/media/"
           "6911fc96663088df8f54f499/RP_COMBINED_TOOL_2025_FINAL_V1.1.xlsx",
}
# before the editions exist (a preview against live): the held rows came from
# that file, whose Introduction says November 2025, version 1.1
LEGACY_RANK = (date(2025, 11, 1), (1, 1))


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _json(x):
    return json.loads(x) if isinstance(x, (str, bytes)) else (x or {})


def _norm(s) -> str:
    return " ".join(str("" if s is None else s).split())


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or \
        (isinstance(v, str) and not v.strip())


def _text(v) -> "str | None":
    """A text cell as published (outer whitespace stripped); None if
    blank."""
    if _blank(v):
        return None
    return str(v).strip()


def _zeros() -> dict:
    return dict.fromkeys(STOCK_COLUMNS, 0)  # not a source value


def _vkey(text) -> tuple:
    """A version's sort key: its numbers, trailing zeros dropped ('1.0' and
    '1' are equal; '1.10' sorts after '1.9')."""
    parts = [int(x) for x in str(text).split(".")]
    while len(parts) > 1 and parts[-1] == 0:
        parts.pop()
    return tuple(parts)


def _vtext(key) -> str:
    return ".".join(str(x) for x in key)


def rank_text(rank) -> str:
    """'version 1.1, dated 2025-11' of a (month, version key) rank."""
    if not rank:
        return "no file version"
    return f"version {_vtext(rank[1])}, dated {rank[0]:%Y-%m}"


def _month(text, where) -> date:
    try:
        d = datetime.strptime(_norm(text), "%B %Y").date()
    except ValueError:
        raise ValueError(f"{where}: {text!r} is not a month 'Month yyyy'") \
            from None
    return d.replace(day=1)


def _version_text(v) -> "str | None":
    """The text of a Version History version cell; None if the cell is not
    a version number."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, numbers.Integral):
        return str(int(v))
    if isinstance(v, numbers.Real):
        return None if math.isnan(v) else repr(float(v))
    s = _norm(v)
    return s if re.fullmatch(r"[0-9]+(?:\.[0-9]+)*", s) else None


# ---------------------------------------------------------------------------
# Discovery: the GOV.UK content API
# ---------------------------------------------------------------------------

DOC_TITLE_RE = re.compile(r"Registered provider social housing stock and rents "
                          r"in England ([0-9]{4}) to ([0-9]{4})")
SERIES_RE = re.compile(r"Registered provider social housing", re.I)
TOOL_TITLE = "Registered providers look-up tool"


def _documents(collection_json) -> list:
    d = _json(collection_json)
    return list(((d.get("links") or {}).get("documents")) or [])


def _titles(collection_json) -> list:
    return [_norm(d.get("title")) for d in _documents(collection_json)]


def release_candidates(collection_json) -> dict:
    """{(y1, y2): [(title, base_path)]} of the collection's documents titled
    'Registered provider social housing stock and rents in England <y1> to
    <y2>'."""
    out = {}
    for doc in _documents(collection_json):
        t = _norm(doc.get("title"))
        mm = DOC_TITLE_RE.fullmatch(t)
        if mm and doc.get("base_path"):
            out.setdefault((int(mm.group(1)), int(mm.group(2))), []).append(
                (t, doc["base_path"]))
    return out


def _years_in(title) -> list:
    """The years a title names: four-digit years, and the closing year of a
    'yyyy/yy' span."""
    ys = [int(y) for y in re.findall(r"(?<![0-9])([0-9]{4})(?![0-9])", title)]
    for a, b in re.findall(r"([0-9]{4})/([0-9]{2})(?![0-9])", title):
        ys.append(int(a) // 100 * 100 + int(b))
    return ys


def latest_release(collection_json) -> tuple:
    """(y1, y2, base_path) of the newest release by its closing year.
    ValueError (listing the titles seen) when no document has the expected
    title (a renamed series), when two documents claim the newest years, when
    a matching title does not span one year, or when a document of the series
    names a later year in a title that does not fit (not passed over for an
    older release)."""
    cands = release_candidates(collection_json)
    titles = _titles(collection_json)
    if not cands:
        raise ValueError("no document titled 'Registered provider social "
                         "housing stock and rents in England <yyyy> to "
                         f"<yyyy>' in the collection; titles seen: {titles[:20]}")
    odd = sorted(k for k in cands if k[1] != k[0] + 1)
    if odd:
        raise ValueError(f"release titles that do not span one year: "
                         f"{[cands[k][0][0] for k in odd]}")
    newest = max(cands, key=lambda k: k[1])
    stray = [t for t in titles if not DOC_TITLE_RE.fullmatch(t)
             and SERIES_RE.match(t)
             and any(y > newest[1] for y in _years_in(t))]
    if stray:
        raise ValueError(f"the newest matching release is {newest[0]} to "
                         f"{newest[1]}, but the collection lists a document of "
                         f"the series with a later year whose title does not "
                         f"fit: {stray}; not choosing an older release over it")
    if len(cands[newest]) != 1:
        raise ValueError(f"{len(cands[newest])} documents for {newest[0]} to "
                         f"{newest[1]}: {[t for t, _ in cands[newest]]}; "
                         "refusing to choose one")
    return newest[0], newest[1], cands[newest][0][1]


def release_for_years(collection_json, y1, y2) -> tuple:
    """(y1, y2, base_path) of the given years' release (--release Y1-Y2, or
    the years of a --file); ValueError unless exactly one document is
    titled for them."""
    cands = release_candidates(collection_json).get((int(y1), int(y2)), [])
    if len(cands) != 1:
        raise ValueError(f"{len(cands)} documents titled for {y1} to {y2} in "
                         f"the collection: {[t for t, _ in cands]}; titles "
                         f"seen: {_titles(collection_json)[:20]}")
    return int(y1), int(y2), cands[0][1]


def page_years(release_json) -> "tuple | None":
    mm = DOC_TITLE_RE.fullmatch(_norm(_json(release_json).get("title")))
    return (int(mm.group(1)), int(mm.group(2))) if mm else None


def _attachments(page_json) -> list:
    d = _json(page_json)
    return list(((d.get("details") or {}).get("attachments")) or [])


def lookup_tool_attachment(release_json) -> dict:
    """{title, url, content_type, file_size} of the release page's one
    attachment titled 'Registered providers look-up tool' with an .xlsx URL.
    ValueError listing every attachment title otherwise (none, several, or
    not .xlsx)."""
    atts = _attachments(release_json)
    hits = [a for a in atts if _norm(a.get("title")) == TOOL_TITLE]
    names = [_norm(a.get("title")) for a in atts]
    if len(hits) != 1 or not hits[0].get("url"):
        raise ValueError(f"{len(hits)} attachments titled {TOOL_TITLE!r}; "
                         f"attachments: {names}")
    a = hits[0]
    if Path(urlparse(a["url"]).path).suffix.lower() != ".xlsx":
        raise ValueError(f"the {TOOL_TITLE!r} attachment is not an .xlsx "
                         f"file ({a['url']}); attachments: {names}")
    return {"title": _norm(a["title"]), "url": a["url"],
            "content_type": a.get("content_type"),
            "file_size": a.get("file_size")}


def change_history(page_json) -> list:
    """[(public_timestamp date, note)] of the page, as published."""
    d = _json(page_json)
    ch = ((d.get("details") or {}).get("change_history")) or []
    return [(str(c.get("public_timestamp") or "")[:10], _norm(c.get("note")))
            for c in ch]


# ---------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------

def stock_cell(v, where) -> int:
    """A published stock count: a non-negative whole number (an int, or a
    float equal to one). Anything else raises ValueError naming `where`: a
    blank (the publisher documents no marker for these columns, so a blank
    halts and is never read as 0), any text, a boolean, a negative or a
    non-integer."""
    if v is None or (isinstance(v, str) and not v.strip()):
        raise ValueError(f"{where}: a blank cell (the publisher documents no "
                         "marker for the stock columns; a blank halts and is "
                         "never read as 0)")
    if isinstance(v, bool):
        raise ValueError(f"{where}: {v!r} is not a count")
    if isinstance(v, numbers.Integral):
        n = int(v)
    elif isinstance(v, numbers.Real):
        if math.isnan(v):
            raise ValueError(f"{where}: a blank cell (NaN); never read as 0")
        if not float(v).is_integer():
            raise ValueError(f"{where}: non-integer {v!r}")
        n = int(v)
    else:
        raise ValueError(f"{where}: {v!r} is not a whole number")
    if n < 0:
        raise ValueError(f"{where}: negative count {n}")
    return n


# ---------------------------------------------------------------------------
# Reading the look-up tool
# ---------------------------------------------------------------------------

def _sheet_rows(wb, name) -> list:
    rows = [tuple(r) for r in wb[name].iter_rows(values_only=True)]
    width = max((len(r) for r in rows), default=1)
    return [r + (None,) * (width - len(r)) for r in rows]


def _cells(rows) -> list:
    return [_norm(v) for r in rows for v in r if not _blank(v)]


def _one(cells, rx, what, sheet):
    hits = [mm for mm in (rx.fullmatch(c) for c in cells) if mm]
    if len(hits) != 1:
        raise ValueError(f"{sheet}: {len(hits)} cells like {what!r}, "
                         f"expected one; cells seen: {cells[:24]}")
    return hits[0]


def _read_history(rows, name) -> tuple:
    """([(version text, month date, note)], tail cells) of the Version
    History sheet: the rows under its 'Version, Publication Date, Changes'
    header (found in any column) while the cell under 'Version' is a version
    number, wholly empty rows skipped; the cells after them."""
    heads = [(i, j) for i, r in enumerate(rows) for j in range(len(r))
             if [_norm(v) for v in r[j:j + 3]] == HISTORY_HEADER]
    if len(heads) != 1:
        raise ValueError(f"{name}: {HISTORY_SHEET} has {len(heads)} header "
                         f"rows {HISTORY_HEADER}; cells seen: "
                         f"{_cells(rows)[:24]}")
    (i, j), out = heads[0], []
    i += 1
    while i < len(rows):
        r = rows[i][j:j + 3] + (None,) * 3
        if all(_blank(v) for v in rows[i]):
            i += 1
            continue
        vt = _version_text(r[0])
        if vt is None:
            break
        out.append((vt, _month(r[1], f"{name}: {HISTORY_SHEET} row {i + 1}"),
                    _norm(r[2])))
        i += 1
    if not out:
        raise ValueError(f"{name}: {HISTORY_SHEET} has no version rows")
    keys = [_vkey(v) for v, _, _ in out]
    if any(b <= a for a, b in zip(keys, keys[1:])):
        raise ValueError(f"{name}: {HISTORY_SHEET} versions are not in "
                         f"order: {[v for v, _, _ in out]}")
    return out, _cells(rows[i:])


def read_tool(path) -> dict:
    """The facts of one look-up tool workbook: {path, title, year, y1, y2,
    stock_date (ISO, 31 March of y2), edition ('<y1> to <y2>'), month_text,
    published (the publication month's first day), version, rank
    ((published, version key)), history [(version, month, note)],
    providers [{row, rp_name, rp_code, rp_type, sdr_size, survey_status,
    la_name, la_code, values}], subtotals {la_code: {row, name, region,
    values}}, regions {name: {row, values}}, empty_rows}. values: {stock
    column: int}.

    Raises ValueError (naming what was seen) on a missing sheet; an
    Introduction without exactly one title, source line, publication month
    and version; a title year other than the source line's closing year, or
    a source line not spanning one year; a Version History whose last row is
    not the Introduction's version and month (or whose own publication and
    version lines differ); a STOCK_BY_LA header that is not exactly the 2025
    set (unknown and missing headers listed); a non-empty row without an
    RP_Code, or one after an empty row (empty rows only at the end, wholly
    empty); an unknown RP_Type; a provider row without an E06-E09 LA_Code; a
    repeated subtotal or region row; a provider without a name or authority
    name; or any stock cell that is not a non-negative whole number."""
    import openpyxl
    p = Path(path)
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 (not a workbook)
        raise ValueError(f"{p.name}: not a readable .xlsx workbook ({e})") \
            from None
    try:
        missing = [s for s in (INTRO_SHEET, HISTORY_SHEET, SHEET)
                   if s not in wb.sheetnames]
        if missing:
            raise ValueError(f"{p.name}: sheets {missing} missing; sheets "
                             f"{wb.sheetnames}")
        intro = _sheet_rows(wb, INTRO_SHEET)
        hist_rows = _sheet_rows(wb, HISTORY_SHEET)
        stock = _sheet_rows(wb, SHEET)
    finally:
        wb.close()
    cells = _cells(intro)
    t = _one(cells, TITLE_RE, "RP social housing by local authority area "
             "(SDR and LADR data) <yyyy>", INTRO_SHEET)
    s = _one(cells, SOURCE_RE, "Statistical Data Return (SDR)/Local "
             "Authority Data Return (LADR) 1 April <yyyy> to 31 March <yyyy>",
             INTRO_SHEET)
    pub = _one(cells, PUB_RE, "Publication date: <Month yyyy>", INTRO_SHEET)
    ver = _one(cells, VERSION_RE, "Version: <n>", INTRO_SHEET)
    year, y1, y2 = int(t.group(1)), int(s.group(1)), int(s.group(2))
    if y2 != y1 + 1:
        raise ValueError(f"{p.name}: the source line spans {y1} to {y2}, not "
                         "one year (1 April to 31 March)")
    if year != y2:
        raise ValueError(f"{p.name}: the Introduction title says {year}, the "
                         f"source line says 1 April {y1} to 31 March {y2}")
    month_text = pub.group(1)
    published = _month(month_text, f"{p.name}: {INTRO_SHEET}")
    version = ver.group(1)
    history, tail = _read_history(hist_rows, p.name)
    lv, lm, _ = history[-1]
    if _vkey(lv) != _vkey(version) or lm != published:
        raise ValueError(f"{p.name}: the last {HISTORY_SHEET} row is version "
                         f"{lv} of {lm:%B %Y}, the Introduction says version "
                         f"{version} of {month_text}")
    for rx, want, what in ((PUB_RE, published, "publication month"),
                           (VERSION_RE, _vkey(version), "version")):
        for c in tail:
            mm = rx.fullmatch(c)
            if not mm:
                continue
            got = (_month(mm.group(1), f"{p.name}: {HISTORY_SHEET}")
                   if rx is PUB_RE else _vkey(mm.group(1)))
            if got != want:
                raise ValueError(f"{p.name}: {HISTORY_SHEET} states {c!r}; "
                                 f"the Introduction's {what} is "
                                 f"{month_text if rx is PUB_RE else version}")
    out = _read_stock(stock, p.name)
    out.update({"path": p, "title": t.group(0), "year": year, "y1": y1,
                "y2": y2, "stock_date": date(y2, 3, 31).isoformat(),
                "edition": f"{y1} to {y2}", "month_text": month_text,
                "published": published, "version": version,
                "rank": (published, _vkey(version)), "history": history})
    return out


def _read_stock(rows, name) -> dict:
    if not rows:
        raise ValueError(f"{name}: {SHEET} is empty")
    header = [_norm(v) for v in rows[0]]
    width = len(header)
    while width and not header[width - 1]:
        width -= 1
    header = header[:width]
    if tuple(header) != HEADERS:
        unknown = [h for h in header if h not in HEADERS]
        missing = [h for h in HEADERS if h not in header]
        raise ValueError(f"{name}: {SHEET} header {header} is not the 2025 "
                         f"set: unknown {unknown}, missing {missing} (expected "
                         f"exactly {list(HEADERS)}, in order); the mapping "
                         "must be corrected deliberately, with the evidence")
    fields = [f for _, f in HEADER_MAP]
    providers, subtotals, regions = [], {}, {}
    empty = 0  # trailing empty rows seen; not a source value
    for n, r in enumerate(rows[1:], start=2):
        if all(_blank(v) for v in r):
            empty += 1
            continue
        where = f"{name}: {SHEET} row {n}"
        if empty:
            raise ValueError(f"{where}: data after {empty} empty row(s); "
                             "empty rows are allowed only at the end")
        if any(not _blank(v) for v in r[width:]):
            raise ValueError(f"{where}: a cell beyond the {width} headed "
                             "columns")
        c = dict(zip(fields, r[:width]))
        code = _text(c["rp_code"])
        if code is None:
            raise ValueError(f"{where}: the row is not empty but has no "
                             f"RP_Code ({[v for v in r if not _blank(v)][:6]})")
        typ = _text(c["rp_type"])
        la_code = _text(c["la_code"])
        values = {f: stock_cell(c[f], f"{where} {HEADER_OF[f]}")
                  for f in STOCK_COLUMNS}
        if typ in PROVIDER_TYPES:
            if not la_code or not DISTRICT_RE.fullmatch(la_code):
                raise ValueError(f"{where}: provider row {code} ({typ}) has "
                                 f"LA_Code {la_code!r}, not an E06-E09 "
                                 "authority code (the 2024 file's provider-"
                                 "by-region rows look like this)")
            nm, la = _text(c["rp_name"]), _text(c["la_nm"])
            if nm is None or la is None:
                raise ValueError(f"{where}: provider row {code} has no "
                                 "RP_Name or LA_Nm")
            providers.append({"row": n, "rp_name": nm, "rp_code": code,
                              "rp_type": typ,
                              "sdr_size": _text(c["sdr_size"]),
                              "survey_status": _text(c["survey_status"]),
                              "la_name": la, "la_code": la_code,
                              "values": values})
        elif typ == SUBTOTAL_TYPE:
            if not la_code or not DISTRICT_RE.fullmatch(la_code):
                raise ValueError(f"{where}: LA subtotal row with LA_Code "
                                 f"{la_code!r}")
            if la_code in subtotals:
                raise ValueError(f"{where}: a second LA subtotal row for "
                                 f"{la_code} (row {subtotals[la_code]['row']})")
            subtotals[la_code] = {"row": n, "name": _text(c["rp_name"]),
                                  "region": _text(c["la_nm"]),
                                  "values": values}
        elif typ == REGION_TYPE:
            reg = _text(c["rp_name"])
            if reg is None or reg in regions:
                raise ValueError(f"{where}: region row {reg!r} unnamed or "
                                 "repeated")
            regions[reg] = {"row": n, "values": values}
        else:
            raise ValueError(f"{where}: unknown RP_Type {typ!r} (known: "
                             f"{sorted(PROVIDER_TYPES)}, {SUBTOTAL_TYPE!r}, "
                             f"{REGION_TYPE!r})")
    return {"providers": providers, "subtotals": subtotals,
            "regions": regions, "empty_rows": empty, "header": header}


def tool_codes(tool) -> set:
    """Every LA code of the provider and LA subtotal rows."""
    return ({p["la_code"] for p in tool["providers"]}
            | set(tool["subtotals"]))


# ---------------------------------------------------------------------------
# Reconciliation inside the file
# ---------------------------------------------------------------------------

def _add(acc, values) -> None:
    for c in STOCK_COLUMNS:
        acc[c] += values[c]


def england_totals(tool) -> dict:
    """{stock column: the sum of the region rows}."""
    out = _zeros()
    for r in tool["regions"].values():
        _add(out, r["values"])
    return out


def reconcile(tool) -> list:
    """Problems (empty = fine), exact: EXPECTED_AREAS LA subtotal rows and
    EXPECTED_REGIONS region rows; for every authority the provider rows sum
    to its LA subtotal row on all five measures (an authority with only one
    of the two is a problem); for every region the LA subtotals of its
    authorities sum to its region row."""
    g = globals()
    out = []
    subs, regs = tool["subtotals"], tool["regions"]
    if len(subs) != g["EXPECTED_AREAS"]:
        out.append(f"{len(subs)} LA subtotal rows, expected "
                   f"{g['EXPECTED_AREAS']}")
    if len(regs) != g["EXPECTED_REGIONS"]:
        out.append(f"{len(regs)} region rows, expected "
                   f"{g['EXPECTED_REGIONS']}")
    sums = {}
    for p in tool["providers"]:
        _add(sums.setdefault(p["la_code"], _zeros()), p["values"])
    for code in sorted(set(sums) | set(subs)):
        if code not in subs:
            out.append(f"{code}: provider rows but no LA subtotal row")
            continue
        if code not in sums:
            out.append(f"{code}: an LA subtotal row but no provider rows")
            continue
        for c in STOCK_COLUMNS:
            a, b = sums[code][c], subs[code]["values"][c]
            if a != b:
                out.append(f"{code} {c}: the provider rows sum to {a:,}, the "
                           f"LA subtotal row says {b:,}")
    rsum = {}
    for code, s in subs.items():
        _add(rsum.setdefault(s["region"], _zeros()), s["values"])
    for reg in sorted(set(rsum) | set(regs), key=str):
        if reg not in regs:
            out.append(f"{reg}: LA subtotal rows but no region row")
            continue
        if reg not in rsum:
            out.append(f"{reg}: a region row but no LA subtotal rows")
            continue
        for c in STOCK_COLUMNS:
            a, b = rsum[reg][c], regs[reg]["values"][c]
            if a != b:
                out.append(f"{reg} {c}: the LA subtotals sum to {a:,}, the "
                           f"region row says {b:,}")
    return out


# ---------------------------------------------------------------------------
# Records
# ---------------------------------------------------------------------------

def _getter(resolve):
    return resolve if callable(resolve) else (lambda c: resolve[c])


def build_rows(tool, resolve, provenance) -> list:
    """The stock date's records from a read tool (read_tool): one per
    provider row, sorted by key. resolve: {publisher code: lad24cd} or a
    callable. provenance: the five file_* values (every record carries
    them). provider_type PRP (Large, Small) or LARP; a blank SDR_Size or
    Survey_Status is NULL (as read). ValueError when a total is not the sum
    of its four parts, a code resolves to nothing, or a key (rp_code,
    lad24cd) repeats (two publisher codes reaching one lad24cd included)."""
    get = _getter(resolve)
    gone = [c for c in PROVENANCE if c not in provenance]
    if gone:
        raise ValueError(f"build_rows: provenance lacks {gone}")
    prov = {c: provenance[c] for c in PROVENANCE}
    out, seen = [], {}
    for p in tool["providers"]:
        v = p["values"]
        parts = sum(v[c] for c in PARTS)
        where = f"{SHEET} row {p['row']} {p['rp_code']}/{p['la_code']}"
        if v["total_social_stock"] != parts:
            raise ValueError(f"{where}: Total Social Stock "
                             f"{v['total_social_stock']:,} is not the sum of "
                             f"its four parts ({parts:,})")
        lad = get(p["la_code"])
        if not lad:
            raise ValueError(f"{where}: {p['la_code']} resolves to no lad24cd")
        key = (p["rp_code"], lad)
        if key in seen:
            raise ValueError(f"{where}: duplicate key {key} (also row "
                             f"{seen[key]}); refusing to choose a winner")
        seen[key] = p["row"]
        rec = {"rp_code": p["rp_code"], "lad24cd": lad,
               "stock_date": tool["stock_date"]}
        rec.update(v)
        rec.update({"rp_name": p["rp_name"],
                    "provider_type": PROVIDER_TYPES[p["rp_type"]],
                    "rp_size_band": p["sdr_size"],
                    "survey_status": p["survey_status"],
                    "publisher_la_code": p["la_code"],
                    "la_name": p["la_name"]})
        rec.update(prov)
        out.append(rec)
    out.sort(key=lambda r: (r["rp_code"], r["lad24cd"]))
    return out


def rows_content_sha(records) -> str:
    """sha256 of the records' content: one line per record sorted by key,
    rp_code|lad24cd|the 11 compared columns (NULL as the empty string),
    lines joined by LF, UTF-8. Provenance is not in it, so it never makes a
    new edition."""
    def cell(v):
        return "" if v is None else str(v)
    ordered = sorted(records, key=lambda r: (str(r["rp_code"]),
                                             str(r["lad24cd"])))
    lines = ["|".join([cell(r[k]) for k in KEY]
                      + [cell(r.get(c)) for c in COMPARED]) for r in ordered]
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Geography
# ---------------------------------------------------------------------------

def resolve_codes(cur, codes, period) -> tuple:
    """({publisher code: lad24cd}, problems). geography.resolve(cur, '23',
    codes, by_period={period: codes}) (source 23 is declared 'old': a new
    Barnsley or Sheffield code is a problem); then each code's canonical
    form must be in la_boundaries, else 'UNEXPLAINED'. Two codes reaching
    one lad24cd are a problem."""
    codes = set(codes)
    recode, problems = geography.resolve(cur, RUN_SOURCE, codes,
                                         by_period={period: codes})
    problems = list(problems)
    cur.execute("SELECT lad24cd FROM public.la_boundaries")
    valid = {r[0] for r in cur.fetchall()}
    out, hit = {}, {}
    for c in sorted(codes):
        r = geography.canonical(c, recode)
        if r in valid:
            out[c] = r
            hit.setdefault(r, []).append(c)
        else:
            problems.append(f"UNEXPLAINED {c}: not in la_boundaries and not a "
                            "la_code_lookup recode")
    two = {k: v for k, v in hit.items() if len(v) > 1}
    if two:
        problems.append(f"{period}: codes {two} reach one lad24cd; refusing "
                        "to choose a winner")
    return out, problems


# ---------------------------------------------------------------------------
# Ranks, ledger and planning
# ---------------------------------------------------------------------------

_RANK_RE = re.compile(r"\bversion ([0-9]+(?:\.[0-9]+)*); dated ([0-9]{4})-"
                      r"([0-9]{2})\b")
_LEDGER_RE = re.compile(r"(.*) \(version ([0-9]+(?:\.[0-9]+)*); dated "
                        r"([0-9]{4})-([0-9]{2}); stock date "
                        r"([0-9]{4}-[0-9]{2}-[0-9]{2})\)")


def release_rank(source_file_or_tool) -> "tuple | None":
    """(publication month, version key) of a file: a read tool's own rank,
    or the 'version <v>; dated <yyyy-mm>' an edition's or ledger row's
    source_file carries; None if it names none (a bare file name is never
    read for a rank)."""
    x = source_file_or_tool
    if isinstance(x, dict):
        return x.get("rank")
    mm = _RANK_RE.search(str(x or ""))
    if not mm:
        return None
    return (date(int(mm.group(2)), int(mm.group(3)), 1), _vkey(mm.group(1)))


def ledger_source(where, rank, version, period) -> str:
    """A file's ledger source_file: where it was read (the final URL, or the
    path of a --file), its own version and month, and the stock date it
    covers, so a later run knows the file's period without reading it."""
    return (f"{where} (version {version}; dated {rank[0]:%Y-%m}; stock date "
            f"{period})")


def parse_ledger_source(s) -> "tuple | None":
    """(where, rank, [stock date]) of a ledger source_file, or None."""
    mm = _LEDGER_RE.fullmatch(str(s or ""))
    if not mm:
        return None
    rank = (date(int(mm.group(3)), int(mm.group(4)), 1), _vkey(mm.group(2)))
    return mm.group(1), rank, [mm.group(5)]


def content_sha256(path) -> str:
    """SHA-256 of the file's bytes."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def tip_ranks(tips, checked) -> dict:
    """{held period: the newest file rank that stated it}: the tip
    edition's, or a later file the ledger records for the period (a file
    found unchanged stores no edition)."""
    out = {}
    for p, t in tips.items():
        ranks = [t.get("rank")] + [release_rank(s)
                                   for s, _ in checked.get(p, ())]
        ranks = [r for r in ranks if r is not None]
        out[p] = max(ranks) if ranks else None
    return out


def ledger_complete(checked, where, sha, ranks, *, allow_older=False,
                    stranded=()) -> "list | None":
    """The periods the ledger records this file's (where, sha256) pair for,
    if that covers every period the file needs (its ledger source declares
    them: here its one stock date); else None (the file is read and
    planned). A held period whose tip comes from a newer file is not needed
    (skipped as older) unless allow_older; a stranded period (editions, no
    live rows) always makes the file be read. So a run that failed or was
    rejected leaves no ledger row and the rerun reads the file again."""
    have, declared = set(), None
    for p, pairs in checked.items():
        for s, h in pairs:
            parsed = parse_ledger_source(s)
            if h == sha and parsed and parsed[0] == where:
                have.add(str(p))
                declared = parsed
    if declared is None:
        return None
    _, rank, need = declared
    need = set(need)
    if not allow_older:
        need = {p for p in need if not (ranks.get(p) is not None
                                        and ranks[p] > rank)}
    if not need or need & set(stranded):
        return None
    return sorted(have) if need <= have else None


def plan_periods(periods, rank, ranks, checked, source, sha, *, recheck,
                 allow_older) -> tuple:
    """(new, revised, skipped). periods: the file's stock dates; rank: its
    own; ranks: {held period: tip_ranks}; checked: {period: {(source,
    sha)}} of the ledger (stranded periods left out by the caller, so they
    are always read). new: not held; revised: held, compared with the tip;
    skipped {period: reason}: held periods whose tip comes from a newer file
    (unless allow_older: the older-file guard, on every path) or whose
    (source, sha) the ledger already records (unless --recheck names it).
    --recheck must be a held period of the file (ValueError)."""
    periods = sorted(str(p) for p in periods)
    if recheck is not None and (recheck not in periods or recheck not in ranks):
        raise ValueError(f"--recheck {recheck}: not a held stock date of this "
                         f"file ({', '.join(periods)})")
    new, revised, skipped = [], [], {}
    for p in periods:
        if p not in ranks:
            new.append(p)
            continue
        t = ranks[p]
        if t is not None and t > rank and not allow_older:
            skipped[p] = (f"older: its tip comes from a file of "
                          f"{rank_text(t)}, this file is {rank_text(rank)}")
            continue
        if recheck != p and (source, sha) in checked.get(p, ()):
            skipped[p] = "checked: this file is already in the ledger"
            continue
        revised.append(p)
    return new, revised, skipped


# ---------------------------------------------------------------------------
# Stop conditions
# ---------------------------------------------------------------------------

def _area_sums(records) -> dict:
    out = {}
    for r in records or ():
        _add(out.setdefault(r["lad24cd"], _zeros()), r)
    return out


def _national(records) -> dict:
    out = _zeros()
    for r in records or ():
        _add(out, r)
    return out


def _over(old, new, pct) -> bool:
    """True when new moves from old by more than pct percent (exact)."""
    if old == new:
        return False
    if old == 0:
        return True
    return abs(new - old) * 100 > pct * old


def period_problems(new, tip, prev, *, kind) -> list:
    """The stop conditions of one stock date (empty = fine). new: the
    file's records; tip: (kind 'revised') the tip edition's records of the
    same stock date; prev: (kind 'new') the latest earlier held stock date's
    records, or None. Both kinds: fewer than EXPECTED_AREAS authorities.
    New: national total_social_stock moving by more than NEW_TOTAL_PCT or
    supported_housing_and_older_people by more than NEW_SUPPORTED_PCT, or
    any authority's total by more than NEW_AREA_PCT, against prev. Revised:
    an authority of the tip missing (a partial file never replaces a fuller
    edition); the national total of any stock column moving by more than
    REV_NATIONAL_PCT; more than REV_MAX_AREAS authorities whose totals (any
    stock column) change; any authority's total moving by more than
    REV_AREA_PCT; more than REV_KEYS_PCT of the tip's provider rows added,
    or removed."""
    g = globals()
    out = []
    an = _area_sums(new)
    want = g["EXPECTED_AREAS"]
    if len(an) < want:
        out.append(f"{len(an)} authorities, expected {want}")
    if kind == "new":
        if not prev:
            return out
        a, b = _national(prev), _national(new)
        for c, lim in (("total_social_stock", g["NEW_TOTAL_PCT"]),
                       ("supported_housing_and_older_people",
                        g["NEW_SUPPORTED_PCT"])):
            if _over(a[c], b[c], lim):
                out.append(f"national {c} {a[c]:,} -> {b[c]:,} against the "
                           f"previous stock date (limit {lim}%)")
        ap = _area_sums(prev)
        big = [f"{k} {ap[k]['total_social_stock']:,}->"
               f"{an[k]['total_social_stock']:,}"
               for k in sorted(set(ap) & set(an))
               if _over(ap[k]["total_social_stock"],
                        an[k]["total_social_stock"], g["NEW_AREA_PCT"])]
        if big:
            out.append(f"total_social_stock moves by more than "
                       f"{g['NEW_AREA_PCT']}% against the previous stock date "
                       f"in {len(big)} authorit(ies): {', '.join(big[:6])}")
        return out
    tip = tip or []
    at = _area_sums(tip)
    gone = sorted(set(at) - set(an))
    if gone:
        out.append(f"{len(gone)} authorit(ies) of the held edition missing "
                   f"from the file (a partial file never replaces a fuller "
                   f"edition): {', '.join(gone[:8])}")
    a, b = _national(tip), _national(new)
    for c in STOCK_COLUMNS:
        if _over(a[c], b[c], g["REV_NATIONAL_PCT"]):
            out.append(f"national {c} {a[c]:,} -> {b[c]:,} (limit "
                       f"{g['REV_NATIONAL_PCT']}%)")
    changed = [k for k in sorted(set(at) | set(an)) if at.get(k) != an.get(k)]
    if len(changed) > g["REV_MAX_AREAS"]:
        out.append(f"{len(changed)} authorities' totals change (limit "
                   f"{g['REV_MAX_AREAS']})")
    big = [f"{k} {at[k]['total_social_stock']:,}->"
           f"{an[k]['total_social_stock']:,}"
           for k in sorted(set(at) & set(an))
           if _over(at[k]["total_social_stock"], an[k]["total_social_stock"],
                    g["REV_AREA_PCT"])]
    if big:
        out.append(f"total_social_stock moves by more than "
                   f"{g['REV_AREA_PCT']}% in {len(big)} authorit(ies): "
                   f"{', '.join(big[:6])}")
    tk = {(r["rp_code"], r["lad24cd"]) for r in tip}
    nk = {(r["rp_code"], r["lad24cd"]) for r in new}
    for what, keys in (("added", nk - tk), ("removed", tk - nk)):
        if tk and len(keys) * 100 > g["REV_KEYS_PCT"] * len(tk):
            out.append(f"{len(keys)} of {len(tk)} provider rows {what} "
                       f"({len(keys) * 100 / len(tk):.1f}%, limit "
                       f"{g['REV_KEYS_PCT']}%)"
                       + ("; a file with fewer providers than the held "
                          "edition is a partial file and never replaces it"
                          if what == "removed" else "")
                       + ": " + ", ".join("/".join(k)
                                          for k in sorted(keys)[:6]))
    return out


def _differs(new, old) -> bool:
    nk = {(r["rp_code"], r["lad24cd"]): r for r in new or ()}
    ok = {(r["rp_code"], r["lad24cd"]): r for r in old or ()}
    if set(nk) != set(ok):
        return True
    return any(nk[k].get(c) != ok[k].get(c) for k in nk for c in COMPARED)


# ---------------------------------------------------------------------------
# Spec and profile
# ---------------------------------------------------------------------------

def tip_row_count(editions_table: str):
    """expected_rows_per_period for status: the tip edition's row count (a
    stock date's providers are its own)."""
    core._ident(editions_table)

    def count(cur, period):
        cur.execute(f"SELECT DISTINCT edition, supersedes FROM "
                    f"public.{editions_table} WHERE stock_date = %s",
                    (period,))
        try:
            tip = core._chain_tip(cur.fetchall(), str(period))
        except (LookupError, ValueError):
            return None
        cur.execute(f"SELECT COUNT(*) FROM public.{editions_table} WHERE "
                    "stock_date = %s AND edition = %s", (period, tip))
        return cur.fetchone()[0]
    return count


SPEC = core.EditionSpec(
    name="s23",
    live_table="rsh_rp_stock_by_la",
    editions_table="rsh_rp_stock_by_la_editions",
    key_cols=KEY,
    period_col="stock_date",
    value_cols=VALUE_TYPES,
    extra_cols=EXTRA_TYPES + PROVENANCE_TYPES,
    refresh_cols=COMPARED + LIVE_PROVENANCE + ("loaded_at",),
    refresh_from=PROVENANCE_PAIRS + (("loaded_at", "loaded_at"),),
    whole_period_cols=LIVE_PROVENANCE,     # one file's provenance per period
    key_types=(("rp_code", "text NOT NULL"),
               ("lad24cd", "varchar(9) NOT NULL"),
               ("stock_date", "date NOT NULL")),
    fk_la_boundaries=True,
    table_constraints=(
        "CONSTRAINT rsh_rp_stock_by_la_editions_provider_type_chk CHECK "
        "(provider_type IN ('PRP', 'LARP'))",
        "CONSTRAINT rsh_rp_stock_by_la_editions_components_sum_chk CHECK "
        "(total_social_stock = general_needs_self_contained + "
        "general_needs_bedspaces + supported_housing_and_older_people + "
        "low_cost_home_ownership)"),
    refresh_key_changes=True,
    expected_rows_per_period=tip_row_count("rsh_rp_stock_by_la_editions"),
)


def _no_label(fetched_on) -> str:
    raise ValueError("an S23 release label comes from the file; store "
                     "through apply_stock_period(...)")


def _provenance_of(records, where="") -> tuple:
    vals = {tuple(r[c] for c in PROVENANCE) for r in records}
    if len(vals) != 1:
        raise ValueError(f"{where}: {len(vals)} provenance values over the "
                         "records; one file's provenance per stock date")
    return vals.pop()


def check_records(records, period) -> None:
    """ValueError unless the stock date's records are whole: each of this
    stock date, one per (rp_code, lad24cd), one provenance."""
    seen = set()
    for r in records:
        if r.get("stock_date") != period:
            raise ValueError(f"{period}: a record of stock date "
                             f"{r.get('stock_date')!r}")
        k = (r["rp_code"], r["lad24cd"])
        if k in seen:
            raise ValueError(f"{period}: {k} appears twice")
        seen.add(k)
    _provenance_of(records, period)


PROFILE = pe.Profile(
    spec=SPEC, value_cols=COMPARED, run_agent=RUN_AGENT,
    run_source=RUN_SOURCE,
    heading="S23 RSH registered provider stock by local authority "
            "(rsh_rp_stock_by_la)",
    default_source_file="RSH registered providers look-up tool",
    # the engine counts areas on the first key column, here rp_code; the 296
    # authorities are enforced by reconcile and period_problems instead
    expected_areas=None,
    release_label=_no_label, content_sha256=rows_content_sha,
    check_records=check_records, file_checks=True, savepoint="s23_period",
    example_label="(stock values)")


def profile() -> pe.Profile:
    """PROFILE bound to the current SPEC (looked up when called, so tests
    can swap the spec)."""
    return PROFILE.with_spec(globals()["SPEC"])


def _late(name: str):
    """A call through this module's attribute `name`, looked up when called
    (so mock.patch.object(s23_rsh_stock_editions, name, ...) reaches the
    engine)."""
    return lambda *a, **kw: globals()[name](*a, **kw)


def create_all(cur) -> None:
    """ddl: the editions table, its append-only triggers and the file-check
    ledger. Idempotent; the live table is not touched."""
    prof = profile()
    core.create_schema(cur, prof.spec)
    pe.create_file_checks(cur, prof)


def insert_live(cur, profile, period: str, records: list) -> None:
    """A new stock date's live rows: the key, the period, the 11 compared
    columns and the five provenance columns from the records' file_* values
    (one value over the records, else ValueError); loaded_at takes its
    default."""
    from psycopg2.extras import execute_values
    spec = profile.spec
    prov = _provenance_of(records, period)
    cols = (tuple(spec.key_cols) + (spec.period_col,)
            + tuple(profile.value_cols) + LIVE_PROVENANCE)
    execute_values(
        cur, f"INSERT INTO public.{spec.live_table} ({', '.join(cols)}) "
        "VALUES %s",
        [tuple(r[k] for k in spec.key_cols) + (period,)
         + tuple(r.get(c) for c in profile.value_cols) + prov
         for r in records],
        page_size=1000)


@dataclass(frozen=True)
class PeriodInfo:
    """What a stock date's edition records: source_file (file name,
    version, month, where read), the release label and the file read
    ((ledger source, sha256),)."""
    source_file: str
    label: str
    files: tuple


def _tip_provenance(cur, spec, period) -> dict:
    tip = core.chain_tip(cur, spec, period)
    cur.execute(f"SELECT DISTINCT {', '.join(PROVENANCE)} FROM "
                f"public.{spec.editions_table} WHERE {spec.period_col} = %s "
                "AND edition = %s", (period, tip))
    rows = cur.fetchall()
    if len(rows) != 1:
        halt(f"{period}: edition {tip} carries {len(rows)} provenance values")
    return dict(zip(PROVENANCE, rows[0]))


def apply_stock_period(cur, profile, period, records, *, fetched_on,
                       info) -> str:
    """pe.apply_period for one stock date, then its ledger row, inside the
    caller's per-period savepoint (s23_period): a new stock date's edition
    1, live rows and ledger row commit or roll back together. A stranded
    period (editions, no live rows) gets its live rows with the tip
    edition's provenance."""
    kind = pe.classify_period(cur, profile, period, records)
    if kind == LIVE_MISSING:
        prov = _tip_provenance(cur, profile.spec, period)
        records = [dict(r, **prov) for r in records]
    prof = dataclasses.replace(profile, release_label=lambda d: info.label)
    kind = pe.apply_period(cur, prof, period, records, fetched_on=fetched_on,
                           source_file=info.source_file,
                           classify=lambda *a, **kw: kind,
                           insert=_late("insert_live"))
    tip = core.chain_tip(cur, prof.spec, period)
    for src, sha in info.files:
        pe.record_file_check(cur, prof, period, src, sha, kind, tip)
    return kind


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

def _session(session):
    if session is not None:
        return session
    import requests
    return requests


def _get(session, url, **kw):
    try:
        r = session.get(url, headers={"User-Agent": USER_AGENT},
                        timeout=300, **kw)
        r.raise_for_status()
    except Exception as e:  # noqa: BLE001 (any failure halts)
        halt(f"GET {url} failed, nothing downloaded: {e}")
    return r


def fetch_json(path, session=None) -> dict:
    """The GOV.UK content API reply for a page path (a read-only GET)."""
    return json.loads(_get(_session(session), API + path).text)


def _abs_url(url) -> str:
    return url if urlparse(url).scheme else GOV_UK + url


def _opens(path) -> None:
    """Raise unless the file opens as an .xlsx workbook with STOCK_BY_LA."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise ValueError(f"no {SHEET} sheet; sheets {wb.sheetnames}")
    finally:
        wb.close()


def fetch(url, dest, session=None) -> tuple:
    """Download url into the directory dest (redirects followed) and return
    (path to read, final URL). The file name is the final URL's. Checked
    before it is kept: more than MIN_FILE_BYTES, an .xlsx name, and it opens
    as an .xlsx workbook with a STOCK_BY_LA sheet. A same-named file with the
    same content is kept as it is; one with different content is never
    replaced: the download is saved beside it as <stem>-<sha8><suffix> and
    that is the one read. An existing name is never a reason to skip the
    download."""
    dest = Path(dest)
    r = _get(_session(session), url)
    final = str(getattr(r, "url", None) or url)
    if final != url:
        print(f"redirected: {url} -> {final}")
    body = r.content
    if len(body) <= MIN_FILE_BYTES:
        halt(f"{final}: {len(body)} bytes, expected more than "
             f"{MIN_FILE_BYTES}; nothing written; saw {body[:100]!r}")
    name = Path(urlparse(final).path).name
    suffix = Path(name).suffix.lower()
    if suffix != ".xlsx":
        halt(f"{final}: not an .xlsx file; nothing written")
    dest.mkdir(parents=True, exist_ok=True)
    target = dest / name
    new_sha = hashlib.sha256(body).hexdigest()
    tmp = dest / (Path(name).stem + ".download.tmp" + suffix)
    tmp.write_bytes(body)
    try:
        _opens(tmp)
    except Exception as e:  # noqa: BLE001
        tmp.unlink()
        halt(f"{final}: does not open as an .xlsx look-up tool ({e}); "
             "nothing kept")
    path = target
    if target.exists():
        old_sha = content_sha256(target)
        if old_sha == new_sha:
            tmp.unlink()
            print(f"{name} already held with the same content (sha256 "
                  f"{new_sha[:16]}); kept as it is")
        else:
            path = dest / f"{target.stem}-{new_sha[:8]}{target.suffix}"
            if path.exists() and content_sha256(path) == new_sha:
                tmp.unlink()
                print(f"NOTE: {name} differs from this download; "
                      f"{path.name} already holds it (sha256 "
                      f"{new_sha[:16]})")
            else:
                tmp.replace(path)
                print(f"NOTE: {name} already exists with different content "
                      f"(sha256 {old_sha[:16]}, new {new_sha[:16]}); kept it "
                      f"untouched and saved the new download as {path.name}")
    else:
        tmp.replace(target)
        print(f"downloaded {name}: {len(body):,} bytes, sha256 "
              f"{new_sha[:16]}")
    return path, final


def _file_source(path: Path) -> str:
    """Where a --file was read: the path under the repo (data/raw/...),
    else the absolute path."""
    p = Path(path).resolve()
    try:
        return p.relative_to(REPO).as_posix()
    except ValueError:
        return p.as_posix()


# ---------------------------------------------------------------------------
# Reading held data
# ---------------------------------------------------------------------------

table_exists = pe.table_exists


def _conn(writing):
    return pe.connect(writing)


def records(cur, spec, period, edition=None) -> list:
    """The period's rows as records (key, period as a string, the 11
    compared columns and the five file_* provenance values): the live
    table's (edition None; its provenance columns read as file_*) or the
    given edition's."""
    names = tuple(spec.key_cols) + COMPARED + PROVENANCE
    if edition is None:
        live = dict(PROVENANCE_PAIRS)
        back = {ec: lc for lc, ec in live.items()}
        sel = [back.get(c, c) for c in names]
        cur.execute(f"SELECT {', '.join(sel)} FROM public.{spec.live_table} "
                    f"WHERE {spec.period_col} = %s", (period,))
    else:
        cur.execute(f"SELECT {', '.join(names)} FROM "
                    f"public.{spec.editions_table} WHERE {spec.period_col} = "
                    "%s AND edition = %s", (period, edition))
    out = [dict(zip(names, r)) for r in cur.fetchall()]
    for r in out:
        r[spec.period_col] = str(period)
    return out


def tip_info(cur, spec) -> dict:
    """{period (string): {edition, source_file, rank}} of every period with
    editions."""
    cur.execute(f"SELECT DISTINCT {spec.period_col} FROM "
                f"public.{spec.editions_table}")
    out = {}
    for (p,) in cur.fetchall():
        tip = core.latest_edition(cur, spec, p)
        cur.execute(f"SELECT DISTINCT source_file FROM "
                    f"public.{spec.editions_table} WHERE {spec.period_col} = "
                    "%s AND edition = %s", (p, tip))
        src = {r[0] for r in cur.fetchall()}
        sf = src.pop() if len(src) == 1 else None
        out[pe._p(p)] = {"edition": tip, "source_file": sf,
                         "rank": release_rank(sf)}
    return out


def legacy_tips(cur, spec) -> dict:
    """Before the editions table exists: {held live period: {rank}}, the
    rank of the held file (LEGACY_RANK)."""
    cur.execute(f"SELECT DISTINCT {spec.period_col} FROM "
                f"public.{spec.live_table}")
    return {pe._p(p): {"edition": None, "source_file": None,
                       "rank": LEGACY_RANK} for (p,) in cur.fetchall()}


def ledger_checks(cur, prof) -> dict:
    """{period (string): {(source_file, file_sha256)}} of the ledger."""
    pc = prof.spec.period_col
    cur.execute(f"SELECT DISTINCT {pc}, source_file, file_sha256 "
                f"FROM public.{pe.file_checks_table(prof)}")
    out = {}
    for p, f, s in cur.fetchall():
        out.setdefault(pe._p(p), set()).add((f, s))
    return out


# ---------------------------------------------------------------------------
# ddl, status, run log, refresh-latest
# ---------------------------------------------------------------------------

def log_run(cur, rows_written: int, notes: str, started_at=None) -> None:
    """The pipeline_run_log row for a committed run (agent RUN_AGENT,
    source '23'). Called only on committed load, migrate-legacy and
    restore-edition runs, partial runs included."""
    pe.log_run(cur, profile(), rows_written, notes, started_at)


def cmd_ddl(args) -> int:
    return pe.run_ddl(profile(), args, connect=_late("_conn"),
                      table_exists=_late("table_exists"),
                      create_schema=lambda cur: create_all(cur))


def status_lines(cur, prof) -> tuple:
    """(lines, ok): per period the tip's edition and file version and date,
    whether live provenance is uniform and the tip's, and the ledger's
    latest file. ok is false when a period's live provenance is not
    uniform, or uniform but not the tip's while live equals the tip."""
    s = prof.spec
    lines, ok = [], True
    lt = pe.file_checks_table(prof)
    for p, t in sorted(tip_info(cur, s).items()):
        line = (f"  {p}: tip edition {t['edition']} ({rank_text(t['rank'])})")
        cur.execute(f"SELECT DISTINCT {', '.join(PROVENANCE)} FROM "
                    f"public.{s.editions_table} WHERE {s.period_col} = %s AND "
                    "edition = %s", (p, t["edition"]))
        tipp = set(cur.fetchall())
        cur.execute(f"SELECT DISTINCT {', '.join(LIVE_PROVENANCE)} FROM "
                    f"public.{s.live_table} WHERE {s.period_col} = %s", (p,))
        livep = set(cur.fetchall())
        current = core.rows_differing(cur, s, p, t["edition"]) == 0
        if len(livep) == 1 and livep == tipp:
            line += "; live provenance uniform, the tip's"
        elif len(livep) == 1:
            line += "; live provenance uniform, not the tip's"
            if current:
                ok = False
                line += " (ACTION NEEDED)"
            else:
                line += " (refresh-latest sets it)"
        elif livep:
            ok = False
            line += (f"; live provenance not uniform ({len(livep)} values): "
                     "ACTION NEEDED")
        cur.execute(f"SELECT source_file, outcome FROM public.{lt} WHERE "
                    f"{s.period_col} = %s ORDER BY id DESC LIMIT 1", (p,))
        last = cur.fetchone()
        line += (f"; ledger latest: {last[1]} {last[0]}" if last
                 else "; no ledger row")
        lines.append(line)
    return lines, ok


def cmd_status(args) -> int:
    """status: what needs action; exit 1 if anything, or (with a clean
    message, nothing created) if the editions table or ledger does not
    exist."""
    prof = profile()
    conn = _conn(False)
    try:
        with conn.cursor() as cur:
            missing = [t for t in (prof.spec.editions_table,
                                   pe.file_checks_table(prof))
                       if not table_exists(cur, t)]
            if missing:
                print(f"status: {' and '.join(missing)} "
                      f"{'does' if len(missing) == 1 else 'do'} not exist "
                      "yet; run `ddl --commit`, then `migrate-legacy FILE "
                      "--commit` (S23 has no sync-new)")
                return 1
            st = pe.status(cur, prof)
            lines, prov_ok = status_lines(cur, prof)
    finally:
        conn.rollback()
        conn.close()
    print(pe.format_status(prof, st, lines))
    return 0 if st["ok"] and prov_ok else 1


def live_equals_tip(cur, spec, periods) -> list:
    """Problems (empty = fine): for each period, the live rows equal the
    tip edition on the key, the stock date, the 11 compared columns and the
    five provenance columns (live column against its file_* column),
    NULL-safe, both ways. (load_checks.check_latest_equals_live compares
    every data column by name, and the file_* columns are not live
    columns.)"""
    ed_cols = ", ".join(tuple(spec.key_cols) + (spec.period_col,) + COMPARED
                        + PROVENANCE)
    lv_cols = ", ".join(tuple(spec.key_cols) + (spec.period_col,) + COMPARED
                        + LIVE_PROVENANCE)
    bad = []
    for p in periods:
        try:
            ed = core.chain_tip(cur, spec, p)
        except (LookupError, ValueError) as e:
            bad.append(f"{pe._p(p)}: {e}")
            continue
        a = (f"SELECT {ed_cols} FROM public.{spec.editions_table} "
             f"WHERE {spec.period_col} = %s AND edition = %s")
        b = (f"SELECT {lv_cols} FROM public.{spec.live_table} "
             f"WHERE {spec.period_col} = %s")
        counts = []
        for x, xa, y, ya in ((a, (p, ed), b, (p,)), (b, (p,), a, (p, ed))):
            cur.execute(f"SELECT COUNT(*) FROM (({x}) EXCEPT ALL ({y})) q",
                        xa + ya)
            counts.append(cur.fetchone()[0])
        if any(counts):
            bad.append(f"{pe._p(p)} ed{ed}: {counts[0]} edition-only, "
                       f"{counts[1]} live-only rows")
    return bad


def cmd_refresh_latest(args) -> int:
    """refresh-latest [--accept-drift P] [--accept-key-changes P]
    [--commit | --simulate]: the engine's refresh (period_editions.
    run_refresh_latest's steps: preview of the rows and the keys added and
    removed, the refusals, editions_core.refresh_latest under its
    before/after guard, provenance set on every row of each refreshed period
    and loaded_at copied), with S23's after-check live_equals_tip."""
    prof = profile()
    spec = prof.spec
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            for t in (spec.editions_table, pe.file_checks_table(prof)):
                if not table_exists(cur, t):
                    halt(f"{t} does not exist yet; run `ddl --commit`, then "
                         "`migrate-legacy FILE --commit`")
            accept = pe._accept_periods(cur, spec, tuple(args.accept_drift
                                                         or ()))
            accept_kc = pe._accept_key_change_periods(
                cur, spec, pe._accept_key_change_args(args))
            counts = core.refresh_counts(cur, spec, accept)
            print("rows refresh-latest would write: "
                  + (", ".join(f"{pe._p(p)}={n}" for p, n in counts.items())
                     or "none")
                  + f" (total {sum(counts.values())})")
            plan, _ = core._plan(cur, spec, accept)
            lines = core.key_change_lines(plan, accept_kc, name=pe._p)
            print("keys added and removed:" + ("" if lines else " none"))
            for line in lines:
                print(f"  {line}")
            if not writing:
                print("DRY RUN: nothing written (use --commit or --simulate)")
                return 0
            pe._refresh_guard(cur, spec, accept, accept_kc)
            res = core.refresh_latest(cur, spec, accept, accept_kc)
            if res["updated"]:
                bad = live_equals_tip(cur, spec, sorted(res["updated"]))
                if bad:
                    halt("refresh-latest: live differs from the latest edition "
                         "after the refresh, rolled back: " + "; ".join(bad[:6]))
        keys = ""
        if res["inserted"] or res["deleted"]:
            keys = (f"; key changes {sum(res['inserted'].values())} inserted, "
                    f"{sum(res['deleted'].values())} deleted in "
                    f"{sorted(pe._p(x) for x in res['inserted'])}")
        pe.finish(conn, args, f"{res['rows']} live rows refreshed in "
                              f"{sorted(pe._p(x) for x in res['updated'])}"
                              f"{keys}; provenance set on every row of each "
                              "refreshed stock date; before/after guard "
                              "passed")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# load
# ---------------------------------------------------------------------------

def _end(conn, args, writing, rc) -> int:
    conn.rollback()  # commit mode has committed already
    if not writing:
        print("PREVIEW: nothing written to the database (use --commit or "
              "--simulate)")
    elif args.simulate:
        print("SIMULATION: ROLLED BACK (nothing persisted)")
    return rc


def _release_arg(text) -> tuple:
    mm = re.fullmatch(r"([0-9]{4})-([0-9]{4})", str(text or ""))
    if not mm or int(mm.group(2)) != int(mm.group(1)) + 1:
        raise argparse.ArgumentTypeError(f"not a release Y1-Y2 (e.g. "
                                         f"2024-2025): {text!r}")
    return int(mm.group(1)), int(mm.group(2))


def _period_arg(text) -> str:
    try:
        return date.fromisoformat(str(text)).isoformat()
    except ValueError:
        raise argparse.ArgumentTypeError(f"not a stock date YYYY-MM-DD: "
                                         f"{text!r}") from None


def _read(path) -> dict:
    try:
        return read_tool(path)
    except ValueError as e:
        halt(f"{Path(path).name}: {e}; nothing stored")


def identity_problems(tool, page_years=None, release=None) -> list:
    """Problems (empty = fine) between a read tool and where it came from:
    the release page's years (when the page was read) and --release must
    both equal the file's own source line years."""
    out = []
    own = (tool["y1"], tool["y2"])
    if page_years is not None and tuple(page_years) != own:
        out.append(f"the release page is {page_years[0]} to {page_years[1]}, "
                   f"the file says {tool['edition']}")
    if release is not None and tuple(release) != own:
        out.append(f"--release {release[0]}-{release[1]}, the file says "
                   f"{tool['edition']}")
    return out


def _print_page(page, base):
    print(f"release: {page.get('title')} ({GOV_UK}{base}); first published "
          f"{str(page.get('first_published_at'))[:10]}; updated "
          f"{str(page.get('public_updated_at'))[:10]}")
    for ts, note in change_history(page)[:4]:
        print(f"  change_history {ts}: {note[:200]}")


def _page(base, years):
    page = fetch_json(base)
    py = page_years(page)
    if py != tuple(years):
        halt(f"{GOV_UK}{base}: the page title {page.get('title')!r} is not "
             f"the {years[0]} to {years[1]} release")
    _print_page(page, base)
    return page


def _source(args) -> dict:
    """The file the run reads: {path, where, how, page_years, page_url,
    source_url, file_name, tool}. --file reads a local file (nothing
    downloaded; it is read here, tool set) and finds its own years' release
    page in the collection unless --no-page; otherwise the content API finds
    the release (the newest, or --release Y1-Y2) and its look-up tool is
    downloaded to RAW_DIR (also in a preview)."""
    out = {"path": None, "where": None, "how": None, "page_years": None,
           "page_url": None, "source_url": None, "file_name": None,
           "tool": None}
    if args.file:
        path = Path(args.file)
        if not path.is_file():
            halt(f"--file {args.file}: no such file")
        print(f"file given: {path}; nothing downloaded; identity is read "
              "from the file")
        tool = _read(path)
        bad = identity_problems(tool, None, args.release)
        if bad:
            halt("file identity check failed, nothing stored: "
                 + "; ".join(bad))
        out.update(path=path, where=_file_source(path), how="local file",
                   source_url=f"file:{path.name}", file_name=path.name,
                   tool=tool)
        if args.no_page:
            out.update(page_url=NO_PAGE, how="local file, --no-page: the "
                       "release page was not read")
            print("NOTE: --no-page: the release page is not read; "
                  f"source_url is recorded as file:{path.name} and "
                  f"release_page_url as {NO_PAGE!r}")
            return out
        coll = fetch_json(COLLECTION_PATH)
        try:
            y1, y2, base = release_for_years(coll, tool["y1"], tool["y2"])
        except ValueError as e:
            halt(f"{GOV_UK}{COLLECTION_PATH}: {e}; the release page of the "
                 f"file's own years ({tool['edition']}) was not found. Give "
                 "--no-page to load the file without it (recorded as such)")
        _page(base, (y1, y2))
        out.update(page_years=(y1, y2), page_url=GOV_UK + base)
        return out
    coll = fetch_json(COLLECTION_PATH)
    try:
        if args.release:
            y1, y2, base = release_for_years(coll, *args.release)
            how = f"release {y1} to {y2} given"
        else:
            y1, y2, base = latest_release(coll)
            how = "newest release (content API)"
    except ValueError as e:
        halt(f"{GOV_UK}{COLLECTION_PATH}: {e}")
    page = _page(base, (y1, y2))
    try:
        att = lookup_tool_attachment(page)
    except ValueError as e:
        halt(f"{GOV_UK}{base}: {e}")
    print(f"  attachment: {att['title']!r} {att['url']}")
    print(f"downloading to {RAW_DIR} (also in a preview; the database is "
          "not written)")
    path, final = fetch(_abs_url(att["url"]), RAW_DIR)
    print(f"  final URL: {final}")
    out.update(path=path, where=final, how=how, page_years=(y1, y2),
               page_url=GOV_UK + base, source_url=final,
               file_name=Path(urlparse(final).path).name)
    return out


def provenance(tool, src) -> dict:
    """The five file_* values a load records: the file's own edition and
    publication month (its first day), where it was read, its name, the
    release page read."""
    return {"file_edition": tool["edition"],
            "file_publication_date": tool["published"],
            "file_source_url": src["source_url"],
            "file_name": src["file_name"],
            "file_release_page_url": src["page_url"]}


def edition_source(tool, src) -> str:
    return (f"{src['file_name']}; version {tool['version']}; dated "
            f"{tool['published']:%Y-%m}; {src['where']}")


def _label(tool, src, sha, how, *, older="") -> str:
    return (f"RSH registered providers look-up tool: {tool['title']}; "
            f"version {tool['version']}; dated {tool['published']:%Y-%m}; "
            f"{src['where']} sha256 {sha[:16]}; {how}"
            + (f"; {older}" if older else ""))


def _report(tool) -> None:
    prp = sum(1 for p in tool["providers"]
              if PROVIDER_TYPES[p["rp_type"]] == "PRP")
    print(f"{tool['path'].name}: {tool['title']!r}; source 1 April "
          f"{tool['y1']} to 31 March {tool['y2']} (stock date "
          f"{tool['stock_date']}); publication {tool['month_text']}; version "
          f"{tool['version']}; Version History "
          + ", ".join(f"{v} {d:%B %Y}" for v, d, _ in tool["history"]))
    print(f"  {SHEET}: {len(tool['providers']):,} provider rows (PRP {prp:,}, "
          f"LARP {len(tool['providers']) - prp:,}), "
          f"{len(tool['subtotals'])} LA subtotal rows, "
          f"{len(tool['regions'])} region rows, {tool['empty_rows']:,} empty "
          "rows at the end")


def _reconciled(tool) -> None:
    problems = reconcile(tool)
    if problems:
        halt("reconciliation inside the file failed, nothing stored: "
             + "; ".join(problems[:8]))
    eng = england_totals(tool)
    print(f"reconciled: the provider rows sum to the {len(tool['subtotals'])} "
          "LA subtotal rows on all five measures, and the subtotals to the "
          f"{len(tool['regions'])} region rows; England (the regions) "
          + " / ".join(f"{eng[c]:,}" for c in STOCK_COLUMNS)
          + f" ({', '.join(STOCK_COLUMNS)})")


def _build(cur, tool, prov) -> list:
    try:
        rmap, problems = resolve_codes(cur, tool_codes(tool),
                                       tool["stock_date"])
    except ValueError as e:
        halt(f"geography: {e}; nothing stored")
    if problems:
        halt("geography check failed, nothing stored: " + "; ".join(problems))
    recoded = sorted((c, r) for c, r in rmap.items() if r != c)
    if recoded:
        print("codes resolved (geography.resolve, la_code_lookup recode): "
              + ", ".join(f"{c} -> {r}" for c, r in recoded))
    try:
        return build_rows(tool, rmap, prov)
    except ValueError as e:
        halt(f"{e}; nothing stored")


def cmd_load(args) -> int:
    if args.no_page and not args.file:
        halt("--no-page is only for --file (a local file whose release page "
             "is not read)")
    prof = profile()
    spec = prof.spec
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            lt = pe.file_checks_table(prof)
            ready = {spec.editions_table: table_exists(cur, spec.editions_table),
                     lt: table_exists(cur, lt)}
            has_ed = all(ready.values())
            missing = [t for t, ok in ready.items() if not ok]
            if writing and not has_ed:
                halt(f"missing {', '.join(missing)}; run `ddl --commit` and "
                     "`migrate-legacy FILE --commit` first")
            if has_ed:
                _, new_live, errors = core.latest_map(cur, spec)
                if errors:
                    halt("invalid edition chain: " + "; ".join(
                        f"{pe._p(p)}: {msg}" for p, msg in errors.items()))
                if new_live:
                    halt(f"{spec.live_table}: live stock dates with no "
                         f"editions {[pe._p(x) for x in new_live]}; run "
                         "migrate-legacy first")
                tips = tip_info(cur, spec)
                checked = ledger_checks(cur, prof)
                stranded = pe.live_missing_periods(cur, prof)
            else:
                print(f"NOTE: {', '.join(missing)} not present: no editions "
                      "table yet; this preview compares the file with the "
                      "LIVE table")
                tips, checked, stranded = legacy_tips(cur, spec), {}, []
            ranks = tip_ranks(tips, checked)
            held = sorted(ranks)
            print("held stock dates: " + (", ".join(held) or "none"))
            src = _source(args)
            path = src["path"]
            sha = content_sha256(path)
            print(f"{path.name}: sha256 {sha[:16]}"
                  + ("; byte-identical to the held file the old build read"
                     if sha == LEGACY_FILE["sha256"] else ""))
            tool = src["tool"]
            if has_ed and args.recheck is None:
                done = ledger_complete(checked, src["where"], sha, ranks,
                                       allow_older=args.allow_older_file,
                                       stranded=stranded)
                if done:
                    print(f"the file's (source, sha256) is already in the "
                          f"ledger for {', '.join(done)}; "
                          + ("nothing compared" if tool else "nothing parsed")
                          + " (--recheck YYYY-MM-DD reads it again)")
                    print("nothing to do")
                    return _end(conn, args, writing, rc=0)
            if tool is None:
                tool = _read(path)
                bad = identity_problems(tool, src["page_years"], args.release)
                if bad:
                    halt("file identity check failed, nothing stored: "
                         + "; ".join(bad))
            _report(tool)
            _reconciled(tool)
            period = tool["stock_date"]
            recs = _build(cur, tool, provenance(tool, src))
            lsrc = ledger_source(src["where"], tool["rank"], tool["version"],
                                 period)
            chk = {p: s for p, s in checked.items() if p not in stranded}
            try:
                new, revised, skipped = plan_periods(
                    [period], tool["rank"], ranks, chk, lsrc, sha,
                    recheck=args.recheck, allow_older=args.allow_older_file)
            except ValueError as e:
                halt(str(e))
            rc = _run_period(args, conn, cur, prof, has_ed, src, tool, recs,
                             new, revised, skipped, tips, ranks, held, sha,
                             lsrc)
        return _end(conn, args, writing, rc)
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def _tip_records(cur, spec, p, tips, has_ed) -> list:
    if has_ed:
        return records(cur, spec, p, tips[p]["edition"])
    return records(cur, spec, p)


def _run_period(args, conn, cur, prof, has_ed, src, tool, recs, new, revised,
                skipped, tips, ranks, held, sha, lsrc) -> int:
    started = datetime.now(timezone.utc)
    spec = prof.spec
    rank = tool["rank"]
    older_note = {}
    for p in revised:
        t = ranks.get(p)
        if t is not None and t > rank:
            older_note[p] = (f"NOTE: --allow-older-file given; {p}'s tip "
                             f"comes from a file of {rank_text(t)}, this file "
                             f"is {rank_text(rank)}")
            print(older_note[p])
    print(f"planned: new {new or 'none'}; compare {revised or 'none'}; "
          "skipped " + (", ".join(f"{p} ({r})" for p, r in sorted(
              skipped.items())) or "none"))
    periods = sorted(new + revised)
    if not periods:
        if skipped and all(r.startswith("older") for r in skipped.values()):
            halt(f"older file: every stock date of the file is skipped; the "
                 f"file is {rank_text(rank)}, older than the tip of "
                 f"{', '.join(sorted(skipped))}; storing it would record an "
                 "older file. If this is deliberate, re-run with "
                 "--allow-older-file")
        print("nothing to do")
        return 0
    problems, infos = {}, {}
    for p in periods:
        if p in revised:
            kind = "revised"
            tip = _tip_records(cur, spec, p, tips, has_ed)
            prev = None
        else:
            kind = "new"
            tip = None
            earlier = [h for h in held if h < p]
            prev = (_tip_records(cur, spec, max(earlier), tips, has_ed)
                    if earlier else None)
        bad = period_problems(recs, tip, prev, kind=kind)
        if kind == "revised" and ranks.get(p) == rank and _differs(recs, tip):
            bad.append(f"the file is {rank_text(rank)}, the same version and "
                       "month as the file of the held tip, but its content "
                       "differs: two files claim the same version")
        if bad:
            problems[p] = bad
            continue
        how = (("new stock date" if kind == "new" else "republished stock "
                "date") + f" ({src['how']})")
        infos[p] = PeriodInfo(
            edition_source(tool, src),
            _label(tool, src, sha, how,
                   older=older_note.get(p, "").removeprefix("NOTE: ")),
            ((lsrc, sha),))
    for p, msgs in sorted(problems.items()):
        for msg in msgs:
            print(f"  {p}: REJECTED, not stored: STOP CONDITION: {msg}")
    ok = [p for p in periods if p not in problems]
    rc = 0
    stats = {"periods": [], "kinds": {}, "stored_rows": 0, "live_rows": 0}
    if ok and has_ed:
        def apply(c, prof_, per, rs, *, fetched_on, source_file=None):
            return _late("apply_stock_period")(c, prof_, per, rs,
                                               fetched_on=fetched_on,
                                               info=infos[per])
        rc = pe.load_periods(cur, prof, ok, lambda per: recs,
                             tool["published"], args.commit,
                             simulate=args.simulate, against="editions",
                             stats=stats, apply=apply)
    elif ok:
        rc = pe.load_periods(cur, prof, ok, lambda per: recs,
                             tool["published"], False, against="live")
        print("  (no editions table yet: preview against live only)")
    if problems:
        print(f"REJECTED {len(problems)} stock date(s), nothing stored for "
              f"them: {', '.join(sorted(problems))}; exit 1")
        rc = 1
    if args.commit:
        stored = live = 0  # rows written by this run; not a source value
        for p in stats["periods"]:
            k = stats["kinds"][p]
            if k in ("new", "revised"):
                stored += len(recs)
            if k in ("new", LIVE_MISSING):
                live += len(recs)
        left = sorted(p for p in ok if p not in stats["periods"])
        partial = (f"PARTIAL RUN (exit 1): rejected {sorted(problems) or 'none'}"
                   f"; failed or not attempted {left or 'none'}; a stock date "
                   "that failed or was rejected stored nothing and is not in "
                   "the counts below; " if rc else "")
        notes = (partial + f"look-up tool {src['where']} sha256 {sha[:16]}, "
                 f"{rank_text(rank)} ({src['how']}): "
                 + ("; ".join(f"{p} {k}" for p, k in stats["kinds"].items())
                    or "nothing stored")
                 + (f"; skipped {sorted(skipped)}" if skipped else "")
                 + "".join(f"; {older_note[p]}" for p in sorted(older_note))
                 + f". Rows stored: {stored} edition rows, {live} live rows.")
        log_run(cur, stored + live, notes, started)
        conn.commit()
        print("pipeline_run_log row written"
              + (" (partial run: see its notes)" if rc else ""))
    return rc


# ---------------------------------------------------------------------------
# One-off: migrate the held table
# ---------------------------------------------------------------------------

def live_state(cur, spec=None) -> tuple:
    """(rows, survey hash, distinct loaded_at, distinct stock dates) of the
    live table. The survey hash is the spec's: md5 of every stored column
    but loaded_at, '|'-joined (NULL size band and status as '~'), rows
    ordered by rp_code, lad24cd."""
    spec = spec or globals()["SPEC"]
    cur.execute(f"""SELECT COUNT(*), md5(string_agg(stock_date||'|'||rp_code
        ||'|'||lad24cd||'|'||rp_name||'|'||provider_type||'|'
        ||coalesce(rp_size_band,'~')||'|'||coalesce(survey_status,'~')||'|'
        ||publisher_la_code||'|'||la_name||'|'||total_social_stock||'|'
        ||general_needs_self_contained||'|'||general_needs_bedspaces||'|'
        ||supported_housing_and_older_people||'|'||low_cost_home_ownership
        ||'|'||publication_date||'|'||edition||'|'||source_url||'|'
        ||source_file||'|'||release_page_url, E'\\n' ORDER BY rp_code,
        lad24cd)), COUNT(DISTINCT loaded_at), COUNT(DISTINCT stock_date)
        FROM public.{spec.live_table}""")
    return tuple(cur.fetchone())


def _proof(held, built) -> tuple:
    """(cells compared, differences) of the held rows against the rebuilt
    rows on the key, the stock date and the 11 compared columns."""
    hk = {(r["rp_code"], r["lad24cd"]): r for r in held}
    bk = {(r["rp_code"], r["lad24cd"]): r for r in built}
    diffs = [("only held", k) for k in sorted(set(hk) - set(bk))]
    diffs += [("only rebuilt", k) for k in sorted(set(bk) - set(hk))]
    cells = 0  # cells compared; not a source value
    for k in sorted(set(hk) & set(bk)):
        for c in ("stock_date",) + COMPARED:
            cells += 1
            if hk[k].get(c) != bk[k].get(c):
                diffs.append((k, c, hk[k].get(c), bk[k].get(c)))
    return cells, diffs


def migrate_legacy(cur, file, *, write, legacy=None, files=None) -> dict:
    """One-off. Preconditions: on write the editions table and ledger exist
    and are empty (in a preview they may be absent; if present they must be
    empty); live as surveyed (legacy: (rows, survey hash), default
    LEGACY_LIVE; one loaded_at; one stock date); the file's sha256 as
    files['sha256'] (default LEGACY_FILE). Proof, any failure halts: the
    file's identity and in-file reconciliation; the live provenance names
    the file's edition, name and URL; this parser on the file reproduces
    every live row on the key, the stock date and the 11 compared columns
    (publication_date and release_page_url are the release page's, not in
    the file, and are carried as held). Edition 1 'as loaded' from live
    (release label the engine's as-loaded label, published_date the load
    date, source_file 'as loaded: <name>; version <v>; dated <yyyy-mm>',
    file_* the live provenance), and a ledger row for the file (outcome
    'unchanged', edition 1) so the next load skips it by (URL, sha). Live
    untouched. Nothing commits here; the caller rolls everything back on a
    halt."""
    g = globals()
    legacy = tuple(legacy or g["LEGACY_LIVE"])
    files = files or g["LEGACY_FILE"]
    prof = profile()
    spec = prof.spec
    lt = pe.file_checks_table(prof)
    ready = {spec.editions_table: table_exists(cur, spec.editions_table),
             lt: table_exists(cur, lt)}
    if write and not all(ready.values()):
        halt(f"missing {[t for t, ok in ready.items() if not ok]}; run "
             "`ddl --commit` first")
    for t, ok in ready.items():
        if ok:
            cur.execute(f"SELECT COUNT(*) FROM public.{t}")
            n = cur.fetchone()[0]
            if n:
                halt(f"{t} already holds {n} rows; migrate-legacy runs once, "
                     "on empty tables")
    n, h, k, nper = live_state(cur, spec)
    if (n, h) != legacy or k != 1 or nper != 1:
        halt(f"{spec.live_table} is not as surveyed: {n} rows, hash {h}, {k} "
             f"loaded_at value(s), {nper} stock date(s); expected "
             f"{legacy[0]} rows, hash {legacy[1]}, one loaded_at, one stock "
             "date; nothing stored")
    path = Path(file)
    sha = content_sha256(path)
    if sha != files["sha256"]:
        halt(f"{path.name}: sha256 {sha[:16]}, expected "
             f"{files['sha256'][:16]} (the held file); nothing stored")
    tool = _read(path)
    _report(tool)
    _reconciled(tool)
    period = tool["stock_date"]
    cur.execute(f"SELECT DISTINCT {spec.period_col} FROM "
                f"public.{spec.live_table}")
    live_periods = [pe._p(r[0]) for r in cur.fetchall()]
    if live_periods != [period]:
        halt(f"live stock dates {live_periods}, the file holds {period}; "
             "nothing stored")
    held = records(cur, spec, period)
    prov = dict(zip(PROVENANCE, _provenance_of(held, period)))
    want = {"file_edition": tool["edition"], "file_name": files["name"],
            "file_source_url": files["url"]}
    off = {c: (prov[c], v) for c, v in want.items() if prov[c] != v}
    if off:
        halt(f"proof failed: the live provenance does not name the file: "
             f"{off}; nothing stored")
    built = _build(cur, tool, prov)
    cells, diffs = _proof(held, built)
    if diffs:
        halt(f"proof failed: {len(diffs)} differences between the held rows "
             f"and the held file read again, e.g. {diffs[:5]}; nothing "
             "stored")
    cur.execute(f"SELECT MIN((loaded_at AT TIME ZONE 'UTC')::date) FROM "
                f"public.{spec.live_table}")
    loaded = cur.fetchone()[0]
    plan = {"period": period, "rows": len(held), "cells": cells,
            "records": held, "sha": sha, "loaded": loaded,
            "source_file": (f"as loaded: {files['name']}; version "
                            f"{tool['version']}; dated "
                            f"{tool['published']:%Y-%m}"),
            "ledger": ledger_source(files["url"], tool["rank"],
                                    tool["version"], period)}
    print(f"proof: the held file read again reproduces every held row: "
          f"{len(held):,} rows matched by key, {cells:,} cells (the stock "
          "date and the 11 compared columns), 0 differences; live edition "
          f"{tool['edition']!r} is the file's own, source_file and "
          "source_url name the held file; publication_date "
          f"{prov['file_publication_date']} and release_page_url are the "
          "release page's (not in the file), carried as held")
    print(f"  {period}: edition 1 as loaded, {len(held):,} rows, published "
          f"{loaded} (the load date); source_file {plan['source_file']!r}; "
          f"ledger {plan['ledger']}")
    if not write:
        return plan
    cur.execute(f"SAVEPOINT {prof.savepoint}")
    ed = core.insert_edition(
        cur, spec, held, period, release_label=core.AS_LOADED_LABEL,
        published_date=loaded, source_file=plan["source_file"],
        source_sha256=rows_content_sha(held), supersedes=None, strict=True)
    if ed != 1 or core.rows_differing(cur, spec, period, ed):
        halt(f"{spec.editions_table} {period}: edition 1 does not equal "
             "live; rolled back")
    pe.record_file_check(cur, prof, period, plan["ledger"], sha,
                         "unchanged", 1)
    cur.execute(f"RELEASE SAVEPOINT {prof.savepoint}")
    return plan


def cmd_migrate_legacy(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    try:
        with conn.cursor() as cur:
            plan = migrate_legacy(cur, args.file, write=writing)
            if args.commit:
                log_run(cur, plan["rows"], f"migrate-legacy: edition 1 as "
                        f"loaded for {plan['period']} ({plan['rows']} rows) "
                        f"from the live table; proof {plan['cells']} cells "
                        "equal the held file read again "
                        f"(sha256 {plan['sha'][:16]}); ledger row for the "
                        "file. Live untouched.", started)
                conn.commit()
                print("migrate-legacy: COMMITTED; pipeline_run_log row "
                      "written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# restore-edition
# ---------------------------------------------------------------------------

def restore_edition(cur, period, edition) -> int:
    """Store edition `edition` of `period` as the next edition ('restored
    from edition N') for refresh-latest to apply; return its number. Halts
    if the edition does not exist or is the tip."""
    spec = globals()["SPEC"]
    p = str(period)
    tip = core.latest_edition(cur, spec, p)
    recs = records(cur, spec, p, edition)
    if not recs:
        halt(f"{spec.editions_table} {p}: no edition {edition}")
    if edition == tip:
        halt(f"{p}: edition {edition} is the tip; nothing to restore")
    cur.execute(f"SELECT DISTINCT source_file FROM public.{spec.editions_table}"
                f" WHERE {spec.period_col} = %s AND edition = %s", (p, edition))
    src = cur.fetchone()[0]
    new = core.insert_edition(
        cur, spec, recs, p, release_label=f"restored from edition {edition}",
        published_date=date.today(), source_file=src,
        source_sha256=rows_content_sha(recs), supersedes=tip, strict=True,
        allow_revert=True)
    if new == tip:
        halt(f"{p}: edition {edition} was not stored as the next edition; "
             "nothing to restore")
    return new


def cmd_restore_edition(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    spec = globals()["SPEC"]
    try:
        with conn.cursor() as cur:
            if not table_exists(cur, spec.editions_table):
                halt(f"{spec.editions_table} does not exist")
            new = restore_edition(cur, args.period, args.edition)
            print(f"{args.period}: edition {args.edition} "
                  + ("stored" if writing else "would be stored")
                  + f" as edition {new}; run refresh-latest to apply it")
            if args.commit:
                n = len(records(cur, spec, args.period, new))
                log_run(cur, n, f"restore-edition: {args.period} edition "
                        f"{args.edition} stored as edition {new}", started)
                conn.commit()
                print("COMMITTED; pipeline_run_log row written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="S23 RSH registered provider stock by local authority: "
        "editions loader on the period-editions engine. There is no "
        "sync-new; migrate-legacy records the held stock date once.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ddl", help="create the editions table, its triggers "
                       "and the file-check ledger (preview by default)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_ddl)
    sub.add_parser("status", help="what needs action; exit 1 if anything"
                   ).set_defaults(func=cmd_status)
    p = sub.add_parser("load", help="find the newest release's look-up tool, "
                       "compare and store its stock date (preview by "
                       "default)")
    p.add_argument("--release", type=_release_arg, metavar="Y1-Y2",
                   help="that release's page (e.g. 2024-2025; default the "
                   "newest)")
    p.add_argument("--file", metavar="PATH",
                   help="a local look-up tool (nothing downloaded; identity "
                   "is read from the file)")
    p.add_argument("--no-page", action="store_true",
                   help="with --file only: do not read the release page "
                   "(recorded as such)")
    p.add_argument("--recheck", type=_period_arg, metavar="YYYY-MM-DD",
                   help="read this held stock date again although the file "
                   "is in the ledger")
    p.add_argument("--allow-older-file", action="store_true",
                   help="compare and store a stock date whose tip comes from "
                   "a newer file (logged)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_load)
    p = sub.add_parser("refresh-latest", help="copy each stock date's latest "
                       "edition into the live table, provenance and "
                       "loaded_at included (preview by default)")
    p.add_argument("--accept-drift", action="append", metavar="YYYY-MM-DD",
                   help="overwrite this stock date although its live rows "
                   "equal no stored edition (repeatable)")
    p.add_argument("--accept-key-changes", action="append",
                   metavar="YYYY-MM-DD",
                   help="apply this stock date although providers are added "
                   "or removed (repeatable; each must be named)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_refresh_latest)
    p = sub.add_parser("migrate-legacy", help="one-off: edition 1 as loaded "
                       "for the held stock date, with the proof (preview by "
                       "default)")
    p.add_argument("file", metavar="FILE")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_migrate_legacy)
    p = sub.add_parser("restore-edition", help="store an earlier edition's "
                       "rows as the next edition (preview by default)")
    p.add_argument("period", type=_period_arg, metavar="YYYY-MM-DD")
    p.add_argument("edition", type=int)
    pe.mode_parser(p)
    p.set_defaults(func=cmd_restore_edition)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
