"""S18 (ONS Price Index of Private Rents): edition history and loader on the
period-editions engine.

la_private_rents holds one row per authority, month, breakdown type and
category (nine breakdown blocks): the latest layer. This module owns
la_private_rents_editions, which keeps every fetch of a month that differed
from the one before, so a revision (including the provisional-to-final change
of the latest month) in a new PIPR edition never overwrites what was held.
The editions machinery is editions_core driven by SPEC, through
period_editions (pe) bound to PROFILE. The pure part (edition names, the
landing-page link, parsing Table 1, records, identity) comes first and moves
the retired s18_pipr_transform.py without a change in what it computes; the
loader follows. Importing this module needs no database, no network and no
key.

Periods are the live table's DATE column, handled as ISO strings
('2026-07-01'), which sort as dates (docs/RULES.md rule 8).

Blanks and zeros (docs/RULES.md rule 1): a blank or a [x]/[z] marker is NULL,
never 0; a published 0 stays 0 (but a zero or negative rent stops the load);
NULL versus a number counts as a revision. provisional is a compared value: a
month going from provisional to final is a revision and is stored.

Files: each run reads the ONS landing page, takes the first (newest) xlsx
link and downloads it to data/raw/pipr_<edition>.xlsx (git-ignored), or reads
the workbook given with --file (offline). Identity is read from the file: the
edition named by the link or file name, the publication date on the Cover
sheet and the latest month in Table 1 must agree (the latest month is the one
before the publication month); anything else halts before anything is stored.

Barnsley and Sheffield (docs/RULES.md rule 4): the workbook publishes
E08000038/39 on the whole back series (declared 'new' for source '18' in
scripts/geography.py); the canonical codes come only from geography.resolve.

Subcommands (every writing command previews by default; --commit and
--simulate are mutually exclusive; --simulate runs the --commit path and
always rolls back):
    python scripts/s18_pipr_editions.py ddl [--commit | --simulate]
    python scripts/s18_pipr_editions.py status
    python scripts/s18_pipr_editions.py sync-new [--expected-areas N]
                                                 [--commit | --simulate]
    python scripts/s18_pipr_editions.py load [--file F] [--min-period D]
                                             [--recheck-n N]
                                             [--expected-areas N]
                                             [--allow-older-file]
                                             [--commit | --simulate]
        # every held month present in the file is rechecked, every month in
        # the file later than the earliest held month and not held is new.
        # Each month is checked whole first: the expected number of areas
        # (held modal count or --expected-areas), every area of the live
        # table's latest month, every area with all nine breakdown blocks, no
        # zero or negative rent; a month that fails, or whose revisions break
        # a stop condition (NULL replacing a number in more than 5 areas, a
        # rent or index revised by more than 50% in more than 10 areas), is
        # REJECTED (not stored, exit 1). Then new (edition 1 AND its live
        # rows, one transaction), unchanged (nothing stored) or revised (next
        # edition; live changes only through refresh-latest). A file whose
        # edition or latest month is earlier than what is held halts unless
        # --allow-older-file.
    python scripts/s18_pipr_editions.py refresh-latest [--commit | --simulate]
                                                       [--accept-drift PERIOD]
"""
import argparse
import dataclasses
import hashlib
import html as html_lib
import re
import sys
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import urlsplit

sys.path.insert(0, str(Path(__file__).resolve().parent))

import editions_core as core  # noqa: E402
import geography  # noqa: E402
import load_checks  # noqa: E402
import period_editions as pe  # noqa: E402
from editions_core import halt  # noqa: E402
from statxplore_months import plan_months  # noqa: E402

ENGLISH_LA_PREFIXES = ("E06", "E07", "E08", "E09")
MIN_PERIOD = date(2024, 3, 1)

# (column suffix, breakdown_type, category): the nine blocks of Table 1.
BLOCKS = [
    ("",                   "all",           "all"),
    (" one bed",           "bedroom",       "1_bed"),
    (" two bed",           "bedroom",       "2_bed"),
    (" three bed",         "bedroom",       "3_bed"),
    (" four or more bed",  "bedroom",       "4_plus_bed"),
    (" detached",          "property_type", "detached"),
    (" semidetached",      "property_type", "semi_detached"),
    (" terraced",          "property_type", "terraced"),
    (" flat maisonette",   "property_type", "flat_maisonette"),
]
BLOCK_KEYS = frozenset((bt, cat) for _, bt, cat in BLOCKS)

# Column scales: mean_rent numeric(8,2), rent_index numeric(8,2),
# annual_pct_change numeric(6,2).
_Q = Decimal("0.01")
COLUMN_QUANTA = {"mean_rent": _Q, "rent_index": _Q, "annual_pct_change": _Q}
NUMERIC_COLUMNS = tuple(COLUMN_QUANTA)
VALUE_COLUMNS = NUMERIC_COLUMNS + ("provisional",)
RATIO_COLUMNS = ("mean_rent", "rent_index")   # revisions are ranked on these
MAX_NULL_AREAS = 5           # stop: NULL replacing a number, areas per period
MAX_BIG_REVISION_AREAS = 10  # stop: revision above 50%, areas per period

LANDING = ("https://www.ons.gov.uk/economy/inflationandpriceindices/datasets/"
           "priceindexofprivaterentsukmonthlypricestatistics")
BASE_URL = "https://www.ons.gov.uk"

_MONTHS = ("january", "february", "march", "april", "may", "june", "july",
           "august", "september", "october", "november", "december")
_SLUG_RE = re.compile(r"(?<![0-9a-z])([0-9]{1,2})(" + "|".join(_MONTHS)
                      + r")([0-9]{4})(?![0-9a-z])")
_ISO_DAY = re.compile(r"[0-9]{4}-[0-9]{2}-[0-9]{2}")


def edition_from_link(url_or_name: str) -> str:
    """The edition slug ('16september2026') named by an ONS file link
    ('.../datasets/priceindexofprivaterentsukmonthlypricestatistics/
    16september2026/<file>.xlsx') or a local name ('pipr_16september2026.
    xlsx'). Anything else raises ValueError naming the input."""
    text = url_or_name if isinstance(url_or_name, str) else ""
    m = _SLUG_RE.search(text.lower())
    if not m:
        raise ValueError(f"not a recognised PIPR edition link or file name: "
                         f"{url_or_name!r}")
    return f"{int(m.group(1))}{m.group(2)}{m.group(3)}"


def edition_date(edition: str) -> date:
    """'16september2026' -> date(2026, 9, 16). ValueError otherwise."""
    m = _SLUG_RE.fullmatch(edition) if isinstance(edition, str) else None
    if not m:
        raise ValueError(f"not a PIPR edition slug: {edition!r}")
    return date(int(m.group(3)), _MONTHS.index(m.group(2)) + 1,
                int(m.group(1)))


def release_label(edition: str) -> str:
    """'16september2026' -> 'ONS PIPR 16september2026 edition' (the text the
    old pipeline put in the live source column)."""
    edition_date(edition)
    return f"ONS PIPR {edition} edition"


_HREF_RE = re.compile(
    r"""href=["'](/file\?uri=/economy/inflationandpriceindices/datasets/"""
    r"""priceindexofprivaterentsukmonthlypricestatistics/([^/"']+)/"""
    r"""([^"']+?\.xlsx))["']""")


def find_workbook_link(html: str) -> str:
    """The absolute URL of the first (newest) xlsx link on the landing page.
    Never hard-coded: the file name carries a number that changes. ValueError
    (with an excerpt) if there is none."""
    m = _HREF_RE.search(html or "")
    if not m:
        raise ValueError("no PIPR xlsx edition link found on the landing "
                         f"page; page excerpt: {(html or '')[:300]!r}")
    return BASE_URL + html_lib.unescape(m.group(1))


def _numeric(val):
    """Cell -> float, or None for a blank or a marker ([x], [z], text)."""
    if val is None or isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        v = float(val)
        return None if v != v or v in (float("inf"), float("-inf")) else v
    if isinstance(val, str):
        s = val.strip()
        if not s or s.startswith("["):
            return None
        try:
            v = float(s)
        except ValueError:
            return None
        return None if v != v or v in (float("inf"), float("-inf")) else v
    return None


def _dec(val, quantum):
    """Cell -> Decimal at the column scale, or None (never 0 for a blank).
    Rounded half up on the published decimal (5.845 -> 5.85, -1.505 ->
    -1.51): that is what the live table holds (checked cell for cell against
    the September 2026 workbook), whatever the float arithmetic of a
    pandas round would give."""
    v = _numeric(val)
    if v is None:
        return None
    return Decimal(repr(v)).quantize(quantum, rounding=ROUND_HALF_UP)


def _iso(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and _ISO_DAY.fullmatch(value[:10]):
        return value[:10]
    raise ValueError(f"Time period is not a date: {value!r}")


def _at(r, i):
    return r[i] if i < len(r) else None


def _header_row(rows):
    """(index of the header row, header list) within the first rows."""
    for i, r in enumerate(rows):
        cells = [("" if c is None else str(c).strip()) for c in r]
        if "Time period" in cells and "Area code" in cells:
            return i, cells
    raise ValueError("Table 1 has no header row holding 'Time period' and "
                     "'Area code' in its first rows")


def parse_workbook(path) -> list:
    """Table 1 of the workbook as long rows: one dict per English area
    (E06/E07/E08/E09), month and breakdown block, with the publisher's own
    area_code (not recoded), period (ISO date), breakdown_type, category,
    mean_rent, rent_index, annual_pct_change (Decimal or None) and
    provisional (True for the file's latest month only: the workbook has no
    marker). Columns are found by header name, so a block that is missing
    raises ValueError naming it. All months are returned; the window is
    build_records'."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Table 1" not in wb.sheetnames:
            raise ValueError(f"{Path(path).name}: no sheet 'Table 1' "
                             f"(sheets {wb.sheetnames})")
        it = wb["Table 1"].iter_rows(values_only=True)
        head_rows = []
        for r in it:
            head_rows.append(r)
            if len(head_rows) >= 10 or (
                    "Time period" in [str(c).strip() for c in r if c]):
                break
        hi, header = _header_row(head_rows)
        pos = {}
        for i, h in enumerate(header):
            pos.setdefault(h, i)
        wanted = {"Time period", "Area code"}
        missing = [c for c in wanted if c not in pos]
        cols = []
        for suffix, bt, cat in BLOCKS:
            names = (f"Rental price{suffix}", f"Index{suffix}",
                     f"Annual change{suffix}")
            gone = [n for n in names if n not in pos]
            if gone:
                missing += gone
                continue
            cols.append((bt, cat) + tuple(pos[n] for n in names))
        if missing:
            raise ValueError(f"{Path(path).name}: Table 1 lacks column(s) "
                             f"{missing}; a breakdown block is missing")
        ip, ia = pos["Time period"], pos["Area code"]
        rows, latest = [], None
        for r in it:
            if r is None or len(r) <= max(ip, ia):
                continue
            code = r[ia]
            if not isinstance(code, str) or not code.startswith(
                    ENGLISH_LA_PREFIXES):
                continue
            period = _iso(r[ip])
            if latest is None or period > latest:
                latest = period
            for bt, cat, i_rent, i_idx, i_chg in cols:
                rows.append({
                    "area_code": code, "period": period,
                    "breakdown_type": bt, "category": cat,
                    "mean_rent": _dec(_at(r, i_rent), _Q),
                    "rent_index": _dec(_at(r, i_idx), _Q),
                    "annual_pct_change": _dec(_at(r, i_chg), _Q),
                    "provisional": False})
    finally:
        wb.close()
    for row in rows:
        row["provisional"] = row["period"] == latest
    return rows


_COVER_RE = re.compile(r"published\s+at\s+[^.]*?\bon\s+([0-9]{1,2})\s+"
                       r"([A-Za-z]+)\s+([0-9]{4})")


def cover_edition(path):
    """The edition slug named by the Cover sheet's publication date ('...
    originally published at 9:30am on 16 September 2026'), or None."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Cover sheet" not in wb.sheetnames:
            return None
        for r in wb["Cover sheet"].iter_rows(values_only=True):
            for v in r:
                m = _COVER_RE.search(v) if isinstance(v, str) else None
                if m and m.group(2).lower() in _MONTHS:
                    return f"{int(m.group(1))}{m.group(2).lower()}{m.group(3)}"
    finally:
        wb.close()
    return None


def build_records(rows, recodes, *, min_period=MIN_PERIOD) -> list:
    """Long rows -> records: area_code resolved through recodes (publisher
    code -> canonical code; any other code unchanged), months before
    min_period dropped, sorted by (period, lad24cd, breakdown_type,
    category). ValueError on a duplicate key (two source codes meeting, or a
    repeated row): no winner is chosen."""
    floor = (min_period if isinstance(min_period, str)
             else min_period.isoformat())
    out, seen = [], {}
    for r in rows:
        if r["period"] < floor:
            continue
        lad = recodes.get(r["area_code"], r["area_code"])
        key = (lad, r["period"], r["breakdown_type"], r["category"])
        if key in seen:
            raise ValueError(
                f"duplicate record for lad24cd {lad}, period {r['period']}, "
                f"{r['breakdown_type']}/{r['category']}: source codes "
                f"{seen[key]} and {r['area_code']}; refusing to choose a "
                "winner")
        seen[key] = r["area_code"]
        out.append({"lad24cd": lad, "period": r["period"],
                    "breakdown_type": r["breakdown_type"],
                    "category": r["category"],
                    **{c: r[c] for c in VALUE_COLUMNS}})
    out.sort(key=lambda x: (x["period"], x["lad24cd"], x["breakdown_type"],
                            x["category"]))
    return out


def group_by_period(records) -> dict:
    by = {}
    for r in records:
        by.setdefault(r["period"], []).append(r)
    return by


def check_identity(edition, latest_period_in_file, link_name,
                   cover=...) -> list:
    """Problems (empty list = fine). edition: the slug of the link or file
    name; link_name: the file or link name (it must name the same edition);
    latest_period_in_file: 'YYYY-MM-DD' (or 'YYYY-MM'): it must be the month
    before the publication month; cover: the Cover sheet's edition (from
    cover_edition; None = the cover names none, a problem; omit to skip)."""
    problems = []
    try:
        pub = edition_date(edition)
    except ValueError as e:
        return [str(e)]
    try:
        named = edition_from_link(link_name)
        if named != edition:
            problems.append(f"{link_name} names edition {named}, not "
                            f"{edition}")
    except ValueError as e:
        problems.append(str(e))
    if cover is not ...:
        if cover is None:
            problems.append("the Cover sheet names no publication date")
        elif cover != edition:
            problems.append(f"the Cover sheet names edition {cover}, the "
                            f"file name {edition}")
    want = (f"{pub.year - 1}-12" if pub.month == 1
            else f"{pub.year}-{pub.month - 1:02d}")
    got = (latest_period_in_file or "")[:7]
    if got != want:
        problems.append(f"edition {edition} should end in {want} (the month "
                        f"before its publication) but the latest month in "
                        f"Table 1 is {got or 'none'}")
    return problems


def content_sha256(path) -> str:
    """SHA-256 of the file's bytes."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _cell(r, col):
    v = r[col]
    if v is None:
        return "NULL"
    if col == "provisional":
        if not isinstance(v, bool):
            raise ValueError(f"provisional must be a bool, got {v!r}")
        return "t" if v else "f"
    if isinstance(v, bool) or not isinstance(v, Decimal):
        raise ValueError(f"{col} must be a Decimal or None, got {v!r} for "
                         f"{r['lad24cd']}")
    return format(v.quantize(COLUMN_QUANTA[col], rounding=ROUND_HALF_UP), "f")


def records_sha256(records) -> str:
    """SHA-256 of one period's records sorted by key; NULL is distinct from
    0 (the edition's source_sha256)."""
    rows = sorted(records, key=lambda r: (r["lad24cd"], r["breakdown_type"],
                                          r["category"]))
    lines = ["|".join([r["lad24cd"], r["period"], r["breakdown_type"],
                       r["category"]] + [_cell(r, c) for c in VALUE_COLUMNS])
             for r in rows]
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Spec and profile
# ---------------------------------------------------------------------------

LIVE = "la_private_rents"
TABLE = "la_private_rents_editions"

# Live table (checked 2026-10-09 in information_schema): lad24cd varchar(9),
# period date, breakdown_type varchar(20), category varchar(30) (all NOT
# NULL, primary key together), mean_rent numeric(8,2), rent_index
# numeric(8,2), annual_pct_change numeric(6,2), provisional boolean default
# false, source text, loaded_at timestamptz default now(). The editions table
# matches it on the key and values, with the core's lad24cd -> la_boundaries
# foreign key. Live `source` is set on the live rows a new month inserts
# (the edition's name); refresh-latest does not touch it.
SPEC = core.EditionSpec(
    name="s18",
    live_table=LIVE,
    editions_table=TABLE,
    key_cols=("lad24cd", "breakdown_type", "category"),
    period_col="period",
    value_cols=(("mean_rent", "numeric(8,2)"),
                ("rent_index", "numeric(8,2)"),
                ("annual_pct_change", "numeric(6,2)"),
                ("provisional", "boolean")),
    refresh_cols=VALUE_COLUMNS,
    key_types=(("period", "date NOT NULL"),
               ("breakdown_type", "varchar(20) NOT NULL"),
               ("category", "varchar(30) NOT NULL")),
    # a period's live rows have one load date today, but a repaired period
    # may carry several: take the latest, as S15 does
    as_loaded_date="latest",
)

RUN_AGENT = "Source 18 - ONS Private Rents"
RUN_SOURCE = "18"
SOURCE_FILE = "ONS PIPR workbook (Table 1)"
LIVE_MISSING = pe.LIVE_MISSING

REPO = Path(__file__).resolve().parent.parent
RAW_DIR = REPO / "data" / "raw"          # git-ignored
USER_AGENT = ("ucws-pipeline S18 loader (read-only download of the public "
              "ONS PIPR workbook)")
MIN_FILE_BYTES = 10 * 1024 * 1024


def _no_label(fetched_on) -> str:
    raise ValueError("an S18 release label comes from the file edition; "
                     "store through run_profile(...)")


def _late(name: str):
    """A call through this module's attribute `name`, looked up when called
    (so mock.patch.object(s18_pipr_editions, name, ...) reaches the engine)."""
    return lambda *a, **kw: globals()[name](*a, **kw)


PROFILE = pe.Profile(
    spec=SPEC, value_cols=VALUE_COLUMNS, run_agent=RUN_AGENT,
    run_source=RUN_SOURCE, heading="S18 ONS private rents",
    default_source_file=SOURCE_FILE,
    expected_areas=None,   # derived at run time from the held modal count
    release_label=_no_label,
    content_sha256=_late("records_sha256"),
    savepoint="s18_period",
    example_label="(rent, index, annual change, provisional)")


def _profile(spec) -> pe.Profile:
    return PROFILE.with_spec(spec)


def check_period(records: list, period: str, expected_areas,
                 required_areas) -> None:
    """ValueError unless the period is whole: an ISO date, every record of
    it, no key repeated, every area holding all nine breakdown blocks, no
    zero or negative rent, one provisional flag, exactly expected_areas
    areas (None: not checked) and every area of required_areas (the live
    table's latest period)."""
    if not isinstance(period, str) or not _ISO_DAY.fullmatch(period):
        raise ValueError(f"not an ISO date period: {period!r}")
    seen, blocks, flags = set(), {}, set()
    for r in records:
        if r.get("period") != period:
            raise ValueError(f"{period}: a record of period {r.get('period')!r}")
        k = (r["lad24cd"], r["breakdown_type"], r["category"])
        if k in seen:
            raise ValueError(f"{period}: {k} appears twice")
        seen.add(k)
        blocks.setdefault(r["lad24cd"], set()).add(k[1:])
        flags.add(r["provisional"])
        rent = r["mean_rent"]
        if rent is not None and rent <= 0:
            raise ValueError(f"{period}: {r['lad24cd']} {k[1]}/{k[2]} has a "
                             f"zero or negative rent ({rent})")
    if len(flags) > 1:
        raise ValueError(f"{period}: mixed provisional flags")
    short = {a: sorted(BLOCK_KEYS - b) for a, b in blocks.items()
             if b != BLOCK_KEYS}
    if short:
        a = sorted(short)[0]
        raise ValueError(f"{period}: {len(short)} area(s) lack breakdown "
                         f"block(s), e.g. {a} lacks {short[a]}; a period "
                         "missing a block is never stored")
    n = len(blocks)
    if expected_areas is not None and n != expected_areas:
        why = ("a short period is never stored" if n < expected_areas else
               "more areas than held; review them, then give "
               "--expected-areas N")
        raise ValueError(f"{period}: {n} areas, expected {expected_areas}; "
                         f"{why}")
    missing = sorted(set(required_areas) - set(blocks))
    if missing:
        raise ValueError(f"{period}: {len(missing)} area(s) of the live "
                         f"table's latest period missing {missing[:10]}; a "
                         "period missing an area is never stored")


def period_problems(by_period: dict, periods, expected_areas,
                    required_areas) -> dict:
    """{period: message} for each period that fails check_period."""
    out = {}
    for p in periods:
        try:
            check_period(by_period[p], p, expected_areas, required_areas)
        except ValueError as e:
            out[p] = str(e)
    return out


def run_profile(spec, edition: str, expected_areas,
                required_areas=()) -> pe.Profile:
    """PROFILE for one run: the spec, the expected area count, the release
    label of the file edition and the per-period check."""
    label = release_label(edition)
    req = frozenset(required_areas)
    return dataclasses.replace(
        _profile(spec), expected_areas=expected_areas,
        release_label=lambda d: label,
        check_records=lambda r, p: check_period(r, p, expected_areas, req))


def insert_live(cur, profile, period: str, records: list) -> None:
    """The period's live rows with `source` set to the edition's name (the
    engine's insert_live writes the value columns only; loaded_at takes its
    default)."""
    from psycopg2.extras import execute_values
    spec = profile.spec
    label = profile.release_label(date.today())
    cols = (tuple(spec.key_cols) + (spec.period_col,)
            + tuple(profile.value_cols) + ("source",))
    execute_values(
        cur, f"INSERT INTO public.{spec.live_table} ({', '.join(cols)}) "
        "VALUES %s",
        [tuple(r[k] for k in spec.key_cols) + (period,)
         + tuple(r[c] for c in profile.value_cols) + (label,)
         for r in records], page_size=1000)


def apply_period(cur, profile, period, records, *, fetched_on,
                 source_file=None):
    """pe.apply_period with the live rows inserted through insert_live."""
    return pe.apply_period(cur, profile, period, records,
                           fetched_on=fetched_on, source_file=source_file,
                           insert=_late("insert_live"))


# ---------------------------------------------------------------------------
# Files: download and read
# ---------------------------------------------------------------------------

def source_file_text(path) -> str:
    """An edition's source_file: the file name and the first 16 hex
    characters of its sha256."""
    return f"{Path(path).name} (sha256 {content_sha256(path)[:16]})"


def fetch_latest_file(session=None, dest=None) -> tuple:
    """(path, edition slug, link): the landing page -> its first xlsx link ->
    the workbook, downloaded (read-only GETs with a User-Agent) to dest
    (default data/raw/, git-ignored) as pipr_<edition>.xlsx. The download is
    checked (more than MIN_FILE_BYTES, openable as a workbook with a Table 1)
    before it is kept. Any parse failure halts printing an excerpt; nothing
    is guessed. session: anything with requests' get(url, headers=, timeout=,
    stream=) (default requests)."""
    if session is None:
        import requests
        session = requests
    dest = Path(dest) if dest is not None else RAW_DIR
    headers = {"User-Agent": USER_AGENT}

    def get(url, **kw):
        try:
            r = session.get(url, headers=headers, timeout=300, **kw)
            r.raise_for_status()
        except Exception as e:  # noqa: BLE001 (any failure halts)
            halt(f"GET {url} failed, nothing downloaded: {e}")
        return r

    page = get(LANDING).text
    try:
        link = find_workbook_link(page)
        edition = edition_from_link(link)
    except ValueError as e:
        halt(f"the landing page {LANDING} is not laid out as expected, "
             f"nothing downloaded: {e}")
    body = get(link).content
    if len(body) <= MIN_FILE_BYTES:
        halt(f"{link}: {len(body)} bytes, expected more than "
             f"{MIN_FILE_BYTES}; nothing written; saw {body[:100]!r}")
    dest.mkdir(parents=True, exist_ok=True)
    target = dest / f"pipr_{edition}.xlsx"
    new_sha = hashlib.sha256(body).hexdigest()
    tmp = dest / f"pipr_{edition}.download.tmp.xlsx"
    tmp.write_bytes(body)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(tmp, read_only=True)
        ok = "Table 1" in wb.sheetnames
        wb.close()
        if not ok:
            raise ValueError("no sheet 'Table 1'")
    except Exception as e:  # noqa: BLE001
        tmp.unlink()
        halt(f"{link}: not a usable workbook ({e}); nothing kept")
    path = target
    if target.exists():
        old_sha = content_sha256(target)
        if old_sha == new_sha:
            tmp.unlink()
            print(f"{target.name} already held with the same content "
                  f"(sha256 {new_sha[:16]}); kept as it is")
        else:
            # ONS re-issued a workbook under the same edition name. The file
            # already on disk may be cited as evidence, so it is never
            # replaced: the new one is saved beside it and is the one read.
            path = dest / f"pipr_{edition}-{new_sha[:8]}.xlsx"
            tmp.replace(path)
            print(f"NOTE: {target.name} already exists with different "
                  f"content (sha256 {old_sha[:16]}, new {new_sha[:16]}); "
                  f"kept it untouched and saved the new download as "
                  f"{path.name}")
    else:
        tmp.replace(target)
    print(f"downloaded {path.name}: {len(body):,} bytes, sha256 "
          f"{hashlib.sha256(body).hexdigest()[:16]} -> {path}")
    print(f"link: {link} (edition {edition})")
    return path, edition, link


def reference_codes(cur) -> tuple:
    """(valid lad24cd set from la_boundaries, {old_code: new_code} from
    la_code_lookup). Only a lookup with exactly one target in la_boundaries
    is used. Barnsley and Sheffield go through geography.resolve."""
    cur.execute("SELECT lad24cd FROM public.la_boundaries")
    valid = {r[0] for r in cur.fetchall()}
    cur.execute("SELECT old_code, new_code FROM public.la_code_lookup")
    targets = {}
    for old, new in cur.fetchall():
        targets.setdefault(old, set()).add(new)
    lookup = {}
    for old, ts in targets.items():
        ok = ts & valid
        if len(ok) == 1:
            lookup[old] = next(iter(ok))
    return valid, lookup


def geography_codes(rows, min_period=MIN_PERIOD) -> dict:
    """{ISO period: set of area codes} of the in-window rows: what
    geography.resolve checks against S18's declared form."""
    floor = (min_period if isinstance(min_period, str)
             else min_period.isoformat())
    out = {}
    for r in rows:
        if r["period"] >= floor:
            out.setdefault(r["period"], set()).add(r["area_code"])
    return out


def load_file_records(path, conn, *, min_period=MIN_PERIOD,
                      link_name=None) -> tuple:
    """(records_by_period, edition) from the workbook: identity checked
    (check_identity: the name, the Cover sheet and the latest month agree),
    a hard stop if Barnsley's or Sheffield's codes disagree with S18's form
    in geography.DATASET_FORM (the recode map comes from geography.resolve),
    records built (build_records), and a hard stop on any unresolved code
    (reported UNEXPLAINED through load_checks.check_codes). Writes nothing."""
    p = Path(path)
    name = link_name or p.name
    try:
        edition = edition_from_link(name)
    except ValueError as e:
        halt(f"file identity: {e}; nothing stored")
    try:
        rows = parse_workbook(p)
    except ValueError as e:
        halt(f"{p.name}: {e}; nothing stored")
    if not rows:
        halt(f"{p.name}: no English area rows in Table 1; nothing stored")
    latest = max(r["period"] for r in rows)
    problems = check_identity(edition, latest, name, cover_edition(p))
    if problems:
        halt("file identity check failed, nothing stored: "
             + "; ".join(problems))
    with conn.cursor() as cur:
        valid, lookup = reference_codes(cur)
        seen = geography_codes(rows, min_period)
        try:
            recode_map, problems = geography.resolve(
                cur, RUN_SOURCE, set().union(*seen.values()) if seen else (),
                by_period=seen)
        except ValueError as e:
            halt(f"Barnsley/Sheffield geography: {e}; nothing stored")
        if problems:
            halt("Barnsley/Sheffield geography check failed: "
                 + "; ".join(problems))
        note = geography.confirm_note(RUN_SOURCE)
        if note:
            print(note)
        recodes, unresolved = {}, {}
        for c in sorted(set().union(*seen.values()) if seen else ()):
            if c in recode_map:
                recodes[c] = recode_map[c]
            elif c in valid:
                continue
            elif c in lookup:
                recodes[c] = lookup[c]
            else:
                unresolved[c] = sum(1 for r in rows
                                    if r["area_code"] == c
                                    and r["period"] >= min_period.isoformat())
        if unresolved:
            codes = sorted(unresolved)
            found = load_checks.check_codes(cur, codes)
            found += [f"UNEXPLAINED {c}" for c in codes
                      if f"UNEXPLAINED {c}" not in found]
            halt(f"unresolved PIPR area codes {found} (rows: "
                 f"{dict(sorted(unresolved.items()))}); explain them in "
                 "la_code_lookup before loading; nothing stored")
        try:
            records = build_records(rows, {k: v for k, v in recodes.items()
                                           if k != v}, min_period=min_period)
        except ValueError as e:
            halt(f"{p.name}: {e}; nothing stored")
    by_period = group_by_period(records)
    print(f"file: {p.name} ({len(rows):,} English area-month-block rows); "
          f"edition {edition}; latest month {latest}")
    return by_period, edition


# ---------------------------------------------------------------------------
# Window, expected areas, cell summary
# ---------------------------------------------------------------------------

def plan_window(held: list, available: list, *, widen: bool,
                recheck_n=None) -> tuple:
    """(new, recheck, earlier), each ascending. Default: recheck every held
    period present in the file (the latest recheck_n if given), new every
    period in the file not held and later than the earliest held, earlier
    the file's periods before the earliest held (never loaded). widen
    (--min-period given): every period in the file is new or rechecked."""
    if widen:
        new = [p for p in sorted(available) if p not in held]
        recheck = [p for p in sorted(available) if p in held]
        if recheck_n is not None:
            recheck = recheck[-recheck_n:] if recheck_n > 0 else []
        return new, recheck, []
    new, recheck = plan_months(held, available,
                               recheck_n if recheck_n is not None else 0,  # not a source value
                               recheck_n is None)
    earliest = min(held) if held else None
    earlier = [p for p in sorted(available)
               if earliest is not None and p < earliest]
    return new, recheck, earlier


def held_area_count(cur, spec=SPEC):
    """The most common per-period area count of the live table (None if it
    is empty)."""
    return core.modal_count(core.authority_counts(cur, spec))


def latest_live_areas(cur, spec=SPEC) -> set:
    """The areas of the live table's latest period (empty if none)."""
    cur.execute(f"SELECT DISTINCT lad24cd FROM public.{spec.live_table} "
                f"WHERE {spec.period_col} = (SELECT MAX({spec.period_col}) "
                f"FROM public.{spec.live_table})")
    return {r[0] for r in cur.fetchall()}


def cell_summary(cur, profile, by_period: dict, periods, against: str) -> dict:
    """Cells that differ between each period's records and what it is
    compared with (NULL equals NULL, NULL differs from 0): per column, NULL
    replacing a number and the reverse, rent or index revisions above 50%,
    the five largest relative revisions; provisional flips; per period the
    areas with a NULL replacing a number and with a revision above 50% (the
    stop conditions). Read-only."""
    cols = profile.value_cols
    out = {"periods_changed": 0, "cells": 0, "by_col": {c: 0 for c in cols},
           "null_for_number": 0, "number_for_null": 0, "over_50pct": 0,
           "areas_one_side": 0, "new_periods": 0, "new_cells": 0,
           "nonpositive": 0, "largest": [], "per_period": {}}
    big = []
    for p in periods:
        recs = by_period[p]
        out["nonpositive"] += sum(1 for r in recs
                                  if r["mean_rent"] is not None
                                  and r["mean_rent"] <= 0)
        old, _ = pe.stored(cur, profile, p, against)
        if not old:
            out["new_periods"] += 1
            out["new_cells"] += sum(1 for r in recs for c in cols
                                    if r[c] is not None)
            continue
        new = {pe._key(profile, r): tuple(r[c] for c in cols) for r in recs}
        out["areas_one_side"] += len(set(old) ^ set(new))
        changed, null_areas, big_areas = 0, set(), set()
        for k in sorted(set(old) & set(new)):
            for c, o, n in zip(cols, old[k], new[k]):
                if (o is None) == (n is None) and (o is None or o == n):
                    continue
                changed += 1
                out["by_col"][c] += 1
                if n is None:
                    out["null_for_number"] += 1
                    null_areas.add(k[0])
                elif o is None:
                    out["number_for_null"] += 1
                elif c in RATIO_COLUMNS and o != 0:
                    rel = abs(Decimal(n) - Decimal(o)) / abs(Decimal(o))
                    if rel > Decimal("0.5"):
                        out["over_50pct"] += 1
                        big_areas.add(k[0])
                    big.append((rel, p, k, c, o, n))
        out["per_period"][p] = {"null_areas": len(null_areas),
                                "big_areas": len(big_areas)}
        if changed:
            out["periods_changed"] += 1
            out["cells"] += changed
    big.sort(key=lambda x: (x[0], x[1], x[2], x[3]), reverse=True)
    out["largest"] = big[:5]
    return out


def stop_problems(summary: dict) -> dict:
    """{period: message} for periods whose revisions break a stop condition:
    NULL replacing a number in more than MAX_NULL_AREAS areas, a rent or
    index revised above 50% in more than MAX_BIG_REVISION_AREAS areas."""
    out = {}
    for p, d in summary["per_period"].items():
        msgs = []
        if d["null_areas"] > MAX_NULL_AREAS:
            msgs.append(f"NULL replaces a number in {d['null_areas']} areas "
                        f"(limit {MAX_NULL_AREAS})")
        if d["big_areas"] > MAX_BIG_REVISION_AREAS:
            msgs.append(f"rent or index revised above 50% in "
                        f"{d['big_areas']} areas (limit "
                        f"{MAX_BIG_REVISION_AREAS})")
        if msgs:
            out[p] = "STOP CONDITION: " + "; ".join(msgs)
    return out


def format_cell_summary(s: dict, against: str) -> str:
    what = "the LIVE table" if against == "live" else "the latest editions"
    cols = ", ".join(f"{c} {n}" for c, n in s["by_col"].items() if n) or "none"
    lines = [f"cells (compared with {what}): {s['cells']} changed in "
             f"{s['periods_changed']} period(s) (by column: {cols}); NULL "
             f"replacing a number {s['null_for_number']}; a number replacing "
             f"NULL {s['number_for_null']}; rent/index revisions above 50% "
             f"{s['over_50pct']}; keys on one side only "
             f"{s['areas_one_side']}; zero or negative rents in the file "
             f"{s['nonpositive']}; new periods {s['new_periods']} "
             f"({s['new_cells']} non-NULL cells)"]
    if s["largest"]:
        lines.append("largest rent/index revisions (relative):")
        for rel, p, k, c, o, n in s["largest"]:
            lines.append(f"  {p} {'/'.join(k)} {c}: {o} -> {n} "
                         f"({'+' if n >= o else '-'}{rel * 100:.2f}%)")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Status, run log, commands
# ---------------------------------------------------------------------------

table_exists = pe.table_exists


def _conn(writing):
    return pe.connect(writing)


def status(cur, spec=SPEC) -> dict:
    """period_editions.status: editions_core.status plus 'live_missing'."""
    return pe.status(cur, _profile(spec))


def format_status(st: dict) -> str:
    return pe.format_status(PROFILE, st)


def log_run(cur, rows_written: int, notes: str, started_at=None) -> None:
    """The pipeline_run_log row for a committed run (agent RUN_AGENT, source
    '18', status 'success'). Called only on `sync-new --commit` and
    `load --commit`."""
    pe.log_run(cur, PROFILE, rows_written, notes, started_at)


def load_run_notes(stats: dict, edition: str, files: str, held: list,
                   available: list) -> str:
    """Words for the run log: what was checked, what was stored."""
    kinds = stats["kinds"]
    done = stats["periods"]
    by = {k: [p for p in done if kinds[p] == k]
          for k in ("new", "revised", "unchanged", LIVE_MISSING)}
    stored = by["new"] + by["revised"]
    head = ("Stored: " + "; ".join(f"{k} {', '.join(by[k])}"
                                   for k in ("new", "revised") if by[k])
            if stored else "Nothing new: no new month and no revision")
    if by[LIVE_MISSING]:
        head += (". Live rows inserted (no new edition) for editions months "
                 f"missing from live: {', '.join(by[LIVE_MISSING])}")
    span = f"{done[0]} .. {done[-1]}" if done else "none"
    live_rows = stats.get("live_rows", 0)  # not a source value
    return (f"{head}. Checked {len(done)} month(s) ({span}); unchanged "
            f"{len(by['unchanged'])}. {release_label(edition)}; file "
            f"{files}. Held latest {max(held) if held else '-'}; file latest "
            f"{max(available) if available else '-'}. Rows stored: "
            f"{stats['stored_rows']} edition rows, {live_rows} live rows "
            "inserted. A run that finds nothing new is logged because the "
            "check is the run.")


def cmd_ddl(args) -> int:
    return pe.run_ddl(_profile(SPEC), args, connect=_late("_conn"),
                      table_exists=_late("table_exists"))


def cmd_status(args) -> int:
    return pe.run_status(
        _profile(SPEC), args, connect=_late("_conn"),
        table_exists=_late("table_exists"),
        status=lambda cur: status(cur, SPEC),
        format_status=_late("format_status"))


def cmd_sync_new(args) -> int:
    return pe.run_sync_new(
        _profile(SPEC), args, connect=_late("_conn"),
        table_exists=_late("table_exists"),
        status=lambda cur: status(cur, SPEC),
        format_status=_late("format_status"))


def cmd_refresh_latest(args) -> int:
    return pe.run_refresh_latest(_profile(SPEC), args, connect=_late("_conn"),
                                 table_exists=_late("table_exists"))


def _derive_expected(cur, spec, override):
    """The area count a period must hold: --expected-areas N, else the held
    modal count; halts if neither exists."""
    derived = held_area_count(cur, spec)
    if override is not None and derived is not None and override != derived:
        print(f"NOTE: --expected-areas {override} overrides the held modal "
              f"count {derived}")
    expected = override if override is not None else derived
    if expected is None:
        halt("no live period to derive the area count from; give "
             "--expected-areas N")
    return expected


_LABEL_RE = re.compile(r"ONS PIPR ([0-9a-z]+) edition")


def _newest_edition(labels):
    found = []
    for label in labels:
        m_ = _LABEL_RE.fullmatch(label or "")
        if m_:
            try:
                found.append((edition_date(m_.group(1)), m_.group(1)))
            except ValueError:
                pass
    return max(found)[1] if found else None


def latest_held_edition(cur, spec=SPEC, has_editions=True):
    """The newest edition slug named by a stored release_label or by the live
    source column ('as loaded' labels name none); None if none is named."""
    labels = []
    if has_editions:
        cur.execute(f"SELECT DISTINCT release_label FROM "
                    f"public.{spec.editions_table}")
        labels += [r[0] for r in cur.fetchall()]
    cur.execute(f"SELECT DISTINCT source FROM public.{spec.live_table}")
    labels += [r[0] for r in cur.fetchall()]
    return _newest_edition(labels)


def older_file_problem(edition, available, held, held_edition):
    """A message if the file is older than what is held, else None. The file
    is older when its edition is before the newest held edition, or its
    latest month is before the latest held month."""
    if held_edition and edition_date(edition) < edition_date(held_edition):
        return (f"the file is the {release_label(edition)} but the "
                f"{release_label(held_edition)} is already held")
    if held and available and available[-1][:7] < held[-1][:7]:
        return (f"the file's latest month is {available[-1][:7]} but "
                f"{held[-1][:7]} is already held")
    return None


def cmd_load(args) -> int:
    spec = SPEC
    base = _profile(spec)
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            has_ed = table_exists(cur, spec.editions_table)
            if writing:
                if not has_ed:
                    halt(pe.no_table(base))
                _, new, errors = core.latest_map(cur, spec)
                if errors:
                    halt("invalid edition chain: " + "; ".join(
                        f"{pe._p(p)}: {msg}" for p, msg in errors.items()))
                if new:
                    halt(f"live months with no editions "
                         f"{[pe._p(x) for x in new]}; run sync-new first so "
                         "the history starts from what is held")
            against = "editions" if has_ed else "live"
            if not has_ed:
                print(f"NOTE: {spec.editions_table} does not exist yet; this "
                      "preview compares each period with the LIVE table "
                      f"{spec.live_table}")
            held = pe.held_periods(cur, base, has_ed)
            print("held periods: " + (f"{held[0]} .. {held[-1]} ({len(held)})"
                                      if held else "none"))
            if not held and args.min_period is None:
                halt("nothing is held, so no window can be planned; give "
                     "--min-period YYYY-MM-DD (and --expected-areas N)")
            expected = _derive_expected(cur, spec, args.expected_areas)
            required = latest_live_areas(cur, spec)
            min_period = args.min_period or MIN_PERIOD
            if args.file:
                path, fetched, link = Path(args.file), None, None
                print("file given: nothing downloaded")
            else:
                path, fetched, link = fetch_latest_file()
            by_period, edition = load_file_records(
                path, conn, min_period=min_period, link_name=None)
            if fetched is not None and fetched != edition:
                halt(f"file identity: the landing page named edition "
                     f"{fetched}, the file is {edition}")
            available = sorted(by_period)
            if not available:
                halt(f"no month from {min_period} in the file; nothing "
                     "stored")
            held_ed = latest_held_edition(cur, spec, has_ed)
            older = older_file_problem(edition, available, held, held_ed)
            if older:
                if args.allow_older_file:
                    print(f"NOTE: --allow-older-file given; {older}")
                else:
                    halt(older + "; storing it would record an older "
                         "edition as a revision. If this is deliberate, "
                         "re-run with --allow-older-file")
            print(f"{release_label(edition)}: file periods {available[0]} .. "
                  f"{available[-1]} ({len(available)} from {min_period})")
            new_p, recheck, earlier = plan_window(
                held, available, widen=args.min_period is not None,
                recheck_n=args.recheck_n)
            print(f"planned: new {', '.join(new_p) or 'none'}; recheck "
                  f"{len(recheck)} held period(s)"
                  + (f" ({recheck[0]} .. {recheck[-1]})" if recheck else ""))
            if earlier:
                print(f"ignored, earlier than the earliest held period "
                      f"{held[0]} (give --min-period to widen): "
                      f"{', '.join(earlier)}")
            periods = sorted(set(new_p) | set(recheck))
            print(f"expected areas per period: {expected}"
                  + (" (--expected-areas)" if args.expected_areas is not None
                     else " (held modal count)")
                  + "; every area of the live table's latest period "
                  f"({len(required)}) required")
            profile = run_profile(spec, edition, expected, required)
            problems = period_problems(by_period, periods, expected, required)
            for p, msg in problems.items():
                print(f"  {p}: REJECTED, not stored: {msg}")
            ok = [p for p in periods if p not in problems]
            summary = cell_summary(cur, profile, by_period, ok, against)
            print(format_cell_summary(summary, against))
            stops = stop_problems(summary)
            for p, msg in stops.items():
                print(f"  {p}: REJECTED, not stored: {msg}")
            problems.update(stops)
            ok = [p for p in ok if p not in stops]
            stats = {"periods": [], "kinds": {}, "stored_rows": 0,
                     "live_rows": 0}
            started = datetime.now(timezone.utc)
            rc = 0
            if ok:
                rc = pe.load_periods(
                    cur, profile, ok, lambda p: by_period[p], date.today(),
                    args.commit, simulate=args.simulate, against=against,
                    stats=stats, source_file=source_file_text(path),
                    apply=_late("apply_period"))
            else:
                print("nothing to compare")
            if problems:
                print(f"REJECTED {len(problems)} period(s), not stored: "
                      f"{', '.join(sorted(problems))}; exit 1")
                rc = 1
            if args.commit:
                if rc == 0:
                    log_run(cur, stats["stored_rows"] + stats["live_rows"],  # not a source value
                            load_run_notes(stats, edition,
                                           source_file_text(path), held,
                                           available), started)
                    conn.commit()
                    print("pipeline_run_log row written")
                else:
                    print("pipeline_run_log: no row written (the run failed "
                          "or rejected a period; periods committed before "
                          "are in the editions table)")
        conn.rollback()  # commit mode has committed period by period already
        if not writing:
            print("PREVIEW: nothing written (use --commit or --simulate)")
        elif args.simulate:
            print("SIMULATION: ROLLED BACK (nothing persisted)")
        return rc
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def _iso_day(text: str) -> date:
    try:
        if not _ISO_DAY.fullmatch(text):
            raise ValueError(text)
        return date.fromisoformat(text)
    except ValueError:
        raise argparse.ArgumentTypeError(f"not a YYYY-MM-DD date: {text!r}")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="S18 ONS private rents: "
                                 "editions loader")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ddl", help="create the editions table (preview by "
                       "default)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_ddl)
    sub.add_parser("status", help="what needs action; exit 1 if anything"
                   ).set_defaults(func=cmd_status)
    p = sub.add_parser("sync-new", help="edition 1 'as loaded' for live months "
                       "with no editions (preview by default)")
    p.add_argument("--expected-areas", "--expected-authorities", type=int,
                   dest="expected_authorities", metavar="N",
                   help="areas per month; only when nothing can be derived "
                   "(no month has editions yet)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_sync_new)
    p = sub.add_parser("load", help="read the latest workbook, compare and "
                       "store months (preview by default)")
    p.add_argument("--file", metavar="FILE",
                   help="a local pipr_<edition>.xlsx (nothing is downloaded)")
    p.add_argument("--min-period", type=_iso_day, metavar="YYYY-MM-DD",
                   help="read the file from this month and load every month "
                   "in it (widens or narrows the window; required when "
                   "nothing is held)")
    p.add_argument("--recheck-n", type=int, metavar="N",
                   help="recheck only the latest N held months (default all)")
    p.add_argument("--allow-older-file", action="store_true",
                   help="load a file whose edition or latest month is "
                   "earlier than what is held (default: halt)")
    p.add_argument("--expected-areas", type=int, metavar="N",
                   help="areas per month (default the held modal count)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_load)
    p = sub.add_parser("refresh-latest", help="copy each month's latest "
                       "edition into the live table (preview by default)")
    pe.mode_parser(p)
    p.add_argument("--accept-drift", action="append", metavar="PERIOD",
                   help="overwrite this month (YYYY-MM-DD) although its live "
                   "rows equal no stored edition (repeatable)")
    p.set_defaults(func=cmd_refresh_latest)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
