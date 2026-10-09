"""S11 (CQC Care directory with filters): dated snapshots of the CQC register
on the period-editions engine.

Each CQC "Care directory with filters" file (HSCA_Active_Locations .ods or
.xlsx) is a point-in-time copy of the register "as at" a date: one row per
registered location. Every file becomes a snapshot that is kept for good:

    period = the file's own "as at" date (snapshot_date, DATE; read from the
             file's README sheet, never from the name or the link)
    key    = location_id (Adult social care directorate rows only)

cqc_location_snapshots is the live layer (one row per location per snapshot),
cqc_location_snapshots_editions the append-only editions, and
cqc_location_snapshots_editions_file_checks the ledger of every file read.
cqc_location_snapshots_unresolved (append-only) records, per snapshot and
edition, the in-scope locations that could not be given a lad24cd; it is
written in the same savepoint as the snapshot. The machinery is editions_core
driven by SPEC, through period_editions (pe) bound to PROFILE. A published CQC
file is not revised; a later edition of a held snapshot only arises if CQC
reissued a file with the same "as at" date, or our parse or mapping of a held
file changed (load --recheck).

History (what the retired scripts did, so the wording here stays truthful):
s11_cqc_load.py upserted with INSERT ... ON CONFLICT (location_id) DO UPDATE,
overwriting each location's previous snapshot with the next file's values, so
the table never held more than one snapshot of a location; s11_cqc_map.py
wrote cqc_unresolved_locations on every run with no preview. migrate-legacy
(one-off) rebuilds the July and August 2026 snapshots from the files that were
loaded, checked against the legacy rows that survive of them, records
September 2026 as loaded, renames cqc_locations to cqc_locations_legacy (kept)
and makes cqc_locations a view over the snapshots with the same columns.

Blanks and zeros (docs/RULES.md rule 1): a blank cell is NULL, never 0; a
published 0 (care home beds on a non-care-home row) stays 0. The 14 service
and band flags are 'Y' or blank (the publisher's yes/no: README "filter for
'Y'"); 'Care home?' and 'Dormant (Y/N)' are 'Y'/'N'; any other marker, an
unparseable number, coordinate or date halts (the old errors="coerce" is gone:
no silent NULL). 0 to NULL or NULL to 0 against the held edition of a
snapshot needs --acknowledge PERIOD.

Identity (rule 3): the README sheet's title must be the HSCA title and its
"Source: CQC database as at DD Month YYYY" date must equal the date in the
file name.

Barnsley and Sheffield (rule 4): source 11 is declared 'none' in
scripts/geography.py (the file carries no GSS code for them); a file holding
any of E08000016/19/38/39 in any cell halts. lad24cd comes from
point-in-polygon against la_boundaries; the postcodes.io fallback codes go
through geography.canonical, then la_boundaries, then la_code_lookup, and an
unknown code halts (UNEXPLAINED).

Subcommands (every writing command previews by default; --commit and
--simulate are mutually exclusive; --simulate runs the --commit path and
always rolls back). There is no sync-new: the legacy table holds no
per-snapshot live rows to record "as loaded"; migrate-legacy does that job.
    python scripts/s11_cqc_editions.py ddl [--commit | --simulate]
    python scripts/s11_cqc_editions.py status
    python scripts/s11_cqc_editions.py load [--file PATH] [--recheck PERIOD]
                                            [--allow-older-file]
                                            [--acknowledge PERIOD]
                                            [--commit | --simulate]
        # the page's filters file (downloaded to data/raw/, also in a
        # preview; postcodes.io is read for rows without coordinates; a
        # preview writes nothing to the database), or --file PATH. A
        # snapshot breaking a stop condition is REJECTED (nothing stored,
        # exit 1).
    python scripts/s11_cqc_editions.py refresh-latest [--commit | --simulate]
                                                      [--accept-drift PERIOD]
    python scripts/s11_cqc_editions.py migrate-legacy FILE FILE FILE
                                                      [--commit | --simulate]
"""
import argparse
import dataclasses
import hashlib
import html as html_lib
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import unquote, urljoin, urlsplit

sys.path.insert(0, str(Path(__file__).resolve().parent))

import editions_core as core  # noqa: E402
import geography  # noqa: E402
import period_editions as pe  # noqa: E402
from editions_core import halt  # noqa: E402

LIVE = "cqc_location_snapshots"
TABLE = "cqc_location_snapshots_editions"
LEGACY = "cqc_locations"
LEGACY_RENAMED = "cqc_locations_legacy"
LEGACY_UNRESOLVED = "cqc_unresolved_locations"
RUN_AGENT = "Source 11 - CQC Care directory"
LEGACY_RUN_AGENT = "Source 11 - CQC Care Providers"
RUN_SOURCE = "11"
SOURCE_FILE = "CQC Care directory with filters (HSCA_Active_Locations)"
LIVE_MISSING = pe.LIVE_MISSING

PAGE = "https://www.cqc.org.uk/about-us/transparency/using-cqc-data"
REPO = Path(__file__).resolve().parent.parent
RAW_DIR = REPO / "data" / "raw"          # git-ignored
USER_AGENT = ("ucws-pipeline S11 loader (read-only download of the public "
              "CQC Care directory with filters)")
POSTCODES_IO = "https://api.postcodes.io"
MIN_FILE_BYTES = 5 * 1024 * 1024

TITLE = ("Active locations for providers registered under the Health and "
         "Social Care Act (HSCA)")
README_SHEET = "README"
DATA_SHEET = "HSCA_Active_Locations"
SCOPE = "Adult social care"
FILTERS_MARK = "HSCA_Active_Locations"
EXCLUDED_MARKS = ("Latest_ratings", "Deactivated_Locations", "CQC_directory")
FORBIDDEN_RE = re.compile(r"E080000(16|19|38|39)(?![0-9])")
# England with a margin (Scilly about -6.45, Lowestoft 1.76, Berwick 55.81)
LAT_RANGE = (Decimal("49.8"), Decimal("55.9"))
LON_RANGE = (Decimal("-6.5"), Decimal("1.9"))
COORD_PLACES = 6                         # numeric(9,6): more would be rounded

# stop conditions (calibrated on Jul->Aug, Aug->Sep, Sep->Oct 2026)
MAX_ROW_CHANGE = Decimal("0.02")         # seen at most 0.25%
MAX_GONE = 1000                          # seen 106..179
MAX_NEW = 1000                           # seen 106..199
MAX_UNRESOLVED = 10                      # seen 5
SL_MIN_COUNT = 10                        # supported living: authorities
SL_MAX_CHANGE = Decimal("0.5")           # ... changing by more than 50%
SL_MAX_AUTHORITIES = 5                   # ... in more than five of them
MAX_NEAREST_M = 2000                     # nearest polygon (seen under 700 m)

EXPECTED_HEADERS = (
    "Location ID",
    "Location HSCA start date",
    "Dormant (Y/N)",
    "Care home?",
    "Location Name",
    "Location ODS Code",
    "Location Telephone Number",
    "Registered manager",
    "Location Web Address",
    "Care homes beds",
    "Location Type/Sector",
    "Location Inspection Directorate",
    "Location Primary Inspection Category",
    "Location Latest Overall Rating",
    "Publication Date",
    "Inherited Rating (Y/N)",
    "Location Region",
    "Location NHS Region",
    "Location Local Authority",
    "Location ONSPD CCG Code",
    "Location ONSPD CCG",
    "Location Commissioning CCG Code",
    "Location Commissioning CCG",
    "Location Street Address",
    "Location Address Line 2",
    "Location City",
    "Location County",
    "Location Postal Code",
    "Location PAF ID",
    "Location UPRN ID",
    "Location Latitude",
    "Location Longitude",
    "Location Parliamentary Constituency",
    "Brand ID",
    "Brand Name",
    "Provider Companies House Number",
    "Provider Charity Number",
    "Provider ID",
    "Provider Name",
    "Provider HSCA start date",
    "Provider Type/Sector",
    "Provider Inspection Directorate",
    "Provider Primary Inspection Category",
    "Provider Ownership Type",
    "Provider Telephone Number",
    "Provider Web Address",
    "Provider Street Address",
    "Provider Address Line 2",
    "Provider City",
    "Provider County",
    "Provider Postal Code",
    "Provider PAF ID",
    "Provider UPRN ID",
    "Provider Local Authority",
    "Provider Region",
    "Provider NHS Region",
    "Provider Latitude",
    "Provider Longitude",
    "Provider Parliamentary Constituency",
    "Provider Nominated Individual Name",
    "Provider Main Partner Name",
    "Regulated activity - Accommodation for persons who require nursing or "
    "personal care",
    "Regulated activity - Accommodation for persons who require treatment "
    "for substance misuse",
    "Regulated activity - Assessment or medical treatment for persons "
    "detained under the Mental Health Act 1983",
    "Regulated activity - Diagnostic and screening procedures",
    "Regulated activity - Family planning",
    "Regulated activity - Management of supply of blood and blood derived "
    "products",
    "Regulated activity - Maternity and midwifery services",
    "Regulated activity - Nursing care",
    "Regulated activity - Personal care",
    "Regulated activity - Services in slimming clinics",
    "Regulated activity - Surgical procedures",
    "Regulated activity - Termination of pregnancies",
    "Regulated activity - Transport services, triage and medical advice "
    "provided remotely",
    "Regulated activity - Treatment of disease, disorder or injury",
    "Service type - Acute services with overnight beds",
    "Service type - Acute services without overnight beds / listed acute "
    "services with or without overnight beds",
    "Service type - Ambulance service",
    "Service type - Blood and Transplant service",
    "Service type - Care home service with nursing",
    "Service type - Care home service without nursing",
    "Service type - Community based services for people who misuse "
    "substances",
    "Service type - Community based services for people with a learning "
    "disability",
    "Service type - Community based services for people with mental health "
    "needs",
    "Service type - Community health care services - Independent Midwives",
    "Service type - Community health care services - Nurses Agency only",
    "Service type - Community healthcare service",
    "Service type - Dental service",
    "Service type - Diagnostic and/or screening service",
    "Service type - Diagnostic and/or screening service - single handed "
    "sessional providers",
    "Service type - Doctors consultation service",
    "Service type - Doctors treatment service",
    "Service type - Domiciliary care service",
    "Service type - Extra Care housing services",
    "Service type - Hospice services",
    "Service type - Hospice services at home",
    "Service type - Hospital services for people with mental health needs, "
    "learning disabilities and problems with substance misuse",
    "Service type - Hyperbaric Chamber",
    "Service type - Long term conditions services",
    "Service type - Mobile doctors service",
    "Service type - Prison Healthcare Services",
    "Service type - Rehabilitation services",
    "Service type - Remote clinical advice service",
    "Service type - Residential substance misuse treatment and/or "
    "rehabilitation service",
    "Service type - Shared Lives",
    "Service type - Specialist college service",
    "Service type - Supported living service",
    "Service type - Urgent care services",
    "Service user band - Children 0-18 years",
    "Service user band - Dementia",
    "Service user band - Learning disabilities or autistic spectrum disorder",
    "Service user band - Mental Health",
    "Service user band - Older People",
    "Service user band - People detained under the Mental Health Act",
    "Service user band - People who misuse drugs and alcohol",
    "Service user band - People with an eating disorder",
    "Service user band - Physical Disability",
    "Service user band - Sensory Impairment",
    "Service user band - Whole Population",
    "Service user band - Younger Adults",
    "Location Dual Registered",
    "Primary ID (Dual registration locations)",
)

# source column -> legacy column (the retired s11_cqc_process.py names)
TEXT_COLS = (
    ("Location ID", "location_id"),
    ("Provider ID", "provider_id"),
    ("Provider Name", "provider_name"),
    ("Brand Name", "brand_name"),
    ("Location Name", "location_name"),
    ("Location Postal Code", "postcode"),
    ("Location Region", "region"),
    ("Location Local Authority", "la_name_cqc"),
    ("Location Latest Overall Rating", "latest_overall_rating"),
    ("Location Inspection Directorate", "inspection_directorate"),
    ("Location Primary Inspection Category", "primary_inspection_category"),
    ("Primary ID (Dual registration locations)", "dual_primary_id"),
)
# 'Y' = registered for the service, blank = not (README: filter for 'Y')
Y_BLANK_FLAGS = (
    ("Service type - Supported living service", "supported_living"),
    ("Regulated activity - Personal care", "personal_care"),
    ("Service type - Domiciliary care service", "domiciliary_care"),
    ("Service type - Extra Care housing services", "extra_care_housing"),
    ("Service type - Shared Lives", "shared_lives"),
    ("Regulated activity - Accommodation for persons who require nursing or "
     "personal care", "accommodation_nursing_personal_care"),
    ("Service user band - Learning disabilities or autistic spectrum "
     "disorder", "band_learning_disabilities_autism"),
    ("Service user band - Mental Health", "band_mental_health"),
    ("Service user band - Younger Adults", "band_younger_adults"),
    ("Service user band - Older People", "band_older_people"),
    ("Service user band - Dementia", "band_dementia"),
    ("Service user band - People who misuse drugs and alcohol",
     "band_substance_misuse"),
    ("Service user band - Physical Disability", "band_physical_disability"),
    ("Service user band - People detained under the Mental Health Act",
     "band_detained_mha"),
)
# 'Y' / 'N', never blank
YN_FLAGS = (("Care home?", "care_home"), ("Dormant (Y/N)", "dormant"))
INHERITED = ("Inherited Rating (Y/N)", "inherited_rating")      # Y/N/blank
DUAL = ("Location Dual Registered", "dual_registered")
DUAL_MARK = "Dual Registration"
BRAND_NONE = "-"
COORDS = (("Location Latitude", "latitude"), ("Location Longitude",
                                               "longitude"))
BEDS = ("Care homes beds", "care_homes_beds")
DATES = (("Location HSCA start date", "location_hsca_start_date"),
         ("Publication Date", "rating_publication_date"))
DIRECTORATE = "Location Inspection Directorate"

# the value columns: the legacy cqc_locations columns and types, in its order
# (checked 2026-10-09 in information_schema)
_B = "boolean NOT NULL"
COLUMN_TYPES = (
    ("provider_id", "varchar(20)"),
    ("provider_name", "text"),
    ("brand_name", "text"),
    ("location_name", "text"),
    ("postcode", "varchar(10)"),
    ("latitude", "numeric(9,6)"),
    ("longitude", "numeric(9,6)"),
    ("lad24cd", "varchar(9) NOT NULL"),
    ("region", "varchar(30)"),
    ("la_name_cqc", "varchar(60)"),
    ("mapping_method", "varchar(30) NOT NULL"),
    ("supported_living", _B),
    ("personal_care", _B),
    ("care_home", _B),
    ("care_homes_beds", "integer"),
    ("domiciliary_care", _B),
    ("extra_care_housing", _B),
    ("shared_lives", _B),
    ("accommodation_nursing_personal_care", _B),
    ("band_learning_disabilities_autism", _B),
    ("band_mental_health", _B),
    ("band_younger_adults", _B),
    ("band_older_people", _B),
    ("band_dementia", _B),
    ("band_substance_misuse", _B),
    ("band_physical_disability", _B),
    ("band_detained_mha", _B),
    ("dormant", _B),
    ("dual_registered", _B),
    ("dual_primary_id", "varchar(20)"),
    ("latest_overall_rating", "varchar(40)"),
    ("rating_publication_date", "date"),
    ("inherited_rating", "boolean"),
    ("inspection_directorate", "varchar(40)"),
    ("primary_inspection_category", "varchar(60)"),
    ("location_hsca_start_date", "date"),
)
VALUE_COLUMNS = tuple(c for c, _ in COLUMN_TYPES)
MAPPING_COLUMNS = ("lad24cd", "mapping_method")
PUBLISHER_COLUMNS = tuple(c for c in VALUE_COLUMNS
                          if c not in MAPPING_COLUMNS)
LEGACY_COLUMNS = (("location_id",) + VALUE_COLUMNS
                  + ("is_active", "deregistered_seen_date",
                     "source_file_date", "loaded_at"))

_MONTHS = ("january", "february", "march", "april", "may", "june", "july",
           "august", "september", "october", "november", "december")
_MONTH_NUM = {m: i for i, m in enumerate(_MONTHS, 1)}
_NAME_DATE = re.compile(r"([0-9]{1,2})_([A-Za-z]+)_([0-9]{4})_"
                        + FILTERS_MARK, re.I)
_AS_AT = re.compile(r"Source:\s*CQC database as at\s+([0-9]{1,2})\s+"
                    r"([A-Za-z]+)\s+([0-9]{4})\s*$", re.I)
_LABEL_DATE = re.compile(r"\(([0-9]{1,2})\s+([A-Za-z]+)\s+([0-9]{4})\)")
_ISO_DATE = re.compile(r"([0-9]{4})-([0-9]{2})-([0-9]{2})"
                       r"(?:T00:00:00(?:\.0+)?)?")
_WHOLE = re.compile(r"[0-9]+")
_DECIMAL = re.compile(r"-?[0-9]+(?:\.[0-9]+)?")


def _dmy(day, month, year) -> date:
    m = _MONTH_NUM.get(month.lower())
    if m is None:
        raise ValueError(f"not a month name: {month!r}")
    return date(int(year), m, int(day))


# ---------------------------------------------------------------------------
# The page
# ---------------------------------------------------------------------------

def _name(url: str) -> str:
    return unquote(urlsplit(url).path.rsplit("/", 1)[-1])


_A_RE = re.compile(r"""<a\b[^>]*?href=["']([^"']+)["'][^>]*>(.*?)</a>""",
                   re.I | re.S)


def _anchors(html: str) -> list:
    """[(absolute url, anchor text)] in page order."""
    out = []
    for href, text in _A_RE.findall(html or ""):
        text = " ".join(re.sub(r"<[^>]+>", " ", html_lib.unescape(text))
                        .split())
        out.append((urljoin(PAGE, html_lib.unescape(href)), text))
    return out


def _is_filters(url: str) -> bool:
    name = _name(url)
    return (FILTERS_MARK.lower() in name.lower()
            and name.lower().endswith((".ods", ".xlsx"))
            and not any(x.lower() in name.lower() for x in EXCLUDED_MARKS))


def find_filters_link(html) -> str:
    """The one "Care directory with filters" link on the page: a file named
    HSCA_Active_Locations, .ods or .xlsx. The ratings, deactivated-locations
    and weekly directory files are excluded by name. ValueError if there is
    none (naming the spreadsheet links seen) or more than one (naming them):
    no winner is guessed."""
    seen, hits = [], []
    for url, _ in _anchors(html):
        if _name(url).lower().endswith((".ods", ".xlsx", ".csv", ".zip")):
            seen.append(_name(url))
        if _is_filters(url) and url not in hits:
            hits.append(url)
    if not hits:
        raise ValueError("no Care directory with filters link "
                         f"({FILTERS_MARK} .ods/.xlsx) on {PAGE}; spreadsheet "
                         f"links seen: {seen or 'none'} (excluded by name: "
                         f"{', '.join(EXCLUDED_MARKS)})")
    if len(hits) > 1:
        raise ValueError("more than one Care directory with filters link on "
                         f"the page, refusing to choose: {', '.join(hits)}")
    return hits[0]


def page_date(html) -> "date | None":
    """The date in the filters link's label, e.g. "Care directory with
    filters (02 October 2026)": the publication date (the file's own "as at"
    date is read from the file). None if the label carries none."""
    for url, text in _anchors(html):
        if _is_filters(url):
            m = _LABEL_DATE.search(text)
            if m:
                try:
                    return _dmy(*m.groups())
                except ValueError:
                    return None
    return None


def name_date(name) -> date:
    """The date in a file name like 01_October_2026_HSCA_Active_Locations.ods
    (also with a -<sha8> suffix). ValueError if there is none."""
    m = _NAME_DATE.search(Path(str(name)).name)
    if not m:
        raise ValueError(f"no DD_Month_YYYY_{FILTERS_MARK} date in the file "
                         f"name {Path(str(name)).name!r}")
    return _dmy(*m.groups())


# ---------------------------------------------------------------------------
# Reading the file (streamed: the ODS content.xml is about 440 MB)
# ---------------------------------------------------------------------------

_T = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
_O = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
_X = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def _ods_cell(cell) -> str:
    """A cell's value as text (the retired s11_cqc_fetch.py cell_value):
    numbers, booleans and dates from their office: value attribute, any
    other cell its text paragraphs joined by a newline."""
    vt = cell.get(_O + "value-type")
    if vt in ("float", "currency", "percentage"):
        return cell.get(_O + "value")
    if vt == "boolean":
        return cell.get(_O + "boolean-value")
    if vt == "date":
        return cell.get(_O + "date-value")
    parts = ["".join(p.itertext()) for p in cell.iter(_X + "p")]
    return "\n".join(parts) if parts else ""


def _ods_row(elem) -> list:
    row = []
    for cell in elem:
        if cell.tag not in (_T + "table-cell", _T + "covered-table-cell"):
            continue
        rep = int(cell.get(_T + "number-columns-repeated", "1"))
        val = _ods_cell(cell)
        if rep > 1000 and val == "":
            rep = 1  # trailing filler columns
        row.extend([val] * rep)
    while row and row[-1] == "":
        row.pop()
    return row


def _iter_ods(path, sheet):
    found = False
    with zipfile.ZipFile(path) as z, z.open("content.xml") as f:
        current = None
        for event, elem in ET.iterparse(f, events=("start", "end")):
            if event == "start" and elem.tag == _T + "table":
                current = elem.get(_T + "name")
                found = found or current == sheet
            elif event == "end" and elem.tag == _T + "table-row":
                if current == sheet:
                    row = _ods_row(elem)
                    if row:
                        rrep = int(elem.get(_T + "number-rows-repeated", "1"))
                        for _ in range(min(rrep, 1000)):
                            yield list(row)
                elem.clear()
            elif event == "end" and elem.tag == _T + "table":
                done = current == sheet
                current = None
                elem.clear()
                if done:
                    return
    if not found:
        raise ValueError(f"{Path(path).name}: no sheet {sheet!r}")


def _xlsx_text(v) -> str:
    """An openpyxl value as the text the ODS route gives."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v)


def _iter_xlsx(path, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"{Path(path).name}: no sheet {sheet!r} "
                             f"(sheets {wb.sheetnames})")
        for r in wb[sheet].iter_rows(values_only=True):
            row = [_xlsx_text(v) for v in r]
            while row and row[-1] == "":
                row.pop()
            if row:
                yield row
    finally:
        wb.close()


def iter_sheet_rows(path, sheet):
    """The non-empty rows of one sheet as lists of text (trailing blanks
    dropped): .ods streamed with xml.etree.iterparse (the retired
    s11_cqc_fetch.py conversion, the same repeat handling), .xlsx through
    openpyxl read-only. Any other format raises ValueError."""
    suffix = Path(path).suffix.lower()
    if suffix == ".ods":
        return _iter_ods(path, sheet)
    if suffix == ".xlsx":
        return _iter_xlsx(path, sheet)
    raise ValueError(f"{Path(path).name}: not an .ods or .xlsx file")


def first_sheet(path) -> "str | None":
    """The name of the file's first sheet (reads only until it starts)."""
    suffix = Path(path).suffix.lower()
    if suffix == ".ods":
        with zipfile.ZipFile(path) as z, z.open("content.xml") as f:
            for event, elem in ET.iterparse(f, events=("start",)):
                if elem.tag == _T + "table":
                    return elem.get(_T + "name")
        return None
    if suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            return wb.sheetnames[0] if wb.sheetnames else None
        finally:
            wb.close()
    raise ValueError(f"{Path(path).name}: not an .ods or .xlsx file")


def read_readme(path) -> dict:
    """The file's own identity from its README sheet: {'title': first line,
    'as_at': the date of the "Source: CQC database as at DD Month YYYY" line
    (None if there is none)}. ValueError if the sheet is missing."""
    title, as_at = None, None
    for i, row in enumerate(iter_sheet_rows(path, README_SHEET)):
        text = " ".join(row[0].split()) if row else ""
        if i == 0:
            title = text or None
        m = _AS_AT.search(text)
        if m and as_at is None:
            as_at = _dmy(*m.groups())
        if i > 40 or (as_at is not None and i > 2):
            break
    return {"title": title, "as_at": as_at}


def check_identity(readme: dict, name) -> list:
    """Problems (empty = fine): the README title must be the HSCA title
    exactly, and its "as at" date must be present and equal the date in the
    file name."""
    problems = []
    if readme.get("title") != TITLE:
        problems.append(f"README title is {readme.get('title')!r}, not "
                        f"{TITLE!r}")
    as_at = readme.get("as_at")
    if as_at is None:
        problems.append("README has no 'Source: CQC database as at DD Month "
                        "YYYY' line")
    try:
        nd = name_date(name)
    except ValueError as e:
        problems.append(str(e))
    else:
        if as_at is not None and as_at != nd:
            problems.append(f"README says as at {as_at} but the file name "
                            f"{Path(str(name)).name} says {nd}")
    return problems


def _header_problems(header) -> list:
    got, want = list(header), list(EXPECTED_HEADERS)
    out = []
    for i in range(max(len(got), len(want))):
        g = got[i] if i < len(got) else "(missing)"
        w = want[i] if i < len(want) else "(none expected)"
        if g != w:
            out.append(f"column {i + 1}: {g!r}, expected {w!r}")
    return out


def _flag(v, col, where, allowed):
    if v in allowed:
        return allowed[v]
    raise ValueError(f"{where}: {col!r} is {v!r}; only "
                     f"{', '.join(repr(k) for k in allowed)} known")


_YB = {"Y": True, "": False}
_YN = {"Y": True, "N": False}
_YNB = {"Y": True, "N": False, "": None}


def _coord(v, col, where):
    if v == "":
        return None
    if not _DECIMAL.fullmatch(v):
        raise ValueError(f"{where}: {col!r} is {v!r}, not a number")
    d = Decimal(v)
    if -d.normalize().as_tuple().exponent > COORD_PLACES:
        raise ValueError(f"{where}: {col!r} is {v!r}, more than "
                         f"{COORD_PLACES} decimal places (would be rounded)")
    lo, hi = LAT_RANGE if col == "Location Latitude" else LON_RANGE
    if not lo <= d <= hi:
        raise ValueError(f"{where}: {col!r} is {v!r}, outside England "
                         f"({lo} .. {hi})")
    return d


def _beds(v, where):
    if v == "":
        return None
    if not _WHOLE.fullmatch(v):
        raise ValueError(f"{where}: {BEDS[0]!r} is {v!r}, not a whole number")
    return int(v)


def _date(v, col, where):
    if v == "":
        return None
    m = _ISO_DATE.fullmatch(v)
    if not m:
        raise ValueError(f"{where}: {col!r} is {v!r}, not an ISO date")
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError as e:
        raise ValueError(f"{where}: {col!r} is {v!r} ({e})") from None


def _record(cells, idx, where) -> dict:
    """One in-scope row as a record with the legacy column names."""
    def get(col):
        return cells[idx[col]].strip()
    rec = {}
    for src, dst in TEXT_COLS:
        rec[dst] = get(src) or None
    where = f"{where} ({rec['location_id']})"
    if rec["brand_name"] == BRAND_NONE:
        rec["brand_name"] = None
    for src, dst in Y_BLANK_FLAGS:
        rec[dst] = _flag(get(src), src, where, _YB)
    for src, dst in YN_FLAGS:
        rec[dst] = _flag(get(src), src, where, _YN)
    rec[INHERITED[1]] = _flag(get(INHERITED[0]), INHERITED[0], where, _YNB)
    rec[DUAL[1]] = _flag(get(DUAL[0]), DUAL[0], where,
                         {DUAL_MARK: True, "": False})
    for src, dst in COORDS:
        rec[dst] = _coord(get(src), src, where)
    rec[BEDS[1]] = _beds(get(BEDS[0]), where)
    for src, dst in DATES:
        rec[dst] = _date(get(src), src, where)
    return rec


def parse_locations(path) -> tuple:
    """(rows, meta) of the HSCA_Active_Locations sheet: the Adult social care
    rows as records {location_id, <the publisher columns>} with the legacy
    column names (lad24cd and mapping_method are added by map_locations),
    and meta {'rows_total', 'in_scope', 'directorates': {name: rows},
    'markers': {column: {value: rows}}} (in-scope rows). Raises ValueError,
    naming the row and the value, on a header that is not EXPECTED_HEADERS
    (listing the changed columns), any unknown flag marker, a blank
    'Dormant (Y/N)' or 'Care home?', a non-whole bed count, an unparseable
    or out-of-England coordinate, an unparseable date, a blank or repeated
    location_id, a cell beyond the header, or any cell holding E08000016/19/
    38/39 (source 11 is declared 'none' in scripts/geography.py)."""
    name = Path(path).name
    it = iter_sheet_rows(path, DATA_SHEET)
    header = next(it, None)
    if header is None:
        raise ValueError(f"{name}: {DATA_SHEET} is empty")
    bad = _header_problems(header)
    if bad:
        raise ValueError(f"{name}: the {DATA_SHEET} header changed, nothing "
                         f"read: " + "; ".join(bad[:12])
                         + (f"; and {len(bad) - 12} more" if len(bad) > 12
                            else ""))
    idx = {h: i for i, h in enumerate(EXPECTED_HEADERS)}
    width = len(EXPECTED_HEADERS)
    rows, seen = [], set()
    directorates = Counter()
    markers = {}
    total = 0
    for n, raw in enumerate(it, start=2):
        total += 1
        where = f"{name} row {n}"
        if len(raw) > width:
            raise ValueError(f"{where}: {len(raw)} cells, the header has "
                             f"{width}")
        hit = FORBIDDEN_RE.search("\x1f".join(raw))
        if hit:
            raise ValueError(f"{where}: a cell holds {hit.group(0)}; source "
                             f"{RUN_SOURCE} is declared 'none' in "
                             "scripts/geography.py (no Barnsley/Sheffield "
                             "codes); nothing read")
        cells = raw + [""] * (width - len(raw))
        directorate = cells[idx[DIRECTORATE]].strip()
        directorates[directorate] += 1
        if directorate != SCOPE:
            continue
        rec = _record(cells, idx, where)
        lid = rec["location_id"]
        if lid is None:
            raise ValueError(f"{where}: blank Location ID in scope")
        if lid in seen:
            raise ValueError(f"{where}: Location ID {lid} appears twice in "
                             "scope; refusing to choose a winner")
        seen.add(lid)
        for col in ([c for c, _ in Y_BLANK_FLAGS + YN_FLAGS]
                    + [INHERITED[0], DUAL[0], BEDS[0], "Brand Name"]):
            v = cells[idx[col]].strip()
            if col == BEDS[0] and v not in ("", "0"):
                v = "(a number)"
            if col == "Brand Name" and v != BRAND_NONE:
                v = "(a name)"
            m = markers.setdefault(col, {})
            m[v] = m.get(v, 0) + 1
        rows.append(rec)
    meta = {"rows_total": total, "in_scope": len(rows),
            "directorates": dict(sorted(directorates.items())),
            "markers": markers}
    return rows, meta


def content_sha256(path) -> str:
    """SHA-256 of the file's bytes."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


# ---------------------------------------------------------------------------
# Mapping a location to a district
# ---------------------------------------------------------------------------

class PostcodesIO:
    """The read-only postcodes.io lookups the mapping needs (injected into
    map_locations; tests pass a stub with the same two methods)."""

    def __init__(self, session=None):
        self.session = session

    def _s(self):
        if self.session is not None:
            return self.session
        import requests
        return requests

    def lookup(self, postcodes) -> dict:
        """{postcode as given: admin_district code} for the live postcodes
        found (bulk, 100 per call). A failing call raises."""
        out = {}
        pcs = list(postcodes)
        for i in range(0, len(pcs), 100):
            r = self._s().post(f"{POSTCODES_IO}/postcodes",
                               json={"postcodes": pcs[i:i + 100]},
                               headers={"User-Agent": USER_AGENT}, timeout=60)
            r.raise_for_status()
            for item in r.json()["result"]:
                res = item.get("result")
                code = (res or {}).get("codes", {}).get("admin_district")
                if code:
                    out[item["query"]] = code
        return out

    def terminated(self, postcode) -> "tuple | None":
        """(latitude, longitude) of a terminated postcode, or None if
        postcodes.io does not know it as terminated (404)."""
        r = self._s().get(f"{POSTCODES_IO}/terminated_postcodes/"
                          + str(postcode).replace(" ", ""),
                          headers={"User-Agent": USER_AGENT}, timeout=30)
        if r.status_code == 404:
            return None
        r.raise_for_status()
        res = r.json().get("result")
        if not res or res.get("latitude") is None:
            return None
        return (res["latitude"], res["longitude"])


def _geom(g):
    from shapely.geometry import shape
    if isinstance(g, str):
        g = json.loads(g)
    if isinstance(g, dict):
        return shape(g)
    return g


class _Areas:
    """la_boundaries polygons with a spatial index; distances in metres on
    British National Grid (EPSG:27700), the retired map step's method."""

    def __init__(self, boundaries):
        from shapely.strtree import STRtree
        self.codes = [c for c, _ in boundaries]
        if len(set(self.codes)) != len(self.codes):
            raise ValueError("la_boundaries repeats a lad24cd")
        self.geoms = [_geom(g) for _, g in boundaries]
        self.tree = STRtree(self.geoms)
        self._bng = None

    def within(self, lon, lat) -> list:
        from shapely.geometry import Point
        hits = self.tree.query(Point(float(lon), float(lat)),
                               predicate="within")
        return sorted(self.codes[i] for i in hits)

    def _to_bng(self):
        from pyproj import Transformer
        return Transformer.from_crs("EPSG:4326", "EPSG:27700",
                                    always_xy=True).transform

    def nearest(self, lon, lat) -> tuple:
        """(lad24cd, metres) of the nearest polygon."""
        from shapely.geometry import Point
        from shapely.ops import transform
        tf = self._to_bng()
        if self._bng is None:
            self._bng = [transform(tf, g) for g in self.geoms]
        p = transform(tf, Point(float(lon), float(lat)))
        dist = [g.distance(p) for g in self._bng]
        best = min(range(len(dist)), key=lambda i: (dist[i], self.codes[i]))
        return self.codes[best], dist[best]


def _point_code(areas, lon, lat, where) -> "str | None":
    hits = areas.within(lon, lat)
    if len(hits) > 1:
        raise ValueError(f"{where}: the point ({lat}, {lon}) is inside "
                         f"{len(hits)} la_boundaries polygons {hits}; "
                         "refusing to choose")
    return hits[0] if hits else None


def map_locations(rows, boundaries, recodes, lookup, *, api) -> tuple:
    """(mapped, unresolved). Each record with both coordinates is placed by
    point-in-polygon into `boundaries` [(lad24cd, GeoJSON or shapely
    geometry)] ('point_in_polygon'); a point outside every polygon goes to
    the nearest one ('nearest_fallback') unless that is over MAX_NEAREST_M
    metres (raises). A record without both coordinates is looked up by
    postcode through api.lookup (postcodes.io): the code goes through
    geography.canonical(code, recodes), then must be in `boundaries`, else
    through `lookup` ({old_code: new_code} of la_code_lookup); a code still
    unknown raises (UNEXPLAINED) ('postcode_api_fallback'). A postcode the
    API does not know as live is tried with api.terminated: its coordinates
    are placed by point-in-polygon ('postcode_terminated_fallback'). Any
    other record is unresolved: {location_id, location_name, postcode,
    reason} with what was established. mapped: the records with lad24cd and
    mapping_method added (nearest and API rows also carry 'mapping_note')."""
    areas = _Areas(boundaries)
    valid = set(areas.codes)
    mapped, unresolved, pending = [], [], []
    for r in rows:
        lat, lon = r.get("latitude"), r.get("longitude")
        where = f"location {r['location_id']}"
        if lat is None or lon is None:
            pending.append(r)
            continue
        code = _point_code(areas, lon, lat, where)
        if code is not None:
            mapped.append({**r, "lad24cd": code,
                           "mapping_method": "point_in_polygon"})
            continue
        code, metres = areas.nearest(lon, lat)
        if metres > MAX_NEAREST_M:
            raise ValueError(f"{where} {r.get('location_name')!r}: the point "
                             f"({lat}, {lon}) is outside every polygon and "
                             f"{metres:.0f} m from the nearest ({code}), "
                             f"over {MAX_NEAREST_M} m; nothing stored")
        mapped.append({**r, "lad24cd": code,
                       "mapping_method": "nearest_fallback",
                       "mapping_note": f"nearest polygon {code}, "
                                       f"{metres:.0f} m"})
    postcodes = sorted({r["postcode"] for r in pending if r.get("postcode")})
    found = api.lookup(postcodes) if postcodes else {}
    for r in pending:
        where = f"location {r['location_id']}"
        pc = r.get("postcode")
        if not pc:
            unresolved.append(_unresolved(r, "the file carries no usable "
                                             "coordinates and no postcode"))
            continue
        api_code = found.get(pc)
        if api_code:
            code = geography.canonical(api_code, recodes)
            if code not in valid:
                code = lookup.get(code)
                if code is not None:
                    code = geography.canonical(code, recodes)
            if code not in valid:
                raise ValueError(f"{where}: postcodes.io gives {api_code} for "
                                 f"{pc}, which resolves to no la_boundaries "
                                 f"code (UNEXPLAINED {api_code}); explain it "
                                 "in la_code_lookup; nothing stored")
            mapped.append({**r, "lad24cd": code,
                           "mapping_method": "postcode_api_fallback",
                           "mapping_note": f"postcodes.io {pc} -> {api_code}"
                                           + (f" -> {code}"
                                              if code != api_code else "")})
            continue
        term = api.terminated(pc)
        if term is not None:
            tlat, tlon = term
            code = _point_code(areas, tlon, tlat, where)
            if code is not None:
                mapped.append({**r, "lad24cd": code,
                               "mapping_method":
                                   "postcode_terminated_fallback",
                               "mapping_note": f"terminated postcode {pc} "
                                               f"at ({tlat}, {tlon})"})
                continue
            unresolved.append(_unresolved(
                r, f"postcode {pc} is terminated; its coordinates ({tlat}, "
                   f"{tlon}) fall outside every la_boundaries polygon, and "
                   "the file carries no usable coordinates"))
            continue
        unresolved.append(_unresolved(
            r, f"postcode {pc} absent from ONS data (not live, not "
               "terminated) and the file carries no usable coordinates"
               + ("" if r.get("la_name_cqc") else " and no LA name")))
    return mapped, unresolved


def _unresolved(r, reason) -> dict:
    return {"location_id": r["location_id"],
            "location_name": r.get("location_name"),
            "postcode": r.get("postcode"), "reason": reason}


# ---------------------------------------------------------------------------
# Stop conditions
# ---------------------------------------------------------------------------

def sl_counts(records) -> Counter:
    """Supported-living (non-dormant) locations per lad24cd: W1's count."""
    return Counter(r["lad24cd"] for r in records
                   if r["supported_living"] and not r["dormant"])


def snapshot_problems(new: dict, prev: "dict | None") -> list:
    """The stop conditions for a snapshot (empty = fine). new: {'records',
    'unresolved' (count), 'authorities' (every lad24cd that must hold a
    location)}; prev: {'records'} of the snapshot it is compared with, or
    None (nothing held). In-scope rows changing by more than 2%; more than
    1,000 locations gone or more than 1,000 new; more than 10 unresolved; an
    authority with no location; the supported-living (non-dormant) count
    changing by more than 50% in more than 5 authorities with at least 10
    (in either snapshot)."""
    out = []
    recs = new["records"]
    if new["unresolved"] > MAX_UNRESOLVED:
        out.append(f"{new['unresolved']} unresolved locations (limit "
                   f"{MAX_UNRESOLVED})")
    empty = sorted(set(new["authorities"]) - {r["lad24cd"] for r in recs})
    if empty:
        out.append(f"{len(empty)} authorities with no location "
                   f"{empty[:10]}")
    if prev is None:
        return out
    old = prev["records"]
    if old:
        change = abs(Decimal(len(recs) - len(old))) / Decimal(len(old))
        if change > MAX_ROW_CHANGE:
            out.append(f"in-scope rows {len(old):,} -> {len(recs):,} "
                       f"({change * 100:.2f}%, limit "
                       f"{MAX_ROW_CHANGE * 100:.0f}%)")
    a = {r["location_id"] for r in old}
    b = {r["location_id"] for r in recs}
    if len(a - b) > MAX_GONE:
        out.append(f"{len(a - b):,} locations gone (limit {MAX_GONE:,})")
    if len(b - a) > MAX_NEW:
        out.append(f"{len(b - a):,} locations new (limit {MAX_NEW:,})")
    big = sl_big_changes(sl_counts(old), sl_counts(recs))
    if len(big) > SL_MAX_AUTHORITIES:
        out.append(f"supported-living count changes by more than "
                   f"{SL_MAX_CHANGE * 100:.0f}% in {len(big)} authorities "
                   f"with at least {SL_MIN_COUNT} (limit "
                   f"{SL_MAX_AUTHORITIES}): "
                   + ", ".join(f"{c} {p}->{n}" for c, p, n in big[:8]))
    return out


def sl_big_changes(old: Counter, new: Counter) -> list:
    """[(lad24cd, old, new)] where the count is at least SL_MIN_COUNT on one
    side and moves by more than SL_MAX_CHANGE of the old count (Counter
    gives an absent authority a count of none, which is a true count here:
    no row)."""
    out = []
    for c in sorted(set(old) | set(new)):
        p, n = old[c], new[c]
        if max(p, n) < SL_MIN_COUNT:
            continue
        if abs(n - p) > SL_MAX_CHANGE * p:
            out.append((c, p, n))
    return out


def zero_null_cells(old_records, new_records) -> list:
    """[(location_id, column, old, new)] where a 0 becomes NULL or a NULL
    becomes 0 (rule 1.10); booleans are not numbers here."""
    old = {r["location_id"]: r for r in old_records}
    out = []
    for r in new_records:
        o = old.get(r["location_id"])
        if o is None:
            continue
        for c in VALUE_COLUMNS:
            a, b = o.get(c), r.get(c)
            if isinstance(a, bool) or isinstance(b, bool):
                continue
            if (a is None and b is not None and b == 0) or (
                    b is None and a is not None and a == 0):
                out.append((r["location_id"], c, a, b))
    return out


def compare_snapshots(old_records, new_records) -> dict:
    """Read-only summary: rows, gone, new, common rows changed, changed
    cells by column, supported-living authorities changed."""
    old = {r["location_id"]: r for r in old_records}
    new = {r["location_id"]: r for r in new_records}
    by_col = Counter()
    changed = 0
    for k in set(old) & set(new):
        diff = [c for c in VALUE_COLUMNS if old[k].get(c) != new[k].get(c)]
        if diff:
            changed += 1
            by_col.update(diff)
    so, sn = sl_counts(old_records), sl_counts(new_records)
    moved = [(c, so[c], sn[c]) for c in sorted(set(so) | set(sn))
             if so[c] != sn[c]]
    return {"rows_old": len(old), "rows_new": len(new),
            "gone": len(set(old) - set(new)), "new": len(set(new) - set(old)),
            "changed": changed, "by_col": by_col, "sl_moved": moved,
            "sl_max": max((abs(n - p) for _, p, n in moved), default=None)}


def format_comparison(s: dict, what: str) -> str:
    cols = ", ".join(f"{c} {n}" for c, n in s["by_col"].most_common(8))
    return (f"compared with {what}: rows {s['rows_old']:,} -> "
            f"{s['rows_new']:,}; gone {s['gone']:,}; new {s['new']:,}; common "
            f"rows changed {s['changed']:,} (by column: {cols or 'none'}); "
            f"supported-living (non-dormant) count changes in "
            f"{len(s['sl_moved'])} authorities"
            + (f", largest {s['sl_max']}" if s["sl_moved"] else ""))


# ---------------------------------------------------------------------------
# Spec and profile
# ---------------------------------------------------------------------------

def tip_row_count(editions_table: str):
    """expected_rows_per_period for status: the tip edition's row count (a
    snapshot's size is its own, so the engine's modal count does not
    apply)."""
    core._ident(editions_table)

    def count(cur, period):
        cur.execute(f"SELECT DISTINCT edition, supersedes FROM "
                    f"public.{editions_table} WHERE snapshot_date = %s",
                    (period,))
        try:
            tip = core._chain_tip(cur.fetchall(), str(period))
        except (LookupError, ValueError):
            return None
        cur.execute(f"SELECT COUNT(*) FROM public.{editions_table} WHERE "
                    "snapshot_date = %s AND edition = %s", (period, tip))
        return cur.fetchone()[0]
    return count


SPEC = core.EditionSpec(
    name="s11",
    live_table=LIVE,
    editions_table=TABLE,
    key_cols=("location_id",),
    period_col="snapshot_date",
    value_cols=COLUMN_TYPES,
    refresh_cols=VALUE_COLUMNS + ("source_file",),
    refresh_from=(("source_file", "source_file"),),
    as_loaded_source_col="source_file",
    refresh_source_whole_period=True,   # one source per refreshed snapshot
    key_types=(("location_id", "varchar(20) NOT NULL"),
               ("snapshot_date", "date NOT NULL")),
    fk_la_boundaries=False,             # lad24cd is not a key; a gate checks
    expected_rows_per_period=tip_row_count(TABLE),
)


def _no_label(fetched_on) -> str:
    raise ValueError("an S11 release label comes from the file; store "
                     "through apply_snapshot(...)")


def _late(name: str):
    """A call through this module's attribute `name`, looked up when called
    (so mock.patch.object(s11_cqc_editions, name, ...) reaches the engine)."""
    return lambda *a, **kw: globals()[name](*a, **kw)


PROFILE = pe.Profile(
    spec=SPEC, value_cols=VALUE_COLUMNS, run_agent=RUN_AGENT,
    run_source=RUN_SOURCE, heading="S11 CQC Care directory with filters",
    default_source_file=SOURCE_FILE,
    expected_areas=None,   # a snapshot's size is its own (tip row count)
    release_label=_no_label, file_checks=True,
    savepoint="s11_snapshot",
    example_label="(location values)")


def _profile(spec) -> pe.Profile:
    return PROFILE.with_spec(spec)


def unresolved_table(spec=SPEC) -> str:
    return f"{spec.live_table}_unresolved"


def _grant(cur, table):
    cur.execute(f"GRANT SELECT ON public.{table} TO pipeline_user, "
                "ucws_readonly")


def create_live(cur, spec=SPEC) -> None:
    """The live table (one row per location per snapshot, the latest-edition
    layer), its primary key (snapshot_date, location_id), the index the view
    needs (location_id, snapshot_date DESC) and SELECT grants. Idempotent."""
    t = spec.live_table
    cols = ",\n        ".join(
        ["location_id varchar(20) NOT NULL", "snapshot_date date NOT NULL"]
        + [f"{c} {sql}" for c, sql in spec.value_cols]
        + ["source_file text NOT NULL",
           "loaded_at timestamptz DEFAULT now()",
           "PRIMARY KEY (snapshot_date, location_id)"])
    cur.execute(f"CREATE TABLE IF NOT EXISTS public.{t} (\n        {cols}\n)")
    cur.execute(f"CREATE INDEX IF NOT EXISTS {t}_location_idx ON public.{t} "
                "(location_id, snapshot_date DESC)")
    cur.execute(f"COMMENT ON TABLE public.{t} IS 'CQC Care directory with "
                "filters: one row per Adult social care location per "
                "snapshot (the file''s own as-at date), every snapshot kept; "
                f"the latest edition of each snapshot (editions in "
                f"{spec.editions_table}). Loaded by scripts/"
                "s11_cqc_editions.py. cqc_locations is a view over it.'")
    _grant(cur, t)


def create_unresolved(cur, spec=SPEC) -> None:
    """The append-only record of each snapshot edition's unresolved
    locations (triggers as the file-check ledger's). Idempotent."""
    t = unresolved_table(spec)
    fn = f"{spec.name}_unresolved_immutable"
    cur.execute(f"""CREATE TABLE IF NOT EXISTS public.{t} (
        snapshot_date date NOT NULL,
        edition integer NOT NULL,
        location_id varchar(20) NOT NULL,
        location_name text,
        postcode text,
        reason text NOT NULL,
        source_file text NOT NULL,
        recorded_at timestamptz NOT NULL DEFAULT now(),
        PRIMARY KEY (snapshot_date, edition, location_id)
    )""")
    cur.execute(f"""
    CREATE OR REPLACE FUNCTION public.{fn}() RETURNS trigger AS $f$
    BEGIN
        RAISE EXCEPTION '{t} is append-only: % is not permitted', TG_OP;
    END
    $f$ LANGUAGE plpgsql""")
    cur.execute(f"""CREATE OR REPLACE TRIGGER {fn}
        BEFORE UPDATE OR DELETE ON public.{t}
        FOR EACH ROW EXECUTE FUNCTION public.{fn}()""")
    cur.execute(f"""CREATE OR REPLACE TRIGGER {spec.name}_unresolved_no_truncate
        BEFORE TRUNCATE ON public.{t}
        FOR EACH STATEMENT EXECUTE FUNCTION public.{fn}()""")
    cur.execute(f"COMMENT ON TABLE public.{t} IS 'In-scope CQC locations of "
                "each snapshot edition that could not be given a lad24cd and "
                f"so are absent from {spec.live_table} (UNEXPLAINED, never "
                "benign: each understates provision by one location). "
                "Append-only; written with the snapshot.'")
    _grant(cur, t)


def create_all(cur, spec=SPEC) -> None:
    """ddl: editions table and triggers, live table, unresolved table (the
    engine adds the file-check ledger)."""
    core.create_schema(cur, spec)
    create_live(cur, spec)
    create_unresolved(cur, spec)
    _grant(cur, spec.editions_table)


def insert_live(cur, profile, period: str, records: list) -> None:
    """The snapshot's live rows with source_file (each record's
    'source_file'); loaded_at takes its default."""
    from psycopg2.extras import execute_values
    spec = profile.spec
    cols = (tuple(spec.key_cols) + (spec.period_col,)
            + tuple(profile.value_cols) + ("source_file",))
    execute_values(
        cur, f"INSERT INTO public.{spec.live_table} ({', '.join(cols)}) "
        "VALUES %s",
        [tuple(r[k] for k in spec.key_cols) + (period,)
         + tuple(r[c] for c in profile.value_cols) + (r["source_file"],)
         for r in records],
        page_size=1000)


def insert_unresolved(cur, spec, period, edition, unresolved,
                      source_file) -> None:
    if not unresolved:
        return
    from psycopg2.extras import execute_values
    execute_values(
        cur, f"INSERT INTO public.{unresolved_table(spec)} (snapshot_date, "
        "edition, location_id, location_name, postcode, reason, source_file) "
        "VALUES %s",
        [(period, edition, u["location_id"], u.get("location_name"),
          u.get("postcode"), u["reason"], source_file) for u in unresolved])


@dataclass(frozen=True)
class FileInfo:
    """One file: what the edition records as source_file (the URL, or the
    path), the path read, the file sha256 and the release label."""
    source_file: str
    path: Path
    sha256: str
    label: str


def release_label(as_at, published, sha: str, how: str,
                  acknowledged: str = "") -> str:
    return (f"CQC Care directory with filters as at {as_at}; published "
            f"{published or '-'}; file sha256 {sha[:16]}; {how}"
            + (f"; ACKNOWLEDGED: {acknowledged}" if acknowledged else ""))


def apply_snapshot(cur, profile, period, records, *, fetched_on, info,
                   unresolved):
    """pe.apply_period for one snapshot, then, in the same savepoint, its
    unresolved rows (for a new edition) and the file-check ledger row: a new
    snapshot's edition 1, live rows, unresolved rows and ledger row commit
    or roll back together."""
    prof = dataclasses.replace(profile, release_label=lambda d: info.label)
    kind = pe.apply_period(cur, prof, period, records, fetched_on=fetched_on,
                           source_file=info.source_file,
                           insert=_late("insert_live"))
    tip = core.chain_tip(cur, prof.spec, period)
    if kind in ("new", "revised"):
        insert_unresolved(cur, prof.spec, period, tip, unresolved,
                          info.source_file)
    pe.record_file_check(cur, prof, period, info.source_file, info.sha256,
                         kind, tip)
    return kind


# ---------------------------------------------------------------------------
# Files: page and download
# ---------------------------------------------------------------------------

def _session(session):
    if session is not None:
        return session
    import requests
    return requests


def _get(session, url, **kw):
    try:
        r = session.get(url, headers={"User-Agent": USER_AGENT},
                        timeout=600, **kw)
        r.raise_for_status()
    except Exception as e:  # noqa: BLE001 (any failure halts)
        halt(f"GET {url} failed, nothing downloaded: {e}")
    return r


def fetch_page(session=None) -> str:
    """The CQC data page's HTML (a read-only GET)."""
    return _get(_session(session), PAGE).text


def fetch_file(url, dest=None, session=None) -> Path:
    """Download the filters file to dest (default data/raw/) under the
    link's file name and return the path to read. Checked before it is
    kept: more than MIN_FILE_BYTES and opening with a README sheet. A
    same-named file with the same content is kept as it is; one with
    different content is never replaced: the new download is saved beside it
    as <stem>-<sha8><ext> and that is the one read."""
    dest = Path(dest) if dest is not None else RAW_DIR
    body = _get(_session(session), url).content
    if len(body) <= MIN_FILE_BYTES:
        halt(f"{url}: {len(body)} bytes, expected more than "
             f"{MIN_FILE_BYTES}; nothing written; saw {body[:100]!r}")
    name = _name(url)
    if Path(name).suffix.lower() not in (".ods", ".xlsx"):
        halt(f"{url}: not an .ods or .xlsx file; nothing written")
    dest.mkdir(parents=True, exist_ok=True)
    target = dest / name
    new_sha = hashlib.sha256(body).hexdigest()
    tmp = dest / (Path(name).stem + ".download.tmp" + Path(name).suffix)
    tmp.write_bytes(body)
    try:
        sheet = first_sheet(tmp)
        if sheet != README_SHEET:
            raise ValueError(f"the first sheet is {sheet!r}, not "
                             f"{README_SHEET!r}")
    except Exception as e:  # noqa: BLE001
        tmp.unlink()
        halt(f"{url}: not a usable CQC filters file ({e}); nothing kept")
    path = target
    if target.exists():
        old_sha = content_sha256(target)
        if old_sha == new_sha:
            tmp.unlink()
            print(f"{name} already held with the same content (sha256 "
                  f"{new_sha[:16]}); kept as it is")
        else:
            path = dest / f"{target.stem}-{new_sha[:8]}{target.suffix}"
            tmp.replace(path)
            print(f"NOTE: {name} already exists with different content "
                  f"(sha256 {old_sha[:16]}, new {new_sha[:16]}); kept it "
                  f"untouched and saved the new download as {path.name}")
    else:
        tmp.replace(target)
        print(f"downloaded {name}: {len(body):,} bytes, sha256 "
              f"{new_sha[:16]}")
    return path


def _file_source(path: Path) -> str:
    """source_file for a --file or migrated file: the path under the repo
    (data/raw/...), else the absolute path."""
    p = Path(path).resolve()
    try:
        return p.relative_to(REPO).as_posix()
    except ValueError:
        return p.as_posix()


def _api():
    return PostcodesIO()


def load_boundaries(cur) -> list:
    """[(lad24cd, geojson)] of la_boundaries, ordered by code."""
    cur.execute("SELECT lad24cd, geojson FROM public.la_boundaries "
                "ORDER BY lad24cd")
    return cur.fetchall()


def load_lookup(cur) -> dict:
    """{old_code: new_code} of la_code_lookup rows that are not 'current'
    (the retired map step's reconciliation, now after canonical)."""
    cur.execute("SELECT old_code, new_code FROM public.la_code_lookup "
                "WHERE change_type <> 'current'")
    return dict(cur.fetchall())


# ---------------------------------------------------------------------------
# Planning a load
# ---------------------------------------------------------------------------

def plan_load(held_latest, as_at, tips, checked, url, sha, *, recheck,
              allow_older, stranded=()) -> tuple:
    """(action, note). Before the download (as_at and sha None): 'skip' when
    url is already in the ledger (checked: {period: {(source_file, sha)}})
    for a snapshot whose live rows are present and no --recheck is given,
    else 'fetch'. After reading the file's own as-at date and sha: 'skip'
    when that file was already checked for that held snapshot (tips: {held
    period: its tip edition's source_file}) and no --recheck; 'recheck'
    for a held snapshot; 'new' otherwise. A file earlier than held_latest
    raises ValueError unless allow_older (then 'new', a backfill, or
    'recheck', with a note) or it is the --recheck snapshot. --recheck
    PERIOD must be the file's snapshot and a held one. Stranded snapshots
    (editions but no live rows) are always read."""
    stranded = set(stranded)
    if as_at is None:
        if url is None or recheck:
            return "fetch", ""
        hit = sorted(p for p, pairs in checked.items()
                     if any(f == url for f, _ in pairs))
        if hit and not stranded & set(hit):
            return "skip", (f"{url} is already in the ledger for {hit[-1]}; "
                            f"nothing downloaded (--recheck {hit[-1]} reads "
                            "it again)")
        return "fetch", ""
    if recheck and recheck != as_at:
        raise ValueError(f"--recheck {recheck}: the file is as at {as_at}")
    held = as_at in tips
    if recheck and not held:
        raise ValueError(f"--recheck {recheck}: not a held snapshot")
    if (held and not recheck and as_at not in stranded
            and any(s == sha for _, s in checked.get(as_at, ()))):
        return "skip", (f"this file (sha256 {sha[:16]}) was already checked "
                        f"for {as_at}; nothing to do (--recheck {as_at} "
                        "reads it again)")
    older = held_latest is not None and as_at < held_latest
    note = ""
    if older and not (held and recheck):
        msg = (f"the file is as at {as_at}, earlier than the latest held "
               f"snapshot {held_latest}")
        if not allow_older:
            raise ValueError(msg + "; storing it would record an older file "
                             "as the latest. If this is deliberate (a "
                             "backfill), re-run with --allow-older-file")
        note = f"NOTE: --allow-older-file given; {msg}"
    return ("recheck" if held else "new"), note


# ---------------------------------------------------------------------------
# Reading held snapshots
# ---------------------------------------------------------------------------

def _dicts(cur, cols) -> list:
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def tip_records(cur, spec, period) -> list:
    """The records of a snapshot's tip edition."""
    tip = core.latest_edition(cur, spec, period)
    cols = ("location_id",) + VALUE_COLUMNS
    cur.execute(f"SELECT {', '.join(cols)} FROM public.{spec.editions_table} "
                "WHERE snapshot_date = %s AND edition = %s", (period, tip))
    return _dicts(cur, cols)


def legacy_records(cur, period, legacy=LEGACY) -> list:
    """The legacy table's rows with source_file_date = period."""
    cols = ("location_id",) + VALUE_COLUMNS
    cur.execute(f"SELECT {', '.join(cols)} FROM public.{legacy} WHERE "
                "source_file_date = %s", (period,))
    return _dicts(cur, cols)


def tip_files(cur, spec) -> dict:
    """{period: the source_file of its latest edition}."""
    cur.execute(f"SELECT DISTINCT ON (snapshot_date) snapshot_date, "
                f"source_file FROM public.{spec.editions_table} "
                "ORDER BY snapshot_date, edition DESC")
    return {pe._p(p): s for p, s in cur.fetchall()}


def ledger_checks(cur, profile) -> dict:
    """{period: {(source_file, file_sha256)}} of the file-check ledger."""
    cur.execute(f"SELECT DISTINCT snapshot_date, source_file, file_sha256 "
                f"FROM public.{pe.file_checks_table(profile)}")
    out = {}
    for p, f, s in cur.fetchall():
        out.setdefault(pe._p(p), set()).add((f, s))
    return out


def relkind(cur, name) -> "str | None":
    """'r' (table), 'v' (view) or None."""
    cur.execute("SELECT c.relkind FROM pg_class c JOIN pg_namespace n ON "
                "n.oid = c.relnamespace WHERE n.nspname = 'public' AND "
                "c.relname = %s", (name,))
    r = cur.fetchone()
    return r[0] if r else None


def _tables(cur, spec) -> dict:
    prof = _profile(spec)
    return {t: table_exists(cur, t)
            for t in (spec.live_table, spec.editions_table,
                      pe.file_checks_table(prof), unresolved_table(spec))}


# ---------------------------------------------------------------------------
# Status, run log, commands
# ---------------------------------------------------------------------------

table_exists = pe.table_exists


def _conn(writing):
    return pe.connect(writing)


def status(cur, spec=SPEC) -> dict:
    return pe.status(cur, _profile(spec))


def status_lines(cur, spec=SPEC, view=LEGACY) -> list:
    """The unresolved count of the latest snapshot and whether cqc_locations
    is the view."""
    lines = []
    cur.execute(f"SELECT MAX(snapshot_date) FROM public.{spec.editions_table}")
    latest = cur.fetchone()[0]
    if latest is not None:
        tip = core.latest_edition(cur, spec, latest)
        cur.execute(f"SELECT COUNT(*) FROM public.{unresolved_table(spec)} "
                    "WHERE snapshot_date = %s AND edition = %s",
                    (latest, tip))
        lines.append(f"  latest snapshot {latest} (edition {tip}): "
                     f"{cur.fetchone()[0]} unresolved location(s)")
    kind = relkind(cur, view)
    lines.append(f"  {view}: " + {"v": "the view over the snapshots",
                                  "r": "still the legacy base table (run "
                                       "migrate-legacy)"}.get(kind,
                                                              "absent"))
    return lines


def format_status(st: dict, extra=()) -> str:
    return pe.format_status(PROFILE, st, extra)


def log_run(cur, rows_written: int, notes: str, started_at=None) -> None:
    """The pipeline_run_log row for a committed run (agent RUN_AGENT, source
    '11'). Called only on `load --commit` and `migrate-legacy --commit`."""
    pe.log_run(cur, PROFILE, rows_written, notes, started_at)


def cmd_ddl(args) -> int:
    return pe.run_ddl(_profile(SPEC), args, connect=_late("_conn"),
                      table_exists=_late("table_exists"),
                      create_schema=lambda cur: create_all(cur, SPEC))


def cmd_status(args) -> int:
    """status: what needs action; exit 1 if anything, or (with a clean
    message, nothing created) if the editions table does not exist."""
    spec = SPEC
    conn = _conn(False)
    try:
        with conn.cursor() as cur:
            if not table_exists(cur, spec.editions_table):
                print(f"status: {spec.editions_table} does not exist yet; "
                      "run `ddl --commit`, then `migrate-legacy FILE FILE "
                      "FILE --commit` (S11 has no sync-new)")
                return 1
            lines = status_lines(cur, spec, LEGACY)
            st = status(cur, spec)
    finally:
        conn.close()
    print(format_status(st, lines))
    return 0 if st["ok"] else 1


def cmd_refresh_latest(args) -> int:
    return pe.run_refresh_latest(_profile(SPEC), args, connect=_late("_conn"),
                                 table_exists=_late("table_exists"))


def _iso_day(text: str) -> str:
    try:
        return date.fromisoformat(text).isoformat()
    except (TypeError, ValueError):
        raise argparse.ArgumentTypeError(
            f"not a date YYYY-MM-DD: {text!r}") from None


def _read_file(path: Path) -> dict:
    try:
        readme = read_readme(path)
    except ValueError as e:
        halt(f"{path.name}: {e}; nothing stored")
    problems = check_identity(readme, path.name)
    if problems:
        halt(f"file identity check failed for {path.name}, nothing stored: "
             + "; ".join(problems))
    return readme


def cmd_load(args) -> int:
    spec = SPEC
    base = _profile(spec)
    writing = args.commit or args.simulate
    conn = _conn(writing)
    try:
        with conn.cursor() as cur:
            tables = _tables(cur, spec)
            ready = all(tables.values())
            if writing and not ready:
                halt("missing " + ", ".join(t for t, ok in tables.items()
                                           if not ok)
                     + "; run `ddl --commit` first")
            if ready:
                _, new_live, errors = core.latest_map(cur, spec)
                if errors:
                    halt("invalid edition chain: " + "; ".join(
                        f"{pe._p(p)}: {msg}" for p, msg in errors.items()))
                if new_live:
                    halt(f"live snapshots with no editions "
                         f"{[pe._p(x) for x in new_live]}; changed outside "
                         "the loader")
                tips = tip_files(cur, spec)
                checked = ledger_checks(cur, base)
                stranded = pe.live_missing_periods(cur, base)
                held = sorted(set(tips) | set(pe.held_periods(cur, base,
                                                              True)))
            else:
                print(f"NOTE: {', '.join(t for t, ok in tables.items() if not ok)} "
                      "do not exist yet; this preview compares with the "
                      f"legacy {LEGACY} table")
                tips, checked, stranded = {}, {}, []
                held = []
                if relkind(cur, LEGACY) is not None:
                    cur.execute(f"SELECT DISTINCT source_file_date FROM "
                                f"public.{LEGACY} ORDER BY 1")
                    held = [pe._p(r[0]) for r in cur.fetchall()]
            held_latest = held[-1] if held else None
            print("held snapshots: " + (", ".join(held) or "none"))
            published = None
            url = None
            if args.file:
                path = Path(args.file)
                source_file = _file_source(path)
                print("file given: nothing downloaded")
            else:
                html = fetch_page()
                try:
                    url = find_filters_link(html)
                except ValueError as e:
                    halt(str(e))
                published = page_date(html)
                print(f"page: {url} (label date {published or '-'})")
                action, note = plan_load(
                    held_latest, None, tips, checked, url, None,
                    recheck=args.recheck, allow_older=args.allow_older_file,
                    stranded=stranded)
                if action == "skip":
                    print(note)
                    print("nothing to fetch")
                    return _end(conn, args, writing, rc=0)
                path = fetch_file(url, RAW_DIR)
                source_file = url
            readme = _read_file(path)
            as_at = readme["as_at"].isoformat()
            sha = content_sha256(path)
            print(f"{path.name}: README as at {as_at}, sha256 {sha[:16]}")
            try:
                action, note = plan_load(
                    held_latest, as_at, tips, checked, source_file, sha,
                    recheck=args.recheck, allow_older=args.allow_older_file,
                    stranded=stranded)
            except ValueError as e:
                halt(str(e))
            if note:
                print(note)
            if action == "skip":
                print("nothing to do")
                return _end(conn, args, writing, rc=0)
            ack = set(args.acknowledge or ())
            stray = sorted(ack - {as_at})
            if stray:
                halt(f"--acknowledge {stray}: not the snapshot of this run "
                     f"({as_at})")
            rc = _run_snapshot(args, conn, cur, spec, base, ready, held,
                               as_at, path, source_file, sha, published,
                               action, note, ack)
        return _end(conn, args, writing, rc)
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def _end(conn, args, writing, rc) -> int:
    conn.rollback()  # commit mode has committed already
    if not writing:
        print("PREVIEW: nothing written to the database (use --commit or "
              "--simulate)")
    elif args.simulate:
        print("SIMULATION: ROLLED BACK (nothing persisted)")
    return rc


ENGINE_GAP = ("the location set differs from the held edition of this "
              "snapshot; refresh-latest cannot move a changed set of "
              "locations into live (editions_core._unrepairable), so a "
              "reissued same-date file that adds or drops locations is not "
              "stored (a whole-snapshot replace in the engine is the "
              "follow-up)")


def _map(cur, rows):
    boundaries = load_boundaries(cur)
    try:
        recodes = geography.load_recodes(cur)
    except ValueError as e:
        halt(f"Barnsley/Sheffield recodes: {e}; nothing stored")
    lookup = load_lookup(cur)
    try:
        mapped, unresolved = map_locations(rows, boundaries, recodes, lookup,
                                           api=_api())
    except ValueError as e:
        halt(f"{e}")
    return mapped, unresolved, {c for c, _ in boundaries}


def _print_mapping(mapped, unresolved):
    methods = Counter(r["mapping_method"] for r in mapped)
    print("mapping: " + ", ".join(f"{k} {v}" for k, v in
                                  sorted(methods.items())))
    for r in mapped:
        if r.get("mapping_note"):
            print(f"  {r['mapping_method']}: {r['location_id']} "
                  f"{r.get('location_name')!r}: {r['mapping_note']}")
    print(f"unresolved: {len(unresolved)}")
    for u in unresolved:
        print(f"  {u['location_id']} {u.get('location_name')!r} "
              f"{u.get('postcode')}: {u['reason']}")


def _run_snapshot(args, conn, cur, spec, base, ready, held, as_at, path,
                  source_file, sha, published, action, note, ack) -> int:
    started = datetime.now(timezone.utc)
    try:
        rows, meta = parse_locations(path)
    except ValueError as e:
        halt(f"{e}")
    print(f"{path.name}: {meta['rows_total']:,} rows, {meta['in_scope']:,} "
          f"Adult social care; directorates {meta['directorates']}")
    mapped, unresolved, authorities = _map(cur, rows)
    _print_mapping(mapped, unresolved)
    records = [{**r, "snapshot_date": as_at, "source_file": source_file}
               for r in mapped]
    # what the snapshot is compared with
    if action == "recheck":
        prev, what = tip_records(cur, spec, as_at), f"the held {as_at}"
    else:
        earlier = [p for p in held if p < as_at]
        prev, what = None, "nothing (no earlier snapshot held)"
        if earlier:
            p = earlier[-1]
            if ready:
                prev, what = tip_records(cur, spec, p), f"snapshot {p}"
            else:
                prev, what = (legacy_records(cur, p, LEGACY),
                                  f"legacy {LEGACY} {p}")
    if prev is not None:
        print(format_comparison(compare_snapshots(prev, records), what))
    problems = snapshot_problems(
        {"records": records, "unresolved": len(unresolved),
         "authorities": authorities},
        {"records": prev} if prev is not None else None)
    acked = ""
    if action == "recheck":
        if {r["location_id"] for r in prev} != {r["location_id"]
                                                for r in records}:
            problems.append(ENGINE_GAP)
        zn = zero_null_cells(prev, records)
        if zn:
            msg = (f"{len(zn)} cell(s) go from 0 to NULL or NULL to 0, e.g. "
                   f"{zn[:3]} (rule 1.10)")
            if as_at in ack:
                acked = msg
                print(f"  {as_at}: ACKNOWLEDGED (--acknowledge): {msg}")
            else:
                problems.append(msg + f"; read them, then --acknowledge "
                                f"{as_at}")
    if problems:
        for p in problems:
            print(f"  {as_at}: REJECTED, not stored: STOP CONDITION: {p}")
        print(f"REJECTED {as_at}, nothing stored; exit 1")
        if args.commit:
            print("pipeline_run_log: no row written (the snapshot was "
                  "rejected)")
        return 1
    how = ("recheck of a held snapshot" if action == "recheck"
           else "new snapshot")
    if note:
        how += "; " + note.removeprefix("NOTE: ")
    info = FileInfo(source_file, path, sha,
                    release_label(as_at, published, sha, how, acked))
    if not ready:
        print(f"  {as_at}: {action}; would store edition 1 and insert "
              f"{len(records):,} live rows and {len(unresolved)} unresolved "
              "rows (tables not created yet: preview only)")
        return 0
    stats = {}
    rc = pe.load_periods(
        cur, base, [as_at], lambda p: records, published or date.today(),
        args.commit,
        simulate=args.simulate, against="editions", stats=stats,
        apply=lambda c, prof, p, recs, *, fetched_on, source_file=None:
        _late("apply_snapshot")(c, prof, p, recs, fetched_on=fetched_on,
                                info=info, unresolved=unresolved))
    if args.commit:
        if rc == 0:
            kind = stats["kinds"].get(as_at)
            notes = (f"{as_at}: {kind}; {len(records):,} locations, "
                     f"{len(unresolved)} unresolved; file {source_file} "
                     f"(sha256 {sha[:16]})"
                     + (f"; {note}" if note else "")
                     + (f"; ACKNOWLEDGED: {acked}" if acked else "")
                     + f". Rows stored: {stats['stored_rows']} edition rows, "
                     f"{stats['live_rows']} live rows.")
            log_run(cur, stats["stored_rows"] + stats["live_rows"],  # not a source value
                    notes, started)
            conn.commit()
            print("pipeline_run_log row written")
        else:
            print("pipeline_run_log: no row written (the run failed)")
    return rc


# ---------------------------------------------------------------------------
# One-off: migrate the legacy table
# ---------------------------------------------------------------------------

VIEW_SQL = """CREATE VIEW public.{view} AS
WITH snaps AS (SELECT DISTINCT snapshot_date FROM public.{live}),
latest AS (SELECT MAX(snapshot_date) AS d FROM snaps),
last_seen AS (
    SELECT DISTINCT ON (location_id) *
    FROM public.{live}
    ORDER BY location_id, snapshot_date DESC
)
SELECT {cols},
    r.snapshot_date = (SELECT d FROM latest) AS is_active,
    CASE WHEN r.snapshot_date = (SELECT d FROM latest) THEN NULL::date
         ELSE (SELECT MIN(s.snapshot_date) FROM snaps s
               WHERE s.snapshot_date > r.snapshot_date)
    END AS deregistered_seen_date,
    r.snapshot_date AS source_file_date,
    r.loaded_at
FROM last_seen r"""

W1_S11_SQL = ("SELECT lad24cd, COUNT(*) AS sl_count FROM public.{table} "
              "WHERE supported_living AND is_active AND NOT dormant "
              "GROUP BY lad24cd")


def view_sql(view=LEGACY, live=LIVE) -> str:
    cols = ", ".join(f"r.{c}" for c in ("location_id",) + VALUE_COLUMNS)
    return VIEW_SQL.format(view=view, live=live, cols=cols)


def view_differences(cur, view, table) -> tuple:
    """(rows only in the view, rows only in the table) on every legacy
    column but loaded_at (EXCEPT ALL both ways)."""
    cols = ", ".join(c for c in LEGACY_COLUMNS if c != "loaded_at")
    cur.execute(f"""SELECT
        (SELECT COUNT(*) FROM (SELECT {cols} FROM public.{view} EXCEPT ALL
                               SELECT {cols} FROM public.{table}) a),
        (SELECT COUNT(*) FROM (SELECT {cols} FROM public.{table} EXCEPT ALL
                               SELECT {cols} FROM public.{view}) b)""")
    return tuple(cur.fetchone())


def w1_counts(cur, table) -> dict:
    cur.execute(W1_S11_SQL.format(table=table))
    return dict(cur.fetchall())


_FILE_DATE_RE = re.compile(r"file date ([0-9]{4}-[0-9]{2}-[0-9]{2})")


def legacy_run_dates(cur, agent=LEGACY_RUN_AGENT) -> dict:
    """{file date: [(load date (UTC), rows_written)]} from the legacy run-log
    rows, oldest first."""
    cur.execute("SELECT notes, (completed_at AT TIME ZONE 'UTC')::date, "
                "rows_written FROM pipeline_run_log WHERE agent_name = %s "
                "ORDER BY completed_at", (agent,))
    out = {}
    for notes, d, n in cur.fetchall():
        m = _FILE_DATE_RE.search(notes or "")
        if m:
            out.setdefault(m.group(1), []).append((d.isoformat(), n))
    return out


def migrate_legacy(cur, files, *, write, spec=SPEC, legacy=LEGACY,
                   renamed=LEGACY_RENAMED,
                   legacy_unresolved=LEGACY_UNRESOLVED) -> dict:
    """One-off: every held snapshot from its file, checked against the
    legacy table, then (write) stored and the legacy table swapped for the
    view. Raises (halt) on any failure; the caller rolls everything back.

    Preconditions: `legacy` is a base table; the files' README dates are
    exactly the legacy source_file_date set and the legacy run-log file
    dates; on write the snapshot tables exist and are empty. Per file,
    oldest first: identity; parse; the legacy rows of that date must all be
    in the file and equal it on every publisher column (their lad24cd and
    mapping_method are taken from legacy, as loaded); the other rows are
    mapped now (reconstructed). The latest snapshot must be entirely the
    legacy rows and its unresolved set the open rows of legacy_unresolved.
    On write each snapshot goes through apply_snapshot (edition 1, live,
    unresolved, ledger); then `legacy` is renamed `renamed`, the view
    `legacy` is created and must equal `renamed` on every column but
    loaded_at, W1's S11 counts must be equal, and legacy_unresolved gets a
    comment saying it is frozen."""
    if relkind(cur, legacy) != "r":
        halt(f"{legacy} is not a base table ({relkind(cur, legacy)!r}); "
             "migrate-legacy has run already or the table is missing")
    if write:
        tables = _tables(cur, spec)
        if not all(tables.values()):
            halt("missing " + ", ".join(t for t, ok in tables.items()
                                       if not ok)
                 + "; run `ddl --commit` first")
        for t in (spec.editions_table, spec.live_table,
                  unresolved_table(spec)):
            cur.execute(f"SELECT COUNT(*) FROM public.{t}")
            n = cur.fetchone()[0]
            if n:
                halt(f"{t} already holds {n} rows; migrate-legacy runs once, "
                     "on empty tables")
        if relkind(cur, renamed) is not None:
            halt(f"{renamed} already exists")
    cur.execute(f"SELECT DISTINCT source_file_date FROM public.{legacy} "
                "ORDER BY 1")
    legacy_dates = [pe._p(r[0]) for r in cur.fetchall()]
    runs = legacy_run_dates(cur)
    info = []
    for f in files:
        p = Path(f)
        readme = _read_file(p)
        info.append((readme["as_at"].isoformat(), p))
    info.sort()
    dates = [d for d, _ in info]
    if dates != legacy_dates:
        halt(f"the files are as at {dates}; the legacy table holds "
             f"{legacy_dates}; give exactly one file per legacy snapshot")
    if sorted(runs) != legacy_dates:
        halt(f"the legacy run log names file dates {sorted(runs)}, the "
             f"table {legacy_dates}")
    cur.execute(f"SELECT location_id, reason, resolved_at IS NULL FROM "
                f"public.{legacy_unresolved}")
    recorded = {i: (reason, is_open) for i, reason, is_open in cur.fetchall()}
    open_unresolved = {i for i, (_, o) in recorded.items() if o}
    cur.execute(f"SELECT location_id FROM public.{legacy}")
    in_legacy = {r[0] for r in cur.fetchall()}
    boundaries = load_boundaries(cur)
    recodes = geography.load_recodes(cur)
    lookup = load_lookup(cur)
    api = _api()
    latest = legacy_dates[-1]
    report = []
    prof = _profile(spec)
    for d, path in info:
        try:
            rows, meta = parse_locations(path)
        except ValueError as e:
            halt(f"{e}")
        by_id = {r["location_id"]: r for r in rows}
        residue = legacy_records(cur, d, legacy)
        missing = sorted(r["location_id"] for r in residue
                         if r["location_id"] not in by_id)
        if missing:
            halt(f"{d}: {len(missing)} legacy rows are not in the file "
                 f"{path.name}: {missing[:5]}")
        diffs = []
        for lr in residue:
            fr = by_id[lr["location_id"]]
            for c in PUBLISHER_COLUMNS:
                if fr.get(c) != lr[c]:
                    diffs.append((lr["location_id"], c, lr[c], fr.get(c)))
        if diffs:
            halt(f"{d}: {len(diffs)} cells of the legacy residue differ from "
                 f"the file {path.name} (legacy -> file): {diffs[:5]}; "
                 "nothing stored")
        res_ids = {r["location_id"] for r in residue}
        # locations the retired loader recorded as unresolved (and so never
        # stored) stay unresolved as loaded: re-asking postcodes.io today
        # would put into a past snapshot a row that was not in it
        kept = [r for r in rows if r["location_id"] not in res_ids
                and r["location_id"] in recorded
                and r["location_id"] not in in_legacy]
        kept_ids = {r["location_id"] for r in kept}
        others = [r for r in rows if r["location_id"] not in res_ids
                  and r["location_id"] not in kept_ids]
        try:
            mapped_now, unresolved = map_locations(others, boundaries,
                                                   recodes, lookup, api=api)
            today, _ = map_locations(kept, boundaries, recodes, lookup,
                                     api=api)
        except ValueError as e:
            halt(f"{d}: {e}")
        unresolved = unresolved + [
            _unresolved(r, f"{recorded[r['location_id']][0]} (recorded "
                           f"unresolved in {legacy_unresolved} by the retired "
                           "loader; kept as loaded)") for r in kept]
        if d == latest:
            if mapped_now:
                halt(f"{d}: {len(mapped_now)} in-scope rows not in the legacy "
                     "table now map ("
                     + ", ".join(f"{r['location_id']} {r['mapping_method']}"
                                 for r in mapped_now[:5])
                     + "); the latest snapshot must be the legacy rows as "
                     "loaded")
            got = {u["location_id"] for u in unresolved}
            if got != open_unresolved:
                halt(f"{d}: unresolved {sorted(got)} differ from the open "
                     f"rows of {legacy_unresolved} "
                     f"{sorted(open_unresolved)}")
        loads = runs.get(d, [])
        want = sorted({n for _, n in loads})
        n_rows = len(residue) + len(mapped_now)
        if want != [n_rows]:
            halt(f"{d}: {n_rows} rows rebuilt, the legacy run log loaded "
                 f"{want} for this file; nothing stored")
        for lr in residue:
            fr = by_id[lr["location_id"]]
            fr["lad24cd"] = lr["lad24cd"]
            fr["mapping_method"] = lr["mapping_method"]
        source_file = _file_source(path)
        records = ([{**by_id[i], "snapshot_date": d,
                     "source_file": source_file}
                    for i in sorted(res_ids)]
                   + [{**r, "snapshot_date": d, "source_file": source_file}
                      for r in mapped_now])
        sha = content_sha256(path)
        on = " / ".join(sorted({x for x, _ in loads}))
        if d == latest:
            how = (f"as loaded on {on} (the legacy rows, {len(residue)} "
                   "checked against the file)")
        else:
            how = (f"reconstructed from the file loaded on {on}; "
                   f"{len(residue)} rows checked against the legacy residue; "
                   f"{len(mapped_now)} rows mapped again on {date.today()}; "
                   f"{len(kept)} kept unresolved as recorded")
        methods = Counter(r["mapping_method"] for r in mapped_now)
        entry = {"snapshot": d, "file": path.name, "rows": len(records),
                 "residue": len(residue), "reconstructed": len(mapped_now),
                 "unresolved": sorted(u["location_id"] for u in unresolved),
                 "methods": dict(methods), "label": release_label(
                     d, None, sha, how),
                 "records": records, "unresolved_rows": unresolved,
                 "path": path, "sha": sha, "source_file": source_file,
                 "loaded_on": loads[0][0] if loads else None}
        report.append(entry)
        print(f"  {d}: {path.name}: {meta['in_scope']:,} in scope; "
              f"{len(records):,} rows ({len(residue):,} legacy rows checked, "
              f"0 differences; {len(mapped_now):,} mapped now "
              f"{dict(methods)}); unresolved {len(unresolved)} "
              f"{entry['unresolved']}; run log loaded {want}")
        for r in mapped_now:
            if r.get("mapping_note"):
                print(f"    {r['mapping_method']}: {r['location_id']}: "
                      f"{r['mapping_note']}")
        for r in today:
            print(f"    NOTE: {r['location_id']} kept unresolved as recorded, "
                  f"although it now maps: {r.get('mapping_note') or r['mapping_method']} "
                  f"({r['lad24cd']})")
    if not write:
        print("view check (the view equals the legacy table on every column "
              "but loaded_at; W1's S11 counts unchanged): would be run in the "
              "commit")
        return {"snapshots": report}
    for e in report:
        cur.execute(f"SAVEPOINT {prof.savepoint}")
        kind = apply_snapshot(
            cur, prof, e["snapshot"], e["records"],
            fetched_on=date.fromisoformat(e["loaded_on"] or e["snapshot"]),
            info=FileInfo(e["source_file"], e["path"], e["sha"], e["label"]),
            unresolved=e["unresolved_rows"])
        if kind != "new":
            halt(f"{e['snapshot']}: stored as {kind}, expected new")
        cur.execute(f"RELEASE SAVEPOINT {prof.savepoint}")
    before = w1_counts(cur, legacy)
    cur.execute(f"ALTER TABLE public.{legacy} RENAME TO {renamed}")
    cur.execute(view_sql(legacy, spec.live_table))
    _grant(cur, legacy)
    diff = view_differences(cur, legacy, renamed)
    if diff != (0, 0):
        halt(f"the view {legacy} differs from {renamed}: {diff[0]} rows only "
             f"in the view, {diff[1]} only in the table; rolled back")
    after = w1_counts(cur, legacy)
    if after != before:
        moved = sorted(c for c in set(before) | set(after)
                       if before.get(c) != after.get(c))
        halt(f"W1's S11 counts differ after the swap in {moved[:10]}; "
             "rolled back")
    cur.execute(f"COMMENT ON TABLE public.{legacy_unresolved} IS 'FROZEN "
                f"{date.today()}: no longer written. The unresolved "
                "locations of each CQC snapshot are in "
                f"{unresolved_table(spec)} (append-only, written with the "
                "snapshot by scripts/s11_cqc_editions.py).'")
    cur.execute(f"COMMENT ON VIEW public.{legacy} IS 'Each CQC location''s "
                f"row from the latest snapshot in {spec.live_table} that "
                "holds it; is_active = in the latest snapshot; "
                "deregistered_seen_date = the first snapshot after its last "
                f"appearance. The old table is {renamed} (kept).'")
    print(f"view {legacy} created over {spec.live_table}; equals {renamed} "
          "on every column but loaded_at (0 and 0 rows); W1's S11 counts "
          f"unchanged ({len(after)} authorities)")
    return {"snapshots": report, "w1": after}


def cmd_migrate_legacy(args) -> int:
    writing = args.commit or args.simulate
    conn = _conn(writing)
    started = datetime.now(timezone.utc)
    try:
        with conn.cursor() as cur:
            rep = migrate_legacy(
                cur, args.files, write=writing, spec=SPEC, legacy=LEGACY,
                renamed=LEGACY_RENAMED, legacy_unresolved=LEGACY_UNRESOLVED)
            if args.commit:
                snaps = rep["snapshots"]
                n = sum(len(e["records"]) for e in snaps)
                log_run(cur, 2 * n,  # edition rows + live rows
                        "migrate-legacy: snapshots " + "; ".join(
                            f"{e['snapshot']} {e['rows']} rows "
                            f"({e['residue']} legacy, {e['reconstructed']} "
                            f"reconstructed, {len(e['unresolved'])} "
                            "unresolved)" for e in snaps)
                        + f"; {LEGACY} is now a view, the table is "
                        f"{LEGACY_RENAMED}.", started)
                conn.commit()
                print("migrate-legacy: COMMITTED; pipeline_run_log row "
                      "written")
            else:
                conn.rollback()
                print("PREVIEW: nothing written to the database (use "
                      "--commit or --simulate)" if not writing else
                      "SIMULATION: ROLLED BACK (nothing persisted)")
        return 0
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="S11 CQC Care directory with filters: dated snapshots on "
        "the period-editions engine. There is no sync-new: the legacy table "
        "holds no per-snapshot live rows; migrate-legacy records the held "
        "snapshots once.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ddl", help="create the editions table, live table, "
                       "unresolved table and file-check ledger (preview by "
                       "default)")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_ddl)
    sub.add_parser("status", help="what needs action; exit 1 if anything"
                   ).set_defaults(func=cmd_status)
    p = sub.add_parser("load", help="read the page's filters file (or "
                       "--file), compare and store its snapshot (preview by "
                       "default)")
    p.add_argument("--file", metavar="PATH",
                   help="one local filters file (nothing is downloaded; "
                   "identity and the as-at date are read from the file)")
    p.add_argument("--recheck", type=_iso_day, metavar="PERIOD",
                   help="read the file of this held snapshot again "
                   "(YYYY-MM-DD)")
    p.add_argument("--allow-older-file", action="store_true",
                   help="load a file as at a date earlier than the latest "
                   "held snapshot (a backfill; logged)")
    p.add_argument("--acknowledge", action="append", metavar="PERIOD",
                   type=_iso_day,
                   help="release a 0/NULL change of a recheck after reading "
                   "the preview")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_load)
    p = sub.add_parser("refresh-latest", help="copy each snapshot's latest "
                       "edition into the live table, source_file included "
                       "(preview by default)")
    pe.mode_parser(p)
    p.add_argument("--accept-drift", action="append", metavar="PERIOD",
                   help="overwrite this snapshot (YYYY-MM-DD) although its "
                   "live rows equal no stored edition (repeatable)")
    p.set_defaults(func=cmd_refresh_latest)
    p = sub.add_parser("migrate-legacy", help="one-off: record the held "
                       "snapshots from their files and make cqc_locations a "
                       "view (preview by default)")
    p.add_argument("files", nargs="+", metavar="FILE")
    pe.mode_parser(p)
    p.set_defaults(func=cmd_migrate_legacy)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
