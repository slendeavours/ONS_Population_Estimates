"""Gates for the S24 (RSH register of registered providers, regulatory
judgements and enforcement notices) editions tables.

Mirrors scripts/s23_rsh_stock_editions_verify.py and
s12_financial_stress_editions_verify.py (several specs). It replaces the old
scripts/historical/s24_rsh_register_verify.py (retired 2026-10-11; it stops
with a RETIRED message), whose gate 5 compared the '-' markers of one file
with the NULL grades of the whole mixed table (the defect this build ends).
Prints `GATE n name: PASS|FAIL`; exits 1 on any FAIL. It writes nothing, not
even a run-log row. The real-table gates read rsh_registered_providers,
rsh_regulatory_judgements, rsh_enforcement_notices, their three editions
tables and the three <editions>_file_checks ledgers. Until the editions
tables exist (ddl and migrate-legacy not yet run) the gates that need them
print `FAIL (pending migration)` with the words "not yet migrated", so the
script exits non-zero until then, like the other verifiers; nothing is
created to make them pass.

The seeded gates never need the real editions tables: they run on the
throwaway tables zz_s24_reg / zz_s24_jud / zz_s24_not (_live, _editions and
_editions_file_checks; copies of the loader's three specs), created inside
the transaction, and call the loader's own writer and commands
(s24_rsh_register_editions.main with the SPECS pointed at the copies, the
content API and the downloads stubbed, every workbook written by the test
fixtures with a veryHidden sentinel sheet; migrate_legacy; restore_edition)
rather than a re-implementation. The real tables are only ever read.

Gates:
  1  three editions tables, three ledgers, live shape, append-only triggers
     (UPDATE, DELETE, TRUNCATE refused)               (throwaway + real)
  2  every live register period has edition 1; live judgements and notices
     hold exactly one period, the newest edition's; the latest edition
     equals live on the compared columns                      (seeded + real)
  3  geography absent: no lad24cd, la_code or region column on any S24 table,
     no S24 column in staging_la_signals, DATASET_FORM['24'] stays 'none'
  4  per-row provenance populated, uniform per live period and the tip's
  5  unassessed stored distinctly, per edition: for each judgements edition
     with its file on disk the file's '-' grade cells equal that edition's
     NULL grades, each with a reason; no NULL value in any edition of any
     table without a null_reasons entry; no grade stored as '' or '-'; live
     NULL grades equal the newest edition's; the 30 September split
     206/66/332 and the 12 August 600 (200/74/326)
  6  the hidden sheet: the module's AST opens sheets only by exact title (no
     worksheets iteration, no sheetnames loop, no other zip member), a proxy
     workbook records every access on the fixture and the held files, the
     sentinel reaches no output, record, log or exception text
  7  identity checks halt on seeded mismatches
  8  older-file guard on the page, --file and --file --no-page paths; the
     rank is the file's own, never its URL
  9  stop conditions: each seeded REJECTED, stores nothing, no ledger row,
     partial run-log row
 10  a new period in one transaction (both tables of a judgements file)
 11  ledger rule: a skip needs the (final URL, sha256) pair
 12  preview and simulate write nothing, every writing command
 13  a byte-identical re-read is `unchanged` on every path: the register
     against edition 1 'as loaded', null_reasons included; the judgements and
     notices need edition 2 first (edition 1 marks its empties
     reason_not_recorded_as_loaded), then are unchanged against it
 14  a newer judgements edition replaces live and keeps history
 15  refresh-latest changes only the named live period and copies loaded_at
 16  edition 1 equals the decision note's before-state lines
 17  the migration proofs re-run read-only on the held files
 18  rerun idempotent
 19  stranded repair
 20  restore-edition round trip, all three tables
 21  no network, no secret in source or output, nothing left committed, a
     download never overwrites a same-named file with different content
 22  a notices-free judgements issue (Scott's Option C, 2026-10-12): REJECTED
     without --acknowledge-no-notices; with it the judgements load and the
     notices period is recorded empty (one ledger row, no editions row);
     refresh-latest --accept-key-changes PERIOD empties live notices, the
     history stays; a re-read is skipped; a later issue restores live; a
     missing or changed header still halts

The decision note's hash lines are scan-safe: a 64-hex run would be flagged
by the credential scan, so a sha256 is written as two 32-hex halves in two
labelled fields:
    before-state <table> <period> rows=<n> sha256-first32=<32 hex> sha256-last32=<32 hex>
    before-state-all <table> <period> rows=<n> sha256-first32=<32 hex> sha256-last32=<32 hex>
<table> is the live table (rsh_registered_providers,
rsh_regulatory_judgements, rsh_enforcement_notices); <period> a live period
before the migration. A before-state line hashes that period's live rows on
the key and the value columns (NULL as '', dates ISO; no null_reasons, which
live has not: m.rows_content_sha(reasons=False)); before-state-all hashes
every stored column but loaded_at (key, period, value columns and the three
provenance columns). Edition 1 'as loaded' must hash to the same: the
register and the 30 September judgements and notices equal live; the 15 live
judgements still tagged 12 August are the rows of the 12 August file (edition
1, 308 rows) whose key is not a key of the 30 September edition 1, so that
line is checked against exactly those rows; the 12 August notices had no live
rows and need no line (the migration proof, gate 17, covers them).
`python scripts/s24_rsh_register_editions_verify.py --print-note-lines`
prints the lines for the live tables as they stand (read-only); run it before
migrate-legacy and paste the lines into the decision note.

Usage:
    python scripts/s24_rsh_register_editions_verify.py
    python scripts/s24_rsh_register_editions_verify.py --print-note-lines

Everything runs in one transaction that ends in a rollback; each seeded step
also runs in a savepoint that is rolled back. The cursor handed to the gates
refuses commit and rollback on its connection, and the commands get a
stand-in connection whose commit only counts. No network, no real-table
writes, no backend is ever terminated. After the rollback a second,
read-only connection checks that no zz% table or function was left behind.
"""
import ast
import contextlib
import dataclasses
import hashlib
import io
import os
import re
import socket
import sys
import tempfile
import unittest
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from unittest import mock

import psycopg2

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE / "tests"))
from _db import ENV, get_conn, get_readonly_conn  # noqa: E402
import editions_core as core  # noqa: E402
import geography  # noqa: E402
import period_editions as pe  # noqa: E402
import s24_rsh_register_editions as m  # noqa: E402
import test_s24_rsh_loader as tl  # noqa: E402
from test_s24_rsh_pure import (HIDDEN_TITLE, JROWS, MEDIA, NROWS,  # noqa: E402
                               REG_ROWS, SENTINEL, jrow, judgements_page,
                               register_page, write_book, write_judgements,
                               write_register)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# the real specs and constants, taken before anything is patched; the
# real-table gates use only these (never m.SPECS, which the seeded gates point
# at the copies)
REAL = dict(m.SPECS)
ZZ = tl.ZZ
NAMES = ("register", "judgements", "notices")
LOADER = HERE / "s24_rsh_register_editions.py"
NOTE = (HERE.parent / "docs" / "decisions"
        / "2026-10-11-s24-editions-first-load.md")
NOT_YET = "the S24 editions tables do not exist yet (not yet migrated)"
RESULTS = []
OUTPUT = []  # every line the gates print, for the secret-leak gate (21)
EMPTY = "no live periods to check (an empty state is not a pass)"
JUL, AUG, SEP = tl.JUL, tl.AUG, tl.SEP
J_AUG, J_SEP = tl.J_AUG, tl.J_SEP
CORR = "s24-aug-notes-none-2026-10"
# the real split of the NULL grades by reason (Task 2, counted from the two
# held files; the sum is the file's '-' grade cells): (not_applicable_local_
# authority, not_assessed_yet, marked_dash_no_reason)
EXPECTED_SPLIT = {"2026-09-30": (206, 66, 332), "2026-08-12": (200, 74, 326)}
KNOWN_REASONS = frozenset((
    m.NOT_APPLICABLE_LA, m.NOT_ASSESSED_YET, m.MARKED_DASH_NO_REASON,
    m.MARKED_DASH, m.BLANK, m.PUBLISHER_TEXT_NONE, m.STORED_TEXT_NONE,
    m.NOT_RECORDED))
GRADE_HEADERS = {"consumer_grade": "Consumer grade",
                 "governance_grade": "Governance Grade",
                 "viability_grade": "Viability Grade", "rent_grade": "Rent"}


def report(n, name, ok, detail="", pending=False):
    RESULTS.append(ok)
    verdict = "PASS" if ok else ("FAIL (pending migration)" if pending
                                 else "FAIL")
    line = f"GATE {n} {name}: {verdict}" + (f"  [{detail}]" if detail else "")
    OUTPUT.append(line)
    print(line)


def _first(e):
    return str(e).splitlines()[0] if str(e) else type(e).__name__


class _NoCommit:
    """The .connection of a gate cursor: it has an encoding (psycopg2's
    helpers read it) but commit and rollback are refused, so nothing run
    from a gate can reach the real transaction."""

    def __init__(self, conn):
        self.encoding = conn.encoding

    def commit(self):
        raise RuntimeError("commit refused by the verify script")

    def rollback(self):
        raise RuntimeError("rollback refused by the verify script")


class SafeCur:
    def __init__(self, cur):
        self._cur = cur
        self.connection = _NoCommit(cur.connection)

    def __getattr__(self, name):
        return getattr(self._cur, name)


def _in_savepoint(cur, fn):
    cur.execute("SAVEPOINT h")
    try:
        return fn(cur)
    finally:
        cur.execute("ROLLBACK TO SAVEPOINT h")
        cur.execute("RELEASE SAVEPOINT h")


def ledger_name(spec):
    return f"{spec.editions_table}_file_checks"


def exists(cur, specs=None):
    specs = specs or REAL
    return all(pe.table_exists(cur, t) for s in specs.values()
               for t in (s.editions_table, ledger_name(s)))


def _need(cur, specs):
    miss = [t for s in specs.values()
            for t in (s.editions_table, ledger_name(s))
            if not pe.table_exists(cur, t)]
    if miss:
        raise ValueError(f"{', '.join(miss)} does not exist")


def cnt(cur, table, where="", args=None):
    cur.execute(f"SELECT COUNT(*) FROM public.{table} {where}", args)
    return cur.fetchone()[0]


def _live_periods(cur, spec):
    cur.execute(f"SELECT DISTINCT {spec.period_col} FROM "
                f"public.{spec.live_table} ORDER BY 1")
    return [str(r[0]) for r in cur.fetchall()]


def _edition_periods(cur, spec):
    cur.execute(f"SELECT DISTINCT {spec.period_col} FROM "
                f"public.{spec.editions_table} ORDER BY 1")
    return [str(r[0]) for r in cur.fetchall()]


def empty_latest(cur, spec):
    """The period recorded empty that live holds as no rows (the notices
    after a notices-free judgements issue, Scott's Option C of 2026-10-12):
    the newest empty period, newer than every editions period, with live
    empty. None otherwise (also for a spec that does not opt in), so an
    empty live table without that record is still not a pass."""
    es = core.empty_state(cur, spec)
    if es is None or es["latest"] is None or es["live"]:
        return None
    return pe._p(es["latest"])


def empty_pending(cur, spec):
    """The period recorded empty whose empty state live does not hold yet
    (refresh-latest --accept-key-changes PERIOD not run), else None."""
    es = core.empty_state(cur, spec)
    if es is None or es["pending"] is None:
        return None
    return pe._p(es["pending"])


# ------------------------------------------------------------ throwaway data

def setup_throwaway(cur):
    names = tl.tables()
    cur.execute("SELECT " + ", ".join(f"to_regclass('public.{n}')"
                                      for n in names))
    if any(cur.fetchone()):
        sys.exit("HARD STOP: a zz_s24 table already exists as a real table; "
                 "refusing to run")
    for name, s in ZZ.items():
        cur.execute(f"CREATE TABLE public.{s.live_table} (LIKE "
                    f"public.{REAL[name].live_table} INCLUDING ALL)")
    with tl.specs():
        m.create_all(cur)


def state(cur):
    """Everything a refused or previewed run must leave alone: the live and
    editions tables and the ledgers of all three parts."""
    out = []
    for t in tl.tables():
        cur.execute(f"SELECT md5(string_agg(t::text, ',' ORDER BY t::text)) "
                    f"FROM public.{t} t")
        out.append(cur.fetchone()[0])
    return tuple(out)


def tables_state(cur):
    """state() without the three ledgers (a re-read adds 'unchanged' rows)."""
    return tuple(x for i, x in enumerate(state(cur)) if i % 3 != 2)


def counts_of(cur, names, period):
    """(live, editions, ledger) row counts of a period over `names`."""
    out = [0, 0, 0]
    for n in names:
        s = ZZ[n]
        w = f"WHERE {s.period_col} = %s"
        for i, t in enumerate((s.live_table, s.editions_table,
                               ledger_name(s))):
            out[i] += cnt(cur, t, w, (period,))
    return tuple(out)


# ------------------------------------------------------------ files and Env

class Env(tl.Held):
    """The loader's main(argv) on the throwaway tables, the content API and
    the downloads stubbed; every workbook is written here by the loader
    tests' fixtures."""

    def __init__(self, cur, root):
        unittest.TestCase.__init__(self)
        self.cur = cur
        self.root = Path(root)
        self.api, self.urls, self.final, self.fetched = {}, {}, {}, []
        self.n = 0

    def run_main(self, cur, argv):
        res = super().run_main(cur, argv)
        OUTPUT.extend(res[1].splitlines())
        return res

    def cmd(self, argv):
        return self.run_main(self.cur, argv)

    def ok(self, cur, argv):
        rc, text, logged = self.run_main(cur, argv)
        if rc != 0:
            raise AssertionError(f"{' '.join(map(str, argv[:4]))}: rc {rc}: "
                                 f"{text[-300:]}")
        return text, logged

    def halts(self, argv, before, needle=None, rc_ok=("halt", 1)):
        """(True, text) if the command halted or exited 1 and left
        everything as `before`, with no run-log row."""
        rc, text, logged = self.cmd(argv)
        ok = (rc in rc_ok and state(self.cur) == before
              and not logged.called
              and (needle is None or needle in text))
        return ok, text

    # -- worlds ---------------------------------------------------------
    def sep_page(self):
        """The register page offering the held 17 September file."""
        url = self.urls_of[("register", SEP)]
        self.api[m.REGISTER_PATH] = register_page("17 September 2026", url)
        self.urls[url] = self.paths[("register", SEP)]

    def jsep_page(self):
        """The judgements page offering the held 30 September file."""
        url = self.urls_of[("judgements", J_SEP)]
        self.api[m.JUDGEMENTS_PATH] = judgements_page(url)
        self.urls[url] = self.paths[("judgements", J_SEP)]

    def pages(self):
        self.sep_page()
        self.jsep_page()

    def migrated(self, plant=False):
        self.seed(self.cur, plant=plant)
        self.migrate(self.cur)

    @contextmanager
    def correction(self, cells=4):
        c = dict(m.CORRECTIONS[CORR], period=AUG, cells=cells)
        with mock.patch.object(m, "CORRECTIONS", {CORR: c}):
            yield

    def completed(self):
        """The state the real run ends in: migrated, edition 2 of both
        judgements editions (and their notices), the August notes corrected,
        refresh-latest done."""
        self.migrated()
        for p in (J_AUG, J_SEP):
            self.ok(self.cur, ["load", "--only", "judgements", "--file",
                               self.paths[("judgements", p)], "--no-page",
                               "--commit", "--acknowledge", p])
        self.pages()
        with self.correction():
            self.ok(self.cur, ["load", "--only", "register", "--correction",
                               CORR, "--acknowledge", AUG, "--commit"])
        self.ok(self.cur, ["refresh-latest", "--accept-key-changes", J_SEP,
                           "--commit"])

    @contextmanager
    def legacy_cli(self):
        """What the migrate-legacy command needs to see the seeded legacy
        state as the surveyed one: the held files in a RAW_DIR, the
        surveyed live state, the fixture's run-log count."""
        raw = self.root / "raw"
        raw.mkdir(exist_ok=True)
        for k, p in self.paths.items():
            (raw / self.held[k]["name"]).write_bytes(p.read_bytes())
        with mock.patch.object(m, "RAW_DIR", raw),                 mock.patch.object(m, "LEGACY_LIVE", self.legacy),                 mock.patch.object(m, "LEGACY_FILES", self.held),                 mock.patch.object(m, "AUGUST_REGISTER",
                                  dict(m.AUGUST_REGISTER, period=AUG)),                 mock.patch.object(m, "LEGACY_RUN",
                                  dict(m.LEGACY_RUN, rows=self.run_rows,
                                       register=JUL, judgements=J_AUG)),                 mock.patch.object(m, "legacy_run_rows",
                                  return_value=self.run_rows):
            yield

    def load_world(self):
        """A plain loaded world (no migration): a register and a judgements
        file loaded by the loader's own load."""
        self.register()
        self.judgements()
        self.ok(self.cur, ["load", "--commit"])


@contextmanager
def env(cur):
    with tempfile.TemporaryDirectory() as tmp:
        yield Env(cur, tmp)


def scenario(cur, fn):
    """Run fn(env) in a savepoint that is rolled back."""
    def body(c):
        with env(c) as e, tl.specs():
            return fn(e)
    return _in_savepoint(cur, body)


def mixed(n, name, seeded_ok, seeded_detail, cur, real_fn):
    """A gate with a seeded part (always) and a real part (once the editions
    tables exist, else pending)."""
    if not seeded_ok:
        return report(n, name, False, f"seeded: {seeded_detail}")
    if not exists(cur):
        return report(n, name, False, f"seeded part passes; real: {NOT_YET}",
                      pending=True)
    try:
        ok, detail = real_fn(cur)
    except (psycopg2.Error, SystemExit, ValueError, LookupError,
            AssertionError) as e:
        return report(n, name, False, _first(e))
    report(n, name, ok, f"{seeded_detail}; real: {detail}")


def rejects(e, argv, needle, names, period, *, commit=True):
    """None if the command was REJECTED (exit 1) naming `needle`, stored
    nothing for the period, wrote no ledger row for it, and (with --commit)
    wrote one partial run-log row; else the problem."""
    before = counts_of(e.cur, names, period)
    rc, text, logged = e.cmd(argv)
    if rc != 1 or "REJECTED" not in text or needle not in text:
        return f"rc {rc}, not REJECTED naming {needle!r}: {text[-160:]}"
    if counts_of(e.cur, names, period) != before:
        return "something was stored or a ledger row written"
    if commit:
        if logged.call_count != 1:
            return "no partial run-log row"
        notes = logged.call_args[0][2]
        if "PARTIAL RUN (exit 1)" not in notes or \
                f"rejected ['" not in notes:
            return f"run-log notes: {notes[:140]}"
    elif logged.called:
        return "a preview wrote a run-log row"
    return None


# ---------------------------------------------------------------- gate 1

_DTYPE = {"integer": "integer", "varchar": "character varying",
          "text": "text", "date": "date", "jsonb": "jsonb",
          "bigint": "bigint"}


def _family(sql):
    return _DTYPE[sql.split()[0].split("(")[0].lower()]


def _columns(cur, table):
    cur.execute("""SELECT column_name, data_type, is_nullable
                   FROM information_schema.columns
                   WHERE table_schema='public' AND table_name=%s""", (table,))
    return {r[0]: r[1:] for r in cur.fetchall()}


def _pk(cur, table):
    cur.execute("""SELECT a.attname FROM pg_index i
                   JOIN pg_attribute a ON a.attrelid = i.indrelid
                    AND a.attnum = ANY(i.indkey)
                   WHERE i.indrelid = %s::regclass AND i.indisprimary""",
                (f"public.{table}",))
    return {r[0] for r in cur.fetchall()}


def _triggers(cur, table, row_trigger, trunc_trigger):
    cur.execute("""SELECT t.tgname, t.tgtype FROM pg_trigger t
                   WHERE t.tgrelid = %s::regclass AND NOT t.tgisinternal""",
                (f"public.{table}",))
    trg = dict(cur.fetchall())
    # tgtype bits: 1 row, 2 before, 4 insert, 8 delete, 16 update, 32 truncate
    ud, tr = trg.get(row_trigger, 0), trg.get(trunc_trigger, 0)
    if not (ud & 1 and ud & 2 and ud & 8 and ud & 16 and not ud & 4
            and tr & 2 and tr & 32):
        return [f"append-only triggers wrong on {table}: {sorted(trg)}"]
    return []


def _constraint_names(cur, table, kind):
    cur.execute("SELECT conname FROM pg_constraint WHERE conrelid = "
                "%s::regclass AND contype = %s", (f"public.{table}", kind))
    return {r[0] for r in cur.fetchall()}


def empty_ledger_problems(cur, spec):
    """An empty_periods_allowed spec's ledger must take an 'empty' row (the
    acknowledgement column, 'empty' in the outcome check, the empty-row
    check that wants no edition and a named acknowledgement); [] for any
    other spec."""
    if not spec.empty_periods_allowed:
        return []
    led = ledger_name(spec)
    lc = _columns(cur, led)
    bad = []
    if tuple(lc.get("acknowledgement") or ()) != ("text", "YES"):
        bad.append(f"{led}: acknowledgement is {lc.get('acknowledgement')}, "
                   "expected ('text', 'YES')")
    defs = pe._check_defs(cur, led)
    if "'empty'" not in defs.get(f"{led}_outcome_check", ""):
        bad.append(f"{led}: the outcome check does not allow 'empty'")
    chk = defs.get(f"{led}_empty_chk", "")
    if not ("edition IS NULL" in chk and "acknowledgement IS NOT NULL" in chk):
        bad.append(f"{led}: empty-row check is {chk!r}")
    return bad


def shape_problems(cur, spec, *, empty_required=True):
    """Problems with the editions table, the live table and the ledger of a
    spec (empty_required: also the 'empty' row of an opted-in ledger)."""
    bad = []
    led = ledger_name(spec)
    for t in (spec.editions_table, spec.live_table, led):
        if not pe.table_exists(cur, t):
            bad.append(f"{t} does not exist")
    if bad:
        return bad
    meta = {"edition": ("integer", "NO"), "supersedes": ("integer", "YES"),
            "release_label": ("text", None), "published_date": ("date", None),
            "source_file": ("text", None), "source_sha256": ("text", None),
            "loaded_at": ("timestamp with time zone", None)}
    cols, lcols = _columns(cur, spec.editions_table), _columns(
        cur, spec.live_table)
    need = dict(meta)
    for c, ty in tuple(spec.key_types):
        need[c] = (_family(ty), "NO")
    for c, ty in spec.value_cols + spec.extra_cols:
        need[c] = (_family(ty), "NO" if "NOT NULL" in ty.upper() else "YES")
    for c, (ty, nullable) in need.items():
        got = cols.get(c)
        if got is None:
            bad.append(f"{spec.editions_table}: {c} missing")
        elif got[0] != ty or (nullable and got[1] != nullable):
            bad.append(f"{spec.editions_table}: {c} is {got}, expected "
                       f"{ty} {nullable}")
    if _pk(cur, spec.editions_table) != set(spec.key_cols) | {
            spec.period_col, "edition"}:
        bad.append(f"{spec.editions_table}: primary key is not key + "
                   "period + edition")
    if _constraint_names(cur, spec.editions_table, "f"):
        bad.append(f"{spec.editions_table}: a foreign key (S24 has no "
                   "geography to reference)")
    have = _constraint_names(cur, spec.editions_table, "c")
    for c in spec.table_constraints:
        name = c.split()[1]
        if name not in have:
            bad.append(f"{spec.editions_table}: CHECK {name} missing")
    bad += _triggers(cur, spec.editions_table, spec.trigger,
                     spec.truncate_trigger)
    kind = m._spec_kind(spec)
    keys, pc, values = m.KINDS[kind]
    for c in keys + (pc,) + values + m.LIVE_PROVENANCE + ("loaded_at",):
        if c not in lcols:
            bad.append(f"{spec.live_table}: {c} missing")
    if "null_reasons" in lcols:
        bad.append(f"{spec.live_table}: has a null_reasons column (live "
                   "keeps none; the editions table does)")
    lc = _columns(cur, led)
    for c, ty, nullable in (("id", "bigint", "NO"), (pc, "date", "NO"),
                            ("source_file", "text", "NO"),
                            ("file_sha256", "text", "NO"),
                            ("outcome", "text", "NO"),
                            ("edition", "integer", "YES"),
                            ("checked_at", "timestamp with time zone", "NO")):
        got = lc.get(c)
        if got is None or tuple(got) != (ty, nullable):
            bad.append(f"{led}: {c} is {got}, expected {(ty, nullable)}")
    cur.execute("""SELECT pg_get_constraintdef(oid) FROM pg_constraint
                   WHERE conrelid = %s::regclass AND contype = 'c'""",
                (f"public.{led}",))
    chk = " ".join(r[0] for r in cur.fetchall())
    if not all(f"'{o}'" in chk for o in pe.FILE_CHECK_OUTCOMES):
        bad.append(f"{led}: outcome check is {chk!r}")
    bad += _triggers(cur, led, f"{spec.name}_file_checks_immutable",
                     f"{spec.name}_file_checks_no_truncate")
    if empty_required:
        bad += empty_ledger_problems(cur, spec)
    return bad


def immutability(cur, stmts, seed_fn):
    """{label: error text or None}; all in savepoints."""
    def body(cur):
        seed_fn(cur)
        out = {}
        for k, stmt in stmts.items():
            cur.execute("SAVEPOINT stmt")
            try:
                cur.execute(stmt)
                out[k] = None
            except psycopg2.Error as e:
                cur.execute("ROLLBACK TO SAVEPOINT stmt")
                out[k] = str(e).splitlines()[0]
        return out
    return _in_savepoint(cur, body)


def gate_1_table_shape_and_immutability(cur):
    name = ("three editions tables, three ledgers, live shape (UPDATE, "
            "DELETE and TRUNCATE are refused on the editions and the "
            "ledgers)")
    bad = [p for s in ZZ.values() for p in shape_problems(cur, s)]

    def seed(cur):
        with env(cur) as e, tl.specs():
            e.load_world()
    stmts = {}
    for n, s in ZZ.items():
        for t in (s.editions_table, ledger_name(s)):
            col = "edition" if t == s.editions_table else "outcome"
            stmts[f"{t} UPDATE"] = f"UPDATE public.{t} SET {col} = {col}"
            stmts[f"{t} DELETE"] = f"DELETE FROM public.{t}"
            stmts[f"{t} TRUNCATE"] = f"TRUNCATE public.{t}"
    for k, msg in immutability(cur, stmts, seed).items():
        if not (msg and "append-only" in msg):
            bad.append(f"{k} not refused ({msg})")
    # ddl brings a ledger made before the empty-period opt-in up to date, in
    # place and idempotently, keeping its rows and its triggers
    def upgrade(cur):
        s = ZZ["notices"]
        old = m.PROFILES["notices"].with_spec(
            dataclasses.replace(s, empty_periods_allowed=False))
        led = ledger_name(s)
        cur.execute(f"DROP TABLE public.{led}")
        pe.create_file_checks(cur, old)
        pe.record_file_check(cur, old, J_SEP, "f (judgements x)", "ab" * 32,
                             "new", 1)
        got = [bool(empty_ledger_problems(cur, s))]
        with tl.specs():
            m.create_all(cur)
            m.create_all(cur)
        got.append(empty_ledger_problems(cur, s))
        got.append(cnt(cur, led))
        got.append(_triggers(cur, led, f"{s.name}_file_checks_immutable",
                             f"{s.name}_file_checks_no_truncate"))
        return got
    old_bad, after, rows, trig = _in_savepoint(cur, upgrade)
    if not old_bad or after or rows != 1 or trig:
        bad.append(f"ddl upgrade of an old notices ledger: before flagged "
                   f"{old_bad}, after {after}, rows {rows}, triggers {trig}")
    scope = "throwaway copies"
    note = ""
    if exists(cur):
        scope += " and real tables"
        bad += [f"real: {p}" for s in REAL.values()
                for p in shape_problems(cur, s, empty_required=False)]
        pend = [p for s in REAL.values() for p in empty_ledger_problems(cur, s)]
        note = ("; NOTE real notices ledger not yet upgraded for an empty "
                "issue (run ddl --commit before a notices-free issue is "
                "loaded; the load rejects it until then)" if pend else
                "; the real notices ledger takes an 'empty' row")
    report(1, name, not bad, "; ".join(bad[:4]) if bad else
           f"checked on the {scope}; UPDATE/DELETE/TRUNCATE raise "
           "append-only on the three editions tables and the three ledgers; "
           "ddl upgrades an old notices ledger in place" + note)


# ---------------------------------------------------------------- gate 2

def real_edition1_and_latest(cur, specs=None):
    """Every live register period has edition 1 and every editions period
    has live rows; judgements and notices live hold exactly one period, the
    newest edition's, which has edition 1, the older periods being history;
    the latest edition equals live on the key, the period and the value
    columns (NULL-safe, both ways), for all three tables. Run after
    refresh-latest: a revision not yet refreshed shows here."""
    specs = specs or REAL
    _need(cur, specs)
    bad, parts = [], []
    for name in NAMES:
        spec = specs[name]
        periods = _live_periods(cur, spec)
        held = _edition_periods(cur, spec)
        if not periods and empty_latest(cur, spec):
            parts.append(f"{spec.live_table} empty: {empty_latest(cur, spec)}"
                         f" recorded empty, {len(held)} held as history")
            continue
        if not periods:
            bad.append(f"{spec.live_table}: {EMPTY}")
            continue
        if empty_pending(cur, spec):
            bad.append(f"{spec.live_table}: {empty_pending(cur, spec)} is "
                       "recorded empty but live still holds "
                       f"{periods} (refresh-latest --accept-key-changes "
                       f"{empty_pending(cur, spec)})")
        for p in periods:
            if not cnt(cur, spec.editions_table,
                       f"WHERE {spec.period_col} = %s AND edition = 1", (p,)):
                bad.append(f"{spec.live_table} {p}: no edition 1")
        if spec.live_newest_period_only:
            if len(periods) != 1:
                bad.append(f"{spec.live_table}: live holds {len(periods)} "
                           f"periods {periods}, the newest edition only "
                           "(refresh-latest --accept-key-changes removes "
                           "the old build's left-over rows)")
            elif periods[0] != held[-1]:
                bad.append(f"{spec.live_table}: live holds {periods[0]}, "
                           f"not the newest edition {held[-1]}")
        else:
            extra = sorted(set(held) - set(periods))
            if extra:
                bad.append(f"{spec.editions_table}: periods with no live "
                           f"rows {extra}")
        bad += m.live_equals_tip(cur, spec, periods)
        parts.append(f"{spec.live_table} {len(periods)} live period(s) of "
                     f"{len(held)} held")
    return not bad, "; ".join(bad[:3]) if bad else (
        "every live period has edition 1 and equals its latest edition ("
        + ", ".join(parts) + ")")


def gate_2_edition1_and_latest(cur):
    name = ("every live register period has edition 1; live judgements and "
            "notices hold the newest edition only; the latest edition "
            "equals live")
    problems = []

    def body(e):
        ok, d = real_edition1_and_latest(e.cur, ZZ)
        if ok or "no live periods" not in d:
            problems.append(f"empty state: {ok}, {d[:80]}")
        e.completed()
        ok, d = real_edition1_and_latest(e.cur, ZZ)
        if not ok:
            problems.append(f"complete state fails: {d}")

        def plant(sql, needle, label):
            e.cur.execute("SAVEPOINT p")
            e.cur.execute(sql)
            ok, d = real_edition1_and_latest(e.cur, ZZ)
            e.cur.execute("ROLLBACK TO SAVEPOINT p")
            if ok or needle not in d:
                problems.append(f"{label} passes or is not named "
                                f"({needle!r}): {d[:90]}")
        plant(f"UPDATE public.{ZZ['register'].live_table} SET designation = "
              "'x' WHERE registration_number = '4001' AND snapshot_date = "
              f"'{SEP}'", "live-only", "a drifted register cell")
        plant(f"UPDATE public.{ZZ['register'].live_table} SET snapshot_date "
              "= '2030-01-01' WHERE snapshot_date = " f"'{JUL}'",
              "no edition 1", "a register period without edition 1")
        plant(f"DELETE FROM public.{ZZ['register'].live_table} WHERE "
              f"snapshot_date = '{JUL}'", "no live rows",
              "an editions period with no live rows")
        plant(f"UPDATE public.{ZZ['judgements'].live_table} SET edition_date "
              f"= '{J_AUG}'", "not the newest edition",
              "live judgements on an older edition")
        plant(f"UPDATE public.{ZZ['notices'].live_table} SET provider_name = "
              "'x'", "edition-only", "a drifted notice")
        # the migrated, not yet refreshed state: two live judgements periods
        e.cur.execute("SAVEPOINT q")
        e.cur.execute(f"INSERT INTO public.{ZZ['judgements'].live_table} "
                      "(registration_number, publication_date, edition_date, "
                      "landlord_name, source_url, source_file, "
                      "release_page_url) VALUES ('L9', '2026-01-01', "
                      f"'{J_AUG}', 'x', 'u', 'f', 'p')")
        ok, d = real_edition1_and_latest(e.cur, ZZ)
        e.cur.execute("ROLLBACK TO SAVEPOINT q")
        if ok or "holds 2 periods" not in d:
            problems.append(f"two live judgements periods pass: {d[:90]}")
    scenario(cur, body)
    mixed(2, name, not problems,
          "an empty state is not a pass; the completed state passes; a "
          "drifted cell, a period without edition 1, an editions period "
          "with no live rows, live on an older edition, a drifted notice and "
          "two live judgements periods fail" if not problems
          else "; ".join(problems[:3]), cur, real_edition1_and_latest)


# ---------------------------------------------------------------- gate 3

GEO_COLUMNS = ("lad24cd", "la_code", "region", "publisher_la_code")
STAGING_LIKE = ("rsh%", "%registered_provider%", "%regulatory_judg%",
                "%enforcement_notice%")
# columns peculiar to the S24 tables (generic names such as status or notes
# are not listed: another table may use them)
STAGING_NAMES = ("registration_number", "organisation_name", "corporate_form",
                 "landlord_name", "provider_name", "consumer_grade",
                 "governance_grade", "viability_grade", "rent_grade",
                 "null_reasons", "file_source_url")


def geography_problems(cur, tables, staging_cols=None):
    """Geography columns on the S24 tables, S24 columns in
    staging_la_signals, a DATASET_FORM['24'] other than 'none'."""
    bad = []
    cur.execute("""SELECT table_name, column_name FROM
                   information_schema.columns WHERE table_schema = 'public'
                   AND table_name = ANY(%s) AND column_name = ANY(%s)
                   ORDER BY 1, 2""", (list(tables), list(GEO_COLUMNS)))
    bad += [f"{t}.{c} is a geography column" for t, c in cur.fetchall()]
    if staging_cols is None:
        cur.execute("""SELECT column_name FROM information_schema.columns
                       WHERE table_name = 'staging_la_signals' AND
                       (column_name LIKE ANY(%s) OR column_name = ANY(%s))""",
                    (list(STAGING_LIKE), list(STAGING_NAMES)))
        staging_cols = [r[0] for r in cur.fetchall()]
    bad += [f"staging_la_signals.{c} is an S24 column (this source is a "
            "watchlist and must not reach W1 or the map)"
            for c in staging_cols]
    form = geography.DATASET_FORM.get("24", ("?",))[0]
    if form != "none":
        bad.append(f"DATASET_FORM['24'] is {form!r}, not 'none'")
    return bad


def real_geography(cur, specs=None):
    specs = specs or REAL
    _need(cur, specs)
    tables = [t for s in specs.values()
              for t in (s.live_table, s.editions_table, ledger_name(s))]
    bad = geography_problems(cur, tables)
    src = LOADER.read_text(encoding="utf-8")
    for needle in ("import geography", "la_code_lookup", "lad24cd"):
        if needle in src:
            bad.append(f"the loader mentions {needle!r}")
    return not bad, "; ".join(bad[:3]) if bad else (
        f"none of {GEO_COLUMNS} on the {len(tables)} S24 tables, no S24 "
        "column in staging_la_signals, DATASET_FORM['24'] is 'none', the "
        "loader names no geography")


def gate_3_geography(cur):
    name = ("geography absent: no lad24cd, la_code or region column on an "
            "S24 table, no S24 column in staging_la_signals, "
            "DATASET_FORM['24'] stays 'none'")
    problems = []

    def body(e):
        tables = tl.tables()
        got = geography_problems(e.cur, tables)
        if got:
            problems.append(f"the throwaway tables: {got[:2]}")
        e.cur.execute(f"ALTER TABLE public.{ZZ['judgements'].live_table} ADD "
                      "COLUMN lad24cd text")
        got = geography_problems(e.cur, tables)
        if not any("lad24cd" in x for x in got):
            problems.append("a lad24cd column passes")
        e.cur.execute(f"ALTER TABLE public.{ZZ['notices'].editions_table} "
                      "ADD COLUMN region text")
        if not any("region" in x for x in geography_problems(e.cur, tables)):
            problems.append("a region column passes")
        if not geography_problems(e.cur, tables[:1],
                                  staging_cols=["rsh_grade"]):
            problems.append("an S24 column in staging_la_signals passes")
        with mock.patch.dict(geography.DATASET_FORM,
                             {"24": ("lad24cd", "x")}):
            if not geography_problems(e.cur, tables[:1], staging_cols=[]):
                problems.append("a DATASET_FORM other than none passes")
    scenario(cur, body)
    mixed(3, name, not problems,
          "no geography column on the throwaway tables; a lad24cd column, a "
          "region column, an S24 column in staging_la_signals and a "
          "DATASET_FORM other than none each fail" if not problems
          else "; ".join(problems[:3]), cur, real_geography)


# ---------------------------------------------------------------- gate 4

def real_provenance(cur, specs=None):
    """Per-row provenance populated in live (source_url, source_file,
    release_page_url) and in every edition (file_source_url, file_name,
    file_release_page_url); live provenance uniform per period and equal to
    the tip's (live column against its file_* column, as refresh-latest's
    own after-check maps them)."""
    specs = specs or REAL
    _need(cur, specs)
    bad, parts = [], []
    for name in NAMES:
        spec = specs[name]
        for t, cols in ((spec.live_table, m.LIVE_PROVENANCE),
                        (spec.editions_table, m.PROVENANCE)):
            w = " OR ".join(f"{c} IS NULL OR {c} = ''" for c in cols)
            n = cnt(cur, t, f"WHERE {w}")
            if n:
                bad.append(f"{t}: {n} row(s) missing a provenance field")
        if not _live_periods(cur, spec) and not empty_latest(cur, spec):
            bad.append(f"{spec.live_table}: {EMPTY}")
        for p in _live_periods(cur, spec):
            cur.execute(f"SELECT DISTINCT {', '.join(m.LIVE_PROVENANCE)} "
                        f"FROM public.{spec.live_table} WHERE "
                        f"{spec.period_col} = %s", (p,))
            live = cur.fetchall()
            if len(live) != 1:
                bad.append(f"{spec.live_table} {p}: provenance is not "
                           f"uniform ({len(live)} values)")
                continue
            try:
                tip = m._tip_provenance(cur, spec, p)
            except SystemExit as e:
                bad.append(f"{spec.editions_table} {p}: {_first(e.code)}")
                continue
            if tuple(live[0]) != tuple(tip[c] for c in m.PROVENANCE):
                bad.append(f"{spec.live_table} {p}: live provenance differs "
                           "from the tip's")
            for c, v in zip(m.LIVE_PROVENANCE, live[0]):
                if c == "source_url" and not str(v).startswith(
                        ("https://", "file:")):
                    bad.append(f"{spec.live_table} {p}: source_url {v!r} is "
                               "neither a URL nor file:<name>")
        parts.append(spec.live_table)
    return not bad, "; ".join(bad[:3]) if bad else (
        "every live and editions row carries its provenance; each live "
        "period's is one value, the tip's (" + ", ".join(parts) + ")")


def gate_4_provenance(cur):
    name = ("per-row provenance populated, uniform per live period and the "
            "tip's")
    problems = []

    def body(e):
        e.completed()
        ok, d = real_provenance(e.cur, ZZ)
        if not ok:
            problems.append(f"complete state fails: {d}")

        def plant(sql, needle, label):
            e.cur.execute("SAVEPOINT p")
            e.cur.execute(sql)
            ok, d = real_provenance(e.cur, ZZ)
            e.cur.execute("ROLLBACK TO SAVEPOINT p")
            if ok or needle not in d:
                problems.append(f"{label} passes or is not named: {d[:90]}")
        live = ZZ["register"].live_table
        plant(f"UPDATE public.{live} SET source_url = '' WHERE "
              "registration_number = '4001'", "missing a provenance",
              "an empty source_url")
        plant(f"UPDATE public.{live} SET source_file = 'other' WHERE "
              f"registration_number = '4001' AND snapshot_date = '{SEP}'",
              "not uniform", "a non-uniform period")
        plant(f"UPDATE public.{ZZ['judgements'].live_table} SET "
              "release_page_url = 'x'", "differs from the tip's",
              "uniform but not the tip's")
        plant(f"UPDATE public.{live} SET source_url = 'ftp://x' WHERE "
              f"snapshot_date = '{SEP}'", "neither a URL", "a bad source_url")
    scenario(cur, body)
    mixed(4, name, not problems,
          "a completed world passes; an empty provenance field, a "
          "non-uniform period, a uniform value that is not the tip's and a "
          "bad source_url fail" if not problems else "; ".join(problems[:3]),
          cur, real_provenance)


# ---------------------------------------------------------------- gate 5

def files_by_sha(dirs):
    """{sha256: path} of every .xlsx under the directories."""
    out = {}
    for d in dirs:
        d = Path(d)
        if d.is_dir():
            for p in sorted(d.rglob("*.xlsx")):
                out.setdefault(m.content_sha256(p), p)
    return out


def file_dashes(path):
    """(edition date, {grade column: '-' cells}) of a judgements file, read
    by the loader's identity function (the visible sheet by its exact
    title)."""
    ident = m.judgements_identity(path, Path(path).name, None)
    idx = {c: m.JUDGEMENT_HEADERS.index(h) for c, h in GRADE_HEADERS.items()}
    out = {c: sum(1 for _, cells in ident["judgement_rows"]
                  if str(cells[i]).strip() == "-")
           for c, i in idx.items()}
    return ident["period"], out


def null_reason_problems(cur, specs):
    """Problems in every edition of the three tables: a NULL value column
    with no null_reasons entry; a reason for a column that is not NULL; a
    reason outside the known set."""
    bad = []
    for name in NAMES:
        spec = specs[name]
        values = m.KINDS[name][2]
        for c in values:
            n = cnt(cur, spec.editions_table,
                    f"WHERE {c} IS NULL AND NOT (null_reasons ? %s)", (c,))
            if n:
                bad.append(f"{spec.editions_table}: {n} NULL {c} value(s) "
                           "without a null_reasons entry")
            n = cnt(cur, spec.editions_table,
                    f"WHERE {c} IS NOT NULL AND (null_reasons ? %s)", (c,))
            if n:
                bad.append(f"{spec.editions_table}: {n} reason(s) for a "
                           f"non-NULL {c}")
        cur.execute(f"SELECT DISTINCT r.reason FROM public."
                    f"{spec.editions_table} e, jsonb_each_text("
                    "e.null_reasons) r(col, reason)")
        odd = sorted(r[0] for r in cur.fetchall()
                     if r[0] not in KNOWN_REASONS)
        if odd:
            bad.append(f"{spec.editions_table}: unknown reason(s) {odd}")
    return bad


def real_unassessed(cur, specs=None, dirs=None, split=None):
    """Unassessed grades stored distinctly, counted per edition. For each
    judgements edition whose file is on disk (found by the sha256 the ledger
    records for that edition): the file's '-' grade cells per column equal
    the edition's NULL grades, and every NULL grade has a reason. In every
    edition of the three tables no NULL value lacks a null_reasons entry. No
    grade is stored as '' or '-' (or outside the allowed set) in an edition
    or in live. Live NULL grades per column equal the live period's tip
    edition's. A period in `split` has, in its newest edition, the real
    reasons in exactly that split (not applicable local authority, not
    assessed yet, marked dash no reason), none 'reason_not_recorded_as_
    loaded'; any other newest edition past edition 1 as loaded carries none
    of those either."""
    specs = specs or REAL
    dirs = [m.RAW_DIR] if dirs is None else dirs
    split = EXPECTED_SPLIT if split is None else split
    _need(cur, specs)
    J = specs["judgements"]
    pc = J.period_col
    bad = null_reason_problems(cur, specs)
    nulls = ", ".join(f"COUNT(*) FILTER (WHERE {c} IS NULL)" for c in m.GRADES)
    cur.execute(f"SELECT {pc}, edition, {nulls} FROM public."
                f"{J.editions_table} GROUP BY 1, 2 ORDER BY 1, 2")
    editions = {(str(p), ed): dict(zip(m.GRADES, ns))
                for p, ed, *ns in cur.fetchall()}
    if not editions:
        bad.append(f"{J.editions_table}: no edition")
    cur.execute(f"SELECT e.{pc}, e.edition, r.col, r.reason, COUNT(*) FROM "
                f"public.{J.editions_table} e, jsonb_each_text(e.null_reasons)"
                " r(col, reason) WHERE r.col = ANY(%s) GROUP BY 1, 2, 3, 4",
                (list(m.GRADES),))
    reasons = {}
    for p, ed, col, why, n in cur.fetchall():
        reasons.setdefault((str(p), ed), {}).setdefault(col, {})[why] = n
    cur.execute(f"SELECT DISTINCT {pc}, edition, release_label FROM public."
                f"{J.editions_table}")
    labels = {(str(p), ed): lab for p, ed, lab in cur.fetchall()}
    cur.execute(f"SELECT {pc}, edition, file_sha256 FROM public."
                f"{ledger_name(J)} WHERE edition IS NOT NULL")
    shas = {}
    for p, ed, s in cur.fetchall():
        shas.setdefault((str(p), ed), set()).add(s)
    files = files_by_sha(dirs)
    checked, unchecked = [], []
    for (p, ed), n_null in sorted(editions.items()):
        per_col = reasons.get((p, ed), {})
        for c in m.GRADES:
            got = sum(per_col.get(c, {}).values())
            if got != n_null[c]:
                bad.append(f"{p} edition {ed}: {n_null[c]} NULL {c} but "
                           f"{got} with a reason")
        found = [files[s] for s in sorted(shas.get((p, ed), ()))
                 if s in files]
        done = False
        for path in found:
            try:
                fp, dashes = file_dashes(path)
            except ValueError as e:
                bad.append(f"{path.name}: {_first(e)}")
                continue
            if fp != p:
                continue
            done = True
            if dashes != n_null:
                bad.append(f"{p} edition {ed}: the file's '-' grade cells "
                           f"{dashes} differ from the NULL grades {n_null}")
        (checked if done else unchecked).append((p, ed, sum(n_null.values())))
    if editions and not checked:
        bad.append("no judgements edition's file was found on disk, so no "
                   "'-' cell was counted against its NULL grades")
    # the newest edition of each period
    tips = {}
    for p, ed in editions:
        tips[p] = max(tips.get(p, ed), ed)
    for p, ed in sorted((p, ed) for p, ed in tips.items()):
        flat = {}
        for col, d in reasons.get((p, ed), {}).items():
            for why, n in d.items():
                flat[why] = flat.get(why, 0) + n
        kept = str(labels.get((p, ed)) or "")
        if flat.get(m.NOT_RECORDED) and not (
                ed == 1 and kept.startswith(core.AS_LOADED_LABEL)) and \
                not kept.startswith("restored from edition"):
            bad.append(f"{p} edition {ed}: still carries "
                       f"{m.NOT_RECORDED} past edition 1 as loaded")
        if p in split:
            want = split[p]
            got = (flat.get(m.NOT_APPLICABLE_LA, 0),
                   flat.get(m.NOT_ASSESSED_YET, 0),
                   flat.get(m.MARKED_DASH_NO_REASON, 0))
            total = sum(editions[(p, ed)].values())
            if flat.get(m.NOT_RECORDED):
                bad.append(f"{p}: the newest edition {ed} has only edition "
                           f"1's {m.NOT_RECORDED}; the re-read edition with "
                           "the real reasons is not stored")
            elif got != want or sum(got) != total:
                bad.append(f"{p} edition {ed}: reasons "
                           f"(not applicable LA, not assessed yet, no "
                           f"reason) {got}, expected {want}; NULL grades "
                           f"{total}")
    for p in split:
        if p not in tips:
            bad.append(f"{p}: no judgements edition stored")
    # no '' / '-' / unknown grade; live NULLs equal the tip's
    for t in (J.live_table, J.editions_table):
        for c in m.GRADES:
            pat = "^(" + m.GRADE_RE[c].pattern + ")$"
            n = cnt(cur, t, f"WHERE {c} IS NOT NULL AND {c} !~ %s", (pat,))
            if n:
                bad.append(f"{t}: {n} {c} value(s) stored as '', '-' or "
                           "outside the allowed set")
    for p in _live_periods(cur, J)[-1:]:   # the newest live period; an
        # older one is the old build's left-over (gate 2 reports it)
        try:
            tip = core.chain_tip(cur, J, p)
        except (LookupError, ValueError) as e:
            bad.append(f"{p}: {_first(e)}")
            continue
        for c in m.GRADES:
            a = cnt(cur, J.live_table,
                    f"WHERE {pc} = %s AND {c} IS NULL", (p,))
            b = cnt(cur, J.editions_table,
                    f"WHERE {pc} = %s AND edition = %s AND {c} IS NULL",
                    (p, tip))
            if a != b:
                bad.append(f"{p}: live has {a} NULL {c}, the newest edition "
                           f"{tip} has {b}")
    pieces = ", ".join(f"{p} ed{ed}={n}" for p, ed, n in checked)
    return not bad, "; ".join(bad[:3]) if bad else (
        f"{len(checked)} edition(s) counted against their files ({pieces}); "
        f"{len(unchecked)} with no file on disk; every NULL in every "
        "edition has a reason; no '-' or '' grade; live NULL grades equal "
        "the newest edition's"
        + ("; newest edition split " + ", ".join(
            f"{p} {split[p]}" for p in sorted(split) if p in tips)
           if split else ""))


def clone_edition(cur, spec, where, **over):
    """Insert a copy of one editions row with SQL literals overridden."""
    cols = list(_columns(cur, spec.editions_table))
    sel = ", ".join(over.get(c, c) for c in cols)
    cur.execute(f"INSERT INTO public.{spec.editions_table} "
                f"({', '.join(cols)}) SELECT {sel} FROM public."
                f"{spec.editions_table} WHERE {where} LIMIT 1")


def held_split_problems(files=None, dirs=None, split=None):
    """The expected split constants checked against the held judgements
    files read now (no table needed): their '-' grade cells equal the NULL
    grades of the new parse, each with a reason, in exactly the split."""
    files = m.LEGACY_FILES if files is None else files
    split = EXPECTED_SPLIT if split is None else split
    found = files_by_sha([m.RAW_DIR] if dirs is None else dirs)
    bad, parts = [], []
    for (part, p), facts in sorted(files.items()):
        if part != "judgements" or p not in split:
            continue
        path = found.get(facts["sha256"])
        if path is None:
            bad.append(f"judgements {p}: no file with sha256 "
                       f"{facts['sha256'][:16]} on disk")
            continue
        try:
            ident = m.judgements_identity(path, facts["name"], None)
            recs = m.build_judgements(ident["judgement_rows"], ident)
            dashes = sum(file_dashes(path)[1].values())
        except ValueError as e:
            bad.append(f"judgements {p}: {_first(e)}")
            continue
        flat = {}
        for r in recs:
            for c in m.GRADES:
                if r.get(c) is None:
                    why = r["null_reasons"].get(c)
                    flat[why] = flat.get(why, 0) + 1
        got = (flat.pop(m.NOT_APPLICABLE_LA, 0),
               flat.pop(m.NOT_ASSESSED_YET, 0),
               flat.pop(m.MARKED_DASH_NO_REASON, 0))
        if flat or got != split[p] or sum(got) != dashes:
            bad.append(f"judgements {p}: {dashes} '-' grade cells, new-parse "
                       f"reasons {got} (+ other {flat}), expected {split[p]}")
        else:
            parts.append(f"{p}: {dashes} '-' cells = {got[0]} + {got[1]} + "
                         f"{got[2]}")
    return bad, parts


def gate_5_unassessed(cur):
    name = ("unassessed stored distinctly, counted per edition: the file's "
            "'-' grade cells equal that edition's NULL grades, each with a "
            "reason; no NULL without a reason")
    problems = []
    SPLIT = (2, 1, 8)    # the fixtures' LA / not assessed / no reason cells

    def migrated(e):
        e.migrated()
        ok, d = real_unassessed(e.cur, ZZ, [e.root], split={})
        if not ok or "2 edition(s) counted" not in d:
            problems.append(f"migrated state: {ok} {d[:100]}")
        ok, d = real_unassessed(e.cur, ZZ, [e.root],
                                split={J_SEP: SPLIT, J_AUG: SPLIT})
        if ok or "only edition 1's" not in d:
            problems.append("edition 1 alone passes the real-reasons split")

    def complete(e):
        e.completed()
        both = {J_SEP: SPLIT, J_AUG: SPLIT}
        ok, d = real_unassessed(e.cur, ZZ, [e.root], split=both)
        if not ok or "4 edition(s) counted" not in d:
            problems.append(f"completed state: {ok} {d[:120]}")
        ok, d = real_unassessed(e.cur, ZZ, [e.root],
                                split={J_SEP: (206, 66, 332)})
        if ok or "expected (206, 66, 332)" not in d:
            problems.append(f"a wrong split passes: {d[:100]}")
        ok, d = real_unassessed(e.cur, ZZ, [], split=both)
        if ok or "no judgements edition's file" not in d:
            problems.append(f"no file on disk passes: {d[:100]}")

        def plant(label, fn, needle):
            e.cur.execute("SAVEPOINT p")
            fn()
            ok, d = real_unassessed(e.cur, ZZ, [e.root], split={})
            e.cur.execute("ROLLBACK TO SAVEPOINT p")
            if ok or needle not in d:
                problems.append(f"{label} passes or is not named "
                                f"({needle!r}): {d[:100]}")
        J = ZZ["judgements"]
        plant("a grade stored as '-'", lambda: clone_edition(
            e.cur, J, "registration_number = 'LA01' AND edition = 2",
            edition="9", consumer_grade="'-'"), "outside the allowed set")
        plant("a grade stored as ''", lambda: clone_edition(
            e.cur, J, "registration_number = 'LA01' AND edition = 2",
            edition="9", consumer_grade="''"), "outside the allowed set")
        plant("a NULL grade with no reason", lambda: clone_edition(
            e.cur, J, "registration_number = 'LA01' AND edition = 2",
            edition="9", consumer_grade="NULL",
            null_reasons="'{}'::jsonb"), "without a null_reasons entry")
        plant("a NULL register value with no reason", lambda: clone_edition(
            e.cur, ZZ["register"], "registration_number = '4001' AND "
            "edition = 1", edition="9", designation="NULL",
            null_reasons="'{}'::jsonb"), "without a null_reasons entry")
        plant("an unknown reason", lambda: clone_edition(
            e.cur, ZZ["notices"], "edition = 1", edition="9",
            publication_type="NULL",
            null_reasons='\'{"publication_type": "whatever"}\'::jsonb'),
            "unknown reason")
        plant("live NULL grades that are not the tip's", lambda:
              e.cur.execute(f"UPDATE public.{J.live_table} SET rent_grade = "
                            "'Y' WHERE rent_grade IS NULL"),
              "live has")
        plant("live with a '-' grade", lambda:
              e.cur.execute(f"UPDATE public.{J.live_table} SET consumer_grade"
                            " = '-' WHERE consumer_grade IS NULL"),
              "outside the allowed set")

    def mutated(label, patch, needle):
        def body(e):
            e.migrated()
            orig = m.grade_cell
            before = state(e.cur)
            try:
                with mock.patch.object(
                        m, "grade_cell",
                        side_effect=lambda *a, **k: patch(orig, *a, **k)):
                    e.ok(e.cur, ["load", "--only", "judgements", "--file",
                                 e.paths[("judgements", J_SEP)], "--no-page",
                                 "--commit", "--acknowledge", J_SEP])
            except (KeyError, AssertionError, SystemExit):
                # the loader's own summary refused the NULL with no reason:
                # nothing may have been stored
                if state(e.cur) != before:
                    problems.append(f"{label}: the loader refused but left "
                                    "something stored")
                return
            ok, d = real_unassessed(e.cur, ZZ, [e.root], split={})
            if ok or needle not in d:
                problems.append(f"{label} passes or is not named "
                                f"({needle!r}): {d[:100]}")
        scenario(cur, body)

    def drop_reason(orig, value, column, lt, change, where=""):
        v, why = orig(value, column, lt, change, where)
        return (v, None) if v is None else (v, why)

    def lose_dash(orig, value, column, lt, change, where=""):
        if column == "consumer_grade" and str(value).strip() == "-":
            return "C3", None
        return orig(value, column, lt, change, where)
    scenario(cur, migrated)
    scenario(cur, complete)
    held_bad, held_parts = held_split_problems()
    problems += [f"held files: {x}" for x in held_bad]
    mutated("a grade_cell that drops the reason", drop_reason,
            "with a reason")
    mutated("a grade_cell that turns a '-' into a grade", lose_dash,
            "differ from the NULL grades")
    mixed(5, name, not problems,
          "the held files' '-' cells (" + "; ".join(held_parts) + "); "
          "edition 1 as loaded and the re-read edition 2 of both judgements "
          "files each count their file's '-' cells against their NULL "
          "grades, every one with a reason; a '-' or '' grade, a NULL "
          "without a reason, an unknown reason, a drifted live count, a "
          "wrong split, a missing file, a grade_cell that drops the reason "
          "and one that turns a '-' into a grade each fail"
          if not problems else "; ".join(problems[:3]), cur,
          real_unassessed)


# ---------------------------------------------------------------- gate 6

SHEET_ATTRS = {"worksheets", "sheetnames", "active", "get_sheet_by_name",
               "get_sheet_names", "chartsheets", "_sheets", "worksheet",
               "get_active_sheet"}
ZIP_ALLOWED = {"xl/workbook.xml", "docProps/core.xml"}
ZIP_NAMES = {"z", "zf", "zipf", "zip_file", "zfile"}
BOOK_NAMES = {"wb", "book", "workbook"}


def sheet_access_problems(source=None):
    """Problems in the loader source (AST): any way of reaching a sheet
    other than wb[title] inside _read_sheet (the only place a workbook is
    opened), any iteration of a workbook, any zip member read other than
    xl/workbook.xml and docProps/core.xml."""
    src = LOADER.read_text(encoding="utf-8") if source is None else source
    tree = ast.parse(src)
    bad = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Attribute) and node.attr in SHEET_ATTRS:
            bad.append(f"uses .{node.attr} (line {node.lineno})")
        if isinstance(node, (ast.For, ast.comprehension)):
            it = node.iter
            while isinstance(it, ast.Call) and it.args:
                it = it.args[0]
            if isinstance(it, ast.Name) and it.id in BOOK_NAMES:
                bad.append("iterates a workbook")
        if isinstance(node, ast.Call) and isinstance(node.func,
                                                     ast.Attribute) \
                and node.func.attr == "read" \
                and isinstance(node.func.value, ast.Name) \
                and node.func.value.id in ZIP_NAMES and node.args:
            a = node.args[0]
            if not (isinstance(a, ast.Constant) and isinstance(a.value, str)):
                bad.append("reads a zip member by a computed name")
            elif a.value not in ZIP_ALLOWED:
                bad.append(f"reads the zip member {a.value!r}")
    loaders = set()
    subs = []
    for f in (n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef)):
        for n in ast.walk(f):
            if isinstance(n, ast.Call):
                fn = n.func
                name = (fn.attr if isinstance(fn, ast.Attribute) else
                        fn.id if isinstance(fn, ast.Name) else None)
                if name == "load_workbook":
                    loaders.add(f.name)
            if isinstance(n, ast.Subscript) and isinstance(n.value, ast.Name) \
                    and n.value.id in BOOK_NAMES:
                subs.append((f.name, n.slice))
    if loaders != {"_read_sheet"}:
        bad.append(f"load_workbook is called in {sorted(loaders)}, only "
                   "_read_sheet may open a workbook")
    if not subs:
        bad.append("no wb[title] access found in _read_sheet")
    for fname, sl in subs:
        if not (isinstance(sl, ast.Name) and sl.id == "title"):
            bad.append(f"{fname}: a sheet is opened by something other than "
                       "the title parameter")
    return sorted(set(bad))


@contextmanager
def watched_workbooks():
    """openpyxl.load_workbook returning a proxy that records every sheet
    opened by title and refuses (and records) any other access."""
    import openpyxl
    log = {"opened": [], "violations": []}
    real = openpyxl.load_workbook

    class Proxy:
        def __init__(self, wb):
            self._wb = wb

        def __getitem__(self, title):
            log["opened"].append(title)
            return self._wb[title]

        def close(self):
            self._wb.close()

        def __iter__(self):
            log["violations"].append("__iter__")
            raise AssertionError("a workbook was iterated")

        def __getattr__(self, name):
            log["violations"].append(name)
            raise AssertionError(f"workbook attribute {name} touched")

    with mock.patch.object(openpyxl, "load_workbook",
                           lambda *a, **k: Proxy(real(*a, **k))):
        yield log


def non_visible_sheets(path):
    """How many sheets of a workbook are hidden or veryHidden (counted from
    xl/workbook.xml; no name is returned or printed)."""
    import zipfile
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("xl/workbook.xml"))
    n = 0
    for node in root:
        if node.tag.endswith("}sheets"):
            n += sum(1 for s in node if s.get("state", "visible")
                     != "visible")
    return n


def hidden_sheet_problems(register_path, period, judgements_path=None):
    """Open the workbooks through the loader's identity functions with the
    proxy: only the wanted visible sheets by exact title, nothing else, no
    sentinel in the identity or in an exception text."""
    bad = []
    with watched_workbooks() as log:
        try:
            ident = m.register_identity(register_path, period)
        except ValueError as e:
            return [f"{Path(register_path).name}: {_first(e)}"]
    if log["opened"] != [ident["sheet"]] or log["violations"]:
        bad.append(f"{Path(register_path).name}: opened {len(log['opened'])}"
                   f" sheet(s), violations {log['violations']}")
    if not str(ident["sheet"]).startswith(m.REGISTER_SHEET_PREFIX):
        bad.append("the sheet opened is not the visible register sheet")
    if judgements_path is not None:
        with watched_workbooks() as log:
            m.judgements_identity(judgements_path,
                                  Path(judgements_path).name, None)
        if log["opened"] != [m.JUDGEMENT_SHEET, m.NOTICE_SHEET] or \
                log["violations"]:
            bad.append(f"judgements workbook: opened {log['opened']}")
    return bad


def real_held_hidden(files=None, dirs=None):
    """The held register workbooks on disk: opened only through their one
    visible sheet by exact title; the count of non-visible sheets each
    carries is reported (never their names)."""
    files = m.LEGACY_FILES if files is None else files
    found = files_by_sha([m.RAW_DIR] if dirs is None else dirs)
    bad, parts = [], []
    for (part, p), facts in sorted(files.items()):
        if part != "register":
            continue
        path = found.get(facts["sha256"])
        if path is None:
            bad.append(f"register {p}: no file with sha256 "
                       f"{facts['sha256'][:16]} on disk")
            continue
        bad += hidden_sheet_problems(path, p)
        parts.append(f"{p}: {non_visible_sheets(path)} non-visible sheet(s) "
                     "never opened")
    return bad, parts


def gate_6_hidden_sheet(cur):
    name = ("the hidden sheet: opened sheets only by exact title, the "
            "sentinel reaches no output, record, log or exception text")
    problems = []
    src = sheet_access_problems()
    problems += [f"loader: {x}" for x in src]
    # the AST check itself
    for label, code in (
            ("wb.worksheets", "import openpyxl\ndef _read_sheet(path, "
             "title):\n    wb = openpyxl.load_workbook(path)\n    "
             "for ws in wb.worksheets:\n        pass\n    return wb[title]\n"),
            ("a sheetnames loop", "import openpyxl\ndef _read_sheet(path, "
             "title):\n    wb = openpyxl.load_workbook(path)\n    "
             "x = [n for n in wb.sheetnames]\n    return wb[title]\n"),
            ("wb.active", "import openpyxl\ndef _read_sheet(path, title):\n"
             "    wb = openpyxl.load_workbook(path)\n    return wb.active\n"),
            ("a second opener", "import openpyxl\ndef _read_sheet(path, "
             "title):\n    wb = openpyxl.load_workbook(path)\n    "
             "return wb[title]\ndef other(p):\n    return "
             "openpyxl.load_workbook(p)\n"),
            ("a sheet by position", "import openpyxl\ndef _read_sheet(path, "
             "title):\n    wb = openpyxl.load_workbook(path)\n    "
             "return wb[0]\n"),
            ("another zip member", "import zipfile\ndef f(p):\n    "
             "with zipfile.ZipFile(p) as z:\n        return "
             "z.read('xl/worksheets/sheet2.xml')\n")):
        if not sheet_access_problems(code):
            problems.append(f"the AST check does not catch {label}")
    if sheet_access_problems(
            "import openpyxl\ndef _read_sheet(path, title):\n    "
            "wb = openpyxl.load_workbook(path)\n    return "
            "list(wb[title].iter_rows())\n"):
        problems.append("the AST check rejects the exact-title form")

    def body(e):
        reg = write_register(e.root / "w" / "r.xlsx")
        jud = write_judgements(e.root / "w" / "20260930_J.xlsx")
        if non_visible_sheets(reg) != 1:
            problems.append("the fixture register has no veryHidden sheet")
        bad = hidden_sheet_problems(reg, SEP, jud)
        problems.extend(bad)
        ident = m.register_identity(reg, SEP)
        for x in (repr(ident), repr(m.build_register(ident["rows"], ident))):
            if SENTINEL in x or HIDDEN_TITLE in x:
                problems.append("the hidden sheet's content is in the "
                                "identity or the records")
        # the hidden sheet is never accepted in place of the visible one,
        # and exception text never names it
        both = write_book(e.root / "w" / "hid.xlsx", [
            ("Notes", [["x"]], "visible"),
            (m.REGISTER_SHEET_PREFIX + "...", [list(m.REGISTER_HEADERS)],
             "veryHidden")], datetime(2026, 9, 17, 14, 47))
        try:
            m.register_identity(both, SEP)
            problems.append("a veryHidden register sheet was accepted")
        except ValueError as ex:
            if SENTINEL in str(ex) or HIDDEN_TITLE in str(ex):
                problems.append("an exception names the hidden sheet")
        wrong = write_register(e.root / "w" / "bad.xlsx",
                               header=["Org", "Reg"] + list(
                                   m.REGISTER_HEADERS[2:]))
        try:
            m.register_identity(wrong, SEP)
        except ValueError as ex:
            if SENTINEL in str(ex) or HIDDEN_TITLE in str(ex):
                problems.append("a header exception names the hidden sheet")
        # end to end: a load stores, logs and prints nothing of it
        e.load_world()
        for s in ZZ.values():
            for t in (s.live_table, s.editions_table, ledger_name(s)):
                for needle in (SENTINEL, HIDDEN_TITLE):
                    n = cnt(e.cur, f"{t} AS t", "WHERE t::text LIKE %s",
                            (f"%{needle}%",))
                    if n:
                        problems.append(f"{t} holds the hidden sheet's text")
        if any(SENTINEL in x or HIDDEN_TITLE in x for x in OUTPUT):
            problems.append("the hidden sheet's text reached the output")
        rc, text, logged = e.cmd(["load", "--only", "register", "--recheck",
                                  SEP, "--commit"])
        for c in logged.call_args_list:
            if SENTINEL in str(c):
                problems.append("a run-log row holds the hidden sheet")
    scenario(cur, body)
    held_bad, parts = real_held_hidden()
    problems += [f"held files: {x}" for x in held_bad]
    report(6, name, not problems, "; ".join(problems[:3]) if problems else (
        "the loader opens a workbook only in _read_sheet, by wb[title], reads "
        "no other zip member; the AST check catches wb.worksheets, a "
        "sheetnames loop, wb.active, a second opener, a sheet by position "
        "and another zip member; a proxy workbook saw only the visible "
        "register sheet and the two judgements sheets opened by title; the "
        "sentinel is in no output, record, ledger, log or exception text; "
        "held register files: " + "; ".join(parts)))


# ---------------------------------------------------------------- gate 7

def raises(problems, label, fn, needle=None):
    try:
        fn()
    except ValueError as ex:
        if needle and needle not in str(ex):
            problems.append(f"{label}: raised without {needle!r}: "
                            f"{_first(ex)[:90]}")
        return
    problems.append(f"{label}: did not raise")


def real_identity(files=None, dirs=None):
    """Identity re-run on the held files by sha256: each register file
    states its own date (the held period) with the docProps modified time
    the loader surveyed, each judgements file its edition date from its
    name and the same modified time."""
    files = m.LEGACY_FILES if files is None else files
    found = files_by_sha([m.RAW_DIR] if dirs is None else dirs)
    bad, parts = [], []
    for (part, p), facts in sorted(files.items()):
        path = found.get(facts["sha256"])
        if path is None:
            bad.append(f"{part} {p}: no file with sha256 "
                       f"{facts['sha256'][:16]} on disk")
            continue
        try:
            ident = (m.register_identity(path, p) if part == "register" else
                     m.judgements_identity(path, facts["name"], None))
        except ValueError as e:
            bad.append(f"{part} {p}: {_first(e)}")
            continue
        if ident["period"] != p:
            bad.append(f"{path.name}: states {ident['period']}, held as {p}")
        if m._iso(ident["modified"]) != facts["modified"]:
            bad.append(f"{path.name}: docProps modified "
                       f"{m._iso(ident['modified'])}, surveyed "
                       f"{facts['modified']}")
        parts.append(f"{part} {p}")
    return not bad, "; ".join(bad[:3]) if bad else (
        f"{len(parts)} held files state the identity they are held under: "
        + ", ".join(parts))


def gate_7_identity(cur):
    name = ("identity from the file and the publisher's page: seeded "
            "mismatches halt")
    problems = []

    def body(e):
        R = e.root / "id"
        D = datetime
        # --- register identity
        def reg(label, needle, **kw):
            title = kw.pop("title_date", SEP)
            path = write_register(R / f"{label.replace(' ', '_')}.xlsx", **kw)
            raises(problems, label,
                   lambda: m.register_identity(path, title), needle)
        reg("title date later than docProps", "this is not the",
            modified=D(2026, 9, 1, 10))
        reg("docProps long after the title date", "expected between",
            modified=D(2026, 10, 5, 10))
        reg("a changed header", "header changed",
            header=["Organisation", "Registration number",
                    "Registration date", "Designation", "Corporate form",
                    "Notes "])
        reg("no register sheet", "0 visible sheets", title="Notes")
        two = write_book(R / "two.xlsx", [
            (m.REGISTER_SHEET_PREFIX + " A", [list(m.REGISTER_HEADERS)],
             "visible"),
            (m.REGISTER_SHEET_PREFIX + " B", [list(m.REGISTER_HEADERS)],
             "visible")], D(2026, 9, 17, 10))
        raises(problems, "two visible register sheets",
               lambda: m.register_identity(two, SEP), "expected one")
        # --- judgements identity
        good = write_judgements(R / "20260930_J.xlsx")
        raises(problems, "a file name that is not a date",
               lambda: m.judgements_identity(good, "Judgements.xlsx", None),
               "does not start with a date")
        raises(problems, "a file name date the page does not state",
               lambda: m.judgements_identity(
                   good, "20260930_J.xlsx", judgements_page(
                       "u", changes=("2026-09-09",),
                       updated="2026-09-10T09:30:00+01:00")),
               "not a date the page states")
        far = write_judgements(R / "20260930_far.xlsx",
                               modified=D(2026, 10, 20, 10))
        raises(problems, "docProps eight or more days after the name",
               lambda: m.judgements_identity(far, "20260930_far.xlsx", None),
               "docProps modified")
        one = write_book(R / "20260930_one.xlsx", [
            ("Regulatory Judgements", [list(m.JUDGEMENT_HEADERS)],
             "visible")], D(2026, 9, 30, 11))
        raises(problems, "a missing Enforcement Notices sheet",
               lambda: m.judgements_identity(one, "20260930_one.xlsx", None),
               "missing")
        hj = write_judgements(R / "20260930_hj.xlsx",
                              jheader=list(m.JUDGEMENT_HEADERS[:-1])
                              + ["Process"])
        raises(problems, "a changed judgements header",
               lambda: m.judgements_identity(hj, "20260930_hj.xlsx", None),
               "header changed")
        hn = write_judgements(R / "20260930_hn.xlsx",
                              nheader=list(m.NOTICE_HEADERS[:-1]) + ["Date"])
        raises(problems, "a changed notices header",
               lambda: m.judgements_identity(hn, "20260930_hn.xlsx", None),
               "header changed")
        # --- discovery
        page = register_page()
        atts = page["details"]["attachments"]
        for label, mutate, needle in (
                ("no register attachment", lambda a: [
                    x for x in a if not x["url"].endswith(".xlsx")],
                 "accessible version"),
                ("two register attachments", lambda a: a + [dict(
                    a[1], title="List of registered providers – 17 August "
                    "2026", url=MEDIA.format("o", "old.xlsx"))],
                 "17 August 2026"),
                ("an .xlsx with another title", lambda a: a + [{
                    "title": "Registered providers list",
                    "url": MEDIA.format("o", "x.xlsx")}], "Registered "
                 "providers list")):
            pg = register_page()
            pg["details"]["attachments"] = mutate(
                register_page()["details"]["attachments"])
            raises(problems, label, lambda: m.register_attachment(pg),
                   needle)
        del atts
        raises(problems, "a judgements attachment with another title",
               lambda: m.judgements_attachment(judgements_page(
                   MEDIA.format("j", "x.xlsx"), title="Judgements table")),
               "Judgements table")
        pg = judgements_page(MEDIA.format("j", "a.xlsx"))
        pg["details"]["attachments"].append(
            dict(pg["details"]["attachments"][0],
                 url=MEDIA.format("j", "b.xlsx")))
        raises(problems, "two judgements attachments",
               lambda: m.judgements_attachment(pg), "expected exactly one")
        pg["details"]["attachments"] = []
        raises(problems, "no judgements attachment",
               lambda: m.judgements_attachment(pg), "expected exactly one")
        # --- cells
        def build_j(label, needle, rows):
            p = write_judgements(R / f"20260930_{label.replace(' ', '_')}.xlsx",
                                 jrows=rows)
            ident = m.judgements_identity(p, p.name, None)
            raises(problems, label, lambda: m.build_judgements(
                ident["judgement_rows"], ident), needle)
        build_j("an unseen governance grade", "not seen before",
                [jrow(g="G5")])
        build_j("an unseen consumer grade", "not seen before",
                [jrow(c="C2+")])
        build_j("an Excel serial out of range", "not an Excel date serial",
                [jrow(cdate=1000)])
        build_j("a date cell with text", "is not a date",
                [jrow(gdate="tomorrow")])
        build_j("a blank registration number", "blank (a key column)",
                [jrow(reg=None)])
        build_j("a repeated key", "duplicate key", [jrow(), jrow()])
        build_j("a blank landlord", "required column", [jrow(landlord=None)])
        # --- through load: halts, nothing stored, no run-log row
        def through(label, setup, needle, part):
            def inner(c):
                with env(c) as ee, tl.specs():
                    setup(ee)
                    ok, t = ee.halts(["load", "--only", part, "--commit"],
                                     state(c), needle)
                    if not ok:
                        problems.append(f"load with {label}: not halted "
                                        f"cleanly naming {needle!r}: "
                                        f"{t[-120:]}")
            _in_savepoint(e.cur, inner)
        through("a register whose docProps disagrees with its title",
                lambda ee: ee.register(modified=D(2026, 10, 5, 10)),
                "expected between", "register")
        def undated(ee):
            ee.judgements(changes=("2026-09-09",))
            ee.api[m.JUDGEMENTS_PATH]["public_updated_at"] =                 "2026-09-10T09:30:00+01:00"
        through("a judgements file the page does not date", undated,
                "not a date the page states", "judgements")
        through("an unseen grade", lambda ee: ee.judgements(
            jrows=[jrow(g="G5")]), "not seen before", "judgements")
    scenario(cur, body)
    ok_real, detail = real_identity()
    if not ok_real:
        problems.append(f"held files: {detail}")
    report(7, name, not problems, "; ".join(problems[:3]) if problems else (
        "the register title date against docProps (earlier, much later), a "
        "changed header, no or two visible register sheets, a judgements "
        "name that is no date or that the page does not state, docProps "
        "eight days out, a missing sheet, changed headers, none, two or a "
        "wrongly titled attachment, an unseen grade, an out-of-range serial, "
        "text in a date cell, a blank or repeated key and a blank landlord "
        "each halt, and a load with such a file stores nothing and logs "
        f"nothing; {detail}"))


# ---------------------------------------------------------------- gate 8

def gate_8_older_file(cur):
    name = ("older-file guard on the page, --file and --file --no-page "
            "paths; the rank is the file's own, never its URL")
    problems = []
    D = datetime

    def register(e):
        e.migrated()
        e.pages()
        before = state(e.cur)
        later = MEDIA.format("m99999999", "List_of_registered_providers_"
                             "30_December_2026.xlsx")
        old = e.register("17 September 2026", rows=REG_ROWS[:3],
                         modified=D(2026, 9, 17, 9), url=later)
        for how, argv in (
                ("the page", ["load", "--only", "register", "--commit"]),
                ("--file", ["load", "--only", "register", "--file", old,
                            "--title-date", SEP, "--commit"]),
                ("--file --no-page", ["load", "--only", "register", "--file",
                                      old, "--title-date", SEP, "--no-page",
                                      "--commit"])):
            ok, t = e.halts(argv, before, "older")
            if not ok:
                problems.append(f"register {how}: an earlier-modified file "
                                f"at a later-looking URL did not halt as "
                                f"older: {t[-100:]}")
        # the file's own rank: later modified, old-looking URL, same
        # content: unchanged, never "older"
        e.register("17 September 2026", rows=REG_ROWS[:3],
                   modified=D(2026, 9, 18, 9),
                   url=MEDIA.format("m1", "List_of_registered_providers_"
                                    "01_January_2020.xlsx"))
        rc, text, _ = e.cmd(["load", "--only", "register"])
        if rc != 0 or f"register {SEP}: unchanged" not in text or \
                "older" in text:
            problems.append(f"a later-modified file at an old-looking URL: "
                            f"rc {rc}: {text[-120:]}")

    def judgements(e):
        e.migrated()
        e.pages()
        before = state(e.cur)
        old = e.judgements("20260701", jrows=JROWS[:4],
                           changes=("2026-07-01",))
        for how, argv in (
                ("the page", ["load", "--only", "judgements", "--commit"]),
                ("--file", ["load", "--only", "judgements", "--file", old,
                            "--commit"]),
                ("--file --no-page", ["load", "--only", "judgements",
                                      "--file", old, "--no-page",
                                      "--commit"])):
            ok, t = e.halts(argv, before, "older")
            if not ok:
                problems.append(f"judgements {how}: an older edition did not"
                                f" halt: {t[-100:]}")
        # same name date, earlier docProps modified time: older
        e.judgements("20260930", modified=D(2026, 9, 30, 8))
        ok, t = e.halts(["load", "--only", "judgements", "--commit"],
                        state(e.cur), "older")
        if not ok:
            problems.append(f"an earlier-modified file of the held edition "
                            f"date did not halt: {t[-100:]}")
        e.judgements("20260701", jrows=JROWS[:4], changes=("2026-07-01",))
        text, logged = e.ok(e.cur, ["load", "--only", "judgements",
                                    "--allow-older-file", "--commit"])
        eds = tl.editions(e.cur, "judgements", "2026-07-01")
        if eds != [(1, None, 4)] or "held-not-live" not in text or \
                not logged.called or "--allow-older-file" not in \
                logged.call_args[0][2]:
            problems.append(f"--allow-older-file: {eds}, logged "
                            f"{logged.called}")
        if tl.live_periods(e.cur, "judgements")[-1][0] != J_SEP:
            problems.append("--allow-older-file displaced the live period")
    scenario(cur, register)
    scenario(cur, judgements)
    mixed(8, name, not problems,
          "an earlier-modified register file at a later-looking URL halts "
          "as older on the page, --file and --file --no-page with nothing "
          "stored or logged, while a later-modified one at an old-looking "
          "URL is read as unchanged; an older judgements edition halts on "
          "all three paths, so does the held edition date with an earlier "
          "modified time; --allow-older-file stores it as history, logged, "
          "live untouched" if not problems else "; ".join(problems[:3]),
          cur, lambda c: real_ranks_from_source(c, REAL))


def real_ranks_from_source(cur, specs=None):
    """Every period whose latest edition is past edition 1 'as loaded' (and
    not restored) carries a rank in its source_file (file date and docProps
    modified time), never one read from a URL or a bare file name; every
    edition 1 'as loaded' except the August register carries one too."""
    specs = specs or REAL
    _need(cur, specs)
    bad, n = [], 0
    for name in NAMES:
        spec = specs[name]
        for p, t in sorted(m.tip_info(cur, spec).items()):
            n += 1
            if t["rank"] is None and not (
                    name == "register" and p == m.AUGUST_REGISTER["period"]):
                bad.append(f"{spec.editions_table} {p}: tip edition "
                           f"{t['edition']} has no rank in its source_file")
    if not n:
        bad.append(EMPTY)
    return not bad, "; ".join(bad[:3]) if bad else (
        f"{n} periods, each tip carries its file's own date and docProps "
        "modified time (the August register, whose file cannot be "
        "obtained, none)")


# ---------------------------------------------------------------- gate 9

def gate_9_stop_conditions(cur):
    name = ("stop conditions: each seeded REJECTED, stores nothing, no "
            "ledger row, partial run-log row")
    problems = []
    D = datetime
    many = [(f"P{i}", f"{9000 + i}", D(2001, 1, 1), "Profit", "Company", None)
            for i in range(40)]
    J = ("judgements", "notices")

    def case(label, setup, argv, needle, names, period, after=None):
        def body(c):
            with env(c) as e, tl.specs():
                setup(e)
                bad = rejects(e, argv(e) if callable(argv) else argv, needle,
                              names, period)
                if bad:
                    problems.append(f"{label}: {bad}")
                elif after:
                    after(e)
        _in_savepoint(cur, body)

    reg_load = ["load", "--only", "register", "--commit"]
    jud_load = ["load", "--only", "judgements", "--commit"]

    def reg_drop(e):
        e.register("17 August 2026", rows=many, modified=D(2026, 8, 17, 10))
        e.ok(e.cur, reg_load)
        e.register()

    def released(e):
        text, logged = e.ok(e.cur, reg_load + ["--acknowledge", SEP])
        if "ACKNOWLEDGED" not in text or tl.editions(
                e.cur, "register", SEP) != [(1, None, 4)] or \
                f"ACKNOWLEDGED {SEP}" not in logged.call_args[0][2]:
            problems.append("--acknowledge did not store and log the period")
    case("a register that drops 40 of 40 providers", reg_drop, reg_load,
         "STOP CONDITION", ("register",), SEP, released)

    def reg_added(e):
        e.register("17 August 2026", modified=D(2026, 8, 17, 10))
        e.ok(e.cur, reg_load)
        e.register(rows=REG_ROWS + many)
    case("a register adding 40 providers", reg_added, reg_load, "added",
         ("register",), SEP)

    def jud_shrink(e):
        e.judgements("20260812", jrows=[jrow(reg=f"Z{i:03d}")
                                        for i in range(20)])
        e.ok(e.cur, jud_load)
        e.judgements("20260930")
    case("a judgements edition with 6 of 20 rows", jud_shrink, jud_load,
         "fewer than 90%", J, J_SEP)

    def reissue(e):
        e.register()
        e.ok(e.cur, reg_load)
        rows = list(REG_ROWS)
        rows[0] = rows[0][:3] + ("Profit",) + rows[0][4:]
        e.register(rows=rows)

    def reissue_after(e):
        rc, _, _ = e.cmd(reg_load + ["--accept-reissue", SEP])
        if rc != 1:
            problems.append("--accept-reissue alone released an "
                            "unacknowledged change")
        text, logged = e.ok(e.cur, reg_load + ["--accept-reissue", SEP,
                                               "--acknowledge", SEP])
        if "ACCEPTED REISSUE" not in text or "--accept-reissue" not in \
                logged.call_args[0][2]:
            problems.append("the accepted reissue was not stored and logged")
    case("a reissue keeping its date and time with other content", reissue,
         reg_load, "--accept-reissue", ("register",), SEP, reissue_after)

    def held_reread(e):
        e.migrated()
        e.pages()
    case("a re-read of a held judgements file (null_reasons change)",
         held_reread, lambda e: ["load", "--only", "judgements", "--file",
                                 e.paths[("judgements", J_SEP)], "--no-page",
                                 "--commit"], "--acknowledge", J, J_SEP)

    def correction_count(e):
        e.migrated()
        e.pages()
    case("the real correction count against the fixture's 4 cells",
         correction_count, ["load", "--only", "register", "--correction",
                            CORR, "--acknowledge", AUG, "--commit"], "1,578",
         ("register",), AUG)

    def no_notices(e):
        e.judgements(nrows=[])

    def no_notices_after(e):
        # neither --acknowledge nor --accept-reissue releases it
        for extra in (["--acknowledge", J_SEP], ["--accept-reissue", J_SEP]):
            bad = rejects(e, jud_load + extra, "--acknowledge-no-notices", J,
                          J_SEP)
            if bad:
                problems.append(f"no notices with {extra[0]}: {bad}")
    case("a judgements file with no enforcement notices, unacknowledged",
         no_notices, jud_load, "--acknowledge-no-notices", J, J_SEP,
         no_notices_after)

    def legacy_two(e):
        e.judgements("20260812", jrows=JROWS[:4] + [
            jrow(reg="L0009", landlord="Gone Homes")])
        e.ok(e.cur, jud_load)
        e.judgements("20260930")
        e.ok(e.cur, jud_load)
        cols = ", ".join(("registration_number", "publication_date",
                          "edition_date") + m.JUDGEMENT_VALUES)
        e.cur.execute(
            f"INSERT INTO public.{ZZ['judgements'].live_table} ({cols}, "
            "source_url, source_file, release_page_url) SELECT "
            f"{cols}, file_source_url, file_name, file_release_page_url FROM "
            f"public.{ZZ['judgements'].editions_table} WHERE "
            "registration_number = 'L0009'")
        e.judgements("20261007", changes=("2026-10-07",))
    case("a new judgements edition while live holds two periods", legacy_two,
         jud_load, "refresh-latest", J, "2026-10-07")

    def halts_case(label, setup, argv, needle):
        def body(c):
            with env(c) as e, tl.specs():
                setup(e)
                ok, t = e.halts(argv, state(c), needle)
                if not ok:
                    problems.append(f"{label}: not halted cleanly naming "
                                    f"{needle!r}: {t[-120:]}")
        _in_savepoint(cur, body)
    halts_case("an acknowledgement naming a period the run does not "
               "compare", lambda e: e.register(),
               reg_load + ["--acknowledge", "2026-01-01"],
               "not periods this run")
    halts_case("a repeated registration number",
               lambda e: e.register(rows=REG_ROWS + [REG_ROWS[0]]), reg_load,
               "duplicate key")
    halts_case("a blank registration number",
               lambda e: e.register(rows=REG_ROWS + [
                   ("Nameless", None, D(2001, 1, 1), "Profit", "Company",
                    None)]), reg_load, "blank (a key column)")
    report(9, name, not problems,
           "a register losing 40 providers, adding 40, a judgements edition "
           "at 30% of the last, a reissue with other content, a held "
           "judgements re-read, a wrong correction count, a file with no "
           "notices (without --acknowledge-no-notices; --acknowledge and "
           "--accept-reissue do not release it) and a new edition while live "
           "holds two periods are each "
           "REJECTED with exit 1, nothing stored, no ledger row and a "
           "partial run-log row; --acknowledge (and --accept-reissue with "
           "it) releases and logs; a stray acknowledgement, a repeated and a "
           "blank key halt with nothing stored"
           if not problems else "; ".join(problems[:3]))


# ---------------------------------------------------------------- gate 10

def gate_10_one_transaction(cur):
    name = ("a new period: edition, live rows and ledger row in one "
            "transaction (both tables of a judgements file)")
    problems = []
    J = ("judgements", "notices")
    real_insert = m.insert_live
    real_record = pe.record_file_check

    def good(e):
        e.register()
        text, logged = e.ok(e.cur, ["load", "--only", "register", "--commit"])
        c = counts_of(e.cur, ("register",), SEP)
        if c != (4, 4, 1) or logged.call_count != 1 or \
                f"register {SEP} new" not in logged.call_args[0][2]:
            problems.append(f"a new register date stored {c}, logged "
                            f"{logged.call_count}")
        e.judgements()
        e.ok(e.cur, ["load", "--only", "judgements", "--commit"])
        n = len(JROWS) + len(NROWS)
        c = counts_of(e.cur, J, J_SEP)
        if c != (n, n, 2):
            problems.append(f"a new judgements file stored {c}, expected "
                            f"{(n, n, 2)}")

    def failing(label, patch, setup, names, period, argv):
        def body(c):
            with env(c) as e, tl.specs():
                setup(e)
                with patch:
                    rc, text, logged = e.cmd(argv)
                if rc != 1 or "FAILED" not in text:
                    problems.append(f"{label}: rc {rc}, not FAILED: "
                                    f"{text[-100:]}")
                    return
                if counts_of(e.cur, names, period) != (0, 0, 0):
                    problems.append(f"{label}: something was left stored: "
                                    f"{counts_of(e.cur, names, period)}")
                notes = logged.call_args[0][2] if logged.called else ""
                if "PARTIAL RUN" not in notes or "failed [" not in notes:
                    problems.append(f"{label}: no partial run-log row")
        _in_savepoint(cur, body)

    def only_notices(real):
        def f(cur_, profile, *a, **k):
            if profile.spec.editions_table == ZZ["notices"].editions_table:
                raise RuntimeError("boom")
            return real(cur_, profile, *a, **k)
        return f

    def reg(e):
        e.register()

    def jud(e):
        e.judgements()
    reg_argv = ["load", "--only", "register", "--commit"]
    jud_argv = ["load", "--only", "judgements", "--commit"]
    failing("a failing register ledger insert",
            mock.patch.object(pe, "record_file_check",
                              side_effect=RuntimeError("boom")),
            reg, ("register",), SEP, reg_argv)
    failing("a failing register live insert",
            mock.patch.object(m, "insert_live",
                              side_effect=RuntimeError("boom")),
            reg, ("register",), SEP, reg_argv)
    failing("a failing notices live insert after the judgements",
            mock.patch.object(m, "insert_live",
                              side_effect=only_notices(real_insert)),
            jud, J, J_SEP, jud_argv)
    failing("a failing notices ledger insert after the judgements",
            mock.patch.object(pe, "record_file_check",
                              side_effect=only_notices(real_record)),
            jud, J, J_SEP, jud_argv)
    scenario(cur, good)
    report(10, name, not problems,
           "a new register date stores edition 1, its live rows and one "
           "ledger row, logged once; a new judgements file does so for both "
           "tables; a failing ledger or live insert (also on the second "
           "table) leaves nothing stored in either, with a partial run-log "
           "row" if not problems else "; ".join(problems[:3]))


# ---------------------------------------------------------------- gate 11

def gate_11_ledger_skip(cur):
    name = ("ledger rule: a skip needs the (final URL, sha256) pair for the "
            "period")
    problems = []
    D = datetime

    def body(e):
        e.register()
        url = e.last_url
        e.ok(e.cur, ["load", "--only", "register", "--commit"])
        base = len(tl.ledger(e.cur, "register"))
        snap = state(e.cur)
        text, logged = e.ok(e.cur, ["load", "--only", "register", "--commit"])
        if "already in the ledger" not in text or logged.called or \
                state(e.cur) != snap:
            problems.append("the same (URL, sha256) was not skipped quietly")
        # the same URL with other bytes is read, not skipped
        e.register(modified=D(2026, 9, 18, 9), url=url)
        text, _ = e.ok(e.cur, ["load", "--only", "register"])
        if "already in the ledger" in text or \
                f"register {SEP}: unchanged" not in text:
            problems.append("other bytes at the same URL were skipped")
        # the same bytes from another final URL are read, not skipped
        e.register(url=url)
        e.final[url] = MEDIA.format("moved", "List.xlsx")
        text, _ = e.ok(e.cur, ["load", "--only", "register", "--commit"])
        if "already in the ledger" in text or \
                len(tl.ledger(e.cur, "register")) != base + 1:
            problems.append("the same bytes from another URL were skipped "
                            "or left no ledger row")
        text, _ = e.ok(e.cur, ["load", "--only", "register", "--commit"])
        if "already in the ledger" not in text:
            problems.append("the new (URL, sha256) pair was not skipped the "
                            "next time")
        # a failed or rejected run leaves no ledger row, so a rerun reads
        e.final.clear()
        e.register("17 October 2026", modified=D(2026, 10, 17, 10))
        with mock.patch.object(pe, "record_file_check",
                               side_effect=RuntimeError("boom")):
            rc, _, _ = e.cmd(["load", "--only", "register", "--commit"])
        text, _ = e.ok(e.cur, ["load", "--only", "register", "--commit"])
        if rc != 1 or "already in the ledger" in text or tl.editions(
                e.cur, "register", "2026-10-17") != [(1, None, 4)]:
            problems.append(f"a failed run was not read again: rc {rc}")
    scenario(cur, body)
    # the pure rule
    t0 = D(2026, 9, 17, 14, 47, tzinfo=m.timezone.utc)
    src = m.ledger_source("https://x/a.xlsx", "register", SEP, t0)
    sha = "ab" * 32
    checked = {SEP: {(src, sha)}}
    rk = (date(2026, 9, 17), t0)
    ranks = {SEP: rk}
    if m.ledger_complete(checked, "https://x/a.xlsx", sha, ranks) != [SEP]:
        problems.append("the recorded pair is not skipped")
    for label, args in (("another URL", ("https://x/b.xlsx", sha)),
                        ("other bytes", ("https://x/a.xlsx", "cd" * 32))):
        if m.ledger_complete(checked, args[0], args[1], ranks) is not None:
            problems.append(f"{label} was skipped")
    if m.ledger_complete(checked, "https://x/a.xlsx", sha, ranks,
                         stranded=(SEP,)) is not None:
        problems.append("a stranded period was skipped")
    newer = {SEP: (rk[0], rk[1].replace(hour=20))}
    if m.ledger_complete(checked, "https://x/a.xlsx", sha, newer) is not None:
        problems.append("a period whose tip came from a newer file was "
                        "skipped as complete")
    report(11, name, not problems,
           "the recorded (final URL, sha256) pair is skipped quietly (no "
           "log, no write); other bytes at the same URL and the same bytes "
           "from another final URL are read; a failed run leaves no row so "
           "the rerun reads; a stranded period or a newer tip is not "
           "skipped as complete" if not problems
           else "; ".join(problems[:3]))


# ---------------------------------------------------------------- gate 12

def gate_12_preview(cur):
    name = ("preview and simulate write nothing, every writing command")
    problems = []
    D = datetime

    def body(e):
        e.completed()
        # something pending for every command: a revised September register
        # (edition 2, live not yet refreshed), a new October register and a
        # new judgements edition
        rows = list(REG_ROWS[:3])
        rows[0] = rows[0][:3] + ("Profit",) + rows[0][4:]
        e.register("17 September 2026", rows=rows, modified=D(2026, 9, 18, 9))
        e.ok(e.cur, ["load", "--only", "register", "--acknowledge", SEP,
                     "--commit"])
        e.register("17 October 2026", modified=D(2026, 10, 17, 10))
        e.judgements("20261007", changes=("2026-10-07", J_SEP))
        snap = state(e.cur)
        reg_live = ZZ["register"].live_table
        for argv in (["status"], ["ddl"], ["ddl", "--simulate"], ["load"],
                     ["load", "--simulate"],
                     ["load", "--only", "judgements"],
                     ["refresh-latest"], ["refresh-latest", "--simulate"],
                     ["restore-edition", reg_live, SEP, "1"],
                     ["restore-edition", reg_live, SEP, "1", "--simulate"]):
            rc, text, logged = e.cmd(argv)
            if state(e.cur) != snap or logged.called:
                problems.append(f"{' '.join(map(str, argv[:3]))}: wrote "
                                f"({'run-log' if logged.called else 'rows'})")
            if argv[0] != "status":
                want = "SIMULATION" if "--simulate" in argv else (
                    "DRY RUN" if argv[0] in ("ddl", "refresh-latest")
                    else "PREVIEW")
                if want not in text:
                    problems.append(f"{' '.join(map(str, argv[:3]))}: no "
                                    f"{want} line")
        text, _ = e.ok(e.cur, ["load"])
        if "would store edition 1" not in text:
            problems.append("the preview had nothing pending to compare")

    def legacy(e):
        e.seed(e.cur)
        snap = state(e.cur)
        with e.legacy_cli():
            for argv, want in ((["migrate-legacy"], "PREVIEW"),
                               (["migrate-legacy", "--simulate"],
                                "SIMULATION")):
                rc, text, logged = e.cmd(argv)
                if rc != 0 or want not in text or state(e.cur) != snap or \
                        logged.called:
                    problems.append(f"{' '.join(argv)}: rc {rc}, wrote or "
                                    f"no {want}: {text[-120:]}")
    scenario(cur, body)
    scenario(cur, legacy)
    report(12, name, not problems,
           "status, ddl, load, refresh-latest and restore-edition as a "
           "preview and with --simulate, and migrate-legacy likewise, change "
           "no live, editions or ledger row and write no run-log row, with "
           "a revision, a new register date and a new judgements edition "
           "pending" if not problems else "; ".join(problems[:3]))


# ---------------------------------------------------------------- gate 13

def real_held_reread(cur, specs=None, files=None, dirs=None):
    """The held files read now with the new rules equal the stored tip of
    their period, null_reasons included (so a byte-identical re-read is
    `unchanged`): the register files equal edition 1 'as loaded'; the
    judgements and notices equal the correction edition 2 (edition 1 marks
    its empties reason_not_recorded_as_loaded, so until edition 2 is stored
    this fails, naming it). The August register, which has no file, holds no
    stored 'None' notes in its tip."""
    specs = specs or REAL
    files = m.LEGACY_FILES if files is None else files
    _need(cur, specs)
    found = files_by_sha([m.RAW_DIR] if dirs is None else dirs)
    bad, parts = [], []
    if not files:
        bad.append("no held files to read (an empty list is not a pass)")
    for (part, p), facts in sorted(files.items()):
        path = found.get(facts["sha256"])
        if path is None:
            bad.append(f"{part} {p}: no file with sha256 "
                       f"{facts['sha256'][:16]} on disk")
            continue
        try:
            if part == "register":
                ident = m.register_identity(path, p)
                built = {"register": m.build_register(ident["rows"], ident)}
            else:
                ident = m.judgements_identity(path, facts["name"], None)
                built = {"judgements": m.build_judgements(
                    ident["judgement_rows"], ident),
                    "notices": m.build_notices(ident["notice_rows"], ident)}
        except ValueError as e:
            bad.append(f"{part} {p}: {_first(e)}")
            continue
        for n, new in built.items():
            spec = specs[n]
            try:
                tip = core.chain_tip(cur, spec, p)
            except (LookupError, ValueError) as e:
                bad.append(f"{n} {p}: {_first(e)}")
                continue
            a, r, c = m.row_diff(new, m.records(cur, spec, p, tip), n)
            if a or r or c:
                bad.append(f"{n} {p}: the file read now differs from the "
                           f"tip edition {tip} ({len(a)} added, {len(r)} "
                           f"removed, {len(c)} changed"
                           + ("; edition 2, with the real reasons, is not "
                              "stored" if tip == 1 and n != "register"
                              else "") + ")")
            else:
                parts.append(f"{n} {p} edition {tip}")
            if n == "register":
                a, r, c = m.row_diff(new, m.records(cur, spec, p, 1), n)
                if a or r or c:
                    bad.append(f"register {p}: edition 1 as loaded differs "
                               f"from the file read now ({len(c)} changed)")
    reg = specs["register"]
    aug = m.AUGUST_REGISTER["period"]
    if aug in _edition_periods(cur, reg):
        tip = core.chain_tip(cur, reg, aug)
        n = cnt(cur, reg.editions_table, f"WHERE {reg.period_col} = %s AND "
                "edition = %s AND notes = 'None'", (aug, tip))
        if n:
            bad.append(f"register {aug}: the tip (edition {tip}) still "
                       f"holds the stored text 'None' in {n} notes")
        else:
            parts.append(f"register {aug} edition {tip}: no stored 'None' "
                         "notes")
    return not bad, "; ".join(bad[:3]) if bad else (
        "each held file read now equals the stored tip, null_reasons "
        "included: " + ", ".join(parts))


def gate_13_reread(cur):
    name = ("a byte-identical re-read is `unchanged` on every path: the "
            "register against edition 1 as loaded; judgements and notices "
            "against edition 2")
    problems = []
    A = []

    def register_paths(e):
        e.migrated()
        e.pages()
        sep, jul = e.paths[("register", SEP)], e.paths[("register", JUL)]
        before = tables_state(e.cur)
        text, logged = e.ok(e.cur, ["load", "--only", "register", "--commit"])
        if "already in the ledger" not in text or logged.called:
            problems.append("register page path: not skipped by the ledger")
        for label, argv, period in (
                ("--recheck", ["--recheck", SEP], SEP),
                ("--file", ["--file", sep, "--title-date", SEP], SEP),
                ("--file --no-page", ["--file", jul, "--title-date", JUL,
                                      "--no-page"], JUL)):
            text, _ = e.ok(e.cur, ["load", "--only", "register"] + argv
                           + ["--commit"])
            if f"register {period}: unchanged" not in text:
                problems.append(f"register {label}: not unchanged against "
                                f"edition 1: {text[-120:]}")
        if tables_state(e.cur) != before:
            problems.append("a register re-read changed the live or "
                            "editions tables")
        for p in (JUL, SEP):
            if tl.editions(e.cur, "register", p)[-1][0] != 1:
                problems.append(f"register {p}: a re-read stored an edition")
        ok, d = real_held_reread(e.cur, ZZ, e.held, [e.root])
        if ok or "edition 2" not in d:
            problems.append(f"the real check passes before edition 2: "
                            f"{d[:100]}")

    def judgements_paths(e):
        e.migrated()
        e.pages()
        jsep = e.paths[("judgements", J_SEP)]
        base = ["load", "--only", "judgements", "--file", jsep, "--no-page",
                "--commit"]
        before = state(e.cur)
        rc, text, logged = e.cmd(base)
        if rc != 1 or "REJECTED" not in text or "--acknowledge" not in text \
                or "--accept-reissue" in text or state(e.cur) != before:
            problems.append(f"edition 1's empties re-read without an "
                            f"acknowledgement: rc {rc}: {text[-120:]}")
        e.ok(e.cur, base + ["--acknowledge", J_SEP])
        for n in ("judgements", "notices"):
            if tl.editions(e.cur, n, J_SEP)[-1][:2] != (2, 1):
                problems.append(f"{n}: edition 2 not stored")
        if core.rows_differing(e.cur, ZZ["judgements"], J_SEP, 2):
            problems.append("live differs from edition 2 on the values")
        snap = tables_state(e.cur)
        for label, argv in (
                ("--recheck", ["--recheck", J_SEP]),
                ("--file", ["--file", jsep, "--recheck", J_SEP]),
                ("--file --no-page", ["--file", jsep, "--no-page",
                                      "--recheck", J_SEP])):
            text, _ = e.ok(e.cur, ["load", "--only", "judgements"] + argv
                           + ["--commit"])
            for n in ("judgements", "notices"):
                if f"{n} {J_SEP}: unchanged" not in text:
                    problems.append(f"{n} {label}: not unchanged against "
                                    f"edition 2: {text[-120:]}")
        text, logged = e.ok(e.cur, ["load", "--only", "judgements",
                                    "--commit"])
        if "already in both ledgers" not in text or logged.called:
            problems.append("judgements page path: not skipped by the "
                            "ledgers")
        text, _ = e.ok(e.cur, ["load", "--only", "judgements", "--recheck",
                               J_SEP, "--commit"])
        if f"judgements {J_SEP}: unchanged" not in text:
            problems.append("judgements page --recheck: not unchanged")
        if tables_state(e.cur) != snap:
            problems.append("a judgements re-read after edition 2 changed "
                            "the live or editions tables")

    def completed(e):
        e.completed()
        ok, d = real_held_reread(e.cur, ZZ, e.held, [e.root])
        if not ok:
            problems.append(f"the completed state fails the real check: {d}")
        A.append(d)
        e.cur.execute("SAVEPOINT x")
        # a held file that is not the stored one: its re-read differs
        clone_edition(e.cur, ZZ["register"], f"snapshot_date = '{SEP}' AND "
                      "edition = 1 AND registration_number = '4001'",
                      edition="2", supersedes="1", designation="'Other'")
        ok, d = real_held_reread(e.cur, ZZ, e.held, [e.root])
        e.cur.execute("ROLLBACK TO SAVEPOINT x")
        if ok or "differs from the tip edition" not in d:
            problems.append(f"a changed tip passes: {d[:100]}")
    scenario(cur, register_paths)
    scenario(cur, judgements_paths)
    scenario(cur, completed)
    mixed(13, name, not problems,
          "the held register files re-read through the page ledger, "
          "--recheck, --file and --file --no-page are unchanged against "
          "edition 1 with no edition stored; the judgements re-read is "
          "REJECTED without --acknowledge (its edition 1 has only "
          "reason_not_recorded_as_loaded), stores edition 2 with it, then "
          "is unchanged on every path; the real check needs edition 2 and "
          "catches a changed tip" if not problems
          else "; ".join(problems[:3]), cur, real_held_reread)


# ---------------------------------------------------------------- gate 14

def real_history(cur, specs=None):
    """Judgements and notices: live holds the newest edition's period only,
    every older period is history in the editions (status ok, history
    listed); the register keeps every period live."""
    specs = specs or REAL
    _need(cur, specs)
    bad, parts = [], []
    for name in NAMES:
        spec = specs[name]
        prof = m.PROFILES[name].with_spec(spec)
        live, held = _live_periods(cur, spec), _edition_periods(cur, spec)
        e = empty_latest(cur, spec)
        if not live and held and e:
            st = pe.status(cur, prof)
            if not st["ok"] or st.get("history") != held:
                bad.append(f"{name}: recorded empty {e}, status ok "
                           f"{st['ok']}, history {st.get('history')}, "
                           f"expected {held}")
            parts.append(f"{name} live empty ({e} recorded empty), history "
                         f"{held}")
            continue
        if not live or not held:
            bad.append(f"{spec.live_table}: {EMPTY}")
            continue
        st = pe.status(cur, prof)
        if not st["ok"]:
            bad.append(f"{name}: status not clean ({ {k: v for k, v in st.items() if v and k != 'ok'} })")
        if name == "register":
            if live != held:
                bad.append(f"register: live {live}, held {held}")
            parts.append(f"register {len(live)} live periods")
            continue
        if live != held[-1:]:
            bad.append(f"{name}: live {live}, the newest edition is "
                       f"{held[-1:]}")
        if st.get("history") != held[:-1]:
            bad.append(f"{name}: history {st.get('history')}, expected "
                       f"{held[:-1]}")
        parts.append(f"{name} live {live}, history {held[:-1]}")
    return not bad, "; ".join(bad[:3]) if bad else "; ".join(parts)


def gate_14_newer_judgements(cur):
    name = ("a newer judgements edition replaces live and keeps history "
            "(notices too)")
    problems = []

    def body(e):
        e.completed()
        ok, d = real_history(e.cur, ZZ)
        if not ok:
            problems.append(f"before the new edition: {d}")
        e.judgements("20261007", changes=("2026-10-07", J_SEP, J_AUG))
        text, logged = e.ok(e.cur, ["load", "--only", "judgements",
                                    "--commit"])
        for n in ("judgements", "notices"):
            if f"{n} 2026-10-07: new" not in text:
                problems.append(f"{n}: the new edition is not 'new'")
            if tl.live_periods(e.cur, n) != [("2026-10-07", len(JROWS)
                                              if n == "judgements" else 1)]:
                problems.append(f"{n}: live is {tl.live_periods(e.cur, n)}")
            for p in (J_AUG, J_SEP):
                if not tl.editions(e.cur, n, p):
                    problems.append(f"{n} {p}: dropped from the editions")
            if [r[2] for r in tl.ledger(e.cur, n)][-1] != "new":
                problems.append(f"{n}: no 'new' ledger row")
        if tl.editions(e.cur, "judgements", J_SEP) != [(1, None, len(JROWS)),
                                                       (2, 1, len(JROWS))]:
            problems.append("the older edition's chain changed")
        ok, d = real_history(e.cur, ZZ)
        if not ok or f"history ['{J_AUG}', '{J_SEP}']" not in d:
            problems.append(f"after the new edition: {ok} {d[:120]}")
        ok, d = real_edition1_and_latest(e.cur, ZZ)
        if not ok:
            problems.append(f"live differs from the newest edition: {d}")
        if not logged.called or "judgements 2026-10-07 new" not in \
                logged.call_args[0][2]:
            problems.append("the new edition was not logged")
        # live holding a second period is not a pass
        e.cur.execute("SAVEPOINT x")
        e.cur.execute(f"UPDATE public.{ZZ['notices'].live_table} SET "
                      f"edition_date = '{J_SEP}'")
        ok, _ = real_history(e.cur, ZZ)
        e.cur.execute("ROLLBACK TO SAVEPOINT x")
        if ok:
            problems.append("live on an older notices edition passes")
    scenario(cur, body)
    mixed(14, name, not problems,
          "a newer judgements file stores its edition 1 and replaces the "
          "live judgements and notices with its own rows; the two older "
          "editions stay in the editions tables as history, status clean "
          "with history listed; live on an older edition fails"
          if not problems else "; ".join(problems[:3]), cur, real_history)


# ---------------------------------------------------------------- gate 15

def real_refreshed(cur, specs=None):
    """Nothing waits for refresh-latest: no left-over live rows of an older
    judgements edition, every live period 'current' (equal to its tip)."""
    specs = specs or REAL
    _need(cur, specs)
    bad = []
    for name in NAMES:
        spec = specs[name]
        newest, legacy = m.legacy_live_rows(cur, spec)
        if legacy:
            bad.append(f"{spec.live_table}: left-over live rows of "
                       f"{sorted(legacy)} (refresh-latest --accept-key-"
                       f"changes {newest})")
        try:
            tips, new, errors = core.latest_map(cur, spec)
        except SystemExit as e:       # two live periods: the engine halts
            bad.append(f"{spec.live_table}: {_first(e.code)[:120]}")
            continue
        if not tips and not new and not empty_latest(cur, spec):
            bad.append(f"{spec.live_table}: {EMPTY}")
        if empty_pending(cur, spec):
            bad.append(f"{spec.live_table}: {empty_pending(cur, spec)} is "
                       "recorded empty, live not yet emptied (pending: "
                       "refresh-latest --accept-key-changes "
                       f"{empty_pending(cur, spec)})")
        if new or errors:
            bad.append(f"{spec.live_table}: periods with no editions "
                       f"{sorted(map(str, new))}, chain errors {errors}")
        for p, tip in tips.items():
            kind = core.classify_period(cur, spec, p, tip)[0]
            if kind != "current":
                bad.append(f"{spec.live_table} {pe._p(p)}: {kind} (live is "
                           "not the latest edition)")
    return not bad, "; ".join(bad[:3]) if bad else (
        "no refresh pending: every live period equals its latest edition, "
        "no left-over rows of an older judgements edition")


def gate_15_refresh(cur):
    name = ("refresh-latest changes only the named live period and copies "
            "loaded_at; key changes need --accept-key-changes")
    problems = []
    D = datetime

    def body(e):
        e.completed()
        ok, d = real_refreshed(e.cur, ZZ)
        if not ok:
            problems.append(f"the completed state is not refreshed: {d}")
        reg = ZZ["register"]

        def live_hash(period):
            e.cur.execute(f"SELECT md5(string_agg(t::text, ',' ORDER BY "
                          f"t::text)) FROM public.{reg.live_table} t WHERE "
                          "snapshot_date = %s", (period,))
            return e.cur.fetchone()[0]
        rows = list(REG_ROWS[:3])
        rows[0] = rows[0][:3] + ("Profit",) + rows[0][4:]
        e.register("17 September 2026", rows=rows, modified=D(2026, 9, 18, 9))
        e.ok(e.cur, ["load", "--only", "register", "--acknowledge", SEP,
                     "--commit"])
        ok, d = real_refreshed(e.cur, ZZ)
        if ok or "pending" not in d:
            problems.append(f"a revision waiting for refresh passes: {d[:90]}")
        other = {p: live_hash(p) for p in (JUL, AUG)}
        e.cur.execute(f"SELECT registration_number, loaded_at FROM public."
                      f"{reg.live_table} WHERE snapshot_date = %s", (SEP,))
        loaded_before = dict(e.cur.fetchall())
        snap = state(e.cur)
        rc, text, logged = e.cmd(["refresh-latest"])
        if rc != 0 or state(e.cur) != snap or logged.called:
            problems.append("the refresh-latest preview wrote")
        e.ok(e.cur, ["refresh-latest", "--commit"])
        if {p: live_hash(p) for p in (JUL, AUG)} != other:
            problems.append("refresh-latest changed another live period")
        if core.rows_differing(e.cur, reg, SEP, 2):
            problems.append("live is not the revised edition")
        e.cur.execute(f"SELECT COUNT(*) FROM public.{reg.live_table} WHERE "
                      "snapshot_date = %s AND designation = 'Profit' AND "
                      "registration_number = '4001'", (SEP,))
        if e.cur.fetchone()[0] != 1:
            problems.append("the revised cell is not live")
        e.cur.execute(f"SELECT registration_number, loaded_at FROM public."
                      f"{reg.live_table} WHERE snapshot_date = %s", (SEP,))
        after = dict(e.cur.fetchall())
        e.cur.execute(f"SELECT DISTINCT loaded_at FROM public."
                      f"{reg.editions_table} WHERE snapshot_date = %s AND "
                      "edition = 2", (SEP,))
        ed2 = [r[0] for r in e.cur.fetchall()]
        moved = sorted(k for k in after if after[k] != loaded_before[k])
        if moved != ["4001"] or ed2 != [after["4001"]]:
            problems.append(f"loaded_at: rows whose loaded_at moved {moved}"
                            f" (only the revised row 4001 should, to the "
                            f"edition's {ed2})")
        ok, d = real_provenance(e.cur, ZZ)
        if not ok:
            problems.append(f"provenance after the refresh: {d}")
        # a judgements reissue that adds a row: a key change, applied only
        # with its period named
        e.judgements("20260930", jrows=JROWS + [jrow(
            reg="L0006", landlord="Fresh Homes")],
            modified=D(2026, 9, 30, 15))
        e.ok(e.cur, ["load", "--only", "judgements", "--acknowledge", J_SEP,
                     "--commit"])
        jud = ZZ["judgements"]
        snap = state(e.cur)
        for label, argv, needle in (
                ("no named period", ["refresh-latest", "--commit"],
                 "accept-key-changes"),
                ("another period named", ["refresh-latest",
                                          "--accept-key-changes", J_AUG,
                                          "--commit"], "no keys added")):
            ok, t = e.halts(argv, snap, needle)
            if not ok:
                problems.append(f"key changes with {label}: not halted "
                                f"cleanly: {t[-100:]}")
        if cnt(e.cur, jud.live_table) != len(JROWS):
            problems.append("the live judgements changed before the "
                            "refresh")
        e.ok(e.cur, ["refresh-latest", "--accept-key-changes", J_SEP,
                     "--commit"])
        if cnt(e.cur, jud.live_table) != len(JROWS) + 1:
            problems.append("the added judgement is not live")
        ok, d = real_refreshed(e.cur, ZZ)
        if not ok:
            problems.append(f"after the key change: {d}")
        # a register revision that adds a provider can be stored but not
        # refreshed (the register spec does not opt in to key changes):
        # refresh-latest halts and leaves live as it was
        e.register("17 September 2026", rows=rows + [(
            "New Homes", "5300", D(2026, 9, 1), "Profit", "Company", None)],
            modified=D(2026, 9, 19, 9))
        e.ok(e.cur, ["load", "--only", "register", "--acknowledge", SEP,
                     "--commit"])
        ok, t = e.halts(["refresh-latest", "--commit"], state(e.cur),
                        "only one of them")
        if not ok:
            problems.append(f"a register key change was not refused: "
                            f"{t[-100:]}")
    scenario(cur, body)
    mixed(15, name, not problems,
          "a revised register date waits for refresh-latest (preview writes "
          "nothing); the commit changes that period only, copies the "
          "edition's loaded_at to the revised row only and its provenance; "
          "an added judgement halts "
          "without --accept-key-changes naming its period (another period "
          "named is refused) and is then inserted; an added register "
          "provider is refused (the register does not opt in to key "
          "changes)" if not problems
          else "; ".join(problems[:3]), cur, real_refreshed)


# ---------------------------------------------------------------- gate 16

def hash_fields(h):
    """The note's scan-safe form of a sha256: two 32-hex halves in two
    labelled fields (the credential scan flags a 64-hex run)."""
    return f"sha256-first32={h[:32]} sha256-last32={h[32:]}"


def all_columns_sha(recs, name):
    """sha256 over every stored column but loaded_at: key, period, value
    columns and the three provenance columns (live's source_url,
    source_file, release_page_url are the records' file_*), NULL as ''."""
    keys, pc, values = m.KINDS[name]
    cols = keys + (pc,) + values + m.PROVENANCE
    ordered = sorted(recs, key=lambda r: tuple(m._cell(r[k]) for k in keys))
    lines = ["|".join(m._cell(r.get(c)) for c in cols) for r in ordered]
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


def values_sha(recs, name):
    return m.rows_content_sha(recs, reasons=False) if recs else \
        hashlib.sha256(b"").hexdigest()


HASH_KINDS = (("before-state", values_sha), ("before-state-all",
                                              all_columns_sha))


def note_lines_for(cur, specs=None):
    """The decision note's hash lines for the live tables as they stand
    (read-only): a before-state and a before-state-all line per live period
    of each table. Print them before migrate-legacy."""
    specs = specs or REAL
    out = []
    for name in NAMES:
        spec = specs[name]
        for p in _live_periods(cur, spec):
            recs = m.records(cur, spec, p)
            for label, fn in HASH_KINDS:
                out.append(f"{label} {spec.live_table} {p} rows={len(recs)} "
                           f"{hash_fields(fn(recs, name))}")
    return out


def note_hashes(note, label):
    """{(table, period): (rows, sha256)} from the decision note's
    `<label> <table> <period> rows=<n> sha256-first32=<32 hex>
    sha256-last32=<32 hex>` lines (the two halves joined), or {}."""
    p = Path(note)
    if not p.exists():
        return {}
    return {(mt.group(1), mt.group(2)): (int(mt.group(3)),
                                         mt.group(4) + mt.group(5))
            for mt in re.finditer(
                rf"(?m)^{re.escape(label)} ([a-z0-9_]+) ([0-9]{{4}}-[0-9]{{2}}"
                r"-[0-9]{2}) rows=([0-9]+) sha256-first32=([0-9a-f]{32}) "
                r"sha256-last32=([0-9a-f]{32})\s*$",
                p.read_text(encoding="utf-8"))}


def real_edition1_hash(cur, specs=None, note=NOTE, survey=None):
    """Edition 1 'as loaded' equals the before-state lines of the decision
    note: for every surveyed live period (survey {table name: [periods]},
    default m.LEGACY_LIVE) the note has a before-state line whose row count
    and hash (key + value columns) equal edition 1's (a before-state-all
    line, if present, must equal it on every stored column too); the 15
    live judgements still tagged with the 12 August date are the rows of its
    edition 1 whose key is not in the next edition's, so that line is
    checked against that subset; no line is for a period without edition 1;
    the lines are scan-safe."""
    specs = specs or REAL
    _need(cur, specs)
    survey = survey or {n: sorted(m.LEGACY_LIVE[n]["periods"])
                        for n in NAMES}
    bad, parts = [], []
    have_note = Path(note).exists()
    if not have_note:
        bad.append(f"decision note {Path(note).name} not written")
    recorded = {label: note_hashes(note, label) for label, _ in HASH_KINDS}
    tables = {specs[n].live_table for n in NAMES}
    for label in recorded:
        for (t, p) in recorded[label]:
            if t not in tables:
                bad.append(f"the note has a {label} line for {t}, which is "
                           "not an S24 table")
    if have_note and re.search(r"[0-9a-fA-F]{40,}",
                               Path(note).read_text(encoding="utf-8")):
        bad.append("the note holds a run of 40 or more hex characters (the "
                   "credential scan would flag it)")
    for name in NAMES:
        spec = specs[name]
        cur.execute(f"SELECT DISTINCT {spec.period_col} FROM public."
                    f"{spec.editions_table} WHERE edition = 1 AND "
                    "release_label LIKE %s ORDER BY 1",
                    (core.AS_LOADED_LABEL + "%",))
        migrated = [str(r[0]) for r in cur.fetchall()]
        if not migrated:
            bad.append(f"{spec.editions_table}: no edition 1 'as loaded' "
                       "(migrate-legacy not run)")
        for t, p in list(recorded["before-state"]) + list(
                recorded["before-state-all"]):
            if t == spec.live_table and p not in migrated:
                bad.append(f"the note has a line for {t} {p}, which has no "
                           "edition 1 as loaded")
        for p in survey[name]:
            if p not in migrated:
                bad.append(f"{spec.live_table} {p}: surveyed live period "
                           "with no edition 1 as loaded")
                continue
            recs = m.records(cur, spec, p, 1)
            subset = None
            later = [q for q in migrated if q > p]
            if spec.live_newest_period_only and later:
                nxt = {tuple(r[k] for k in spec.key_cols)
                       for r in m.records(cur, spec, min(later), 1)}
                subset = [r for r in recs
                          if tuple(r[k] for k in spec.key_cols) not in nxt]
            for label, fn in HASH_KINDS:
                got = recorded[label].get((spec.live_table, p))
                if got is None:
                    if label == "before-state" and have_note:
                        bad.append(f"the decision note records no "
                                   f"'before-state {spec.live_table} {p} "
                                   "rows=.. sha256-first32=.. "
                                   "sha256-last32=..' line")
                    continue
                cands = [(len(recs), fn(recs, name))]
                if subset is not None:
                    cands.append((len(subset), fn(subset, name)))
                if got not in cands:
                    bad.append(f"{label} {spec.live_table} {p}: the note's "
                               f"{got[0]} rows {got[1][:16]}.. equal "
                               f"neither edition 1's {cands[0][0]} rows "
                               f"{cands[0][1][:16]}.."
                               + (f" nor the {cands[1][0]} left-over rows "
                                  f"{cands[1][1][:16]}.." if subset is not None
                                  else ""))
            parts.append(f"{spec.live_table} {p} {len(recs):,} rows")
    return not bad, "; ".join(bad[:3]) if bad else (
        "edition 1 of every surveyed live period hashes to the note's "
        "before-state line (" + "; ".join(parts) + ")")


def gate_16_before_state_hash(cur):
    name = ("edition 1 equals the before-state hash lines in the decision "
            "note")
    problems = []

    def body(e):
        e.seed(e.cur)
        lines = note_lines_for(e.cur, ZZ)
        survey = {n: sorted(e.legacy[n]["periods"]) for n in NAMES}
        if len(lines) != 2 * sum(len(v) for v in survey.values()):
            problems.append(f"{len(lines)} note lines for "
                            f"{sum(len(v) for v in survey.values())} live "
                            "periods")
        for ln in lines:
            if re.search(r"[0-9a-fA-F]{40,}", ln) or not re.fullmatch(
                    r"before-state(-all)? [a-z0-9_]+ \d{4}-\d{2}-\d{2} "
                    r"rows=\d+ sha256-first32=[0-9a-f]{32} "
                    r"sha256-last32=[0-9a-f]{32}", ln):
                problems.append(f"a note line is not scan-safe: {ln[:60]}")
        e.migrate(e.cur)
        with tempfile.TemporaryDirectory() as tmp:
            note = Path(tmp) / "note.md"

            def check(text, label, needle=None, ok_expected=False):
                note.write_text(text, encoding="utf-8")
                ok, d = real_edition1_hash(e.cur, ZZ, note, survey)
                if ok != ok_expected or (needle and needle not in d):
                    problems.append(f"{label}: {ok} {d[:110]}")
            full = "intro\n" + "\n".join(lines) + "\n"
            check(full, "the note written before the migration",
                  ok_expected=True)
            h = next(x for x in lines if x.startswith("before-state ")
                     and "_reg_live" in x)
            sha = h.split("sha256-first32=")[1][:32]
            check(full.replace(sha, "0" * 32), "a wrong hash", "equal "
                  "neither")
            rows = re.search(r"rows=(\d+)", h).group(1)
            check(full.replace(h, h.replace(f"rows={rows}",
                                            f"rows={int(rows) + 1}")),
                  "a wrong row count", "equal neither")
            jaug = next(x for x in lines if x.startswith("before-state ")
                        and f"_jud_live {J_AUG}" in x)
            check(full.replace(jaug + "\n", ""), "a missing 12 August line",
                  "records no")
            check(full.replace(jaug, jaug.replace("before-state ",
                                                  "before-state-all ", 1))
                  .replace(next(x for x in lines if x.startswith(
                      "before-state-all ") and f"_jud_live {J_AUG}" in x)
                      + "\n", ""),
                  "a before-state-all line standing for a before-state one",
                  "records no")
            check(full + re.sub(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", "2030-01-01", h) + "\n",
                  "a line for a period with no edition 1", "no edition 1")
            check(full + h.replace("zz_s24_reg_live", "other_table") + "\n",
                  "a line for another table", "not an S24 table")
            check(full + "x" * 3 + "a" * 64 + "\n", "a 64-hex run",
                  "40 or more hex")
            ok, d = real_edition1_hash(e.cur, ZZ, Path(tmp) / "none.md",
                                       survey)
            if ok or "not written" not in d:
                problems.append("a missing note passes")
            note.write_text(full, encoding="utf-8")
            ok, d = real_edition1_hash(e.cur, ZZ, note, {
                **survey, "register": survey["register"] + ["2026-12-01"]})
            if ok or "no edition 1 as loaded" not in d:
                problems.append("a surveyed period without edition 1 passes")
        e.cur.execute("SAVEPOINT x")
        e.cur.execute(f"UPDATE public.{ZZ['register'].live_table} SET notes "
                      "= 'x' WHERE registration_number = '4001'")
        if note_lines_for(e.cur, ZZ) == lines:
            problems.append("a changed live cell does not change the lines")
        e.cur.execute("ROLLBACK TO SAVEPOINT x")
    scenario(cur, body)
    mixed(16, name, not problems,
          "the lines written from a seeded legacy state before the "
          "migration are scan-safe and equal edition 1 afterwards (the "
          "12 August judgements line against the 1 left-over row of that "
          "edition's file); a wrong hash or row count, a missing line "
          "(before-state-all does not stand in), a line for a period or "
          "table with no edition 1, a 64-hex run, a missing note and a "
          "surveyed period without edition 1 fail" if not problems
          else "; ".join(problems[:3]), cur, real_edition1_hash)


def print_note_lines():
    """The decision note's hash lines for the live tables as they stand,
    read on a read-only connection (nothing is written)."""
    conn = get_readonly_conn()
    try:
        with conn.cursor() as cur:
            for line in note_lines_for(cur, REAL):
                print(line)
    finally:
        conn.rollback()
        conn.close()


# ---------------------------------------------------------------- gate 17

def real_migration_proof(cur, specs=None, files=None, dirs=None, run=None,
                         run_rows=None):
    """The migration proofs re-run read-only from the held files on disk
    (found by sha256): each file read with the old build's rules reproduces
    every cell of its edition 1 'as loaded' (key and value columns; 0
    differences) and the new parse gives the same values; the register's
    edition 1 null_reasons are the new parse's; every held period's edition
    1 is labelled as loaded; the August register's states its file cannot be
    obtained; run-log 96's rows_written equals the 24 July register + the
    12 August judgements + notices of edition 1."""
    specs = specs or REAL
    files = m.LEGACY_FILES if files is None else files
    run = m.LEGACY_RUN if run is None else run
    run_rows = m.legacy_run_rows if run_rows is None else run_rows
    _need(cur, specs)
    found = files_by_sha([m.RAW_DIR] if dirs is None else dirs)
    bad, cells, parts = [], 0, []
    for (part, p), facts in sorted(files.items()):
        path = found.get(facts["sha256"])
        if path is None:
            bad.append(f"{part} {p}: no file with sha256 "
                       f"{facts['sha256'][:16]} on disk")
            continue
        try:
            if part == "register":
                ident = m.register_identity(path, p)
                built = {"register": (
                    m.build_register(ident["rows"], ident, rules="old"),
                    m.build_register(ident["rows"], ident))}
            else:
                ident = m.judgements_identity(path, facts["name"], None)
                built = {
                    "judgements": (
                        m.build_judgements(ident["judgement_rows"], ident,
                                           rules="old"),
                        m.build_judgements(ident["judgement_rows"], ident)),
                    "notices": (
                        m.build_notices(ident["notice_rows"], ident,
                                        rules="old"),
                        m.build_notices(ident["notice_rows"], ident))}
        except ValueError as e:
            bad.append(f"{part} {p}: {_first(e)}")
            continue
        for n, (old, new) in built.items():
            spec = specs[n]
            ed1 = m.records(cur, spec, p, 1)
            if not ed1:
                bad.append(f"{n} {p}: no edition 1")
                continue
            diffs = m._proof(ed1, old, n) + m._proof(old, new, n)
            if n == "register":
                why = {r["registration_number"]: r["null_reasons"]
                       for r in new}
                diffs += [("null_reasons", r["registration_number"])
                          for r in ed1
                          if r["null_reasons"] != why.get(
                              r["registration_number"])]
            if diffs:
                bad.append(f"{n} {p}: {len(diffs)} differences between "
                           f"edition 1, {path.name} read with the old rules "
                           f"and the new parse, e.g. {diffs[:2]}")
            cells += len(ed1) * len(m.KINDS[n][2])
            parts.append(f"{n} {p} {len(ed1):,} rows")
    reg = specs["register"]
    cur.execute(f"SELECT DISTINCT {reg.period_col}, release_label, "
                f"source_file FROM public.{reg.editions_table} WHERE "
                "edition = 1")
    for p, label, src in cur.fetchall():
        if not str(label).startswith(core.AS_LOADED_LABEL):
            bad.append(f"register {p}: edition 1 is not the 'as loaded' one")
        if str(p) == m.AUGUST_REGISTER["period"] and \
                "file not obtainable" not in str(label):
            bad.append(f"register {p}: edition 1 does not say the file "
                       "cannot be obtained")
    got = run_rows(cur)
    want = sum(len(m.records(cur, specs[n], p, 1)) for n, p in (
        ("register", run["register"]), ("judgements", run["judgements"]),
        ("notices", run["judgements"])))
    if got != run["rows"] or got != want:
        bad.append(f"run-log {run.get('id')} wrote {got} rows; edition 1 of "
                   f"the 24 July register and the 12 August judgements and "
                   f"notices holds {want}; surveyed {run['rows']}")
    return not bad, "; ".join(bad[:3]) if bad else (
        f"{cells:,} cells of edition 1 ({', '.join(parts)}) equal the held "
        f"files read now with the old rules and the new parse (0 "
        f"differences); run-log {run.get('id')} wrote {got:,} rows = edition "
        "1 of the 24 July register + the 12 August judgements and notices")


def gate_17_migration_proof(cur):
    name = ("the migration proofs re-run read-only from the held files on "
            "disk; seeded: a planted difference stops migrate-legacy and "
            "stores nothing")
    problems = []

    def ok_world(e):
        e.migrated()
        run = dict(m.LEGACY_RUN, rows=e.run_rows, register=JUL,
                   judgements=J_AUG)
        kw = dict(run=run, run_rows=lambda c: e.run_rows)
        ok, d = real_migration_proof(e.cur, ZZ, e.held, [e.root], **kw)
        if not ok or "0 differences" not in d:
            problems.append(f"the proof does not pass: {d[:120]}")
        raw = e.root / "raw"
        raw.mkdir()
        for k, p in e.paths.items():
            (raw / p.name).write_bytes(p.read_bytes())
        ok, d = real_migration_proof(e.cur, ZZ, e.held, [raw], **kw)
        if not ok:
            problems.append(f"the proof does not re-run from a directory: "
                            f"{d[:100]}")
        ok, d = real_migration_proof(e.cur, ZZ, {
            **e.held, ("register", SEP): {**e.held[("register", SEP)],
                                          "sha256": "0" * 64}}, [raw], **kw)
        if ok or "no file" not in d:
            problems.append("a file with another sha256 was accepted")
        victim = raw / e.paths[("judgements", J_SEP)].name
        victim.write_bytes(victim.read_bytes() + b"\0")
        ok, d = real_migration_proof(e.cur, ZZ, e.held, [raw], **kw)
        if ok or "no file" not in d:
            problems.append("a changed file on disk was accepted")
        other = write_register(e.root / "other" / "o.xlsx", rows=REG_ROWS[:2])
        held = {**e.held, ("register", SEP): {
            **e.held[("register", SEP)], "sha256": m.content_sha256(other)}}
        ok, d = real_migration_proof(e.cur, ZZ, held, [other.parent, e.root],
                                     **kw)
        if ok or "differences" not in d:
            problems.append("a file that is not edition 1 passed")
        ok, d = real_migration_proof(e.cur, ZZ, e.held, [e.root],
                                     run=dict(run, rows=run["rows"] + 1),
                                     run_rows=kw["run_rows"])
        if ok or "run-log" not in d:
            problems.append("a wrong run-log count passed")

    def stops(label, mutate, needle, **kw):
        def body(e):
            e.seed(e.cur, plant=kw.pop("plant", False))
            held = mutate(e) or e.held
            before = state(e.cur)
            try:
                e.migrate(e.cur, held_files=held, **kw)
                problems.append(f"{label}: migrate-legacy did not stop")
            except SystemExit as ex:
                if needle not in str(ex.code):
                    problems.append(f"{label}: stopped without {needle!r}: "
                                    f"{_first(ex.code)[:90]}")
            if state(e.cur) != before:
                problems.append(f"{label}: something was stored")
        scenario(cur, body)

    def wrong_sha(e):
        held = {k: dict(v) for k, v in e.held.items()}
        held[("judgements", J_SEP)]["sha256"] = "0" * 32 + "f" * 32
        return held
    scenario(cur, ok_world)
    stops("a planted live difference", lambda e: None, "proof failed",
          plant=True)
    stops("a wrong file sha256", wrong_sha, "sha256")

    def short_live(e):
        e.cur.execute(f"DELETE FROM public.{ZZ['register'].live_table} WHERE "
                      f"snapshot_date = '{AUG}' AND registration_number = "
                      "'4001'")
    stops("a live table not as surveyed", short_live, "not as surveyed")

    def twice(e):
        e.migrate(e.cur)
    stops("editions that already hold rows", twice, "already holds")
    mixed(17, name, not problems,
          "the proof passes on the seeded held files (0 differences, "
          "run-log count equal) and re-runs from a directory by sha256; "
          "another sha256, a changed file, a file that is not edition 1 and "
          "a wrong run-log count fail; a planted live difference, a wrong "
          "held sha256, a live table not as surveyed and a second "
          "migration each stop migrate-legacy with nothing stored"
          if not problems else "; ".join(problems[:3]), cur,
          real_migration_proof)


# ---------------------------------------------------------------- gate 18

def gate_18_rerun(cur):
    name = "rerun idempotent: a second load changes nothing"
    problems = []

    def body(e):
        e.completed()
        sep = e.paths[("register", SEP)]
        jsep = e.paths[("judgements", J_SEP)]
        snap, led = tables_state(e.cur), [
            len(tl.ledger(e.cur, n)) for n in NAMES]
        for argv in (["load", "--commit"],
                     ["load", "--only", "register", "--file", sep,
                      "--title-date", SEP, "--commit"],
                     ["load", "--only", "register", "--file", sep,
                      "--title-date", SEP, "--no-page", "--commit"],
                     ["load", "--only", "judgements", "--file", jsep,
                      "--commit"],
                     ["load", "--only", "judgements", "--file", jsep,
                      "--no-page", "--commit"],
                     ["refresh-latest", "--commit"]):
            rc, text, logged = e.cmd(argv)
            if rc != 0 or tables_state(e.cur) != snap:
                problems.append(f"{' '.join(map(str, argv[:3]))}: changed "
                                f"the editions or live: rc {rc}")
        for n, before in zip(NAMES, led):
            if any(o != "unchanged" for _, _, o, _ in
                   tl.ledger(e.cur, n)[before:]):
                problems.append(f"{n}: a re-read ledger row is not "
                                "'unchanged'")
        before = state(e.cur)
        rc, text, logged = e.cmd(["load", "--commit"])
        if rc != 0 or state(e.cur) != before or logged.called:
            problems.append("a ledger-skipped load changed state or logged")
        rc, text, logged = e.cmd(["load", "--only", "register", "--recheck",
                                  SEP, "--commit"])
        if rc != 0 or tl.editions(e.cur, "register", SEP) != [(1, None, 3)]:
            problems.append(f"a recheck stored an edition: "
                            f"{tl.editions(e.cur, 'register', SEP)}")
        ok, d = real_refreshed(e.cur, ZZ)
        if not ok:
            problems.append(f"after the reruns: {d}")
    scenario(cur, body)
    report(18, name, not problems,
           "a second load by the page, --file and --file --no-page (the "
           "register and the judgements) and a refresh-latest stored no "
           "edition and changed no live row (re-reads add 'unchanged' "
           "ledger rows only); a ledger-skipped load logs nothing; a "
           "recheck stores no edition" if not problems
           else "; ".join(problems[:3]))


# ---------------------------------------------------------------- gate 19

def gate_19_stranded(cur):
    name = "stranded period repair (edition held, live rows missing)"
    problems = []

    def body(e):
        e.completed()
        reg = ZZ["register"]
        e.cur.execute(f"DELETE FROM public.{reg.live_table} WHERE "
                      f"snapshot_date = '{SEP}'")
        prof = m.profile("register")
        if pe.status(e.cur, prof)["live_missing"] != [SEP]:
            problems.append("status does not name the stranded period")
        rc, text, _ = e.cmd(["load", "--only", "register", "--commit"])
        if rc != 0 or "live-missing" not in text or \
                tl.editions(e.cur, "register", SEP) != [(1, None, 3)] or \
                tl.count(e.cur, reg.live_table, "WHERE snapshot_date = %s",
                         (SEP,)) != 3:
            problems.append(f"repair failed: rc {rc}: {text[-120:]}")
        if "live-missing" not in [o for _, _, o, _ in tl.ledger(
                e.cur, "register")]:
            problems.append("no live-missing ledger row")
        ok, d = real_provenance(e.cur, ZZ)
        if not ok or not pe.status(e.cur, prof)["ok"] or \
                m.live_equals_tip(e.cur, reg, [SEP]):
            problems.append(f"the repaired rows do not equal the tip: {d}")
    scenario(cur, body)
    report(19, name, not problems,
           f"a stranded {SEP} register date was rebuilt from its edition 1 "
           "with no new edition, the tip's provenance and a live-missing "
           "ledger row" if not problems else "; ".join(problems[:3]))


# ---------------------------------------------------------------- gate 20

def gate_20_restore(cur):
    name = "restore-edition round trip, all three tables"
    problems = []

    def body(e):
        e.completed()
        for n, p in (("register", AUG), ("judgements", J_SEP),
                     ("notices", J_SEP)):
            s = ZZ[n]
            label = f"{n} {p}"
            snap = state(e.cur)
            for extra in ([], ["--simulate"]):
                rc, text, logged = e.cmd(["restore-edition", s.live_table, p,
                                          "1"] + extra)
                if rc != 0 or state(e.cur) != snap or logged.called:
                    problems.append(f"{label}: restore-edition "
                                    f"{extra or 'preview'} wrote")
            live_before = tables_state(e.cur)[::2]
            rc, text, logged = e.cmd(["restore-edition", s.live_table, p, "1",
                                      "--commit"])
            eds = tl.editions(e.cur, n, p)
            if rc != 0 or eds[-1][:2] != (3, 2) or not logged.called:
                problems.append(f"{label}: restore --commit: rc {rc}, "
                                f"{eds}")
            e.cur.execute(f"SELECT DISTINCT release_label FROM public."
                          f"{s.editions_table} WHERE {s.period_col} = %s AND "
                          "edition = 3", (p,))
            if e.cur.fetchall() != [("restored from edition 1",)]:
                problems.append(f"{label}: the restored edition is not "
                                "labelled")
            if tables_state(e.cur)[::2] != live_before:
                problems.append(f"{label}: restore-edition touched live "
                                "before refresh-latest")
            e.ok(e.cur, ["refresh-latest", "--commit"])
            if core.rows_differing(e.cur, s, p, 3):
                problems.append(f"{label}: refresh-latest did not apply the "
                                "restored edition")
            if m.rows_content_sha(m.records(e.cur, s, p, 1)) != \
                    m.rows_content_sha(m.records(e.cur, s, p, 3)):
                problems.append(f"{label}: edition 3 does not hold edition "
                                "1's rows and reasons")
            for argv in (["restore-edition", s.live_table, p, "3", "--commit"],
                         ["restore-edition", s.live_table, p, "9",
                          "--commit"]):
                ok, _ = e.halts(argv, state(e.cur))
                if not ok:
                    problems.append(f"{label}: {' '.join(argv[2:4])} did "
                                    "not halt")
        for n in NAMES:
            if not pe.status(e.cur, m.profile(n))["ok"]:
                problems.append(f"{n}: status not clean after the restores")
    scenario(cur, body)
    report(20, name, not problems,
           "edition 1 stored as edition 3 ('restored from edition 1') for "
           "the register (August), the judgements and the notices, "
           "applied by refresh-latest with the same rows and reasons; the "
           "tip and a missing edition are refused; a preview and a "
           "simulation write nothing" if not problems
           else "; ".join(problems[:3]))


# ---------------------------------------------------------------- gate 21

def secret_values():
    """Values of environment/.env entries that look like secrets (names with
    PASSWORD, KEY, TOKEN or SECRET), for the leak check only. Never printed."""
    out = {}
    for src in (ENV, os.environ):
        for k, v in src.items():
            if (re.search(r"PASSWORD|KEY|TOKEN|SECRET", k.upper())
                    and isinstance(v, str) and len(v.strip()) >= 8):
                out[k] = v.strip()
    return out


def gate_21_no_network_no_secrets(cur):
    name = ("no network, no secret in source or output, nothing written "
            "outside the rolled-back transaction; a download never "
            "overwrites a same-named file with different content")
    files = [Path(__file__).resolve(), LOADER, HERE / "period_editions.py",
             HERE / "tests" / "test_s24_rsh_loader.py",
             HERE / "tests" / "test_s24_rsh_pure.py"]
    attempts = []

    def refuse(*a, **k):
        attempts.append(a[:1])
        raise OSError("network blocked by the verify script")
    cur.execute("SELECT COUNT(*) FROM public.pipeline_run_log")
    runs = cur.fetchone()[0]
    probe = [f"public.{s.editions_table}" for s in REAL.values()] + [
        f"public.{ledger_name(s)}" for s in REAL.values()]
    cur.execute("SELECT " + ", ".join("to_regclass(%s)" for _ in probe),
                probe)
    real_before = cur.fetchone()

    def resp(body, url):
        return mock.Mock(content=body, url=url, raise_for_status=lambda: None)

    def body(c):
        with env(c) as e, tl.specs():
            with mock.patch.object(socket.socket, "connect", refuse), \
                    mock.patch.object(socket, "create_connection", refuse), \
                    mock.patch.object(socket, "getaddrinfo", refuse):
                e.completed()
                rcs = [e.cmd(a)[0] for a in (
                    ["status"], ["refresh-latest"], ["load"],
                    ["load", "--commit"], ["ddl"])]
                a = write_register(e.root / "src" / "a.xlsx").read_bytes()
                b = write_register(e.root / "src" / "b.xlsx",
                                   rows=REG_ROWS[:2]).read_bytes()
                dest = e.root / "dl"
                url = "https://assets.example/media/1/List.xlsx"
                sess = mock.Mock()
                with contextlib.redirect_stdout(io.StringIO()), \
                        mock.patch.object(m, "MIN_FILE_BYTES", 1000):
                    sess.get.return_value = resp(a, url)
                    p1, _ = m.fetch(url, dest, session=sess)
                    sha1 = m.content_sha256(p1)
                    p2, _ = m.fetch(url, dest, session=sess)
                    sess.get.return_value = resp(b, url)
                    p3, _ = m.fetch(url, dest, session=sess)
                    sess.get.return_value = resp(b"<html>no</html>", url)
                    try:
                        m.fetch(url, dest, session=sess)
                        html_halts = False
                    except SystemExit:
                        html_halts = True
                kept = (p1 == dest / "List.xlsx" and
                        m.content_sha256(p1) == sha1 and p2 == p1
                        and p3 != p1 and p3.name.startswith("List-")
                        and sess.get.call_count == 4 and html_halts)
            return rcs, kept
    try:
        rcs, kept = _in_savepoint(cur, body)
    except Exception as ex:  # a blocked socket or any other failure
        return report(21, name, False, f"{type(ex).__name__}: {ex}")
    cur.execute("SELECT COUNT(*) FROM public.pipeline_run_log")
    runs_after = cur.fetchone()[0]
    cur.execute("SELECT " + ", ".join("to_regclass(%s)" for _ in probe),
                probe)
    real_after = cur.fetchone()
    pat = re.compile(r"""(?i)(api[_-]?key|password|token|secret)['"]?\s*[:=]"""
                     r"""\s*['"][A-Za-z0-9]{16,}""")
    texts = {f.name: f.read_text(encoding="utf-8") for f in files}
    literal = [n for n, t in texts.items() if pat.search(t)]
    secrets = secret_values()
    in_src = sorted({k for k, v in secrets.items()
                     for t in texts.values() if v in t})
    in_out = sorted({k for k, v in secrets.items()
                     for line in OUTPUT if v in line})
    ok = (not attempts and rcs[0] == 0 and rcs[1:] == [0, 0, 0, 0] and kept
          and runs == runs_after and real_before == real_after
          and not literal and not in_src and not in_out)
    report(21, name, ok, f"socket attempts {len(attempts)}; status, "
           f"refresh-latest, load, load --commit and ddl through stubs "
           f"returned {rcs}; a same-named different download was saved "
           f"beside it, the first untouched, a re-download kept, an HTML "
           f"error page refused={kept}; run-log rows {runs}->{runs_after}; "
           f"real editions tables and ledgers unchanged="
           f"{real_before == real_after}; secret-like literal in source "
           f"{literal or 'none'}; {len(secrets)} secret-named settings "
           f"checked against {len(files)} files and {len(OUTPUT)} output "
           f"lines: in source {in_src or 'none'}, in output "
           f"{in_out or 'none'}")


# ---------------------------------------------------------------- gate 22

def _empty_rows(cur, period):
    led = ledger_name(ZZ["notices"])
    cur.execute(f"SELECT outcome, edition, acknowledgement, source_file, "
                f"file_sha256 FROM public.{led} WHERE edition_date = %s "
                "ORDER BY id", (period,))
    return cur.fetchall()


def gate_22_empty_notices(cur):
    name = ("a notices-free judgements issue (Scott's Option C, 2026-10-12): "
            "stored only with --acknowledge-no-notices, the notices recorded "
            "empty, live notices emptied by refresh-latest, history kept")
    problems = []
    OCT, NOV = "2026-10-07", "2026-10-21"
    J = ("judgements", "notices")
    NOT = ZZ["notices"]
    load = ["load", "--only", "judgements", "--commit"]
    ack = ["--acknowledge-no-notices", OCT]
    D = datetime

    def body(e):
        e.completed()
        path = e.judgements("20261007", nrows=[],
                            changes=(OCT, J_SEP, J_AUG))
        url = e.last_url
        # unacknowledged: REJECTED, nothing stored, a partial run-log row
        bad = rejects(e, load, "--acknowledge-no-notices", J, OCT)
        if bad:
            problems.append(f"unacknowledged: {bad}")
        # a stray flag halts with nothing stored
        ok, t = e.halts(load + ["--acknowledge-no-notices", "2026-10-08"],
                        state(e.cur), "--acknowledge-no-notices")
        if not ok:
            problems.append(f"a stray --acknowledge-no-notices: {t[-100:]}")
        # the preview and --simulate with the flag write nothing
        snap = state(e.cur)
        for extra in ([], ["--simulate"]):
            rc, text, logged = e.cmd(["load", "--only", "judgements"] + ack
                                     + extra)
            if rc != 0 or state(e.cur) != snap or logged.called or \
                    f"notices {OCT}: empty" not in text:
                problems.append(f"load {' '.join(extra) or 'preview'} with "
                                f"the flag wrote or did not say empty: rc "
                                f"{rc}: {text[-120:]}")
        ed = core.period_hashes(e.cur, NOT, "editions")
        text, logged = e.ok(e.cur, load + ack)
        if tl.editions(e.cur, "judgements", OCT) != [(1, None, len(JROWS))] \
                or tl.live_periods(e.cur, "judgements") != [(OCT, len(JROWS))]:
            problems.append("the judgements of the notices-free issue did not "
                            "load as usual")
        if tl.editions(e.cur, "notices", OCT):
            problems.append("an empty notices issue stored editions rows")
        got = _empty_rows(e.cur, OCT)
        want = ("empty", None)
        if len(got) != 1 or got[0][:2] != want or \
                f"--acknowledge-no-notices {OCT}" not in (got[0][2] or "") or \
                not str(got[0][3]).startswith(url) or \
                got[0][4] != m.content_sha256(path):
            problems.append(f"the empty ledger row is {got}")
        if core.period_hashes(e.cur, NOT, "editions") != ed:
            problems.append("the notices editions changed at the load")
        notes = logged.call_args[0][2] if logged.called else ""
        if f"notices {OCT} empty" not in notes or \
                f"ACKNOWLEDGED NO NOTICES {OCT}" not in notes:
            problems.append(f"run-log notes: {notes[:160]}")
        # pending until refresh-latest: live notices are still September's
        if tl.live_periods(e.cur, "notices") != [(J_SEP, 1)]:
            problems.append("the load changed the live notices")
        ok, d = real_refreshed(e.cur, ZZ)
        if ok or "recorded empty" not in d:
            problems.append(f"the pending empty state passes: {d[:100]}")
        rc, text, _ = e.cmd(["status"])
        if rc != 1 or f"--accept-key-changes {OCT}" not in text:
            problems.append("status does not name the pending refresh")
        snap = state(e.cur)
        rc, text, logged = e.cmd(["refresh-latest"])
        if rc != 0 or state(e.cur) != snap or logged.called or \
                f"{OCT} is recorded empty" not in text:
            problems.append(f"the refresh preview: rc {rc}: {text[-120:]}")
        ok, t = e.halts(["refresh-latest", "--commit"], snap,
                        f"--accept-key-changes {OCT}")
        if not ok:
            problems.append(f"refresh without the period named: {t[-100:]}")
        e.ok(e.cur, ["refresh-latest", "--accept-key-changes", OCT,
                     "--commit"])
        if tl.live_periods(e.cur, "notices"):
            problems.append("live notices not empty after the refresh")
        if core.period_hashes(e.cur, NOT, "editions") != ed or not all(
                tl.editions(e.cur, "notices", p) for p in (J_AUG, J_SEP)):
            problems.append("the notices history changed")
        for fn in (real_edition1_and_latest, real_provenance, real_history,
                   real_refreshed):
            ok, d = fn(e.cur, ZZ)
            if not ok:
                problems.append(f"{fn.__name__} on the empty state: {d}")
        st = pe.status(e.cur, m.profile("notices"))
        if not st["ok"] or st["history"] != [J_AUG, J_SEP] or \
                st["empty"] != [OCT]:
            problems.append(f"notices status after the refresh: {st}")
        # a byte-identical re-read is unchanged
        snap = state(e.cur)
        text, logged = e.ok(e.cur, load)
        if "already in both ledgers" not in text or logged.called or \
                state(e.cur) != snap:
            problems.append("a byte-identical re-read was not skipped")
        # a later issue with notices restores live
        e.judgements("20261021", changes=(NOV, OCT, J_SEP))
        text, _ = e.ok(e.cur, load)
        if tl.live_periods(e.cur, "notices") != [(NOV, 1)] or \
                f"notices {NOV}: new" not in text:
            problems.append(f"a later issue did not restore live: "
                            f"{tl.live_periods(e.cur, 'notices')}")
        ok, d = real_history(e.cur, ZZ)
        st = pe.status(e.cur, m.profile("notices"))
        if not ok or st["empty"] != [OCT] or st["history"] != [J_AUG, J_SEP]:
            problems.append(f"after the later issue: {ok} {d[:100]} {st}")

    def headers(e):
        e.completed()
        snap = state(e.cur)
        nohead = write_book(e.root / "nh" / "20261007_RJN_published.xlsx", [
            ("Regulatory Judgements", [list(m.JUDGEMENT_HEADERS)]
             + list(JROWS), "visible"),
            ("Enforcement Notices", [], "visible")], D(2026, 10, 7, 11))
        changed = write_judgements(
            e.root / "ch" / "20261007_RJN_published.xlsx", nrows=[],
            nheader=["Reg Code", "Provider"], modified=D(2026, 10, 7, 11))
        for label, p, needle in (("no header", nohead, "is empty"),
                                 ("a changed header", changed,
                                  "header changed")):
            ok, t = e.halts(["load", "--only", "judgements", "--file", p,
                             "--no-page", "--commit"] + ack, snap, needle,
                            rc_ok=("halt",))
            if not ok:
                problems.append(f"{label}: not an identity halt: {t[-100:]}")
        # an empty reissue of a held issue whose notices have rows
        e.judgements("20260930", nrows=[], modified=D(2026, 9, 30, 15))
        bad = rejects(e, load + ["--acknowledge", J_SEP,
                                 "--acknowledge-no-notices", J_SEP],
                      "held with rows", J, J_SEP)
        if bad:
            problems.append(f"an empty reissue of a held issue: {bad}")
    scenario(cur, body)
    scenario(cur, headers)
    report(22, name, not problems,
           "unacknowledged it is REJECTED (nothing stored, partial run-log "
           "row; --acknowledge and --accept-reissue do not release it); the "
           "preview and --simulate write nothing; acknowledged, the "
           "judgements load as usual and the notices period is one 'empty' "
           "ledger row naming the acknowledgement, the file's URL and sha256 "
           "(no editions row, the run log names it); status and the real "
           "checks name the pending refresh, which halts unless "
           "--accept-key-changes names the issue, then empties live with the "
           "notices history unchanged and every real check passing; a "
           "byte-identical re-read is skipped; a later issue with notices "
           "restores live; a missing or changed header halts as an identity "
           "failure; an empty reissue of a held issue is rejected"
           if not problems else "; ".join(problems[:3]))


# --------------------------------------------------------------------- main

GATES = (gate_1_table_shape_and_immutability, gate_2_edition1_and_latest,
         gate_3_geography, gate_4_provenance, gate_5_unassessed,
         gate_6_hidden_sheet, gate_7_identity, gate_8_older_file,
         gate_9_stop_conditions, gate_10_one_transaction,
         gate_11_ledger_skip, gate_12_preview, gate_13_reread,
         gate_14_newer_judgements, gate_15_refresh,
         gate_16_before_state_hash, gate_17_migration_proof, gate_18_rerun,
         gate_19_stranded, gate_20_restore, gate_21_no_network_no_secrets,
         gate_22_empty_notices)


def run_gate(cur, gate):
    """One gate in its own savepoint. A gate that raises (the loader halts
    inside a seeded scenario, a database error) is a FAIL of that gate, not
    the end of the run."""
    n = int(re.match(r"gate_(\d+)_", gate.__name__).group(1))
    seen = len(RESULTS)
    cur.execute("SAVEPOINT verify_gate")
    try:
        gate(cur)
        cur.execute("RELEASE SAVEPOINT verify_gate")
    except (Exception, SystemExit) as e:  # noqa: BLE001
        cur.execute("ROLLBACK TO SAVEPOINT verify_gate")
        cur.execute("RELEASE SAVEPOINT verify_gate")
        if len(RESULTS) == seen:
            report(n, gate.__name__, False,
                   f"gate raised {type(e).__name__}: {_first(e)}")


def leftovers():
    """zz% tables, functions and triggers in the real database, read on a
    separate read-only connection after the rollback."""
    conn = get_readonly_conn()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT relname FROM pg_class WHERE relname LIKE "
                        "'zz%%' AND relnamespace = 'public'::regnamespace")
            out = [r[0] for r in cur.fetchall()]
            cur.execute("SELECT 'function ' || proname FROM pg_proc WHERE "
                        "proname LIKE 'zz%%' AND pronamespace = "
                        "'public'::regnamespace")
            out += [r[0] for r in cur.fetchall()]
            cur.execute("SELECT 'trigger ' || tgname FROM pg_trigger WHERE "
                        "tgname LIKE 'zz%%' AND NOT tgisinternal")
            out += [r[0] for r in cur.fetchall()]
            return sorted(out)
    finally:
        conn.rollback()
        conn.close()


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if argv == ["--print-note-lines"]:
        return print_note_lines()
    if argv:
        sys.exit("usage: s24_rsh_register_editions_verify.py "
                 "[--print-note-lines]")
    conn = get_conn()
    try:
        with conn.cursor() as raw:
            cur = SafeCur(raw)
            setup_throwaway(cur)
            for gate in GATES:
                run_gate(cur, gate)
    finally:
        conn.rollback()
        conn.close()
    left = leftovers()
    mine = [t for t in left if "zz_s24" in t]
    if mine:
        print(f"LEFTOVER: committed zz_s24 objects in the database: {mine}")
        RESULTS.append(False)
    else:
        print("no zz_s24 table, function or trigger left in the database "
              "after the rollback")
    others = [t for t in left if t not in mine]
    if others:
        print(f"NOTE: other zz% objects exist in the database (not made by "
              f"this run): {others}")
    sys.exit(0 if all(RESULTS) else 1)


if __name__ == "__main__":
    main()
